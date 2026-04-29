import sys
import unittest
from pathlib import Path
from datetime import date, datetime

from openpyxl import Workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
SALES_ROOT = PROJECT_ROOT / "Sales"
OILSTATION_ROOT = PROJECT_ROOT / "OilStation"
OILSTATION_SCRIPTS = OILSTATION_ROOT / "scripts"
for path in (str(OILSTATION_ROOT), str(OILSTATION_SCRIPTS), str(SALES_ROOT)):
    if path not in sys.path:
        sys.path.insert(0, path)

import generate_accounting_handover_reports as handover


def workbook_with_main_rows(rows: list[dict]):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        row_number = row["row"]
        ws.cell(row_number, 1).value = row.get("name")
        ws.cell(row_number, 3).value = "л"
        ws.cell(row_number, 5).value = row.get("tank_begin")
        ws.cell(row_number, 10).value = row.get("meter_end")
        ws.cell(row_number, 13).value = row.get("meter_begin")
        ws.cell(row_number, 17).value = row.get("tank_number")
        ws.cell(row_number, 26).value = row.get("tank_end")
    ws.cell(30, 1).value = "Расшифровка поступления"
    ws.cell(40, 1).value = "Движение наличных денег"
    return wb, ws


class AccountingHandoverReportTests(unittest.TestCase):
    def test_counter_matching_falls_back_to_resource_row_order(self):
        wb, ws = workbook_with_main_rows([
            {"row": 7, "name": "АИ-92", "meter_begin": 90, "meter_end": 100},
            {"row": 8, "meter_begin": 190, "meter_end": 200},
        ])
        try:
            assignments = handover.assign_counter_rows(
                ws,
                [
                    {"BalanceShopShiftKey": 1, "resource": 103, "BeginBalance": 0, "EndBalance": 101},
                    {"BalanceShopShiftKey": 2, "resource": 103, "BeginBalance": 0, "EndBalance": 202},
                ],
            )
        finally:
            wb.close()

        self.assertEqual(assignments[7]["EndBalance"], 101)
        self.assertEqual(assignments[8]["EndBalance"], 202)

    def test_zero_template_counter_row_does_not_take_leftover_idle_counter(self):
        wb, ws = workbook_with_main_rows([
            {"row": 7, "name": "Аи-80", "meter_begin": 0, "meter_end": 0},
            {"row": 8, "meter_begin": 710, "meter_end": 710},
        ])
        try:
            assignments = handover.assign_counter_rows(
                ws,
                [
                    {"BalanceShopShiftKey": 1, "resource": 102, "BeginBalance": 710, "EndBalance": 710},
                    {"BalanceShopShiftKey": 2, "resource": 102, "BeginBalance": 690, "EndBalance": 690},
                ],
            )
        finally:
            wb.close()

        self.assertIsNone(assignments[7])
        self.assertEqual(assignments[8]["EndBalance"], 710)

    def test_tank_matching_falls_back_to_resource_row_order(self):
        wb, ws = workbook_with_main_rows([
            {
                "row": 7,
                "name": "Аи 80",
                "tank_number": 1,
                "tank_begin": 19698.73,
                "tank_end": "19430,774709",
            },
        ])
        try:
            assignments = handover.assign_tank_rows(
                ws,
                [
                    {
                        "BalanceShopShiftKey": 1,
                        "resource": 102,
                        "WarehouseCode": 8,
                        "BeginBalance": 19425.774709,
                        "EndBalance": 19204.260004,
                    },
                ],
            )
        finally:
            wb.close()

        self.assertEqual(assignments[7]["WarehouseCode"], 8)
        self.assertEqual(assignments[7]["EndBalance"], 19204.260004)

    def test_supply_volume_can_be_written_to_non_product_tank_row(self):
        wb, ws = workbook_with_main_rows([
            {"row": 7, "name": "ДТ", "tank_number": 3, "meter_begin": 0, "meter_end": 0},
            {"row": 8, "tank_number": 4, "meter_begin": 0, "meter_end": 0},
        ])
        try:
            warnings = handover.fill_main_table(
                ws,
                handover.ShiftRange(4, (994,), None, None),
                [
                    {"BalanceShopShiftKey": 1, "resource": 106, "BeginBalance": 0, "EndBalance": 0},
                    {"BalanceShopShiftKey": 2, "resource": 106, "BeginBalance": 0, "EndBalance": 0},
                ],
                [
                    {
                        "BalanceShopShiftKey": 3,
                        "resource": 106,
                        "WarehouseCode": 3,
                        "BeginBalance": 0,
                        "EndBalance": 0,
                        "EndBalanceCalculation": 0,
                    },
                    {
                        "BalanceShopShiftKey": 4,
                        "resource": 106,
                        "WarehouseCode": 4,
                        "BeginBalance": 100,
                        "EndBalance": 500,
                        "EndBalanceCalculation": 500,
                    },
                ],
                {
                    106: {
                        "volume": 5030,
                        "entries": [{"warehouse_code": 4, "volume": 5030}],
                    },
                },
            )
            row_7_supply = ws.cell(7, 6).value
            row_8_supply = ws.cell(8, 6).value
        finally:
            wb.close()

        self.assertEqual(warnings, [])
        self.assertEqual(row_7_supply, 0.0)
        self.assertEqual(row_8_supply, 5030.0)

    def test_zero_amount_words_preserve_template_zero_sign(self):
        positive_zero = handover.amount_words_ru(0)
        negative_zero = handover.amount_words_ru(-0.000000001)

        self.assertEqual(
            handover.zero_amount_words_like_template(positive_zero, negative_zero),
            positive_zero,
        )
        self.assertEqual(
            handover.zero_amount_words_like_template(negative_zero, positive_zero),
            negative_zero,
        )

    def test_fill_detail_table_clears_aj_blank_column(self):
        wb = Workbook()
        ws = wb.active
        ws.cell(10, 1).value = "Расшифровка поступления"
        ws.cell(14, 1).value = "AI-92"
        ws.cell(14, 36).value = 0
        ws.cell(20, 1).value = "Движение наличных денег"
        try:
            handover.fill_detail_table(ws, {}, {})
            blank_value = ws.cell(14, 36).value
        finally:
            wb.close()

        self.assertEqual(blank_value, "")

    def test_fill_main_table_clears_aj_blank_column(self):
        wb, ws = workbook_with_main_rows([
            {"row": 7, "name": "AI-92", "meter_begin": 0, "meter_end": 0},
            {"row": 8, "meter_begin": 0, "meter_end": 0},
            {"row": 9, "name": "DT", "meter_begin": 0, "meter_end": 0},
        ])
        ws.cell(7, 36).value = 0
        ws.cell(8, 36).value = 0
        ws.cell(9, 36).value = 0
        try:
            handover.fill_main_table(
                ws,
                handover.ShiftRange(26, (733,), None, None),
                [],
                [],
                {},
            )
            product_value = ws.cell(7, 36).value
            continuation_value = ws.cell(8, 36).value
            second_product_value = ws.cell(9, 36).value
        finally:
            wb.close()

        self.assertEqual(product_value, 0)
        self.assertEqual(continuation_value, "")
        self.assertEqual(second_product_value, "")

    def test_shift_ranges_use_shift_begin_date_for_accounting_reports(self):
        class FakeCursor:
            def __init__(self):
                self.sql = ""
                self.params = None

            def execute(self, sql, params):
                self.sql = sql
                self.params = params

            def fetchall(self):
                return [{
                    "ShopKey": 5,
                    "ShiftKey": 922,
                    "DatetimeShiftBegin": datetime(2026, 7, 27, 18, 43, 13),
                    "DatetimeShiftEnd": datetime(2026, 7, 28, 8, 21, 37),
                    "ShiftMinutes": 818,
                    "SellingRows": 0,
                }]

        cursor = FakeCursor()
        shifts = handover.load_shift_ranges(cursor, date(2026, 7, 27))

        self.assertIn("DatetimeShiftBegin >= %s", cursor.sql)
        self.assertNotIn("CASE", cursor.sql)
        self.assertEqual(cursor.params, ("2026-07-27", "2026-07-28"))
        self.assertEqual(shifts[5][0].shift_keys, (922,))


if __name__ == "__main__":
    unittest.main()
