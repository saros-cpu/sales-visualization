"""月度日销表工作簿定位的回归测试。

背景：导出服务把月度日销表写入 ``exports/monthly/<年份>/`` 子目录，而
``_find_report_workbook_for_ym`` 历史上只做非递归的根目录 ``glob`` 扫描，
于是每个自然月的「首个月份」（根目录还没有该月历史文件）都找不到工作簿，
快照刷新直接以 skip 失败。此处锁定索引优先 + 递归兜底两条路径。
"""

import json
import os
import tempfile
import time
import unittest
from pathlib import Path

import openpyxl

from retail_internal import import_oilstation

common = import_oilstation("endpoints.v1.retail.analysis_report.service.common")
export_excel = import_oilstation("export_excel")


def _title(year: int, month: int) -> str:
    return f"Ежедневный отчет по АЗС - {export_excel.MONTH_NAMES_RU[month]} {year}"


class RetailReportWorkbookLookupTests(unittest.TestCase):
    def setUp(self):
        self._original_export_dir = common.EXPORT_DIR
        self._tmp = tempfile.TemporaryDirectory()
        self.export_dir = Path(self._tmp.name)
        common.EXPORT_DIR = str(self.export_dir)
        self.addCleanup(self._restore)

    def _restore(self):
        common.EXPORT_DIR = self._original_export_dir
        self._tmp.cleanup()

    def _write_workbook(self, relative_path: str, year: int, month: int, daily_sheets=("01日",)) -> Path:
        path = self.export_dir / relative_path
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "总表"
        sheet.cell(1, 1, _title(year, month))
        for name in daily_sheets:
            workbook.create_sheet(name)
        workbook.save(path)
        workbook.close()
        return path

    def _write_index(self, year_month: str, stored_path: str) -> None:
        payload = {year_month: {"file": Path(stored_path).name, "path": stored_path}}
        (self.export_dir / ".monthly_report_index.json").write_text(
            json.dumps(payload, ensure_ascii=False), encoding="utf-8"
        )

    def test_index_entry_under_monthly_subdir_is_found(self):
        """索引记录指向 monthly/<年份>/ 时必须命中（月初首月的回归点）。"""
        path = self._write_workbook("monthly/2026/加油站九月份日销表01.09.xlsx", 2026, 9)
        self._write_index("2026-09", "monthly/2026/加油站九月份日销表01.09.xlsx")

        self.assertEqual(common._find_report_workbook_for_ym("2026-09"), path)

    def test_fallback_scans_monthly_subdir_when_index_missing(self):
        """索引缺失/未写入时，递归兜底也要能在 monthly/<年份>/ 找到当月工作簿。"""
        path = self._write_workbook("monthly/2026/加油站九月份日销表01.09.xlsx", 2026, 9)

        self.assertEqual(common._find_report_workbook_for_ym("2026-09"), path)

    def test_fallback_ignores_other_months(self):
        self._write_workbook("monthly/2026/加油站八月份日销表31.08.xlsx", 2026, 8)

        self.assertIsNone(common._find_report_workbook_for_ym("2026-09"))

    def test_fallback_prefers_workbook_with_daily_sheets(self):
        """跨月请求产出的空壳（只有总表）不得抢占带日分页的完整工作簿。"""
        shell = self._write_workbook(
            "monthly/2026/加油站九月份日销表31.08.xlsx", 2026, 9, daily_sheets=()
        )
        path = self._write_workbook(
            "monthly/2026/加油站九月份日销表01.09.xlsx", 2026, 9, daily_sheets=("01日",)
        )
        # 让空壳的 mtime 更新，确保命中顺序不是靠时间碰巧正确的
        os.utime(shell, (time.time() + 10, time.time() + 10))

        self.assertEqual(common._find_report_workbook_for_ym("2026-09"), path)

    def test_fallback_falls_back_to_shell_workbook(self):
        """只有空壳可用时仍然返回它，由上层 db supplement 补数。"""
        shell = self._write_workbook(
            "monthly/2026/加油站九月份日销表31.08.xlsx", 2026, 9, daily_sheets=()
        )

        self.assertEqual(common._find_report_workbook_for_ym("2026-09"), shell)

    def test_index_entry_outside_export_dir_is_rejected(self):
        """索引被写坏并指向 exports 之外的路径时不得被采信。"""
        outside = self.export_dir.parent / "outside.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "总表"
        sheet.cell(1, 1, _title(2026, 9))
        workbook.save(outside)
        workbook.close()
        try:
            self._write_index("2026-09", "../outside.xlsx")

            self.assertIsNone(common._find_report_workbook_for_ym("2026-09"))
        finally:
            outside.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
