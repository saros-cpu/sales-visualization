import sys
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


PROJECT_ROOT = Path(__file__).resolve().parents[2]
OILSTATION_ROOT = PROJECT_ROOT / "OilStation"
if str(OILSTATION_ROOT) not in sys.path:
    sys.path.insert(0, str(OILSTATION_ROOT))

from endpoints.v1.retail.export.controller import _write_self_use_station_sheet


def _styles():
    return {
        "header_font": Font(name="Times New Roman", bold=True, size=11, color="FFFFFF"),
        "header_fill": PatternFill(start_color="CC0000", end_color="CC0000", fill_type="solid"),
        "center_align": Alignment(horizontal="center", vertical="center"),
        "thin_border": Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        ),
        "bold_font": Font(name="Times New Roman", bold=True, size=11),
        "subtotal_fill": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
        "total_fill": PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid"),
    }


class SelfUseStationSheetTests(unittest.TestCase):
    def test_station_sheet_adds_ton_totals_by_vehicle_type(self):
        wb = Workbook()
        records = [
            {
                "shopKey": 1,
                "stationName": "Station 1",
                "stationShort": "AZS-1",
                "plateNumber": "CAR-1",
                "productName": "AI-92",
                "volume": 10,
                "mass": 7,
                "amount": 100,
                "personKey": 1,
                "personGroup": 1,
                "selfUseNumber": 0,
            },
            {
                "shopKey": 1,
                "stationName": "Station 1",
                "stationShort": "AZS-1",
                "plateNumber": "CAR-2",
                "productName": "AI-92",
                "volume": 4,
                "mass": 2.8,
                "amount": 40,
                "personKey": 2,
                "personGroup": 1,
                "selfUseNumber": 0,
            },
            {
                "shopKey": 1,
                "stationName": "Station 1",
                "stationShort": "AZS-1",
                "plateNumber": "EMP-1",
                "productName": "AI-92",
                "volume": 5,
                "mass": 3.5,
                "amount": 50,
                "personKey": 9999,
                "personGroup": 1,
                "selfUseNumber": 0,
            },
            {
                "shopKey": 1,
                "stationName": "Station 1",
                "stationShort": "AZS-1",
                "plateNumber": "ГЕНЕРАТОР",
                "productName": "AI-92",
                "volume": 3,
                "mass": 2.1,
                "amount": 30,
                "personKey": 9999,
                "personGroup": 9002,
                "selfUseNumber": 1,
            },
        ]

        ws = _write_self_use_station_sheet(wb, records, _styles())

        headers = [ws.cell(row=1, column=col).value for col in range(1, 10)]
        self.assertEqual(headers, [
            "加油站",
            "加油站俄语名称",
            "车牌号",
            "油品名称",
            "加油量公斤(kg)",
            "加油量升(L)",
            "合计吨位(t)",
            "金额(сом)",
            "备注",
        ])
        self.assertEqual(ws.cell(row=2, column=7).value, 0.0098)
        self.assertEqual(ws.cell(row=3, column=7).value, None)
        self.assertEqual(ws.cell(row=4, column=7).value, 0.0035)
        self.assertEqual(ws.cell(row=5, column=7).value, 0.0021)
        self.assertEqual(ws.cell(row=5, column=9).value, "发电机")
        merged_ranges = {str(cell_range) for cell_range in ws.merged_cells.ranges}
        self.assertIn("D2:D3", merged_ranges)
        self.assertIn("G2:G3", merged_ranges)
        self.assertNotIn("D2:D4", merged_ranges)
        self.assertNotIn("G2:G4", merged_ranges)


if __name__ == "__main__":
    unittest.main()
