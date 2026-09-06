import sys
import unittest
from pathlib import Path


PROJECT_ROOT = Path(__file__).resolve().parents[2]
SALES_ROOT = PROJECT_ROOT / "Sales"
if str(SALES_ROOT) not in sys.path:
    sys.path.insert(0, str(SALES_ROOT))

from analysis_reports import _fmt as sales_report_fmt
from analysis_reports import _round_half_up as sales_report_round
from openpyxl import Workbook

from retail_internal import import_oilstation

_oil_export_excel = import_oilstation("export_excel")
BalanceMap = _oil_export_excel.BalanceMap
_excel_round_formula = _oil_export_excel._excel_round_formula
oil_report_round = _oil_export_excel._round_half_up
_write_station_block = _oil_export_excel._write_station_block


class ReportRoundingTests(unittest.TestCase):
    def test_report_rounding_uses_half_up_not_bankers_rounding(self):
        self.assertEqual(oil_report_round(2.675, 2), 2.68)
        self.assertEqual(oil_report_round(2.665, 2), 2.67)
        self.assertEqual(oil_report_round(-2.675, 2), -2.68)
        self.assertEqual(sales_report_round(2.5, 0), 3.0)
        self.assertEqual(sales_report_fmt(2.675, 2), "2.68")

    def test_excel_total_formula_rounds_formula_result(self):
        self.assertEqual(_excel_round_formula("SUM(A1:A3)", 3), "=ROUND(SUM(A1:A3),3)")

    def test_station_block_subtotal_formula_rounds_written_cells(self):
        wb = Workbook()
        ws = wb.active
        start_balance = BalanceMap({(1, 103): 100.005})
        start_balance.mass = {(1, 103): 74.505}
        start_balance.density = {(1, 103): 0.745}
        end_balance = BalanceMap({(1, 103): 200.005})
        end_balance.mass = {(1, 103): 148.505}
        end_balance.density = {(1, 103): 0.745}
        sales_data = {
            1: {
                103: {
                    "cash_vol": 1.005,
                    "cash_amt": 2.675,
                    "total_vol": 1.005,
                    "total_amt": 2.675,
                    "avg_price": 2.675,
                    "cash_mass": 0.505,
                    "total_mass": 0.505,
                    "density": 0.745,
                }
            }
        }

        _write_station_block(
            ws,
            5,
            1,
            {"name": "Station", "short": "ST"},
            sales_data,
            {},
            {},
            start_balance,
            end_balance,
            fuel_keys=[103],
            include_contract_ton=True,
            ton_decimals=4,
        )

        self.assertEqual(ws.cell(row=6, column=6).value, "=ROUND(SUM(F5:F5),3)")
        self.assertEqual(ws.cell(row=6, column=7).value, "=ROUND(SUM(G5:G5),4)")


if __name__ == "__main__":
    unittest.main()
