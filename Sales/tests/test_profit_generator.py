import sys
import unittest
from pathlib import Path

from openpyxl import Workbook


from retail_internal import import_oilstation

profit = import_oilstation("endpoints.v1.retail.export.service.profit_generator")
REPORT_HIDDEN_STATION_NOTE = import_oilstation("constants").REPORT_HIDDEN_STATION_NOTE


class ProfitGeneratorTests(unittest.TestCase):
    def test_profit_period_normalizers_validate_inputs(self):
        self.assertEqual(profit.normalize_profit_year_month("2026-7"), "2026-07")
        self.assertEqual(profit.normalize_profit_year_month("2026-07"), "2026-07")
        self.assertIsNone(profit.normalize_profit_year_month("2026-13"))
        self.assertIsNone(profit.normalize_profit_year_month("not-a-month"))
        self.assertEqual(profit.normalize_profit_year("2026"), "2026")
        self.assertIsNone(profit.normalize_profit_year("1999"))
        self.assertIsNone(profit.normalize_profit_year("year"))

    def test_month_layout_uses_retail_station_and_fuel_sources(self):
        defaults = profit._profit_defaults_by_shop(profit._load_profit_config())
        stations = profit._build_profit_station_layout(defaults, {}, {}, {}, {}, {})

        self.assertEqual(len(stations), 30)
        self.assertEqual(stations[0]["name"], "S01\nS01")
        self.assertEqual(
            [row["product"] for row in stations[0]["rows"]],
            ["АИ 95", "АИ 92", "АИ 80", "Д/Т", "СПБТ"],
        )

        kemin_circle = next(station for station in stations if station["shopKey"] == 28)
        self.assertEqual(
            [row["product"] for row in kemin_circle["rows"]],
            ["АИ 95", "АИ 92", "АИ 80", "Д/Т", "ДТ -35", "СПБТ"],
        )

    def test_alexandrovka_ai80_ai92_purchase_prices_are_not_swapped(self):
        defaults = profit._profit_defaults_by_shop(profit._load_profit_config())
        stations = profit._build_profit_station_layout(defaults, {}, {}, {}, {}, {})
        alexandrovka = next(station for station in stations if station["shopKey"] == 3)
        by_product = {row["product"]: row for row in alexandrovka["rows"]}

        self.assertEqual(by_product["АИ 80"]["H"], "АЙ80")
        self.assertEqual(by_product["АИ 80"]["I"], 50)
        self.assertEqual(by_product["АИ 92"]["H"], "АЙ92")
        self.assertEqual(by_product["АИ 92"]["I"], 54)

    def test_month_sheet_hides_closed_station_rows_and_adds_note(self):
        wb = Workbook()
        context = {
            "month": 7,
            "stations": [
                {
                    "shopKey": 11,
                    "name": "S-22\nSTN-21",
                    "rows": [{"rk": 102, "product": "АИ 80", "D": 60, "H": "АЙ80", "I": 50}],
                    "M": 100,
                    "O": 10,
                    "P": 5,
                    "Q": 2,
                    "n_rate": 0.3,
                    "n_rks": [102],
                    "diesel_rks": [],
                },
                {
                    "shopKey": 1,
                    "name": "S01\nS01",
                    "rows": [{"rk": 102, "product": "АИ 80", "D": 60, "H": "АЙ80", "I": 50}],
                    "M": 100,
                    "O": 10,
                    "P": 5,
                    "Q": 2,
                    "n_rate": 0.3,
                    "n_rks": [102],
                    "diesel_rks": [],
                },
            ],
            "cMap": {(11, 102): 1, (1, 102): 1},
            "eMap": {(11, 102): 60, (1, 102): 60},
        }

        total_row = profit._write_profit_sheet(wb.active, context)
        ws = wb.active

        self.assertTrue(bool(ws.row_dimensions[2].hidden))
        self.assertFalse(bool(ws.row_dimensions[3].hidden))
        self.assertEqual(ws.cell(row=total_row + 2, column=1).value, REPORT_HIDDEN_STATION_NOTE)

    def test_confirmed_month_purchase_prices_override_defaults(self):
        overrides = profit._load_month_purchase_overrides("2026-01", include_db=False)

        self.assertEqual(overrides[(1, 102)]["H"], "АЙ80")
        self.assertEqual(overrides[(1, 102)]["I"], 52.5)
        self.assertEqual(overrides[(1, 106)]["H"], "Дт")
        self.assertEqual(overrides[(1, 106)]["I"], 64.0)

        defaults = profit._profit_defaults_by_shop(profit._load_profit_config())
        stations = profit._build_profit_station_layout(defaults, {}, {}, {}, {}, {}, overrides)
        kara_balta = next(station for station in stations if station["shopKey"] == 1)
        by_product = {row["product"]: row for row in kara_balta["rows"]}

        self.assertEqual(by_product["АИ 80"]["I"], 52.5)
        self.assertEqual(by_product["Д/Т"]["I"], 64.0)

    def test_db_purchase_prices_override_template_values(self):
        original = profit._load_db_month_purchase_overrides
        try:
            profit._load_db_month_purchase_overrides = lambda year_month: {
                (1, 102): {"H": "АЙ80", "I": 60.0},
            }
            overrides = profit._load_month_purchase_overrides("2026-01")
        finally:
            profit._load_db_month_purchase_overrides = original

        self.assertEqual(overrides[(1, 102)]["H"], "АЙ80")
        self.assertEqual(overrides[(1, 102)]["I"], 60.0)

    def test_alexandrovka_crossed_cells_are_corrected_in_month_overrides(self):
        overrides = profit._load_month_purchase_overrides("2026-07", include_db=False)
        defaults = profit._profit_defaults_by_shop(profit._load_profit_config())
        stations = profit._build_profit_station_layout(defaults, {}, {}, {}, {}, {}, overrides)
        alexandrovka = next(station for station in stations if station["shopKey"] == 3)
        by_product = {row["product"]: row for row in alexandrovka["rows"]}

        self.assertEqual(overrides[(3, 102)]["H"], "АЙ80")
        self.assertEqual(overrides[(3, 102)]["I"], 50.0)
        self.assertEqual(overrides[(3, 103)]["H"], "АЙ92")
        self.assertEqual(overrides[(3, 103)]["I"], 54.0)
        self.assertEqual(by_product["АИ 80"]["I"], 50.0)
        self.assertEqual(by_product["АИ 92"]["I"], 54.0)

    def test_fixed_costs_are_full_month_for_existing_operating_station(self):
        stations = [{
            "shopKey": 1,
            "M": 3100,
            "O": 310,
            "P": 31,
            "Q": 0,
        }]

        profit._apply_operating_cost_ratios(stations, {1: 15}, 31, "2026-01-01", {1: "2025-12-01"})

        self.assertEqual(stations[0]["operatingDays"], 15)
        self.assertEqual(stations[0]["periodDays"], 31)
        self.assertEqual(stations[0]["operatingRatio"], 1.0)
        self.assertEqual(stations[0]["costProrationReason"], "full_month")
        self.assertEqual(stations[0]["MBase"], 3100)
        self.assertEqual(stations[0]["M"], 3100)
        self.assertEqual(stations[0]["O"], 310)
        self.assertEqual(stations[0]["P"], 31)
        self.assertEqual(stations[0]["Q"], 0)

    def test_fixed_costs_are_prorated_for_new_station_first_month(self):
        stations = [{
            "shopKey": 29,
            "M": 3100,
            "O": 310,
            "P": 31,
            "Q": 0,
        }]

        profit._apply_operating_cost_ratios(stations, {29: 15}, 31, "2026-01-01", {29: "2026-01-10"})

        self.assertEqual(stations[0]["operatingDays"], 15)
        self.assertAlmostEqual(stations[0]["operatingRatio"], 0.4839)
        self.assertEqual(stations[0]["costProrationReason"], "new_station")
        self.assertEqual(stations[0]["M"], 1500)
        self.assertEqual(stations[0]["O"], 150)
        self.assertEqual(stations[0]["P"], 15)
        self.assertEqual(stations[0]["Q"], 0)

    def test_month_payload_includes_operating_time_and_product_rows(self):
        station = {
            "shopKey": 1,
            "name": "S01\nS01",
            "rows": [{"rk": 102, "product": "АИ 80", "D": None, "I": 52.5, "H": "АЙ80"}],
            "M": 1500,
            "MBase": 3100,
            "O": 150,
            "OBase": 310,
            "P": 15,
            "PBase": 31,
            "Q": 0,
            "QBase": 0,
            "n_rate": 0.3,
            "n_rks": [102],
            "diesel_rks": [],
            "operatingDays": 15,
            "periodDays": 31,
            "operatingRatio": 0.4839,
            "firstOperatingDay": "2026-01-10",
            "costProrationReason": "new_station",
        }
        payload = profit._profit_month_payload({
            "year": 2026,
            "month": 1,
            "yearMonth": "2026-01",
            "queryStart": "2026-01-01",
            "queryEnd": "2026-02-01",
            "periodDays": 31,
            "summary": {"volume": 100, "salesAmount": 6200},
            "stations": [station],
            "cMap": {(1, 102): 100},
            "eMap": {(1, 102): 6200},
            "massMap": {(1, 102): 72.1},
        })

        self.assertEqual(payload["stations"][0]["operatingDays"], 15)
        self.assertEqual(payload["stations"][0]["firstOperatingDay"], "2026-01-10")
        self.assertEqual(payload["stations"][0]["costProrationReason"], "new_station")
        self.assertEqual(payload["stations"][0]["salary"], 1500)
        self.assertEqual(payload["stations"][0]["salaryBase"], 3100)
        self.assertEqual(payload["stations"][0]["rows"][0]["purchasePrice"], 52.5)
        self.assertEqual(payload["stations"][0]["rows"][0]["purchaseAmount"], 5250)
        self.assertEqual(payload["stations"][0]["rows"][0]["salesPrice"], 62)

    def test_station_summary_includes_diesel_volume(self):
        station = {
            "shopKey": 1,
            "name": "站点",
            "rows": [
                {"rk": 102, "I": 50},
                {"rk": 106, "I": 68},
            ],
            "M": 100,
            "O": 10,
            "P": 5,
            "Q": 2,
            "n_rate": 0.3,
            "n_rks": [102, 106],
            "diesel_rks": [106],
            "operatingDays": 31,
            "periodDays": 31,
            "operatingRatio": 1,
        }
        summary = profit._calculate_profit_station_summary(
            station,
            {(1, 102): 60, (1, 106): 40},
            {(1, 102): 3600, (1, 106): 2800},
            {(1, 102): 45, (1, 106): 33},
        )

        self.assertEqual(summary["volume"], 100)
        self.assertEqual(summary["dieselVolume"], 40)
        self.assertEqual(summary["dieselTons"], 0.033)
        self.assertEqual(summary["dieselSalesAmount"], 2800)

    def test_month_data_uses_snapshot_unless_forced(self):
        original_load = profit._load_profit_snapshot
        original_save = profit._save_profit_snapshot
        original_fresh = profit._fresh_profit_month_data
        calls = []
        try:
            profit._load_profit_snapshot = lambda snapshot_type, period_key: (
                {"yearMonth": period_key, "summary": {"netProfit": 12, "dieselVolume": 3, "dieselTons": 0.002}, "stations": []},
                "2026-08-08 16:00:00",
            )
            profit._save_profit_snapshot = lambda snapshot_type, period_key, payload: None
            profit._fresh_profit_month_data = lambda conn, year_month: calls.append(year_month) or {
                "yearMonth": year_month,
                "summary": {"netProfit": 99},
                "stations": [],
            }

            cached = profit.get_profit_month_data(None, "2026-07")
            refreshed = profit.get_profit_month_data(None, "2026-07", force_refresh=True)
        finally:
            profit._load_profit_snapshot = original_load
            profit._save_profit_snapshot = original_save
            profit._fresh_profit_month_data = original_fresh

        self.assertEqual(cached["summary"]["netProfit"], 12)
        self.assertTrue(cached["snapshotCached"])
        self.assertEqual(refreshed["summary"]["netProfit"], 99)
        self.assertFalse(refreshed["snapshotCached"])
        self.assertEqual(calls, ["2026-07"])

    def test_month_data_attaches_comparison_from_snapshots(self):
        original_load = profit._load_profit_snapshot
        original_save = profit._save_profit_snapshot
        original_fresh = profit._fresh_profit_month_data
        fresh_calls = []
        snapshots = {
            ("month", "2026-07"): {
                "yearMonth": "2026-07",
                "summary": {"salesAmount": 200, "volume": 20, "dieselVolume": 8, "dieselTons": 0.006},
                "stations": [],
            },
            ("month", "2026-06"): {
                "yearMonth": "2026-06",
                "summary": {"salesAmount": 150, "volume": 15, "dieselVolume": 6, "dieselTons": 0.004},
                "stations": [],
            },
            ("month", "2025-07"): {
                "yearMonth": "2025-07",
                "summary": {"salesAmount": 100, "volume": 10, "dieselVolume": 5, "dieselTons": 0.003},
                "stations": [],
            },
        }
        try:
            profit._load_profit_snapshot = lambda snapshot_type, period_key: (
                (snapshots[(snapshot_type, str(period_key))], "2026-08-08 16:00:00")
                if (snapshot_type, str(period_key)) in snapshots else None
            )
            profit._save_profit_snapshot = lambda snapshot_type, period_key, payload: None
            profit._fresh_profit_month_data = lambda conn, year_month: fresh_calls.append(year_month) or {}

            data = profit.get_profit_month_data(None, "2026-07")
        finally:
            profit._load_profit_snapshot = original_load
            profit._save_profit_snapshot = original_save
            profit._fresh_profit_month_data = original_fresh

        self.assertEqual(data["comparison"]["mom"]["prevMonth"], "2026-06")
        self.assertEqual(data["comparison"]["mom"]["summary"]["salesAmount"], 150)
        self.assertEqual(data["comparison"]["yoy"]["yoyMonth"], "2025-07")
        self.assertEqual(data["comparison"]["yoy"]["summary"]["volume"], 10)
        self.assertEqual(fresh_calls, [])

    def test_year_data_attaches_ytd_comparison_from_snapshot(self):
        original_load = profit._load_profit_snapshot
        original_save = profit._save_profit_snapshot
        original_fresh = profit._fresh_profit_year_data
        snapshots = {
            ("year", "2026"): {
                "year": 2026,
                "includedMonths": 2,
                "summary": {"salesAmount": 300, "volume": 30, "dieselVolume": 10, "dieselTons": 0.01},
                "months": [
                    {"month": 1, "active": True, "salesAmount": 100, "volume": 10, "dieselVolume": 4, "dieselTons": 0.004},
                    {"month": 2, "active": True, "salesAmount": 200, "volume": 20, "dieselVolume": 6, "dieselTons": 0.006},
                ],
            },
            ("year", "2025"): {
                "year": 2025,
                "includedMonths": 12,
                "summary": {"salesAmount": 1200, "volume": 120, "dieselVolume": 50, "dieselTons": 0.05},
                "months": [
                    {"month": 1, "active": True, "salesAmount": 80, "volume": 8, "dieselVolume": 3, "dieselTons": 0.003},
                    {"month": 2, "active": True, "salesAmount": 90, "volume": 9, "dieselVolume": 4, "dieselTons": 0.004},
                ],
            },
        }
        try:
            profit._load_profit_snapshot = lambda snapshot_type, period_key: (
                (snapshots[(snapshot_type, str(period_key))], "2026-08-08 16:00:00")
                if (snapshot_type, str(period_key)) in snapshots else None
            )
            profit._save_profit_snapshot = lambda snapshot_type, period_key, payload: None
            profit._fresh_profit_year_data = lambda conn, year, force_refresh=False: {}

            data = profit.get_profit_year_data(None, "2026")
        finally:
            profit._load_profit_snapshot = original_load
            profit._save_profit_snapshot = original_save
            profit._fresh_profit_year_data = original_fresh

        self.assertEqual(data["comparison"]["mom"]["summary"]["salesAmount"], 100)
        self.assertEqual(data["comparison"]["yoy"]["summary"]["salesAmount"], 170)
        self.assertEqual(data["comparison"]["yoy"]["summary"]["volume"], 17)

    def test_year_summary_is_single_sheet_with_formula_notes(self):
        wb = Workbook()
        profit._write_profit_year_sheet(
            wb.active,
            2026,
            {
                7: {
                    "hasActivity": True,
                    "volume": 100,
                    "salesAmount": 1300,
                    "purchaseAmount": 700,
                    "salary": 100,
                    "freight": 20,
                    "otherFees": 10,
                    "dieselSalesAmount": 300,
                    "tons": 1.2,
                    "avgCnyKgs": 12.5,
                }
            },
            include_through_month=7,
        )
        ws = wb.active

        self.assertEqual(ws.title, "年度汇总")
        self.assertEqual(ws["A1"].value, "加油站年销量利润核算")
        self.assertEqual(ws["C9"].value, 100.0)
        self.assertEqual(ws["O2"].value, "销量(吨)")
        self.assertEqual(ws["P2"].value, "当月汇率")
        self.assertEqual(ws["R2"].value, "柴油销售金额")
        self.assertEqual(ws["I9"].value, "=($D9-$R9)/1.12+$R9/1.13")
        self.assertEqual(ws["L9"].value, "=$R9*0.01")
        self.assertEqual(ws["O9"].value, 1.2)
        self.assertEqual(ws["P9"].value, 12.5)
        self.assertEqual(ws["Q9"].value, '=IFERROR(N9/P9,"")')
        self.assertEqual(ws["R9"].value, 300.0)
        self.assertEqual(ws["I10"].value, '=IF($D10="","",($D10-$R10)/1.12+$R10/1.13)')
        self.assertEqual(ws["N10"].value, '=IF($D10="","",M10-F10-H10)')
        self.assertFalse(bool(ws.column_dimensions["R"].hidden))
        self.assertEqual(ws["B19"].value, "销售金额=数据库当月销售金额")


if __name__ == "__main__":
    unittest.main()
