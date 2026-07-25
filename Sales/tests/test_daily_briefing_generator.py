import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Border, Side

PROJECT_ROOT = Path(__file__).resolve().parents[2]
OILSTATION_ROOT = PROJECT_ROOT / "OilStation"
if str(OILSTATION_ROOT) not in sys.path:
    sys.path.insert(0, str(OILSTATION_ROOT))

from endpoints.v1.retail.export.service import daily_briefing as briefing


class DailyBriefingGeneratorTests(unittest.TestCase):
    def test_normalizes_report_date_formats(self):
        self.assertEqual(briefing.normalize_daily_briefing_date("2026-08-10"), date(2026, 8, 10))
        self.assertEqual(briefing.normalize_daily_briefing_date("2026.08.10"), date(2026, 8, 10))
        self.assertEqual(briefing.normalize_daily_briefing_date("20260810"), date(2026, 8, 10))
        self.assertIsNone(briefing.normalize_daily_briefing_date("2026/08/10"))

    def test_diesel_variants_are_grouped_into_diesel_row(self):
        totals = briefing._template_totals_from_rows([
            {"resource_key": 106, "total_volume": 100, "total_amount": 1000, "total_mass": 800},
            {"resource_key": 109, "total_volume": 200, "total_amount": 2000, "total_mass": 1600},
            {"resource_key": 110, "total_volume": 300, "total_amount": 3000, "total_mass": 2400},
            {"resource_key": 104, "total_volume": 50, "total_amount": 500, "total_mass": 37},
        ])

        self.assertEqual(totals[8]["volume"], 600)
        self.assertEqual(totals[8]["amount"], 6000)
        self.assertEqual(totals[8]["mass"], 4800)
        self.assertEqual(totals[5]["mass"], 37)
        self.assertEqual(totals[9]["mass"], 0)

    def test_monthly_snapshot_must_match_expected_period_end(self):
        original_get_cached = briefing.get_cached

        def fake_get_cached(_key):
            return {
                "summary": {"queryEnd": "2026-09-01"},
                "byResource": [{"resource_key": 106, "total_mass": 1000}],
            }

        try:
            briefing.get_cached = fake_get_cached
            self.assertIsNone(briefing._monthly_snapshot_tons("2026-08", expected_end=date(2026, 8, 10)))
            self.assertEqual(briefing._monthly_snapshot_tons("2026-08", expected_end=date(2026, 9, 1))[8], 1.0)
        finally:
            briefing.get_cached = original_get_cached

    def test_partial_month_prefers_matching_monthly_snapshot(self):
        original_get_cached = briefing.get_cached
        original_daily_snapshots = briefing._period_tons_from_daily_snapshots

        def fake_get_cached(key):
            if key == "monthly:2026-08":
                return {
                    "summary": {"queryEnd": "2026-08-10"},
                    "byResource": [
                        {"resourceKey": 102, "mass": 500000},
                        {"resourceKey": 106, "mass": 20000},
                        {"resourceKey": 110, "mass": 5000},
                    ],
                }
            return None

        try:
            briefing.get_cached = fake_get_cached
            briefing._period_tons_from_daily_snapshots = lambda *_args: ({8: 25.0}, [])
            tons = briefing._period_tons_by_cache_or_query(None, date(2026, 8, 1), date(2026, 8, 10))

            self.assertEqual(tons[7], 500.0)
            self.assertEqual(tons[8], 25.0)
        finally:
            briefing.get_cached = original_get_cached
            briefing._period_tons_from_daily_snapshots = original_daily_snapshots

    def test_yearly_snapshot_is_used_without_two_decimal_accumulation(self):
        original_get_cached = briefing.get_cached

        def fake_get_cached(key):
            if key == "resources_yearly:2026:year":
                return {
                    "summary": {"totalTons": 19000.0},
                    "resources": [
                        {"resourceKey": 102, "mass": 12000000},
                        {"resourceKey": 106, "mass": 2000000},
                        {"resourceKey": 110, "mass": 25000},
                    ],
                }
            if key == "kpi_yearly:2026:summary":
                return {"summary": {"queryEndExclusive": "2026-08-10"}}
            return None

        try:
            briefing.get_cached = fake_get_cached
            tons = briefing._year_tons_by_cache_or_query(None, date(2026, 8, 10))

            self.assertEqual(tons[7], 12000.0)
            self.assertEqual(tons[8], 2025.0)
        finally:
            briefing.get_cached = original_get_cached

    def test_merge_tons_keeps_retail_three_decimal_precision(self):
        totals = briefing._merge_tons({8: 20.0}, {8: 5.0})
        self.assertEqual(totals[8], 25.0)

    def test_direct_query_excludes_self_use_and_test_sales(self):
        class FakeCursor:
            def __init__(self):
                self.query = ""

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            def execute(self, query, _params):
                self.query = query

            def fetchall(self):
                return []

        class FakeConn:
            def __init__(self):
                self.cursor_obj = FakeCursor()

            def cursor(self):
                return self.cursor_obj

        conn = FakeConn()
        briefing._query_retail_resource_totals(conn, "2026-08-01", "2026-08-02")

        self.assertIn("DerivedPaymentTypeKey NOT IN (2, 4)", conn.cursor_obj.query)

    def test_sheet_layout_and_retail_cells_match_template(self):
        wb = Workbook()
        ws = wb.active
        day_totals = briefing._empty_template_totals()
        day_totals[5]["mass"] = 370
        day_totals[5]["amount"] = 32375
        day_totals[8]["mass"] = 4800
        day_totals[8]["amount"] = 420000
        day_totals[11]["mass"] = 550
        day_totals[11]["amount"] = 48125
        briefing._write_daily_briefing_sheet(
            ws,
            date(2026, 8, 10),
            day_totals,
            {5: 3.7, 8: 48, 11: 5.5},
            {5: 37, 8: 480, 11: 55},
            {5: 0.3, 8: 0.8, 11: 0.1},
            {5: 3.0, 8: 8.0, 11: 1.0},
            {5: 30.0, 8: 80.0, 11: 10.0},
            {8: (1160, 1182), 10: (337, None), 11: (720, None)},
        )

        self.assertEqual(ws.title, "报表")
        self.assertEqual(ws["A1"].value, "ProductTeam销售日报表")
        self.assertEqual(ws["H2"].value, "汇率：87.50")
        self.assertEqual(ws["J2"].value, date(2026, 8, 10))
        self.assertEqual(ws["J2"].number_format, "mm-dd-yy")
        self.assertEqual(ws["F8"].number_format, "0.000")
        self.assertEqual(ws["G8"].number_format, "0.00_ ")
        self.assertEqual(ws["I8"].number_format, "0.000")
        self.assertEqual(ws["K8"].number_format, "0.000")
        self.assertEqual(ws["C13"].number_format, "0.000")
        self.assertEqual(ws.column_dimensions["A"].width, 10.375)
        self.assertEqual(ws.column_dimensions["B"].width, 12.5)
        for column in "CDEFGHIJK":
            self.assertAlmostEqual(ws.column_dimensions[column].width, 16.7109375)
        self.assertEqual(ws["D9"].font.name, "黑体")
        self.assertEqual(ws["D9"].font.sz, 12)
        self.assertEqual(ws["E9"].font.name, "黑体")
        self.assertEqual(ws["E9"].font.sz, 12)
        self.assertEqual(ws["A1"].border.left.style, "thin")
        self.assertEqual(ws["A1"].border.top.style, "thin")
        self.assertEqual(ws["K1"].border.right.style, "thin")
        self.assertEqual(ws["K1"].border.top.style, "thin")
        self.assertEqual(ws["A2"].border.left.style, "thin")
        self.assertEqual(ws["A2"].border.bottom.style, "thin")
        self.assertEqual(ws["K2"].border.right.style, "thin")
        self.assertEqual(ws["K2"].border.bottom.style, "thin")
        self.assertIsNone(ws["A1"].border.bottom.style)
        self.assertIsNone(ws["K1"].border.bottom.style)
        self.assertIsNone(ws["A2"].border.top.style)
        self.assertIsNone(ws["K2"].border.top.style)
        self.assertIsNone(ws["B1"].border.left.style)
        self.assertIsNone(ws["J1"].border.right.style)
        self.assertIsNone(ws["B2"].border.left.style)
        self.assertIsNone(ws["G2"].border.right.style)
        self.assertIsNone(ws["H2"].border.left.style)
        self.assertIsNone(ws["H2"].border.right.style)
        self.assertIsNone(ws["I2"].border.left.style)
        self.assertIsNone(ws["I2"].border.right.style)
        self.assertIsNone(ws["J2"].border.left.style)
        self.assertTrue(briefing._top_rows_have_outline_borders(ws))
        self.assertEqual(ws.row_dimensions[14].height, 87)

        expected_merges = {
            "A1:K1", "A2:G2", "H2:I2", "J2:K2", "A3:A4", "B3:B4",
            "C3:G3", "H3:I3", "J3:K3",
            "A13:B13", "C13:G13", "H13:I13", "J13:K13", "A14:K14",
        }
        self.assertEqual({str(item) for item in ws.merged_cells.ranges}, expected_merges)

        self.assertEqual(ws["B8"].value, "柴油")
        self.assertEqual(ws["C8"].value, 0.8)
        self.assertEqual(ws["H8"].value, 8.0)
        self.assertEqual(ws["J8"].value, 80.0)
        self.assertEqual(ws["D8"].value, 1160)
        self.assertEqual(ws["E8"].value, 1182)
        self.assertEqual(ws["D10"].value, 337)
        self.assertEqual(ws["D11"].value, 720)
        self.assertEqual(ws["F8"].value, 4.8)
        self.assertEqual(ws["G8"].value, 1000.0)
        self.assertEqual(ws["I8"].value, 48)
        self.assertEqual(ws["K8"].value, 480)
        self.assertIsNone(ws["F9"].value)
        self.assertEqual(ws["C12"].value, "=SUM(C5:C11)")
        self.assertEqual(ws["F12"].value, "=SUM(F5:F11)")
        self.assertEqual(ws["I12"].value, "=SUM(I5:I11)")
        self.assertEqual(ws["K12"].value, "=SUM(K5:K11)")
        self.assertEqual(ws["C13"].value, "=C12+F12")
        self.assertEqual(ws["H13"].value, "=H12+I12")
        self.assertEqual(ws["J13"].value, "=J12+K12")
        self.assertIn("销量单位：吨；价格单位：美元/吨", ws["A14"].value)
        self.assertTrue(str(ws["A14"].value).startswith("备注："))

    def test_daily_retail_price_uses_fixed_usd_exchange_rate(self):
        self.assertEqual(
            briefing._retail_price_usd_per_ton({"mass": 4800, "amount": 420000}),
            1000.0,
        )
        self.assertIsNone(briefing._retail_price_usd_per_ton({"mass": 0, "amount": 420000}))

    def test_existing_old_layout_file_is_not_reused_as_current_template(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "old.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "报表"
            ws.merge_cells("A1:H1")
            ws["A1"] = "ProductTeam销售日报表"
            ws["D4"] = "加油站零售量（吨）"
            wb.save(path)

            self.assertFalse(briefing._daily_briefing_file_matches_current_template(path))

    def test_existing_file_with_top_row_inside_borders_is_not_reused(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "inside-border.xlsx"
            wb = Workbook()
            ws = wb.active
            briefing._write_daily_briefing_sheet(
                ws,
                date(2026, 8, 10),
                briefing._empty_template_totals(),
                briefing._empty_tons(),
                briefing._empty_tons(),
                briefing._empty_tons(),
                briefing._empty_tons(),
                briefing._empty_tons(),
                {},
            )
            thin = Side(style="thin", color="000000")
            ws["A1"].border = Border(left=thin, top=thin, bottom=thin)
            wb.save(path)

            self.assertFalse(briefing._daily_briefing_file_matches_current_template(path))


if __name__ == "__main__":
    unittest.main()
