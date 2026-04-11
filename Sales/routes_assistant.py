from __future__ import annotations

from calendar import monthrange
from datetime import date, datetime, timedelta
from email.utils import parsedate_to_datetime
from functools import wraps
import html
import json
import logging
import os
import re
import sys
from pathlib import Path
from urllib.parse import urlencode
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen
import xml.etree.ElementTree as ET
import pandas as pd
import pymysql
from flask import jsonify, request
from assistant_agent import is_action_request, read_only_refusal
from assistant_quality import enrich_assistant_payload
from config import DB_CONFIG
from db import get_connection
from shared_utils import (
    ensure_market_watch_schema, OILSTATION_ROOT,
    active_station_statuses, retail_station_catalog_cached,
)
from text_utils import clean_str, normalize_station_name, parse_date, parse_decimal
from retail_internal import open_oilstation
from services.weighing_service import get_weighing_list
from summary_service import available_periods, period_bounds, summarize_combined, summarize_wholesale
from wholesale_service import summarize_wholesale_page
from routes_data import chinese_market_item, market_source_display


_ASSISTANT_QUERY_EXECUTOR = None


def execute_assistant_query_payload(payload):
    if not callable(_ASSISTANT_QUERY_EXECUTOR):
        return {"ok": False, "error": "assistant_service_not_ready", "answer": "AI 助手业务服务尚未初始化。"}, 503
    return _ASSISTANT_QUERY_EXECUTOR(payload if isinstance(payload, dict) else {})


_ASSISTANT_DIGIT_TRANSLATION = str.maketrans("０１２３４５６７８９", "0123456789")
_ASSISTANT_MONTH_NAME_PATTERN = r"\d{1,2}|一|二|两|三|四|五|六|七|八|九|十|十一|十二"
_ASSISTANT_QUARTER_NAME_PATTERN = r"[1-4]|一|二|三|四"
_ASSISTANT_COUNT_NAME_PATTERN = r"\d{1,3}|[一二两俩三四五六七八九]?十[一二两俩三四五六七八九]?|[一二两俩三四五六七八九]"
_ASSISTANT_SPRING_FESTIVAL_DATES = {
    2024: (2, 10),
    2025: (1, 29),
    2026: (2, 17),
    2027: (2, 6),
    2028: (1, 26),
    2029: (2, 13),
    2030: (2, 3),
}
_ASSISTANT_DATE_PATTERN = (
    r"(20\d{2}|\d{2})\s*(?:年|[-/.])\s*"
    r"(\d{1,2})\s*(?:月|[-/.])\s*"
    r"(\d{1,2})\s*(?:日|号)?"
)
_ASSISTANT_RANGE_PATTERN = r"(?:至|到|--?|—|~|～)"
_ASSISTANT_TOOL_CATALOG = {
    "retail_summary": {
        "name": "retail_summary",
        "description": "零售区间汇总",
        "readOnly": True,
        "dataSource": "retail_snapshots",
    },
    "retail_ranking": {
        "name": "retail_ranking",
        "description": "零售站点、油品、区域排行",
        "readOnly": True,
        "dataSource": "retail_snapshots",
    },
    "retail_trend": {
        "name": "retail_trend",
        "description": "零售日趋势",
        "readOnly": True,
        "dataSource": "retail_daily_snapshots",
    },
    "station_trend_anomaly": {
        "name": "station_trend_anomaly",
        "description": "站点销量下滑和异常波动识别",
        "readOnly": True,
        "dataSource": "retail_daily_snapshots",
    },
    "station_data_check": {
        "name": "station_data_check",
        "description": "站点数据上传和补传检查",
        "readOnly": True,
        "dataSource": "retail_snapshots",
    },
    "station_status": {
        "name": "station_status",
        "description": "站点营业状态查询",
        "readOnly": True,
        "dataSource": "station_status_catalog",
    },
    "accounting_report_status": {
        "name": "accounting_report_status",
        "description": "吉方财务交接日报生成和缺站状态",
        "readOnly": True,
        "dataSource": "accounting_handover_reports",
    },
    "wholesale_summary": {
        "name": "wholesale_summary",
        "description": "批发业务汇总和客户排行",
        "readOnly": True,
        "dataSource": "wholesale_snapshots",
    },
    "wholesale_review": {
        "name": "wholesale_review",
        "description": "批发复核异常提示",
        "readOnly": True,
        "dataSource": "wholesale_snapshots",
    },
    "combined_summary": {
        "name": "combined_summary",
        "description": "零售与批发合并经营汇总",
        "readOnly": True,
        "dataSource": "combined_snapshots",
    },
    "exchange_summary": {
        "name": "exchange_summary",
        "description": "汇率查询",
        "readOnly": True,
        "dataSource": "exchange_rates",
    },
    "market_summary": {
        "name": "market_summary",
        "description": "市场观察摘要",
        "readOnly": True,
        "dataSource": "market_watch",
    },
    "contract_summary": {
        "name": "contract_summary",
        "description": "合同只读查询",
        "readOnly": True,
        "dataSource": "contracts",
    },
    "density_summary": {
        "name": "density_summary",
        "description": "油品密度摘要",
        "readOnly": True,
        "dataSource": "density_records",
    },
    "retail_inventory": {
        "name": "retail_inventory",
        "description": "零售库存查询",
        "readOnly": True,
        "dataSource": "retail_snapshots",
    },
    "retail_price": {
        "name": "retail_price",
        "description": "零售均价查询",
        "readOnly": True,
        "dataSource": "retail_snapshots",
    },
    "retail_resource_share": {
        "name": "retail_resource_share",
        "description": "零售油品结构和占比",
        "readOnly": True,
        "dataSource": "retail_snapshots",
    },
    "retail_comparison": {
        "name": "retail_comparison",
        "description": "零售同比环比对比",
        "readOnly": True,
        "dataSource": "retail_snapshots",
    },
    "retail_kpi": {
        "name": "retail_kpi",
        "description": "零售年度 KPI 查询",
        "readOnly": True,
        "dataSource": "retail_kpi_snapshots",
    },
    "self_use_summary": {
        "name": "self_use_summary",
        "description": "自用油 report24-only 汇总",
        "readOnly": True,
        "dataSource": "self_use_report24_snapshots",
    },
    "procurement_summary": {
        "name": "procurement_summary",
        "description": "原料采购明细和供应商汇总",
        "readOnly": True,
        "dataSource": "wholesale_detail_records",
    },
    "weighing_summary": {
        "name": "weighing_summary",
        "description": "称重明细、分类和审核状态汇总",
        "readOnly": True,
        "dataSource": "weighing_service",
    },
    "materials_summary": {
        "name": "materials_summary",
        "description": "物资字典、入库、调拨单和库存汇总",
        "readOnly": True,
        "dataSource": "materials",
    },
    "material_transfer_summary": {
        "name": "material_transfer_summary",
        "description": "物资调拨单只读汇总",
        "readOnly": True,
        "dataSource": "material_transfer_orders",
    },
    "transfer_loss_summary": {
        "name": "transfer_loss_summary",
        "description": "调拨损耗只读摘要",
        "readOnly": True,
        "dataSource": "transfer_loss",
    },
    "system_health_summary": {
        "name": "system_health_summary",
        "description": "系统健康、自动任务和运行治理摘要",
        "readOnly": True,
        "dataSource": "system_health",
    },
    "knowledge_summary": {
        "name": "knowledge_summary",
        "description": "知识库文章和附件索引摘要",
        "readOnly": True,
        "dataSource": "knowledge_base",
    },
    "audit_log_summary": {
        "name": "audit_log_summary",
        "description": "审计日志只读摘要",
        "readOnly": True,
        "dataSource": "audit_logs",
    },
    "alias_summary": {
        "name": "alias_summary",
        "description": "客户和油品别名匹配摘要",
        "readOnly": True,
        "dataSource": "entity_product_aliases",
    },
    "user_permission_summary": {
        "name": "user_permission_summary",
        "description": "用户、角色、部门和权限只读摘要",
        "readOnly": True,
        "dataSource": "organization_permissions",
    },
    "settings_summary": {
        "name": "settings_summary",
        "description": "客户端和系统设置只读说明",
        "readOnly": True,
        "dataSource": "runtime_settings",
    },
    "compound_query": {
        "name": "compound_query",
        "description": "复合自然语言问题拆解与只读工具编排",
        "readOnly": True,
        "dataSource": "assistant_tool_orchestration",
    },
}


def assistant_tool_catalog():
    return {
        key: {
            **value,
            "route": key,
            "mode": "read_only",
            "canWrite": False,
        }
        for key, value in sorted(_ASSISTANT_TOOL_CATALOG.items())
    }


_ASSISTANT_FOLLOWUP_WORDS = (
    "那",
    "它",
    "他",
    "该站",
    "这个站",
    "这家站",
    "刚才",
    "上面",
    "前面",
    "继续",
    "为什么",
    "原因",
    "只看",
    "换成",
    "改成",
    "按月",
    "按天",
    "按日",
    "呢",
)
_ASSISTANT_METRIC_WORDS = (
    "金额",
    "销售额",
    "营业额",
    "收入",
    "索姆",
    "人民币",
    "元",
    "吨",
    "吨数",
    "吨位",
    "重量",
    "销量",
    "销售量",
    "升",
    "升数",
    "油量",
    "卖得",
)
_ASSISTANT_BUSINESS_LEXICON = {
    "domains": {
        "self_use": ("自用油", "自用油数据", "自用油量", "自用", "内部用油", "公司用油", "公司自用油"),
        "retail": ("零售", "加油站", "站点", "油站", "站里", "销售站"),
        "wholesale": ("批发", "外运", "出口", "本地", "客户"),
        "accounting": ("吉方", "交接日报", "会计日报", "财务交接", "日报缺", "缺站"),
        "combined": ("合并", "总销售", "整体", "公司", "全部业务", "经营情况"),
        "exchange": ("汇率", "美元", "人民币", "美金", "索姆/人民币"),
        "market": ("市场", "行情", "观察", "原油", "成品油", "俄罗斯", "乌兹别克", "国际"),
        "contract": ("合同", "协议"),
        "density": ("密度",),
        "procurement": ("采购", "原料采购", "进货", "采购数据", "供应商", "供货商"),
        "weighing": ("称重", "地磅", "过磅", "磅单", "待审核", "待复核", "审核称重"),
        "materials": ("物资", "物料", "办公用品", "物资库存", "物资入库", "调拨单"),
        "transfer_loss": ("调拨损耗", "损耗率", "损耗情况", "运输损耗", "途中损耗"),
        "system_health": ("系统健康", "健康检查", "运行治理", "提醒中心", "告警", "自动任务", "任务状态", "服务状态"),
        "knowledge": ("知识库", "知识文章", "附件索引", "资料库", "文档库", "口径说明"),
        "audit": ("审计", "审计日志", "操作日志", "登录日志"),
        "alias": ("别名", "别名匹配", "未匹配名称", "客户别名", "油品别名"),
        "users": ("用户", "账号", "权限", "角色", "部门", "组织权限"),
        "management": ("管理中心", "驾驶舱配置", "看板配置", "组件配置"),
        "settings": ("设置", "系统设置", "客户端设置", "api地址", "API地址", "更新源"),
    },
    "metrics": {
        "amount": ("金额", "销售额", "营业额", "收入", "钱", "索姆", "人民币", "元", "款"),
        "volume": ("销量", "销售量", "升", "升数", "油量", "卖得", "卖了多少", "出油量"),
        "tons": ("吨", "吨数", "吨位", "重量", "质量"),
        "count": ("笔数", "单数", "交易数", "车次", "次数"),
    },
        "dimensions": {
            "station": ("站点", "加油站", "哪家站", "哪个站", "各站", "站排行", "站点排行", "油站", "按站点", "分站", "站点拆分"),
            "resource": ("油品", "产品", "品类", "资源", "汽油", "柴油", "д/т", "аи", "按油品", "油品拆分"),
            "region": ("区域", "地区", "州", "各州", "按区域", "区域拆分"),
            "customer": ("客户", "客户排行", "哪家客户", "按客户", "客户拆分"),
            "supplier": ("供应商", "供货商", "发货单位"),
            "vehicle": ("车辆", "车号", "车牌"),
            "status": ("状态", "待审核", "待复核", "待审批", "告警", "异常"),
            "depot": ("油库", "仓库", "库房", "库区", "库点"),
            "material": ("物资", "物料", "办公用品", "材料"),
            "date": ("每日", "日趋势", "按天", "每天", "按日", "逐日"),
            "month": ("按月", "每月", "月度", "逐月"),
        },
        "intents": {
            "summary": ("概况", "情况", "怎么样", "汇总", "总览", "统计", "数据"),
            "ranking": ("排行", "排名", "top", "哪个", "哪家", "最高", "最多", "最好", "前几", "第一"),
            "trend": ("趋势", "走势", "每日", "日销量", "波动"),
            "trend_anomaly": ("掉", "下降", "下滑", "跌", "减少", "掉得厉害", "异常"),
            "comparison": ("同比", "环比", "对比", "较去年", "较上月", "增长率", "变化"),
            "status": ("状态", "停业", "暂停营业", "维修", "停电", "缺站"),
            "data_availability": ("数据", "传回来", "补回来", "补传", "上传", "更新", "到了", "有没有"),
            "visualization": ("图", "图表", "趋势图", "折线图", "柱状图", "柱形图", "饼图", "环图", "可视化", "画图", "做图"),
        },
    }
_ASSISTANT_DOMAIN_OPTIONS = ("零售", "自用油", "批发", "合并", "吉方日报", "采购", "称重", "物资", "系统健康")
_ASSISTANT_METRIC_OPTIONS = ("销量", "销售额", "吨数")
_ASSISTANT_RANKING_DIMENSION_OPTIONS = ("站点", "油品", "区域", "客户", "供应商", "车辆", "状态")
_ASSISTANT_COMPARISON_SIGNAL_WORDS = (
    "同比", "环比", "对比", "较去年", "比去年", "较上月", "比上月", "比上个月",
    "和去年", "跟去年", "与去年", "去年和今年", "今年和去年",
    "上个月同期", "上月同期", "去年同期", "上年同期", "增长率", "变化",
)
_ASSISTANT_RANKING_SIGNAL_WORDS = (
    "排行", "排名", "top", "TOP", "哪个", "哪家", "最高", "最多", "最好", "卖得最好", "前几", "第一名",
)
_ASSISTANT_TREND_SIGNAL_WORDS = (
    "趋势", "走势", "每日", "日销量", "波动", "掉", "下降", "下滑", "跌", "减少", "掉得厉害", "异常",
)
_ASSISTANT_CHART_SIGNAL_WORDS = (
    "图", "图表", "趋势图", "折线图", "柱状图", "柱形图", "饼图", "环图", "可视化", "画图", "做图",
)
_ASSISTANT_EXPLICIT_TIME_PATTERN = re.compile(
    rf"(?:20\d{{2}}|\d{{2}})\s*年|"
    rf"(?:今年|本年|当前年|当年|这个年|去年|上年|上一年|前年|大前年|全年|年度|年初|月初|本月|这个月|上月|上个月|"
    rf"本周|这周|上周|昨天|今天|今日|上半年|下半年|前半年度|后半年度|季度|春节|五一|劳动节)|"
    rf"(?:{_ASSISTANT_MONTH_NAME_PATTERN})\s*月(?:份)?|"
    rf"(?:最近|近|过去|前)\s*(?:半|{_ASSISTANT_COUNT_NAME_PATTERN})\s*(?:个\s*)?(?:天|日|周|星期|月|年)"
)


def _assistant_question_score(value):
    text = str(value or "")
    score = len(re.findall(r"\d{1,2}", text))
    score += min(len(re.findall(r"[\u4e00-\u9fff]", text)), 24)
    score -= len(re.findall(r"[�]", text)) * 5
    for word in (
        "年", "月", "周", "今天", "昨天", "零售", "批发", "销量", "销售", "金额", "数据", "统计",
        "自用油", "采购", "称重", "物资", "调拨", "损耗", "知识库", "审计", "权限", "别名",
        "市场", "观察", "行情", "俄罗斯", "系统健康", "管理中心", "设置",
    ):
        if word in text:
            score += 3
    return score


def _assistant_prepare_question_text(value):
    text = str(value or "").strip().translate(_ASSISTANT_DIGIT_TRANSLATION)
    candidates = [text]
    for encoding in ("cp1252", "latin1", "gbk", "cp936"):
        try:
            candidate = text.encode(encoding, errors="ignore").decode("utf-8", errors="ignore")
        except (LookupError, UnicodeError):
            continue
        if candidate and candidate not in candidates:
            candidates.append(candidate.translate(_ASSISTANT_DIGIT_TRANSLATION))
    return max(candidates, key=_assistant_question_score)


def _assistant_month_number(value):
    text = str(value or "").strip()
    if not text:
        return None
    try:
        month = int(text)
        return month if 1 <= month <= 12 else None
    except ValueError:
        pass
    month_map = {
        "一": 1,
        "二": 2,
        "两": 2,
        "三": 3,
        "四": 4,
        "五": 5,
        "六": 6,
        "七": 7,
        "八": 8,
        "九": 9,
        "十": 10,
        "十一": 11,
        "十二": 12,
    }
    return month_map.get(text)


def _assistant_count_number(value):
    text = str(value or "").strip()
    if not text:
        return None
    try:
        number = int(text)
        return number if number > 0 else None
    except ValueError:
        pass

    digit_map = {
        "一": 1,
        "二": 2,
        "两": 2,
        "俩": 2,
        "三": 3,
        "四": 4,
        "五": 5,
        "六": 6,
        "七": 7,
        "八": 8,
        "九": 9,
    }
    if text in digit_map:
        return digit_map[text]
    if "十" in text:
        left, _, right = text.partition("十")
        tens = digit_map.get(left, 1) if left else 1
        ones = digit_map.get(right, 0) if right else 0
        number = tens * 10 + ones
        return number if number > 0 else None
    return None


def _assistant_year_number(value):
    try:
        year = int(str(value).strip())
    except (TypeError, ValueError):
        return None
    if 0 <= year < 100:
        year += 2000
    return year if 1900 <= year <= 2099 else None


def _assistant_context_year_number(value, today):
    text = str(value or "").strip()
    if not text:
        return today.year
    if text in {"今年", "本年", "当前年", "当年", "这个年"}:
        return today.year
    if text in {"去年", "上年", "上一年"}:
        return today.year - 1
    if text == "前年":
        return today.year - 2
    if text == "大前年":
        return today.year - 3
    return _assistant_year_number(text)


def _assistant_quarter_number(value):
    text = str(value or "").strip().upper().lstrip("Q")
    quarter_map = {"一": 1, "二": 2, "三": 3, "四": 4}
    if text in quarter_map:
        return quarter_map[text]
    try:
        quarter = int(text)
    except (TypeError, ValueError):
        return None
    return quarter if 1 <= quarter <= 4 else None


def _assistant_safe_date(year, month, day):
    normalized_year = _assistant_year_number(year)
    if normalized_year is None:
        return None
    try:
        return date(normalized_year, int(month), int(day))
    except (TypeError, ValueError):
        return None


def _assistant_shift_months(value, months):
    month_index = value.year * 12 + value.month - 1 + months
    year = month_index // 12
    month = month_index % 12 + 1
    day = min(value.day, monthrange(year, month)[1])
    return date(year, month, day)


def _assistant_month_to_date(year, month, anchor_day):
    day = min(int(anchor_day), monthrange(year, month)[1])
    start = date(year, month, 1)
    return start, date(year, month, day) + timedelta(days=1)


def _assistant_calendar_period(text, today):
    if _assistant_contains_any(text, ("月初到现在", "月初至今", "本月以来", "这个月以来")):
        return "本月以来", today.replace(day=1), today + timedelta(days=1)
    if _assistant_contains_any(text, ("年初到现在", "年初至今", "今年以来")):
        return f"{today.year}年以来", date(today.year, 1, 1), today + timedelta(days=1)

    if _assistant_contains_any(text, ("上个月同期", "上月同期")):
        if _assistant_contains_any(text, ("比", "较", "对比", "环比")):
            return "本月以来", today.replace(day=1), today + timedelta(days=1)
        previous_month = _assistant_shift_months(today.replace(day=1), -1)
        start, end = _assistant_month_to_date(previous_month.year, previous_month.month, today.day)
        return "上个月同期", start, end

    if _assistant_contains_any(text, ("去年同期", "上年同期")):
        if _assistant_contains_any(text, ("比", "较", "对比", "同比")):
            return "本月以来", today.replace(day=1), today + timedelta(days=1)
        start, end = _assistant_month_to_date(today.year - 1, today.month, today.day)
        return "去年同期", start, end

    if _assistant_contains_any(text, ("春节期间", "春节假期")):
        month_day = _ASSISTANT_SPRING_FESTIVAL_DATES.get(today.year)
        if month_day:
            start = date(today.year, month_day[0], month_day[1])
            return f"{today.year}年春节期间", start, start + timedelta(days=7)
    if "春节后" in text:
        month_day = _ASSISTANT_SPRING_FESTIVAL_DATES.get(today.year)
        if month_day:
            start = date(today.year, month_day[0], month_day[1])
            return f"{today.year}年春节后", start, today + timedelta(days=1)

    if _assistant_contains_any(text, ("五一期间", "劳动节期间")):
        return f"{today.year}年五一期间", date(today.year, 5, 1), date(today.year, 5, 6)
    if _assistant_contains_any(text, ("五一后", "劳动节后")):
        return f"{today.year}年五一后", date(today.year, 5, 1), today + timedelta(days=1)

    return None


def _assistant_rolling_period(text, today):
    half_year = re.search(r"(?:最近|近|过去)\s*半\s*年", text)
    if half_year:
        return "最近半年", _assistant_shift_months(today, -6), today + timedelta(days=1)

    match = re.search(
        rf"(?:最近|近|过去|前)\s*({_ASSISTANT_COUNT_NAME_PATTERN})\s*(?:个\s*)?(天|日|周|星期|月|年)",
        text,
    )
    if not match:
        return None

    count = _assistant_count_number(match.group(1))
    unit = match.group(2)
    if not count:
        return None
    if unit in {"天", "日"}:
        if count > 3650:
            return None
        return f"最近{count}天", today - timedelta(days=count), today + timedelta(days=1)
    if unit in {"周", "星期"}:
        if count > 520:
            return None
        return f"最近{count}周", today - timedelta(days=count * 7), today + timedelta(days=1)
    if unit == "月":
        if count > 120:
            return None
        return f"最近{count}个月", _assistant_shift_months(today, -count), today + timedelta(days=1)
    if unit == "年":
        if count > 10:
            return None
        return f"最近{count}年", _assistant_shift_months(today, -count * 12), today + timedelta(days=1)
    return None


def _assistant_explicit_date_range(text):
    full_range = re.search(
        rf"{_ASSISTANT_DATE_PATTERN}\s*{_ASSISTANT_RANGE_PATTERN}\s*{_ASSISTANT_DATE_PATTERN}",
        text,
    )
    if full_range:
        start = _assistant_safe_date(*full_range.group(1, 2, 3))
        end_day = _assistant_safe_date(*full_range.group(4, 5, 6))
        if start and end_day and end_day >= start:
            label = f"{start.year}年{start.month}月{start.day}日至{end_day.year}年{end_day.month}月{end_day.day}日"
            return label, start, end_day + timedelta(days=1)

    same_year_range = re.search(
        rf"{_ASSISTANT_DATE_PATTERN}\s*{_ASSISTANT_RANGE_PATTERN}\s*"
        r"(\d{1,2})\s*月\s*(\d{1,2})\s*(?:日|号)?",
        text,
    )
    if same_year_range:
        start = _assistant_safe_date(*same_year_range.group(1, 2, 3))
        end_day = _assistant_safe_date(same_year_range.group(1), same_year_range.group(4), same_year_range.group(5))
        if start and end_day and end_day >= start:
            label = f"{start.year}年{start.month}月{start.day}日至{end_day.month}月{end_day.day}日"
            return label, start, end_day + timedelta(days=1)

    return None


def _assistant_explicit_day_period(text, today):
    full_day = re.search(_ASSISTANT_DATE_PATTERN, text)
    if full_day:
        target = _assistant_safe_date(*full_day.group(1, 2, 3))
        if target:
            return f"{target.year}年{target.month}月{target.day}日", target, target + timedelta(days=1)

    year_prefix = r"20\d{2}|\d{2}|今年|本年|当前年|当年|这个年|去年|上年|上一年|前年|大前年"
    relative_day = re.search(
        rf"(?<!\d)(?:({year_prefix})\s*年?\s*)?({_ASSISTANT_MONTH_NAME_PATTERN})\s*月\s*"
        rf"({_ASSISTANT_COUNT_NAME_PATTERN})\s*(?:日|号)",
        text,
    )
    if not relative_day:
        return None
    year = _assistant_context_year_number(relative_day.group(1), today)
    month = _assistant_month_number(relative_day.group(2))
    day = _assistant_count_number(relative_day.group(3))
    if year is None or not month or not day:
        return None
    target = _assistant_safe_date(year, month, day)
    if not target:
        return None
    return f"{target.year}年{target.month}月{target.day}日", target, target + timedelta(days=1)


def _assistant_half_year_period(text, today):
    year_prefix = r"20\d{2}|\d{2}|今年|本年|当前年|当年|这个年|去年|上年|上一年|前年|大前年"
    match = re.search(rf"(?<!\d)(?:({year_prefix})\s*年?\s*)?(上半年度?|前半年度?)", text)
    if match:
        year = _assistant_context_year_number(match.group(1), today)
        if year is not None:
            return f"{year}年上半年", date(year, 1, 1), date(year, 7, 1)

    match = re.search(rf"(?<!\d)(?:({year_prefix})\s*年?\s*)?(下半年度?|后半年度?)", text)
    if match:
        year = _assistant_context_year_number(match.group(1), today)
        if year is not None:
            return f"{year}年下半年", date(year, 7, 1), date(year + 1, 1, 1)

    return None


def _assistant_quarter_period(text, today):
    current_quarter = (today.month - 1) // 3 + 1
    if "本季度" in text or "这个季度" in text or "当前季度" in text:
        start_month = (current_quarter - 1) * 3 + 1
        start = date(today.year, start_month, 1)
        return f"{today.year}年第{current_quarter}季度", start, today + timedelta(days=1)
    if "上季度" in text or "上个季度" in text or "上一季度" in text:
        quarter = current_quarter - 1
        year = today.year
        if quarter == 0:
            quarter = 4
            year -= 1
        start_month = (quarter - 1) * 3 + 1
        start = date(year, start_month, 1)
        end = date(year + 1, 1, 1) if quarter == 4 else date(year, start_month + 3, 1)
        return f"{year}年第{quarter}季度", start, end

    year_prefix = r"20\d{2}|\d{2}|今年|本年|当前年|当年|这个年|去年|上年|上一年|前年|大前年"
    match = re.search(
        rf"(?<!\d)(?:({year_prefix})\s*年?\s*)?(?:第?\s*({_ASSISTANT_QUARTER_NAME_PATTERN})\s*季度|[Qq]\s*([1-4]))",
        text,
    )
    if not match:
        return None
    year = _assistant_context_year_number(match.group(1), today)
    quarter = _assistant_quarter_number(match.group(2) or match.group(3))
    if year is None or quarter is None:
        return None
    start_month = (quarter - 1) * 3 + 1
    start = date(year, start_month, 1)
    end = date(year + 1, 1, 1) if quarter == 4 else date(year, start_month + 3, 1)
    return f"{year}年第{quarter}季度", start, end


def parse_assistant_period(question, today=None):
    today = today or date.today()
    text = _assistant_prepare_question_text(question)
    explicit_range = _assistant_explicit_date_range(text)
    if explicit_range:
        return explicit_range
    explicit_day = _assistant_explicit_day_period(text, today)
    if explicit_day:
        return explicit_day
    calendar_period = _assistant_calendar_period(text, today)
    if calendar_period:
        return calendar_period
    rolling = _assistant_rolling_period(text, today)
    if rolling:
        return rolling
    half_year = _assistant_half_year_period(text, today)
    if half_year:
        return half_year
    quarter = _assistant_quarter_period(text, today)
    if quarter:
        return quarter
    if "昨天" in text:
        start = today - timedelta(days=1)
        return "昨天", start, today
    if "今天" in text or "今日" in text:
        return "今天", today, today + timedelta(days=1)
    current_monday = today - timedelta(days=today.weekday())
    if "上周" in text or "上星期" in text:
        start = current_monday - timedelta(days=7)
        return "上周", start, current_monday
    if "本周" in text or "这周" in text:
        return "本周", current_monday, today + timedelta(days=1)
    if "上月" in text:
        first_this_month = today.replace(day=1)
        last_month_end = first_this_month
        if first_this_month.month == 1:
            start = date(first_this_month.year - 1, 12, 1)
        else:
            start = date(first_this_month.year, first_this_month.month - 1, 1)
        return "上月", start, last_month_end
    match = re.search(r"(?<!\d)(20\d{2}|\d{2})\s*[年/-]\s*(\d{1,2})\s*月?", text)
    if match:
        year = _assistant_year_number(match.group(1))
        month = int(match.group(2))
        if year is not None and 1 <= month <= 12:
            start = date(year, month, 1)
            end = date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)
            return f"{year}年{month}月", start, end
    relative_year_month = re.search(
        rf"(去年|上年|上一年|前年|大前年|今年|本年|当前年|当年|这个年)\s*({_ASSISTANT_MONTH_NAME_PATTERN})\s*月",
        text,
    )
    if relative_year_month:
        year = _assistant_context_year_number(relative_year_month.group(1), today)
        month = _assistant_month_number(relative_year_month.group(2))
        if year is not None and month:
            start = date(year, month, 1)
            end = date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)
            return f"{year}年{month}月", start, end
    relative_month = re.search(rf"(?:当年|今年|本年|当前年|这个年)\s*({_ASSISTANT_MONTH_NAME_PATTERN})\s*月", text)
    if not relative_month:
        relative_month = re.search(rf"(?<!\d)({_ASSISTANT_MONTH_NAME_PATTERN})\s*月(?:份)?(?=.*(?:零售|批发|销量|销售|金额|数据|统计))", text)
    if not relative_month:
        relative_month = re.search(rf"(?<!\d)({_ASSISTANT_MONTH_NAME_PATTERN})\s*月(?:份)?(?!\s*\d)", text)
    if relative_month:
        month = _assistant_month_number(relative_month.group(1))
        if month:
            start = date(today.year, month, 1)
            end = date(today.year + 1, 1, 1) if month == 12 else date(today.year, month + 1, 1)
            return f"{today.year}年{month}月", start, end
    if "本月" in text or "这个月" in text:
        start = today.replace(day=1)
        return "本月", start, today + timedelta(days=1)
    year_match = re.search(r"(?<!\d)(20\d{2}|\d{2})\s*年(?!\s*\d)", text)
    if year_match:
        year = _assistant_year_number(year_match.group(1))
        if year is not None:
            return f"{year}年", date(year, 1, 1), date(year + 1, 1, 1)
    if "大前年" in text:
        year = today.year - 3
        return f"{year}年", date(year, 1, 1), date(year + 1, 1, 1)
    if "前年" in text:
        year = today.year - 2
        return f"{year}年", date(year, 1, 1), date(year + 1, 1, 1)
    if _assistant_contains_any(text, ("去年", "上年", "上一年")):
        year = today.year - 1
        return f"{year}年", date(year, 1, 1), date(today.year, 1, 1)
    if _assistant_contains_any(text, ("今年", "本年", "当前年", "当年", "这个年")):
        return f"{today.year}年", date(today.year, 1, 1), today + timedelta(days=1)
    if "年销量" in text or "年度" in text or "全年" in text:
        return f"{today.year}年", date(today.year, 1, 1), today + timedelta(days=1)
    start = today - timedelta(days=7)
    return "最近7天", start, today + timedelta(days=1)


def _assistant_parse_iso_date(value):
    if not value:
        return None
    try:
        return date.fromisoformat(str(value)[:10])
    except (TypeError, ValueError):
        return None


def _assistant_read_snapshot_json(path):
    try:
        if not path.exists():
            return None
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        return data if isinstance(data, dict) else None
    except (OSError, json.JSONDecodeError):
        return None


def _assistant_snapshot_query_end(payload):
    if not isinstance(payload, dict):
        return None
    summary = payload.get("summary") if isinstance(payload.get("summary"), dict) else {}
    for source in (summary, payload):
        for key in ("queryEnd", "queryEndExclusive"):
            parsed = _assistant_parse_iso_date(source.get(key))
            if parsed:
                return parsed
    daily_dates = [
        _assistant_parse_iso_date(item.get("date"))
        for item in (payload.get("daily") or [])
        if isinstance(item, dict)
    ]
    daily_dates = [item for item in daily_dates if item]
    return max(daily_dates) + timedelta(days=1) if daily_dates else None


def _assistant_latest_retail_operation_end(today):
    """Return the exclusive end date of the latest completed retail operation day."""
    limit = today + timedelta(days=1)
    year_start = date(today.year, 1, 1)
    candidates = []

    monthly_root = OILSTATION_ROOT / "cache_snapshots" / "monthly" / str(today.year)
    for month in range(min(today.month, 12), 0, -1):
        payload = _assistant_read_snapshot_json(monthly_root / f"monthly_{month:02d}.json")
        query_end = _assistant_snapshot_query_end(payload)
        if query_end and year_start < query_end <= limit:
            candidates.append(query_end)
            break

    daily_root = OILSTATION_ROOT / "cache_snapshots" / "daily" / str(today.year)
    if daily_root.exists():
        for path in sorted(daily_root.glob("*/day_*.json"), reverse=True):
            try:
                month = int(path.parent.name)
                day = int(path.stem.split("_")[-1])
                snapshot_day = date(today.year, month, day)
            except (TypeError, ValueError):
                continue
            query_end = snapshot_day + timedelta(days=1)
            if year_start < query_end <= limit:
                candidates.append(query_end)
                break

    return max(candidates) if candidates else None


def _assistant_cap_period_to_operational_cutoff(label, start, end, today, operational_end=None):
    if not start or not end or end <= start:
        return label, start, end
    cutoff = operational_end or _assistant_latest_retail_operation_end(today) or (today + timedelta(days=1))
    cutoff = min(cutoff, today + timedelta(days=1))
    if start <= today < end and start < cutoff < end:
        return label, start, cutoff
    return label, start, end


def _assistant_route_uses_retail_operation_cutoff(domain, route):
    if domain in {"retail", "self_use", "combined"}:
        return True
    return route in {
        "retail_summary",
        "retail_ranking",
        "retail_trend",
        "retail_comparison",
        "retail_kpi",
        "retail_inventory",
        "retail_price",
        "retail_resource_share",
        "station_trend_anomaly",
        "station_data_check",
        "self_use_summary",
        "combined_summary",
    }


def _assistant_contains_any(text, words):
    folded = str(text or "").casefold()
    return any(str(word).casefold() in folded for word in words)


def _assistant_matching_terms(text, words):
    folded = str(text or "").casefold()
    matches = []
    for word in words:
        item = str(word)
        if item and item.casefold() in folded:
            matches.append(item)
    return matches


def _assistant_business_terms(text):
    result = {}
    for category, groups in _ASSISTANT_BUSINESS_LEXICON.items():
        category_matches = {}
        for name, words in groups.items():
            matches = _assistant_matching_terms(text, words)
            if matches:
                category_matches[name] = matches
        result[category] = category_matches
    return result


def _assistant_has_explicit_domain(business_terms):
    domains = business_terms.get("domains") if isinstance(business_terms, dict) else {}
    return bool(domains)


def _assistant_metric_from_terms(business_terms):
    metrics = business_terms.get("metrics") if isinstance(business_terms, dict) else {}
    for metric in ("amount", "tons", "volume"):
        if metrics.get(metric):
            return metric
    return None


def _assistant_domain_from_terms(business_terms, default="retail"):
    domain_terms = business_terms.get("domains") if isinstance(business_terms, dict) else {}
    weighing_terms = set(domain_terms.get("weighing") or [])
    if weighing_terms.intersection({"称重", "地磅", "过磅", "磅单", "审核称重"}):
        return "weighing"
    for domain_name in (
        "system_health",
        "knowledge",
        "audit",
        "alias",
        "users",
        "management",
        "settings",
        "transfer_loss",
        "materials",
        "procurement",
    ):
        if domain_terms.get(domain_name):
            return domain_name
    if domain_terms.get("market"):
        return "market"
    if domain_terms.get("exchange"):
        return "exchange"
    if domain_terms.get("density"):
        return "density"
    if domain_terms.get("contract"):
        return "contract"
    if domain_terms.get("accounting"):
        return "accounting"
    if domain_terms.get("self_use"):
        return "self_use"
    if domain_terms.get("combined") and not domain_terms.get("retail"):
        return "combined"
    if domain_terms.get("wholesale"):
        return "wholesale"
    if domain_terms.get("retail"):
        return "retail"
    if domain_terms.get("weighing"):
        return "weighing"
    return default


def _assistant_dimensions_from_terms(business_terms):
    dimensions = []
    dimension_terms = business_terms.get("dimensions") if isinstance(business_terms, dict) else {}
    for dimension_name in (
        "station", "resource", "region", "customer", "supplier", "vehicle",
        "status", "depot", "material", "date", "month",
    ):
        if dimension_terms.get(dimension_name):
            dimensions.append(dimension_name)
    return dimensions


def _assistant_chart_requested(text):
    return _assistant_contains_any(text, _ASSISTANT_CHART_SIGNAL_WORDS)


def _assistant_chart_kind(text, default="auto"):
    if _assistant_contains_any(text, ("饼图", "环图")):
        return "pie"
    if _assistant_contains_any(text, ("折线图", "曲线图", "趋势图")):
        return "line"
    if _assistant_contains_any(text, ("柱状图", "柱形图", "条形图")):
        return "bar"
    if _assistant_contains_any(text, ("趋势", "走势", "每日", "按天", "按日", "逐日")):
        return "line"
    if _assistant_contains_any(text, ("占比", "结构", "构成")):
        return "pie"
    return default


def _assistant_resource_from_text(text):
    folded = str(text or "").casefold()
    patterns = [
        (103, "АИ 92", (r"(?:аи|ai|а[\s-]*и)\s*92", r"92\s*号?(?:汽油|油|油品)?")),
        (104, "АИ 95", (r"(?:аи|ai|а[\s-]*и)\s*95", r"95\s*号?(?:汽油|油|油品)?")),
        (102, "АИ 80", (r"(?:аи|ai|а[\s-]*и)\s*80", r"80\s*号?(?:汽油|油|油品)?")),
        (106, "Д/Т", (r"(?:д/т|дт|柴油|дизель)",)),
        (108, "СПБТ", (r"(?:спбт|液化气|燃气|газ)",)),
        (110, "ДТ -35", (r"(?:дт\s*-?\s*35|-35\s*柴油)",)),
    ]
    for key, name, regexes in patterns:
        if any(re.search(pattern, folded, flags=re.IGNORECASE) for pattern in regexes):
            return {"resourceKey": key, "name": name}
    return None


def _assistant_clarification_payload(
    text,
    *,
    domain,
    route,
    explicit_dimensions,
    explicit_metric,
    business_terms,
    followup,
    context_domain=False,
):
    if followup:
        return None

    needs_domain = (
        domain == "retail"
        and route == "retail_summary"
        and not _assistant_has_explicit_domain(business_terms)
        and not context_domain
        and not explicit_dimensions
        and _assistant_contains_any(text, ("数据", "情况", "怎么样", "统计", "分析", "看一下", "查一下"))
    )
    if needs_domain:
        return {
            "question": "我需要先确认业务范围：你想看零售、批发、合并，还是吉方日报？",
            "missing": ["domain"],
            "options": list(_ASSISTANT_DOMAIN_OPTIONS),
        }

    if route == "retail_ranking" and not explicit_dimensions:
        return {
            "question": "你想按哪个维度排行：站点、油品、区域，还是客户？",
            "missing": ["dimension"],
            "options": list(_ASSISTANT_RANKING_DIMENSION_OPTIONS),
        }

    metric_ambiguous = (
        route == "retail_ranking"
        and not explicit_metric
        and _assistant_contains_any(text, ("最高", "最多", "第一", "排行", "排名", "top", "TOP"))
        and not _assistant_contains_any(text, ("卖得", "销量", "销售量", "金额", "销售额", "营业额", "吨", "吨数"))
    )
    if metric_ambiguous:
        return {
            "question": "这个排行要按哪个指标看：销量、销售额，还是吨数？",
            "missing": ["metric"],
            "options": list(_ASSISTANT_METRIC_OPTIONS),
        }

    return None


def _assistant_has_explicit_time_signal(text):
    return bool(_ASSISTANT_EXPLICIT_TIME_PATTERN.search(str(text or "")))


def _assistant_self_check(text, plan):
    checks = []

    def add(code, label, status, detail="", severity=None):
        item = {
            "code": code,
            "label": label,
            "status": status,
            "severity": severity or ("error" if status == "fail" else "warning" if status == "warn" else "info"),
        }
        if detail:
            item["detail"] = detail
        checks.append(item)

    period = plan.get("period") if isinstance(plan.get("period"), dict) else {}
    comparison = plan.get("comparison") if isinstance(plan.get("comparison"), dict) else None
    tool = plan.get("tool") if isinstance(plan.get("tool"), dict) else {}
    explicit = plan.get("explicit") if isinstance(plan.get("explicit"), dict) else {}
    dimensions = plan.get("dimensions") if isinstance(plan.get("dimensions"), list) else []
    route = str(plan.get("route") or "")
    needs_clarification = bool(plan.get("needsClarification"))
    period_label = str(period.get("label") or "")

    if _assistant_has_explicit_time_signal(text):
        if period and period_label != "最近7天":
            add("period_explicit", "时间条件识别", "pass", f"已解析为 {period_label}")
        else:
            add("period_explicit", "时间条件识别", "fail", "问题里有明确时间词，但计划仍落到最近7天")
    elif period_label == "最近7天":
        add("period_default", "默认时间假设", "warn", "问题未明确时间，默认使用最近7天")
    else:
        add("period_inferred", "时间条件识别", "pass", f"已解析为 {period_label}")

    if _assistant_contains_any(text, _ASSISTANT_COMPARISON_SIGNAL_WORDS):
        if comparison and isinstance(comparison.get("period"), dict):
            add("comparison_baseline", "对比基准区间", "pass", comparison.get("description") or "已生成对比基准")
        else:
            add("comparison_baseline", "对比基准区间", "fail", "问题要求同环比/对比，但没有生成基准区间")

    if _assistant_contains_any(text, _ASSISTANT_RANKING_SIGNAL_WORDS):
        if route in {"retail_ranking", "wholesale_summary", "procurement_summary", "weighing_summary"} and dimensions:
            add("ranking_dimension", "排行维度", "pass", "、".join(dimensions))
        elif needs_clarification:
            add("ranking_dimension", "排行维度", "warn", "排行问题需要用户补充维度或指标")
        elif route in {"station_trend_anomaly", "compound_query"}:
            add("ranking_dimension", "排行/下钻意图", "pass", route)
        else:
            add("ranking_dimension", "排行维度", "fail", "问题像排行，但未选中排行工具或维度")

    if _assistant_contains_any(text, _ASSISTANT_TREND_SIGNAL_WORDS):
        if route in {"retail_trend", "station_trend_anomaly", "compound_query"}:
            add("trend_route", "趋势/异常工具", "pass", route)
        elif needs_clarification:
            add("trend_route", "趋势/异常工具", "warn", "趋势问题需要先澄清")
        else:
            add("trend_route", "趋势/异常工具", "warn", "问题包含趋势/下滑词，但当前工具不是趋势工具")

    visual = plan.get("visual") if isinstance(plan.get("visual"), dict) else {}
    if _assistant_chart_requested(text):
        if visual.get("requested"):
            add("chart_request", "图表回答", "pass", str(visual.get("preferredType") or "auto"))
        else:
            add("chart_request", "图表回答", "warn", "问题要求图表，但计划未记录图表偏好")

    if tool.get("readOnly") is True:
        add("readonly_tool", "只读工具约束", "pass", str(tool.get("name") or route))
    else:
        add("readonly_tool", "只读工具约束", "fail", "计划未声明只读工具")

    if route == "compound_query":
        subplans = plan.get("subPlans") if isinstance(plan.get("subPlans"), list) else []
        if len(subplans) >= 2:
            add("compound_decomposition", "复合问题拆解", "pass", f"已拆解 {len(subplans)} 个子问题")
        else:
            add("compound_decomposition", "复合问题拆解", "fail", "复合问题没有形成足够子计划")

    if explicit.get("period") is False and period_label == "最近7天":
        assumptions = plan.get("assumptions") if isinstance(plan.get("assumptions"), list) else []
        if "未明确时间时默认查询最近7天" not in assumptions:
            assumptions.append("未明确时间时默认查询最近7天")
            plan["assumptions"] = assumptions

    failures = [item for item in checks if item["status"] == "fail"]
    warnings = [item for item in checks if item["status"] == "warn"]
    status = "fail" if failures else "warn" if warnings or needs_clarification else "pass"
    next_action = "clarify" if failures or needs_clarification else "execute_with_assumptions" if warnings else "execute"
    score = max(0.1, round(1 - len(failures) * 0.35 - len(warnings) * 0.12, 2))
    return {
        "status": status,
        "score": score,
        "nextAction": next_action,
        "checks": checks,
        "risks": [item["detail"] for item in [*failures, *warnings] if item.get("detail")],
    }


def _assistant_self_check_clarification(self_check):
    checks = self_check.get("checks") if isinstance(self_check, dict) else []
    failed = [item for item in checks if isinstance(item, dict) and item.get("status") == "fail"]
    if not failed:
        return None
    missing = []
    if any(item.get("code") == "period_explicit" for item in failed):
        missing.append("period")
    if any(item.get("code") == "comparison_baseline" for item in failed):
        missing.append("comparison")
    if any(item.get("code") == "ranking_dimension" for item in failed):
        missing.append("dimension")
    if not missing:
        missing.append("condition")
    return {
        "question": "我识别到你的问题里有明确条件，但计划自检没有通过。请补充或改写查询条件后我再查，避免答错口径。",
        "missing": list(dict.fromkeys(missing)),
        "options": ["今天", "昨天", "本月", "今年", "26年上半年", "最近30天"],
    }


def _assistant_finalize_plan(plan, text):
    self_check = _assistant_self_check(text, plan)
    plan["selfCheck"] = self_check
    if self_check.get("status") == "fail" and not plan.get("needsClarification"):
        clarification = _assistant_self_check_clarification(self_check)
        if clarification:
            plan["needsClarification"] = True
            plan["clarification"] = clarification
            plan["confidence"] = min(float(plan.get("confidence") or 0.5), 0.42)
    elif self_check.get("status") == "warn":
        plan["confidence"] = min(float(plan.get("confidence") or 0.5), 0.82)
    return plan


def _assistant_station_from_text(text):
    catalog = retail_station_catalog_cached()
    folded = str(text or "").casefold()
    for match in re.finditer(r"(?<!\d)(\d{1,3})\s*(?:号\s*)?(?:站|加油站)", folded):
        shop_key = int(match.group(1))
        if shop_key in catalog:
            return dict(catalog[shop_key])
    for shop_key, item in catalog.items():
        names = [
            str(item.get("name") or "").strip(),
            str(item.get("short") or "").strip(),
            f"{shop_key}号站",
            f"{shop_key}站",
        ]
        if any(name and name.casefold() in folded for name in names):
            return dict(item)
    return None


def _assistant_period_payload(label, start, end):
    end_day = end - timedelta(days=1)
    return {
        "label": label,
        "start": start.isoformat(),
        "endExclusive": end.isoformat(),
        "endInclusive": end_day.isoformat(),
    }


def _assistant_period_from_payload(period):
    if not isinstance(period, dict):
        return None
    try:
        label = str(period.get("label") or "").strip() or "区间"
        start = date.fromisoformat(str(period.get("start")))
        end = date.fromisoformat(str(period.get("endExclusive")))
    except (TypeError, ValueError):
        return None
    return (label, start, end) if end > start else None


def _assistant_context_page(context):
    context = _assistant_context_dict(context)
    page = context.get("pageContext")
    return page if isinstance(page, dict) else {}


def _assistant_context_reference_requested(text):
    return _assistant_contains_any(
        text,
        ("这个", "该", "当前", "本页", "这里", "它", "这个页面", "这个站", "该站", "这家站", "这个油品", "该油品"),
    )


def _assistant_station_from_context(context):
    page = _assistant_context_page(context)
    catalog = retail_station_catalog_cached()
    candidates = []
    for source in (page, _assistant_context_dict(context)):
        search = source.get("searchParams") if isinstance(source.get("searchParams"), dict) else {}
        for key in ("shopKey", "stationKey", "stationId", "station", "shop", "id"):
            if source.get(key) is not None:
                candidates.append(source.get(key))
            if search.get(key) is not None:
                candidates.append(search.get(key))
    for value in candidates:
        try:
            shop_key = int(float(str(value).strip()))
        except (TypeError, ValueError):
            continue
        if shop_key in catalog:
            return dict(catalog[shop_key])
    return None


def _assistant_resource_from_context(context):
    page = _assistant_context_page(context)
    candidates = []
    for source in (page, _assistant_context_dict(context)):
        search = source.get("searchParams") if isinstance(source.get("searchParams"), dict) else {}
        for key in ("resourceKey", "productKey", "resource", "product", "oilKey"):
            if source.get(key) is not None:
                candidates.append(source.get(key))
            if search.get(key) is not None:
                candidates.append(search.get(key))
    resource_names = {
        103: "АИ 92",
        104: "АИ 95",
        102: "АИ 80",
        106: "Д/Т",
        108: "СПБТ",
        110: "ДТ -35",
    }
    for value in candidates:
        try:
            resource_key = int(float(str(value).strip()))
        except (TypeError, ValueError):
            resource_key = None
        if resource_key in resource_names:
            return {"resourceKey": resource_key, "name": resource_names[resource_key]}
    return None


def _assistant_page_hints(page_context):
    path = str(page_context.get("path") or page_context.get("pathname") or "").casefold()
    page = str(page_context.get("page") or "").casefold()
    module = str(page_context.get("module") or "").casefold()
    search = page_context.get("searchParams") if isinstance(page_context.get("searchParams"), dict) else {}
    hints = {"contextUsed": []}
    if "retail" in path or module == "retail":
        hints["domain"] = "retail"
    if "wholesale" in path or module == "wholesale":
        hints["domain"] = "wholesale"
    if "contract" in path:
        hints["domain"] = "contract"
    if "procurement" in path:
        hints["domain"] = "procurement"
    if "weighing" in path:
        hints["domain"] = "weighing"
    if "materials" in path:
        hints["domain"] = "materials"
    if "transfer-loss" in path:
        hints["domain"] = "transfer_loss"
    if "knowledge" in path:
        hints["domain"] = "knowledge"
    if "audit-log" in path:
        hints["domain"] = "audit"
    if "aliases" in path:
        hints["domain"] = "alias"
    if "users" in path or "organization" in path:
        hints["domain"] = "users"
    if "settings" in path:
        hints["domain"] = "settings"
    if "management-center" in path:
        hints["domain"] = "management"
    if "exchange" in path:
        hints["domain"] = "exchange"
    if "system-health" in path:
        hints["domain"] = "accounting" if str(search.get("tab") or "").casefold() == "accounting" else "system_health"
    if "station" in path or page == "station":
        hints["dimensions"] = ["station"]
    elif "resource" in path or page == "resource":
        hints["dimensions"] = ["resource"]
    elif "region" in path or page == "region":
        hints["dimensions"] = ["region"]
    elif "customer" in path or "customers" in path:
        hints["dimensions"] = ["customer"]
    elif "procurement" in path:
        hints["dimensions"] = ["supplier"]
    elif "weighing" in path:
        hints["dimensions"] = ["status"]
    elif "materials" in path:
        hints["dimensions"] = ["material"]
    if hints.get("domain") or hints.get("dimensions"):
        hints["contextUsed"].append("page")
    return hints


def _assistant_shift_period(label, start, end, *, months=0, days=0, shifted_label=None):
    if months:
        shifted_start = _assistant_shift_months(start, months)
        shifted_end = _assistant_shift_months(end, months)
    else:
        shifted_start = start + timedelta(days=days)
        shifted_end = end + timedelta(days=days)
    return {
        "mode": shifted_label or label,
        "period": _assistant_period_payload(shifted_label or label, shifted_start, shifted_end),
    }


def _assistant_comparison_payload(text, period_label, start, end):
    if _assistant_contains_any(text, ("同比", "去年同期", "上年同期", "较去年", "比去年", "对比去年", "和去年", "跟去年", "与去年", "去年和今年", "今年和去年")):
        payload = _assistant_shift_period(period_label, start, end, months=-12, shifted_label="去年同期")
        payload["type"] = "year_over_year"
        payload["description"] = "当前区间与去年同期对比"
        return payload
    if _assistant_contains_any(text, ("环比", "上月同期", "上个月同期", "较上月", "比上月", "比上个月", "对比上月")):
        payload = _assistant_shift_period(period_label, start, end, months=-1, shifted_label="上月同期")
        payload["type"] = "month_over_month"
        payload["description"] = "当前区间与上月同期对比"
        return payload
    if _assistant_contains_any(text, ("较上周", "比上周", "上周同期")):
        payload = _assistant_shift_period(period_label, start, end, days=-7, shifted_label="上周同期")
        payload["type"] = "week_over_week"
        payload["description"] = "当前区间与上周同期对比"
        return payload
    return None


def _assistant_compound_segments(text):
    normalized = str(text or "").strip()
    if not normalized:
        return []
    if not _assistant_contains_any(normalized, ("、", "；", ";", "，", ",", "同时", "另外", "再看", "再查", "还有")):
        return []
    prepared = re.sub(r"(?:同时|另外|再看|再查|还有)", "；", normalized)
    prepared = re.sub(r"[，,]\s*(?=(?:哪个|哪家|哪些|油品|站点|区域|客户|同比|环比|为什么|原因|排行|下滑|掉|总额|销量|营业额))", "；", prepared)
    parts = [part.strip(" ，,、;；") for part in re.split(r"[;；、]", prepared) if part.strip(" ，,、;；")]
    intentful = [
        part for part in parts
        if part not in {"为什么", "原因", "怎么回事"}
        if _assistant_contains_any(part, (
            "数据", "总额", "汇总", "销量", "销售额", "营业额", "排行", "最高", "最多",
            "哪个", "哪家", "油品", "站", "区域", "客户", "下滑", "掉", "同比", "环比", "为什么", "原因", "缺站"
        ))
    ]
    return intentful if len(intentful) >= 2 else []


def _assistant_segment_has_period(text):
    label, _, _ = parse_assistant_period(text)
    return label != "最近7天"


def _assistant_segment_has_domain(text):
    return _assistant_has_explicit_domain(_assistant_business_terms(text))


def _assistant_enrich_segment(segment, base_plan):
    parts = []
    period = base_plan.get("period") if isinstance(base_plan, dict) else {}
    if isinstance(period, dict) and period.get("label") and not _assistant_segment_has_period(segment):
        parts.append(str(period.get("label")))
    if base_plan.get("domain") == "retail" and not _assistant_segment_has_domain(segment):
        parts.append("零售")
    parts.append(segment)
    return "".join(parts)


def _assistant_metric_from_text(text):
    if _assistant_contains_any(text, ("金额", "销售额", "营业额", "收入", "索姆", "人民币", "元")):
        return "amount"
    if _assistant_contains_any(text, ("吨", "吨数", "吨位", "重量")):
        return "tons"
    if _assistant_contains_any(text, ("销量", "销售量", "升", "升数", "油量", "卖得")):
        return "volume"
    return "amount"


def _assistant_tool_payload(route):
    tool = _ASSISTANT_TOOL_CATALOG.get(route) or {
        "name": route or "unknown",
        "description": "业务只读查询",
        "readOnly": True,
        "dataSource": "business_snapshots",
    }
    return dict(tool)


def _assistant_context_dict(context):
    return context if isinstance(context, dict) else {}


def _assistant_previous_plan(context):
    context = _assistant_context_dict(context)
    plan = context.get("lastAssistantPlan") or context.get("assistantPlan")
    if isinstance(plan, dict):
        return plan
    recent = context.get("recentMessages")
    if isinstance(recent, list):
        for item in reversed(recent):
            if not isinstance(item, dict):
                continue
            candidate = item.get("assistantPlan")
            if isinstance(candidate, dict):
                return candidate
    return {}


def _assistant_context_station(plan):
    entities = plan.get("entities") if isinstance(plan, dict) else {}
    station = entities.get("station") if isinstance(entities, dict) else None
    if not isinstance(station, dict) or not station.get("shopKey"):
        return None
    return dict(station)


def _assistant_period_from_plan(plan):
    period = plan.get("period") if isinstance(plan, dict) else None
    if not isinstance(period, dict):
        return None
    try:
        label = str(period.get("label") or "").strip() or "上下文区间"
        start = date.fromisoformat(str(period.get("start")))
        end = date.fromisoformat(str(period.get("endExclusive")))
    except (TypeError, ValueError):
        return None
    if end <= start:
        return None
    return label, start, end


def _assistant_is_followup_question(text):
    stripped = str(text or "").strip()
    if not stripped:
        return False
    if stripped.startswith(("那", "那么", "还有", "继续")):
        return True
    if stripped.startswith(("只看", "换成", "改成", "改为", "按月", "按天", "按日")):
        return True
    if stripped.startswith(("按站点", "按油品", "按区域", "按客户", "分站", "拆分")):
        return True
    if len(stripped) <= 10 and (_assistant_chart_requested(stripped) or _assistant_contains_any(stripped, ("异常点", "原因分析", "为什么"))):
        return True
    if len(stripped) <= 8 and _assistant_contains_any(stripped, _ASSISTANT_FOLLOWUP_WORDS):
        return True
    return _assistant_contains_any(stripped, ("该站", "这个站", "这家站", "刚才", "上面", "前面", "异常点"))


def _assistant_metric_clarification():
    return {
        "question": "这个排行要按哪个指标看：销量、销售额，还是吨数？",
        "missing": ["metric"],
        "options": list(_ASSISTANT_METRIC_OPTIONS),
    }


def _assistant_route_for_domain(domain, previous_route="retail_summary", previous_intent="summary"):
    if domain in {"market", "exchange", "density", "contract"}:
        return f"{domain}_summary", "summary"
    if domain == "accounting":
        return "accounting_report_status", "status"
    if domain == "combined":
        return "combined_summary", "summary"
    if domain == "wholesale":
        return "wholesale_summary", "summary"
    if domain == "self_use":
        return "self_use_summary", "summary"
    if str(previous_route or "").startswith("retail_") or previous_route in {"station_data_check", "station_trend_anomaly", "station_status"}:
        return previous_route, previous_intent or "summary"
    return "retail_summary", "summary"


def _assistant_resolve_clarification_plan(text, previous_plan, today):
    if not isinstance(previous_plan, dict) or not previous_plan.get("needsClarification"):
        return None
    previous_clarification = previous_plan.get("clarification") if isinstance(previous_plan.get("clarification"), dict) else {}
    missing = previous_clarification.get("missing") if isinstance(previous_clarification, dict) else []
    if not isinstance(missing, list) or not missing:
        return None

    business_terms = _assistant_business_terms(text)
    current_domain = _assistant_domain_from_terms(business_terms, default=None)
    current_dimensions = _assistant_dimensions_from_terms(business_terms)
    current_metric = _assistant_metric_from_terms(business_terms)
    previous_explicit = previous_plan.get("explicit") if isinstance(previous_plan.get("explicit"), dict) else {}

    domain = str(previous_plan.get("domain") or "retail")
    route = str(previous_plan.get("route") or "retail_summary")
    intent = str(previous_plan.get("intent") or "summary")
    metric = previous_plan.get("metric") if previous_plan.get("metric") in {"amount", "volume", "tons"} else "amount"
    dimensions = [str(item) for item in (previous_plan.get("dimensions") or []) if item]
    station = _assistant_station_from_text(text) or _assistant_context_station(previous_plan)
    period = _assistant_period_from_plan(previous_plan)
    if period:
        period_label, start, end = period
    else:
        period_label, start, end = parse_assistant_period(str(previous_plan.get("normalizedQuestion") or text), today=today)

    resolved = []
    context_used = ["clarification", "period"]
    domain_explicit = bool(previous_explicit.get("domain"))
    dimensions_explicit = bool(previous_explicit.get("dimensions"))
    metric_explicit = bool(previous_explicit.get("metric"))

    if "domain" in missing:
        if not current_domain:
            return None
        domain = current_domain
        route, intent = _assistant_route_for_domain(domain, route, intent)
        domain_explicit = True
        resolved.append("domain")
        context_used.append("domain")
    elif current_domain:
        domain = current_domain
        route, intent = _assistant_route_for_domain(domain, route, intent)
        domain_explicit = True

    if "dimension" in missing:
        if not current_dimensions:
            return None
        dimensions = current_dimensions
        dimensions_explicit = True
        resolved.append("dimensions")
        context_used.append("dimensions")
        if "customer" in dimensions and domain == "retail" and not current_domain:
            domain = "wholesale"
            route = "wholesale_summary"
            intent = "ranking"
        elif route == "retail_summary":
            route = "retail_ranking"
            intent = "ranking"
    elif current_dimensions:
        dimensions = current_dimensions
        dimensions_explicit = True

    if "metric" in missing:
        if not current_metric:
            return None
        metric = current_metric
        metric_explicit = True
        resolved.append("metric")
        context_used.append("metric")
    elif current_metric:
        metric = current_metric
        metric_explicit = True

    if route == "retail_ranking" and not dimensions:
        dimensions = ["station"]

    clarification = None
    if route == "retail_ranking" and dimensions_explicit and not metric_explicit:
        clarification = _assistant_metric_clarification()
    if _assistant_route_uses_retail_operation_cutoff(domain, route):
        period_label, start, end = _assistant_cap_period_to_operational_cutoff(period_label, start, end, today)

    confidence = 0.86 if not clarification else 0.48
    plan = {
        "normalizedQuestion": f"{previous_plan.get('normalizedQuestion') or ''} / 澄清：{text}".strip(" /"),
        "domain": domain,
        "intent": intent,
        "route": route,
        "metric": metric,
        "dimensions": dimensions,
        "entities": {"station": station} if station else {},
        "period": _assistant_period_payload(period_label, start, end),
        "tool": _assistant_tool_payload(route),
        "contextUsed": list(dict.fromkeys(context_used)),
        "businessTerms": business_terms,
        "needsClarification": bool(clarification),
        "clarification": clarification,
        "clarificationResolved": {
            "fromQuestion": previous_plan.get("normalizedQuestion"),
            "answer": text,
            "resolved": list(dict.fromkeys(resolved)),
        },
        "explicit": {
            "domain": domain_explicit,
            "dimensions": dimensions if dimensions_explicit else [],
            "metric": metric_explicit,
            "period": True,
        },
        "confidence": confidence,
    }
    return _assistant_finalize_plan(plan, text)


def build_assistant_query_plan(question, today=None, context=None, allow_compound=True):
    today = today or date.today()
    text = _assistant_prepare_question_text(question)
    previous_plan = _assistant_previous_plan(context)
    clarified_plan = _assistant_resolve_clarification_plan(text, previous_plan, today)
    if clarified_plan:
        return clarified_plan
    followup = _assistant_is_followup_question(text) and bool(previous_plan)
    context_used = []
    business_terms = _assistant_business_terms(text)
    page_hints = _assistant_page_hints(_assistant_context_page(context))
    period_label, start, end = parse_assistant_period(text, today=today)
    if (
        _assistant_contains_any(text, ("今年", "本年", "当前年", "当年", "这个年"))
        and _assistant_contains_any(text, ("去年", "上年", "上一年"))
        and _assistant_contains_any(text, _ASSISTANT_COMPARISON_SIGNAL_WORDS)
    ):
        period_label, start, end = f"{today.year}年", date(today.year, 1, 1), today + timedelta(days=1)
    if followup and period_label == "最近7天":
        previous_period = _assistant_period_from_plan(previous_plan)
        if previous_period:
            period_label, start, end = previous_period
            context_used.append("period")
    station = _assistant_station_from_text(text)
    if not station and (followup or _assistant_context_reference_requested(text)):
        station = _assistant_station_from_context(context)
        if station:
            context_used.append("pageStation")
    if followup and not station:
        previous_station = _assistant_context_station(previous_plan)
        if previous_station:
            station = previous_station
            context_used.append("station")
    resource = None if station else _assistant_resource_from_text(text)
    if not resource and not station and (followup or _assistant_context_reference_requested(text)):
        resource = _assistant_resource_from_context(context)
        if resource:
            context_used.append("pageResource")
    if followup and not resource:
        previous_entities = previous_plan.get("entities") if isinstance(previous_plan, dict) else {}
        previous_resource = previous_entities.get("resource") if isinstance(previous_entities, dict) else None
        if isinstance(previous_resource, dict) and previous_resource.get("resourceKey"):
            resource = dict(previous_resource)
            context_used.append("resource")
    metric = _assistant_metric_from_terms(business_terms) or _assistant_metric_from_text(text)
    explicit_metric = bool(_assistant_metric_from_terms(business_terms)) or _assistant_contains_any(text, _ASSISTANT_METRIC_WORDS)

    dimensions = []
    if station:
        dimensions.append("station")
    if resource:
        dimensions.append("resource")
    dimensions.extend(_assistant_dimensions_from_terms(business_terms))
    if not dimensions and page_hints.get("dimensions") and _assistant_contains_any(text, ("这个", "该", "当前", "本页", "这里", "它", "这个页面")):
        dimensions.extend(page_hints.get("dimensions") or [])
        context_used.extend(page_hints.get("contextUsed") or [])
    dimensions = list(dict.fromkeys(dimensions))
    explicit_dimensions = list(dimensions)

    domain = _assistant_domain_from_terms(business_terms)
    if not _assistant_has_explicit_domain(business_terms) and page_hints.get("domain"):
        domain = str(page_hints.get("domain"))
        context_used.extend(page_hints.get("contextUsed") or [])
    elif (
        page_hints.get("domain") == domain
        and _assistant_contains_any(text, ("这个", "该", "当前", "本页", "这里", "它", "这个页面"))
    ):
        context_used.extend(page_hints.get("contextUsed") or [])
    if followup and not _assistant_has_explicit_domain(business_terms):
        previous_domain = previous_plan.get("domain") if isinstance(previous_plan, dict) else None
        if previous_domain:
            domain = str(previous_domain)
            context_used.append("domain")
    if domain == "accounting":
        if period_label == "最近7天":
            start = today - timedelta(days=1)
            end = today
            period_label = "昨天"

    route = "retail_summary"
    intent = "summary"

    status_words = ("停业", "站点状态", "异常站点", "暂停营业", "维修", "停电", "装修", "修路", "迁址")
    data_recovery_words = ("数据", "传回来", "补回来", "回来", "补传", "上传", "更新", "到了", "有没有")
    trend_words = ("趋势", "走势", "每日", "日销量", "波动", "掉", "下降", "下滑", "跌", "减少", "掉得厉害", "异常")
    comparison_words = _ASSISTANT_COMPARISON_SIGNAL_WORDS
    ranking_words = _ASSISTANT_RANKING_SIGNAL_WORDS

    if domain in {"market", "exchange", "density", "contract"}:
        route = f"{domain}_summary"
        intent = "summary"
    elif domain == "procurement":
        route = "procurement_summary"
        intent = "ranking" if _assistant_contains_any(text, ranking_words) else "summary"
        if not explicit_metric:
            metric = "tons"
        if _assistant_contains_any(text, ("供应商", "供货商", "发货单位")) and "supplier" not in dimensions:
            dimensions.append("supplier")
        elif _assistant_contains_any(text, ("油品", "产品", "品类")) and "resource" not in dimensions:
            dimensions.append("resource")
    elif domain == "weighing":
        route = "weighing_summary"
        intent = "status" if _assistant_contains_any(text, ("待审核", "待复核", "未审核", "状态")) else "summary"
        if not explicit_metric:
            metric = "count" if intent == "status" else "tons"
        if _assistant_contains_any(text, ("车辆", "车号", "车牌")) and "vehicle" not in dimensions:
            dimensions.append("vehicle")
        if _assistant_contains_any(text, ("待审核", "待复核", "状态")) and "status" not in dimensions:
            dimensions.append("status")
    elif domain == "materials":
        if _assistant_contains_any(text, ("调拨单", "调拨", "审批", "待审批")):
            route = "material_transfer_summary"
            intent = "status" if _assistant_contains_any(text, ("状态", "审批", "待审批")) else "summary"
        else:
            route = "materials_summary"
            intent = "inventory" if _assistant_contains_any(text, ("库存", "现有", "剩余")) else "summary"
        if not explicit_metric:
            metric = "amount" if _assistant_contains_any(text, ("金额", "价格", "费用")) else "count"
    elif domain == "transfer_loss":
        route = "transfer_loss_summary"
        intent = "summary"
        if not explicit_metric:
            metric = "loss_rate" if _assistant_contains_any(text, ("损耗率", "率")) else "tons"
    elif domain == "system_health":
        route = "system_health_summary"
        intent = "status"
        if not explicit_metric:
            metric = "status"
    elif domain == "knowledge":
        route = "knowledge_summary"
        intent = "summary"
        if not explicit_metric:
            metric = "count"
    elif domain == "audit":
        route = "audit_log_summary"
        intent = "summary"
        if not explicit_metric:
            metric = "count"
    elif domain == "alias":
        route = "alias_summary"
        intent = "summary"
        if not explicit_metric:
            metric = "count"
    elif domain == "users":
        route = "user_permission_summary"
        intent = "summary"
        if not explicit_metric:
            metric = "count"
    elif domain in {"settings", "management"}:
        route = "settings_summary"
        intent = "summary"
        if not explicit_metric:
            metric = "status"
    elif domain == "accounting":
        route = "accounting_report_status"
        intent = "status"
    elif domain == "combined":
        route = "combined_summary"
        intent = "summary"
    elif domain == "wholesale":
        if _assistant_contains_any(text, ("审核", "复核", "异常", "提示", "未匹配")):
            route = "wholesale_review"
            intent = "review"
        else:
            route = "wholesale_summary"
            intent = "ranking" if _assistant_contains_any(text, ranking_words) else "summary"
    elif domain == "self_use":
        route = "self_use_summary"
        intent = "summary"
        if (_assistant_chart_kind(text) == "line" or _assistant_contains_any(text, trend_words)) and "date" not in dimensions:
            dimensions.append("date")
    elif station and _assistant_contains_any(text, status_words) and _assistant_contains_any(text, data_recovery_words):
        route = "station_data_check"
        intent = "data_availability"
    elif _assistant_contains_any(text, status_words):
        route = "station_status"
        intent = "status"
    elif _assistant_contains_any(text, ("库存", "期末")):
        route = "retail_inventory"
        intent = "inventory"
    elif _assistant_contains_any(text, ("价格", "单价", "售价", "零售价", "均价")):
        route = "retail_price"
        intent = "price"
    elif _assistant_contains_any(text, ("占比", "结构", "油品结构", "油品占比", "构成")):
        route = "retail_resource_share"
        intent = "structure"
    elif _assistant_contains_any(text, trend_words) and ("station" in dimensions or _assistant_contains_any(text, ("哪个站", "哪家站", "站点"))):
        route = "station_trend_anomaly"
        intent = "trend_anomaly"
        if not explicit_metric:
            metric = "volume"
    elif _assistant_contains_any(text, trend_words):
        route = "retail_trend"
        intent = "trend"
        if not explicit_metric:
            metric = "volume"
    elif _assistant_contains_any(text, comparison_words):
        route = "retail_comparison"
        intent = "comparison"
        if period_label == "最近7天":
            period_label = "本月以来"
            start = today.replace(day=1)
            end = today + timedelta(days=1)
    elif _assistant_contains_any(text, ("KPI", "kpi", "关键指标", "年度指标", "年度目标", "目标完成")):
        route = "retail_kpi"
        intent = "kpi"
    elif _assistant_contains_any(text, ranking_words):
        route = "retail_ranking"
        intent = "ranking"

    previous_route = previous_plan.get("route") if isinstance(previous_plan, dict) else None
    if followup and previous_route:
        if route == "retail_summary" and previous_route != "retail_summary":
            route = str(previous_route)
            intent = str(previous_plan.get("intent") or intent)
            context_used.append("route")
        if route == previous_route and previous_plan.get("domain") and domain == "retail":
            domain = str(previous_plan.get("domain"))
            context_used.append("domain")
        previous_dimensions = previous_plan.get("dimensions")
        if route == previous_route and not dimensions and isinstance(previous_dimensions, list):
            dimensions = [str(item) for item in previous_dimensions if item]
            context_used.append("dimensions")
        previous_metric = previous_plan.get("metric")
        if route == previous_route and not explicit_metric and previous_metric in {"amount", "volume", "tons"}:
            metric = previous_metric
            context_used.append("metric")

    if route == "retail_ranking" and not dimensions:
        dimensions.append("station")
    if route == "retail_resource_share" and "resource" not in dimensions:
        dimensions.append("resource")
    if _assistant_route_uses_retail_operation_cutoff(domain, route):
        period_label, start, end = _assistant_cap_period_to_operational_cutoff(period_label, start, end, today)
    comparison = _assistant_comparison_payload(text, period_label, start, end)
    execution_hints = []
    if route in {"retail_comparison", "station_trend_anomaly", "retail_trend"}:
        execution_hints.append("先读快照汇总，再按时间/维度解释变化，不直接拼 SQL")
    if route == "station_trend_anomaly":
        execution_hints.append("下滑原因按数据完整性、站点状态、油品结构、库存价格顺序排查")
    assumptions = []
    if route in {"retail_trend", "station_trend_anomaly"} and not explicit_metric:
        assumptions.append("趋势、下滑、波动类问题未明确指标时默认按销量分析")

    clarification = _assistant_clarification_payload(
        text,
        domain=domain,
        route=route,
        explicit_dimensions=explicit_dimensions,
        explicit_metric=explicit_metric,
        business_terms=business_terms,
        followup=followup,
        context_domain=bool(page_hints.get("domain")),
    )

    confidence = 0.55
    if domain != "retail" or route != "retail_summary":
        confidence += 0.2
    if period_label != "最近7天":
        confidence += 0.1
    if station:
        confidence += 0.1
    if dimensions:
        confidence += 0.05
    if clarification:
        confidence = min(confidence, 0.48)
    confidence = min(confidence, 0.95)

    plan = {
        "normalizedQuestion": text,
        "domain": domain,
        "intent": intent,
        "route": route,
        "metric": metric,
        "dimensions": dimensions,
        "entities": {key: value for key, value in (("station", station), ("resource", resource)) if value},
        "period": _assistant_period_payload(period_label, start, end),
        "tool": _assistant_tool_payload(route),
        "contextUsed": list(dict.fromkeys(context_used)),
        "businessTerms": business_terms,
        "comparison": comparison,
        "visual": {
            "requested": _assistant_chart_requested(text),
            "preferredType": _assistant_chart_kind(text),
        },
        "executionHints": execution_hints,
        "assumptions": assumptions,
        "needsClarification": bool(clarification),
        "clarification": clarification,
        "explicit": {
            "domain": _assistant_has_explicit_domain(business_terms),
            "dimensions": explicit_dimensions,
            "metric": explicit_metric,
            "period": period_label != "最近7天",
        },
        "confidence": round(confidence, 2),
    }
    if allow_compound and not clarification:
        segments = _assistant_compound_segments(text)
        if segments:
            subplans = []
            for segment in segments[:5]:
                enriched = _assistant_enrich_segment(segment, plan)
                subplan = build_assistant_query_plan(enriched, today=today, context=context, allow_compound=False)
                subplan["sourceSegment"] = segment
                subplans.append(subplan)
            if len(subplans) >= 2:
                plan["intent"] = "compound"
                plan["route"] = "compound_query"
                plan["tool"] = _assistant_tool_payload("compound_query")
                plan["subPlans"] = subplans
                plan["executionHints"] = list(dict.fromkeys([*execution_hints, "按子问题顺序调用只读业务工具并合并结论"]))
                plan["confidence"] = max(plan["confidence"], 0.86)
    return _assistant_finalize_plan(plan, text)


def register_assistant_routes(app):
    """Register assistant routes."""

    def _assistant_period(question):
        return parse_assistant_period(question)


    def _assistant_period_display(label, start, end):
        try:
            end_day = end - timedelta(days=1)
            return f"{label}（{start.isoformat()} 至 {end_day.isoformat()}）"
        except Exception:
            return label


    def _retail_density_map():
        oilstation_path = str(OILSTATION_ROOT)
        sys.path.insert(0, oilstation_path)
        try:
            from export_excel import RESOURCE_DENSITY
            return {int(key): float(value) for key, value in RESOURCE_DENSITY.items()}
        except Exception:
            return {102: 0.76, 103: 0.74, 104: 0.75, 106: 0.85, 108: 0.54, 109: 0.85, 110: 0.85}
        finally:
            sys.path = [item for item in sys.path if item != oilstation_path]


    def _retail_snapshot_file(kind, key):
        prefix = {"daily": "day", "monthly": "monthly"}.get(kind, kind)
        if kind == "daily":
            year, month, day = key.split("-")
            return OILSTATION_ROOT / "cache_snapshots" / "daily" / year / month / f"{prefix}_{int(day):02d}.json"
        year, month = key.split("-")
        return OILSTATION_ROOT / "cache_snapshots" / "monthly" / year / f"{prefix}_{int(month):02d}.json"


    def _read_retail_snapshot(kind, key):
        path = _retail_snapshot_file(kind, key)
        if not path.exists():
            return None
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return None


    def _report_extras_snapshot_path(ym):
        year, month = ym.split("-")
        return OILSTATION_ROOT / "cache_snapshots" / "report_extras" / year / f"report_extras_{int(month):02d}.json"


    def _retail_report_index():
        path = OILSTATION_ROOT / "exports" / ".monthly_report_index.json"
        if not path.exists():
            return {}
        try:
            data = json.loads(path.read_text(encoding="utf-8-sig"))
            return data if isinstance(data, dict) else {}
        except (OSError, json.JSONDecodeError):
            return {}


    def _retail_report_workbook_from_index(ym):
        item = _retail_report_index().get(ym) or {}
        filename = item.get("file")
        if not filename:
            return None
        path = OILSTATION_ROOT / "exports" / filename
        return path if path.exists() else None


    def _station_key_from_cell(value):
        try:
            station_key = int(float(value))
        except (TypeError, ValueError):
            return None
        return station_key if station_key in retail_station_catalog_cached() else None


    def _float_cell(value):
        if value is None or value == "":
            return 0.0
        try:
            return float(value)
        except (TypeError, ValueError):
            return 0.0


    def _optional_float(value):
        if value is None or value == "":
            return None
        try:
            return float(value)
        except (TypeError, ValueError):
            return None


    def _retail_mass_is_plausible(volume, mass):
        try:
            volume_value = float(volume or 0)
            mass_value = float(mass or 0)
        except (TypeError, ValueError):
            return False
        if volume_value <= 0:
            return mass_value >= 0
        if mass_value <= 0:
            return False
        density = mass_value / volume_value
        return 0.4 <= density <= 1.1


    def _retail_mass_from_item(item):
        volume = item.get("volume") if isinstance(item, dict) else 0
        tons_value = _optional_float(item.get("tons") if isinstance(item, dict) else None)
        if tons_value is not None:
            mass_value = tons_value * 1000
            if _retail_mass_is_plausible(volume, mass_value):
                return mass_value
        mass_value = _optional_float(item.get("mass") if isinstance(item, dict) else None)
        if mass_value is not None and _retail_mass_is_plausible(volume, mass_value):
            return mass_value
        return None


    def _retail_resource_key_from_name(value):
        name = str(value or "").strip().lower()
        if not name:
            return 0
        if "-35" in name:
            return 110
        if "спбт" in name or "сжиж" in name:
            return 108
        if "95" in name:
            return 104
        if "92" in name:
            return 103
        if "80" in name:
            return 102
        if "д/т" in name or "диз" in name:
            return 106
        return 0


    def _retail_report_column_map(sheet):
        header = next(sheet.iter_rows(min_row=4, max_row=4, values_only=True), ())
        labels = [str(value or "").strip() for value in header]

        def find(*candidates, default=0):
            for candidate in candidates:
                for index, label in enumerate(labels):
                    if candidate in label:
                        return index
            return default

        return {
            "product": find("油品名", default=3),
            "sales_tons": find("销售总吨", "销售量吨", default=20),
            "end_inventory": find("期末量", default=23),
            "end_inventory_tons": find("期末吨", default=24),
        }


    def _retail_resource_tons_from_report(ym):
        workbook_path = _retail_report_workbook_from_index(ym)
        if not workbook_path:
            return {}
        try:
            from openpyxl import load_workbook
        except Exception:
            return {}
        workbook = None
        resource_tons = {}
        try:
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            sheet_name = "总表" if "总表" in workbook.sheetnames else workbook.sheetnames[0]
            sheet = workbook[sheet_name]
            columns = _retail_report_column_map(sheet)
            product_idx = columns["product"]
            tons_idx = columns["sales_tons"]
            for row in sheet.iter_rows(min_row=5, values_only=True):
                product_name = str(row[product_idx] if len(row) > product_idx and row[product_idx] is not None else "").strip()
                if not product_name or "合计" in product_name or "итого" in product_name.lower():
                    continue
                resource_key = _retail_resource_key_from_name(product_name)
                if not resource_key:
                    continue
                resource_tons[resource_key] = resource_tons.get(resource_key, 0.0) + _float_cell(row[tons_idx] if len(row) > tons_idx else None)
        except Exception:
            return {}
        finally:
            if workbook is not None:
                workbook.close()
        return {key: round(value, 3) for key, value in resource_tons.items()}


    def _retail_report_extras_from_workbook(ym):
        workbook_path = _retail_report_workbook_from_index(ym)
        if not workbook_path:
            return None
        try:
            from openpyxl import load_workbook
        except Exception:
            return None
        workbook = None
        try:
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            sheet_name = "总表" if "总表" in workbook.sheetnames else workbook.sheetnames[0]
            sheet = workbook[sheet_name]
            columns = _retail_report_column_map(sheet)
            product_idx = columns["product"]
            inventory_idx = columns["end_inventory"]
            inventory_tons_idx = columns["end_inventory_tons"]
            result = {
                "totalEndInventory": 0.0,
                "totalEndInventoryTons": 0.0,
                "stations": {},
                "source": "daily_report_workbook",
                "sourceFile": workbook_path.name,
                "sourceMtime": workbook_path.stat().st_mtime,
            }
            current_shop_key = None
            for row in sheet.iter_rows(min_row=5, values_only=True):
                station_key = _station_key_from_cell(row[0] if len(row) > 0 else None)
                if station_key is not None:
                    current_shop_key = station_key
                if current_shop_key is None:
                    continue
                product_name = str(row[product_idx] if len(row) > product_idx and row[product_idx] is not None else "").strip()
                if not product_name or "合计" in product_name or "итого" in product_name.lower():
                    continue
                end_inventory = _float_cell(row[inventory_idx] if len(row) > inventory_idx else None)
                end_inventory_tons = _float_cell(row[inventory_tons_idx] if len(row) > inventory_tons_idx else None)
                if not end_inventory and not end_inventory_tons:
                    continue
                item = result["stations"].setdefault(str(current_shop_key), {
                    "endInventory": 0.0,
                    "endInventoryTons": 0.0,
                })
                item["endInventory"] += end_inventory
                item["endInventoryTons"] += end_inventory_tons
                result["totalEndInventory"] += end_inventory
                result["totalEndInventoryTons"] += end_inventory_tons
            result["totalEndInventory"] = round(result["totalEndInventory"], 2)
            result["totalEndInventoryTons"] = round(result["totalEndInventoryTons"], 3)
            for item in result["stations"].values():
                item["endInventory"] = round(item["endInventory"], 2)
                item["endInventoryTons"] = round(item["endInventoryTons"], 3)
            return result
        except Exception:
            return None
        finally:
            if workbook is not None:
                workbook.close()


    def _retail_station_tons_from_report(ym):
        workbook_path = _retail_report_workbook_from_index(ym)
        if not workbook_path:
            return {}
        try:
            from openpyxl import load_workbook
        except Exception:
            return {}
        workbook = None
        station_tons = {}
        try:
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            sheet_name = "总表" if "总表" in workbook.sheetnames else workbook.sheetnames[0]
            sheet = workbook[sheet_name]
            columns = _retail_report_column_map(sheet)
            product_idx = columns["product"]
            tons_idx = columns["sales_tons"]
            current_shop_key = None
            for row in sheet.iter_rows(min_row=5, values_only=True):
                station_key = _station_key_from_cell(row[0] if len(row) > 0 else None)
                if station_key is not None:
                    current_shop_key = station_key
                if current_shop_key is None:
                    continue
                product_name = str(row[product_idx] if len(row) > product_idx and row[product_idx] is not None else "").strip()
                if not product_name or "合计" in product_name or "итого" in product_name.lower():
                    continue
                station_tons[current_shop_key] = station_tons.get(current_shop_key, 0.0) + _float_cell(row[tons_idx] if len(row) > tons_idx else None)
        except Exception:
            return {}
        finally:
            if workbook is not None:
                workbook.close()
        return {key: round(value, 3) for key, value in station_tons.items()}


    def _read_retail_report_extras(ym):
        path = _report_extras_snapshot_path(ym)
        data = None
        if path.exists():
            try:
                loaded = json.loads(path.read_text(encoding="utf-8-sig"))
                data = loaded if isinstance(loaded, dict) else None
            except (OSError, json.JSONDecodeError):
                data = None
        workbook_path = _retail_report_workbook_from_index(ym)
        if data and workbook_path and data.get("source") == "daily_report_workbook":
            stale = data.get("sourceFile") != workbook_path.name
            source_mtime = data.get("sourceMtime")
            try:
                stale = stale or (source_mtime is not None and workbook_path.stat().st_mtime > float(source_mtime) + 0.5)
            except (OSError, TypeError, ValueError):
                stale = True
            if stale:
                data = None
        if data is None:
            data = _retail_report_extras_from_workbook(ym)
        return data


    def _month_after(year, month):
        return (year + 1, 1) if month == 12 else (year, month + 1)


    def _full_month_ym(start, end):
        if start.day != 1:
            return None
        next_year, next_month = _month_after(start.year, start.month)
        if end == date(next_year, next_month, 1):
            return f"{start.year}-{start.month:02d}"
        return None


    def _retail_inventory_for_period(start, end):
        months = []
        cursor = start
        while cursor < end:
            ym = _full_month_ym(cursor, min(end, date(*_month_after(cursor.year, cursor.month), 1)))
            if not ym:
                break
            months.append(ym)
            next_year, next_month = _month_after(cursor.year, cursor.month)
            cursor = date(next_year, next_month, 1)
        if not months:
            ym = _full_month_ym(start, end)
            months = [ym] if ym else []
        result = {
            "totalEndInventory": 0.0,
            "totalEndInventoryTons": 0.0,
            "stations": {},
            "sourceFile": "",
            "month": months[-1] if months else "",
        }
        for ym in months:
            extras = _read_retail_report_extras(ym)
            if not extras:
                continue
            result["totalEndInventory"] = float(extras.get("totalEndInventory") or 0)
            result["totalEndInventoryTons"] = float(extras.get("totalEndInventoryTons") or 0)
            result["stations"] = extras.get("stations") or {}
            result["sourceFile"] = extras.get("sourceFile") or ""
            result["month"] = ym
        if not result["month"]:
            return None
        return result


    def _normalize_retail_snapshot_payload(payload, ym=None):
        if not payload:
            return None
        if ym is None:
            ym = payload.get("month")
        resource_tons_from_report = _retail_resource_tons_from_report(ym) if ym else {}
        station_tons_from_report = _retail_station_tons_from_report(ym) if ym else {}
        summary = payload.get("summary") or {}
        resources = []
        for item in payload.get("byResource") or []:
            resource_key = int(item.get("resourceKey") or item.get("key") or 0)
            volume = float(item.get("volume") or 0)
            amount = float(item.get("amount") or 0)
            mass_value = _retail_mass_from_item(item)
            tons_value = mass_value / 1000 if mass_value is not None else None
            if tons_value is None and resource_key in resource_tons_from_report:
                tons_value = resource_tons_from_report[resource_key]
                mass_value = tons_value * 1000
            resources.append({
                "resourceKey": resource_key,
                "key": resource_key,
                "name": item.get("name") or "--",
                "volume": round(volume, 2),
                "mass": round(float(mass_value if mass_value is not None else (tons_value or 0) * 1000), 3),
                "tons": round(float(tons_value or 0), 3),
                "amount": round(amount, 2),
            })
        resource_total_liters = round(sum(float(item.get("volume") or 0) for item in resources), 2)
        resource_total_tons = round(sum(float(item.get("tons") or 0) for item in resources), 3)
        resource_total_sales = round(sum(float(item.get("amount") or 0) for item in resources), 2)
        stations = []
        for item in payload.get("stations") or []:
            shop_key = item.get("shopKey")
            try:
                shop_key_int = int(shop_key)
            except (TypeError, ValueError):
                shop_key_int = None
            mass_value = _retail_mass_from_item(item)
            tons_value = mass_value / 1000 if mass_value is not None else None
            if tons_value is None and shop_key_int in station_tons_from_report:
                tons_value = station_tons_from_report[shop_key_int]
                mass_value = tons_value * 1000
            stations.append({
                "shopKey": shop_key,
                "name": item.get("name") or "--",
                "short": item.get("short") or "",
                "region": item.get("region") or "--",
                "volume": round(float(item.get("volume") or 0), 2),
                "mass": round(float(mass_value if mass_value is not None else (tons_value or 0) * 1000), 3) if tons_value is not None or mass_value is not None else None,
                "tons": round(float(tons_value or 0), 3) if tons_value is not None else None,
                "amount": round(float(item.get("amount") or 0), 2),
            })
        region_tons_from_stations = {}
        for item in stations:
            if item.get("tons") is not None:
                region = item.get("region") or "--"
                region_tons_from_stations[region] = region_tons_from_stations.get(region, 0.0) + float(item.get("tons") or 0)
        raw_regions = payload.get("regions") or []
        if isinstance(raw_regions, dict):
            region_iter = ({**value, "region": key} for key, value in raw_regions.items())
        else:
            region_iter = raw_regions
        regions = [
            {
                "region": item.get("region") or item.get("name") or "--",
                "volume": round(float(item.get("volume") or 0), 2),
                "tons": round(float(
                    ((_retail_mass_from_item(item) or 0) / 1000)
                    or region_tons_from_stations.get(item.get("region") or item.get("name") or "--")
                    or 0
                ), 3),
                "amount": round(float(item.get("amount") or 0), 2),
            }
            for item in region_iter
        ]
        stations.sort(key=lambda item: item["amount"], reverse=True)
        regions.sort(key=lambda item: item["amount"], reverse=True)
        station_total_tons = round(sum(float(item.get("tons") or 0) for item in stations), 3)
        summary_tons = _optional_float(summary.get("totalTons"))
        total_tons = resource_total_tons or station_total_tons or float(summary_tons or 0)
        return {
            "source": payload.get("source") or "OilStationSnapshot",
            "summary": {
                "totalTxn": int(summary.get("totalTxn") or 0),
                "totalLiters": round(float(summary.get("totalLiters") or resource_total_liters or 0), 2),
                "totalTons": round(total_tons, 3),
                "totalSales": round(float(summary.get("totalSales") or resource_total_sales or 0), 2),
                "totalSalesCny": round(float(summary.get("totalSalesCny") or 0), 2),
            },
            "byResource": resources,
            "stations": stations,
            "regions": regions,
        }


    def _merge_retail_snapshot_payloads(payloads):
        merged = {"summary": {"totalTxn": 0, "totalLiters": 0, "totalTons": 0, "totalSales": 0, "totalSalesCny": 0}, "byResource": [], "stations": [], "regions": []}
        resource_map = {}
        station_map = {}
        region_map = {}
        for payload in payloads:
            normalized = _normalize_retail_snapshot_payload(payload, (payload or {}).get("month"))
            if not normalized:
                continue
            summary = normalized["summary"]
            for key in ("totalTxn", "totalLiters", "totalTons", "totalSales", "totalSalesCny"):
                merged["summary"][key] += float(summary.get(key) or 0)
            for item in normalized.get("byResource") or []:
                bucket = resource_map.setdefault(item.get("resourceKey"), {**item, "volume": 0, "mass": 0, "tons": 0, "amount": 0})
                bucket["volume"] += float(item.get("volume") or 0)
                bucket["mass"] += float(item.get("mass") or 0)
                bucket["tons"] += float(item.get("tons") or 0)
                bucket["amount"] += float(item.get("amount") or 0)
            for item in normalized.get("stations") or []:
                key = item.get("shopKey") or item.get("name")
                bucket = station_map.setdefault(key, {**item, "volume": 0, "mass": 0, "tons": 0, "amount": 0})
                bucket["volume"] += float(item.get("volume") or 0)
                bucket["mass"] += float(item.get("mass") or 0)
                bucket["tons"] += float(item.get("tons") or 0)
                bucket["amount"] += float(item.get("amount") or 0)
            for item in normalized.get("regions") or []:
                key = item.get("region")
                bucket = region_map.setdefault(key, {**item, "volume": 0, "mass": 0, "tons": 0, "amount": 0})
                bucket["volume"] += float(item.get("volume") or 0)
                bucket["mass"] += float(item.get("mass") or 0)
                bucket["tons"] += float(item.get("tons") or 0)
                bucket["amount"] += float(item.get("amount") or 0)
        merged["byResource"] = sorted(
            ({**item, "volume": round(item["volume"], 2), "mass": round(item["mass"], 3), "tons": round(item["tons"], 3), "amount": round(item["amount"], 2)} for item in resource_map.values()),
            key=lambda item: item["volume"],
            reverse=True,
        )
        merged["summary"] = {key: round(value, 3 if key == "totalTons" else 2) for key, value in merged["summary"].items()}
        if merged["byResource"]:
            resource_tons = round(sum(float(item.get("tons") or 0) for item in merged["byResource"]), 3)
            summary_tons = float(merged["summary"].get("totalTons") or 0)
            if resource_tons and (not summary_tons or abs(summary_tons - resource_tons) > max(1.0, resource_tons * 0.05)):
                merged["summary"]["totalTons"] = resource_tons
        merged["stations"] = sorted(
            ({**item, "volume": round(item["volume"], 2), "mass": round(item["mass"], 3), "tons": round(item["tons"], 3), "amount": round(item["amount"], 2)} for item in station_map.values()),
            key=lambda item: item["amount"],
            reverse=True,
        )
        merged["regions"] = sorted(
            ({**item, "volume": round(item["volume"], 2), "mass": round(item["mass"], 3), "tons": round(item["tons"], 3), "amount": round(item["amount"], 2)} for item in region_map.values()),
            key=lambda item: item["amount"],
            reverse=True,
        )
        return merged


    def _empty_retail_snapshot_payload(start, end, source="OilStationSnapshot"):
        return {
            "source": source,
            "summary": {
                "totalTxn": 0,
                "totalLiters": 0.0,
                "totalTons": 0.0,
                "totalSales": 0.0,
                "totalSalesCny": 0.0,
                "queryStart": start.isoformat() if hasattr(start, "isoformat") else "",
                "queryEnd": end.isoformat() if hasattr(end, "isoformat") else "",
            },
            "byResource": [],
            "stations": [],
            "regions": [],
        }


    def _retail_snapshot_for_period(start, end):
        today = date.today()
        if end <= start:
            return _empty_retail_snapshot_payload(start, end)

        if start.day == 1 and start.year == today.year and start.month == today.month and end >= today:
            ym = f"{start.year}-{start.month:02d}"
            current_month = _normalize_retail_snapshot_payload(_read_retail_snapshot("monthly", ym), ym)
            if current_month:
                return current_month
        if start.day == 1 and end == (date(start.year + 1, 1, 1) if start.month == 12 else date(start.year, start.month + 1, 1)):
            ym = f"{start.year}-{start.month:02d}"
            return _normalize_retail_snapshot_payload(_read_retail_snapshot("monthly", ym), ym)
        if start.month == 1 and start.day == 1 and end == date(start.year + 1, 1, 1):
            payloads = [
                _read_retail_snapshot("monthly", f"{start.year}-{month:02d}")
                for month in range(1, 13)
            ]
            payloads = [item for item in payloads if item]
            return _merge_retail_snapshot_payloads(payloads) if payloads else None

        if start.day == 1:
            payloads = []
            cursor = start
            while cursor < end:
                next_year, next_month = _month_after(cursor.year, cursor.month)
                next_month_start = date(next_year, next_month, 1)
                if next_month_start > end:
                    payloads = []
                    break
                ym = f"{cursor.year}-{cursor.month:02d}"
                payload = _read_retail_snapshot("monthly", ym)
                if not payload:
                    payloads = []
                    break
                payloads.append(payload)
                cursor = next_month_start
            if payloads and cursor == end:
                return _merge_retail_snapshot_payloads(payloads)

        # Daily snapshots are generated only after a business day is complete.
        # If the requested range includes today, use completed days only and avoid
        # falling back to the live retail DB for incomplete shifts.
        effective_end = min(end, today)
        if effective_end <= start:
            return _empty_retail_snapshot_payload(start, effective_end)

        dates = []
        current = start
        while current < effective_end:
            payload = _read_retail_snapshot("daily", current.isoformat())
            if not payload:
                return None
            dates.append(payload)
            current += timedelta(days=1)
        return _merge_retail_snapshot_payloads(dates)


    def _combined_data_from_snapshots(cur, year, month, start, end):
        retail_snapshot = _retail_snapshot_for_period(start, end)
        if not retail_snapshot:
            return summarize_combined(cur, year, month)
        wholesale = summarize_wholesale(cur, year, month)
        retail_summary = retail_snapshot.get("summary") or {}
        retail = {
            "liters": round(float(retail_summary.get("totalLiters") or 0), 2),
            "tons": round(float(retail_summary.get("totalTons") or 0), 3),
            "amountKgs": round(float(retail_summary.get("totalSales") or 0), 2),
            "amountCny": round(float(retail_summary.get("totalSalesCny") or 0), 2),
            "source": "snapshot",
        }
        combined = {
            "salesTons": round(retail["tons"] + wholesale["salesTons"], 3),
            "salesAmountCny": round(retail["amountCny"] + wholesale["salesAmountCny"], 2),
            "transferTons": wholesale["transferTons"],
        }
        return {"retail": retail, "wholesale": wholesale, "combined": combined}


    def _format_assistant_retail_answer(label, payload):
        summary = payload.get("summary") or {}
        question = payload.get("question") or ""
        rows = [
            item for item in payload.get("byResource") or []
            if float(item.get("volume") or 0) or float(item.get("amount") or 0)
        ]
        lines = [
            f"{label}零售销量统计如下：",
            f"合计销量 {float(summary.get('totalLiters') or 0):,.2f} 升，折算 {float(summary.get('totalTons') or 0):,.3f} 吨，合计金额 {float(summary.get('totalSales') or 0):,.2f} 索姆。",
        ]
        if rows:
            lines.append("各油品明细：")
            for item in rows:
                lines.append(
                    f"{item.get('name') or '--'}：{float(item.get('volume') or 0):,.2f} 升，"
                    f"{float(item.get('tons') or 0):,.3f} 吨，"
                    f"{float(item.get('amount') or 0):,.2f} 索姆"
                )
        else:
            lines.append("该区间暂未查询到已结束班组的销售数据。")
        if any(word in question for word in ("站点", "加油站", "排行", "排名", "最好")) and payload.get("stations"):
            lines.append("站点排行：")
            for item in payload.get("stations", [])[:10]:
                lines.append(
                    f"{item.get('name') or '--'}：{float(item.get('volume') or 0):,.2f} 升，"
                    f"{float(item.get('tons') or 0):,.3f} 吨，{float(item.get('amount') or 0):,.2f} 索姆"
                )
        if any(word in question for word in ("区域", "州", "地区")) and payload.get("regions"):
            lines.append("区域排行：")
            for item in payload.get("regions", [])[:8]:
                lines.append(
                    f"{item.get('region') or '--'}：{float(item.get('volume') or 0):,.2f} 升，"
                    f"{float(item.get('tons') or 0):,.3f} 吨，{float(item.get('amount') or 0):,.2f} 索姆"
                )
        return "\n".join(lines)


    def _assistant_has_any(question, words):
        return any(word in question for word in words)


    def _append_rank_lines(lines, title, rows, limit=5, amount_label="元"):
        clean_rows = [item for item in rows or [] if float(item.get("tons") or item.get("volume") or 0)]
        if not clean_rows:
            return
        lines.append(title)
        for item in clean_rows[:limit]:
            name = item.get("name") or item.get("region") or item.get("product") or "--"
            if "volume" in item:
                lines.append(f"{name}：{float(item.get('volume') or 0):,.2f} 升，{float(item.get('amount') or 0):,.2f} 索姆")
            elif "amountCny" in item:
                lines.append(f"{name}：{float(item.get('tons') or 0):,.3f} 吨，{float(item.get('amountCny') or 0):,.2f} {amount_label}")
            else:
                lines.append(f"{name}：{float(item.get('tons') or 0):,.3f} 吨")


    def _assistant_table(title, columns, rows):
        return {
            "title": title,
            "columns": columns,
            "rows": [
                ["" if value is None else str(value) for value in row]
                for row in rows
            ],
        }


    def _fmt_amount(value, unit=""):
        suffix = f" {unit}" if unit else ""
        return f"{float(value or 0):,.2f}{suffix}"


    def _fmt_volume(value):
        return f"{float(value or 0):,.2f} 升"


    def _fmt_tons(value):
        return f"{float(value or 0):,.3f} 吨"


    def _chart_number(value, digits=2):
        try:
            return round(float(value or 0), digits)
        except (TypeError, ValueError):
            return 0.0


    def _assistant_chart(title, chart_type, *, x_axis=None, series=None, data=None, unit="", summary="", source="assistant"):
        x_axis = [str(item) for item in (x_axis or [])]
        clean_series = []
        for item in series or []:
            if not isinstance(item, dict):
                continue
            values = [_chart_number(value, 3 if item.get("unit") == "吨" else 2) for value in item.get("data") or []]
            if values:
                clean_series.append({
                    "name": str(item.get("name") or "数值"),
                    "type": str(item.get("type") or chart_type or "bar"),
                    "unit": str(item.get("unit") or unit or ""),
                    "data": values,
                })
        clean_data = []
        for item in data or []:
            if not isinstance(item, dict):
                continue
            name = str(item.get("name") or item.get("label") or "--")
            value = _chart_number(item.get("value"), 3 if unit == "吨" else 2)
            if value:
                clean_data.append({"name": name, "value": value})
        if chart_type == "pie" and not clean_data:
            return None
        if chart_type != "pie" and (not x_axis or not clean_series):
            return None
        return {
            "title": title,
            "type": chart_type,
            "xAxis": x_axis,
            "series": clean_series,
            "data": clean_data,
            "unit": unit,
            "summary": summary,
            "source": source,
        }


    def _assistant_metric_unit(metric):
        return {"volume": "升", "tons": "吨", "amount": "索姆"}.get(metric, "索姆")


    def _assistant_metric_name(metric):
        return {"volume": "销量", "tons": "吨位", "amount": "销售金额"}.get(metric, "销售金额")


    def _assistant_axis_labels(rows, name_key="name", limit=12):
        labels = []
        for item in rows[:limit]:
            labels.append(str(item.get(name_key) or item.get("name") or item.get("region") or item.get("date") or "--"))
        return labels


    def _assistant_rank_chart(title, rows, metric, *, name_key="name", limit=12):
        clean_rows = [
            item for item in rows or []
            if _chart_number(item.get(metric if metric != "amount" else "amount")) or _chart_number(item.get("volume")) or _chart_number(item.get("tons"))
        ][:limit]
        if not clean_rows:
            return None
        value_key = "amount" if metric == "amount" else metric
        return _assistant_chart(
            title,
            "bar",
            x_axis=_assistant_axis_labels(clean_rows, name_key=name_key, limit=limit),
            series=[{
                "name": _assistant_metric_name(metric),
                "type": "bar",
                "unit": _assistant_metric_unit(metric),
                "data": [_chart_number(item.get(value_key), 3 if metric == "tons" else 2) for item in clean_rows],
            }],
            unit=_assistant_metric_unit(metric),
            summary=f"按{_assistant_metric_name(metric)}降序展示前 {len(clean_rows)} 项。",
        )


    def _assistant_daily_mixed_chart(title, daily, *, volume_label="销量", amount_label="金额"):
        rows = [item for item in daily or [] if item.get("date")]
        if not rows:
            return None
        return _assistant_chart(
            title,
            "mixed",
            x_axis=[str(item.get("date") or "") for item in rows],
            series=[
                {"name": volume_label, "type": "bar", "unit": "升", "data": [_chart_number(item.get("volume")) for item in rows]},
                {"name": amount_label, "type": "line", "unit": "索姆", "data": [_chart_number(item.get("amount")) for item in rows]},
            ],
            summary="柱状为销量，折线为金额。",
        )


    def _assistant_resource_pie_chart(title, rows, *, value_key="volume", unit="升"):
        clean_rows = [
            {"name": item.get("name") or item.get("product") or "--", "value": item.get(value_key)}
            for item in rows or []
            if _chart_number(item.get(value_key))
        ]
        return _assistant_chart(title, "pie", data=clean_rows, unit=unit, summary="按油品占比展示。")


    def _assistant_retail_tables(label, payload, question):
        tables = []
        summary = payload.get("summary") or {}
        resources = [item for item in payload.get("byResource") or [] if float(item.get("volume") or 0) or float(item.get("amount") or 0)]
        if resources:
            resource_rows = [
                [item.get("name") or "--", _fmt_volume(item.get("volume")), _fmt_tons(item.get("tons")), _fmt_amount(item.get("amount"), "索姆")]
                for item in resources
            ]
            resource_rows.append([
                "合计",
                _fmt_volume(summary.get("totalLiters")),
                _fmt_tons(summary.get("totalTons")),
                _fmt_amount(summary.get("totalSales"), "索姆"),
            ])
            tables.append(_assistant_table(
                f"{label}零售油品明细",
                ["油品", "销量", "吨位", "销售金额"],
                resource_rows,
            ))
        if _assistant_has_any(question, ("站点", "加油站", "排行", "排名", "最好")) and payload.get("stations"):
            tables.append(_assistant_table(
                f"{label}站点排行",
                ["站点", "区域", "销量", "吨位", "销售金额"],
                [
                    [item.get("name") or "--", item.get("region") or "--", _fmt_volume(item.get("volume")), _fmt_tons(item.get("tons")), _fmt_amount(item.get("amount"), "索姆")]
                    for item in payload.get("stations", [])[:12]
                ],
            ))
        if _assistant_has_any(question, ("区域", "州", "地区")) and payload.get("regions"):
            tables.append(_assistant_table(
                f"{label}区域排行",
                ["区域", "销量", "吨位", "销售金额"],
                [
                    [item.get("region") or "--", _fmt_volume(item.get("volume")), _fmt_tons(item.get("tons")), _fmt_amount(item.get("amount"), "索姆")]
                    for item in payload.get("regions", [])[:10]
                ],
            ))
        return tables


    def _assistant_retail_charts(label, payload, question, *, visual=None):
        charts = []
        resources = [item for item in payload.get("byResource") or [] if float(item.get("volume") or 0) or float(item.get("amount") or 0)]
        stations = [item for item in payload.get("stations") or [] if float(item.get("volume") or 0) or float(item.get("amount") or 0)]
        regions = [item for item in payload.get("regions") or [] if float(item.get("volume") or 0) or float(item.get("amount") or 0)]
        preferred = (visual or {}).get("preferredType")

        if resources:
            if preferred == "pie" or _assistant_has_any(question, ("占比", "结构", "构成")):
                chart = _assistant_resource_pie_chart(f"{label}零售油品销量占比", resources)
            else:
                chart = _assistant_chart(
                    f"{label}零售油品销量与金额",
                    "mixed",
                    x_axis=[item.get("name") or "--" for item in resources],
                    series=[
                        {"name": "销量", "type": "bar", "unit": "升", "data": [item.get("volume") for item in resources]},
                        {"name": "金额", "type": "line", "unit": "索姆", "data": [item.get("amount") for item in resources]},
                    ],
                    summary="按油品展示销量和金额。",
                )
            if chart:
                charts.append(chart)
        if stations and (_assistant_has_any(question, ("站点", "加油站", "排行", "排名", "最高", "最多", "最好")) or (visual or {}).get("requested")):
            chart = _assistant_rank_chart(f"{label}零售站点销量排行", stations, "volume", limit=12)
            if chart:
                charts.append(chart)
        if regions and _assistant_has_any(question, ("区域", "州", "地区")):
            chart = _assistant_rank_chart(f"{label}零售区域销量排行", regions, "volume", name_key="region", limit=10)
            if chart:
                charts.append(chart)
        return charts[:3]


    def _assistant_self_use_product_name(resource_key, fallback=""):
        if fallback:
            return str(fallback)
        names = {
            102: "АИ 80",
            103: "АИ 92",
            104: "АИ 95",
            106: "Д/Т",
            108: "СПБТ",
            109: "Д/Т",
            110: "ДТ -35",
        }
        try:
            key = int(resource_key or 0)
        except (TypeError, ValueError):
            key = 0
        return names.get(key, f"Other-{key}" if key else "--")


    def _assistant_float(value):
        try:
            return float(value or 0)
        except (TypeError, ValueError):
            return 0.0


    def _assistant_int(value):
        try:
            return int(value or 0)
        except (TypeError, ValueError):
            return 0


    def _assistant_empty_self_use_data(start, end, source_error=""):
        return {
            "source": "OilStationSelfUse",
            "sourceVersion": "report24-only-v1",
            "date_from": start.isoformat(),
            "date_to": (end - timedelta(days=1)).isoformat() if end > start else start.isoformat(),
            "records": [],
            "summary": {
                "totalVolume": 0.0,
                "totalWeightKg": 0.0,
                "totalTons": 0.0,
                "totalAmount": 0.0,
                "totalRecords": 0,
                "activeStations": 0,
            },
            "byStation": [],
            "byResource": [],
            "byRegion": [],
            "daily": [],
            "sourceError": source_error,
        }


    def _assistant_self_use_data_for_period(start, end, resource=None):
        if end <= start:
            return _assistant_empty_self_use_data(start, end)

        query_end = end - timedelta(days=1)
        try:
            response = open_oilstation(
                "/api/retail/self-use",
                query_string={"date_from": start.isoformat(), "date_to": query_end.isoformat()},
            )
            raw_data = response.get_json() if response.status_code < 400 else {}
        except Exception as exc:
            return _assistant_empty_self_use_data(start, end, str(exc))

        records = list((raw_data or {}).get("records") or [])
        if isinstance(resource, dict) and resource.get("resourceKey"):
            resource_key = _assistant_int(resource.get("resourceKey"))
            records = [record for record in records if _assistant_int(record.get("resourceKey")) == resource_key]
        catalog = retail_station_catalog_cached()
        stations = {}
        resources = {}
        regions = {}
        daily = {}
        total_volume = 0.0
        total_mass = 0.0
        total_amount = 0.0

        for record in records:
            volume = _assistant_float(record.get("volume"))
            mass = _assistant_float(record.get("mass"))
            amount = _assistant_float(record.get("amount"))
            total_volume += volume
            total_mass += mass
            total_amount += amount

            shop_key = _assistant_int(record.get("shopKey"))
            station_info = catalog.get(shop_key, {}) if shop_key else {}
            station_name = record.get("stationName") or station_info.get("name") or (f"{shop_key}号站" if shop_key else "--")
            region_name = station_info.get("region") or record.get("region") or "--"
            station_bucket = stations.setdefault(shop_key or station_name, {
                "shopKey": shop_key,
                "name": station_name,
                "region": region_name,
                "volume": 0.0,
                "mass": 0.0,
                "amount": 0.0,
                "records": 0,
            })
            station_bucket["volume"] += volume
            station_bucket["mass"] += mass
            station_bucket["amount"] += amount
            station_bucket["records"] += 1

            resource_key = _assistant_int(record.get("resourceKey"))
            product_name = _assistant_self_use_product_name(resource_key, record.get("productName") or "")
            resource_bucket = resources.setdefault(resource_key or product_name, {
                "resourceKey": resource_key,
                "name": product_name,
                "volume": 0.0,
                "mass": 0.0,
                "amount": 0.0,
                "records": 0,
            })
            resource_bucket["volume"] += volume
            resource_bucket["mass"] += mass
            resource_bucket["amount"] += amount
            resource_bucket["records"] += 1

            region_bucket = regions.setdefault(region_name, {
                "region": region_name,
                "volume": 0.0,
                "mass": 0.0,
                "amount": 0.0,
                "records": 0,
            })
            region_bucket["volume"] += volume
            region_bucket["mass"] += mass
            region_bucket["amount"] += amount
            region_bucket["records"] += 1

            date_key = str(record.get("date") or "")
            if date_key:
                day_bucket = daily.setdefault(date_key, {
                    "date": date_key,
                    "volume": 0.0,
                    "mass": 0.0,
                    "amount": 0.0,
                    "records": 0,
                })
                day_bucket["volume"] += volume
                day_bucket["mass"] += mass
                day_bucket["amount"] += amount
                day_bucket["records"] += 1

        def rounded_items(values, *, name_key):
            result = []
            for item in values:
                copied = dict(item)
                copied["volume"] = round(float(copied.get("volume") or 0), 2)
                copied["mass"] = round(float(copied.get("mass") or 0), 2)
                copied["tons"] = round(copied["mass"] / 1000, 3)
                copied["amount"] = round(float(copied.get("amount") or 0), 2)
                copied["records"] = int(copied.get("records") or 0)
                result.append(copied)
            result.sort(key=lambda item: (float(item.get("volume") or 0), float(item.get("amount") or 0), str(item.get(name_key) or "")), reverse=True)
            return result

        return {
            "source": "OilStationSelfUse",
            "sourceVersion": (raw_data or {}).get("sourceVersion") or "report24-only-v1",
            "date_from": start.isoformat(),
            "date_to": query_end.isoformat(),
            "records": records,
            "summary": {
                "totalVolume": round(total_volume, 2),
                "totalWeightKg": round(total_mass, 2),
                "totalTons": round(total_mass / 1000, 3),
                "totalAmount": round(total_amount, 2),
                "totalRecords": len(records),
                "activeStations": len([item for item in stations.values() if item.get("volume") or item.get("mass") or item.get("amount")]),
            },
            "byStation": rounded_items(stations.values(), name_key="name"),
            "byResource": rounded_items(resources.values(), name_key="name"),
            "byRegion": rounded_items(regions.values(), name_key="region"),
            "daily": sorted(rounded_items(daily.values(), name_key="date"), key=lambda item: item.get("date") or ""),
            "sourceError": "",
        }


    def _assistant_self_use_answer(label, data):
        summary = data.get("summary") or {}
        total_weight_tons = _assistant_float(summary.get("totalTons")) or _assistant_float(summary.get("totalWeightKg")) / 1000
        lines = [
            f"{label}自用油数据如下：",
            (
                f"合计自用油 {_fmt_volume(summary.get('totalVolume'))}，"
                f"折算 {_fmt_tons(total_weight_tons)}，"
                f"金额 {_fmt_amount(summary.get('totalAmount'), '索姆')}，"
                f"记录 {int(summary.get('totalRecords') or 0)} 笔，"
                f"覆盖 {int(summary.get('activeStations') or 0)} 个站点。"
            ),
            "口径：原系统报表24自用油，条件为 BasePaymentTypeKey=2 AND DerivedPaymentTypeKey=2；销量、金额、重量分别使用 selling.Volume、selling.ShopCost、selling.Mass。",
        ]
        if not data.get("records"):
            lines.append("该区间暂未查询到自用油记录。")
        elif data.get("byResource"):
            lines.append("油品明细：")
            for item in data.get("byResource", [])[:8]:
                lines.append(
                    f"{item.get('name') or '--'}：{_fmt_volume(item.get('volume'))}，"
                    f"{_fmt_tons(_assistant_float(item.get('tons')) or _assistant_float(item.get('mass')) / 1000)}，"
                    f"{_fmt_amount(item.get('amount'), '索姆')}"
                )
        if data.get("sourceError"):
            lines.append(f"数据读取提示：{data.get('sourceError')}")
        return "\n".join(lines)


    def _assistant_self_use_tables(label, data):
        summary = data.get("summary") or {}
        resource_rows = [
            [
                item.get("name") or "--",
                _fmt_volume(item.get("volume")),
                _fmt_tons(_assistant_float(item.get("tons")) or _assistant_float(item.get("mass")) / 1000),
                _fmt_amount(item.get("amount"), "索姆"),
                item.get("records") or 0,
            ]
            for item in data.get("byResource", [])
        ]
        resource_rows.append([
            "合计",
            _fmt_volume(summary.get("totalVolume")),
            _fmt_tons(_assistant_float(summary.get("totalTons")) or _assistant_float(summary.get("totalWeightKg")) / 1000),
            _fmt_amount(summary.get("totalAmount"), "索姆"),
            summary.get("totalRecords") or 0,
        ])
        tables = [_assistant_table(
            f"{label}自用油油品汇总",
            ["油品", "自用油量", "重量", "金额", "记录数"],
            resource_rows,
        )]
        station_rows = [
            [
                item.get("name") or "--",
                item.get("region") or "--",
                _fmt_volume(item.get("volume")),
                _fmt_tons(_assistant_float(item.get("tons")) or _assistant_float(item.get("mass")) / 1000),
                _fmt_amount(item.get("amount"), "索姆"),
                item.get("records") or 0,
            ]
            for item in data.get("byStation", [])[:12]
        ]
        if station_rows:
            station_rows.append([
                "合计",
                "全部区域",
                _fmt_volume(summary.get("totalVolume")),
                _fmt_tons(_assistant_float(summary.get("totalTons")) or _assistant_float(summary.get("totalWeightKg")) / 1000),
                _fmt_amount(summary.get("totalAmount"), "索姆"),
                summary.get("totalRecords") or 0,
            ])
            tables.append(_assistant_table(
                f"{label}自用油站点汇总",
                ["站点", "区域", "自用油量", "重量", "金额", "记录数"],
                station_rows,
            ))
        return tables


    def _assistant_self_use_charts(label, data, question, *, visual=None):
        charts = []
        daily = data.get("daily") or []
        resources = data.get("byResource") or []
        stations = data.get("byStation") or []
        preferred = (visual or {}).get("preferredType")
        if daily and (preferred in {"line", "auto"} or _assistant_has_any(question, ("趋势", "走势", "每日", "按天", "日趋势", "图"))):
            chart = _assistant_daily_mixed_chart(f"{label}自用油每日趋势", daily, volume_label="自用油量", amount_label="金额")
            if chart:
                charts.append(chart)
        if resources:
            chart = (
                _assistant_resource_pie_chart(f"{label}自用油油品占比", resources)
                if preferred == "pie" or _assistant_has_any(question, ("占比", "结构"))
                else _assistant_rank_chart(f"{label}自用油油品排行", resources, "volume", limit=10)
            )
            if chart:
                charts.append(chart)
        if stations and (_assistant_has_any(question, ("站点", "加油站", "排行", "排名", "哪个")) or (visual or {}).get("requested")):
            chart = _assistant_rank_chart(f"{label}自用油站点排行", stations, "volume", limit=12)
            if chart:
                charts.append(chart)
        return charts[:3]


    def _assistant_retail_inventory_answer(label, data):
        return (
            f"{label}零售期末库存如下：\n"
            f"合计库存 {float(data.get('totalEndInventory') or 0):,.2f} 升，"
            f"折算 {float(data.get('totalEndInventoryTons') or 0):,.3f} 吨。\n"
            f"数据来源：{data.get('sourceFile') or '日销表'}。"
        )


    def _assistant_retail_inventory_tables(label, data):
        catalog = retail_station_catalog_cached()
        station_rows = []
        for raw_shop_key, item in (data.get("stations") or {}).items():
            try:
                shop_key = int(raw_shop_key)
            except (TypeError, ValueError):
                continue
            info = catalog.get(shop_key, {})
            inventory = float(item.get("endInventory") or 0)
            tons = float(item.get("endInventoryTons") or 0)
            if not inventory and not tons:
                continue
            station_rows.append([
                info.get("name") or f"Station-{shop_key}",
                info.get("region") or "--",
                _fmt_volume(inventory),
                _fmt_tons(tons),
            ])
        station_rows.sort(key=lambda row: float(str(row[2]).replace(",", "").split()[0] or 0), reverse=True)
        station_rows.append([
            "合计",
            "全部区域",
            _fmt_volume(data.get("totalEndInventory")),
            _fmt_tons(data.get("totalEndInventoryTons")),
        ])
        return [_assistant_table(
            f"{label}零售期末库存",
            ["加油站", "区域", "期末库存", "折算吨位"],
            station_rows,
        )]


    def _assistant_retail_inventory_charts(label, data):
        catalog = retail_station_catalog_cached()
        rows = []
        for raw_shop_key, item in (data.get("stations") or {}).items():
            try:
                shop_key = int(raw_shop_key)
            except (TypeError, ValueError):
                continue
            info = catalog.get(shop_key, {})
            inventory = float(item.get("endInventory") or 0)
            tons = float(item.get("endInventoryTons") or 0)
            if inventory or tons:
                rows.append({
                    "name": info.get("name") or f"Station-{shop_key}",
                    "volume": inventory,
                    "tons": tons,
                })
        rows.sort(key=lambda item: item["volume"], reverse=True)
        chart = _assistant_rank_chart(f"{label}零售期末库存排行", rows, "volume", limit=15)
        return [chart] if chart else []


    def _assistant_wholesale_tables(label, data, question):
        cards = data.get("cards") or {}
        labels = {"export_wholesale": "出口批发", "local_wholesale": "本地批发", "unknown": "待识别", "transfer": "外运", "processing": "代加工", "procurement": "采购"}
        sales_keys = ("export_wholesale", "local_wholesale", "unknown")
        sales_total_tons = sum(float((cards.get(key) or {}).get("tons") or 0) for key in sales_keys)
        sales_total_amount = sum(float((cards.get(key) or {}).get("amountCny") or 0) for key in ("export_wholesale", "local_wholesale"))
        sales_total_review = sum(int((cards.get(key) or {}).get("reviewRows") or 0) for key in sales_keys)
        category_rows = [
            [
                labels.get(key, key),
                _fmt_tons((cards.get(key) or {}).get("tons")),
                "--" if key in ("unknown", "transfer") else _fmt_amount((cards.get(key) or {}).get("amountCny"), "元"),
                (cards.get(key) or {}).get("entities") or 0,
                (cards.get(key) or {}).get("reviewRows") or 0,
            ]
            for key in ("export_wholesale", "local_wholesale", "unknown", "transfer")
        ]
        category_rows.append(["销售合计", _fmt_tons(sales_total_tons), _fmt_amount(sales_total_amount, "元"), "--", sales_total_review])
        tables = [_assistant_table(
            f"{label}批发分类统计",
            ["分类", "吨数", "金额", "客户/去向数", "待复核"],
            category_rows,
        )]
        rankings = data.get("customerRankings") or {}
        if _assistant_has_any(question, ("客户", "排行", "排名", "top", "TOP")):
            for key in ("export_wholesale", "local_wholesale", "unknown"):
                rows = rankings.get(key) or []
                if rows:
                    tables.append(_assistant_table(
                        f"{labels[key]}客户排行",
                        ["客户", "吨数", "金额", "记录数"],
                        [[item.get("name") or "--", _fmt_tons(item.get("tons")), _fmt_amount(item.get("amountCny"), "元"), item.get("rows") or 0] for item in rows[:10]],
                    ))
        if _assistant_has_any(question, ("油品", "产品", "结构", "品类")):
            products = data.get("productRankings") or {}
            for key in ("export_wholesale", "local_wholesale", "unknown", "transfer"):
                rows = products.get(key) or []
                if rows:
                    tables.append(_assistant_table(
                        f"{labels[key]}油品排行",
                        ["油品", "吨数", "金额"],
                        [[item.get("name") or "--", _fmt_tons(item.get("tons")), "--" if key in ("unknown", "transfer") else _fmt_amount(item.get("amountCny"), "元")] for item in rows[:10]],
                    ))
        if "外运" in question:
            transfer_summary = data.get("transferSummary") or {}
            if transfer_summary.get("byRegion"):
                tables.append(_assistant_table(
                    f"{label}外运区域汇总",
                    ["区域", "吨数", "记录数"],
                    [[item.get("region") or "--", _fmt_tons(item.get("tons")), item.get("rows") or 0] for item in transfer_summary.get("byRegion", [])[:10]],
                ))
        return tables


    def _assistant_wholesale_charts(label, data, question):
        cards = data.get("cards") or {}
        labels = {"export_wholesale": "出口批发", "local_wholesale": "本地批发", "unknown": "待识别", "transfer": "外运", "processing": "代加工", "procurement": "采购"}
        category_rows = [
            {"name": labels.get(key, key), "tons": float((cards.get(key) or {}).get("tons") or 0)}
            for key in ("export_wholesale", "local_wholesale", "unknown", "transfer")
        ]
        charts = []
        chart = _assistant_rank_chart(f"{label}批发分类吨数", category_rows, "tons", limit=6)
        if chart:
            charts.append(chart)
        rankings = data.get("customerRankings") or {}
        if _assistant_has_any(question, ("客户", "排行", "排名", "top", "TOP")):
            customer_rows = []
            for key in ("export_wholesale", "local_wholesale"):
                for item in (rankings.get(key) or [])[:8]:
                    customer_rows.append({
                        "name": f"{labels.get(key, key)}-{item.get('name') or '--'}",
                        "tons": item.get("tons"),
                    })
            chart = _assistant_rank_chart(f"{label}批发客户吨数排行", sorted(customer_rows, key=lambda item: float(item.get("tons") or 0), reverse=True), "tons", limit=12)
            if chart:
                charts.append(chart)
        products = data.get("productRankings") or {}
        if _assistant_has_any(question, ("油品", "产品", "结构", "品类")):
            product_rows = []
            for key in ("export_wholesale", "local_wholesale", "transfer"):
                for item in (products.get(key) or [])[:8]:
                    product_rows.append({
                        "name": f"{labels.get(key, key)}-{item.get('name') or item.get('product') or '--'}",
                        "tons": item.get("tons"),
                    })
            chart = _assistant_rank_chart(f"{label}批发油品吨数排行", sorted(product_rows, key=lambda item: float(item.get("tons") or 0), reverse=True), "tons", limit=12)
            if chart:
                charts.append(chart)
        return charts[:3]


    def _assistant_combined_tables(label, data):
        retail = data.get("retail") or {}
        wholesale = data.get("wholesale") or {}
        combined = data.get("combined") or {}
        return [_assistant_table(
            f"{label}合并统计",
            ["板块", "销量/销售吨数", "销售金额"],
            [
                ["合计", _fmt_tons(combined.get("salesTons")), _fmt_amount(combined.get("salesAmountCny"), "元")],
                ["零售", _fmt_tons(retail.get("tons")), _fmt_amount(retail.get("amountCny"), "元")],
                ["批发销售", _fmt_tons(wholesale.get("salesTons")), _fmt_amount(wholesale.get("salesAmountCny"), "元")],
                ["外运", _fmt_tons(combined.get("transferTons")), "--"],
            ],
        )]


    def _assistant_combined_charts(label, data):
        retail = data.get("retail") or {}
        wholesale = data.get("wholesale") or {}
        combined = data.get("combined") or {}
        chart = _assistant_chart(
            f"{label}合并经营结构",
            "bar",
            x_axis=["合计", "零售", "批发销售", "外运"],
            series=[{
                "name": "吨数",
                "type": "bar",
                "unit": "吨",
                "data": [
                    combined.get("salesTons"),
                    retail.get("tons"),
                    wholesale.get("salesTons"),
                    combined.get("transferTons"),
                ],
            }],
            unit="吨",
            summary="按板块展示合并吨数结构。",
        )
        return [chart] if chart else []


    def _assistant_review_tables(label, data):
        review = data.get("review") or {}
        status = review.get("status") or {}
        tables = []
        if status:
            tables.append(_assistant_table(
                f"{label}匹配状态",
                ["状态", "记录数", "吨数"],
                [[key or "--", value.get("rows") or 0, _fmt_tons(value.get("tons"))] for key, value in status.items()],
            ))
        if review.get("topNeedsReview"):
            tables.append(_assistant_table(
                f"{label}待复核数据",
                ["客户/去向", "油品", "记录数", "吨数"],
                [[item.get("entity") or "--", item.get("product") or "--", item.get("rows") or 0, _fmt_tons(item.get("tons"))] for item in review["topNeedsReview"][:10]],
            ))
        issues = data.get("issues") or []
        if issues:
            tables.append(_assistant_table(
                f"{label}数据审核",
                ["类型", "级别", "日期", "提示"],
                [[item.get("type") or "--", item.get("severity") or "--", item.get("date") or "--", item.get("message") or ""] for item in issues[:12]],
            ))
        return tables


    def _assistant_year_month_from_period(start, end):
        if start.month == 1 and start.day == 1 and end == date(start.year + 1, 1, 1):
            return start.year, None
        if start.day == 1:
            next_month = date(start.year + 1, 1, 1) if start.month == 12 else date(start.year, start.month + 1, 1)
            if end == next_month:
                return start.year, start.month
        return start.year, start.month if start.year == end.year and start.month == end.month else None


    def _format_assistant_wholesale_answer(label, data, question):
        cards = data.get("cards") or {}
        export_card = cards.get("export_wholesale") or {}
        local_card = cards.get("local_wholesale") or {}
        pending_card = cards.get("unknown") or {}
        transfer_card = cards.get("transfer") or {}
        lines = [f"{label}批发数据如下："]
        if "出口" in question:
            lines.append(f"出口批发：{float(export_card.get('tons') or 0):,.3f} 吨，金额 {float(export_card.get('amountCny') or 0):,.2f} 元，客户数 {export_card.get('entities') or 0}。")
        elif "本地" in question:
            lines.append(f"本地批发：{float(local_card.get('tons') or 0):,.3f} 吨，金额 {float(local_card.get('amountCny') or 0):,.2f} 元，客户数 {local_card.get('entities') or 0}。")
        elif "外运" in question:
            lines.append(f"外运：{float(transfer_card.get('tons') or 0):,.3f} 吨，去向数 {transfer_card.get('entities') or 0}；外运仅统计吨数，不计算金额。")
        else:
            total_tons = float(export_card.get("tons") or 0) + float(local_card.get("tons") or 0) + float(pending_card.get("tons") or 0)
            total_amount = float(export_card.get("amountCny") or 0) + float(local_card.get("amountCny") or 0)
            lines.append(f"批发销售合计：{total_tons:,.3f} 吨，金额 {total_amount:,.2f} 元。")
            lines.append(f"出口批发：{float(export_card.get('tons') or 0):,.3f} 吨，{float(export_card.get('amountCny') or 0):,.2f} 元。")
            lines.append(f"本地批发：{float(local_card.get('tons') or 0):,.3f} 吨，{float(local_card.get('amountCny') or 0):,.2f} 元。")
            lines.append(f"待识别：{float(pending_card.get('tons') or 0):,.3f} 吨，数量已统计，金额待合同复核。")
            lines.append(f"外运：{float(transfer_card.get('tons') or 0):,.3f} 吨。")
        rankings = data.get("customerRankings") or {}
        key = "export_wholesale" if "出口" in question else ("local_wholesale" if "本地" in question else None)
        if key and rankings.get(key):
            _append_rank_lines(lines, "客户排行：", rankings[key], limit=8)
        elif _assistant_has_any(question, ("客户", "排行", "排名", "top", "TOP")):
            _append_rank_lines(lines, "出口批发客户排行：", rankings.get("export_wholesale"), limit=5)
            _append_rank_lines(lines, "本地批发客户排行：", rankings.get("local_wholesale"), limit=5)
        if _assistant_has_any(question, ("油品", "产品", "结构", "品类")):
            products = data.get("productRankings") or {}
            if "出口" in question:
                _append_rank_lines(lines, "出口批发油品排行：", products.get("export_wholesale"), limit=8)
            elif "本地" in question:
                _append_rank_lines(lines, "本地批发油品排行：", products.get("local_wholesale"), limit=8)
            elif "外运" in question:
                _append_rank_lines(lines, "外运油品排行：", products.get("transfer"), limit=8)
            else:
                _append_rank_lines(lines, "出口批发油品排行：", products.get("export_wholesale"), limit=5)
                _append_rank_lines(lines, "本地批发油品排行：", products.get("local_wholesale"), limit=5)
                _append_rank_lines(lines, "待识别油品排行：", products.get("unknown"), limit=5)
                _append_rank_lines(lines, "外运油品排行：", products.get("transfer"), limit=5)
        if "外运" in question and _assistant_has_any(question, ("区域", "去向", "详情", "油库", "加油站")):
            transfer_summary = data.get("transferSummary") or {}
            _append_rank_lines(lines, "外运区域汇总：", transfer_summary.get("byRegion"), limit=8)
            _append_rank_lines(lines, "外运去向排行：", data.get("transferTargets"), limit=8)
        return "\n".join(lines)


    def _assistant_exchange_answer(cur, label, start, end):
        cur.execute(
            """
            SELECT rate_date, usd_kgs, cny_kgs, source, is_manual
            FROM exchange_rates
            WHERE rate_date >= %s AND rate_date < %s
            ORDER BY rate_date DESC
            LIMIT 31
            """,
            (start, end),
        )
        rows = cur.fetchall()
        if not rows:
            return f"{label}没有查询到汇率快照。"
        latest = rows[0]
        usd_vals = [float(row[1]) for row in rows if row[1] is not None]
        cny_vals = [float(row[2]) for row in rows if row[2] is not None]
        return "\n".join([
            f"{label}汇率数据如下：",
            f"最新日期 {latest[0]}：USD/KGS {float(latest[1] or 0):.4f}，CNY/KGS {float(latest[2] or 0):.4f}，来源 {'手工录入' if latest[4] else (latest[3] or 'NBKR')}。",
            f"区间平均：USD/KGS {(sum(usd_vals)/len(usd_vals)):.4f}，CNY/KGS {(sum(cny_vals)/len(cny_vals)):.4f}。",
        ])


    def _assistant_exchange_tables(cur, label, start, end):
        cur.execute(
            """
            SELECT rate_date, usd_kgs, cny_kgs, source, is_manual
            FROM exchange_rates
            WHERE rate_date >= %s AND rate_date < %s
            ORDER BY rate_date DESC
            LIMIT 31
            """,
            (start, end),
        )
        rows = cur.fetchall()
        if not rows:
            return []
        usd_vals = [float(row[1]) for row in rows if row[1] is not None]
        cny_vals = [float(row[2]) for row in rows if row[2] is not None]
        table_rows = [
            [
                row[0],
                f"{float(row[1] or 0):.4f}",
                f"{float(row[2] or 0):.4f}",
                f"{(1 / float(row[2])):.4f}" if row[2] else "--",
                "手工录入" if row[4] else (row[3] or "NBKR"),
            ]
            for row in rows
        ]
        if usd_vals or cny_vals:
            cny_avg = (sum(cny_vals) / len(cny_vals)) if cny_vals else 0
            table_rows.append([
                "区间平均",
                f"{(sum(usd_vals) / len(usd_vals)):.4f}" if usd_vals else "--",
                f"{cny_avg:.4f}" if cny_vals else "--",
                f"{(1 / cny_avg):.4f}" if cny_avg else "--",
                "--",
            ])
        return [_assistant_table(
            f"{label}汇率快照",
            ["日期", "美元/索姆", "人民币/索姆", "索姆/人民币", "来源"],
            table_rows,
        )]


    def _assistant_exchange_charts(cur, label, start, end):
        cur.execute(
            """
            SELECT rate_date, usd_kgs, cny_kgs
            FROM exchange_rates
            WHERE rate_date >= %s AND rate_date < %s
            ORDER BY rate_date ASC
            LIMIT 180
            """,
            (start, end),
        )
        rows = cur.fetchall()
        if not rows:
            return []
        chart = _assistant_chart(
            f"{label}汇率走势",
            "line",
            x_axis=[str(row[0]) for row in rows],
            series=[
                {"name": "USD/KGS", "type": "line", "unit": "", "data": [row[1] for row in rows]},
                {"name": "CNY/KGS", "type": "line", "unit": "", "data": [row[2] for row in rows]},
            ],
            summary="按日期展示美元和人民币兑索姆汇率。",
        )
        return [chart] if chart else []


    def _assistant_density_answer(cur):
        cur.execute(
            """
            SELECT p.display_name, d.density, d.effective_date, d.source
            FROM product_density_history d
            JOIN products p ON p.id = d.product_id
            JOIN (
                SELECT product_id, MAX(effective_date) effective_date
                FROM product_density_history
                GROUP BY product_id
            ) latest ON latest.product_id = d.product_id AND latest.effective_date = d.effective_date
            ORDER BY p.product_code
            """
        )
        rows = cur.fetchall()
        if not rows:
            return "暂未查询到密度维护记录。"
        lines = ["当前最新密度如下："]
        for name, density, effective_date, source in rows:
            lines.append(f"{name}：{float(density):.6f}，生效日期 {effective_date}，来源 {source or '--'}")
        return "\n".join(lines)


    def _assistant_density_tables(cur):
        cur.execute(
            """
            SELECT p.display_name, d.density, d.effective_date, d.source
            FROM product_density_history d
            JOIN products p ON p.id = d.product_id
            JOIN (
                SELECT product_id, MAX(effective_date) effective_date
                FROM product_density_history
                GROUP BY product_id
            ) latest ON latest.product_id = d.product_id AND latest.effective_date = d.effective_date
            ORDER BY p.product_code
            """
        )
        rows = cur.fetchall()
        if not rows:
            return []
        return [_assistant_table(
            "当前密度",
            ["油品", "密度", "生效日期", "来源"],
            [[name, f"{float(density):.6f}", effective_date, source or "--"] for name, density, effective_date, source in rows],
        )]


    def _assistant_density_charts(cur):
        cur.execute(
            """
            SELECT p.display_name, d.density
            FROM product_density_history d
            JOIN products p ON p.id = d.product_id
            JOIN (
                SELECT product_id, MAX(effective_date) effective_date
                FROM product_density_history
                GROUP BY product_id
            ) latest ON latest.product_id = d.product_id AND latest.effective_date = d.effective_date
            ORDER BY p.product_code
            """
        )
        rows = cur.fetchall()
        chart = _assistant_chart(
            "当前油品密度",
            "bar",
            x_axis=[str(name) for name, _ in rows],
            series=[{"name": "密度", "type": "bar", "unit": "", "data": [density for _, density in rows]}],
            summary="展示最新生效密度。",
        )
        return [chart] if chart else []


    def _assistant_contract_answer(cur, question):
        where = []
        if "出口" in question:
            where.append("c.sales_type = 'export_wholesale'")
        if "本地" in question:
            where.append("c.sales_type = 'local_wholesale'")
        sql_where = "WHERE " + " AND ".join(where) if where else ""
        cur.execute(
            f"""
            SELECT COUNT(*),
                   SUM(CASE
                       WHEN (c.effective_start IS NULL OR c.effective_start <= CURDATE())
                        AND (c.effective_end IS NULL OR c.effective_end >= CURDATE())
                       THEN 1 ELSE 0 END),
                   MIN(c.effective_start),
                   MAX(c.effective_end)
            FROM contracts c
            {sql_where}
            """
        )
        count, active_count, min_from, max_to = cur.fetchone()
        cur.execute(
            f"""
            SELECT
                c.contract_no,
                COALESCE(cu.standard_name, c.raw_customer_name),
                COALESCE(p.display_name, c.raw_product_name),
                COALESCE(c.price_with_logistics, c.price_without_logistics, 0),
                c.currency_code,
                c.effective_start,
                c.effective_end,
                c.sales_type
            FROM contracts c
            LEFT JOIN customers cu ON cu.id = c.customer_id
            LEFT JOIN products p ON p.id = c.product_id
            {sql_where}
            ORDER BY COALESCE(c.effective_end, c.effective_start, c.signed_date) DESC, c.id DESC
            LIMIT 8
            """
        )
        rows = cur.fetchall()
        lines = [f"合同信息：共 {int(count or 0)} 条，其中有效合同 {int(active_count or 0)} 条；日期范围 {min_from or '--'} 至 {max_to or '--'}。"]
        if rows:
            lines.append("近期合同：")
            for no, customer, product, price, currency, valid_from, valid_to, sales_type in rows:
                label = "出口批发" if sales_type == "export_wholesale" else ("本地批发" if sales_type == "local_wholesale" else "其他")
                lines.append(f"{no or '--'}｜{customer or '--'}｜{product or '--'}｜{float(price or 0):,.2f} {currency}｜{valid_from or '--'} 至 {valid_to or '--'}｜{label}")
        return "\n".join(lines)


    def _assistant_contract_tables(cur, question):
        where = []
        if "出口" in question:
            where.append("c.sales_type = 'export_wholesale'")
        if "本地" in question:
            where.append("c.sales_type = 'local_wholesale'")
        sql_where = "WHERE " + " AND ".join(where) if where else ""
        cur.execute(
            f"""
            SELECT
                c.contract_no,
                COALESCE(cu.standard_name, c.raw_customer_name),
                COALESCE(p.display_name, c.raw_product_name),
                COALESCE(c.price_with_logistics, c.price_without_logistics, 0),
                c.currency_code,
                c.effective_start,
                c.effective_end,
                c.sales_type
            FROM contracts c
            LEFT JOIN customers cu ON cu.id = c.customer_id
            LEFT JOIN products p ON p.id = c.product_id
            {sql_where}
            ORDER BY COALESCE(c.effective_end, c.effective_start, c.signed_date) DESC, c.id DESC
            LIMIT 12
            """
        )
        rows = []
        for no, customer, product, price, currency, valid_from, valid_to, sales_type in cur.fetchall():
            label = "出口批发" if sales_type == "export_wholesale" else ("本地批发" if sales_type == "local_wholesale" else "其他")
            rows.append([no or "--", customer or "--", product or "--", f"{float(price or 0):,.2f} {currency or ''}".strip(), valid_from or "--", valid_to or "--", label])
        return [_assistant_table("合同明细", ["合同号", "客户", "油品", "价格", "生效", "到期", "类型"], rows)] if rows else []


    _CLASSIFICATION_LABELS = {
        "export_wholesale": "出口批发",
        "local_wholesale": "本地批发",
        "transfer": "外运/调拨",
        "procurement": "采购",
        "processing": "代加工",
        "unknown": "未分类",
    }


    def _assistant_classification_label(value):
        return _CLASSIFICATION_LABELS.get(str(value or ""), str(value or "--"))


    def _assistant_table_exists(cur, table_name):
        try:
            cur.execute("SHOW TABLES LIKE %s", (table_name,))
            return bool(cur.fetchone())
        except pymysql.MySQLError:
            return False


    def _assistant_column_exists(cur, table_name, column_name):
        try:
            cur.execute(f"SHOW COLUMNS FROM {table_name} LIKE %s", (column_name,))
            return bool(cur.fetchone())
        except pymysql.MySQLError:
            return False


    def _assistant_procurement_data(start, end):
        conn = get_connection()
        try:
            with conn.cursor(pymysql.cursors.DictCursor) as cur:
                cur.execute(
                    """
                    SELECT COUNT(*) AS records,
                           COALESCE(SUM(weight_tons), 0) AS tons,
                           COUNT(DISTINCT COALESCE(NULLIF(standard_entity_name, ''), NULLIF(sender_name, ''))) AS suppliers,
                           COUNT(DISTINCT NULLIF(raw_product_name, '')) AS products
                    FROM wholesale_detail_records
                    WHERE classification = 'procurement'
                      AND business_date >= %s AND business_date < %s
                    """,
                    (start, end),
                )
                summary = cur.fetchone() or {}
                cur.execute(
                    """
                    SELECT COALESCE(NULLIF(standard_entity_name, ''), NULLIF(sender_name, ''), '--') AS name,
                           COUNT(*) AS records,
                           COALESCE(SUM(weight_tons), 0) AS tons
                    FROM wholesale_detail_records
                    WHERE classification = 'procurement'
                      AND business_date >= %s AND business_date < %s
                    GROUP BY COALESCE(NULLIF(standard_entity_name, ''), NULLIF(sender_name, ''), '--')
                    ORDER BY tons DESC, records DESC
                    LIMIT 12
                    """,
                    (start, end),
                )
                suppliers = [dict(row) for row in cur.fetchall()]
                cur.execute(
                    """
                    SELECT COALESCE(NULLIF(raw_product_name, ''), '--') AS name,
                           COUNT(*) AS records,
                           COALESCE(SUM(weight_tons), 0) AS tons
                    FROM wholesale_detail_records
                    WHERE classification = 'procurement'
                      AND business_date >= %s AND business_date < %s
                    GROUP BY COALESCE(NULLIF(raw_product_name, ''), '--')
                    ORDER BY tons DESC, records DESC
                    LIMIT 12
                    """,
                    (start, end),
                )
                products = [dict(row) for row in cur.fetchall()]
                return {"summary": dict(summary), "suppliers": suppliers, "products": products}
        finally:
            conn.close()


    def _assistant_procurement_answer(label, data):
        summary = data.get("summary") or {}
        lines = [f"{label}原料采购数据如下："]
        lines.append(
            f"合计 {int(summary.get('records') or 0)} 笔，"
            f"{float(summary.get('tons') or 0):,.3f} 吨，"
            f"覆盖供应商 {int(summary.get('suppliers') or 0)} 个、油品 {int(summary.get('products') or 0)} 类。"
        )
        if data.get("suppliers"):
            lines.append("供应商采购量排行：")
            for item in data["suppliers"][:5]:
                lines.append(f"{item.get('name') or '--'}：{float(item.get('tons') or 0):,.3f} 吨，{int(item.get('records') or 0)} 笔")
        return "\n".join(lines)


    def _assistant_procurement_tables(label, data):
        summary = data.get("summary") or {}
        tables = [_assistant_table(
            f"{label}采购汇总",
            ["指标", "数值"],
            [
                ["记录数", int(summary.get("records") or 0)],
                ["采购吨数", f"{float(summary.get('tons') or 0):,.3f}"],
                ["供应商数", int(summary.get("suppliers") or 0)],
                ["油品数", int(summary.get("products") or 0)],
            ],
        )]
        if data.get("suppliers"):
            tables.append(_assistant_table(
                f"{label}供应商采购排行",
                ["供应商", "吨数", "记录数"],
                [[item.get("name") or "--", f"{float(item.get('tons') or 0):,.3f}", int(item.get("records") or 0)] for item in data["suppliers"]],
            ))
        if data.get("products"):
            tables.append(_assistant_table(
                f"{label}采购油品排行",
                ["油品", "吨数", "记录数"],
                [[item.get("name") or "--", f"{float(item.get('tons') or 0):,.3f}", int(item.get("records") or 0)] for item in data["products"]],
            ))
        return tables


    def _assistant_procurement_charts(label, data):
        charts = []
        supplier_chart = _assistant_rank_chart(f"{label}供应商采购吨数排行", data.get("suppliers") or [], "tons", name_key="name")
        product_chart = _assistant_rank_chart(f"{label}采购油品吨数排行", data.get("products") or [], "tons", name_key="name")
        if supplier_chart:
            charts.append(supplier_chart)
        if product_chart:
            charts.append(product_chart)
        return charts


    def _assistant_weighing_status_from_text(text):
        if _assistant_contains_any(text, ("待审核", "待复核", "未审核", "待处理")):
            return "pending"
        if _assistant_contains_any(text, ("已通过", "已审核", "通过")):
            return "approved"
        if _assistant_contains_any(text, ("驳回", "拒绝", "不通过")):
            return "rejected"
        return None


    def _assistant_weighing_classification_from_text(text):
        if _assistant_contains_any(text, ("出口", "出口批发")):
            return "export_wholesale"
        if _assistant_contains_any(text, ("本地", "本地批发")):
            return "local_wholesale"
        if _assistant_contains_any(text, ("外运", "调拨")):
            return "transfer"
        if _assistant_contains_any(text, ("采购", "进货")):
            return "procurement"
        if _assistant_contains_any(text, ("代加工", "加工")):
            return "processing"
        if _assistant_contains_any(text, ("未知", "未分类")):
            return "unknown"
        return None


    def _assistant_weighing_data(start, end, question):
        date_to = end - timedelta(days=1)
        status = _assistant_weighing_status_from_text(question)
        classification = _assistant_weighing_classification_from_text(question)
        data = get_weighing_list(
            date_from=start,
            date_to=date_to,
            status=status,
            classification=classification,
            page=1,
            page_size=5000,
        )
        items = data.get("items") or []
        by_classification = {}
        by_status = {}
        by_source = {}
        for item in items:
            cls = item.get("classification") or "unknown"
            status_key = item.get("status") or "pending"
            source = item.get("source") or "unknown"
            for bucket, key in ((by_classification, cls), (by_status, status_key), (by_source, source)):
                row = bucket.setdefault(key, {"key": key, "records": 0, "tons": 0.0})
                row["records"] += 1
                row["tons"] += float(item.get("netWeight") or 0)
        return {
            "summary": {
                "records": int(data.get("total") or len(items)),
                "tons": float(data.get("totalWeight") or 0),
                "status": status,
                "classification": classification,
            },
            "byClassification": sorted(by_classification.values(), key=lambda row: row["tons"], reverse=True),
            "byStatus": sorted(by_status.values(), key=lambda row: row["records"], reverse=True),
            "bySource": sorted(by_source.values(), key=lambda row: row["records"], reverse=True),
            "items": items[:12],
        }


    def _assistant_weighing_answer(label, data):
        summary = data.get("summary") or {}
        scope = []
        if summary.get("classification"):
            scope.append(_assistant_classification_label(summary.get("classification")))
        if summary.get("status"):
            status_label = {"pending": "待审核", "approved": "已通过", "rejected": "已驳回"}.get(summary.get("status"), summary.get("status"))
            scope.append(status_label)
        suffix = f"（{'、'.join(scope)}）" if scope else ""
        lines = [f"{label}称重数据{suffix}如下："]
        lines.append(f"合计 {int(summary.get('records') or 0)} 笔，净重 {float(summary.get('tons') or 0):,.3f} 吨。")
        if data.get("byClassification"):
            lines.append("分类汇总：")
            for item in data["byClassification"][:6]:
                lines.append(f"{_assistant_classification_label(item.get('key'))}：{item['tons']:,.3f} 吨，{item['records']} 笔")
        return "\n".join(lines)


    def _assistant_weighing_tables(label, data):
        summary = data.get("summary") or {}
        tables = [_assistant_table(
            f"{label}称重汇总",
            ["指标", "数值"],
            [["记录数", int(summary.get("records") or 0)], ["净重(吨)", f"{float(summary.get('tons') or 0):,.3f}"]],
        )]
        if data.get("byClassification"):
            tables.append(_assistant_table(
                f"{label}称重分类",
                ["分类", "吨数", "记录数"],
                [[_assistant_classification_label(item.get("key")), f"{item['tons']:,.3f}", item["records"]] for item in data["byClassification"]],
            ))
        if data.get("byStatus"):
            tables.append(_assistant_table(
                f"{label}审核状态",
                ["状态", "吨数", "记录数"],
                [[{"pending": "待审核", "approved": "已通过", "rejected": "已驳回"}.get(item.get("key"), item.get("key")), f"{item['tons']:,.3f}", item["records"]] for item in data["byStatus"]],
            ))
        if data.get("items"):
            tables.append(_assistant_table(
                f"{label}近期称重明细",
                ["日期", "车号", "发货", "收货", "油品", "净重(吨)", "分类", "状态"],
                [
                    [
                        item.get("businessDate") or "--",
                        item.get("vehicleNo") or "--",
                        item.get("senderName") or "--",
                        item.get("receiverName") or "--",
                        item.get("standardProductName") or item.get("rawProductName") or "--",
                        f"{float(item.get('netWeight') or 0):,.3f}",
                        _assistant_classification_label(item.get("classification")),
                        item.get("status") or "--",
                    ]
                    for item in data["items"]
                ],
            ))
        return tables


    def _assistant_weighing_charts(label, data):
        chart = _assistant_chart(
            f"{label}称重分类吨数",
            "bar",
            x_axis=[_assistant_classification_label(item.get("key")) for item in data.get("byClassification") or []],
            series=[{"name": "净重", "type": "bar", "unit": "吨", "data": [item.get("tons") for item in data.get("byClassification") or []]}],
            unit="吨",
            summary="按称重自动分类汇总净重。",
        )
        return [chart] if chart else []


    def _assistant_materials_data(start, end):
        conn = get_connection()
        try:
            with conn.cursor(pymysql.cursors.DictCursor) as cur:
                data = {"summary": {}, "inventory": [], "inbound": {}, "orderStatus": []}
                if _assistant_table_exists(cur, "materials"):
                    cur.execute("SELECT COUNT(*) AS total, SUM(CASE WHEN is_active THEN 1 ELSE 0 END) AS active FROM materials")
                    data["summary"]["materials"] = dict(cur.fetchone() or {})
                if _assistant_table_exists(cur, "material_categories"):
                    cur.execute("SELECT COUNT(*) AS total FROM material_categories WHERE is_active = TRUE")
                    data["summary"]["categories"] = int((cur.fetchone() or {}).get("total") or 0)
                if _assistant_table_exists(cur, "depots"):
                    cur.execute("SELECT COUNT(*) AS total FROM depots WHERE is_active = TRUE")
                    data["summary"]["depots"] = int((cur.fetchone() or {}).get("total") or 0)
                if _assistant_table_exists(cur, "material_inventory"):
                    cur.execute(
                        """
                        SELECT COALESCE(m.name_zh, CONCAT('物资-', i.material_id)) AS material_name,
                               COALESCE(d.name_zh, CONCAT('库-', i.depot_id)) AS depot_name,
                               m.unit,
                               COALESCE(i.quantity, 0) AS quantity,
                               i.updated_at
                        FROM material_inventory i
                        LEFT JOIN materials m ON m.id = i.material_id
                        LEFT JOIN depots d ON d.id = i.depot_id
                        ORDER BY i.quantity DESC
                        LIMIT 20
                        """
                    )
                    data["inventory"] = [dict(row) for row in cur.fetchall()]
                if _assistant_table_exists(cur, "material_inbound"):
                    cur.execute(
                        """
                        SELECT COUNT(*) AS records,
                               COALESCE(SUM(quantity), 0) AS quantity,
                               COALESCE(SUM(total_amount), 0) AS amount
                        FROM material_inbound
                        WHERE inbound_date >= %s AND inbound_date < %s
                        """,
                        (start, end),
                    )
                    data["inbound"] = dict(cur.fetchone() or {})
                if _assistant_table_exists(cur, "material_transfer_orders"):
                    date_expr = "COALESCE(o.order_date, DATE(o.created_at))" if _assistant_column_exists(cur, "material_transfer_orders", "order_date") else "DATE(o.created_at)"
                    cur.execute(
                        f"""
                        SELECT COALESCE(status, 'pending') AS status,
                               COUNT(*) AS records,
                               COALESCE(SUM(requested_qty), 0) AS quantity
                        FROM material_transfer_orders o
                        WHERE {date_expr} >= %s AND {date_expr} < %s
                        GROUP BY COALESCE(status, 'pending')
                        ORDER BY records DESC
                        """,
                        (start, end),
                    )
                    data["orderStatus"] = [dict(row) for row in cur.fetchall()]
                return data
        finally:
            conn.close()


    def _assistant_materials_answer(label, data, transfer_only=False):
        summary = data.get("summary") or {}
        material_summary = summary.get("materials") or {}
        inbound = data.get("inbound") or {}
        lines = [f"{label}物资{'调拨单' if transfer_only else '数据'}如下："]
        if not transfer_only:
            lines.append(
                f"物资字典 {int(material_summary.get('total') or 0)} 项，启用 {int(material_summary.get('active') or 0)} 项；"
                f"分类 {int(summary.get('categories') or 0)} 个，库点 {int(summary.get('depots') or 0)} 个。"
            )
            lines.append(
                f"本区间入库 {int(inbound.get('records') or 0)} 笔，数量 {float(inbound.get('quantity') or 0):,.3f}，金额 {float(inbound.get('amount') or 0):,.2f}。"
            )
        if data.get("orderStatus"):
            lines.append("调拨单状态：")
            for item in data["orderStatus"]:
                lines.append(f"{item.get('status') or '--'}：{int(item.get('records') or 0)} 单，数量 {float(item.get('quantity') or 0):,.3f}")
        elif transfer_only:
            lines.append("本区间暂无调拨单记录。")
        return "\n".join(lines)


    def _assistant_materials_tables(label, data, transfer_only=False):
        tables = []
        if not transfer_only:
            summary = data.get("summary") or {}
            material_summary = summary.get("materials") or {}
            inbound = data.get("inbound") or {}
            tables.append(_assistant_table(
                f"{label}物资汇总",
                ["指标", "数值"],
                [
                    ["物资字典", int(material_summary.get("total") or 0)],
                    ["启用物资", int(material_summary.get("active") or 0)],
                    ["分类数", int(summary.get("categories") or 0)],
                    ["库点数", int(summary.get("depots") or 0)],
                    ["区间入库笔数", int(inbound.get("records") or 0)],
                    ["区间入库金额", f"{float(inbound.get('amount') or 0):,.2f}"],
                ],
            ))
            if data.get("inventory"):
                tables.append(_assistant_table(
                    "当前物资库存",
                    ["物资", "库点", "数量", "单位", "更新时间"],
                    [[item.get("material_name") or "--", item.get("depot_name") or "--", f"{float(item.get('quantity') or 0):,.3f}", item.get("unit") or "", item.get("updated_at") or "--"] for item in data["inventory"]],
                ))
        if data.get("orderStatus"):
            tables.append(_assistant_table(
                f"{label}调拨单状态",
                ["状态", "单数", "数量"],
                [[item.get("status") or "--", int(item.get("records") or 0), f"{float(item.get('quantity') or 0):,.3f}"] for item in data["orderStatus"]],
            ))
        return tables


    def _assistant_materials_charts(label, data, transfer_only=False):
        charts = []
        if not transfer_only and data.get("inventory"):
            chart = _assistant_chart(
                "当前物资库存排行",
                "bar",
                x_axis=[item.get("material_name") or "--" for item in data["inventory"][:12]],
                series=[{"name": "库存数量", "type": "bar", "unit": "", "data": [item.get("quantity") for item in data["inventory"][:12]]}],
                summary="按当前库存数量排序。",
            )
            if chart:
                charts.append(chart)
        if data.get("orderStatus"):
            chart = _assistant_chart(
                f"{label}调拨单状态",
                "bar",
                x_axis=[item.get("status") or "--" for item in data["orderStatus"]],
                series=[{"name": "单数", "type": "bar", "unit": "单", "data": [item.get("records") for item in data["orderStatus"]]}],
                unit="单",
                summary="按调拨单状态汇总。",
            )
            if chart:
                charts.append(chart)
        return charts


    def _assistant_transfer_loss_data(start, end):
        names = retail_station_catalog_cached()
        retail_cfg = dict(DB_CONFIG)
        retail_cfg["database"] = os.getenv("RETAIL_DB_NAME", "ocw")
        retail_cfg["autocommit"] = True
        retail_rows = []
        retail_conn = pymysql.connect(**retail_cfg)
        try:
            with retail_conn.cursor(pymysql.cursors.DictCursor) as cur:
                cur.execute(
                    """
                    SELECT DATE(AcceptDateTime) AS supply_date,
                           ShopKey,
                           COALESCE(SUM(NULLIF(Mass, 0)), SUM(NULLIF(MassTTN, 0)), SUM(NULLIF(TankerMass, 0)), 0) / 1000 AS tons,
                           SUM(Volume) AS liters,
                           COUNT(*) AS rows
                    FROM supply
                    WHERE AcceptDateTime >= %s AND AcceptDateTime < %s
                      AND COALESCE(Deleted, 0) = 0
                    GROUP BY DATE(AcceptDateTime), ShopKey
                    """,
                    (start, end),
                )
                retail_rows = [dict(row) for row in cur.fetchall()]
        finally:
            retail_conn.close()

        conn = get_connection()
        try:
            with conn.cursor(pymysql.cursors.DictCursor) as cur:
                cur.execute(
                    """
                    SELECT business_date,
                           entity_type,
                           entity_code,
                           COALESCE(NULLIF(standard_entity_name, ''), raw_entity_name) AS target_name,
                           SUM(weight_tons) AS tons,
                           COUNT(*) AS rows
                    FROM wholesale_detail_records
                    WHERE business_date >= %s AND business_date < %s
                      AND classification = 'transfer'
                    GROUP BY business_date, entity_type, entity_code, COALESCE(NULLIF(standard_entity_name, ''), raw_entity_name)
                    """,
                    (start - timedelta(days=3), end),
                )
                wholesale_rows = [dict(row) for row in cur.fetchall()]
        finally:
            conn.close()

        direct_index = {}
        for row in wholesale_rows:
            if row.get("entity_type") != "retail_station":
                continue
            try:
                shop_key = int(row.get("entity_code") or 0)
            except (TypeError, ValueError):
                continue
            key = (row.get("business_date"), shop_key)
            item = direct_index.setdefault(key, {"tons": 0.0, "rows": 0})
            item["tons"] += float(row.get("tons") or 0)
            item["rows"] += int(row.get("rows") or 0)

        items = []
        for supply in retail_rows:
            supply_date = supply.get("supply_date")
            shop_key = int(supply.get("ShopKey") or 0)
            matched = None
            for days in range(1, 4):
                key = (supply_date - timedelta(days=days), shop_key)
                if key in direct_index:
                    matched = {**direct_index[key], "date": key[0], "daysBefore": days}
                    break
            actual_qty = float(supply.get("tons") or 0)
            system_qty = float((matched or {}).get("tons") or 0)
            loss_qty = system_qty - actual_qty if matched else 0.0
            loss_rate = loss_qty / system_qty if matched and system_qty else 0.0
            station = names.get(shop_key, {})
            items.append({
                "route": "Example总库-加油站" if matched else "分库-加油站",
                "station": station.get("short") or station.get("name") or f"{shop_key}号站",
                "systemQty": round(system_qty, 3),
                "actualQty": round(actual_qty, 3),
                "lossQty": round(loss_qty, 3),
                "lossRate": round(loss_rate, 4),
                "supplyDate": supply_date,
                "evidence": f"{matched['daysBefore']}天前有外运称重" if matched else "前1-3天未匹配到Example直发称重",
            })

        summary = {
            "records": len(items),
            "matched": sum(1 for item in items if item["systemQty"] > 0),
            "systemQty": sum(item["systemQty"] for item in items),
            "actualQty": sum(item["actualQty"] for item in items),
            "lossQty": sum(item["lossQty"] for item in items),
        }
        items.sort(key=lambda item: abs(float(item.get("lossQty") or 0)), reverse=True)
        return {"summary": summary, "items": items[:20]}


    def _assistant_transfer_loss_answer(label, data):
        summary = data.get("summary") or {}
        lines = [f"{label}调拨损耗摘要如下："]
        lines.append(
            f"收油记录 {int(summary.get('records') or 0)} 条，匹配Example直发 {int(summary.get('matched') or 0)} 条；"
            f"系统吨数 {float(summary.get('systemQty') or 0):,.3f}，实收吨数 {float(summary.get('actualQty') or 0):,.3f}，"
            f"差异 {float(summary.get('lossQty') or 0):,.3f} 吨。"
        )
        if data.get("items"):
            lines.append("差异较大的站点：")
            for item in data["items"][:5]:
                lines.append(f"{item.get('station') or '--'}：{item.get('lossQty'):,.3f} 吨，{item.get('evidence') or '--'}")
        return "\n".join(lines)


    def _assistant_transfer_loss_tables(label, data):
        summary = data.get("summary") or {}
        tables = [_assistant_table(
            f"{label}调拨损耗汇总",
            ["指标", "数值"],
            [
                ["收油记录", int(summary.get("records") or 0)],
                ["匹配Example直发", int(summary.get("matched") or 0)],
                ["系统吨数", f"{float(summary.get('systemQty') or 0):,.3f}"],
                ["实收吨数", f"{float(summary.get('actualQty') or 0):,.3f}"],
                ["差异吨数", f"{float(summary.get('lossQty') or 0):,.3f}"],
            ],
        )]
        if data.get("items"):
            tables.append(_assistant_table(
                f"{label}调拨损耗明细",
                ["站点", "路线", "系统吨数", "实收吨数", "差异吨数", "损耗率", "证据"],
                [
                    [
                        item.get("station") or "--",
                        item.get("route") or "--",
                        f"{float(item.get('systemQty') or 0):,.3f}",
                        f"{float(item.get('actualQty') or 0):,.3f}",
                        f"{float(item.get('lossQty') or 0):,.3f}",
                        _fmt_pct(float(item.get("lossRate") or 0) * 100),
                        item.get("evidence") or "--",
                    ]
                    for item in data["items"]
                ],
            ))
        return tables


    def _assistant_transfer_loss_charts(label, data):
        rows = data.get("items") or []
        chart = _assistant_chart(
            f"{label}调拨差异吨数",
            "bar",
            x_axis=[item.get("station") or "--" for item in rows[:12]],
            series=[{"name": "差异吨数", "type": "bar", "unit": "吨", "data": [item.get("lossQty") for item in rows[:12]]}],
            unit="吨",
            summary="正值表示系统外运吨数高于零售实收吨数。",
        )
        return [chart] if chart else []


    def _assistant_system_health_data():
        conn = get_connection()
        try:
            with conn.cursor(pymysql.cursors.DictCursor) as cur:
                data = {"database": "ok", "automation": [], "ai": {}, "knowledge": {}, "users": {}}
                if _assistant_table_exists(cur, "automation_task_runs"):
                    cur.execute(
                        """
                        SELECT task_name, status, MAX(finished_at) AS last_finished_at, COUNT(*) AS runs
                        FROM automation_task_runs
                        WHERE started_at >= DATE_SUB(NOW(), INTERVAL 7 DAY)
                        GROUP BY task_name, status
                        ORDER BY last_finished_at DESC
                        LIMIT 20
                        """
                    )
                    data["automation"] = [dict(row) for row in cur.fetchall()]
                if _assistant_table_exists(cur, "ai_agent_call_logs"):
                    cur.execute(
                        """
                        SELECT COUNT(*) AS total,
                               SUM(CASE WHEN status = 'ok' THEN 1 ELSE 0 END) AS ok_count,
                               SUM(CASE WHEN status <> 'ok' THEN 1 ELSE 0 END) AS failed_count
                        FROM ai_agent_call_logs
                        WHERE created_at >= DATE_SUB(NOW(), INTERVAL 7 DAY)
                        """
                    )
                    data["ai"] = dict(cur.fetchone() or {})
                if _assistant_table_exists(cur, "knowledge_attachments"):
                    cur.execute(
                        """
                        SELECT extraction_status, COUNT(*) AS records
                        FROM knowledge_attachments
                        GROUP BY extraction_status
                        """
                    )
                    data["knowledge"]["attachments"] = [dict(row) for row in cur.fetchall()]
                if _assistant_table_exists(cur, "system_users"):
                    cur.execute("SELECT COUNT(*) AS total, SUM(CASE WHEN is_active THEN 1 ELSE 0 END) AS active FROM system_users")
                    data["users"] = dict(cur.fetchone() or {})
                return data
        finally:
            conn.close()


    def _assistant_system_health_answer(data):
        lines = ["系统健康只读摘要如下："]
        lines.append("数据库连接：正常。" if data.get("database") == "ok" else "数据库连接：异常。")
        ai = data.get("ai") or {}
        if ai:
            lines.append(f"AI 最近7天调用 {int(ai.get('total') or 0)} 次，成功 {int(ai.get('ok_count') or 0)} 次，失败 {int(ai.get('failed_count') or 0)} 次。")
        if data.get("automation"):
            warning = [row for row in data["automation"] if str(row.get("status") or "").lower() not in {"ok", "success", "completed"}]
            lines.append(f"自动任务最近7天记录 {len(data['automation'])} 组，其中非成功状态 {len(warning)} 组。")
        knowledge = (data.get("knowledge") or {}).get("attachments") or []
        if knowledge:
            lines.append("知识库附件索引：" + "，".join(f"{row.get('extraction_status') or '--'} {int(row.get('records') or 0)}" for row in knowledge))
        users = data.get("users") or {}
        if users:
            lines.append(f"用户：共 {int(users.get('total') or 0)} 个，启用 {int(users.get('active') or 0)} 个。")
        return "\n".join(lines)


    def _assistant_system_health_tables(data):
        tables = [_assistant_table("系统健康摘要", ["项目", "状态/数值"], [["数据库", data.get("database") or "--"]])]
        ai = data.get("ai") or {}
        if ai:
            tables[0]["rows"].extend([["AI最近7天调用", int(ai.get("total") or 0)], ["AI成功", int(ai.get("ok_count") or 0)], ["AI失败", int(ai.get("failed_count") or 0)]])
        if data.get("automation"):
            tables.append(_assistant_table(
                "自动任务最近7天",
                ["任务", "状态", "最近完成", "次数"],
                [[row.get("task_name") or "--", row.get("status") or "--", row.get("last_finished_at") or "--", int(row.get("runs") or 0)] for row in data["automation"]],
            ))
        return tables


    def _assistant_knowledge_data(question):
        query = re.sub(r"知识库|有没有|是否|里|中|说明|资料|文档", " ", str(question or "")).strip()
        conn = get_connection()
        try:
            with conn.cursor(pymysql.cursors.DictCursor) as cur:
                data = {"summary": {}, "articles": [], "attachments": []}
                if _assistant_table_exists(cur, "knowledge_articles"):
                    cur.execute("SELECT COUNT(*) AS total, SUM(CASE WHEN status = 'published' THEN 1 ELSE 0 END) AS published FROM knowledge_articles")
                    data["summary"] = dict(cur.fetchone() or {})
                    if query:
                        like = f"%{query}%"
                        cur.execute(
                            """
                            SELECT title, summary, module_key, updated_at
                            FROM knowledge_articles
                            WHERE status = 'published'
                              AND allow_ai = TRUE
                              AND (title LIKE %s OR summary LIKE %s OR content LIKE %s)
                            ORDER BY updated_at DESC
                            LIMIT 8
                            """,
                            (like, like, like),
                        )
                        data["articles"] = [dict(row) for row in cur.fetchall()]
                if _assistant_table_exists(cur, "knowledge_attachments"):
                    cur.execute(
                        """
                        SELECT extraction_status, COUNT(*) AS records
                        FROM knowledge_attachments
                        GROUP BY extraction_status
                        """
                    )
                    data["attachments"] = [dict(row) for row in cur.fetchall()]
                data["query"] = query
                return data
        finally:
            conn.close()


    def _assistant_knowledge_answer(data):
        summary = data.get("summary") or {}
        lines = [f"知识库摘要：文章 {int(summary.get('total') or 0)} 篇，已发布 {int(summary.get('published') or 0)} 篇。"]
        if data.get("attachments"):
            lines.append("附件索引：" + "，".join(f"{row.get('extraction_status') or '--'} {int(row.get('records') or 0)}" for row in data["attachments"]))
        if data.get("query"):
            lines.append(f"关键词“{data['query']}”命中文章 {len(data.get('articles') or [])} 篇。")
        for item in (data.get("articles") or [])[:5]:
            lines.append(f"《{item.get('title') or '--'}》：{item.get('summary') or '--'}")
        return "\n".join(lines)


    def _assistant_knowledge_tables(data):
        tables = [_assistant_table(
            "知识库摘要",
            ["指标", "数值"],
            [["文章数", int((data.get("summary") or {}).get("total") or 0)], ["已发布", int((data.get("summary") or {}).get("published") or 0)]],
        )]
        if data.get("articles"):
            tables.append(_assistant_table(
                "知识库命中文章",
                ["标题", "摘要", "模块", "更新时间"],
                [[item.get("title") or "--", item.get("summary") or "--", item.get("module_key") or "--", item.get("updated_at") or "--"] for item in data["articles"]],
            ))
        return tables


    def _assistant_audit_log_data(start, end):
        conn = get_connection()
        try:
            with conn.cursor(pymysql.cursors.DictCursor) as cur:
                data = {"summary": {}, "modules": [], "actions": [], "recent": []}
                if not _assistant_table_exists(cur, "audit_logs"):
                    return data
                cur.execute(
                    "SELECT COUNT(*) AS records FROM audit_logs WHERE created_at >= %s AND created_at < %s",
                    (start, end),
                )
                data["summary"] = dict(cur.fetchone() or {})
                for field, key in (("module_key", "modules"), ("action_key", "actions")):
                    if _assistant_column_exists(cur, "audit_logs", field):
                        cur.execute(
                            f"""
                            SELECT {field} AS name, COUNT(*) AS records
                            FROM audit_logs
                            WHERE created_at >= %s AND created_at < %s
                            GROUP BY {field}
                            ORDER BY records DESC
                            LIMIT 10
                            """,
                            (start, end),
                        )
                        data[key] = [dict(row) for row in cur.fetchall()]
                cur.execute(
                    """
                    SELECT created_at, module_key, action_key, username
                    FROM audit_logs
                    WHERE created_at >= %s AND created_at < %s
                    ORDER BY created_at DESC
                    LIMIT 10
                    """,
                    (start, end),
                )
                data["recent"] = [dict(row) for row in cur.fetchall()]
                return data
        finally:
            conn.close()


    def _assistant_audit_log_answer(label, data):
        summary = data.get("summary") or {}
        lines = [f"{label}审计日志：共 {int(summary.get('records') or 0)} 条。"]
        if data.get("modules"):
            lines.append("模块排行：" + "，".join(f"{row.get('name') or '--'} {int(row.get('records') or 0)}" for row in data["modules"][:5]))
        if data.get("actions"):
            lines.append("操作排行：" + "，".join(f"{row.get('name') or '--'} {int(row.get('records') or 0)}" for row in data["actions"][:5]))
        return "\n".join(lines)


    def _assistant_audit_log_tables(label, data):
        tables = [_assistant_table(f"{label}审计日志汇总", ["指标", "数值"], [["记录数", int((data.get("summary") or {}).get("records") or 0)]])]
        if data.get("recent"):
            tables.append(_assistant_table(
                f"{label}近期审计日志",
                ["时间", "模块", "操作", "用户"],
                [[row.get("created_at") or "--", row.get("module_key") or "--", row.get("action_key") or "--", row.get("username") or "--"] for row in data["recent"]],
            ))
        return tables


    def _assistant_alias_data():
        conn = get_connection()
        try:
            with conn.cursor(pymysql.cursors.DictCursor) as cur:
                data = {"productAliases": 0, "entityAliases": 0, "products": 0, "customers": 0}
                if _assistant_table_exists(cur, "product_aliases"):
                    cur.execute("SELECT COUNT(*) AS total FROM product_aliases")
                    data["productAliases"] = int((cur.fetchone() or {}).get("total") or 0)
                if _assistant_table_exists(cur, "entity_aliases"):
                    cur.execute("SELECT COUNT(*) AS total FROM entity_aliases")
                    data["entityAliases"] = int((cur.fetchone() or {}).get("total") or 0)
                if _assistant_table_exists(cur, "products"):
                    cur.execute("SELECT COUNT(*) AS total FROM products WHERE is_active = 1")
                    data["products"] = int((cur.fetchone() or {}).get("total") or 0)
                if _assistant_table_exists(cur, "customers"):
                    cur.execute("SELECT COUNT(*) AS total FROM customers")
                    data["customers"] = int((cur.fetchone() or {}).get("total") or 0)
                return data
        finally:
            conn.close()


    def _assistant_user_permission_data():
        conn = get_connection()
        try:
            with conn.cursor(pymysql.cursors.DictCursor) as cur:
                data = {"users": {}, "roles": 0, "departments": 0, "permissions": 0}
                if _assistant_table_exists(cur, "system_users"):
                    cur.execute("SELECT COUNT(*) AS total, SUM(CASE WHEN is_active THEN 1 ELSE 0 END) AS active FROM system_users")
                    data["users"] = dict(cur.fetchone() or {})
                if _assistant_table_exists(cur, "system_roles"):
                    cur.execute("SELECT COUNT(*) AS total FROM system_roles WHERE is_active = TRUE")
                    data["roles"] = int((cur.fetchone() or {}).get("total") or 0)
                if _assistant_table_exists(cur, "system_departments"):
                    cur.execute("SELECT COUNT(*) AS total FROM system_departments WHERE is_active = TRUE")
                    data["departments"] = int((cur.fetchone() or {}).get("total") or 0)
                if _assistant_table_exists(cur, "system_permissions"):
                    cur.execute("SELECT COUNT(*) AS total FROM system_permissions WHERE is_active = TRUE")
                    data["permissions"] = int((cur.fetchone() or {}).get("total") or 0)
                return data
        finally:
            conn.close()


    def _assistant_settings_answer():
        return (
            "设置页属于只读安全边界：我可以说明当前系统公开入口和配置边界，但不会读取或修改本机隐藏配置、密钥或用户本地缓存。\n"
            "当前生产公开入口为 http://198.51.100.10/，后端生产服务只监听 127.0.0.1:8900，客户端自动更新源为 /download/latest.yml。"
        )


    def _assistant_market_watch_answer(cur):
        ensure_market_watch_schema(cur)
        cur.execute(
            """
            SELECT market_region, title, summary, price_text, source_name, observed_at
            FROM market_watch_items
            WHERE is_active = 1
            ORDER BY observed_at DESC, priority DESC, id DESC
            LIMIT 12
            """
        )
        rows = cur.fetchall()
        if not rows:
            return "暂未查询到市场观察信息。可以在后台执行“立即获取市场观察”后再查看。", []
        latest = max((row[5] for row in rows if row[5]), default=None)
        lines = [f"市场观察信息如下（最新获取时间 {latest or '--'}）："]
        region_names = {
            "russia": "俄罗斯",
            "Russia": "俄罗斯",
            "uzbekistan": "乌兹别克",
            "Uzbekistan": "乌兹别克",
            "global": "国际",
            "Global": "国际",
        }
        table_rows = []
        for region, title, summary, price_text, source, observed_at in rows[:9]:
            item = chinese_market_item({
                "region": region,
                "title": title,
                "summary": summary,
                "sourceName": source,
                "observedAt": observed_at.isoformat() if observed_at else None,
            })
            label = region_names.get(region, region or "市场")
            lines.append(f"{label}｜{item.get('title') or '--'}")
            if item.get("summary"):
                lines.append(f"  {item['summary'][:180]}")
            if price_text:
                lines.append(f"  价格：{price_text}")
            source_label = market_source_display(item.get("sourceName") or source)
            lines.append(f"  来源：{source_label}，时间：{observed_at or '--'}")
            table_rows.append([
                label,
                item.get("title") or "--",
                price_text or "--",
                item.get("summary") or "--",
                source_label,
                observed_at or "--",
            ])
        return "\n".join(lines), [_assistant_table("市场观察", ["区域", "标题", "价格信息", "摘要", "来源", "日期"], table_rows)]


    def _assistant_review_answer(label, data):
        review = data.get("review") or {}
        issues = data.get("issues") or []
        status = review.get("status") or {}
        matched = status.get("matched", {})
        need_review_rows = sum(int(item.get("rows") or 0) for key, item in status.items() if key != "matched")
        lines = [f"{label}审核信息如下："]
        lines.append(
            f"已匹配 {int(matched.get('rows') or 0)} 条，"
            f"待复核 {need_review_rows} 条，"
            f"开放数据问题 {len(issues)} 条。"
        )
        if review.get("topNeedsReview"):
            lines.append("待复核数据排行：")
            for item in review["topNeedsReview"][:8]:
                lines.append(
                    f"{item.get('entity') or '--'}｜{item.get('product') or '--'}："
                    f"{int(item.get('rows') or 0)} 条，{float(item.get('tons') or 0):,.3f} 吨"
                )
        if issues:
            lines.append("待处理问题：")
            for item in issues[:8]:
                lines.append(f"{item.get('type') or '--'}｜{item.get('severity') or '--'}｜{item.get('date') or '--'}：{item.get('message') or ''}")
        return "\n".join(lines)


    def _assistant_combined_answer(label, data):
        retail = data.get("retail") or {}
        wholesale = data.get("wholesale") or {}
        combined = data.get("combined") or {}
        lines = [f"{label}合并统计如下："]
        lines.append(f"合并销售量：{float(combined.get('salesTons') or 0):,.3f} 吨，销售金额：{float(combined.get('salesAmountCny') or 0):,.2f} 元。")
        lines.append(f"零售：{float(retail.get('tons') or 0):,.3f} 吨，{float(retail.get('amountCny') or 0):,.2f} 元。")
        lines.append(f"批发销售：{float(wholesale.get('salesTons') or 0):,.3f} 吨，{float(wholesale.get('salesAmountCny') or 0):,.2f} 元。")
        lines.append(f"外运：{float(combined.get('transferTons') or 0):,.3f} 吨。")
        if wholesale.get("byProductCategory"):
            products = wholesale["byProductCategory"]
            _append_rank_lines(lines, "出口批发油品排行：", products.get("export_wholesale"), limit=3)
            _append_rank_lines(lines, "本地批发油品排行：", products.get("local_wholesale"), limit=3)
        return "\n".join(lines)


    # ── Enhanced assistant: new query types ──

    def _read_daily_trend_snapshot(ym):
        year, month = ym.split("-")
        path = OILSTATION_ROOT / "cache_snapshots" / "monthly" / year / f"daily_{int(month):02d}.json"
        if not path.exists():
            return None
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return None


    def _read_yearly_kpi_snapshot(year):
        path = OILSTATION_ROOT / "cache_snapshots" / "yearly" / f"kpi_yearly_{year}_summary.json"
        if not path.exists():
            return None
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return None


    def _read_yearly_dashboard_snapshot(year):
        path = OILSTATION_ROOT / "cache_snapshots" / "yearly" / f"dashboard_yearly_{year}.json"
        if not path.exists():
            return None
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return None


    def _assistant_daily_trend_answer(label, data):
        daily = data.get("daily") or []
        if not daily:
            return f"{label}趋势数据暂未生成。"
        volumes = [float(d.get("volume") or 0) for d in daily]
        amounts = [float(d.get("amount") or 0) for d in daily]
        total_vol = sum(volumes)
        total_amt = sum(amounts)
        avg_vol = total_vol / len(volumes) if volumes else 0
        max_idx = volumes.index(max(volumes)) if volumes else -1
        min_idx = volumes.index(min(volumes)) if volumes else -1
        lines = [f"{label}每日销量趋势（共 {len(daily)} 天数据）："]
        lines.append(f"合计销量 {_fmt_volume(total_vol)}，合计金额 {_fmt_amount(total_amt, '索姆')}")
        lines.append(f"日均销量 {_fmt_volume(avg_vol)}")
        if max_idx >= 0:
            lines.append(f"最高日：{daily[max_idx]['date']} 销量 {_fmt_volume(volumes[max_idx])} 金额 {_fmt_amount(amounts[max_idx], '索姆')}")
        if min_idx >= 0:
            lines.append(f"最低日：{daily[min_idx]['date']} 销量 {_fmt_volume(volumes[min_idx])} 金额 {_fmt_amount(amounts[min_idx], '索姆')}")
        return "\n".join(lines)


    def _assistant_daily_trend_tables(label, data):
        daily = data.get("daily") or []
        if not daily:
            return []
        return [_assistant_table(
            f"{label}每日趋势",
            ["日期", "销量(升)", "金额(索姆)"],
            [
                [d.get("date", ""), _fmt_volume(d.get("volume")), _fmt_amount(d.get("amount"), "索姆")]
                for d in daily
            ],
        )]


    def _assistant_yearly_kpi_answer(label, data):
        summary = data.get("summary") or {}
        comp = data.get("comparison") or {}
        lines = [f"{label}年度KPI如下："]
        lines.append(f"累计销量 {_fmt_volume(summary.get('totalVolume') or summary.get('totalLiters'))}")
        lines.append(f"累计金额 {_fmt_amount(summary.get('totalSales') or summary.get('totalAmount'), '索姆')}")
        lines.append(f"交易笔数 {int(summary.get('transactionCount') or 0)} 笔")
        lines.append(f"活跃站点 {int(summary.get('activeShops') or 0)} 个")
        yoy_vol = comp.get("yoyVolume")
        yoy_sales = comp.get("yoySales")
        if yoy_vol is not None:
            lines.append(f"同比销量：{_fmt_pct(yoy_vol)}")
        if yoy_sales is not None:
            lines.append(f"同比金额：{_fmt_pct(yoy_sales)}")
        return "\n".join(lines)


    def _assistant_yearly_kpi_tables(label, data):
        summary = data.get("summary") or {}
        comp = data.get("comparison") or {}
        rows = [
            ["累计销量", _fmt_volume(summary.get("totalVolume") or summary.get("totalLiters"))],
            ["累计金额", _fmt_amount(summary.get("totalSales") or summary.get("totalAmount"), "索姆")],
            ["交易笔数", str(int(summary.get("transactionCount") or 0))],
            ["活跃站点", str(int(summary.get("activeShops") or 0))],
        ]
        yoy_vol = comp.get("yoyVolume")
        yoy_sales = comp.get("yoySales")
        if yoy_vol is not None:
            rows.append(["同比销量", _fmt_pct(yoy_vol)])
        if yoy_sales is not None:
            rows.append(["同比金额", _fmt_pct(yoy_sales)])
        inventory = summary.get("inventory")
        if inventory is not None and float(inventory) > 0:
            rows.append(["期末库存", _fmt_volume(inventory)])
        return [_assistant_table(f"{label}年度KPI", ["指标", "数值"], rows)]


    def _assistant_yearly_kpi_charts(label, data):
        summary = data.get("summary") or {}
        chart = _assistant_chart(
            f"{label}年度KPI关键指标",
            "bar",
            x_axis=["累计销量", "累计金额", "交易笔数", "活跃站点"],
            series=[{
                "name": "指标值",
                "type": "bar",
                "unit": "",
                "data": [
                    summary.get("totalVolume") or summary.get("totalLiters"),
                    summary.get("totalSales") or summary.get("totalAmount"),
                    summary.get("transactionCount"),
                    summary.get("activeShops"),
                ],
            }],
            summary="展示年度累计关键指标，单位以指标本身为准。",
        )
        return [chart] if chart else []


    def _fmt_pct(v):
        if v is None:
            return "--"
        sign = "+" if v > 0 else ""
        return f"{sign}{v:.1f}%"


    def _assistant_station_status_answer():
        conn = get_connection()
        try:
            cur = conn.cursor()
            statuses = active_station_statuses(cur)
        finally:
            conn.close()
        if not statuses:
            return "当前所有站点正常营业。"
        closed = [s for s in statuses.values() if s.get("status") == "closed"]
        warning = [s for s in statuses.values() if s.get("status") == "warning"]
        lines = []
        if closed:
            lines.append(f"暂停营业/停站点（{len(closed)} 个）：")
            for s in closed:
                reason = s.get("closedReason") or s.get("statusLabel") or "未知"
                lines.append(f"  {s.get('name')}（{s.get('region')}）— {reason}")
        if warning:
            if closed:
                lines.append("")
            lines.append(f"异常状态站点（{len(warning)} 个）：")
            for s in warning:
                reason = s.get("closedReason") or s.get("statusLabel") or "未知"
                lines.append(f"  {s.get('name')}（{s.get('region')}）— {reason}（不影响销售）")
        if not closed and not warning:
            lines.append("当前所有站点状态正常。")
        return "\n".join(lines)


    def _assistant_station_status_tables():
        conn = get_connection()
        try:
            cur = conn.cursor()
            statuses = active_station_statuses(cur)
        finally:
            conn.close()
        if not statuses:
            return []
        closed = [s for s in statuses.values() if s.get("status") == "closed"]
        warning = [s for s in statuses.values() if s.get("status") == "warning"]
        all_items = closed + warning
        if not all_items:
            return []
        return [_assistant_table(
            "当前站点异常状态",
            ["站点", "区域", "状态", "原因", "开始时间"],
            [
                [
                    s.get("name", "--"),
                    s.get("region", "--"),
                    "停业" if s.get("status") == "closed" else "异常",
                    s.get("closedReason") or s.get("statusLabel") or "",
                    s.get("startedAt") or "--",
                ]
                for s in all_items
            ],
        )]


    def _assistant_response(question, plan, **payload):
        charts = payload.get("charts")
        if isinstance(charts, list):
            payload["charts"] = [chart for chart in charts if isinstance(chart, dict)]
        result = {"ok": True, "question": question, **payload}
        result["assistantPlan"] = plan
        result = enrich_assistant_payload(result)
        return result, 200


    def _assistant_scope_module(plan):
        route = str((plan or {}).get("route") or "")
        if route.startswith("retail_") or route.startswith("station_") or route in {"self_use_summary", "accounting_report_status", "density_summary"}:
            return "retail"
        return route.split("_", 1)[0] if "_" in route else route


    def _assistant_scope_values(user_context, module_key, resource_types):
        if not isinstance(user_context, dict) or user_context.get("user", {}).get("role") == "admin":
            return None
        if (user_context.get("dataScopes") or {}).get(module_key) != "custom":
            return None
        resources = (user_context.get("dataScopeResources") or {}).get(module_key) or {}
        ids = set()
        labels = set()
        for resource_type in resource_types:
            for item in resources.get(resource_type) or []:
                if isinstance(item, dict):
                    if item.get("id") not in (None, ""):
                        ids.add(str(item.get("id")))
                    if item.get("label"):
                        labels.add(str(item.get("label")))
        return {"ids": ids, "labels": labels}


    def _assistant_station_allowed(item, allowed):
        if allowed is None:
            return True
        ids = allowed.get("ids") or set()
        labels = allowed.get("labels") or set()
        candidates = {
            str(item.get("shopKey") or ""),
            str(item.get("stationKey") or ""),
            str(item.get("stationId") or ""),
            str(item.get("id") or ""),
            str(item.get("stationNo") or ""),
        }
        label_candidates = {
            str(item.get("name") or ""),
            str(item.get("short") or ""),
            str(item.get("stationName") or ""),
            str(item.get("resourceLabel") or ""),
        }
        return bool((ids and candidates.intersection(ids)) or (labels and label_candidates.intersection(labels)))


    def _assistant_apply_station_scope(data, plan, user_context):
        module_key = _assistant_scope_module(plan)
        allowed = _assistant_scope_values(user_context, module_key, ("station", "depot"))
        if allowed is None:
            return data
        copied = dict(data or {})
        applied = {"type": "custom", "module": module_key, "resourceTypes": ["station", "depot"], "filtered": True}
        stations = copied.get("stations")
        if isinstance(stations, list):
            filtered = [item for item in stations if isinstance(item, dict) and _assistant_station_allowed(item, allowed)]
            copied["stations"] = filtered
            summary = dict(copied.get("summary") or {})
            summary["totalLiters"] = round(sum(float(item.get("volume") or 0) for item in filtered), 2)
            summary["totalTons"] = round(sum(float(item.get("tons") or 0) for item in filtered), 3)
            summary["totalSales"] = round(sum(float(item.get("amount") or 0) for item in filtered), 2)
            summary["totalTxn"] = int(sum(float(item.get("txnCount") or item.get("transactions") or 0) for item in filtered))
            copied["summary"] = summary
            copied["byResource"] = []
            copied["regions"] = []
        elif isinstance(stations, dict):
            filtered_stations = {
                key: value for key, value in stations.items()
                if _assistant_station_allowed({"shopKey": key}, allowed)
            }
            copied["stations"] = filtered_stations
            daily_by_date = {}
            for rows in filtered_stations.values():
                for item in rows or []:
                    if not isinstance(item, dict) or not item.get("date"):
                        continue
                    bucket = daily_by_date.setdefault(str(item.get("date")), {"date": str(item.get("date")), "volume": 0.0, "amount": 0.0})
                    bucket["volume"] += float(item.get("volume") or 0)
                    bucket["amount"] += float(item.get("amount") or 0)
            copied["daily"] = [
                {**item, "volume": round(item["volume"], 2), "amount": round(item["amount"], 2)}
                for _, item in sorted(daily_by_date.items())
            ]
        records = copied.get("records")
        if isinstance(records, list):
            filtered_records = [item for item in records if isinstance(item, dict) and _assistant_station_allowed(item, allowed)]
            copied["records"] = filtered_records
            summary = dict(copied.get("summary") or {})
            summary["totalVolume"] = round(sum(float(item.get("volume") or 0) for item in filtered_records), 2)
            summary["totalAmount"] = round(sum(float(item.get("amount") or 0) for item in filtered_records), 2)
            summary["totalRecords"] = len(filtered_records)
            copied["summary"] = summary
        copied["dataScope"] = applied
        return copied


    def _assistant_filter_retail_resource(data, resource):
        if not isinstance(resource, dict) or not resource.get("resourceKey"):
            return data
        resource_key = _assistant_int(resource.get("resourceKey"))
        if not resource_key:
            return data
        copied = dict(data or {})
        resources = [
            item for item in copied.get("byResource") or []
            if _assistant_int(item.get("resourceKey") or item.get("key")) == resource_key
        ]
        summary = dict(copied.get("summary") or {})
        if resources:
            summary["totalLiters"] = round(sum(float(item.get("volume") or 0) for item in resources), 2)
            summary["totalTons"] = round(sum(float(item.get("tons") or 0) for item in resources), 3)
            summary["totalSales"] = round(sum(float(item.get("amount") or 0) for item in resources), 2)
            copied["byResource"] = resources
            copied["summary"] = summary
            copied["stations"] = []
            copied["regions"] = []
            copied["resourceFilter"] = {"resourceKey": resource_key, "name": resources[0].get("name") or resource.get("name")}
        return copied


    def _retail_payload_needs_mass_refresh(data):
        if not isinstance(data, dict):
            return False
        resources = data.get("byResource") or data.get("resources") or []
        if not resources:
            return False
        resource_tons = 0.0
        has_volume_without_tons = False
        for item in resources:
            if not isinstance(item, dict):
                continue
            volume = float(item.get("volume") or 0)
            if volume <= 0:
                continue
            mass = _retail_mass_from_item(item)
            if mass is None:
                has_volume_without_tons = True
                continue
            resource_tons += mass / 1000
        summary = data.get("summary") or {}
        summary_tons = float(summary.get("totalTons") or 0)
        if has_volume_without_tons and summary_tons > 0:
            return True
        if summary_tons > 0 and resource_tons > 0:
            return abs(summary_tons - resource_tons) > max(1.0, summary_tons * 0.05)
        return False


    def _assistant_retail_data_for_period(start, end, resource=None, *, plan=None, user_context=None):
        data = _retail_snapshot_for_period(start, end)
        needs_refresh = _retail_payload_needs_mass_refresh(data)
        if not data or needs_refresh:
            response = open_oilstation(
                "/api/assistant/retail-summary",
                query_string={
                    "start": start.isoformat(),
                    "end": end.isoformat(),
                    **({"refresh": "1"} if needs_refresh else {}),
                },
            )
            fresh_data = response.get_json() if response.status_code < 400 else {}
            if fresh_data:
                data = fresh_data
        data = _assistant_filter_retail_resource(data or {}, resource)
        return _assistant_apply_station_scope(data, plan or {}, user_context)


    def _assistant_retail_summary_metrics(data):
        summary = data.get("summary") or {}
        return {
            "volume": float(summary.get("totalLiters") or 0),
            "tons": float(summary.get("totalTons") or 0),
            "amount": float(summary.get("totalSales") or 0),
            "transactions": float(summary.get("totalTxn") or 0),
        }


    def _assistant_change_text(metric, current_value, baseline_value):
        delta = current_value - baseline_value
        pct = (delta / baseline_value * 100) if baseline_value else 0.0
        direction = "增加" if delta >= 0 else "减少"
        return f"{_assistant_metric_label(metric)}{direction} {_assistant_metric_format(metric, abs(delta))}（{pct:+.1f}%）"


    def _assistant_retail_comparison_answer(label, current_data, comparison):
        base_period = (comparison or {}).get("period") or {}
        base_label = base_period.get("label") or (comparison or {}).get("mode") or "对比区间"
        current = _assistant_retail_summary_metrics(current_data)
        baseline = _assistant_retail_summary_metrics((comparison or {}).get("data") or {})
        lines = [f"{label}与{base_label}零售对比："]
        for metric in ("volume", "tons", "amount"):
            lines.append(
                f"  {_assistant_metric_label(metric)}：本期 {_assistant_metric_format(metric, current[metric])}，"
                f"{base_label} {_assistant_metric_format(metric, baseline[metric])}，"
                f"{_assistant_change_text(metric, current[metric], baseline[metric])}。"
            )
        lines.append("说明：这里按当前可读零售快照直接计算区间汇总对比；如需深层原因，可继续追问具体站点、油品或日期。")
        return "\n".join(lines)


    def _assistant_retail_comparison_tables(label, current_data, comparison):
        base_period = (comparison or {}).get("period") or {}
        base_label = base_period.get("label") or (comparison or {}).get("mode") or "对比区间"
        current = _assistant_retail_summary_metrics(current_data)
        baseline = _assistant_retail_summary_metrics((comparison or {}).get("data") or {})
        rows = []
        for metric in ("volume", "tons", "amount"):
            delta = current[metric] - baseline[metric]
            pct = (delta / baseline[metric] * 100) if baseline[metric] else 0.0
            rows.append([
                _assistant_metric_label(metric),
                _assistant_metric_format(metric, current[metric]),
                _assistant_metric_format(metric, baseline[metric]),
                _assistant_metric_format(metric, delta),
                f"{pct:+.1f}%",
            ])
        return [_assistant_table(f"{label}零售对比", ["指标", "本期", base_label, "变化", "变化率"], rows)]


    def _assistant_retail_comparison_charts(label, current_data, comparison):
        base_period = (comparison or {}).get("period") or {}
        base_label = base_period.get("label") or (comparison or {}).get("mode") or "对比区间"
        current = _assistant_retail_summary_metrics(current_data)
        baseline = _assistant_retail_summary_metrics((comparison or {}).get("data") or {})
        chart = _assistant_chart(
            f"{label}与{base_label}零售对比",
            "bar",
            x_axis=["销量", "吨位", "金额"],
            series=[
                {"name": "本期", "type": "bar", "unit": "", "data": [current["volume"], current["tons"], current["amount"]]},
                {"name": base_label, "type": "bar", "unit": "", "data": [baseline["volume"], baseline["tons"], baseline["amount"]]},
            ],
            summary="同一图中包含不同单位，精确单位请看下方表格。",
        )
        return [chart] if chart else []


    def _assistant_enrich_comparison_data(plan, user_context=None):
        comparison = plan.get("comparison") if isinstance(plan.get("comparison"), dict) else None
        if not comparison:
            return None
        period = _assistant_period_from_payload(comparison.get("period"))
        if not period:
            return comparison
        _, baseline_start, baseline_end = period
        resource = ((plan.get("entities") or {}).get("resource") if isinstance(plan.get("entities"), dict) else None)
        return {**comparison, "data": _assistant_retail_data_for_period(baseline_start, baseline_end, resource, plan=plan, user_context=user_context)}


    def _assistant_payload_for_plan(original_question, plan, user_context=None):
        period = plan.get("period") or {}
        start = date.fromisoformat(period.get("start"))
        end = date.fromisoformat(period.get("endExclusive"))
        label = _assistant_period_display(period.get("label") or "", start, end)
        match_question = plan.get("normalizedQuestion") or original_question
        route = plan.get("route")
        resource = (plan.get("entities") or {}).get("resource")
        if route == "self_use_summary":
            data = _assistant_self_use_data_for_period(start, end, resource)
            return {
                "periodLabel": label,
                "queryStart": start.isoformat(),
                "queryEndExclusive": end.isoformat(),
                "answer": _assistant_self_use_answer(label, data),
                "tables": _assistant_self_use_tables(label, data),
                "charts": _assistant_self_use_charts(label, data, match_question, visual=plan.get("visual")),
                "data": data,
            }
        if route == "station_trend_anomaly":
            station = (plan.get("entities") or {}).get("station")
            trend_data = _assistant_daily_trend_for_period(start, end)
            return {
                "periodLabel": label,
                "queryStart": start.isoformat(),
                "queryEndExclusive": end.isoformat(),
                "answer": _assistant_station_trend_anomaly_answer(label, trend_data, station),
                "tables": _assistant_station_trend_anomaly_tables(label, trend_data, station),
                "charts": _assistant_station_trend_anomaly_charts(label, trend_data, station),
                "data": trend_data,
            }
        if route == "retail_trend":
            trend_data = _assistant_daily_trend_for_period(start, end)
            return {
                "periodLabel": label,
                "queryStart": start.isoformat(),
                "queryEndExclusive": end.isoformat(),
                "answer": _assistant_daily_trend_answer(label, trend_data),
                "tables": _assistant_daily_trend_tables(label, trend_data),
                "charts": [_assistant_daily_mixed_chart(f"{label}零售每日趋势", trend_data.get("daily") or [])],
                "data": trend_data,
            }
        if route == "retail_ranking":
            dimension = next((item for item in plan.get("dimensions") or [] if item in {"resource", "region", "station"}), "station")
            data = _assistant_retail_data_for_period(start, end, resource, plan=plan, user_context=user_context)
            metric = plan.get("metric") or "amount"
            ranking_rows = _assistant_retail_ranking_rows(data, dimension, metric)
            return {
                "periodLabel": label,
                "queryStart": start.isoformat(),
                "queryEndExclusive": end.isoformat(),
                "answer": _assistant_retail_ranking_answer(label, data, dimension, metric, match_question),
                "tables": _assistant_retail_ranking_tables(label, data, dimension, metric),
                "charts": [_assistant_rank_chart(f"{label}零售{ {'resource': '油品', 'region': '区域', 'station': '站点'}.get(dimension, '站点') }{_assistant_metric_name(metric)}排行", ranking_rows, metric, name_key="region" if dimension == "region" else "name")],
                "data": data,
            }
        if route == "retail_comparison":
            data = _assistant_retail_data_for_period(start, end, resource, plan=plan, user_context=user_context)
            comparison = _assistant_enrich_comparison_data(plan, user_context)
            if comparison:
                return {
                    "periodLabel": label,
                    "queryStart": start.isoformat(),
                    "queryEndExclusive": end.isoformat(),
                    "answer": _assistant_retail_comparison_answer(label, data, comparison),
                    "tables": _assistant_retail_comparison_tables(label, data, comparison),
                    "charts": _assistant_retail_comparison_charts(label, data, comparison),
                    "data": {"current": data, "comparison": comparison},
                }
        data = _assistant_retail_data_for_period(start, end, resource, plan=plan, user_context=user_context)
        if data:
            return {
                "periodLabel": label,
                "queryStart": start.isoformat(),
                "queryEndExclusive": end.isoformat(),
                "answer": _format_assistant_retail_answer(label, data),
                "tables": _assistant_retail_tables(label, data, match_question),
                "charts": _assistant_retail_charts(label, data, match_question, visual=plan.get("visual")),
                "data": data,
            }
        return {
            "periodLabel": label,
            "queryStart": start.isoformat(),
            "queryEndExclusive": end.isoformat(),
            "answer": f"{label}暂时没有查到可用数据。",
            "tables": [],
            "data": {},
        }


    def _assistant_compound_payload(question, plan, user_context=None):
        sub_results = []
        summary_rows = []
        tables = []
        charts = []
        for index, subplan in enumerate((plan.get("subPlans") or [])[:5], start=1):
            payload = _assistant_payload_for_plan(question, subplan, user_context)
            answer = str(payload.get("answer") or "")
            first_line = next((line.strip() for line in answer.splitlines() if line.strip()), "")
            sub_results.append({"plan": subplan, **payload})
            summary_rows.append([
                index,
                subplan.get("sourceSegment") or subplan.get("normalizedQuestion") or "--",
                (subplan.get("tool") or {}).get("description") or subplan.get("route") or "--",
                (subplan.get("period") or {}).get("label") or "--",
                first_line[:180],
            ])
            for table in payload.get("tables") or []:
                copied = dict(table)
                copied["title"] = f"{index}. {copied.get('title') or subplan.get('sourceSegment') or '子问题'}"
                tables.append(copied)
            for chart in payload.get("charts") or []:
                copied = dict(chart)
                copied["title"] = f"{index}. {copied.get('title') or subplan.get('sourceSegment') or '子问题'}"
                charts.append(copied)
        header = f"我把这个问题拆成 {len(sub_results)} 个只读查询，并按顺序完成："
        lines = [header]
        for row in summary_rows:
            lines.append(f"{row[0]}. {row[1]}：{row[4] or '已完成查询。'}")
        summary_table = _assistant_table("复合问题执行摘要", ["序号", "子问题", "工具", "时间", "结论"], summary_rows)
        return {
            "periodLabel": (plan.get("period") or {}).get("label") or "复合查询",
            "answer": "\n".join(lines),
            "tables": [summary_table, *tables[:8]],
            "charts": charts[:6],
            "data": {"subResults": sub_results},
        }


    def _assistant_metric_label(metric):
        return {"amount": "销售金额", "tons": "吨位", "volume": "销量"}.get(metric, "销售金额")


    def _assistant_metric_value(row, metric):
        if metric == "tons":
            return float(row.get("tons") or 0)
        if metric == "volume":
            return float(row.get("volume") or 0)
        return float(row.get("amount") or 0)


    def _assistant_metric_format(metric, value):
        if metric == "tons":
            return _fmt_tons(value)
        if metric == "volume":
            return _fmt_volume(value)
        return _fmt_amount(value, "索姆")


    def _assistant_retail_ranking_rows(data, dimension, metric):
        source_key = {"resource": "byResource", "region": "regions"}.get(dimension, "stations")
        rows = [
            item for item in (data.get(source_key) or [])
            if _assistant_metric_value(item, metric) or float(item.get("volume") or 0) or float(item.get("amount") or 0)
        ]
        rows.sort(key=lambda item: _assistant_metric_value(item, metric), reverse=True)
        return rows


    def _assistant_retail_ranking_answer(label, data, dimension, metric, question):
        dimension_label = {"resource": "油品", "region": "区域", "station": "站点"}.get(dimension, "站点")
        rows = _assistant_retail_ranking_rows(data, dimension, metric)
        if not rows:
            return f"{label}暂无可用于{dimension_label}排行的零售数据。"
        metric_label = _assistant_metric_label(metric)
        total = sum(_assistant_metric_value(item, metric) for item in rows)
        lines = [f"{label}零售{dimension_label}{metric_label}排行："]
        for index, item in enumerate(rows[:10], start=1):
            name = item.get("name") or item.get("region") or "--"
            value = _assistant_metric_value(item, metric)
            share = value / total * 100 if total else 0
            detail = f"，占前列合计 {share:.1f}%" if total else ""
            if dimension == "station":
                detail += f"，区域 {item.get('region') or '--'}"
            lines.append(f"{index}. {name}：{_assistant_metric_format(metric, value)}{detail}")
        if _assistant_contains_any(question, ("为什么", "原因", "怎么回事")):
            leader = rows[0]
            leader_name = leader.get("name") or leader.get("region") or "--"
            lines.append(
                f"判断依据：{leader_name} 在该口径下的{metric_label}最高；"
                "系统先按已结束班组销售数据排序，深层原因还需要结合价格、库存、站点状态和当地车流继续核查。"
            )
        return "\n".join(lines)


    def _assistant_retail_ranking_tables(label, data, dimension, metric):
        dimension_label = {"resource": "油品", "region": "区域", "station": "站点"}.get(dimension, "站点")
        rows = _assistant_retail_ranking_rows(data, dimension, metric)
        table_rows = []
        for index, item in enumerate(rows[:12], start=1):
            name = item.get("name") or item.get("region") or "--"
            table_row = [
                index,
                name,
                _fmt_volume(item.get("volume")),
                _fmt_tons(item.get("tons")),
                _fmt_amount(item.get("amount"), "索姆"),
            ]
            if dimension == "station":
                table_row.insert(2, item.get("region") or "--")
            table_rows.append(table_row)
        columns = ["排名", dimension_label, "销量", "吨位", "销售金额"]
        if dimension == "station":
            columns.insert(2, "区域")
        return [_assistant_table(f"{label}零售{dimension_label}排行", columns, table_rows)] if table_rows else []


    def _assistant_daily_trend_for_period(start, end):
        daily_rows = []
        station_rows = {}
        cursor = date(start.year, start.month, 1)
        while cursor < end:
            ym = f"{cursor.year}-{cursor.month:02d}"
            snapshot = _read_daily_trend_snapshot(ym) or {}
            for item in snapshot.get("daily") or []:
                item_date = parse_date(item.get("date"))
                if item_date and start <= item_date < end:
                    daily_rows.append(item)
            for shop_key, rows in (snapshot.get("stations") or {}).items():
                clean_rows = []
                for item in rows or []:
                    item_date = parse_date(item.get("date"))
                    if item_date and start <= item_date < end:
                        clean_rows.append(item)
                if clean_rows:
                    station_rows.setdefault(str(shop_key), []).extend(clean_rows)
            next_year, next_month = _month_after(cursor.year, cursor.month)
            cursor = date(next_year, next_month, 1)
        daily_rows.sort(key=lambda item: item.get("date") or "")
        for rows in station_rows.values():
            rows.sort(key=lambda item: item.get("date") or "")
        return {"daily": daily_rows, "stations": station_rows}


    def _assistant_average(values):
        values = [float(value or 0) for value in values]
        return sum(values) / len(values) if values else 0.0


    def _assistant_station_trend_rows(trend_data, station=None):
        catalog = retail_station_catalog_cached()
        station_map = trend_data.get("stations") or {}
        target_keys = [str(station.get("shopKey"))] if station else list(station_map.keys())
        result = []
        for shop_key in target_keys:
            rows = station_map.get(str(shop_key)) or []
            if len(rows) < 2:
                continue
            volumes = [float(item.get("volume") or 0) for item in rows]
            amounts = [float(item.get("amount") or 0) for item in rows]
            zero_days = [str(item.get("date") or "") for item in rows if float(item.get("volume") or 0) <= 0]
            lowest_days = sorted(
                ((str(item.get("date") or ""), float(item.get("volume") or 0)) for item in rows),
                key=lambda item: item[1],
            )[:3]
            window = max(1, min(3, len(volumes) // 2))
            first_avg = _assistant_average(volumes[:window])
            last_avg = _assistant_average(volumes[-window:])
            delta = last_avg - first_avg
            pct = (delta / first_avg * 100) if first_avg else 0.0
            info = catalog.get(int(shop_key)) if str(shop_key).isdigit() else {}
            result.append({
                "shopKey": int(shop_key) if str(shop_key).isdigit() else shop_key,
                "name": (info or {}).get("name") or f"{shop_key}号站",
                "region": (info or {}).get("region") or "--",
                "firstAvg": first_avg,
                "lastAvg": last_avg,
                "delta": delta,
                "pct": pct,
                "totalVolume": sum(volumes),
                "totalAmount": sum(amounts),
                "zeroDays": zero_days,
                "lowestDays": lowest_days,
                "days": len(rows),
            })
        result.sort(key=lambda item: (item["pct"], item["delta"]))
        return result


    def _assistant_station_trend_anomaly_answer(label, trend_data, station=None):
        rows = _assistant_station_trend_rows(trend_data, station)
        if station:
            if not rows:
                return f"{label}{station.get('name')}没有足够的日趋势数据，暂时无法判断是否下滑。"
            item = rows[0]
            direction = "下降" if item["delta"] < 0 else "上升"
            return "\n".join([
                f"{label}{item['name']}日销量趋势：",
                f"前段日均 {_fmt_volume(item['firstAvg'])}，后段日均 {_fmt_volume(item['lastAvg'])}，{direction} {_fmt_volume(abs(item['delta']))}（{item['pct']:+.1f}%）。",
                f"区间合计 {_fmt_volume(item['totalVolume'])}，金额 {_fmt_amount(item['totalAmount'], '索姆')}，覆盖 {item['days']} 天。",
                _assistant_station_trend_reason_line(item),
            ])
        falling = [item for item in rows if item["delta"] < 0]
        if not falling:
            return f"{label}没有发现明显下滑站点。"
        lines = [f"{label}下滑最明显的站点（按前后段日均销量变化）："]
        for index, item in enumerate(falling[:8], start=1):
            lines.append(
                f"{index}. {item['name']}（{item['region']}）："
                f"{_fmt_volume(item['firstAvg'])} -> {_fmt_volume(item['lastAvg'])}，"
                f"下降 {_fmt_volume(abs(item['delta']))}（{item['pct']:.1f}%）"
            )
        lines.append("原因排查建议：先看是否存在零销量/缺日，再看停业维修、补传状态、油品结构、库存和价格变化。")
        return "\n".join(lines)


    def _assistant_station_trend_reason_line(item):
        zero_count = len(item.get("zeroDays") or [])
        lowest = [
            f"{day}:{_fmt_volume(value)}"
            for day, value in (item.get("lowestDays") or [])
            if day
        ]
        parts = ["原因排查："]
        if zero_count:
            parts.append(f"区间内有 {zero_count} 天零销量，优先核查停业、停电、补传或快照缺失。")
        elif lowest:
            parts.append("低销量日期：" + "、".join(lowest) + "。")
        parts.append("若数据完整，再继续看油品结构、库存、价格和当地车流变化。")
        return "".join(parts)


    def _assistant_station_trend_anomaly_tables(label, trend_data, station=None):
        rows = _assistant_station_trend_rows(trend_data, station)
        rows = rows if station else [item for item in rows if item["delta"] < 0]
        table_rows = [
            [
                item["name"],
                item["region"],
                _fmt_volume(item["firstAvg"]),
                _fmt_volume(item["lastAvg"]),
                _fmt_volume(item["delta"]),
                f"{item['pct']:+.1f}%",
                _fmt_volume(item["totalVolume"]),
                len(item.get("zeroDays") or []),
            ]
            for item in rows[:12]
        ]
        return [_assistant_table(
            f"{label}站点趋势变化",
            ["站点", "区域", "前段日均", "后段日均", "变化", "变化率", "区间合计", "零销量天数"],
            table_rows,
        )] if table_rows else []


    def _assistant_station_trend_anomaly_charts(label, trend_data, station=None):
        if station:
            station_rows = (trend_data.get("stations") or {}).get(str(station.get("shopKey"))) or []
            chart = _assistant_daily_mixed_chart(f"{label}{station.get('name') or station.get('shopKey')}日趋势", station_rows)
            return [chart] if chart else []
        rows = [item for item in _assistant_station_trend_rows(trend_data, station) if item["delta"] < 0][:12]
        chart = _assistant_chart(
            f"{label}站点销量下滑幅度",
            "bar",
            x_axis=[item["name"] for item in rows],
            series=[{"name": "销量变化", "type": "bar", "unit": "升", "data": [item["delta"] for item in rows]}],
            unit="升",
            summary="负值代表后段日均销量低于前段日均销量。",
        )
        return [chart] if chart else []


    def _assistant_station_data_check_answer(label, data, station):
        rows = [
            item for item in data.get("stations") or []
            if int(item.get("shopKey") or 0) == int(station.get("shopKey") or 0)
        ]
        station_no = station.get("shopKey") or station.get("stationNo") or ""
        station_name = station.get("name") or f"{station_no}号站"
        station_label = f"{station_name}（{station_no}号站）" if station_no and str(station_no) not in str(station_name) else station_name
        if not rows:
            return f"{label}{station_label}未在零售汇总快照中找到销售记录；如果刚补传数据，需要等待快照任务重新生成后再查。"
        item = rows[0]
        volume = float(item.get("volume") or 0)
        amount = float(item.get("amount") or 0)
        if volume or amount:
            return (
                f"{label}{station_label}已查到销售数据："
                f"销量 {_fmt_volume(volume)}，金额 {_fmt_amount(amount, '索姆')}。"
                "这说明该区间数据已经进入当前零售汇总口径。"
            )
        return f"{label}{station_label}有站点记录，但销量和金额为 0；需要继续核查该站是否停业、数据是否仍未补齐或该日确无销售。"


    def _assistant_station_data_check_tables(label, data, station):
        rows = [
            item for item in data.get("stations") or []
            if int(item.get("shopKey") or 0) == int(station.get("shopKey") or 0)
        ]
        if not rows:
            return []
        item = rows[0]
        return [_assistant_table(
            f"{label}{station.get('name')}数据检查",
            ["站点", "区域", "销量", "吨位", "金额", "状态"],
            [[
                item.get("name") or station.get("name") or "--",
                item.get("region") or station.get("region") or "--",
                _fmt_volume(item.get("volume")),
                _fmt_tons(item.get("tons")),
                _fmt_amount(item.get("amount"), "索姆"),
                item.get("status") or "--",
            ]],
        )]


    def _assistant_accounting_report_answer(label, target_day):
        try:
            from api_v2.system import _accounting_report_status
        except Exception:
            return "吉方日报状态模块暂不可用。", [], {}
        conn = get_connection()
        try:
            cur = conn.cursor(pymysql.cursors.DictCursor)
            data = _accounting_report_status(cur, target_day)
        finally:
            conn.close()
        missing_count = int(data.get("missingCount") or 0)
        generated_count = int(data.get("generatedStationCount") or 0)
        expected_count = int(data.get("expectedCount") or 0)
        lines = [f"{label}吉方财务交接日报状态：应生成 {expected_count} 站，已生成 {generated_count} 站，缺少 {missing_count} 站。"]
        if missing_count:
            names = [item.get("name") or item.get("stationNo") or "--" for item in data.get("missingStations") or []]
            lines.append("缺少站点：" + "、".join(names[:12]))
        else:
            lines.append("当前未发现缺少日报的应生成站点。")
        if data.get("outputDir"):
            lines.append(f"输出目录：{data.get('outputDir')}")
        table_rows = [
            [item.get("stationNo") or "--", item.get("name") or "--", item.get("region") or "--"]
            for item in data.get("missingStations") or []
        ]
        tables = [_assistant_table(f"{label}吉方日报缺站", ["站号", "站点", "区域"], table_rows)] if table_rows else []
        return "\n".join(lines), tables, data


    def _assistant_comparison_answer(label, monthly_data):
        comp = monthly_data.get("comparison") or {}
        mom = comp.get("mom") or {}
        yoy = comp.get("yoy") or {}
        lines = [f"{label}同比环比："]
        yoy_amt = yoy.get("sales")
        yoy_vol = yoy.get("volume")
        mom_amt = mom.get("sales")
        mom_vol = mom.get("volume")
        if yoy_amt is not None:
            lines.append(f"同比金额：{_fmt_pct(yoy_amt)}")
        if yoy_vol is not None:
            lines.append(f"同比销量：{_fmt_pct(yoy_vol)}")
        if mom_amt is not None:
            lines.append(f"环比金额：{_fmt_pct(mom_amt)}")
        if mom_vol is not None:
            lines.append(f"环比销量：{_fmt_pct(mom_vol)}")
        if yoy_amt is None and yoy_vol is None and mom_amt is None and mom_vol is None:
            lines.append("暂无对比数据（快照未生成对比字段，刷新后重试）。")
        return "\n".join(lines)


    def _assistant_comparison_tables(label, monthly_data):
        comp = monthly_data.get("comparison") or {}
        mom = comp.get("mom") or {}
        yoy = comp.get("yoy") or {}
        rows = []
        for v, label_k in [(yoy.get("volume"), "同比销量"), (yoy.get("sales"), "同比金额"),
                            (mom.get("volume"), "环比销量"), (mom.get("sales"), "环比金额")]:
            if v is not None:
                rows.append([label_k, _fmt_pct(v)])
        if not rows:
            return []
        return [_assistant_table(f"{label}同比环比", ["指标", "变化"], rows)]


    def _assistant_resource_share_answer(label, data):
        resources = [item for item in (data.get("byResource") or []) if float(item.get("volume") or 0) > 0]
        if not resources:
            return f"{label}无油品数据。"
        total_vol = sum(float(r.get("volume") or 0) for r in resources)
        total_amt = sum(float(r.get("amount") or 0) for r in resources)
        lines = [f"{label}油品结构与占比："]
        for item in sorted(resources, key=lambda r: float(r.get("amount") or 0), reverse=True):
            vol = float(item.get("volume") or 0)
            amt = float(item.get("amount") or 0)
            pct = vol / total_vol * 100 if total_vol else 0
            lines.append(f"  {item.get('name')}：{_fmt_volume(vol)}（占比 {pct:.1f}%），{_fmt_amount(amt, '索姆')}")
        lines.append(f"  合计：{_fmt_volume(total_vol)}，{_fmt_amount(total_amt, '索姆')}")
        return "\n".join(lines)


    def _execute_data_assistant_query_payload(payload):
        payload = payload if isinstance(payload, dict) else {}
        question = clean_str(payload.get("question"))
        if not question:
            return {"error": "question_required"}, 400
        if is_action_request(question):
            return enrich_assistant_payload(read_only_refusal(question, payload.get("language"))), 200

        conversation_context = payload.get("conversationContext")
        user_context = payload.get("userContext") if isinstance(payload.get("userContext"), dict) else None
        plan = build_assistant_query_plan(question, context=conversation_context)
        period = plan.get("period") or {}
        start = date.fromisoformat(period.get("start"))
        end = date.fromisoformat(period.get("endExclusive"))
        label = _assistant_period_display(period.get("label") or "", start, end)
        match_question = plan.get("normalizedQuestion") or question
        route = plan.get("route")
        resource = (plan.get("entities") or {}).get("resource")

        if plan.get("needsClarification"):
            clarification = plan.get("clarification") or {}
            answer = clarification.get("question") or "我需要先确认查询范围后再回答。"
            options = clarification.get("options") or []
            if options:
                answer = f"{answer}\n可选：{'、'.join(str(item) for item in options)}。"
            return _assistant_response(
                question,
                plan,
                periodLabel=label,
                queryStart=start.isoformat(),
                queryEndExclusive=end.isoformat(),
                answer=answer,
                tables=[],
                data={},
                needsClarification=True,
                clarification=clarification,
            )

        if route == "compound_query":
            return _assistant_response(
                question,
                plan,
                **_assistant_compound_payload(question, plan, user_context),
            )

        if route == "self_use_summary":
            data = _assistant_self_use_data_for_period(start, end, resource)
            data = _assistant_apply_station_scope(data, plan, user_context)
            return _assistant_response(
                question,
                plan,
                periodLabel=label,
                queryStart=start.isoformat(),
                queryEndExclusive=end.isoformat(),
                answer=_assistant_self_use_answer(label, data),
                tables=_assistant_self_use_tables(label, data),
                charts=_assistant_self_use_charts(label, data, match_question, visual=plan.get("visual")),
                data=data,
            )

        if route == "station_data_check" and (plan.get("entities") or {}).get("station"):
            station = plan["entities"]["station"]
            data = _assistant_retail_data_for_period(start, end, resource, plan=plan, user_context=user_context)
            return _assistant_response(
                question,
                plan,
                periodLabel=label,
                queryStart=start.isoformat(),
                queryEndExclusive=end.isoformat(),
                answer=_assistant_station_data_check_answer(label, data, station),
                tables=_assistant_station_data_check_tables(label, data, station),
                data=data,
            )

        if route == "station_trend_anomaly":
            station = (plan.get("entities") or {}).get("station")
            trend_data = _assistant_daily_trend_for_period(start, end)
            trend_data = _assistant_apply_station_scope(trend_data, plan, user_context)
            return _assistant_response(
                question,
                plan,
                periodLabel=label,
                queryStart=start.isoformat(),
                queryEndExclusive=end.isoformat(),
                answer=_assistant_station_trend_anomaly_answer(label, trend_data, station),
                tables=_assistant_station_trend_anomaly_tables(label, trend_data, station),
                charts=_assistant_station_trend_anomaly_charts(label, trend_data, station),
                data=trend_data,
            )

        if route == "retail_ranking":
            dimension = next((item for item in plan.get("dimensions") or [] if item in {"resource", "region", "station"}), "station")
            data = _assistant_retail_data_for_period(start, end, resource, plan=plan, user_context=user_context)
            metric = plan.get("metric") or "amount"
            ranking_rows = _assistant_retail_ranking_rows(data, dimension, metric)
            return _assistant_response(
                question,
                plan,
                periodLabel=label,
                queryStart=start.isoformat(),
                queryEndExclusive=end.isoformat(),
                answer=_assistant_retail_ranking_answer(label, data, dimension, metric, match_question),
                tables=_assistant_retail_ranking_tables(label, data, dimension, metric),
                charts=[_assistant_rank_chart(f"{label}零售{ {'resource': '油品', 'region': '区域', 'station': '站点'}.get(dimension, '站点') }{_assistant_metric_name(metric)}排行", ranking_rows, metric, name_key="region" if dimension == "region" else "name")],
                data=data,
            )

        if route == "accounting_report_status":
            answer, tables, data = _assistant_accounting_report_answer(label, start)
            return _assistant_response(
                question,
                plan,
                periodLabel=label,
                queryStart=start.isoformat(),
                queryEndExclusive=end.isoformat(),
                answer=answer,
                tables=tables,
                data=data,
            )

        if route == "station_status" or _assistant_has_any(match_question, ("停业", "站点状态", "异常站点", "暂停营业", "维修", "停电")):
            return _assistant_response(
                question,
                plan,
                periodLabel="当前",
                answer=_assistant_station_status_answer(),
                tables=_assistant_station_status_tables(),
            )
        if route == "procurement_summary":
            data = _assistant_procurement_data(start, end)
            return _assistant_response(
                question,
                plan,
                periodLabel=label,
                queryStart=start.isoformat(),
                queryEndExclusive=end.isoformat(),
                answer=_assistant_procurement_answer(label, data),
                tables=_assistant_procurement_tables(label, data),
                charts=_assistant_procurement_charts(label, data),
                data=data,
            )
        if route == "weighing_summary":
            data = _assistant_weighing_data(start, end, match_question)
            return _assistant_response(
                question,
                plan,
                periodLabel=label,
                queryStart=start.isoformat(),
                queryEndExclusive=end.isoformat(),
                answer=_assistant_weighing_answer(label, data),
                tables=_assistant_weighing_tables(label, data),
                charts=_assistant_weighing_charts(label, data),
                data=data,
            )
        if route in {"materials_summary", "material_transfer_summary"}:
            transfer_only = route == "material_transfer_summary"
            data = _assistant_materials_data(start, end)
            return _assistant_response(
                question,
                plan,
                periodLabel=label,
                queryStart=start.isoformat(),
                queryEndExclusive=end.isoformat(),
                answer=_assistant_materials_answer(label, data, transfer_only=transfer_only),
                tables=_assistant_materials_tables(label, data, transfer_only=transfer_only),
                charts=_assistant_materials_charts(label, data, transfer_only=transfer_only),
                data=data,
            )
        if route == "transfer_loss_summary":
            data = _assistant_transfer_loss_data(start, end)
            return _assistant_response(
                question,
                plan,
                periodLabel=label,
                queryStart=start.isoformat(),
                queryEndExclusive=end.isoformat(),
                answer=_assistant_transfer_loss_answer(label, data),
                tables=_assistant_transfer_loss_tables(label, data),
                charts=_assistant_transfer_loss_charts(label, data),
                data=data,
            )
        if route == "system_health_summary":
            data = _assistant_system_health_data()
            return _assistant_response(
                question,
                plan,
                periodLabel="当前",
                answer=_assistant_system_health_answer(data),
                tables=_assistant_system_health_tables(data),
                data=data,
            )
        if route == "knowledge_summary":
            data = _assistant_knowledge_data(match_question)
            return _assistant_response(
                question,
                plan,
                periodLabel="知识库",
                answer=_assistant_knowledge_answer(data),
                tables=_assistant_knowledge_tables(data),
                data=data,
            )
        if route == "audit_log_summary":
            data = _assistant_audit_log_data(start, end)
            return _assistant_response(
                question,
                plan,
                periodLabel=label,
                queryStart=start.isoformat(),
                queryEndExclusive=end.isoformat(),
                answer=_assistant_audit_log_answer(label, data),
                tables=_assistant_audit_log_tables(label, data),
                data=data,
            )
        if route == "alias_summary":
            data = _assistant_alias_data()
            answer = (
                "别名匹配摘要如下："
                f"油品别名 {data.get('productAliases', 0)} 条，客户/实体别名 {data.get('entityAliases', 0)} 条；"
                f"启用油品 {data.get('products', 0)} 个，客户 {data.get('customers', 0)} 个。"
            )
            return _assistant_response(
                question,
                plan,
                periodLabel="别名匹配",
                answer=answer,
                tables=[_assistant_table("别名匹配摘要", ["项目", "数量"], [
                    ["油品别名", data.get("productAliases", 0)],
                    ["客户/实体别名", data.get("entityAliases", 0)],
                    ["启用油品", data.get("products", 0)],
                    ["客户", data.get("customers", 0)],
                ])],
                data=data,
            )
        if route == "user_permission_summary":
            data = _assistant_user_permission_data()
            users = data.get("users") or {}
            answer = (
                "用户权限只读摘要如下："
                f"用户 {int(users.get('total') or 0)} 个，启用 {int(users.get('active') or 0)} 个；"
                f"角色 {data.get('roles', 0)} 个，部门 {data.get('departments', 0)} 个，权限点 {data.get('permissions', 0)} 个。"
                "我不会通过 AI 助手修改账号、角色或权限。"
            )
            return _assistant_response(
                question,
                plan,
                periodLabel="权限",
                answer=answer,
                tables=[_assistant_table("用户权限摘要", ["项目", "数量"], [
                    ["用户", int(users.get("total") or 0)],
                    ["启用用户", int(users.get("active") or 0)],
                    ["角色", data.get("roles", 0)],
                    ["部门", data.get("departments", 0)],
                    ["权限点", data.get("permissions", 0)],
                ])],
                data=data,
            )
        if route == "settings_summary":
            data = {
                "publicEntry": "http://198.51.100.10/",
                "backendListen": "127.0.0.1:8900",
                "updateFeed": "/download/latest.yml",
                "readOnly": True,
            }
            return _assistant_response(
                question,
                plan,
                periodLabel="系统设置",
                answer=_assistant_settings_answer(),
                tables=[_assistant_table("设置只读摘要", ["项目", "值"], [
                    ["生产公开入口", data["publicEntry"]],
                    ["后端监听", data["backendListen"]],
                    ["客户端更新源", data["updateFeed"]],
                    ["AI写操作", "禁止"],
                ])],
                data=data,
            )
        if route == "retail_trend" or _assistant_has_any(match_question, ("趋势", "每日", "日趋势", "日销量", "走势")):
            trend_data = _assistant_daily_trend_for_period(start, end)
            if trend_data.get("daily"):
                return _assistant_response(
                    question,
                    plan,
                    periodLabel=label,
                    queryStart=start.isoformat(),
                    queryEndExclusive=end.isoformat(),
                    answer=_assistant_daily_trend_answer(label, trend_data),
                    tables=_assistant_daily_trend_tables(label, trend_data),
                    charts=[_assistant_daily_mixed_chart(f"{label}零售每日趋势", trend_data.get("daily") or [])],
                    data=trend_data,
                )
        if route == "retail_comparison" and plan.get("comparison"):
            data = _assistant_retail_data_for_period(start, end, resource, plan=plan, user_context=user_context)
            comparison = _assistant_enrich_comparison_data(plan, user_context)
            return _assistant_response(
                question,
                plan,
                periodLabel=label,
                queryStart=start.isoformat(),
                queryEndExclusive=end.isoformat(),
                answer=_assistant_retail_comparison_answer(label, data, comparison),
                tables=_assistant_retail_comparison_tables(label, data, comparison),
                charts=_assistant_retail_comparison_charts(label, data, comparison),
                data={"current": data, "comparison": comparison},
            )
        if route == "retail_comparison" or _assistant_has_any(match_question, ("同比", "环比", "对比", "较去年", "较上月")):
            ym = f"{start.year}-{start.month:02d}"
            monthly_data = _read_retail_snapshot("monthly", ym)
            if monthly_data:
                return _assistant_response(
                    question,
                    plan,
                    periodLabel=ym,
                    answer=_assistant_comparison_answer(ym, monthly_data),
                    tables=_assistant_comparison_tables(ym, monthly_data),
                    data=monthly_data,
                )
        if route == "retail_resource_share" or _assistant_has_any(match_question, ("占比", "结构", "油品结构", "油品占比", "构成")):
            data = _assistant_retail_data_for_period(start, end, resource, plan=plan, user_context=user_context)
            return _assistant_response(
                question,
                plan,
                periodLabel=label,
                answer=_assistant_resource_share_answer(label, data),
                tables=_assistant_retail_tables(label, data, match_question),
                charts=_assistant_retail_charts(label, data, match_question, visual={"preferredType": "pie", "requested": True}),
                data=data,
            )
        if route == "retail_kpi" or _assistant_has_any(match_question, ("KPI", "kpi", "关键指标", "年度指标", "年度目标", "目标完成")):
            year = start.year
            kpi_data = _read_yearly_kpi_snapshot(year)
            if kpi_data:
                return _assistant_response(
                    question,
                    plan,
                    periodLabel=f"{year}年",
                    answer=_assistant_yearly_kpi_answer(f"{year}年", kpi_data),
                    tables=_assistant_yearly_kpi_tables(f"{year}年", kpi_data),
                    charts=_assistant_yearly_kpi_charts(f"{year}年", kpi_data),
                    data=kpi_data,
                )
        if route == "retail_price" or _assistant_has_any(match_question, ("价格", "单价", "售价", "零售价")):
            data = _assistant_retail_data_for_period(start, end, resource, plan=plan, user_context=user_context)
            resources = [item for item in (data.get("byResource") or []) if float(item.get("volume") or 0) > 0]
            lines = [f"{label}零售价格（均价 = 金额 ÷ 销量）："]
            price_rows = []
            for item in sorted(resources, key=lambda r: float(r.get("amount") or 0), reverse=True):
                vol = float(item.get("volume") or 0)
                amt = float(item.get("amount") or 0)
                price = amt / vol if vol else 0
                lines.append(f"  {item.get('name')}：{price:.2f} 索姆/升（{_fmt_volume(vol)}，{_fmt_amount(amt, '索姆')}）")
                price_rows.append([item.get("name", "--"), f"{price:.2f}", _fmt_volume(vol), _fmt_amount(amt, "索姆")])
            if not lines[1:]:
                lines.append("  该区间暂无销售数据，无法计算均价。")
            total_vol = sum(float(r.get("volume") or 0) for r in resources)
            total_amt = sum(float(r.get("amount") or 0) for r in resources)
            avg_price = total_amt / total_vol if total_vol else 0
            lines.append(f"  综合均价：{avg_price:.2f} 索姆/升")
            price_rows.append(["综合均价", f"{avg_price:.2f}", _fmt_volume(total_vol), _fmt_amount(total_amt, "索姆")])
            return _assistant_response(
                question,
                plan,
                periodLabel=label,
                answer="\n".join(lines),
                tables=[_assistant_table(f"{label}零售价格", ["油品", "均价(索姆/升)", "销量", "金额"], price_rows)],
                charts=[_assistant_chart(
                    f"{label}零售均价",
                    "bar",
                    x_axis=[row[0] for row in price_rows if row[0] != "综合均价"],
                    series=[{"name": "均价", "type": "bar", "unit": "索姆/升", "data": [row[1] for row in price_rows if row[0] != "综合均价"]}],
                    unit="索姆/升",
                    summary="均价按金额除以销量计算。",
                )],
                data=data,
            )
        if route == "market_summary" or _assistant_has_any(match_question, ("市场", "行情", "观察", "原油", "成品油", "俄罗斯", "乌兹别克", "国际")):
            conn = get_connection()
            try:
                cur = conn.cursor()
                answer, tables = _assistant_market_watch_answer(cur)
            finally:
                conn.close()
            return _assistant_response(question, plan, periodLabel="市场观察", answer=answer, tables=tables)
        if route == "contract_summary" or "合同" in match_question:
            conn = get_connection()
            try:
                cur = conn.cursor()
                answer = _assistant_contract_answer(cur, match_question)
                tables = _assistant_contract_tables(cur, match_question)
            finally:
                conn.close()
            return _assistant_response(question, plan, periodLabel="合同", answer=answer, tables=tables)
        if route == "combined_summary" or _assistant_has_any(match_question, ("合并", "总销售", "整体", "公司", "总量", "总金额")):
            year, month = _assistant_year_month_from_period(start, end)
            conn = get_connection()
            try:
                cur = conn.cursor()
                data = _combined_data_from_snapshots(cur, year, month, start, end)
            finally:
                conn.close()
            return _assistant_response(
                question,
                plan,
                periodLabel=label,
                queryStart=start.isoformat(),
                queryEndExclusive=end.isoformat(),
                answer=_assistant_combined_answer(label, data),
                tables=_assistant_combined_tables(label, data),
                charts=_assistant_combined_charts(label, data),
                data=data,
            )
        if route == "wholesale_review" or _assistant_has_any(match_question, ("审核", "复核", "异常", "提示", "未匹配")):
            year, month = _assistant_year_month_from_period(start, end)
            conn = get_connection()
            try:
                cur = conn.cursor()
                data = summarize_wholesale_page(cur, year, month)
                answer = _assistant_review_answer(label, data)
            finally:
                conn.close()
            return _assistant_response(question, plan, periodLabel=label, answer=answer, tables=_assistant_review_tables(label, data), data=data)
        if route == "wholesale_summary" or any(word in match_question for word in ("批发", "外运", "出口", "本地", "客户")):
            year, month = _assistant_year_month_from_period(start, end)
            conn = get_connection()
            try:
                cur = conn.cursor()
                data = summarize_wholesale_page(cur, year, month)
            finally:
                conn.close()
            return _assistant_response(
                question,
                plan,
                periodLabel=label,
                queryStart=start.isoformat(),
                queryEndExclusive=end.isoformat(),
                answer=_format_assistant_wholesale_answer(label, data, match_question),
                tables=_assistant_wholesale_tables(label, data, match_question),
                charts=_assistant_wholesale_charts(label, data, match_question),
                data=data,
            )
        if route == "exchange_summary" or "汇率" in match_question or "美元" in match_question or "人民币" in match_question:
            conn = get_connection()
            try:
                cur = conn.cursor()
                answer = _assistant_exchange_answer(cur, label, start, end)
                tables = _assistant_exchange_tables(cur, label, start, end)
                charts = _assistant_exchange_charts(cur, label, start, end)
            finally:
                conn.close()
            return _assistant_response(question, plan, periodLabel=label, answer=answer, tables=tables, charts=charts)
        if route == "density_summary" or "密度" in match_question:
            conn = get_connection()
            try:
                cur = conn.cursor()
                answer = _assistant_density_answer(cur)
                tables = _assistant_density_tables(cur)
                charts = _assistant_density_charts(cur)
            finally:
                conn.close()
            return _assistant_response(question, plan, periodLabel="当前", answer=answer, tables=tables, charts=charts)
        if route == "retail_inventory" or _assistant_has_any(match_question, ("库存", "期末")):
            inventory_data = _retail_inventory_for_period(start, end)
            if inventory_data:
                return _assistant_response(
                    question,
                    plan,
                    periodLabel=label,
                    queryStart=start.isoformat(),
                    queryEndExclusive=end.isoformat(),
                    answer=_assistant_retail_inventory_answer(label, inventory_data),
                    tables=_assistant_retail_inventory_tables(label, inventory_data),
                    charts=_assistant_retail_inventory_charts(label, inventory_data),
                    data=inventory_data,
                )
        data = _assistant_retail_data_for_period(start, end, resource, plan=plan, user_context=user_context)
        if not data:
            return {"error": "retail_query_failed", "detail": "empty_retail_response"}, 502
        data["question"] = match_question
        return _assistant_response(
            question,
            plan,
            periodLabel=label,
            queryStart=start.isoformat(),
            queryEndExclusive=end.isoformat(),
            answer=_format_assistant_retail_answer(label, data),
            tables=_assistant_retail_tables(label, data, match_question),
            charts=_assistant_retail_charts(label, data, match_question, visual=plan.get("visual")),
            data=data,
        )


    global _ASSISTANT_QUERY_EXECUTOR
    _ASSISTANT_QUERY_EXECUTOR = _execute_data_assistant_query_payload


    @app.route("/api/data-assistant/query", methods=["POST"])
    def api_data_assistant_query():
        result, status = _execute_data_assistant_query_payload(request.get_json(force=True) or {})
        response = jsonify(result)
        response.headers["Deprecation"] = "true"
        response.headers["Link"] = '</api/v2/assistant/query>; rel="successor-version"'
        return response, status



