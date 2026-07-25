from __future__ import annotations

from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from urllib.parse import quote
import html
import os
import zipfile

import pandas as pd
from flask import jsonify, request, send_file
from openpyxl.styles import Alignment, Border, Font, Side

from db import get_connection
from summary_service import period_bounds, summarize_combined
from wholesale_service import summarize_wholesale_page


EXPORT_DIR = Path(__file__).resolve().parent / "exports"


CLASS_LABELS = {
    "export_wholesale": "出口批发",
    "local_wholesale": "本地批发",
    "unknown": "待识别批发",
    "transfer": "外运",
}


def _fmt(value, digits=2):
    try:
        return f"{_round_half_up(value, digits):,.{int(digits or 0)}f}"
    except (TypeError, ValueError):
        return "--"


def _round_half_up(value, digits=0):
    digits = int(digits or 0)
    try:
        quant = Decimal("1").scaleb(-digits)
        return float(Decimal(str(value or 0)).quantize(quant, rounding=ROUND_HALF_UP))
    except (InvalidOperation, TypeError, ValueError):
        return 0.0


def _period_label(year, month=None):
    return f"{year}年{int(month)}月份" if month else f"{year}年年度"


def _parse_year_month(year_month):
    year, month = year_month.split("-", 1)
    return int(year), int(month)


def _sum(rows, key):
    return sum(float((row or {}).get(key) or 0) for row in rows)


def _display_width(value):
    text = str(value or "")
    width = 0
    for char in text:
        width += 2 if ord(char) > 127 else 1
    return width


def _table(headers, rows):
    head = "".join(f"<th>{html.escape(str(item))}</th>" for item in headers)
    body = "".join(
        "<tr>" + "".join(f"<td>{html.escape(str(cell))}</td>" for cell in row) + "</tr>"
        for row in rows
    )
    empty = f'<tr><td colspan="{len(headers)}">暂无数据</td></tr>'
    return f"<table class=\"data-table\"><thead><tr>{head}</tr></thead><tbody>{body or empty}</tbody></table>"


def _report_shell(title, subtitle, cards, sections):
    card_html = "".join(
        f"<div class=\"kpi-card\"><div class=\"kpi-label\">{html.escape(label)}</div><div class=\"kpi-value\">{html.escape(value)}</div><div class=\"kpi-note\">{html.escape(note)}</div></div>"
        for label, value, note in cards
    )
    section_html = "".join(
        f"<div class=\"card\"><div class=\"card-title\">{html.escape(heading)}</div>{content}</div>"
        for heading, content in sections
    )
    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M")
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{html.escape(title)}</title>
  <style>
    * {{ margin:0; padding:0; box-sizing:border-box; }}
    body {{ font-family:'Microsoft YaHei', Arial, sans-serif; background:#f0f2f5; color:#333; }}
    .header {{ background:linear-gradient(135deg,#1a3a5c 0%,#2e6da4 60%,#1e8bc3 100%); color:#fff; padding:32px 0 24px; }}
    .header-inner {{ max-width:1400px; margin:0 auto; padding:0 20px; text-align:center; }}
    .header h1 {{ font-size:26px; font-weight:700; letter-spacing:2px; }}
    .header p {{ margin-top:8px; font-size:14px; opacity:.86; }}
    .container {{ max-width:1400px; margin:24px auto; padding:0 20px; }}
    .kpi-grid {{ display:grid; grid-template-columns:repeat(4,1fr); gap:16px; margin-bottom:24px; }}
    .kpi-card {{ background:#fff; border-radius:12px; padding:20px 22px; box-shadow:0 2px 8px rgba(0,0,0,.07); border-left:4px solid #2e6da4; }}
    .kpi-label {{ color:#667085; font-size:13px; margin-bottom:8px; }}
    .kpi-value {{ color:#1a3a5c; font-size:24px; font-weight:700; }}
    .kpi-note {{ color:#7b8794; font-size:12px; margin-top:6px; line-height:1.5; }}
    .card {{ background:#fff; border-radius:12px; padding:22px; margin-bottom:24px; box-shadow:0 2px 8px rgba(0,0,0,.07); }}
    .card-title {{ font-size:18px; color:#1a3a5c; font-weight:700; margin-bottom:16px; border-bottom:2px solid #e6eef7; padding-bottom:10px; }}
    .analysis-note, .note {{ background:#f0f6ff; border-left:3px solid #2e6da4; padding:14px 18px; border-radius:0 8px 8px 0; font-size:13px; line-height:1.9; }}
    .data-table {{ width:100%; border-collapse:collapse; font-size:13px; }}
    .data-table th, .data-table td {{ border:1px solid #d9e2ec; padding:9px 10px; text-align:right; }}
    .data-table th:first-child, .data-table td:first-child {{ text-align:left; }}
    .data-table th {{ background:#1a3a5c; color:#fff; font-weight:700; }}
    .data-table tr:last-child td {{ font-weight:700; background:#f8fbff; }}
    footer {{ color:#7b8794; font-size:12px; text-align:center; padding:10px 0 26px; }}
    @media(max-width:900px) {{ .kpi-grid {{ grid-template-columns:repeat(2,1fr); }} }}
    @media(max-width:600px) {{ .kpi-grid {{ grid-template-columns:1fr; }} .card {{ padding:16px; }} }}
  </style>
</head>
<body>
  <div class="header"><div class="header-inner"><h1>{html.escape(title)}</h1><p>{html.escape(subtitle)}</p></div></div>
  <main class="container"><div class="kpi-grid">{card_html}</div>{section_html}</main>
  <footer>生成时间：{generated_at}</footer>
</body>
</html>"""


def _docx_xml(title, subtitle, sections):
    def esc(text):
        return html.escape(str(text), quote=False)

    def para(text="", bold=False, size=21, align=None, color="333333", spacing=120):
        ppr = f"<w:pPr><w:spacing w:after=\"{spacing}\"/>"
        if align:
            ppr += f"<w:jc w:val=\"{align}\"/>"
        ppr += "</w:pPr>"
        rpr = (
            f"<w:rPr><w:rFonts w:ascii=\"Microsoft YaHei\" w:hAnsi=\"Microsoft YaHei\" "
            f"w:eastAsia=\"Microsoft YaHei\"/><w:sz w:val=\"{size}\"/><w:color w:val=\"{color}\"/>"
        )
        if bold:
            rpr += "<w:b/>"
        rpr += "</w:rPr>"
        parts = str(text).split("\n")
        text_xml = "<w:br/>".join(f"<w:t>{esc(part)}</w:t>" for part in parts)
        return (
            f"<w:p>{ppr}<w:r>{rpr}{text_xml}</w:r></w:p>"
        )

    def heading_para(text):
        return para(text, bold=True, size=28, color="1A3A5C", spacing=160)

    def cell(text, bold=False, shade=None, align="center"):
        tcpr = "<w:tcPr><w:tcW w:w=\"0\" w:type=\"auto\"/><w:vAlign w:val=\"center\"/>"
        if shade:
            tcpr += f"<w:shd w:fill=\"{shade}\"/>"
        tcpr += "</w:tcPr>"
        color = "FFFFFF" if shade else "333333"
        return f"<w:tc>{tcpr}{para(text, bold=bold, size=18, color=color, spacing=0, align=align)}</w:tc>"

    def table(rows):
        if not rows:
            rows = [["暂无数据"]]
        headers = rows[0]
        body_rows = rows[1:] or [["--"] * len(headers)]
        body = [
            "<w:tbl><w:tblPr><w:tblStyle w:val=\"TableGrid\"/>"
            "<w:tblW w:w=\"9026\" w:type=\"dxa\"/><w:jc w:val=\"center\"/>"
            "<w:tblLayout w:type=\"autofit\"/>"
            "<w:tblBorders>"
            "<w:top w:val=\"single\" w:sz=\"4\" w:color=\"D9E2EC\"/>"
            "<w:left w:val=\"single\" w:sz=\"4\" w:color=\"D9E2EC\"/>"
            "<w:bottom w:val=\"single\" w:sz=\"4\" w:color=\"D9E2EC\"/>"
            "<w:right w:val=\"single\" w:sz=\"4\" w:color=\"D9E2EC\"/>"
            "<w:insideH w:val=\"single\" w:sz=\"4\" w:color=\"D9E2EC\"/>"
            "<w:insideV w:val=\"single\" w:sz=\"4\" w:color=\"D9E2EC\"/>"
            "</w:tblBorders></w:tblPr>"
        ]
        body.append("<w:tr>" + "".join(cell(header, bold=True, shade="1A3A5C") for header in headers) + "</w:tr>")
        for row in body_rows:
            is_total = bool(row) and str(row[0]) == "合计"
            body.append("<w:tr>" + "".join(cell(value, bold=is_total) for value in row) + "</w:tr>")
        body.append("</w:tbl>")
        return "".join(body) + para("", spacing=100)

    body = [
        para("Example石油公司销售板块", bold=True, size=36, align="center", color="1A3A5C", spacing=120),
        para(title, bold=True, size=32, align="center", color="1A3A5C", spacing=120),
        para(subtitle, size=21, align="center", color="667085", spacing=220),
    ]
    for section_heading, rows in sections:
        body.append(heading_para(section_heading))
        body.append(table(rows))
    return f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>{''.join(body)}<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440"/></w:sectPr></w:body>
</w:document>"""


def _write_docx(path, title, subtitle, sections):
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>""")
        zf.writestr("_rels/.rels", """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>""")
        zf.writestr("word/document.xml", _docx_xml(title, subtitle, sections))


def _write_html_docx(base_name, html_content, docx_title, docx_subtitle, docx_sections):
    EXPORT_DIR.mkdir(exist_ok=True)
    html_name = f"{base_name}.html"
    docx_name = f"{base_name}.docx"
    html_path = EXPORT_DIR / html_name
    docx_path = EXPORT_DIR / docx_name
    html_path.write_text(html_content, encoding="utf-8")
    _write_docx(docx_path, docx_title, docx_subtitle, docx_sections)
    return {
        "htmlFile": html_name,
        "docxFile": docx_name,
        "htmlViewUrl": f"/exports/{quote(html_name)}",
        "htmlDownloadUrl": f"/api/export/analysis/sales/download/{quote(html_name)}",
        "docxDownloadUrl": f"/api/export/analysis/sales/download/{quote(docx_name)}",
    }


def _wholesale_report_dataset(cur, year, month=None):
    data = summarize_wholesale_page(cur, year, month)
    cards = data.get("cards") or {}
    sales_tons = sum(float((cards.get(key) or {}).get("tons") or 0) for key in ("export_wholesale", "local_wholesale", "unknown"))
    sales_amount = sum(float((cards.get(key) or {}).get("amountCny") or 0) for key in ("export_wholesale", "local_wholesale"))
    transfer_tons = float((cards.get("transfer") or {}).get("tons") or 0)
    return data, sales_tons, sales_amount, transfer_tons


def _make_wholesale_report(year, month=None):
    conn = get_connection()
    try:
        cur = conn.cursor()
        data, sales_tons, sales_amount, transfer_tons = _wholesale_report_dataset(cur, year, month)
    finally:
        conn.close()

    period = _period_label(year, month)
    title = f"批发{period}分析报告"
    cards = [
        ("批发销售量", f"{_fmt(sales_tons, 3)} 吨", "出口批发 + 本地批发 + 待识别批发"),
        ("批发销售金额", f"{_fmt(sales_amount, 2)} 元", "仅统计已匹配合同金额"),
        ("外运数量", f"{_fmt(transfer_tons, 3)} 吨", "外运只统计数量"),
        ("待复核", f"{_fmt(sum((data.get('review', {}).get('status', {}).get(k, {}).get('rows') or 0) for k in ('needs_review', 'unmatched')), 0)} 条", "合同、客户、油品或重复数据审核"),
    ]
    category_rows = []
    for key in ("export_wholesale", "local_wholesale", "unknown", "transfer"):
        item = (data.get("cards") or {}).get(key) or {}
        category_rows.append([
            CLASS_LABELS[key],
            _fmt(item.get("tons"), 3),
            "--" if key in ("unknown", "transfer") else _fmt(item.get("amountCny"), 2),
            item.get("entities") or 0,
            item.get("rows") or 0,
        ])
    category_rows.append(["合计", _fmt(sales_tons + transfer_tons, 3), _fmt(sales_amount, 2), "--", "--"])

    product_rows = []
    for key, rows in (data.get("productRankings") or {}).items():
        for item in (rows or [])[:12]:
            product_rows.append([CLASS_LABELS.get(key, key), item.get("name") or "--", _fmt(item.get("tons"), 3), _fmt(item.get("amountCny"), 2)])
    product_rows.append(["合计", "--", _fmt(_sum(product_rows_as_dict := [{"tons": row[2].replace(",", "") if row[2] != "--" else 0, "amount": row[3].replace(",", "") if row[3] != "--" else 0} for row in product_rows], "tons"), 3), _fmt(_sum(product_rows_as_dict, "amount"), 2)])

    customers = []
    for key, rows in (data.get("customerRankings") or {}).items():
        for item in (rows or [])[:10]:
            customers.append([CLASS_LABELS.get(key, key), item.get("name") or "--", _fmt(item.get("tons"), 3), _fmt(item.get("amountCny"), 2), item.get("rows") or 0])

    daily = data.get("dailyTrend") or []
    daily_rows = [[item.get("date"), _fmt(item.get("exportTons"), 3), _fmt(item.get("localTons"), 3), _fmt(item.get("pendingTons"), 3), _fmt(item.get("transferTons"), 3), _fmt(item.get("amountCny"), 2)] for item in daily]
    daily_rows.append(["合计", _fmt(_sum(daily, "exportTons"), 3), _fmt(_sum(daily, "localTons"), 3), _fmt(_sum(daily, "pendingTons"), 3), _fmt(_sum(daily, "transferTons"), 3), _fmt(_sum(daily, "amountCny"), 2)])

    note = (
        f"<div class=\"note\">{period}批发销售量合计 {_fmt(sales_tons, 3)} 吨，"
        f"已计算合同金额 {_fmt(sales_amount, 2)} 元；外运 {_fmt(transfer_tons, 3)} 吨，仅做数量统计。"
        "待识别批发数据先计入数量，不提前计算金额。</div>"
    )
    sections = [
        ("一、经营概览", note),
        ("二、批发分类统计", _table(["分类", "吨数", "金额(元)", "客户/去向", "记录数"], category_rows)),
        ("三、油品结构", _table(["分类", "油品", "吨数", "金额(元)"], product_rows)),
        ("四、客户排行", _table(["分类", "客户", "吨数", "金额(元)", "记录数"], customers)),
        ("五、每日趋势", _table(["日期", "出口", "本地", "待审核", "外运", "金额(元)"], daily_rows)),
    ]
    docx_sections = [
        ("批发分类统计", [["分类", "吨数", "金额(元)", "客户/去向", "记录数"], *category_rows]),
        ("油品结构", [["分类", "油品", "吨数", "金额(元)"], *product_rows]),
        ("每日趋势", [["日期", "出口", "本地", "待审核", "外运", "金额(元)"], *daily_rows]),
    ]
    source_note = "批发数据按统一口径汇总：2026-06-01 前取称重生成明细，之后取上传或手工录入明细，并结合合同匹配结果。"
    html_content = _report_shell(title, source_note, cards, sections)
    return title, _write_html_docx(title, html_content, title, source_note, docx_sections)


def _make_sales_report(year, month=None):
    conn = get_connection()
    try:
        cur = conn.cursor()
        data = summarize_combined(cur, year, month)
    finally:
        conn.close()

    period = _period_label(year, month)
    title = f"销售{period}分析报告"
    retail = data.get("retail") or {}
    wholesale = data.get("wholesale") or {}
    combined = data.get("combined") or {}
    cards = [
        ("合并销售量", f"{_fmt(combined.get('salesTons'), 3)} 吨", "零售 + 批发销售"),
        ("合并销售金额", f"{_fmt(combined.get('salesAmountCny'), 2)} 元", "人民币口径"),
        ("零售销售量", f"{_fmt(retail.get('tons'), 3)} 吨", f"{_fmt(retail.get('liters'), 2)} 升"),
        ("批发销售量", f"{_fmt(wholesale.get('salesTons'), 3)} 吨", f"外运 {_fmt(wholesale.get('transferTons'), 3)} 吨另列"),
    ]
    overview_rows = [
        ["零售", _fmt(retail.get("tons"), 3), _fmt(retail.get("amountCny"), 2), "按零售快照/历史汇率/密度换算"],
        ["批发", _fmt(wholesale.get("salesTons"), 3), _fmt(wholesale.get("salesAmountCny"), 2), "出口批发 + 本地批发 + 待审核数量"],
        ["合并", _fmt(combined.get("salesTons"), 3), _fmt(combined.get("salesAmountCny"), 2), "销售口径合计"],
        ["外运", _fmt(wholesale.get("transferTons"), 3), "--", "只统计数量，不计算金额"],
    ]
    product_rows = []
    for item in (retail.get("byProduct") or [])[:10]:
        product_rows.append(["零售", item.get("name") or "--", _fmt(item.get("tons"), 3), _fmt(item.get("amountCny") or item.get("amount"), 2)])
    for item in (wholesale.get("byProduct") or [])[:12]:
        product_rows.append(["批发", item.get("name") or "--", _fmt(item.get("tons"), 3), _fmt(item.get("amountCny"), 2)])
    product_rows.append(["合计", "--", _fmt(combined.get("salesTons"), 3), _fmt(combined.get("salesAmountCny"), 2)])
    note = (
        f"<div class=\"note\">{period}合并销售量 {_fmt(combined.get('salesTons'), 3)} 吨，"
        f"销售金额 {_fmt(combined.get('salesAmountCny'), 2)} 元。外运 {_fmt(wholesale.get('transferTons'), 3)} 吨单独列示，不计入销售金额。</div>"
    )
    sections = [
        ("一、综合经营概览", note),
        ("二、零售与批发对比", _table(["板块", "销售量(吨)", "金额(元)", "口径"], overview_rows)),
        ("三、油品结构", _table(["板块", "油品", "吨数", "金额(元)"], product_rows)),
    ]
    docx_sections = [
        ("零售与批发对比", [["板块", "销售量(吨)", "金额(元)", "口径"], *overview_rows]),
        ("油品结构", [["板块", "油品", "吨数", "金额(元)"], *product_rows]),
    ]
    html_content = _report_shell(title, "涵盖零售、批发与合并统计口径", cards, sections)
    return title, _write_html_docx(title, html_content, title, "涵盖零售、批发与合并统计口径", docx_sections)


def _exchange_export(year, month=None):
    start, end = period_bounds(year, month)
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT rate_date, usd_kgs, cny_kgs, source, is_manual, fetched_at, note
            FROM exchange_rates
            WHERE rate_date >= %s AND rate_date < %s
            ORDER BY rate_date
            """,
            (start, end),
        )
        rows = []
        for rate_date, usd, cny, source, is_manual, fetched_at, note in cur.fetchall():
            cny_value = float(cny) if cny is not None else None
            rows.append({
                "日期": rate_date.isoformat() if rate_date else "",
                "美元/索姆(USD/KGS)": _round_half_up(usd, 4) if usd is not None else None,
                "人民币/索姆(CNY/KGS)": _round_half_up(cny_value, 4) if cny_value is not None else None,
                "索姆/人民币(KGS/CNY)": _round_half_up(1 / cny_value, 6) if cny_value else None,
                "数据来源": "人工录入" if is_manual else (source or "NBKR"),
                "获取时间": fetched_at.isoformat(sep=" ") if fetched_at else "",
            })
    finally:
        conn.close()
    EXPORT_DIR.mkdir(exist_ok=True)
    name = f"汇率快照{_period_label(year, month)}.xlsx"
    path = EXPORT_DIR / name
    df = pd.DataFrame(rows, columns=["日期", "美元/索姆(USD/KGS)", "人民币/索姆(CNY/KGS)", "索姆/人民币(KGS/CNY)", "数据来源", "获取时间"])
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        writer.book._fonts[0] = Font(name='Times New Roman', size=11)
        df.to_excel(writer, index=False, sheet_name="每日汇率")
        worksheet = writer.sheets["每日汇率"]
        black_side = Side(style="thin", color="000000")
        black_border = Border(left=black_side, right=black_side, top=black_side, bottom=black_side)
        center = Alignment(horizontal="center", vertical="center", wrap_text=False)
        header_font = Font(name='Times New Roman', bold=True)
        for row in worksheet.iter_rows(min_row=1, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column):
            for cell in row:
                cell.alignment = center
                cell.border = black_border
                if cell.row == 1:
                    cell.font = header_font
                elif cell.column in (2, 3):
                    cell.number_format = "#,##0.0000"
                elif cell.column == 4:
                    cell.number_format = "#,##0.000000"
        for col in worksheet.columns:
            width = max(_display_width(cell.value) for cell in col) + 3
            worksheet.column_dimensions[col[0].column_letter].width = min(max(width, 10), 60)
    return name, path


def register_analysis_report_routes(app):
    @app.route("/api/export/analysis/wholesale/monthly/<year_month>")
    def api_export_wholesale_monthly(year_month):
        year, month = _parse_year_month(year_month)
        title, files = _make_wholesale_report(year, month)
        return jsonify({"status": "ok", "kind": "monthly", "title": title, **files})

    @app.route("/api/export/analysis/wholesale/yearly/<int:year>")
    def api_export_wholesale_yearly(year):
        title, files = _make_wholesale_report(year)
        return jsonify({"status": "ok", "kind": "yearly", "title": title, **files})

    @app.route("/api/export/analysis/sales/monthly/<year_month>")
    def api_export_sales_monthly(year_month):
        year, month = _parse_year_month(year_month)
        title, files = _make_sales_report(year, month)
        return jsonify({"status": "ok", "kind": "monthly", "title": title, **files})

    @app.route("/api/export/analysis/sales/yearly/<int:year>")
    def api_export_sales_yearly(year):
        title, files = _make_sales_report(year)
        return jsonify({"status": "ok", "kind": "yearly", "title": title, **files})

    @app.route("/api/export/exchange-rates")
    def api_export_exchange_rates_excel():
        year = int(request.args.get("year", str(datetime.today().year)))
        month_raw = request.args.get("month")
        month = int(month_raw) if month_raw else None
        filename, path = _exchange_export(year, month)
        return send_file(path, as_attachment=True, download_name=filename)

    @app.route("/api/export/analysis/sales/download/<path:filename>")
    def api_download_sales_analysis(filename):
        safe_name = os.path.basename(filename)
        path = EXPORT_DIR / safe_name
        if not path.exists():
            return jsonify({"error": "report not found"}), 404
        return send_file(path, as_attachment=True, download_name=safe_name)
