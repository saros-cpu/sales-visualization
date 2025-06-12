from __future__ import annotations

from datetime import date, datetime, timedelta
from email.utils import parsedate_to_datetime
from functools import wraps
import html
import importlib.util
import json
import logging
import os
import re
import subprocess
import sys
import tempfile
import threading
import time
import xml.etree.ElementTree as ET
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

import pandas as pd
from flask import Flask, Response, jsonify, redirect, request, send_from_directory, session
from pymysql.err import IntegrityError
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

from calculations import recalculate_wholesale_batch
from analysis_reports import register_analysis_report_routes
from config import LOGIN_REMEMBER_DAYS, NBKR_DAILY_XML_URL, SECRET_KEY
from db import get_connection
from import_helpers import (
    get_or_create_customer,
    load_entity_lookup,
    load_product_lookup,
    lookup_entity,
    lookup_product,
)
from scripts.import_contracts import import_contracts
from scripts.import_wholesale_excel import SHEET_DETAIL, import_wholesale
from summary_service import available_periods, period_bounds, summarize_combined, summarize_wholesale
from text_utils import clean_str, normalize_station_name, normalize_text, parse_date, parse_decimal
from wholesale_service import exchange_rates, summarize_wholesale_page
from retail_internal import open_oilstation


app = Flask(__name__, static_folder=".", static_url_path="")
register_analysis_report_routes(app)
app.secret_key = SECRET_KEY
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.getenv("SALES_COOKIE_SECURE", "").lower() in {"1", "true", "yes"},
    PERMANENT_SESSION_LIFETIME=timedelta(days=LOGIN_REMEMBER_DAYS),
)
APP_BUILD = "2026-05-11-retail-resource-report-v2"
PROJECT_ROOT = Path(__file__).resolve().parent
OILSTATION_ROOT = PROJECT_ROOT.parent / "OilStation"
SHARED_ASSETS_ROOT = PROJECT_ROOT.parent / "assets"
FRONTEND_DIST = PROJECT_ROOT.parent / "frontend" / "dist"
UPLOAD_ROOT = PROJECT_ROOT / "uploads"
REFRESH_CACHE_STATUS = {
    "running": False,
    "startedAt": None,
    "finishedAt": None,
    "status": "idle",
    "message": "尚未执行",
    "items": [],
}
REFRESH_CACHE_LOCK = threading.Lock()

RETAIL_API_PREFIXES = (
    "years",
    "months",
    "kpi",
    "dashboard",
    "resources",
    "monthly",
    "station",
    "station-map-details",
    "stations",
    "daily",
    "regions",
    "export",
    "integration",
    "assistant",
)

PUBLIC_PATHS = {
    "/login",
    "/api/auth/status",
    "/api/auth/login",
    "/api/health",
}

PUBLIC_PAGE_PREFIXES = ()

PUBLIC_API_PATHS = set()

PUBLIC_API_PREFIXES = ()


def ensure_auth_schema(cur):
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS system_users (
            id BIGINT PRIMARY KEY AUTO_INCREMENT,
            username VARCHAR(64) NOT NULL UNIQUE,
            password_hash VARCHAR(255) NOT NULL,
            display_name VARCHAR(128) NOT NULL,
            role ENUM('admin', 'user') NOT NULL DEFAULT 'user',
            is_active BOOLEAN NOT NULL DEFAULT TRUE,
            created_by BIGINT NULL,
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP NULL ON UPDATE CURRENT_TIMESTAMP,
            last_login_at TIMESTAMP NULL,
            KEY idx_system_users_role (role, is_active)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
        """
    )


def ensure_operation_log_schema(cur):
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS operation_logs (
            id BIGINT PRIMARY KEY AUTO_INCREMENT,
            user_id BIGINT NULL,
            username VARCHAR(64) NULL,
            action VARCHAR(128) NOT NULL,
            target_type VARCHAR(128) NULL,
            target_id VARCHAR(128) NULL,
            detail JSON NULL,
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            KEY idx_operation_logs_created (created_at),
            KEY idx_operation_logs_action (action)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
        """
    )


def ensure_market_watch_schema(cur):
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS market_watch_items (
            id BIGINT PRIMARY KEY AUTO_INCREMENT,
            market_region ENUM('Russia', 'Uzbekistan', 'Europe_US', 'Global') NOT NULL DEFAULT 'Global',
            market_type ENUM('crude', 'refined', 'policy', 'logistics', 'other') NOT NULL DEFAULT 'other',
            title VARCHAR(255) NOT NULL,
            summary TEXT NOT NULL,
            price_text VARCHAR(255) NULL,
            source_name VARCHAR(128) NULL,
            source_url VARCHAR(500) NULL,
            observed_at DATE NOT NULL,
            priority INT NOT NULL DEFAULT 100,
            is_active BOOLEAN NOT NULL DEFAULT TRUE,
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP NULL ON UPDATE CURRENT_TIMESTAMP,
            KEY idx_market_watch_active (is_active, observed_at, priority),
            KEY idx_market_watch_region (market_region, observed_at)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
        """
    )
    cur.execute("SHOW COLUMNS FROM market_watch_items LIKE 'price_text'")
    if not cur.fetchone():
        cur.execute("ALTER TABLE market_watch_items ADD COLUMN price_text VARCHAR(255) NULL AFTER summary")


def ensure_automation_task_schema(cur):
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS automation_task_runs (
            id BIGINT PRIMARY KEY AUTO_INCREMENT,
            task_name VARCHAR(64) NOT NULL,
            status ENUM('success', 'failed', 'running') NOT NULL DEFAULT 'running',
            started_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            finished_at TIMESTAMP NULL,
            summary JSON NULL,
            error TEXT NULL,
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            KEY idx_automation_task_name (task_name, started_at),
            KEY idx_automation_task_status (status, started_at)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
        """
    )


def ensure_data_quality_issue_schema(cur):
    duplicate_issue = "possible_duplicate_conflicting_product"
    cur.execute("SHOW COLUMNS FROM data_quality_issues LIKE 'issue_type'")
    row = cur.fetchone()
    if not row:
        return
    column_type = str(row[1] or "")
    if duplicate_issue in column_type:
        return
    cur.execute(
        """
        ALTER TABLE data_quality_issues
        MODIFY issue_type ENUM(
            'missing_exchange_rate',
            'missing_density',
            'unknown_product',
            'unknown_entity',
            'contract_unmatched',
            'contract_out_of_period',
            'summary_mismatch',
            'possible_duplicate_conflicting_product'
        ) NOT NULL
        """
    )


def record_automation_task(cur, task_name, status, summary=None, error=None):
    ensure_automation_task_schema(cur)
    cur.execute(
        """
        INSERT INTO automation_task_runs (task_name, status, started_at, finished_at, summary, error)
        VALUES (%s, %s, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP, %s, %s)
        """,
        (task_name, status, json.dumps(summary or {}, ensure_ascii=False), error),
    )


def market_region_label(value):
    return {
        "Russia": "俄罗斯",
        "Uzbekistan": "乌兹别克",
        "Global": "国际",
    }.get(value, value or "国际")


def market_type_label(value):
    return {
        "crude": "原料油",
        "refined": "成品油",
        "policy": "政策",
        "logistics": "物流",
        "other": "其他",
    }.get(value, value or "其他")


def default_market_details(region):
    label = market_region_label(region)
    return [
        {
            "typeLabel": "市场行情",
            "title": f"{label}市场行情",
            "summary": "关注供应节奏、批发成交、运输变化、库存压力和价格波动对采购销售的影响。",
        },
        {
            "typeLabel": "原料油",
            "title": f"{label}原料油行情",
            "summary": "关注原油、凝析油、炼厂原料供应、炼厂开工负荷和外部价格传导。",
        },
        {
            "typeLabel": "成品油",
            "title": f"{label}成品油行情",
            "summary": "关注汽油、柴油、液化气供应变化、库存水平、区域价差和短期需求。",
        },
    ]


UZEX_QUOTES_URL = "https://www.uzex.uz/Quote/GetQuotes"
UZEX_SOURCE_URL = "https://www.uzex.uz/ru/pages/public-data"
CBU_KGS_RATE_URL = "https://cbu.uz/ru/arkhiv-kursov-valyut/json/KGS/{rate_date}/"
SPIMEX_INDICES_URL = "https://spimex.global/indices/"
UZEX_PRODUCT_CODES = {
    "汽油(80)": "AI80",
    "柴油": "DIESEL_K5",
    "液化气": "LPG",
}
UZEX_FALLBACK_DENSITIES = {
    "汽油(80)": 0.760,
    "柴油": 0.850,
    "液化气": 0.554,
}

MARKET_BENCHMARKS = [
    {"symbol": "BZ%3DF", "label": "Brent"},
    {"symbol": "CL%3DF", "label": "WTI"},
]


def extract_market_price_text(*values):
    text = " ".join(clean_str(value) or "" for value in values)
    price_parts = []
    for label, value in re.findall(
        r"\b(Brent|WTI|Urals)\b[^\d$]{0,28}\$?\s*(\d+(?:\.\d+)?)",
        text,
        flags=re.IGNORECASE,
    ):
        price_parts.append(f"{label.upper() if label.lower() == 'wti' else label.title()} {value} 美元/桶")
    return "；".join(dict.fromkeys(price_parts))


def fetch_market_benchmark_item():
    values = []
    quote_date = None
    for benchmark in MARKET_BENCHMARKS:
        url = f"https://query1.finance.yahoo.com/v8/finance/chart/{benchmark['symbol']}?range=5d&interval=1d"
        req = Request(url, headers={"User-Agent": "SalesVisualization/1.0", "Accept": "application/json"})
        try:
            with urlopen(req, timeout=20) as response:
                payload = json.loads(response.read().decode("utf-8"))
            result = payload.get("chart", {}).get("result", [None])[0] or {}
            meta = result.get("meta", {})
            price = meta.get("regularMarketPrice")
            if price is None:
                closes = ((result.get("indicators", {}).get("quote") or [{}])[0].get("close") or [])
                price = next((value for value in reversed(closes) if value is not None), None)
            if price is None:
                continue
            market_time = meta.get("regularMarketTime")
            if market_time:
                market_date = datetime.fromtimestamp(market_time).date()
                quote_date = market_date if quote_date is None else max(quote_date, market_date)
            values.append(f"{benchmark['label']} {float(price):,.2f} 美元/桶")
        except (HTTPError, URLError, TimeoutError, OSError, ValueError, KeyError, TypeError) as exc:
            raise RuntimeError(f"国际原油价格获取失败：{exc}") from exc
    if not values:
        raise RuntimeError("国际原油价格获取失败：报价为空")
    price_text = "；".join(values)
    return {
        "region": "Global",
        "type": "crude",
        "title": "国际原油基准价格",
        "summary": f"{price_text}。用于观察国际原料油成本变化与成品油价格传导趋势。",
        "priceText": price_text,
        "sourceName": "Yahoo Finance",
        "sourceUrl": "https://finance.yahoo.com/markets/commodities/",
        "observedAt": quote_date or date.today(),
        "priority": 1,
    }


def fetch_cbu_uzs_per_kgs(rate_date):
    req = Request(
        CBU_KGS_RATE_URL.format(rate_date=rate_date.isoformat()),
        headers={"User-Agent": "SalesVisualization/1.0", "Accept": "application/json"},
    )
    last_error = None
    for timeout in (8, 15):
        try:
            with urlopen(req, timeout=timeout) as response:
                rates = json.loads(response.read().decode("utf-8"))
            break
        except (HTTPError, URLError, TimeoutError, OSError, ValueError, TypeError) as exc:
            last_error = exc
    else:
        raise RuntimeError(f"乌兹别克央行索姆汇率获取失败：{last_error}") from last_error
    row = next(
        (item for item in rates if str(item.get("Ccy") or "").upper() == "KGS"),
        None,
    ) if isinstance(rates, list) else None
    value = parse_decimal(row.get("Rate") if row else None)
    if value is None or value <= 0:
        raise RuntimeError("乌兹别克央行索姆汇率获取失败：未找到 KGS 汇率")
    effective_date = parse_date(row.get("Date") if row else None) or rate_date
    return float(value), effective_date


def format_uzex_raw_price(value):
    return f"{float(value):,.2f} 千乌兹别克索姆/吨"


def build_uzex_price_text(label, value, density, uzs_per_kgs):
    if uzs_per_kgs:
        kgs_per_liter = float(value) * density / uzs_per_kgs
        return f"{label} {kgs_per_liter:,.2f} 索姆/升"
    return f"{label} {format_uzex_raw_price(value)}"


def market_refined_densities(cur, effective_date):
    product_codes = tuple(UZEX_PRODUCT_CODES.values())
    placeholders = ", ".join(["%s"] * len(product_codes))
    cur.execute(
        f"""
        SELECT p.product_code, d.density
        FROM product_density_history d
        JOIN products p ON p.id = d.product_id
        WHERE p.product_code IN ({placeholders}) AND d.effective_date <= %s
        ORDER BY p.product_code, d.effective_date DESC, d.id DESC
        """,
        (*product_codes, effective_date),
    )
    by_code = {}
    for product_code, density in cur.fetchall():
        by_code.setdefault(product_code, float(density))
    return {
        label: by_code.get(product_code, UZEX_FALLBACK_DENSITIES[label])
        for label, product_code in UZEX_PRODUCT_CODES.items()
    }


def fetch_uzex_market_item(cur):
    observed_at = date.today()
    req = Request(UZEX_QUOTES_URL, headers={"User-Agent": "SalesVisualization/1.0", "Accept": "application/json"})
    try:
        with urlopen(req, timeout=20) as response:
            quotes = json.loads(response.read().decode("utf-8"))
    except (HTTPError, URLError, TimeoutError, OSError, ValueError, TypeError) as exc:
        raise RuntimeError(f"UZEX 官方报价获取失败：{exc}") from exc

    quote_names = {
        "\u0411\u0415\u041d\u0417\u0418\u041d (80)": "汽油(80)",
        "\u0414\u0418\u0417\u0415\u041b\u042c \u0401\u049a\u0418\u041b\u0492\u0418\u0421\u0418": "柴油",
        "\u0421\u0423\u042e\u041b\u0422\u0418\u0420\u0418\u041b\u0413\u0410\u041d \u0413\u0410\u0417": "液化气",
    }
    uzs_per_kgs = None
    rate_date = None
    rate_error = None
    try:
        uzs_per_kgs, rate_date = fetch_cbu_uzs_per_kgs(observed_at)
    except RuntimeError as exc:
        rate_error = str(exc)
    densities = market_refined_densities(cur, observed_at)
    values = []
    movements = []
    for row in quotes if isinstance(quotes, list) else []:
        label = quote_names.get(clean_str(row.get("name")) or "")
        if not label:
            continue
        value = parse_decimal(row.get("summa"))
        if value is None:
            continue
        percent = clean_str(row.get("percent")) or "0.00"
        values.append(build_uzex_price_text(label, value, densities[label], uzs_per_kgs))
        movements.append(f"{label} {percent}%")
    if not values:
        raise RuntimeError("UZEX 官方报价获取失败：未找到油品报价")
    price_text = "；".join(values)
    density_text = "、".join(f"{label} {density:.3f}" for label, density in densities.items())
    if uzs_per_kgs:
        conversion_note = (
            f"按 UZEX 千乌兹别克索姆/吨报价、乌兹别克央行 {rate_date.isoformat()} 汇率"
            f"（1 索姆={uzs_per_kgs:,.2f} 乌兹别克索姆）及平台密度（{density_text}）换算"
        )
    else:
        conversion_note = (
            f"乌兹别克央行 KGS 汇率接口暂时不可用，今日先保留 UZEX 原始报价；"
            f"汇率恢复后可在后台重新获取并换算为索姆/升。{rate_error}"
        )
    return {
        "region": "Uzbekistan",
        "type": "refined",
        "title": "乌兹别克成品油公开报价",
        "summary": f"UZEX 官网公开报价折算：{price_text}；涨跌幅：{'；'.join(movements)}。{conversion_note}。",
        "priceText": price_text,
        "sourceName": "UZEX乌兹别克商品交易所",
        "sourceUrl": UZEX_SOURCE_URL,
        "observedAt": observed_at,
        "priority": 2,
    }


def spimex_official_reference_item():
    return {
        "region": "Russia",
        "type": "refined",
        "title": "俄罗斯成品油官方指数参考",
        "summary": "SPIMEX 官方提供成品油、液化气与原油指数。根据其信息使用条款，指数数值的再展示需按许可要求接入，当前平台先提供官方入口供经营复核。",
        "priceText": "",
        "sourceName": "SPIMEX彼得堡国际商品交易所",
        "sourceUrl": SPIMEX_INDICES_URL,
        "observedAt": date.today(),
        "priority": 2,
    }


def upsert_market_source_item(cur, item):
    cur.execute(
        """
        SELECT id FROM market_watch_items
        WHERE market_region = %s AND source_url = %s
        ORDER BY id DESC LIMIT 1
        """,
        (item["region"], item["sourceUrl"]),
    )
    existing = cur.fetchone()
    values = (
        item["title"], item["summary"], item.get("priceText") or None,
        item["sourceName"], item["sourceUrl"], item["observedAt"], item["priority"],
    )
    if existing:
        cur.execute(
            """
            UPDATE market_watch_items
            SET market_type = %s, title = %s, summary = %s, price_text = %s,
                source_name = %s, source_url = %s, observed_at = %s, priority = %s,
                is_active = TRUE
            WHERE id = %s
            """,
            (item["type"], *values, existing[0]),
        )
        cur.execute(
            """
            UPDATE market_watch_items
            SET is_active = FALSE
            WHERE market_region = %s AND source_url = %s AND id <> %s AND is_active = TRUE
            """,
            (item["region"], item["sourceUrl"], existing[0]),
        )
        return "updated"
    cur.execute(
        """
        INSERT INTO market_watch_items
            (market_region, market_type, title, summary, price_text, source_name, source_url, observed_at, priority, is_active)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, TRUE)
        """,
        (
            item["region"], item["type"], item["title"], item["summary"], item.get("priceText") or None,
            item["sourceName"], item["sourceUrl"], item["observedAt"], item["priority"],
        ),
    )
    return "inserted"


def has_chinese(text):
    return bool(re.search(r"[\u4e00-\u9fff]", text or ""))


MARKET_TRANSLATION_TERMS = (
    ("centre for research on energy and clean air", "能源与清洁空气研究中心"),
    ("the times of emerging markets", "中亚时报"),
    ("times of emerging markets", "中亚时报"),
    ("s&p global", "标普全球"),
    ("the economic times", "经济时报"),
    ("tradingkey", "市场资讯"),
    ("oil slip", "油价回落"),
    ("kazakhstan enforces", "哈萨克斯坦实施"),
    ("kazakhstan", "哈萨克斯坦"),
    ("uzbekistan", "乌兹别克斯坦"),
    ("emerging markets", "中亚"),
    ("refined products", "成品油"),
    ("refined product", "成品油"),
    ("refinery margins", "炼厂利润"),
    ("refinery margin", "炼厂利润"),
    ("crack spread", "裂解价差"),
    ("condensate", "凝析油"),
    ("urals", "Urals 原油"),
    ("monthly analysis", "月度分析"),
    ("russian fossil fuel exports and sanctions", "俄罗斯化石燃料出口与制裁"),
    ("russian fossil fuel exports", "俄罗斯化石燃料出口"),
    ("russian", "俄罗斯"),
    ("russia", "俄罗斯"),
    ("fuel market situation in kyrgyzstan remains stable", "吉尔吉斯斯坦燃料市场形势保持稳定"),
    ("situation in kyrgyzstan remains stable", "吉尔吉斯斯坦市场形势保持稳定"),
    ("fuel market situation", "燃料市场形势"),
    ("market situation", "市场形势"),
    ("remains stable", "保持稳定"),
    ("kyrgyzstan", "吉尔吉斯斯坦"),
    ("fossil fuel", "化石燃料"),
    ("exports", "出口"),
    ("export", "出口"),
    ("sanctions", "制裁"),
    ("crude oil", "原油"),
    ("crude futures", "原油期货"),
    ("oil prices", "油价"),
    ("oil price", "油价"),
    ("diesel", "柴油"),
    ("gasoline", "汽油"),
    ("petrol", "汽油"),
    ("refinery", "炼厂"),
    ("refineries", "炼厂"),
    ("fuel", "燃料"),
    ("market", "市场"),
    ("situation", "形势"),
    ("inventory", "库存"),
    ("stocks", "库存"),
    ("stock", "库存"),
    ("enforces", "实施"),
    ("ban", "禁令"),
    ("opec", "OPEC"),
    ("brent", "Brent 原油"),
    ("wti", "WTI 原油"),
    ("april", "4月"),
    ("may", "5月"),
    ("june", "6月"),
    ("january", "1月"),
    ("february", "2月"),
    ("march", "3月"),
    ("july", "7月"),
    ("august", "8月"),
    ("september", "9月"),
    ("october", "10月"),
    ("november", "11月"),
    ("december", "12月"),
    (" of ", "的"),
    (" in ", "在"),
)


def translate_market_text(text):
    value = clean_str(text) or ""
    translated = value
    for raw, zh in MARKET_TRANSLATION_TERMS:
        pattern = rf"\b{raw}\b(?!\s*原油)" if raw in {"urals", "brent", "wti"} else raw
        translated = re.sub(pattern, zh, translated, flags=re.IGNORECASE)
    translated = re.sub(r"\b(Brent|WTI|Urals)\s*原油(?:\s*原油)+", r"\1 原油", translated)
    translated = re.sub(r"\s+-\s+[^-]{3,80}$", "", translated).strip()
    translated = translated.replace(" — ", "：").replace(" - ", "：")
    translated = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])", "", translated)
    translated = re.sub(r"\s{2,}", " ", translated).strip(" -—：:")
    return translated or value


def market_source_display(source):
    source = clean_str(source) or ""
    if source.lower() == "yahoo finance":
        return "雅虎财经"
    return source if has_chinese(source) else "公开资讯"


def strip_market_source_echo(text, source):
    value = clean_str(text) or ""
    source = clean_str(source) or ""
    candidates = [source, translate_market_text(source)]
    for candidate in candidates:
        if candidate:
            value = re.sub(rf"\s*[-—：:]?\s*{re.escape(candidate)}\s*$", "", value, flags=re.IGNORECASE).strip()
    return value


def has_untranslated_market_english(text):
    value = clean_str(text) or ""
    value = re.sub(r"(OPEC|WTI|Brent|USD|CNY|KGS|SPIMEX|UZEX)", "", value, flags=re.IGNORECASE)
    return bool(re.search(r"[A-Za-z]{3,}", value))


def chinese_market_item(item):
    item = dict(item)
    title = item.get("title") or ""
    summary = item.get("summary") or ""
    source = market_source_display(item.get("sourceName"))
    translated_title = strip_market_source_echo(translate_market_text(title), item.get("sourceName"))
    translated_summary = strip_market_source_echo(
        translate_market_text(re.sub(r"\s+", " ", summary or title).strip()),
        item.get("sourceName"),
    )
    item["displayReady"] = not has_untranslated_market_english(translated_title + " " + translated_summary)
    if has_untranslated_market_english(translated_title):
        translated_title = f"{market_region_label(item.get('region'))}市场资讯更新"
    item["title"] = translated_title if translated_title else f"{market_region_label(item.get('region'))}市场资讯"
    translated_summary = re.sub(rf"^(?:{re.escape(source)}报道：)+", "", translated_summary).strip()
    if has_untranslated_market_english(translated_summary):
        translated_summary = item["title"]
    if translated_summary and translated_summary != item["title"]:
        item["summary"] = f"{source}报道：{translated_summary[:260]}"
    else:
        item["summary"] = f"{source}报道：{item['title']}"
    return item


def fetch_market_feed(feed):
    query = urlencode({
        "q": feed["query"],
        "hl": feed.get("hl", "en-US"),
        "gl": feed.get("gl", "KG"),
        "ceid": feed.get("ceid", "KG:en"),
    })
    request_url = f"https://news.google.com/rss/search?{query}"
    req = Request(
        request_url,
        headers={
            "User-Agent": "SalesVisualization/1.0",
            "Accept": "application/rss+xml,application/xml,text/xml,*/*",
        },
    )
    try:
        with urlopen(req, timeout=20) as response:
            raw_xml = response.read()
    except (HTTPError, URLError, TimeoutError, OSError) as exc:
        raise RuntimeError(f"{market_region_label(feed['region'])}市场观察获取失败：{exc}") from exc

    try:
        root = ET.fromstring(raw_xml)
    except ET.ParseError as exc:
        raise RuntimeError(f"{market_region_label(feed['region'])}市场观察返回内容无法解析") from exc

    items = []
    for node in root.findall(".//item")[:3]:
        title = html.unescape(clean_str(node.findtext("title")) or "")
        link = clean_str(node.findtext("link")) or ""
        description = html.unescape(clean_str(node.findtext("description")) or "")
        description = re.sub(r"<.*?>", " ", description)
        description = re.sub(r"\s+", " ", description).strip()
        source_node = node.find("source")
        source_name = clean_str(source_node.text if source_node is not None else None) or "Google News"
        observed_at = date.today()
        pub_date = clean_str(node.findtext("pubDate"))
        if pub_date:
            try:
                observed_at = parsedate_to_datetime(pub_date).date()
            except (TypeError, ValueError, AttributeError):
                pass
        if title and link:
            items.append({
                "region": feed["region"],
                "type": feed["type"],
                "title": title[:255],
                "summary": (description or title)[:900],
                "priceText": extract_market_price_text(title, description),
                "sourceName": source_name[:128],
                "sourceUrl": link[:500],
                "observedAt": observed_at,
                "priority": feed["priority"] + len(items),
            })
    return items


def refresh_market_watch_chinese(cur, limit=120):
    ensure_market_watch_schema(cur)
    cur.execute(
        """
        SELECT id, market_region, title, summary, source_name, observed_at
        FROM market_watch_items
        WHERE is_active = TRUE
        ORDER BY observed_at DESC, priority ASC, id DESC
        LIMIT %s
        """,
        (limit,),
    )
    rows = cur.fetchall()
    updated = 0
    for item_id, region, title, summary, source_name, observed_at in rows:
        zh_item = chinese_market_item({
            "region": region,
            "title": title,
            "summary": summary,
            "sourceName": source_name,
            "observedAt": observed_at.isoformat() if observed_at else None,
        })
        new_title = zh_item.get("title") or title
        new_summary = zh_item.get("summary") or summary
        new_source = market_source_display(source_name)
        if new_title != title or new_summary != summary or new_source != source_name:
            cur.execute(
                """
                UPDATE market_watch_items
                SET title = %s, summary = %s, source_name = %s
                WHERE id = %s
                """,
                (new_title[:255], new_summary[:900], new_source[:128], item_id),
            )
            updated += 1
    return updated


def market_watch_rows(cur, limit=80):
    ensure_market_watch_schema(cur)
    cur.execute(
        """
        SELECT id, market_region, market_type, title, summary,
               price_text, source_name, source_url, observed_at, priority, is_active, created_at
        FROM market_watch_items
        WHERE market_region IN ('Russia', 'Uzbekistan', 'Global')
        ORDER BY observed_at DESC, priority ASC, id DESC
        LIMIT %s
        """,
        (limit,),
    )
    return [
        {
            "id": row[0],
            "region": row[1],
            "regionLabel": market_region_label(row[1]),
            "type": row[2],
            "typeLabel": market_type_label(row[2]),
            "title": row[3],
            "summary": row[4],
            "priceText": row[5] or "",
            "sourceName": row[6] or "",
            "sourceUrl": row[7] or "",
            "observedAt": row[8].isoformat() if row[8] else None,
            "priority": row[9],
            "isActive": bool(row[10]),
            "createdAt": row[11].isoformat() if row[11] else None,
        }
        for row in cur.fetchall()
    ]


def run_exchange_maintenance(cur):
    today = date.today()
    start = date(today.year, today.month, 1)
    end = today + timedelta(days=1)
    dates = missing_exchange_dates(cur, start, end)
    fetched = []
    unresolved = []
    errors = []
    for rate_date in dates:
        try:
            usd, cny, source_url = fetch_nbkr_daily_rates(rate_date)
            resolved = upsert_exchange_rate(cur, rate_date, usd, cny, f"NBKR 历史页面自动获取：{source_url}")
            fetched.append({
                "date": rate_date.isoformat(),
                "usdKgs": float(usd) if usd is not None else None,
                "cnyKgs": float(cny) if cny is not None else None,
                "resolved": resolved,
            })
            if not resolved:
                unresolved.append(rate_date.isoformat())
        except RuntimeError as exc:
            errors.append({"date": rate_date.isoformat(), "error": str(exc)})
    recalculation = recalculate_all_wholesale(cur) if fetched else {"matched": 0, "unmatched": 0, "missing_rates": 0}
    return {
        "requested": len(dates),
        "updated": len(fetched),
        "resolved": sum(1 for item in fetched if item["resolved"]),
        "unresolved": unresolved,
        "errors": errors,
        "recalculation": recalculation,
    }


def run_market_watch_maintenance(cur):
    ensure_market_watch_schema(cur)
    official_items = []
    errors = []
    try:
        official_items.append(fetch_market_benchmark_item())
    except RuntimeError as exc:
        errors.append(str(exc))
    try:
        official_items.append(fetch_uzex_market_item(cur))
    except RuntimeError as exc:
        errors.append(str(exc))
    official_items.append(spimex_official_reference_item())
    inserted = 0
    updated = 0
    for item in official_items:
        status = upsert_market_source_item(cur, item)
        inserted += status == "inserted"
        updated += status == "updated"
    translated = refresh_market_watch_chinese(cur)
    return {"fetched": len(official_items), "inserted": inserted, "updated": updated, "translated": translated, "errors": errors}


def run_retail_snapshot_maintenance():
    script = OILSTATION_ROOT / "scripts" / "refresh_current_snapshots.py"
    if not script.exists():
        raise RuntimeError(f"零售快照脚本不存在：{script}")
    completed = subprocess.run(
        [sys.executable, str(script)],
        cwd=str(OILSTATION_ROOT),
        capture_output=True,
        text=True,
        timeout=900,
    )
    output = (completed.stdout or completed.stderr or "").strip()
    if completed.returncode != 0:
        raise RuntimeError(output[-1000:] or f"零售快照脚本退出码 {completed.returncode}")
    return {
        "returnCode": completed.returncode,
        "outputTail": output[-1200:],
    }


def log_operation(cur, action, target_type=None, target_id=None, detail=None):
    ensure_operation_log_schema(cur)
    user = session_user() or {}
    cur.execute(
        """
        INSERT INTO operation_logs (user_id, username, action, target_type, target_id, detail)
        VALUES (%s, %s, %s, %s, %s, %s)
        """,
        (
            user.get("id"),
            user.get("username"),
            action,
            target_type,
            str(target_id) if target_id is not None else None,
            json.dumps(detail or {}, ensure_ascii=False),
        ),
    )


def log_system_operation(action, target_type=None, target_id=None, detail=None):
    conn = get_connection()
    try:
        cur = conn.cursor()
        log_operation(cur, action, target_type, target_id, detail)
        conn.commit()
    except Exception:
        conn.rollback()
        logging.exception("Failed to write operation log: %s", action)
    finally:
        conn.close()


def retail_station_catalog():
    spec = importlib.util.spec_from_file_location(
        "sales_oilstation_constants",
        OILSTATION_ROOT / "constants.py",
    )
    module = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    spec.loader.exec_module(module)
    return {
        int(shop_key): {
            "shopKey": int(shop_key),
            "name": info.get("name", ""),
            "short": info.get("short", ""),
            "region": info.get("region", ""),
        }
        for shop_key, info in module.SHOP_KEY_MAP.items()
    }


def retail_station_catalog_cached():
    if not hasattr(retail_station_catalog_cached, "_cache"):
        retail_station_catalog_cached._cache = retail_station_catalog()
    return retail_station_catalog_cached._cache


STATION_STATUS_TYPES = {
    "normal": "正常营业",
    "power_outage": "停电",
    "maintenance": "设备故障",
    "renovation": "装修/维修",
    "suspended": "暂停营业",
    "road_blocked": "修路占道",
    "data_anomaly": "数据异常",
    "other": "其他",
}


def ensure_station_status_schema(cur):
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS retail_station_status (
            id BIGINT PRIMARY KEY AUTO_INCREMENT,
            shop_key INT NOT NULL,
            status_type VARCHAR(32) NOT NULL,
            status_label VARCHAR(64) NOT NULL,
            started_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
            expected_resume_at DATETIME NULL,
            ended_at DATETIME NULL,
            affects_sales BOOLEAN NOT NULL DEFAULT FALSE,
            affects_inventory BOOLEAN NOT NULL DEFAULT FALSE,
            suppress_alerts BOOLEAN NOT NULL DEFAULT FALSE,
            note TEXT NULL,
            is_active BOOLEAN NOT NULL DEFAULT TRUE,
            created_by BIGINT NULL,
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP NULL ON UPDATE CURRENT_TIMESTAMP,
            KEY idx_retail_station_status_active (is_active, ended_at, shop_key),
            KEY idx_retail_station_status_shop (shop_key, started_at)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
        """
    )
    cur.execute("SELECT COUNT(*) FROM retail_station_status")
    if cur.fetchone()[0]:
        return
    default_rows = [
        (1, "renovation", "装修中", "装修停业，五月中恢复", True, True, False),
        (2, "normal", "正常营业", "已恢复营业", False, False, False),
        (11, "road_blocked", "修路占道", "修路占道，无法营业", True, True, False),
        (12, "road_blocked", "修路占道", "修路占道，无法营业", True, True, False),
        (29, "renovation", "装修中", "装修停业，五月初开业", True, True, False),
        (30, "renovation", "等待装修", "停业，等待装修", True, True, False),
    ]
    cur.executemany(
        """
        INSERT INTO retail_station_status
            (shop_key, status_type, status_label, note, affects_sales, affects_inventory, suppress_alerts)
        VALUES (%s, %s, %s, %s, %s, %s, %s)
        """,
        default_rows,
    )
    connection = getattr(cur, "connection", None)
    if connection:
        connection.commit()


def active_station_statuses(cur):
    ensure_station_status_schema(cur)
    cur.execute(
        """
        SELECT id, shop_key, status_type, status_label, started_at, expected_resume_at,
               ended_at, affects_sales, affects_inventory, suppress_alerts, note, is_active
        FROM retail_station_status
        WHERE is_active = TRUE AND ended_at IS NULL
          AND status_type <> 'normal'
        ORDER BY shop_key, started_at DESC, id DESC
        """
    )
    catalog = retail_station_catalog()
    statuses = {}
    for row in cur.fetchall():
        shop_key = int(row[1])
        if shop_key in statuses:
            continue
        station = catalog.get(shop_key, {"name": f"Station-{shop_key}", "short": "", "region": ""})
        status_type = row[2]
        affects_sales = bool(row[7])
        public_status = "normal" if status_type == "normal" else ("closed" if affects_sales else "warning")
        reason = "" if status_type == "normal" else (row[10] or row[3] or STATION_STATUS_TYPES.get(status_type, status_type))
        statuses[shop_key] = {
            "id": row[0],
            "shopKey": shop_key,
            "name": station["name"],
            "short": station["short"],
            "region": station["region"],
            "status": public_status,
            "statusType": status_type,
            "statusLabel": row[3],
            "closedReason": reason,
            "note": row[10] or "",
            "startedAt": row[4].isoformat() if row[4] else None,
            "expectedResumeAt": row[5].isoformat() if row[5] else None,
            "endedAt": row[6].isoformat() if row[6] else None,
            "affectsSales": affects_sales,
            "affectsInventory": bool(row[8]),
            "suppressAlerts": bool(row[9]),
            "isActive": bool(row[11]),
        }
    return statuses


def apply_station_status(item, statuses):
    if not isinstance(item, dict):
        return item
    shop_key = item.get("shopKey")
    if shop_key is None:
        return item
    status = statuses.get(int(shop_key))
    if not status:
        return item
    item["status"] = status["status"]
    item["statusType"] = status["statusType"]
    item["statusLabel"] = status["statusLabel"]
    item["closedReason"] = status["closedReason"]
    item["affectsSales"] = status["affectsSales"]
    item["affectsInventory"] = status["affectsInventory"]
    item["suppressAlerts"] = status["suppressAlerts"]
    item["statusNote"] = status["note"]
    item["expectedResumeAt"] = status["expectedResumeAt"]
    return item


def enrich_retail_station_statuses(payload):
    if not isinstance(payload, dict):
        return payload
    conn = get_connection()
    try:
        cur = conn.cursor()
        statuses = active_station_statuses(cur)
    finally:
        conn.close()
    def sync_station(item):
        if not isinstance(item, dict):
            return item
        shop_key = item.get("shopKey")
        if shop_key is None:
            return item
        if int(shop_key) in statuses:
            return apply_station_status(item, statuses)
        item["status"] = "normal"
        item["statusType"] = "normal"
        item["statusLabel"] = "正常营业"
        item["closedReason"] = ""
        item["affectsSales"] = False
        item["affectsInventory"] = False
        item["suppressAlerts"] = False
        item["statusNote"] = ""
        item["expectedResumeAt"] = None
        return item

    if "stations" in payload:
        if isinstance(payload["stations"], list):
            payload["stations"] = [sync_station(item) for item in payload["stations"]]
        elif isinstance(payload["stations"], dict):
            for rows in payload["stations"].values():
                if isinstance(rows, list):
                    for item in rows:
                        sync_station(item)
    if "shopKey" in payload:
        sync_station(payload)
    payload["stationStatuses"] = list(statuses.values())
    return payload


def has_users(cur):
    ensure_auth_schema(cur)
    cur.execute("SELECT COUNT(*) FROM system_users")
    return bool(cur.fetchone()[0])


def session_user():
    user = session.get("user")
    return user if isinstance(user, dict) else None


def wants_json():
    return request.path.startswith("/api/")


@app.after_request
def add_security_headers(response):
    response.headers.setdefault("X-Content-Type-Options", "nosniff")
    response.headers.setdefault("X-Frame-Options", "SAMEORIGIN")
    response.headers.setdefault("Referrer-Policy", "strict-origin-when-cross-origin")
    response.headers.setdefault("Permissions-Policy", "camera=(), microphone=(), geolocation=(), payment=()")
    return response


def login_redirect():
    next_url = request.full_path if request.query_string else request.path
    if not next_url.startswith("/") or next_url.startswith("//"):
        next_url = "/"
    return redirect(f"/login?next={next_url}")


@app.before_request
def require_session():
    path = request.path
    if path in PUBLIC_PATHS or path.startswith("/assets/"):
        return None
    if path in PUBLIC_API_PATHS or any(path.startswith(prefix) for prefix in PUBLIC_API_PREFIXES):
        return None
    if any(path == prefix or path.startswith(prefix + "/") for prefix in PUBLIC_PAGE_PREFIXES):
        return None
    if path.endswith(".css") or path.endswith(".js") or path.endswith(".json"):
        return None
    conn = get_connection()
    try:
        cur = conn.cursor()
        if not has_users(cur):
            if wants_json():
                return jsonify({"error": "bootstrap_required"}), 401
            return redirect("/login")
    finally:
        conn.close()
    if session_user():
        return None
    if wants_json():
        return jsonify({"error": "login_required"}), 401
    return login_redirect()


def admin_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        user = session_user()
        if not user or user.get("role") != "admin":
            return jsonify({"error": "admin_required"}), 403
        return fn(*args, **kwargs)

    return wrapper


def recalculate_all_wholesale(cur):
    cur.execute(
        """
        SELECT id
        FROM import_batches
        WHERE source_type IN ('export_sales_excel', 'manual')
          AND status = 'imported'
        """
    )
    total = {"matched": 0, "unmatched": 0, "missing_rates": 0}
    for (batch_id,) in cur.fetchall():
        result = recalculate_wholesale_batch(cur, batch_id)
        for key in total:
            total[key] += result[key]
    return total


def reconcile_wholesale_product_ids(cur):
    product_lookup = load_product_lookup(cur)
    targets = (
        ("wholesale_detail_records", "raw_product_name", "product_id"),
        ("wholesale_summary_records", "raw_product_name", "product_id"),
        ("contracts", "raw_product_name", "product_id"),
    )
    result = {}
    for table, name_column, id_column in targets:
        cur.execute(f"SELECT id, {name_column}, {id_column} FROM {table}")
        scanned = 0
        recognized = 0
        updated = 0
        for row_id, raw_name, old_product_id in cur.fetchall():
            scanned += 1
            product = lookup_product(product_lookup, raw_name)
            if not product:
                continue
            recognized += 1
            if old_product_id != product["id"]:
                cur.execute(f"UPDATE {table} SET {id_column} = %s WHERE id = %s", (product["id"], row_id))
                updated += 1
        result[table] = {"scanned": scanned, "recognized": recognized, "updated": updated}
    return result


def _nbkr_daily_url(rate_date):
    separator = "&" if "?" in NBKR_DAILY_XML_URL else "?"
    query = urlencode({"date": rate_date.strftime("%d.%m.%Y")})
    return f"{NBKR_DAILY_XML_URL}{separator}{query}"


def _nbkr_history_url(rate_date, currency_id):
    return "https://www.nbkr.kg/index1.jsp?" + urlencode({
        "beg_day": f"{rate_date.day:02d}",
        "beg_month": f"{rate_date.month:02d}",
        "beg_year": rate_date.year,
        "end_day": f"{rate_date.day:02d}",
        "end_month": f"{rate_date.month:02d}",
        "end_year": rate_date.year,
        "item": 1562,
        "lang": "ENG",
        "valuta_id": currency_id,
    })


def fetch_nbkr_history_rate(rate_date, currency_id):
    request_url = _nbkr_history_url(rate_date, currency_id)
    req = Request(
        request_url,
        headers={
            "User-Agent": "SalesVisualization/1.0",
            "Accept": "text/html,*/*",
        },
    )
    try:
        with urlopen(req, timeout=20) as response:
            html = response.read().decode("utf-8", errors="ignore")
    except (HTTPError, URLError, TimeoutError, OSError) as exc:
        raise RuntimeError(f"NBKR 历史汇率页面获取失败：{exc}") from exc

    expected = rate_date.strftime("%d.%m.%Y")
    match = re.search(
        rf"<!--date-->\s*{re.escape(expected)}\s*<!--date-->.*?<!--value-->\s*([^<]+?)\s*<!--value-->",
        html,
        flags=re.S,
    )
    if not match:
        raise RuntimeError(f"NBKR 历史汇率页面未返回 {expected} 的汇率")
    value = parse_decimal(match.group(1))
    if value is None:
        raise RuntimeError(f"NBKR 历史汇率页面中 {expected} 的汇率无法解析")
    return value, request_url


def fetch_nbkr_daily_rates(rate_date):
    usd, usd_url = fetch_nbkr_history_rate(rate_date, 15)
    cny, cny_url = fetch_nbkr_history_rate(rate_date, 24)
    return usd, cny, f"USD:{usd_url}; CNY:{cny_url}"


def fetch_nbkr_latest_daily_rates(rate_date):
    request_url = _nbkr_daily_url(rate_date)
    req = Request(
        request_url,
        headers={
            "User-Agent": "SalesVisualization/1.0",
            "Accept": "application/xml,text/xml,*/*",
        },
    )
    try:
        with urlopen(req, timeout=20) as response:
            raw_xml = response.read()
    except (HTTPError, URLError, TimeoutError, OSError) as exc:
        raise RuntimeError(f"NBKR 汇率获取失败：{exc}") from exc

    try:
        root = ET.fromstring(raw_xml)
    except ET.ParseError as exc:
        raise RuntimeError("NBKR 返回内容无法解析") from exc

    returned_date = parse_date(root.attrib.get("Date"))
    if returned_date and returned_date != rate_date:
        raise RuntimeError(
            f"NBKR daily.xml 返回日期为 {returned_date.isoformat()}，不是请求的 {rate_date.isoformat()}"
        )

    rates = {"USD": None, "CNY": None}
    for node in root.iter():
        iso = (node.attrib.get("ISOCode") or node.attrib.get("isocode") or "").upper()
        if iso not in rates:
            continue
        value_node = node.find("Value")
        if value_node is None:
            value_node = node.find("value")
        value = parse_decimal(value_node.text if value_node is not None else None)
        if value is not None:
            rates[iso] = value
    return rates["USD"], rates["CNY"], request_url


def missing_exchange_dates(cur, start, end):
    fetch_end = min(end, date.today() + timedelta(days=1))
    if start >= fetch_end:
        return []

    cur.execute(
        """
        SELECT DISTINCT rate_date
        FROM (
            SELECT rate_date
            FROM exchange_rate_gaps
            WHERE rate_date >= %s AND rate_date < %s
              AND status = 'open'
            UNION
            SELECT rate_date
            FROM exchange_rates
            WHERE rate_date >= %s AND rate_date < %s
              AND (
                usd_kgs IS NULL
                OR cny_kgs IS NULL
                OR (is_manual = FALSE AND note LIKE '%%/XML/daily.xml?date=%%')
              )
            UNION
            SELECT exchange_rate_date AS rate_date
            FROM calculation_results
            WHERE exchange_rate_date >= %s AND exchange_rate_date < %s
              AND exchange_missing = TRUE
            UNION
            SELECT issue_date AS rate_date
            FROM data_quality_issues
            WHERE issue_type = 'missing_exchange_rate'
              AND issue_date >= %s AND issue_date < %s
        ) missing_dates
        WHERE rate_date IS NOT NULL
        ORDER BY rate_date
        """,
        (start, fetch_end, start, fetch_end, start, fetch_end, start, fetch_end),
    )
    dates = {row[0] for row in cur.fetchall()}

    cur.execute(
        """
        SELECT rate_date
        FROM exchange_rates
        WHERE rate_date >= %s AND rate_date < %s
          AND usd_kgs IS NOT NULL
          AND cny_kgs IS NOT NULL
          AND NOT (is_manual = FALSE AND note LIKE '%%/XML/daily.xml?date=%%')
        """,
        (start, fetch_end),
    )
    complete_dates = {row[0] for row in cur.fetchall()}

    current = start
    while current < fetch_end:
        if current not in complete_dates:
            dates.add(current)
        current += timedelta(days=1)

    return sorted(dates)


def upsert_exchange_rate(cur, rate_date, usd, cny, note):
    cur.execute(
        """
        INSERT INTO exchange_rates
            (rate_date, usd_kgs, cny_kgs, source, is_manual, fetched_at, note)
        VALUES (%s, %s, %s, 'NBKR', FALSE, CURRENT_TIMESTAMP, %s)
        ON DUPLICATE KEY UPDATE
            usd_kgs = IF(is_manual = TRUE AND usd_kgs IS NOT NULL, usd_kgs, COALESCE(VALUES(usd_kgs), usd_kgs)),
            cny_kgs = IF(is_manual = TRUE AND cny_kgs IS NOT NULL, cny_kgs, COALESCE(VALUES(cny_kgs), cny_kgs)),
            source = IF(is_manual = TRUE, source, 'NBKR'),
            fetched_at = CURRENT_TIMESTAMP,
            note = IF(is_manual = TRUE AND note IS NOT NULL, note, VALUES(note))
        """,
        (rate_date, usd, cny, note),
    )
    cur.execute("SELECT usd_kgs, cny_kgs FROM exchange_rates WHERE rate_date = %s", (rate_date,))
    saved = cur.fetchone()
    missing_usd = not saved or saved[0] is None
    missing_cny = not saved or saved[1] is None
    if missing_usd or missing_cny:
        cur.execute(
            """
            INSERT INTO exchange_rate_gaps (rate_date, missing_usd, missing_cny, status, note)
            VALUES (%s, %s, %s, 'open', 'NBKR 自动获取后仍缺少部分汇率。')
            ON DUPLICATE KEY UPDATE
                missing_usd = VALUES(missing_usd),
                missing_cny = VALUES(missing_cny),
                status = 'open',
                note = VALUES(note)
            """,
            (rate_date, missing_usd, missing_cny),
        )
        return False

    cur.execute(
        """
        INSERT INTO exchange_rate_gaps (rate_date, missing_usd, missing_cny, status, resolved_at, note)
        VALUES (%s, FALSE, FALSE, 'resolved', CURRENT_TIMESTAMP, 'Resolved by NBKR automatic fetch.')
        ON DUPLICATE KEY UPDATE
            missing_usd = FALSE,
            missing_cny = FALSE,
            status = 'resolved',
            resolved_at = CURRENT_TIMESTAMP,
            note = VALUES(note)
        """,
        (rate_date,),
    )
    return True


def create_manual_batch(cur, note, period_date=None):
    period_year = period_date.year if period_date else None
    period_month = period_date.month if period_date else None
    cur.execute(
        """
        INSERT INTO import_batches
            (source_type, period_year, period_month, status, imported_at, note)
        VALUES ('manual', %s, %s, 'imported', CURRENT_TIMESTAMP, %s)
        """,
        (period_year, period_month, note),
    )
    return cur.lastrowid


def detect_upload_type(path):
    xl = pd.ExcelFile(path)
    if len(xl.sheet_names) > SHEET_DETAIL:
        return "wholesale"
    return "contracts"


def upload_type_label(upload_type):
    return "外运外销表" if upload_type == "wholesale" else "合同登记表"


def import_source_label(source_type):
    return {
        "export_sales_excel": "外运外销表",
        "contract_excel": "合同登记表",
        "manual": "人工录入",
    }.get(source_type, source_type or "--")


def review_detail_rows(cur, start, end, kind):
    where = ["d.business_date >= %s", "d.business_date < %s"]
    params = [start, end]
    if kind == "pending":
        where.append("d.match_status <> 'matched'")
    elif kind == "contract":
        where.append("d.classification IN ('export_wholesale', 'local_wholesale', 'unknown')")
    query = f"""
        SELECT
            d.id,
            d.batch_id,
            d.business_date,
            d.raw_entity_name,
            d.standard_entity_name,
            d.entity_type,
            d.entity_code,
            d.raw_product_name,
            p.product_code,
            p.display_name,
            d.classification,
            d.weight_tons,
            d.match_status,
            c.contract_no,
            c.effective_start,
            c.effective_end,
            cr.amount_cny
        FROM wholesale_detail_records d
        LEFT JOIN products p ON p.id = d.product_id
        LEFT JOIN contracts c ON c.id = d.contract_id
        LEFT JOIN calculation_results cr ON cr.detail_id = d.id
        WHERE {" AND ".join(where)}
        ORDER BY
            CASE d.match_status WHEN 'needs_review' THEN 0 WHEN 'unmatched' THEN 1 ELSE 2 END,
            d.business_date DESC,
            d.id DESC
        LIMIT 300
    """
    cur.execute(query, params)
    return [
        {
            "id": row[0],
            "batchId": row[1],
            "date": row[2].isoformat() if row[2] else None,
            "entity": row[3],
            "standardEntity": row[4],
            "entityType": row[5],
            "entityCode": row[6],
            "product": row[7],
            "productCode": row[8],
            "standardProduct": row[9],
            "classification": row[10],
            "tons": round(float(row[11] or 0), 3),
            "matchStatus": row[12],
            "contractNo": row[13],
            "effectiveStart": row[14].isoformat() if row[14] else None,
            "effectiveEnd": row[15].isoformat() if row[15] else None,
            "amountCny": round(float(row[16] or 0), 2),
        }
        for row in cur.fetchall()
    ]


def review_issue_rows(cur, start, end):
    cur.execute(
        """
        SELECT
            qi.id,
            qi.issue_type,
            qi.severity,
            qi.issue_date,
            qi.message,
            qi.related_table,
            qi.related_id,
            d.raw_entity_name,
            d.raw_product_name,
            d.classification,
            d.match_status,
            d.vehicle_no,
            d.weight_tons
        FROM data_quality_issues qi
        LEFT JOIN wholesale_detail_records d
          ON qi.related_table = 'wholesale_detail_records'
         AND d.id = qi.related_id
        WHERE qi.status = 'open'
          AND (qi.issue_date IS NULL OR (qi.issue_date >= %s AND qi.issue_date < %s))
        ORDER BY
          CASE qi.severity WHEN 'error' THEN 0 WHEN 'warning' THEN 1 ELSE 2 END,
          qi.issue_date DESC,
          qi.id DESC
        LIMIT 300
        """,
        (start, end),
    )
    return [
        {
            "id": row[0],
            "issueType": row[1],
            "severity": row[2],
            "date": row[3].isoformat() if row[3] else None,
            "message": row[4],
            "relatedTable": row[5],
            "detailId": row[6],
            "entity": row[7],
            "product": row[8],
            "classification": row[9],
            "matchStatus": row[10],
            "vehicleNo": row[11],
            "tons": round(float(row[12] or 0), 3),
        }
        for row in cur.fetchall()
    ]


# ── DEPRECATED: HTML-serving routes ──
# These are now served by Nginx (React SPA).  Keep until Nginx fully verified.
# After confirming Nginx production, remove these routes — Flask should only
# ── React SPA routes ──
# Serve React frontend for all main page routes
@app.route("/")
def index():
    return send_from_directory(str(FRONTEND_DIST), "index.html")


@app.route("/login")
def login_page():
    return send_from_directory(str(FRONTEND_DIST), "index.html")


@app.route("/admin")
def admin_page():
    return send_from_directory(str(FRONTEND_DIST), "index.html")


@app.route("/wholesale")
@app.route("/wholesale/")
@app.route("/wholesale/<section>")
def wholesale_page(section=None):
    return send_from_directory(str(FRONTEND_DIST), "index.html")


@app.route("/retail")
@app.route("/retail/")
@app.route("/retail/<path:page>")
def retail_page(page=None):
    return send_from_directory(str(FRONTEND_DIST), "index.html")


@app.route("/api/retail/pages")
def api_retail_pages():
    base = "/retail"
    return jsonify({
        "baseUrl": base,
        "pages": [
            {"key": "dashboard", "label": "零售主页", "url": f"{base}"},
            {"key": "monthly", "label": "月度", "url": f"{base}/monthly"},
            {"key": "region", "label": "区域", "url": f"{base}/region"},
            {"key": "station", "label": "站点", "url": f"{base}/station"},
            {"key": "resource", "label": "油品", "url": f"{base}/resource"},
            {"key": "report", "label": "报表", "url": f"{base}/report"},
            {"key": "alert", "label": "预警", "url": f"{base}/alert"},
        ],
    })


@app.route("/retail-app/")
@app.route("/retail-app/")
@app.route("/retail-app/<path:filename>")
def retail_app_file(filename=""):
    if not filename or filename.startswith("api/"):
        if filename.startswith("api/"):
            return proxy_oilstation_api(filename[4:])
        return redirect("/retail")
    ext = Path(filename).suffix.lower()
    if ext in {".html", ".htm"}:
        return send_from_directory(str(FRONTEND_DIST), "index.html")
    response = send_from_directory(OILSTATION_ROOT, filename)
    if ext in {".css", ".js", ".json", ".png", ".jpg", ".jpeg", ".webp", ".svg", ".woff", ".woff2"}:
        response.headers["Cache-Control"] = "public, max-age=86400"
    return response


def proxy_oilstation_api(api_path):
    if api_path.split("/", 1)[0] not in RETAIL_API_PREFIXES:
        return jsonify({"error": "not_found"}), 404

    headers = {}
    if request.content_type:
        headers["Content-Type"] = request.content_type
    if request.headers.get("Accept"):
        headers["Accept"] = request.headers["Accept"]

    data = None if request.method in {"GET", "HEAD"} else request.get_data()
    response = open_oilstation(
        f"/api/{api_path}",
        method=request.method,
        query_string=request.query_string.decode("utf-8", errors="ignore"),
        data=data,
        headers=headers,
    )

    content_type = response.headers.get("Content-Type", "")
    if response.status_code < 400 and "application/json" in content_type and request.method == "GET":
        top = api_path.split("/", 1)[0]
        if top in {"dashboard", "monthly", "station", "stations", "regions"}:
            payload = response.get_json(silent=True)
            if payload is not None:
                payload = enrich_retail_station_statuses(payload)
                return jsonify(payload), response.status_code

    proxied = Response(
        response.get_data(),
        status=response.status_code,
        content_type=response.headers.get("Content-Type", "application/octet-stream"),
    )
    for header in ("Content-Disposition", "Cache-Control"):
        value = response.headers.get(header)
        if value:
            proxied.headers[header] = value
    return proxied


@app.route("/api/<path:api_path>", methods=["GET", "POST", "PUT", "PATCH", "DELETE"])
def retail_api_proxy(api_path):
    return proxy_oilstation_api(api_path)


@app.route("/assets/<path:filename>")
def shared_asset(filename):
    allowed = {"echarts.min.js", "kgz_merged.json"}
    if filename.startswith("kgz-map/"):
        response = send_from_directory(SHARED_ASSETS_ROOT, filename)
        response.headers["Cache-Control"] = "public, max-age=604800"
        return response
    if filename not in allowed:
        return jsonify({"error": "not_found"}), 404
    response = send_from_directory(OILSTATION_ROOT, filename)
    response.headers["Cache-Control"] = "public, max-age=604800"
    return response


# ── React SPA static assets ──
@app.route("/_next/<path:filename>")
def react_static(filename):
    response = send_from_directory(str(FRONTEND_DIST / "_next"), filename)
    response.headers["Cache-Control"] = "public, max-age=31536000, immutable"
    return response


# ── Root-level static files used by React SPA (echarts, etc.) ──
@app.route("/echarts.min.js")
def echarts_js():
    response = send_from_directory(str(OILSTATION_ROOT), "echarts.min.js")
    response.headers["Cache-Control"] = "public, max-age=604800"
    return response


# ── SPA catch-all ──
# Any route not matched above serves the React SPA index.html
@app.route("/<path:path>")
def spa_catchall(path):
    # Don't catch API or file routes
    if path.startswith("api/") or path.startswith("retail-app/"):
        return jsonify({"error": "not_found"}), 404
    # Try serving from OilStation for legacy static assets
    oilstation_file = OILSTATION_ROOT / path
    if oilstation_file.is_file():
        response = send_from_directory(str(OILSTATION_ROOT), path)
        response.headers["Cache-Control"] = "public, max-age=86400"
        return response
    index_path = FRONTEND_DIST / "index.html"
    if index_path.exists():
        return send_from_directory(str(FRONTEND_DIST), "index.html")
    return jsonify({"error": "frontend not built"}), 503


@app.route("/api/health")
def health():
    return jsonify({"ok": True, "project": "销售数据可视化"})


@app.route("/api/auth/status")
def api_auth_status():
    conn = get_connection()
    try:
        cur = conn.cursor()
        return jsonify({
            "hasUsers": has_users(cur),
            "authenticated": bool(session_user()),
            "user": session_user(),
        })
    finally:
        conn.close()


_ASSISTANT_DIGIT_TRANSLATION = str.maketrans("０１２３４５６７８９", "0123456789")


def _assistant_question_score(value):
    text = str(value or "")
    score = len(re.findall(r"\d{1,2}", text))
    for word in ("年", "月", "周", "今天", "昨天", "零售", "批发", "销量", "销售", "金额", "统计"):
        if word in text:
            score += 3
    return score


def _assistant_prepare_question_text(value):
    text = str(value or "").strip().translate(_ASSISTANT_DIGIT_TRANSLATION)
    candidates = [text]
    for encoding in ("cp1252", "latin1", "gbk", "cp936"):
        try:
            candidate = text.encode(encoding, errors="ignore").decode("utf-8", errors="ignore")
        except (LookupError, UnicodeError):
            continue
        if candidate and candidate not in candidates:
            candidates.append(candidate.translate(_ASSISTANT_DIGIT_TRANSLATION))
    return max(candidates, key=_assistant_question_score)


def _assistant_month_number(value):
    text = str(value or "").strip()
    if not text:
        return None
    try:
        month = int(text)
        return month if 1 <= month <= 12 else None
    except ValueError:
        pass
    month_map = {
        "一": 1,
        "二": 2,
        "两": 2,
        "三": 3,
        "四": 4,
        "五": 5,
        "六": 6,
        "七": 7,
        "八": 8,
        "九": 9,
        "十": 10,
        "十一": 11,
        "十二": 12,
    }
    return month_map.get(text)


def _assistant_period(question):
    today = date.today()
    text = _assistant_prepare_question_text(question)
    if "昨天" in text:
        start = today - timedelta(days=1)
        return "昨天", start, today
    if "今天" in text or "今日" in text:
        return "今天", today, today + timedelta(days=1)
    current_monday = today - timedelta(days=today.weekday())
    if "上周" in text or "上星期" in text:
        start = current_monday - timedelta(days=7)
        return "上周", start, current_monday
    if "本周" in text or "这周" in text:
        return "本周", current_monday, today + timedelta(days=1)
    if "上月" in text:
        first_this_month = today.replace(day=1)
        last_month_end = first_this_month
        if first_this_month.month == 1:
            start = date(first_this_month.year - 1, 12, 1)
        else:
            start = date(first_this_month.year, first_this_month.month - 1, 1)
        return "上月", start, last_month_end
    match = re.search(r"(20\d{2})\s*[年/-]\s*(\d{1,2})\s*月?", text)
    if match:
        year = int(match.group(1))
        month = int(match.group(2))
        if 1 <= month <= 12:
            start = date(year, month, 1)
            end = date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)
            return f"{year}年{month}月", start, end
    relative_month = re.search(r"(?:当年|今年|本年|当前年|这个年)\s*(\d{1,2}|一|二|两|三|四|五|六|七|八|九|十|十一|十二)\s*月", text)
    if not relative_month:
        relative_month = re.search(r"(?<!\d)(\d{1,2}|一|二|两|三|四|五|六|七|八|九|十|十一|十二)\s*月(?:份)?(?=.*(?:零售|批发|销量|销售|金额|数据|统计))", text)
    if not relative_month:
        relative_month = re.search(r"(?<!\d)(\d{1,2}|一|二|两|三|四|五|六|七|八|九|十|十一|十二)\s*月(?:份)?(?!\s*\d)", text)
    if relative_month:
        month = _assistant_month_number(relative_month.group(1))
        if month:
            start = date(today.year, month, 1)
            end = date(today.year + 1, 1, 1) if month == 12 else date(today.year, month + 1, 1)
            return f"{today.year}年{month}月", start, end
    if "本月" in text or "这个月" in text:
        start = today.replace(day=1)
        return "本月", start, today + timedelta(days=1)
    year_match = re.search(r"(20\d{2})\s*年(?!\s*\d)", text)
    if year_match:
        year = int(year_match.group(1))
        return f"{year}年", date(year, 1, 1), date(year + 1, 1, 1)
    if "年销量" in text or "年度" in text or "全年" in text:
        return f"{today.year}年", date(today.year, 1, 1), today + timedelta(days=1)
    start = today - timedelta(days=7)
    return "最近7天", start, today + timedelta(days=1)


def _assistant_period_display(label, start, end):
    try:
        end_day = end - timedelta(days=1)
        return f"{label}（{start.isoformat()} 至 {end_day.isoformat()}）"
    except Exception:
        return label


def _retail_density_map():
    oilstation_path = str(OILSTATION_ROOT)
    sys.path.insert(0, oilstation_path)
    try:
        from export_excel import RESOURCE_DENSITY
        return {int(key): float(value) for key, value in RESOURCE_DENSITY.items()}
    except Exception:
        return {102: 0.76, 103: 0.74, 104: 0.75, 106: 0.85, 108: 0.54, 109: 0.85, 110: 0.85}
    finally:
        sys.path = [item for item in sys.path if item != oilstation_path]


def _retail_snapshot_file(kind, key):
    prefix = {"daily": "day", "monthly": "monthly"}.get(kind, kind)
    return OILSTATION_ROOT / "cache_snapshots" / kind / f"{prefix}_{key}.json"


def _read_retail_snapshot(kind, key):
    path = _retail_snapshot_file(kind, key)
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None


def _report_extras_snapshot_path(ym):
    return OILSTATION_ROOT / "cache_snapshots" / "report_extras" / f"report_extras_{ym}.json"


def _retail_report_index():
    path = OILSTATION_ROOT / "exports" / ".monthly_report_index.json"
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        return data if isinstance(data, dict) else {}
    except (OSError, json.JSONDecodeError):
        return {}


def _retail_report_workbook_from_index(ym):
    item = _retail_report_index().get(ym) or {}
    filename = item.get("file")
    if not filename:
        return None
    path = OILSTATION_ROOT / "exports" / filename
    return path if path.exists() else None


def _station_key_from_cell(value):
    try:
        station_key = int(float(value))
    except (TypeError, ValueError):
        return None
    return station_key if station_key in retail_station_catalog_cached() else None


def _float_cell(value):
    if value is None or value == "":
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _retail_report_extras_from_workbook(ym):
    workbook_path = _retail_report_workbook_from_index(ym)
    if not workbook_path:
        return None
    try:
        from openpyxl import load_workbook
    except Exception:
        return None
    workbook = None
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        sheet_name = "总表" if "总表" in workbook.sheetnames else workbook.sheetnames[0]
        sheet = workbook[sheet_name]
        result = {
            "totalEndInventory": 0.0,
            "totalEndInventoryTons": 0.0,
            "stations": {},
            "source": "daily_report_workbook",
            "sourceFile": workbook_path.name,
            "sourceMtime": workbook_path.stat().st_mtime,
        }
        current_shop_key = None
        for row in sheet.iter_rows(min_row=5, values_only=True):
            station_key = _station_key_from_cell(row[0] if len(row) > 0 else None)
            if station_key is not None:
                current_shop_key = station_key
            if current_shop_key is None:
                continue
            product_name = str(row[3] if len(row) > 3 and row[3] is not None else "").strip()
            if not product_name or "合计" in product_name or "итого" in product_name.lower():
                continue
            end_inventory = _float_cell(row[23] if len(row) > 23 else None)
            end_inventory_tons = _float_cell(row[24] if len(row) > 24 else None)
            if not end_inventory and not end_inventory_tons:
                continue
            item = result["stations"].setdefault(str(current_shop_key), {
                "endInventory": 0.0,
                "endInventoryTons": 0.0,
            })
            item["endInventory"] += end_inventory
            item["endInventoryTons"] += end_inventory_tons
            result["totalEndInventory"] += end_inventory
            result["totalEndInventoryTons"] += end_inventory_tons
        result["totalEndInventory"] = round(result["totalEndInventory"], 2)
        result["totalEndInventoryTons"] = round(result["totalEndInventoryTons"], 3)
        for item in result["stations"].values():
            item["endInventory"] = round(item["endInventory"], 2)
            item["endInventoryTons"] = round(item["endInventoryTons"], 3)
        return result
    except Exception:
        return None
    finally:
        if workbook is not None:
            workbook.close()


def _retail_station_tons_from_report(ym):
    workbook_path = _retail_report_workbook_from_index(ym)
    if not workbook_path:
        return {}
    try:
        from openpyxl import load_workbook
    except Exception:
        return {}
    workbook = None
    station_tons = {}
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        sheet_name = "总表" if "总表" in workbook.sheetnames else workbook.sheetnames[0]
        sheet = workbook[sheet_name]
        current_shop_key = None
        for row in sheet.iter_rows(min_row=5, values_only=True):
            station_key = _station_key_from_cell(row[0] if len(row) > 0 else None)
            if station_key is not None:
                current_shop_key = station_key
            if current_shop_key is None:
                continue
            product_name = str(row[3] if len(row) > 3 and row[3] is not None else "").strip()
            if not product_name or "合计" in product_name or "итого" in product_name.lower():
                continue
            station_tons[current_shop_key] = station_tons.get(current_shop_key, 0.0) + _float_cell(row[20] if len(row) > 20 else None)
    except Exception:
        return {}
    finally:
        if workbook is not None:
            workbook.close()
    return {key: round(value, 3) for key, value in station_tons.items()}


def _read_retail_report_extras(ym):
    path = _report_extras_snapshot_path(ym)
    data = None
    if path.exists():
        try:
            loaded = json.loads(path.read_text(encoding="utf-8-sig"))
            data = loaded if isinstance(loaded, dict) else None
        except (OSError, json.JSONDecodeError):
            data = None
    workbook_path = _retail_report_workbook_from_index(ym)
    if data and workbook_path and data.get("source") == "daily_report_workbook":
        stale = data.get("sourceFile") != workbook_path.name
        source_mtime = data.get("sourceMtime")
        try:
            stale = stale or (source_mtime is not None and workbook_path.stat().st_mtime > float(source_mtime) + 0.5)
        except (OSError, TypeError, ValueError):
            stale = True
        if stale:
            data = None
    if data is None:
        data = _retail_report_extras_from_workbook(ym)
    return data


def _month_after(year, month):
    return (year + 1, 1) if month == 12 else (year, month + 1)


def _full_month_ym(start, end):
    if start.day != 1:
        return None
    next_year, next_month = _month_after(start.year, start.month)
    if end == date(next_year, next_month, 1):
        return f"{start.year}-{start.month:02d}"
    return None


def _retail_inventory_for_period(start, end):
    months = []
    cursor = start
    while cursor < end:
        ym = _full_month_ym(cursor, min(end, date(*_month_after(cursor.year, cursor.month), 1)))
        if not ym:
            break
        months.append(ym)
        next_year, next_month = _month_after(cursor.year, cursor.month)
        cursor = date(next_year, next_month, 1)
    if not months:
        ym = _full_month_ym(start, end)
        months = [ym] if ym else []
    result = {
        "totalEndInventory": 0.0,
        "totalEndInventoryTons": 0.0,
        "stations": {},
        "sourceFile": "",
        "month": months[-1] if months else "",
    }
    for ym in months:
        extras = _read_retail_report_extras(ym)
        if not extras:
            continue
        result["totalEndInventory"] = float(extras.get("totalEndInventory") or 0)
        result["totalEndInventoryTons"] = float(extras.get("totalEndInventoryTons") or 0)
        result["stations"] = extras.get("stations") or {}
        result["sourceFile"] = extras.get("sourceFile") or ""
        result["month"] = ym
    if not result["month"]:
        return None
    return result


def _normalize_retail_snapshot_payload(payload, ym=None):
    if not payload:
        return None
    density = _retail_density_map()
    if ym is None:
        ym = payload.get("month")
    station_tons_from_report = _retail_station_tons_from_report(ym) if ym else {}
    summary = payload.get("summary") or {}
    resources = []
    for item in payload.get("byResource") or []:
        resource_key = int(item.get("resourceKey") or item.get("key") or 0)
        volume = float(item.get("volume") or 0)
        amount = float(item.get("amount") or 0)
        resources.append({
            "resourceKey": resource_key,
            "key": resource_key,
            "name": item.get("name") or "--",
            "volume": round(volume, 2),
            "tons": round(volume * density.get(resource_key, 0) / 1000, 3),
            "amount": round(amount, 2),
        })
    resource_total_liters = round(sum(float(item.get("volume") or 0) for item in resources), 2)
    resource_total_tons = round(sum(float(item.get("tons") or 0) for item in resources), 3)
    resource_total_sales = round(sum(float(item.get("amount") or 0) for item in resources), 2)
    stations = []
    for item in payload.get("stations") or []:
        shop_key = item.get("shopKey")
        try:
            shop_key_int = int(shop_key)
        except (TypeError, ValueError):
            shop_key_int = None
        tons_value = item.get("tons")
        if (tons_value is None or not float(tons_value or 0)) and shop_key_int in station_tons_from_report:
            tons_value = station_tons_from_report[shop_key_int]
        stations.append({
            "shopKey": shop_key,
            "name": item.get("name") or "--",
            "short": item.get("short") or "",
            "region": item.get("region") or "--",
            "volume": round(float(item.get("volume") or 0), 2),
            "tons": round(float(tons_value or 0), 3) if tons_value is not None else None,
            "amount": round(float(item.get("amount") or 0), 2),
        })
    region_tons_from_stations = {}
    for item in stations:
        if item.get("tons") is not None:
            region = item.get("region") or "--"
            region_tons_from_stations[region] = region_tons_from_stations.get(region, 0.0) + float(item.get("tons") or 0)
    raw_regions = payload.get("regions") or []
    if isinstance(raw_regions, dict):
        region_iter = ({**value, "region": key} for key, value in raw_regions.items())
    else:
        region_iter = raw_regions
    regions = [
        {
            "region": item.get("region") or item.get("name") or "--",
            "volume": round(float(item.get("volume") or 0), 2),
            "tons": round(float(item.get("tons") or region_tons_from_stations.get(item.get("region") or item.get("name") or "--") or 0), 3),
            "amount": round(float(item.get("amount") or 0), 2),
        }
        for item in region_iter
    ]
    stations.sort(key=lambda item: item["amount"], reverse=True)
    regions.sort(key=lambda item: item["amount"], reverse=True)
    return {
        "source": payload.get("source") or "OilStationSnapshot",
        "summary": {
            "totalTxn": int(summary.get("totalTxn") or 0),
            "totalLiters": round(float(summary.get("totalLiters") or resource_total_liters or 0), 2),
            "totalTons": round(float(summary.get("totalTons") or resource_total_tons or 0), 3),
            "totalSales": round(float(summary.get("totalSales") or resource_total_sales or 0), 2),
            "totalSalesCny": round(float(summary.get("totalSalesCny") or 0), 2),
        },
        "byResource": resources,
        "stations": stations,
        "regions": regions,
    }


def _merge_retail_snapshot_payloads(payloads):
    merged = {"summary": {"totalTxn": 0, "totalLiters": 0, "totalTons": 0, "totalSales": 0, "totalSalesCny": 0}, "byResource": [], "stations": [], "regions": []}
    resource_map = {}
    station_map = {}
    region_map = {}
    for payload in payloads:
        normalized = _normalize_retail_snapshot_payload(payload, (payload or {}).get("month"))
        if not normalized:
            continue
        summary = normalized["summary"]
        for key in ("totalTxn", "totalLiters", "totalTons", "totalSales", "totalSalesCny"):
            merged["summary"][key] += float(summary.get(key) or 0)
        for item in normalized.get("byResource") or []:
            bucket = resource_map.setdefault(item.get("resourceKey"), {**item, "volume": 0, "tons": 0, "amount": 0})
            bucket["volume"] += float(item.get("volume") or 0)
            bucket["tons"] += float(item.get("tons") or 0)
            bucket["amount"] += float(item.get("amount") or 0)
        for item in normalized.get("stations") or []:
            key = item.get("shopKey") or item.get("name")
            bucket = station_map.setdefault(key, {**item, "volume": 0, "tons": 0, "amount": 0})
            bucket["volume"] += float(item.get("volume") or 0)
            bucket["tons"] += float(item.get("tons") or 0)
            bucket["amount"] += float(item.get("amount") or 0)
        for item in normalized.get("regions") or []:
            key = item.get("region")
            bucket = region_map.setdefault(key, {**item, "volume": 0, "tons": 0, "amount": 0})
            bucket["volume"] += float(item.get("volume") or 0)
            bucket["tons"] += float(item.get("tons") or 0)
            bucket["amount"] += float(item.get("amount") or 0)
    merged["byResource"] = sorted(
        ({**item, "volume": round(item["volume"], 2), "tons": round(item["tons"], 3), "amount": round(item["amount"], 2)} for item in resource_map.values()),
        key=lambda item: item["volume"],
        reverse=True,
    )
    merged["summary"] = {key: round(value, 3 if key == "totalTons" else 2) for key, value in merged["summary"].items()}
    if merged["byResource"]:
        resource_tons = round(sum(float(item.get("tons") or 0) for item in merged["byResource"]), 3)
        summary_tons = float(merged["summary"].get("totalTons") or 0)
        if resource_tons and (not summary_tons or abs(summary_tons - resource_tons) > max(1.0, resource_tons * 0.05)):
            merged["summary"]["totalTons"] = resource_tons
    merged["stations"] = sorted(
        ({**item, "volume": round(item["volume"], 2), "tons": round(item["tons"], 3), "amount": round(item["amount"], 2)} for item in station_map.values()),
        key=lambda item: item["amount"],
        reverse=True,
    )
    merged["regions"] = sorted(
        ({**item, "volume": round(item["volume"], 2), "tons": round(item["tons"], 3), "amount": round(item["amount"], 2)} for item in region_map.values()),
        key=lambda item: item["amount"],
        reverse=True,
    )
    return merged


def _empty_retail_snapshot_payload(start, end, source="OilStationSnapshot"):
    return {
        "source": source,
        "summary": {
            "totalTxn": 0,
            "totalLiters": 0.0,
            "totalTons": 0.0,
            "totalSales": 0.0,
            "totalSalesCny": 0.0,
            "queryStart": start.isoformat() if hasattr(start, "isoformat") else "",
            "queryEnd": end.isoformat() if hasattr(end, "isoformat") else "",
        },
        "byResource": [],
        "stations": [],
        "regions": [],
    }


def _retail_snapshot_for_period(start, end):
    today = date.today()
    if end <= start:
        return _empty_retail_snapshot_payload(start, end)

    if start.day == 1 and start.year == today.year and start.month == today.month and end >= today:
        ym = f"{start.year}-{start.month:02d}"
        current_month = _normalize_retail_snapshot_payload(_read_retail_snapshot("monthly", ym), ym)
        if current_month:
            return current_month
    if start.day == 1 and end == (date(start.year + 1, 1, 1) if start.month == 12 else date(start.year, start.month + 1, 1)):
        ym = f"{start.year}-{start.month:02d}"
        return _normalize_retail_snapshot_payload(_read_retail_snapshot("monthly", ym), ym)
    if start.month == 1 and start.day == 1 and end == date(start.year + 1, 1, 1):
        payloads = [
            _read_retail_snapshot("monthly", f"{start.year}-{month:02d}")
            for month in range(1, 13)
        ]
        payloads = [item for item in payloads if item]
        return _merge_retail_snapshot_payloads(payloads) if payloads else None

    # Daily snapshots are generated only after a business day is complete.
    # If the requested range includes today, use completed days only and avoid
    # falling back to the live retail DB for incomplete shifts.
    effective_end = min(end, today)
    if effective_end <= start:
        return _empty_retail_snapshot_payload(start, effective_end)

    dates = []
    current = start
    while current < effective_end:
        payload = _read_retail_snapshot("daily", current.isoformat())
        if not payload:
            return None
        dates.append(payload)
        current += timedelta(days=1)
    return _merge_retail_snapshot_payloads(dates)


def _combined_data_from_snapshots(cur, year, month, start, end):
    retail_snapshot = _retail_snapshot_for_period(start, end)
    if not retail_snapshot:
        return summarize_combined(cur, year, month)
    wholesale = summarize_wholesale(cur, year, month)
    retail_summary = retail_snapshot.get("summary") or {}
    retail = {
        "liters": round(float(retail_summary.get("totalLiters") or 0), 2),
        "tons": round(float(retail_summary.get("totalTons") or 0), 3),
        "amountKgs": round(float(retail_summary.get("totalSales") or 0), 2),
        "amountCny": round(float(retail_summary.get("totalSalesCny") or 0), 2),
        "source": "snapshot",
    }
    combined = {
        "salesTons": round(retail["tons"] + wholesale["salesTons"], 3),
        "salesAmountCny": round(retail["amountCny"] + wholesale["salesAmountCny"], 2),
        "transferTons": wholesale["transferTons"],
    }
    return {"retail": retail, "wholesale": wholesale, "combined": combined}


def _format_assistant_retail_answer(label, payload):
    summary = payload.get("summary") or {}
    question = payload.get("question") or ""
    rows = [
        item for item in payload.get("byResource") or []
        if float(item.get("volume") or 0) or float(item.get("amount") or 0)
    ]
    lines = [
        f"{label}零售销量统计如下：",
        f"合计销量 {float(summary.get('totalLiters') or 0):,.2f} 升，折算 {float(summary.get('totalTons') or 0):,.3f} 吨，合计金额 {float(summary.get('totalSales') or 0):,.2f} 索姆。",
    ]
    if rows:
        lines.append("各油品明细：")
        for item in rows:
            lines.append(
                f"{item.get('name') or '--'}：{float(item.get('volume') or 0):,.2f} 升，"
                f"{float(item.get('tons') or 0):,.3f} 吨，"
                f"{float(item.get('amount') or 0):,.2f} 索姆"
            )
    else:
        lines.append("该区间暂未查询到已结束班组的销售数据。")
    if any(word in question for word in ("站点", "加油站", "排行", "排名", "最好")) and payload.get("stations"):
        lines.append("站点排行：")
        for item in payload.get("stations", [])[:10]:
            lines.append(
                f"{item.get('name') or '--'}：{float(item.get('volume') or 0):,.2f} 升，"
                f"{float(item.get('tons') or 0):,.3f} 吨，{float(item.get('amount') or 0):,.2f} 索姆"
            )
    if any(word in question for word in ("区域", "州", "地区")) and payload.get("regions"):
        lines.append("区域排行：")
        for item in payload.get("regions", [])[:8]:
            lines.append(
                f"{item.get('region') or '--'}：{float(item.get('volume') or 0):,.2f} 升，"
                f"{float(item.get('tons') or 0):,.3f} 吨，{float(item.get('amount') or 0):,.2f} 索姆"
            )
    return "\n".join(lines)


def _assistant_has_any(question, words):
    return any(word in question for word in words)


def _append_rank_lines(lines, title, rows, limit=5, amount_label="元"):
    clean_rows = [item for item in rows or [] if float(item.get("tons") or item.get("volume") or 0)]
    if not clean_rows:
        return
    lines.append(title)
    for item in clean_rows[:limit]:
        name = item.get("name") or item.get("region") or item.get("product") or "--"
        if "volume" in item:
            lines.append(f"{name}：{float(item.get('volume') or 0):,.2f} 升，{float(item.get('amount') or 0):,.2f} 索姆")
        elif "amountCny" in item:
            lines.append(f"{name}：{float(item.get('tons') or 0):,.3f} 吨，{float(item.get('amountCny') or 0):,.2f} {amount_label}")
        else:
            lines.append(f"{name}：{float(item.get('tons') or 0):,.3f} 吨")


def _assistant_table(title, columns, rows):
    return {
        "title": title,
        "columns": columns,
        "rows": [
            ["" if value is None else str(value) for value in row]
            for row in rows
        ],
    }


def _fmt_amount(value, unit=""):
    suffix = f" {unit}" if unit else ""
    return f"{float(value or 0):,.2f}{suffix}"


def _fmt_volume(value):
    return f"{float(value or 0):,.2f} 升"


def _fmt_tons(value):
    return f"{float(value or 0):,.3f} 吨"


def _assistant_retail_tables(label, payload, question):
    tables = []
    summary = payload.get("summary") or {}
    resources = [item for item in payload.get("byResource") or [] if float(item.get("volume") or 0) or float(item.get("amount") or 0)]
    if resources:
        resource_rows = [
            [item.get("name") or "--", _fmt_volume(item.get("volume")), _fmt_tons(item.get("tons")), _fmt_amount(item.get("amount"), "索姆")]
            for item in resources
        ]
        resource_rows.append([
            "合计",
            _fmt_volume(summary.get("totalLiters")),
            _fmt_tons(summary.get("totalTons")),
            _fmt_amount(summary.get("totalSales"), "索姆"),
        ])
        tables.append(_assistant_table(
            f"{label}零售油品明细",
            ["油品", "销量", "吨位", "销售金额"],
            resource_rows,
        ))
    if _assistant_has_any(question, ("站点", "加油站", "排行", "排名", "最好")) and payload.get("stations"):
        tables.append(_assistant_table(
            f"{label}站点排行",
            ["站点", "区域", "销量", "吨位", "销售金额"],
            [
                [item.get("name") or "--", item.get("region") or "--", _fmt_volume(item.get("volume")), _fmt_tons(item.get("tons")), _fmt_amount(item.get("amount"), "索姆")]
                for item in payload.get("stations", [])[:12]
            ],
        ))
    if _assistant_has_any(question, ("区域", "州", "地区")) and payload.get("regions"):
        tables.append(_assistant_table(
            f"{label}区域排行",
            ["区域", "销量", "吨位", "销售金额"],
            [
                [item.get("region") or "--", _fmt_volume(item.get("volume")), _fmt_tons(item.get("tons")), _fmt_amount(item.get("amount"), "索姆")]
                for item in payload.get("regions", [])[:10]
            ],
        ))
    return tables


def _assistant_retail_inventory_answer(label, data):
    return (
        f"{label}零售期末库存如下：\n"
        f"合计库存 {float(data.get('totalEndInventory') or 0):,.2f} 升，"
        f"折算 {float(data.get('totalEndInventoryTons') or 0):,.3f} 吨。\n"
        f"数据来源：{data.get('sourceFile') or '日销表'}。"
    )


def _assistant_retail_inventory_tables(label, data):
    catalog = retail_station_catalog_cached()
    station_rows = []
    for raw_shop_key, item in (data.get("stations") or {}).items():
        try:
            shop_key = int(raw_shop_key)
        except (TypeError, ValueError):
            continue
        info = catalog.get(shop_key, {})
        inventory = float(item.get("endInventory") or 0)
        tons = float(item.get("endInventoryTons") or 0)
        if not inventory and not tons:
            continue
        station_rows.append([
            info.get("name") or f"Station-{shop_key}",
            info.get("region") or "--",
            _fmt_volume(inventory),
            _fmt_tons(tons),
        ])
    station_rows.sort(key=lambda row: float(str(row[2]).replace(",", "").split()[0] or 0), reverse=True)
    station_rows.append([
        "合计",
        "全部区域",
        _fmt_volume(data.get("totalEndInventory")),
        _fmt_tons(data.get("totalEndInventoryTons")),
    ])
    return [_assistant_table(
        f"{label}零售期末库存",
        ["加油站", "区域", "期末库存", "折算吨位"],
        station_rows,
    )]


def _assistant_wholesale_tables(label, data, question):
    cards = data.get("cards") or {}
    labels = {"export_wholesale": "出口批发", "local_wholesale": "本地批发", "unknown": "待审核批发", "transfer": "外运"}
    sales_keys = ("export_wholesale", "local_wholesale", "unknown")
    sales_total_tons = sum(float((cards.get(key) or {}).get("tons") or 0) for key in sales_keys)
    sales_total_amount = sum(float((cards.get(key) or {}).get("amountCny") or 0) for key in ("export_wholesale", "local_wholesale"))
    sales_total_review = sum(int((cards.get(key) or {}).get("reviewRows") or 0) for key in sales_keys)
    category_rows = [
        [
            labels.get(key, key),
            _fmt_tons((cards.get(key) or {}).get("tons")),
            "--" if key in ("unknown", "transfer") else _fmt_amount((cards.get(key) or {}).get("amountCny"), "元"),
            (cards.get(key) or {}).get("entities") or 0,
            (cards.get(key) or {}).get("reviewRows") or 0,
        ]
        for key in ("export_wholesale", "local_wholesale", "unknown", "transfer")
    ]
    category_rows.append(["销售合计", _fmt_tons(sales_total_tons), _fmt_amount(sales_total_amount, "元"), "--", sales_total_review])
    tables = [_assistant_table(
        f"{label}批发分类统计",
        ["分类", "吨数", "金额", "客户/去向数", "待复核"],
        category_rows,
    )]
    rankings = data.get("customerRankings") or {}
    if _assistant_has_any(question, ("客户", "排行", "排名", "top", "TOP")):
        for key in ("export_wholesale", "local_wholesale", "unknown"):
            rows = rankings.get(key) or []
            if rows:
                tables.append(_assistant_table(
                    f"{labels[key]}客户排行",
                    ["客户", "吨数", "金额", "记录数"],
                    [[item.get("name") or "--", _fmt_tons(item.get("tons")), _fmt_amount(item.get("amountCny"), "元"), item.get("rows") or 0] for item in rows[:10]],
                ))
    if _assistant_has_any(question, ("油品", "产品", "结构", "品类")):
        products = data.get("productRankings") or {}
        for key in ("export_wholesale", "local_wholesale", "unknown", "transfer"):
            rows = products.get(key) or []
            if rows:
                tables.append(_assistant_table(
                    f"{labels[key]}油品排行",
                    ["油品", "吨数", "金额"],
                    [[item.get("name") or "--", _fmt_tons(item.get("tons")), "--" if key in ("unknown", "transfer") else _fmt_amount(item.get("amountCny"), "元")] for item in rows[:10]],
                ))
    if "外运" in question:
        transfer_summary = data.get("transferSummary") or {}
        if transfer_summary.get("byRegion"):
            tables.append(_assistant_table(
                f"{label}外运区域汇总",
                ["区域", "吨数", "记录数"],
                [[item.get("region") or "--", _fmt_tons(item.get("tons")), item.get("rows") or 0] for item in transfer_summary.get("byRegion", [])[:10]],
            ))
    return tables


def _assistant_combined_tables(label, data):
    retail = data.get("retail") or {}
    wholesale = data.get("wholesale") or {}
    combined = data.get("combined") or {}
    return [_assistant_table(
        f"{label}合并统计",
        ["板块", "销量/销售吨数", "销售金额"],
        [
            ["合计", _fmt_tons(combined.get("salesTons")), _fmt_amount(combined.get("salesAmountCny"), "元")],
            ["零售", _fmt_tons(retail.get("tons")), _fmt_amount(retail.get("amountCny"), "元")],
            ["批发销售", _fmt_tons(wholesale.get("salesTons")), _fmt_amount(wholesale.get("salesAmountCny"), "元")],
            ["外运", _fmt_tons(combined.get("transferTons")), "--"],
        ],
    )]


def _assistant_review_tables(label, data):
    review = data.get("review") or {}
    status = review.get("status") or {}
    tables = []
    if status:
        tables.append(_assistant_table(
            f"{label}匹配状态",
            ["状态", "记录数", "吨数"],
            [[key or "--", value.get("rows") or 0, _fmt_tons(value.get("tons"))] for key, value in status.items()],
        ))
    if review.get("topNeedsReview"):
        tables.append(_assistant_table(
            f"{label}待复核数据",
            ["客户/去向", "油品", "记录数", "吨数"],
            [[item.get("entity") or "--", item.get("product") or "--", item.get("rows") or 0, _fmt_tons(item.get("tons"))] for item in review["topNeedsReview"][:10]],
        ))
    issues = data.get("issues") or []
    if issues:
        tables.append(_assistant_table(
            f"{label}数据审核",
            ["类型", "级别", "日期", "提示"],
            [[item.get("type") or "--", item.get("severity") or "--", item.get("date") or "--", item.get("message") or ""] for item in issues[:12]],
        ))
    return tables


def _assistant_year_month_from_period(start, end):
    if start.month == 1 and start.day == 1 and end == date(start.year + 1, 1, 1):
        return start.year, None
    if start.day == 1:
        next_month = date(start.year + 1, 1, 1) if start.month == 12 else date(start.year, start.month + 1, 1)
        if end == next_month:
            return start.year, start.month
    return start.year, start.month if start.year == end.year and start.month == end.month else None


def _format_assistant_wholesale_answer(label, data, question):
    cards = data.get("cards") or {}
    export_card = cards.get("export_wholesale") or {}
    local_card = cards.get("local_wholesale") or {}
    pending_card = cards.get("unknown") or {}
    transfer_card = cards.get("transfer") or {}
    lines = [f"{label}批发数据如下："]
    if "出口" in question:
        lines.append(f"出口批发：{float(export_card.get('tons') or 0):,.3f} 吨，金额 {float(export_card.get('amountCny') or 0):,.2f} 元，客户数 {export_card.get('entities') or 0}。")
    elif "本地" in question:
        lines.append(f"本地批发：{float(local_card.get('tons') or 0):,.3f} 吨，金额 {float(local_card.get('amountCny') or 0):,.2f} 元，客户数 {local_card.get('entities') or 0}。")
    elif "外运" in question:
        lines.append(f"外运：{float(transfer_card.get('tons') or 0):,.3f} 吨，去向数 {transfer_card.get('entities') or 0}；外运仅统计吨数，不计算金额。")
    else:
        total_tons = float(export_card.get("tons") or 0) + float(local_card.get("tons") or 0) + float(pending_card.get("tons") or 0)
        total_amount = float(export_card.get("amountCny") or 0) + float(local_card.get("amountCny") or 0)
        lines.append(f"批发销售合计：{total_tons:,.3f} 吨，金额 {total_amount:,.2f} 元。")
        lines.append(f"出口批发：{float(export_card.get('tons') or 0):,.3f} 吨，{float(export_card.get('amountCny') or 0):,.2f} 元。")
        lines.append(f"本地批发：{float(local_card.get('tons') or 0):,.3f} 吨，{float(local_card.get('amountCny') or 0):,.2f} 元。")
        lines.append(f"待审核批发：{float(pending_card.get('tons') or 0):,.3f} 吨，数量已统计，金额待合同复核。")
        lines.append(f"外运：{float(transfer_card.get('tons') or 0):,.3f} 吨。")
    rankings = data.get("customerRankings") or {}
    key = "export_wholesale" if "出口" in question else ("local_wholesale" if "本地" in question else None)
    if key and rankings.get(key):
        _append_rank_lines(lines, "客户排行：", rankings[key], limit=8)
    elif _assistant_has_any(question, ("客户", "排行", "排名", "top", "TOP")):
        _append_rank_lines(lines, "出口批发客户排行：", rankings.get("export_wholesale"), limit=5)
        _append_rank_lines(lines, "本地批发客户排行：", rankings.get("local_wholesale"), limit=5)
    if _assistant_has_any(question, ("油品", "产品", "结构", "品类")):
        products = data.get("productRankings") or {}
        if "出口" in question:
            _append_rank_lines(lines, "出口批发油品排行：", products.get("export_wholesale"), limit=8)
        elif "本地" in question:
            _append_rank_lines(lines, "本地批发油品排行：", products.get("local_wholesale"), limit=8)
        elif "外运" in question:
            _append_rank_lines(lines, "外运油品排行：", products.get("transfer"), limit=8)
        else:
            _append_rank_lines(lines, "出口批发油品排行：", products.get("export_wholesale"), limit=5)
            _append_rank_lines(lines, "本地批发油品排行：", products.get("local_wholesale"), limit=5)
            _append_rank_lines(lines, "待审核批发油品排行：", products.get("unknown"), limit=5)
            _append_rank_lines(lines, "外运油品排行：", products.get("transfer"), limit=5)
    if "外运" in question and _assistant_has_any(question, ("区域", "去向", "详情", "油库", "加油站")):
        transfer_summary = data.get("transferSummary") or {}
        _append_rank_lines(lines, "外运区域汇总：", transfer_summary.get("byRegion"), limit=8)
        _append_rank_lines(lines, "外运去向排行：", data.get("transferTargets"), limit=8)
    return "\n".join(lines)


def _assistant_exchange_answer(cur, label, start, end):
    cur.execute(
        """
        SELECT rate_date, usd_kgs, cny_kgs, source, is_manual
        FROM exchange_rates
        WHERE rate_date >= %s AND rate_date < %s
        ORDER BY rate_date DESC
        LIMIT 31
        """,
        (start, end),
    )
    rows = cur.fetchall()
    if not rows:
        return f"{label}没有查询到汇率快照。"
    latest = rows[0]
    usd_vals = [float(row[1]) for row in rows if row[1] is not None]
    cny_vals = [float(row[2]) for row in rows if row[2] is not None]
    return "\n".join([
        f"{label}汇率数据如下：",
        f"最新日期 {latest[0]}：USD/KGS {float(latest[1] or 0):.4f}，CNY/KGS {float(latest[2] or 0):.4f}，来源 {'手工录入' if latest[4] else (latest[3] or 'NBKR')}。",
        f"区间平均：USD/KGS {(sum(usd_vals)/len(usd_vals)):.4f}，CNY/KGS {(sum(cny_vals)/len(cny_vals)):.4f}。",
    ])


def _assistant_exchange_tables(cur, label, start, end):
    cur.execute(
        """
        SELECT rate_date, usd_kgs, cny_kgs, source, is_manual
        FROM exchange_rates
        WHERE rate_date >= %s AND rate_date < %s
        ORDER BY rate_date DESC
        LIMIT 31
        """,
        (start, end),
    )
    rows = cur.fetchall()
    if not rows:
        return []
    usd_vals = [float(row[1]) for row in rows if row[1] is not None]
    cny_vals = [float(row[2]) for row in rows if row[2] is not None]
    table_rows = [
        [
            row[0],
            f"{float(row[1] or 0):.4f}",
            f"{float(row[2] or 0):.4f}",
            f"{(1 / float(row[2])):.4f}" if row[2] else "--",
            "手工录入" if row[4] else (row[3] or "NBKR"),
        ]
        for row in rows
    ]
    if usd_vals or cny_vals:
        cny_avg = (sum(cny_vals) / len(cny_vals)) if cny_vals else 0
        table_rows.append([
            "区间平均",
            f"{(sum(usd_vals) / len(usd_vals)):.4f}" if usd_vals else "--",
            f"{cny_avg:.4f}" if cny_vals else "--",
            f"{(1 / cny_avg):.4f}" if cny_avg else "--",
            "--",
        ])
    return [_assistant_table(
        f"{label}汇率快照",
        ["日期", "美元/索姆", "人民币/索姆", "索姆/人民币", "来源"],
        table_rows,
    )]


def _assistant_density_answer(cur):
    cur.execute(
        """
        SELECT p.display_name, d.density, d.effective_date, d.source
        FROM product_density_history d
        JOIN products p ON p.id = d.product_id
        JOIN (
            SELECT product_id, MAX(effective_date) effective_date
            FROM product_density_history
            GROUP BY product_id
        ) latest ON latest.product_id = d.product_id AND latest.effective_date = d.effective_date
        ORDER BY p.product_code
        """
    )
    rows = cur.fetchall()
    if not rows:
        return "暂未查询到密度维护记录。"
    lines = ["当前最新密度如下："]
    for name, density, effective_date, source in rows:
        lines.append(f"{name}：{float(density):.6f}，生效日期 {effective_date}，来源 {source or '--'}")
    return "\n".join(lines)


def _assistant_density_tables(cur):
    cur.execute(
        """
        SELECT p.display_name, d.density, d.effective_date, d.source
        FROM product_density_history d
        JOIN products p ON p.id = d.product_id
        JOIN (
            SELECT product_id, MAX(effective_date) effective_date
            FROM product_density_history
            GROUP BY product_id
        ) latest ON latest.product_id = d.product_id AND latest.effective_date = d.effective_date
        ORDER BY p.product_code
        """
    )
    rows = cur.fetchall()
    if not rows:
        return []
    return [_assistant_table(
        "当前密度",
        ["油品", "密度", "生效日期", "来源"],
        [[name, f"{float(density):.6f}", effective_date, source or "--"] for name, density, effective_date, source in rows],
    )]


def _assistant_contract_answer(cur, question):
    where = []
    if "出口" in question:
        where.append("c.sales_type = 'export_wholesale'")
    if "本地" in question:
        where.append("c.sales_type = 'local_wholesale'")
    sql_where = "WHERE " + " AND ".join(where) if where else ""
    cur.execute(
        f"""
        SELECT COUNT(*),
               SUM(CASE
                   WHEN (c.effective_start IS NULL OR c.effective_start <= CURDATE())
                    AND (c.effective_end IS NULL OR c.effective_end >= CURDATE())
                   THEN 1 ELSE 0 END),
               MIN(c.effective_start),
               MAX(c.effective_end)
        FROM contracts c
        {sql_where}
        """
    )
    count, active_count, min_from, max_to = cur.fetchone()
    cur.execute(
        f"""
        SELECT
            c.contract_no,
            COALESCE(cu.standard_name, c.raw_customer_name),
            COALESCE(p.display_name, c.raw_product_name),
            COALESCE(c.price_with_logistics, c.price_without_logistics, 0),
            c.currency_code,
            c.effective_start,
            c.effective_end,
            c.sales_type
        FROM contracts c
        LEFT JOIN customers cu ON cu.id = c.customer_id
        LEFT JOIN products p ON p.id = c.product_id
        {sql_where}
        ORDER BY COALESCE(c.effective_end, c.effective_start, c.signed_date) DESC, c.id DESC
        LIMIT 8
        """
    )
    rows = cur.fetchall()
    lines = [f"合同信息：共 {int(count or 0)} 条，其中有效合同 {int(active_count or 0)} 条；日期范围 {min_from or '--'} 至 {max_to or '--'}。"]
    if rows:
        lines.append("近期合同：")
        for no, customer, product, price, currency, valid_from, valid_to, sales_type in rows:
            label = "出口批发" if sales_type == "export_wholesale" else ("本地批发" if sales_type == "local_wholesale" else "其他")
            lines.append(f"{no or '--'}｜{customer or '--'}｜{product or '--'}｜{float(price or 0):,.2f} {currency}｜{valid_from or '--'} 至 {valid_to or '--'}｜{label}")
    return "\n".join(lines)


def _assistant_contract_tables(cur, question):
    where = []
    if "出口" in question:
        where.append("c.sales_type = 'export_wholesale'")
    if "本地" in question:
        where.append("c.sales_type = 'local_wholesale'")
    sql_where = "WHERE " + " AND ".join(where) if where else ""
    cur.execute(
        f"""
        SELECT
            c.contract_no,
            COALESCE(cu.standard_name, c.raw_customer_name),
            COALESCE(p.display_name, c.raw_product_name),
            COALESCE(c.price_with_logistics, c.price_without_logistics, 0),
            c.currency_code,
            c.effective_start,
            c.effective_end,
            c.sales_type
        FROM contracts c
        LEFT JOIN customers cu ON cu.id = c.customer_id
        LEFT JOIN products p ON p.id = c.product_id
        {sql_where}
        ORDER BY COALESCE(c.effective_end, c.effective_start, c.signed_date) DESC, c.id DESC
        LIMIT 12
        """
    )
    rows = []
    for no, customer, product, price, currency, valid_from, valid_to, sales_type in cur.fetchall():
        label = "出口批发" if sales_type == "export_wholesale" else ("本地批发" if sales_type == "local_wholesale" else "其他")
        rows.append([no or "--", customer or "--", product or "--", f"{float(price or 0):,.2f} {currency or ''}".strip(), valid_from or "--", valid_to or "--", label])
    return [_assistant_table("合同明细", ["合同号", "客户", "油品", "价格", "生效", "到期", "类型"], rows)] if rows else []


def _assistant_market_watch_answer(cur):
    ensure_market_watch_schema(cur)
    cur.execute(
        """
        SELECT market_region, title, summary, price_text, source_name, observed_at
        FROM market_watch_items
        WHERE is_active = 1
        ORDER BY observed_at DESC, priority DESC, id DESC
        LIMIT 12
        """
    )
    rows = cur.fetchall()
    if not rows:
        return "暂未查询到市场观察信息。可以在后台执行“立即获取市场观察”后再查看。", []
    latest = max((row[5] for row in rows if row[5]), default=None)
    lines = [f"市场观察信息如下（最新获取时间 {latest or '--'}）："]
    region_names = {
        "russia": "俄罗斯",
        "Russia": "俄罗斯",
        "uzbekistan": "乌兹别克",
        "Uzbekistan": "乌兹别克",
        "global": "国际",
        "Global": "国际",
    }
    table_rows = []
    for region, title, summary, price_text, source, observed_at in rows[:9]:
        item = chinese_market_item({
            "region": region,
            "title": title,
            "summary": summary,
            "sourceName": source,
            "observedAt": observed_at.isoformat() if observed_at else None,
        })
        label = region_names.get(region, region or "市场")
        lines.append(f"{label}｜{item.get('title') or '--'}")
        if item.get("summary"):
            lines.append(f"  {item['summary'][:180]}")
        if price_text:
            lines.append(f"  价格：{price_text}")
        source_label = market_source_display(item.get("sourceName") or source)
        lines.append(f"  来源：{source_label}，时间：{observed_at or '--'}")
        table_rows.append([
            label,
            item.get("title") or "--",
            price_text or "--",
            item.get("summary") or "--",
            source_label,
            observed_at or "--",
        ])
    return "\n".join(lines), [_assistant_table("市场观察", ["区域", "标题", "价格信息", "摘要", "来源", "日期"], table_rows)]


def _assistant_review_answer(label, data):
    review = data.get("review") or {}
    issues = data.get("issues") or []
    status = review.get("status") or {}
    matched = status.get("matched", {})
    need_review_rows = sum(int(item.get("rows") or 0) for key, item in status.items() if key != "matched")
    lines = [f"{label}审核信息如下："]
    lines.append(
        f"已匹配 {int(matched.get('rows') or 0)} 条，"
        f"待复核 {need_review_rows} 条，"
        f"开放数据问题 {len(issues)} 条。"
    )
    if review.get("topNeedsReview"):
        lines.append("待复核数据排行：")
        for item in review["topNeedsReview"][:8]:
            lines.append(
                f"{item.get('entity') or '--'}｜{item.get('product') or '--'}："
                f"{int(item.get('rows') or 0)} 条，{float(item.get('tons') or 0):,.3f} 吨"
            )
    if issues:
        lines.append("待处理问题：")
        for item in issues[:8]:
            lines.append(f"{item.get('type') or '--'}｜{item.get('severity') or '--'}｜{item.get('date') or '--'}：{item.get('message') or ''}")
    return "\n".join(lines)


def _assistant_combined_answer(label, data):
    retail = data.get("retail") or {}
    wholesale = data.get("wholesale") or {}
    combined = data.get("combined") or {}
    lines = [f"{label}合并统计如下："]
    lines.append(f"合并销售量：{float(combined.get('salesTons') or 0):,.3f} 吨，销售金额：{float(combined.get('salesAmountCny') or 0):,.2f} 元。")
    lines.append(f"零售：{float(retail.get('tons') or 0):,.3f} 吨，{float(retail.get('amountCny') or 0):,.2f} 元。")
    lines.append(f"批发销售：{float(wholesale.get('salesTons') or 0):,.3f} 吨，{float(wholesale.get('salesAmountCny') or 0):,.2f} 元。")
    lines.append(f"外运：{float(combined.get('transferTons') or 0):,.3f} 吨。")
    if wholesale.get("byProductCategory"):
        products = wholesale["byProductCategory"]
        _append_rank_lines(lines, "出口批发油品排行：", products.get("export_wholesale"), limit=3)
        _append_rank_lines(lines, "本地批发油品排行：", products.get("local_wholesale"), limit=3)
    return "\n".join(lines)


# ── Enhanced assistant: new query types ──

def _read_daily_trend_snapshot(ym):
    path = OILSTATION_ROOT / "cache_snapshots" / "monthly" / f"daily_{ym}.json"
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None


def _read_yearly_kpi_snapshot(year):
    path = OILSTATION_ROOT / "cache_snapshots" / "monthly" / f"kpi_yearly_{year}_summary.json"
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None


def _read_yearly_dashboard_snapshot(year):
    path = OILSTATION_ROOT / "cache_snapshots" / "monthly" / f"dashboard_yearly_{year}.json"
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None


def _assistant_daily_trend_answer(label, data):
    daily = data.get("daily") or []
    if not daily:
        return f"{label}趋势数据暂未生成。"
    volumes = [float(d.get("volume") or 0) for d in daily]
    amounts = [float(d.get("amount") or 0) for d in daily]
    total_vol = sum(volumes)
    total_amt = sum(amounts)
    avg_vol = total_vol / len(volumes) if volumes else 0
    max_idx = volumes.index(max(volumes)) if volumes else -1
    min_idx = volumes.index(min(volumes)) if volumes else -1
    lines = [f"{label}每日销量趋势（共 {len(daily)} 天数据）："]
    lines.append(f"合计销量 {_fmt_volume(total_vol)}，合计金额 {_fmt_amount(total_amt, '索姆')}")
    lines.append(f"日均销量 {_fmt_volume(avg_vol)}")
    if max_idx >= 0:
        lines.append(f"最高日：{daily[max_idx]['date']} 销量 {_fmt_volume(volumes[max_idx])} 金额 {_fmt_amount(amounts[max_idx], '索姆')}")
    if min_idx >= 0:
        lines.append(f"最低日：{daily[min_idx]['date']} 销量 {_fmt_volume(volumes[min_idx])} 金额 {_fmt_amount(amounts[min_idx], '索姆')}")
    return "\n".join(lines)


def _assistant_daily_trend_tables(label, data):
    daily = data.get("daily") or []
    if not daily:
        return []
    return [_assistant_table(
        f"{label}每日趋势",
        ["日期", "销量(升)", "金额(索姆)"],
        [
            [d.get("date", ""), _fmt_volume(d.get("volume")), _fmt_amount(d.get("amount"), "索姆")]
            for d in daily
        ],
    )]


def _assistant_yearly_kpi_answer(label, data):
    summary = data.get("summary") or {}
    comp = data.get("comparison") or {}
    lines = [f"{label}年度KPI如下："]
    lines.append(f"累计销量 {_fmt_volume(summary.get('totalVolume') or summary.get('totalLiters'))}")
    lines.append(f"累计金额 {_fmt_amount(summary.get('totalSales') or summary.get('totalAmount'), '索姆')}")
    lines.append(f"交易笔数 {int(summary.get('transactionCount') or 0)} 笔")
    lines.append(f"活跃站点 {int(summary.get('activeShops') or 0)} 个")
    yoy_vol = comp.get("yoyVolume")
    yoy_sales = comp.get("yoySales")
    if yoy_vol is not None:
        lines.append(f"同比销量：{_fmt_pct(yoy_vol)}")
    if yoy_sales is not None:
        lines.append(f"同比金额：{_fmt_pct(yoy_sales)}")
    return "\n".join(lines)


def _assistant_yearly_kpi_tables(label, data):
    summary = data.get("summary") or {}
    comp = data.get("comparison") or {}
    rows = [
        ["累计销量", _fmt_volume(summary.get("totalVolume") or summary.get("totalLiters"))],
        ["累计金额", _fmt_amount(summary.get("totalSales") or summary.get("totalAmount"), "索姆")],
        ["交易笔数", str(int(summary.get("transactionCount") or 0))],
        ["活跃站点", str(int(summary.get("activeShops") or 0))],
    ]
    yoy_vol = comp.get("yoyVolume")
    yoy_sales = comp.get("yoySales")
    if yoy_vol is not None:
        rows.append(["同比销量", _fmt_pct(yoy_vol)])
    if yoy_sales is not None:
        rows.append(["同比金额", _fmt_pct(yoy_sales)])
    inventory = summary.get("inventory")
    if inventory is not None and float(inventory) > 0:
        rows.append(["期末库存", _fmt_volume(inventory)])
    return [_assistant_table(f"{label}年度KPI", ["指标", "数值"], rows)]


def _fmt_pct(v):
    if v is None:
        return "--"
    sign = "+" if v > 0 else ""
    return f"{sign}{v:.1f}%"


def _assistant_station_status_answer():
    conn = get_connection()
    try:
        cur = conn.cursor()
        statuses = active_station_statuses(cur)
    finally:
        conn.close()
    if not statuses:
        return "当前所有站点正常营业。"
    closed = [s for s in statuses.values() if s.get("status") == "closed"]
    warning = [s for s in statuses.values() if s.get("status") == "warning"]
    lines = []
    if closed:
        lines.append(f"暂停营业/停站点（{len(closed)} 个）：")
        for s in closed:
            reason = s.get("closedReason") or s.get("statusLabel") or "未知"
            lines.append(f"  {s.get('name')}（{s.get('region')}）— {reason}")
    if warning:
        if closed:
            lines.append("")
        lines.append(f"异常状态站点（{len(warning)} 个）：")
        for s in warning:
            reason = s.get("closedReason") or s.get("statusLabel") or "未知"
            lines.append(f"  {s.get('name')}（{s.get('region')}）— {reason}（不影响销售）")
    if not closed and not warning:
        lines.append("当前所有站点状态正常。")
    return "\n".join(lines)


def _assistant_station_status_tables():
    conn = get_connection()
    try:
        cur = conn.cursor()
        statuses = active_station_statuses(cur)
    finally:
        conn.close()
    if not statuses:
        return []
    closed = [s for s in statuses.values() if s.get("status") == "closed"]
    warning = [s for s in statuses.values() if s.get("status") == "warning"]
    all_items = closed + warning
    if not all_items:
        return []
    return [_assistant_table(
        "当前站点异常状态",
        ["站点", "区域", "状态", "原因", "开始时间"],
        [
            [
                s.get("name", "--"),
                s.get("region", "--"),
                "停业" if s.get("status") == "closed" else "异常",
                s.get("closedReason") or s.get("statusLabel") or "",
                s.get("startedAt") or "--",
            ]
            for s in all_items
        ],
    )]


def _assistant_comparison_answer(label, monthly_data):
    comp = monthly_data.get("comparison") or {}
    mom = comp.get("mom") or {}
    yoy = comp.get("yoy") or {}
    lines = [f"{label}同比环比："]
    yoy_amt = yoy.get("sales")
    yoy_vol = yoy.get("volume")
    mom_amt = mom.get("sales")
    mom_vol = mom.get("volume")
    if yoy_amt is not None:
        lines.append(f"同比金额：{_fmt_pct(yoy_amt)}")
    if yoy_vol is not None:
        lines.append(f"同比销量：{_fmt_pct(yoy_vol)}")
    if mom_amt is not None:
        lines.append(f"环比金额：{_fmt_pct(mom_amt)}")
    if mom_vol is not None:
        lines.append(f"环比销量：{_fmt_pct(mom_vol)}")
    if yoy_amt is None and yoy_vol is None and mom_amt is None and mom_vol is None:
        lines.append("暂无对比数据（快照未生成对比字段，刷新后重试）。")
    return "\n".join(lines)


def _assistant_comparison_tables(label, monthly_data):
    comp = monthly_data.get("comparison") or {}
    mom = comp.get("mom") or {}
    yoy = comp.get("yoy") or {}
    rows = []
    for v, label_k in [(yoy.get("volume"), "同比销量"), (yoy.get("sales"), "同比金额"),
                        (mom.get("volume"), "环比销量"), (mom.get("sales"), "环比金额")]:
        if v is not None:
            rows.append([label_k, _fmt_pct(v)])
    if not rows:
        return []
    return [_assistant_table(f"{label}同比环比", ["指标", "变化"], rows)]


def _assistant_resource_share_answer(label, data):
    resources = [item for item in (data.get("byResource") or []) if float(item.get("volume") or 0) > 0]
    if not resources:
        return f"{label}无油品数据。"
    total_vol = sum(float(r.get("volume") or 0) for r in resources)
    total_amt = sum(float(r.get("amount") or 0) for r in resources)
    lines = [f"{label}油品结构与占比："]
    for item in sorted(resources, key=lambda r: float(r.get("amount") or 0), reverse=True):
        vol = float(item.get("volume") or 0)
        amt = float(item.get("amount") or 0)
        pct = vol / total_vol * 100 if total_vol else 0
        lines.append(f"  {item.get('name')}：{_fmt_volume(vol)}（占比 {pct:.1f}%），{_fmt_amount(amt, '索姆')}")
    lines.append(f"  合计：{_fmt_volume(total_vol)}，{_fmt_amount(total_amt, '索姆')}")
    return "\n".join(lines)


@app.route("/api/data-assistant/query", methods=["POST"])
def api_data_assistant_query():
    payload = request.get_json(force=True) or {}
    question = clean_str(payload.get("question"))
    if not question:
        return jsonify({"error": "question_required"}), 400

    label, start, end = _assistant_period(question)
    label = _assistant_period_display(label, start, end)
    if _assistant_has_any(question, ("停业", "站点状态", "异常站点", "暂停营业", "维修", "停电")):
        return jsonify({
            "ok": True,
            "question": question,
            "periodLabel": "当前",
            "answer": _assistant_station_status_answer(),
            "tables": _assistant_station_status_tables(),
        })
    if _assistant_has_any(question, ("趋势", "每日", "日趋势", "日销量", "走势")):
        ym = f"{start.year}-{start.month:02d}"
        trend_data = _read_daily_trend_snapshot(ym)
        if trend_data:
            return jsonify({
                "ok": True,
                "question": question,
                "periodLabel": ym,
                "answer": _assistant_daily_trend_answer(ym, trend_data),
                "tables": _assistant_daily_trend_tables(ym, trend_data),
                "data": trend_data,
            })
    if _assistant_has_any(question, ("同比", "环比", "对比", "较去年", "较上月")):
        ym = f"{start.year}-{start.month:02d}"
        monthly_data = _read_retail_snapshot("monthly", ym)
        if monthly_data:
            return jsonify({
                "ok": True,
                "question": question,
                "periodLabel": ym,
                "answer": _assistant_comparison_answer(ym, monthly_data),
                "tables": _assistant_comparison_tables(ym, monthly_data),
                "data": monthly_data,
            })
    if _assistant_has_any(question, ("占比", "结构", "油品结构", "油品占比", "构成")):
        data = _retail_snapshot_for_period(start, end)
        if not data:
            response = open_oilstation(
                "/api/assistant/retail-summary",
                query_string={"start": start.isoformat(), "end": end.isoformat()},
            )
            data = response.get_json() if response.status_code < 400 else {}
        return jsonify({
            "ok": True,
            "question": question,
            "periodLabel": label,
            "answer": _assistant_resource_share_answer(label, data),
            "tables": _assistant_retail_tables(label, data, question),
            "data": data,
        })
    if _assistant_has_any(question, ("KPI", "关键指标", "累计", "年度指标", "今年")):
        year = start.year
        kpi_data = _read_yearly_kpi_snapshot(year)
        if kpi_data:
            return jsonify({
                "ok": True,
                "question": question,
                "periodLabel": f"{year}年",
                "answer": _assistant_yearly_kpi_answer(f"{year}年", kpi_data),
                "tables": _assistant_yearly_kpi_tables(f"{year}年", kpi_data),
                "data": kpi_data,
            })
    if _assistant_has_any(question, ("价格", "单价", "售价", "零售价")):
        data = _retail_snapshot_for_period(start, end)
        if not data:
            response = open_oilstation(
                "/api/assistant/retail-summary",
                query_string={"start": start.isoformat(), "end": end.isoformat()},
            )
            data = response.get_json() if response.status_code < 400 else {}
        resources = [item for item in (data.get("byResource") or []) if float(item.get("volume") or 0) > 0]
        lines = [f"{label}零售价格（均价 = 金额 ÷ 销量）："]
        price_rows = []
        for item in sorted(resources, key=lambda r: float(r.get("amount") or 0), reverse=True):
            vol = float(item.get("volume") or 0)
            amt = float(item.get("amount") or 0)
            price = amt / vol if vol else 0
            lines.append(f"  {item.get('name')}：{price:.2f} 索姆/升（{_fmt_volume(vol)}，{_fmt_amount(amt, '索姆')}）")
            price_rows.append([item.get("name", "--"), f"{price:.2f}", _fmt_volume(vol), _fmt_amount(amt, "索姆")])
        if not lines[1:]:
            lines.append("  该区间暂无销售数据，无法计算均价。")
        total_vol = sum(float(r.get("volume") or 0) for r in resources)
        total_amt = sum(float(r.get("amount") or 0) for r in resources)
        avg_price = total_amt / total_vol if total_vol else 0
        lines.append(f"  综合均价：{avg_price:.2f} 索姆/升")
        price_rows.append(["综合均价", f"{avg_price:.2f}", _fmt_volume(total_vol), _fmt_amount(total_amt, "索姆")])
        return jsonify({
            "ok": True,
            "question": question,
            "periodLabel": label,
            "answer": "\n".join(lines),
            "tables": [_assistant_table(f"{label}零售价格", ["油品", "均价(索姆/升)", "销量", "金额"], price_rows)],
            "data": data,
        })
    if _assistant_has_any(question, ("市场", "行情", "观察", "原油", "成品油", "俄罗斯", "乌兹别克", "国际")):
        conn = get_connection()
        try:
            cur = conn.cursor()
            answer, tables = _assistant_market_watch_answer(cur)
        finally:
            conn.close()
        return jsonify({"ok": True, "question": question, "periodLabel": "市场观察", "answer": answer, "tables": tables})
    if "合同" in question:
        conn = get_connection()
        try:
            cur = conn.cursor()
            answer = _assistant_contract_answer(cur, question)
            tables = _assistant_contract_tables(cur, question)
        finally:
            conn.close()
        return jsonify({"ok": True, "question": question, "periodLabel": "合同", "answer": answer, "tables": tables})
    if _assistant_has_any(question, ("合并", "总销售", "整体", "公司", "总量", "总金额")):
        year, month = _assistant_year_month_from_period(start, end)
        conn = get_connection()
        try:
            cur = conn.cursor()
            data = _combined_data_from_snapshots(cur, year, month, start, end)
        finally:
            conn.close()
        return jsonify({
            "ok": True,
            "question": question,
            "periodLabel": label,
            "queryStart": start.isoformat(),
            "queryEndExclusive": end.isoformat(),
            "answer": _assistant_combined_answer(label, data),
            "tables": _assistant_combined_tables(label, data),
            "data": data,
        })
    if _assistant_has_any(question, ("审核", "复核", "异常", "提示", "未匹配")):
        year, month = _assistant_year_month_from_period(start, end)
        conn = get_connection()
        try:
            cur = conn.cursor()
            data = summarize_wholesale_page(cur, year, month)
            answer = _assistant_review_answer(label, data)
        finally:
            conn.close()
        return jsonify({"ok": True, "question": question, "periodLabel": label, "answer": answer, "tables": _assistant_review_tables(label, data), "data": data})
    if any(word in question for word in ("批发", "外运", "出口", "本地", "客户")):
        year, month = _assistant_year_month_from_period(start, end)
        conn = get_connection()
        try:
            cur = conn.cursor()
            data = summarize_wholesale_page(cur, year, month)
        finally:
            conn.close()
        return jsonify({
            "ok": True,
            "question": question,
            "periodLabel": label,
            "queryStart": start.isoformat(),
            "queryEndExclusive": end.isoformat(),
            "answer": _format_assistant_wholesale_answer(label, data, question),
            "tables": _assistant_wholesale_tables(label, data, question),
            "data": data,
        })
    if "汇率" in question or "美元" in question or "人民币" in question:
        conn = get_connection()
        try:
            cur = conn.cursor()
            answer = _assistant_exchange_answer(cur, label, start, end)
            tables = _assistant_exchange_tables(cur, label, start, end)
        finally:
            conn.close()
        return jsonify({"ok": True, "question": question, "periodLabel": label, "answer": answer, "tables": tables})
    if "密度" in question:
        conn = get_connection()
        try:
            cur = conn.cursor()
            answer = _assistant_density_answer(cur)
            tables = _assistant_density_tables(cur)
        finally:
            conn.close()
        return jsonify({"ok": True, "question": question, "periodLabel": "当前", "answer": answer, "tables": tables})
    if _assistant_has_any(question, ("库存", "期末")):
        inventory_data = _retail_inventory_for_period(start, end)
        if inventory_data:
            return jsonify({
                "ok": True,
                "question": question,
                "periodLabel": label,
                "queryStart": start.isoformat(),
                "queryEndExclusive": end.isoformat(),
                "answer": _assistant_retail_inventory_answer(label, inventory_data),
                "tables": _assistant_retail_inventory_tables(label, inventory_data),
                "data": inventory_data,
            })
    data = _retail_snapshot_for_period(start, end)
    if not data:
        response = open_oilstation(
            "/api/assistant/retail-summary",
            query_string={"start": start.isoformat(), "end": end.isoformat()},
        )
        if response.status_code >= 400:
            return jsonify({"error": "retail_query_failed", "detail": response.get_data(as_text=True)}), 502
        data = response.get_json() or {}
    data["question"] = question
    return jsonify({
        "ok": True,
        "question": question,
        "periodLabel": label,
        "queryStart": start.isoformat(),
        "queryEndExclusive": end.isoformat(),
        "answer": _format_assistant_retail_answer(label, data),
        "tables": _assistant_retail_tables(label, data, question),
        "data": data,
    })


@app.route("/api/auth/bootstrap", methods=["POST"])
def api_auth_bootstrap():
    payload = request.get_json(force=True)
    username = clean_str(payload.get("username"))
    password = payload.get("password") or ""
    display_name = clean_str(payload.get("displayName")) or username
    if not username or len(password) < 6:
        return jsonify({"error": "用户名和至少 6 位密码必填"}), 400
    conn = get_connection()
    try:
        cur = conn.cursor()
        if has_users(cur):
            return jsonify({"error": "系统已完成首次管理员创建"}), 409
        cur.execute(
            """
            INSERT INTO system_users (username, password_hash, display_name, role)
            VALUES (%s, %s, %s, 'admin')
            """,
            (username, generate_password_hash(password), display_name),
        )
        user_id = cur.lastrowid
        log_operation(cur, "bootstrap_admin", "system_users", user_id, {"username": username})
        conn.commit()
        session.permanent = True
        session["user"] = {
            "id": user_id,
            "username": username,
            "displayName": display_name,
            "role": "admin",
        }
        return jsonify({"ok": True, "user": session["user"]})
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@app.route("/api/auth/login", methods=["POST"])
def api_auth_login():
    payload = request.get_json(force=True)
    username = clean_str(payload.get("username"))
    password = payload.get("password") or ""
    conn = get_connection()
    try:
        cur = conn.cursor()
        ensure_auth_schema(cur)
        cur.execute(
            """
            SELECT id, username, password_hash, display_name, role, is_active
            FROM system_users
            WHERE username = %s
            """,
            (username,),
        )
        row = cur.fetchone()
        if not row or not row[5] or not check_password_hash(row[2], password):
            log_operation(
                cur,
                "login_failed",
                "system_users",
                username,
                {"username": username, "reason": "invalid_credentials", "remoteAddr": request.remote_addr},
            )
            conn.commit()
            return jsonify({"error": "用户名或密码不正确"}), 401
        cur.execute("UPDATE system_users SET last_login_at = CURRENT_TIMESTAMP WHERE id = %s", (row[0],))
        log_operation(cur, "login_success", "system_users", row[0], {"username": row[1], "remoteAddr": request.remote_addr})
        conn.commit()
        session.permanent = True
        session["user"] = {
            "id": row[0],
            "username": row[1],
            "displayName": row[3],
            "role": row[4],
        }
        return jsonify({"ok": True, "user": session["user"]})
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@app.route("/api/auth/logout", methods=["POST"])
def api_auth_logout():
    user = session_user()
    if user:
        log_system_operation("logout", "system_users", user.get("id"), {"username": user.get("username")})
    session.clear()
    return jsonify({"ok": True})


@app.route("/api/admin/users")
@admin_required
def api_admin_users():
    conn = get_connection()
    try:
        cur = conn.cursor()
        ensure_auth_schema(cur)
        cur.execute(
            """
            SELECT id, username, display_name, role, is_active, created_at, last_login_at
            FROM system_users
            ORDER BY id
            """
        )
        rows = [
            {
                "id": row[0],
                "username": row[1],
                "displayName": row[2],
                "role": row[3],
                "isActive": bool(row[4]),
                "createdAt": row[5].isoformat() if row[5] else None,
                "lastLoginAt": row[6].isoformat() if row[6] else None,
            }
            for row in cur.fetchall()
        ]
        return jsonify({"items": rows})
    finally:
        conn.close()


@app.route("/api/admin/users", methods=["POST"])
@admin_required
def api_admin_create_user():
    payload = request.get_json(force=True)
    username = clean_str(payload.get("username"))
    password = payload.get("password") or ""
    display_name = clean_str(payload.get("displayName")) or username
    role = payload.get("role") if payload.get("role") in ("admin", "user") else "user"
    if not username or len(password) < 6:
        return jsonify({"error": "用户名和至少 6 位密码必填"}), 400
    conn = get_connection()
    try:
        cur = conn.cursor()
        ensure_auth_schema(cur)
        cur.execute(
            """
            INSERT INTO system_users (username, password_hash, display_name, role, created_by)
            VALUES (%s, %s, %s, %s, %s)
            """,
            (username, generate_password_hash(password), display_name, role, (session_user() or {}).get("id")),
        )
        user_id = cur.lastrowid
        log_operation(cur, "create_user", "system_users", user_id, {"username": username, "role": role})
        conn.commit()
        return jsonify({"ok": True, "id": user_id})
    except IntegrityError as exc:
        conn.rollback()
        if exc.args and exc.args[0] == 1062:
            return jsonify({"error": "用户名已存在，请更换用户名"}), 409
        raise
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@app.route("/api/retail/station-catalog")
@admin_required
def api_retail_station_catalog():
    return jsonify({"items": list(retail_station_catalog().values())})


@app.route("/api/retail/station-statuses")
def api_retail_station_statuses():
    conn = get_connection()
    try:
        cur = conn.cursor()
        statuses = active_station_statuses(cur)
        return jsonify({"items": list(statuses.values())})
    finally:
        conn.close()


@app.route("/api/admin/retail-station-statuses")
@admin_required
def api_admin_retail_station_statuses():
    conn = get_connection()
    try:
        cur = conn.cursor()
        ensure_station_status_schema(cur)
        catalog = retail_station_catalog()
        cur.execute(
            """
            SELECT id, shop_key, status_type, status_label, started_at, expected_resume_at,
                   ended_at, affects_sales, affects_inventory, suppress_alerts, note, is_active,
                   created_at, updated_at
            FROM retail_station_status
            WHERE NOT (status_type = 'normal' AND is_active = TRUE AND ended_at IS NULL)
            ORDER BY is_active DESC, ended_at IS NULL DESC, shop_key, started_at DESC, id DESC
            """
        )
        items = []
        for row in cur.fetchall():
            station = catalog.get(int(row[1]), {"name": f"Station-{row[1]}", "short": "", "region": ""})
            items.append({
                "id": row[0],
                "shopKey": int(row[1]),
                "name": station["name"],
                "short": station["short"],
                "region": station["region"],
                "statusType": row[2],
                "statusLabel": row[3],
                "startedAt": row[4].isoformat() if row[4] else None,
                "expectedResumeAt": row[5].isoformat() if row[5] else None,
                "endedAt": row[6].isoformat() if row[6] else None,
                "affectsSales": bool(row[7]),
                "affectsInventory": bool(row[8]),
                "suppressAlerts": bool(row[9]),
                "note": row[10] or "",
                "isActive": bool(row[11]),
                "createdAt": row[12].isoformat() if row[12] else None,
                "updatedAt": row[13].isoformat() if row[13] else None,
            })
        return jsonify({"items": items, "statusTypes": STATION_STATUS_TYPES})
    finally:
        conn.close()


@app.route("/api/admin/market-watch")
@admin_required
def api_admin_market_watch():
    conn = get_connection()
    try:
        cur = conn.cursor()
        return jsonify({"items": market_watch_rows(cur)})
    finally:
        conn.close()


@app.route("/api/admin/market-watch", methods=["POST"])
@admin_required
def api_admin_save_market_watch():
    payload = request.get_json(force=True)
    title = clean_str(payload.get("title"))
    summary = clean_str(payload.get("summary"))
    observed_at = parse_date(payload.get("observedAt")) or date.today()
    region = payload.get("region") if payload.get("region") in {"Russia", "Uzbekistan", "Global"} else "Global"
    market_type = payload.get("type") if payload.get("type") in {"crude", "refined", "policy", "logistics", "other"} else "other"
    source_name = clean_str(payload.get("sourceName"))
    source_url = clean_str(payload.get("sourceUrl"))
    price_text = clean_str(payload.get("priceText"))
    priority = int(payload.get("priority") or 100)
    if not title or not summary:
        return jsonify({"error": "标题和摘要必填"}), 400
    if market_type == "refined" and price_text and "索姆/升" not in price_text:
        return jsonify({"error": "成品油价格请统一按索姆/升录入"}), 400

    conn = get_connection()
    try:
        cur = conn.cursor()
        ensure_market_watch_schema(cur)
        cur.execute(
            """
            INSERT INTO market_watch_items
                (market_region, market_type, title, summary, price_text, source_name, source_url, observed_at, priority, is_active)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, TRUE)
            """,
            (region, market_type, title, summary, price_text, source_name, source_url, observed_at, priority),
        )
        item_id = cur.lastrowid
        log_operation(cur, "save_market_watch", "market_watch_items", item_id, {"title": title})
        conn.commit()
        return jsonify({"ok": True, "id": item_id})
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@app.route("/api/admin/market-watch/fetch", methods=["POST"])
@admin_required
def api_admin_fetch_market_watch():
    conn = get_connection()
    try:
        cur = conn.cursor()
        result = run_market_watch_maintenance(cur)
        log_operation(cur, "fetch_market_watch", "market_watch_items", None, result)
        record_automation_task(cur, "market_watch", "success" if not result["errors"] else "failed", result, "; ".join(result["errors"]) or None)
        conn.commit()
        return jsonify({"ok": True, **result, "items": market_watch_rows(cur, 20)})
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@app.route("/api/admin/automation-tasks")
@admin_required
def api_admin_automation_tasks():
    conn = get_connection()
    try:
        cur = conn.cursor()
        ensure_automation_task_schema(cur)
        cur.execute(
            """
            SELECT r.id, r.task_name, r.status, r.started_at, r.finished_at, r.summary, r.error
            FROM automation_task_runs r
            JOIN (
                SELECT task_name, MAX(id) AS id
                FROM automation_task_runs
                GROUP BY task_name
            ) latest ON latest.id = r.id
            ORDER BY r.task_name
            """
        )
        latest = [
            {
                "id": row[0],
                "taskName": row[1],
                "status": row[2],
                "startedAt": row[3].isoformat() if row[3] else None,
                "finishedAt": row[4].isoformat() if row[4] else None,
                "summary": json.loads(row[5] or "{}"),
                "error": row[6] or "",
            }
            for row in cur.fetchall()
        ]
        cur.execute(
            """
            SELECT id, task_name, status, started_at, finished_at, summary, error
            FROM automation_task_runs
            ORDER BY id DESC
            LIMIT 40
            """
        )
        recent = [
            {
                "id": row[0],
                "taskName": row[1],
                "status": row[2],
                "startedAt": row[3].isoformat() if row[3] else None,
                "finishedAt": row[4].isoformat() if row[4] else None,
                "summary": json.loads(row[5] or "{}"),
                "error": row[6] or "",
            }
            for row in cur.fetchall()
        ]
        return jsonify({"latest": latest, "recent": recent})
    finally:
        conn.close()


@app.route("/api/admin/automation-tasks/run", methods=["POST"])
@admin_required
def api_admin_run_automation_task():
    payload = request.get_json(force=True)
    task = payload.get("task")
    if task not in {"exchange_rates", "market_watch", "retail_snapshots", "all"}:
        return jsonify({"error": "请选择有效的自动任务"}), 400

    conn = get_connection()
    try:
        cur = conn.cursor()
        results = {}
        if task in {"exchange_rates", "all"}:
            result = run_exchange_maintenance(cur)
            status = "success" if not result["errors"] else "failed"
            record_automation_task(cur, "exchange_rates", status, result, json.dumps(result["errors"], ensure_ascii=False) if result["errors"] else None)
            results["exchangeRates"] = result
        if task in {"market_watch", "all"}:
            result = run_market_watch_maintenance(cur)
            status = "success" if not result["errors"] else "failed"
            record_automation_task(cur, "market_watch", status, result, "; ".join(result["errors"]) or None)
            results["marketWatch"] = result
        if task in {"retail_snapshots", "all"}:
            result = _queue_retail_cache_refresh()
            results["retailSnapshots"] = result
        log_operation(cur, "run_automation_task", "automation_task_runs", task, results)
        conn.commit()
        return jsonify({"ok": True, "results": results})
    except Exception as exc:
        conn.rollback()
        try:
            cur = conn.cursor()
            record_automation_task(cur, task if task != "all" else "automation_all", "failed", {}, str(exc))
            conn.commit()
        except Exception:
            conn.rollback()
        raise
    finally:
        conn.close()


def _health_file_info(path):
    path = Path(path)
    if not path.exists():
        return {"exists": False, "size": 0, "updatedAt": None}
    stat = path.stat()
    return {
        "exists": True,
        "size": stat.st_size,
        "updatedAt": date.fromtimestamp(stat.st_mtime).isoformat(),
    }


def _health_month_iter(start_ym, end_ym):
    year, month = [int(part) for part in start_ym.split("-")]
    end_year, end_month = [int(part) for part in end_ym.split("-")]
    while (year, month) <= (end_year, end_month):
        yield f"{year}-{month:02d}"
        if month == 12:
            year += 1
            month = 1
        else:
            month += 1


def _health_completed_month():
    today = date.today()
    year, month = today.year, today.month - 1
    if month == 0:
        year -= 1
        month = 12
    return f"{year}-{month:02d}"


def _health_report_snapshot_status():
    start_ym = "2023-12"
    end_ym = _health_completed_month()
    completed_day = date.today() - timedelta(days=1)
    active_ym = f"{completed_day.year}-{completed_day.month:02d}"
    report_dir = OILSTATION_ROOT / "cache_snapshots" / "report_extras"
    expected = list(_health_month_iter(start_ym, end_ym))
    missing = [
        ym for ym in expected
        if not (report_dir / f"report_extras_{ym}.json").exists()
    ]
    files = list(report_dir.glob("report_extras_*.json")) if report_dir.exists() else []
    day_dir = OILSTATION_ROOT / "cache_snapshots" / "daily"
    day_start = date(completed_day.year, completed_day.month, 1)
    day_end = completed_day
    day_expected = []
    cursor = day_start
    while cursor <= day_end:
        day_expected.append(cursor.isoformat())
        cursor += timedelta(days=1)
    day_missing = [
        item for item in day_expected
        if not (day_dir / f"day_{item}.json").exists()
    ]
    day_files = [
        path for item in day_expected
        if (path := day_dir / f"day_{item}.json").exists()
    ] if day_dir.exists() else []
    current_ym = active_ym
    report_ym = active_ym
    monthly_dir = OILSTATION_ROOT / "cache_snapshots" / "monthly"
    current_files = {
        prefix: _health_file_info(monthly_dir / f"{prefix}_{current_ym}.json")
        for prefix in ("monthly", "stations", "regions", "daily")
    }
    report_workbook = {"exists": False, "file": None, "lastNonemptyDay": None, "updatedAt": None, "size": 0}
    report_index_path = OILSTATION_ROOT / "exports" / ".monthly_report_index.json"
    try:
        report_index = json.loads(report_index_path.read_text(encoding="utf-8-sig")) if report_index_path.exists() else {}
        report_item = report_index.get(report_ym) or {}
        report_file = report_item.get("file")
        report_path = OILSTATION_ROOT / "exports" / report_file if report_file else None
        if report_path and report_path.exists():
            info = _health_file_info(report_path)
            report_workbook.update({
                "exists": True,
                "month": report_ym,
                "file": report_file,
                "lastNonemptyDay": report_item.get("lastNonemptyDay"),
                "updatedAt": info.get("updatedAt"),
                "size": info.get("size", 0),
            })
    except (OSError, json.JSONDecodeError):
        pass
    return {
        "expectedStart": start_ym,
        "expectedEnd": end_ym,
        "expectedCount": len(expected),
        "actualCount": len(files),
        "missing": missing,
        "dailyExpectedStart": day_start.isoformat(),
        "dailyExpectedEnd": day_end.isoformat() if day_expected else None,
        "dailyExpectedCount": len(day_expected),
        "dailyActualCount": len(day_files),
        "dailyMissing": day_missing,
        "currentMonth": current_ym,
        "currentMonthFiles": current_files,
        "reportWorkbook": report_workbook,
    }


def _task_latest_record(automation_records, task_name):
    mapping = {
        "SalesVisualizationDailyExchangeRates": "exchange_rates",
        "SalesVisualizationDailyMarketWatch": "market_watch",
        "SalesVisualizationDailyRetailSnapshots": "retail_snapshots",
    }
    automation_name = mapping.get(task_name)
    if not automation_name:
        return None
    for item in automation_records or []:
        if item.get("taskName") == automation_name:
            return item
    return None


def _health_task_fallback(task_name, detail, automation_records=None, service_running=True):
    if task_name == "SalesVisualizationProduction":
        return {
            "taskName": task_name,
            "ok": bool(service_running),
            "warning": True,
            "status": "服务运行中" if service_running else "服务需检查",
            "detail": f"计划任务读取受限；{detail}",
            "source": "service",
        }
    latest = _task_latest_record(automation_records, task_name)
    if not latest:
        return {
            "taskName": task_name,
            "ok": False,
            "warning": True,
            "status": "无执行记录",
            "detail": f"计划任务读取受限；{detail}",
            "source": "automation_log",
        }
    finished_at = latest.get("finishedAt") or latest.get("startedAt")
    is_success = latest.get("status") == "success"
    fresh = False
    if finished_at:
        try:
            finished_dt = datetime.fromisoformat(finished_at)
            fresh = datetime.now() - finished_dt <= timedelta(hours=36)
        except ValueError:
            fresh = False
    return {
        "taskName": task_name,
        "ok": bool(is_success and fresh),
        "warning": True,
        "status": "最近成功" if is_success and fresh else ("最近失败" if not is_success else "记录过期"),
        "detail": f"计划任务读取受限；最近记录 {finished_at or '--'}；{detail}",
        "source": "automation_log",
        "lastRunAt": finished_at,
    }


def _health_task_status(task_name, automation_records=None, service_running=True):
    task_path = f"\\SalesVisualization\\{task_name}"
    try:
        completed = subprocess.run(
            ["schtasks", "/Query", "/TN", task_path, "/FO", "LIST"],
            cwd=str(PROJECT_ROOT),
            capture_output=True,
            text=True,
            timeout=10,
            check=False,
        )
    except (OSError, subprocess.SubprocessError) as exc:
        return {"taskName": task_name, "ok": False, "status": "unknown", "detail": str(exc)}
    output = (completed.stdout or completed.stderr or "").strip()
    if completed.returncode != 0:
        detail = output.splitlines()[0] if output else "无法读取计划任务"
        return _health_task_fallback(task_name, detail, automation_records, service_running)
    status = "unknown"
    for line in output.splitlines():
        stripped = line.strip()
        if stripped.startswith("Status:") or stripped.startswith("状态:") or stripped.startswith("РЎС‚Р°С‚СѓСЃ:"):
            status = line.split(":", 1)[1].strip()
            break
    return {"taskName": task_name, "ok": True, "warning": False, "status": status, "detail": "", "source": "schtasks"}


def _health_timed(label, fn):
    started = time.perf_counter()
    try:
        payload = fn()
        elapsed = int((time.perf_counter() - started) * 1000)
        return {"name": label, "ok": True, "elapsedMs": elapsed, "detail": payload}
    except Exception as exc:
        elapsed = int((time.perf_counter() - started) * 1000)
        return {"name": label, "ok": False, "elapsedMs": elapsed, "detail": str(exc)}


def _current_month():
    today = date.today()
    return f"{today.year}-{today.month:02d}"


def _refresh_retail_cache_worker():
    current_month = _current_month()
    steps = [
        {
            "label": "零售日快照与当前月汇总",
            "type": "command",
            "command": [
                sys.executable,
                str(OILSTATION_ROOT / "scripts" / "refresh_current_snapshots.py"),
            ],
            "cwd": str(OILSTATION_ROOT),
        },
        {"label": "年度总览预热", "type": "api", "path": "/api/dashboard/yearly", "query": f"year={date.today().year}"},
        {"label": "油品统计预热", "type": "api", "path": "/api/resources/yearly", "query": f"year={date.today().year}"},
        {"label": "月度快照预热", "type": "api", "path": f"/api/monthly/{current_month}", "query": ""},
        {"label": "站点快照预热", "type": "api", "path": f"/api/stations/{current_month}", "query": ""},
        {"label": "区域快照预热", "type": "api", "path": f"/api/regions/{current_month}", "query": ""},
        {"label": "日趋势快照预热", "type": "api", "path": f"/api/daily/{current_month}", "query": ""},
    ]
    items = []
    with REFRESH_CACHE_LOCK:
        REFRESH_CACHE_STATUS.update({
            "running": True,
            "startedAt": datetime.now().isoformat(timespec="seconds"),
            "finishedAt": None,
            "status": "running",
            "message": "正在刷新当前月缓存",
            "items": items,
        })
    try:
        for step in steps:
            started = time.perf_counter()
            item = {
                "label": step["label"],
                "path": step.get("path") or step.get("command", [""])[-1],
                "query": step.get("query", ""),
                "status": "running",
                "elapsedMs": None,
                "error": "",
            }
            items.append(item)
            if step["type"] == "command":
                completed = subprocess.run(
                    step["command"],
                    cwd=step["cwd"],
                    capture_output=True,
                    text=True,
                    timeout=900,
                    check=False,
                )
                item["elapsedMs"] = int((time.perf_counter() - started) * 1000)
                if completed.returncode:
                    item["status"] = "failed"
                    item["error"] = (completed.stderr or completed.stdout or "")[-1000:]
                else:
                    item["status"] = "success"
                continue
            response = open_oilstation(step["path"], query_string=step.get("query") or "")
            item["elapsedMs"] = int((time.perf_counter() - started) * 1000)
            if response.status_code >= 400:
                item["status"] = "failed"
                item["error"] = response.get_data(as_text=True)[:300]
            else:
                item["status"] = "success"
        status = "success" if all(item["status"] == "success" for item in items) else "failed"
        message = "当前月缓存刷新完成" if status == "success" else "当前月缓存刷新存在失败项"
    except Exception as exc:
        status = "failed"
        message = str(exc)
        logging.exception("Retail cache refresh failed")
    with REFRESH_CACHE_LOCK:
        REFRESH_CACHE_STATUS.update({
            "running": False,
            "finishedAt": datetime.now().isoformat(timespec="seconds"),
            "status": status,
            "message": message,
            "items": items,
        })
    conn = get_connection()
    try:
        cur = conn.cursor()
        record_automation_task(
            cur,
            "retail_snapshots",
            status,
            {"items": items, "message": message},
            message if status == "failed" else None,
        )
        conn.commit()
    finally:
        conn.close()


def _queue_retail_cache_refresh():
    with REFRESH_CACHE_LOCK:
        if REFRESH_CACHE_STATUS.get("running"):
            return {"status": "running", "message": "零售快照刷新正在后台运行"}
        REFRESH_CACHE_STATUS.update({
            "running": True,
            "startedAt": datetime.now().isoformat(timespec="seconds"),
            "finishedAt": None,
            "status": "queued",
            "message": "零售快照刷新已加入后台执行",
            "items": [],
        })
    thread = threading.Thread(target=_refresh_retail_cache_worker, name="retail-cache-refresh", daemon=True)
    thread.start()
    return {"status": "queued", "message": "零售快照刷新已加入后台执行"}


@app.route("/api/admin/system-health")
@admin_required
def api_admin_system_health():
    checks = []
    checks.append(_health_timed("主项目服务", lambda: {"ok": True}))
    checks.append(_health_timed("零售嵌入接口", lambda: open_oilstation("/api/health").get_json(silent=True)))

    conn = get_connection()
    try:
        cur = conn.cursor()
        checks.append(_health_timed("销售数据库", lambda: (cur.execute("SELECT 1"), {"ok": True})[1]))
        ensure_automation_task_schema(cur)
        cur.execute(
            """
            SELECT task_name, status, started_at, finished_at, summary, error
            FROM automation_task_runs
            ORDER BY id DESC
            LIMIT 8
            """
        )
        automation = [
            {
                "taskName": row[0],
                "status": row[1],
                "startedAt": row[2].isoformat() if row[2] else None,
                "finishedAt": row[3].isoformat() if row[3] else None,
                "summary": json.loads(row[4] or "{}"),
                "error": row[5] or "",
            }
            for row in cur.fetchall()
        ]
        cur.execute(
            """
            SELECT rate_date, usd_kgs, cny_kgs, source, is_manual, fetched_at
            FROM exchange_rates
            WHERE rate_date <= CURDATE()
            ORDER BY rate_date DESC
            LIMIT 1
            """
        )
        rate_row = cur.fetchone()
        latest_exchange = {
            "date": rate_row[0].isoformat() if rate_row and rate_row[0] else None,
            "usdKgs": float(rate_row[1]) if rate_row and rate_row[1] is not None else None,
            "cnyKgs": float(rate_row[2]) if rate_row and rate_row[2] is not None else None,
            "source": rate_row[3] if rate_row else None,
            "isManual": bool(rate_row[4]) if rate_row else False,
            "fetchedAt": rate_row[5].isoformat() if rate_row and rate_row[5] else None,
        }
        ensure_market_watch_schema(cur)
        cur.execute(
            """
            SELECT observed_at, title, source_name, created_at
            FROM market_watch_items
            WHERE is_active = TRUE
            ORDER BY observed_at DESC, priority ASC, id DESC
            LIMIT 1
            """
        )
        market_row = cur.fetchone()
        latest_market = {
            "observedAt": market_row[0].isoformat() if market_row and market_row[0] else None,
            "title": market_row[1] if market_row else None,
            "sourceName": market_row[2] if market_row else None,
            "createdAt": market_row[3].isoformat() if market_row and market_row[3] else None,
        }
    finally:
        conn.close()

    snapshots = _health_report_snapshot_status()
    logs = []
    for path in (
        PROJECT_ROOT / "logs" / "production.log",
        PROJECT_ROOT / "logs" / "production-error.log",
        PROJECT_ROOT / "logs" / "daily_maintenance.log",
    ):
        info = _health_file_info(path)
        info["name"] = path.name
        if info["exists"]:
            text = path.read_text(encoding="utf-8", errors="ignore")[-2000:]
            lowered = text.lower()
            last_traceback = lowered.rfind("traceback")
            last_success = max(lowered.rfind("finished with success"), lowered.rfind("starting sales visualization"))
            info["hasTraceback"] = last_traceback != -1 and last_traceback > last_success
            info["tail"] = text.splitlines()[-8:]
        logs.append(info)

    service_running = all(item["ok"] for item in checks[:2])
    tasks = [
        _health_task_status("SalesVisualizationProduction", automation, service_running),
        _health_task_status("SalesVisualizationDailyExchangeRates", automation, service_running),
        _health_task_status("SalesVisualizationDailyMarketWatch", automation, service_running),
        _health_task_status("SalesVisualizationDailyRetailSnapshots", automation, service_running),
    ]
    security = {
        "runtimeSecret": (PROJECT_ROOT / ".runtime" / "secret_key.txt").exists(),
        "cookieSecure": bool(app.config.get("SESSION_COOKIE_SECURE")),
        "cookieHttpOnly": bool(app.config.get("SESSION_COOKIE_HTTPONLY")),
        "sameSite": app.config.get("SESSION_COOKIE_SAMESITE"),
    }
    ok = (
        all(item["ok"] for item in checks)
        and not snapshots["missing"]
        and not snapshots["dailyMissing"]
        and all(not item.get("hasTraceback") for item in logs)
        and all(item["ok"] for item in tasks)
    )
    return jsonify({
        "ok": ok,
        "generatedAt": date.today().isoformat(),
        "checks": checks,
        "snapshots": snapshots,
        "logs": logs,
        "tasks": tasks,
        "automation": automation,
        "security": security,
        "latestExchange": latest_exchange,
        "latestMarket": latest_market,
        "cacheRefresh": REFRESH_CACHE_STATUS,
    })


@app.route("/api/admin/system-health/refresh-cache", methods=["POST"])
@admin_required
def api_admin_refresh_retail_cache():
    queued = _queue_retail_cache_refresh()
    log_system_operation("refresh_retail_cache", "cache", _current_month(), {"month": _current_month()})
    return jsonify({"ok": True, "status": REFRESH_CACHE_STATUS, "message": queued["message"]})


@app.route("/api/admin/system/restart", methods=["POST"])
@admin_required
def api_admin_restart_service():
    script = PROJECT_ROOT / "scripts" / "restart_production_task.ps1"
    if not script.exists():
        return jsonify({"error": "重启脚本不存在"}), 500
    creation_flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    subprocess.Popen(
        [
            "powershell.exe",
            "-NoProfile",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(script),
        ],
        cwd=str(PROJECT_ROOT.parent),
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        creationflags=creation_flags,
    )
    log_system_operation("restart_service", "scheduled_task", "SalesVisualizationProduction", {"script": str(script)})
    return jsonify({"ok": True, "message": "已提交重启请求，约 20 秒后刷新页面。"})


@app.route("/api/admin/retail-station-statuses", methods=["POST"])
@admin_required
def api_admin_create_retail_station_status():
    payload = request.get_json(force=True)
    shop_key = int(payload.get("shopKey") or 0)
    status_type = payload.get("statusType") if payload.get("statusType") in STATION_STATUS_TYPES else "other"
    status_label = clean_str(payload.get("statusLabel")) or STATION_STATUS_TYPES[status_type]
    started_at = payload.get("startedAt") or None
    expected_resume_at = payload.get("expectedResumeAt") or None
    if started_at:
        started_at = started_at.replace("T", " ")
    if expected_resume_at:
        expected_resume_at = expected_resume_at.replace("T", " ")
    note = clean_str(payload.get("note"))
    affects_sales = bool(payload.get("affectsSales"))
    affects_inventory = bool(payload.get("affectsInventory"))
    suppress_alerts = bool(payload.get("suppressAlerts"))
    if not shop_key:
        return jsonify({"error": "请选择加油站"}), 400
    if status_type == "normal":
        affects_sales = False
        affects_inventory = False
        suppress_alerts = False
    conn = get_connection()
    try:
        cur = conn.cursor()
        ensure_station_status_schema(cur)
        cur.execute(
            """
            UPDATE retail_station_status
            SET is_active = FALSE, ended_at = COALESCE(ended_at, CURRENT_TIMESTAMP)
            WHERE shop_key = %s AND is_active = TRUE AND ended_at IS NULL
            """,
            (shop_key,),
        )
        if status_type == "normal":
            log_operation(
                cur,
                "restore_station_normal",
                "retail_station_status",
                shop_key,
                {"shopKey": shop_key, "statusType": status_type, "statusLabel": status_label},
            )
            conn.commit()
            return jsonify({"ok": True, "id": None})
        cur.execute(
            """
            INSERT INTO retail_station_status
                (shop_key, status_type, status_label, started_at, expected_resume_at,
                 affects_sales, affects_inventory, suppress_alerts, note, created_by)
            VALUES (%s, %s, %s, COALESCE(%s, CURRENT_TIMESTAMP), %s, %s, %s, %s, %s, %s)
            """,
            (
                shop_key,
                status_type,
                status_label,
                started_at,
                expected_resume_at,
                affects_sales,
                affects_inventory,
                suppress_alerts,
                note,
                session_user()["id"],
            ),
        )
        status_id = cur.lastrowid
        log_operation(
            cur,
            "update_station_status",
            "retail_station_status",
            status_id,
            {"shopKey": shop_key, "statusType": status_type, "statusLabel": status_label},
        )
        conn.commit()
        return jsonify({"ok": True, "id": status_id})
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@app.route("/api/admin/retail-station-statuses/<int:status_id>/close", methods=["POST"])
@admin_required
def api_admin_close_retail_station_status(status_id):
    conn = get_connection()
    try:
        cur = conn.cursor()
        ensure_station_status_schema(cur)
        cur.execute(
            """
            UPDATE retail_station_status
            SET is_active = FALSE, ended_at = CURRENT_TIMESTAMP
            WHERE id = %s
            """,
            (status_id,),
        )
        log_operation(cur, "close_station_status", "retail_station_status", status_id)
        conn.commit()
        return jsonify({"ok": True})
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@app.route("/api/admin/exchange-rates")
@admin_required
def api_admin_exchange_rates():
    year = int(request.args.get("year", str(date.today().year)))
    month_raw = request.args.get("month")
    month = int(month_raw) if month_raw else None
    start, end = period_bounds(year, month)
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT
                dates.rate_date,
                er.usd_kgs,
                er.cny_kgs,
                er.source,
                er.is_manual,
                er.fetched_at,
                er.note,
                COALESCE(g.status, IF(er.rate_date IS NULL OR er.usd_kgs IS NULL OR er.cny_kgs IS NULL, 'open', 'resolved')) AS gap_status,
                COALESCE(g.missing_usd, er.usd_kgs IS NULL) AS missing_usd,
                COALESCE(g.missing_cny, er.cny_kgs IS NULL) AS missing_cny
            FROM (
                SELECT rate_date
                FROM exchange_rates
                WHERE rate_date >= %s AND rate_date < %s
                UNION
                SELECT rate_date
                FROM exchange_rate_gaps
                WHERE rate_date >= %s AND rate_date < %s
                UNION
                SELECT exchange_rate_date AS rate_date
                FROM calculation_results
                WHERE exchange_rate_date >= %s AND exchange_rate_date < %s
                  AND exchange_missing = TRUE
                UNION
                SELECT issue_date AS rate_date
                FROM data_quality_issues
                WHERE issue_type = 'missing_exchange_rate'
                  AND issue_date >= %s AND issue_date < %s
            ) dates
            LEFT JOIN exchange_rates er ON er.rate_date = dates.rate_date
            LEFT JOIN exchange_rate_gaps g ON g.rate_date = dates.rate_date
            WHERE dates.rate_date IS NOT NULL
            ORDER BY dates.rate_date DESC
            """,
            (start, end, start, end, start, end, start, end),
        )
        items = [
            {
                "date": row[0].isoformat(),
                "usdKgs": float(row[1]) if row[1] is not None else None,
                "cnyKgs": float(row[2]) if row[2] is not None else None,
                "source": row[3],
                "isManual": bool(row[4]),
                "fetchedAt": row[5].isoformat() if row[5] else None,
                "note": row[6] or "",
                "gapStatus": row[7],
                "missingUsd": bool(row[8]),
                "missingCny": bool(row[9]),
            }
            for row in cur.fetchall()
        ]
        return jsonify({
            "items": items,
            "total": len(items),
            "openGaps": sum(1 for item in items if item["gapStatus"] == "open"),
        })
    finally:
        conn.close()


@app.route("/api/admin/exchange-rates/fetch-missing", methods=["POST"])
@admin_required
def api_admin_fetch_missing_exchange_rates():
    payload = request.get_json(silent=True) or {}
    year = int(payload.get("year") or date.today().year)
    month_raw = payload.get("month")
    month = int(month_raw) if month_raw else None
    start, end = period_bounds(year, month)
    conn = get_connection()
    try:
        cur = conn.cursor()
        dates = missing_exchange_dates(cur, start, end)
        fetched = []
        unresolved = []
        errors = []
        for rate_date in dates:
            try:
                usd, cny, source_url = fetch_nbkr_daily_rates(rate_date)
                if usd is None or cny is None:
                    unresolved.append({
                        "date": rate_date.isoformat(),
                        "missingUsd": usd is None,
                        "missingCny": cny is None,
                    })
                resolved = upsert_exchange_rate(
                    cur,
                    rate_date,
                    usd,
                    cny,
                    f"NBKR 自动获取：{source_url}",
                )
                fetched.append({
                    "date": rate_date.isoformat(),
                    "usdKgs": float(usd) if usd is not None else None,
                    "cnyKgs": float(cny) if cny is not None else None,
                    "resolved": resolved,
                })
            except RuntimeError as exc:
                errors.append({"date": rate_date.isoformat(), "error": str(exc)})

        recalculation = recalculate_all_wholesale(cur) if fetched else {"matched": 0, "unmatched": 0, "missing_rates": 0}
        log_operation(
            cur,
            "fetch_missing_exchange_rates",
            "exchange_rates",
            f"{start.isoformat()}..{end.isoformat()}",
            {
                "requested": len(dates),
                "fetched": len(fetched),
                "unresolved": len(unresolved),
                "errors": len(errors),
            },
        )
        conn.commit()
        return jsonify({
            "ok": True,
            "requested": len(dates),
            "fetched": fetched,
            "updated": len(fetched),
            "resolved": sum(1 for item in fetched if item["resolved"]),
            "unresolved": unresolved,
            "errors": errors,
            "recalculation": recalculation,
        })
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@app.route("/api/admin/exchange-rates", methods=["POST"])
@admin_required
def api_admin_save_exchange_rate():
    payload = request.get_json(force=True)
    rate_date = parse_date(payload.get("rateDate"))
    usd = parse_decimal(payload.get("usdKgs"))
    cny = parse_decimal(payload.get("cnyKgs"))
    note = clean_str(payload.get("note"))
    if not rate_date:
        return jsonify({"error": "请选择汇率日期"}), 400
    if usd is None and cny is None:
        return jsonify({"error": "至少录入 USD/KGS 或 CNY/KGS 中的一项"}), 400
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO exchange_rates
                (rate_date, usd_kgs, cny_kgs, source, is_manual, fetched_at, note)
            VALUES (%s, %s, %s, 'manual', TRUE, CURRENT_TIMESTAMP, %s)
            ON DUPLICATE KEY UPDATE
                usd_kgs = COALESCE(VALUES(usd_kgs), usd_kgs),
                cny_kgs = COALESCE(VALUES(cny_kgs), cny_kgs),
                source = 'manual',
                is_manual = TRUE,
                fetched_at = CURRENT_TIMESTAMP,
                note = COALESCE(VALUES(note), note)
            """,
            (rate_date, usd, cny, note),
        )
        cur.execute("SELECT usd_kgs, cny_kgs FROM exchange_rates WHERE rate_date = %s", (rate_date,))
        saved = cur.fetchone()
        missing_usd = saved[0] is None
        missing_cny = saved[1] is None
        if missing_usd or missing_cny:
            cur.execute(
                """
                INSERT INTO exchange_rate_gaps (rate_date, missing_usd, missing_cny, status, note)
                VALUES (%s, %s, %s, 'open', 'Manual rate still incomplete.')
                ON DUPLICATE KEY UPDATE
                    missing_usd = VALUES(missing_usd),
                    missing_cny = VALUES(missing_cny),
                    status = 'open',
                    note = VALUES(note)
                """,
                (rate_date, missing_usd, missing_cny),
            )
        else:
            cur.execute(
                """
                INSERT INTO exchange_rate_gaps (rate_date, missing_usd, missing_cny, status, resolved_at, note)
                VALUES (%s, FALSE, FALSE, 'resolved', CURRENT_TIMESTAMP, 'Resolved by manual rate entry.')
                ON DUPLICATE KEY UPDATE
                    missing_usd = FALSE,
                    missing_cny = FALSE,
                    status = 'resolved',
                    resolved_at = CURRENT_TIMESTAMP,
                    note = VALUES(note)
                """,
                (rate_date,),
            )
        result = recalculate_all_wholesale(cur)
        log_operation(
            cur,
            "save_exchange_rate",
            "exchange_rates",
            rate_date.isoformat(),
            {"usdKgs": str(usd) if usd is not None else None, "cnyKgs": str(cny) if cny is not None else None},
        )
        conn.commit()
        return jsonify({"ok": True, "recalculation": result})
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@app.route("/api/admin/products")
@admin_required
def api_admin_products():
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT id, product_code, display_name, category, retail_resource_key, is_active, created_at
            FROM products
            ORDER BY category, product_code
            """
        )
        products = []
        by_id = {}
        for row in cur.fetchall():
            item = {
                "id": row[0],
                "code": row[1],
                "name": row[2],
                "category": row[3],
                "retailResourceKey": row[4],
                "isActive": bool(row[5]),
                "createdAt": row[6].isoformat() if row[6] else None,
                "aliases": [],
            }
            products.append(item)
            by_id[item["id"]] = item
        cur.execute(
            """
            SELECT id, product_id, alias_name, source, created_at
            FROM product_aliases
            ORDER BY product_id, alias_name
            """
        )
        for row in cur.fetchall():
            if row[1] in by_id:
                by_id[row[1]]["aliases"].append({
                    "id": row[0],
                    "name": row[2],
                    "source": row[3],
                    "createdAt": row[4].isoformat() if row[4] else None,
                })
        return jsonify({"items": products})
    finally:
        conn.close()


@app.route("/api/admin/product-aliases", methods=["POST"])
@admin_required
def api_admin_create_product_alias():
    payload = request.get_json(force=True)
    product_code = clean_str(payload.get("productCode"))
    alias_name = clean_str(payload.get("aliasName"))
    if not product_code or not alias_name:
        return jsonify({"error": "请选择油品并填写别名"}), 400
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute("SELECT id FROM products WHERE product_code = %s", (product_code,))
        row = cur.fetchone()
        if not row:
            return jsonify({"error": "油品不存在"}), 404
        normalized = normalize_text(alias_name)
        cur.execute(
            """
            INSERT INTO product_aliases (product_id, alias_name, normalized_alias, source)
            VALUES (%s, %s, %s, 'admin')
            ON DUPLICATE KEY UPDATE
                product_id = VALUES(product_id),
                alias_name = VALUES(alias_name),
                source = 'admin'
            """,
            (row[0], alias_name, normalized),
        )
        log_operation(cur, "create_product_alias", "product_aliases", cur.lastrowid, {"productCode": product_code, "alias": alias_name})
        conn.commit()
        return jsonify({"ok": True})
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@app.route("/api/admin/customers")
@admin_required
def api_admin_customers():
    customer_type = request.args.get("customerType")
    keyword = clean_str(request.args.get("keyword"))
    filters = []
    params = []
    if customer_type:
        filters.append("c.customer_type = %s")
        params.append(customer_type)
    if keyword:
        filters.append("c.standard_name LIKE %s")
        params.append(f"%{keyword}%")
    where = "WHERE " + " AND ".join(filters) if filters else ""
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            f"""
            SELECT
                c.id,
                c.standard_name,
                c.customer_type,
                c.created_at,
                c.updated_at,
                COUNT(DISTINCT ct.id) AS contract_count,
                COUNT(DISTINCT d.id) AS detail_count,
                COALESCE(SUM(d.weight_tons), 0) AS detail_tons
            FROM customers c
            LEFT JOIN contracts ct ON ct.customer_id = c.id
            LEFT JOIN wholesale_detail_records d
                ON d.entity_type = 'wholesale_customer'
               AND d.raw_entity_name IS NOT NULL
               AND d.raw_entity_name <> ''
               AND d.standard_entity_name = c.standard_name
            {where}
            GROUP BY c.id, c.standard_name, c.customer_type, c.created_at, c.updated_at
            ORDER BY c.updated_at DESC, c.created_at DESC, c.standard_name
            LIMIT 500
            """,
            tuple(params),
        )
        items = [
            {
                "id": row[0],
                "name": row[1],
                "customerType": row[2],
                "createdAt": row[3].isoformat() if row[3] else None,
                "updatedAt": row[4].isoformat() if row[4] else None,
                "contractCount": int(row[5] or 0),
                "detailCount": int(row[6] or 0),
                "detailTons": round(float(row[7] or 0), 3),
            }
            for row in cur.fetchall()
        ]
        return jsonify({"items": items, "total": len(items)})
    finally:
        conn.close()


@app.route("/api/admin/customers", methods=["POST"])
@admin_required
def api_admin_save_customer():
    payload = request.get_json(force=True)
    name = clean_str(payload.get("name"))
    customer_type = payload.get("customerType") if payload.get("customerType") in ("local_wholesale", "export_wholesale", "mixed", "unknown") else "unknown"
    if not name:
        return jsonify({"error": "客户名称必填"}), 400
    conn = get_connection()
    try:
        cur = conn.cursor()
        customer_id = get_or_create_customer(cur, name, customer_type)
        cur.execute(
            """
            UPDATE customers
            SET standard_name = %s,
                customer_type = %s
            WHERE id = %s
            """,
            (name, customer_type, customer_id),
        )
        recalculation = recalculate_all_wholesale(cur)
        log_operation(cur, "save_customer", "customers", customer_id, {"name": name, "customerType": customer_type})
        conn.commit()
        return jsonify({"ok": True, "id": customer_id, "recalculation": recalculation})
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@app.route("/api/admin/entity-options")
@admin_required
def api_admin_entity_options():
    catalog = retail_station_catalog()
    retail = [
        {
            "entityType": "retail_station",
            "entityCode": str(item["shopKey"]),
            "name": item["name"],
            "label": f"{item['shopKey']} · {item['name']}（{item['short']}）",
        }
        for item in catalog.values()
    ]
    depots = [
        {"entityType": "own_depot", "entityCode": "north_CTY-22-LAT_depot", "name": "S-66", "label": "S-66"},
        {"entityType": "own_depot", "entityCode": "south_CTY-30_depot", "name": "S-64", "label": "S-64"},
        {"entityType": "own_depot", "entityCode": "hq_depot", "name": "STN-61", "label": "STN-61"},
    ]
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute("SELECT standard_name, customer_type FROM customers ORDER BY standard_name LIMIT 1000")
        customers = [
            {
                "entityType": "wholesale_customer",
                "entityCode": None,
                "name": row[0],
                "label": f"{row[0]} · {row[1]}",
            }
            for row in cur.fetchall()
        ]
        return jsonify({"items": retail + depots + customers})
    finally:
        conn.close()


@app.route("/api/admin/entity-aliases")
@admin_required
def api_admin_entity_aliases():
    entity_type = request.args.get("entityType")
    keyword = clean_str(request.args.get("keyword"))
    filters = []
    params = []
    if entity_type:
        filters.append("entity_type = %s")
        params.append(entity_type)
    if keyword:
        filters.append("(standard_name LIKE %s OR alias_name LIKE %s)")
        params.extend([f"%{keyword}%", f"%{keyword}%"])
    where = "WHERE " + " AND ".join(filters) if filters else ""
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            f"""
            SELECT id, entity_type, entity_code, standard_name, alias_name, confidence, source, created_at
            FROM entity_aliases
            {where}
            ORDER BY entity_type, standard_name, confidence DESC, alias_name
            LIMIT 800
            """,
            tuple(params),
        )
        items = [
            {
                "id": row[0],
                "entityType": row[1],
                "entityCode": row[2],
                "standardName": row[3],
                "aliasName": row[4],
                "confidence": int(row[5] or 0),
                "source": row[6],
                "createdAt": row[7].isoformat() if row[7] else None,
            }
            for row in cur.fetchall()
        ]
        return jsonify({"items": items, "total": len(items)})
    finally:
        conn.close()


@app.route("/api/admin/entity-aliases", methods=["POST"])
@admin_required
def api_admin_save_entity_alias():
    payload = request.get_json(force=True)
    entity_type = payload.get("entityType") if payload.get("entityType") in ("retail_station", "own_depot", "wholesale_customer", "unknown") else "unknown"
    entity_code = clean_str(payload.get("entityCode"))
    standard_name = clean_str(payload.get("standardName"))
    alias_name = clean_str(payload.get("aliasName"))
    confidence = int(payload.get("confidence") or 100)
    confidence = min(100, max(1, confidence))
    if not standard_name or not alias_name:
        return jsonify({"error": "标准名称和别名必填"}), 400
    normalized = normalize_station_name(alias_name) if entity_type == "retail_station" else normalize_text(alias_name)
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO entity_aliases
                (entity_type, entity_code, standard_name, alias_name, normalized_alias, confidence, source)
            VALUES (%s, %s, %s, %s, %s, %s, 'admin')
            ON DUPLICATE KEY UPDATE
                entity_type = VALUES(entity_type),
                entity_code = VALUES(entity_code),
                standard_name = VALUES(standard_name),
                alias_name = VALUES(alias_name),
                confidence = VALUES(confidence),
                source = 'admin'
            """,
            (entity_type, entity_code, standard_name, alias_name, normalized, confidence),
        )
        if entity_type == "wholesale_customer":
            get_or_create_customer(cur, standard_name, "unknown")
            recalculation = recalculate_all_wholesale(cur)
        else:
            recalculation = None
        log_operation(
            cur,
            "save_entity_alias",
            "entity_aliases",
            cur.lastrowid,
            {"entityType": entity_type, "standardName": standard_name, "aliasName": alias_name},
        )
        conn.commit()
        return jsonify({"ok": True, "recalculation": recalculation})
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@app.route("/api/admin/contracts")
@admin_required
def api_admin_contracts():
    year_raw = request.args.get("year")
    month_raw = request.args.get("month")
    sales_type = request.args.get("salesType")
    product_code = clean_str(request.args.get("productCode"))
    customer = clean_str(request.args.get("customer"))
    params = []
    filters = []
    if year_raw:
        start, end = period_bounds(int(year_raw), int(month_raw) if month_raw else None)
        filters.append("(c.effective_end IS NULL OR c.effective_end >= %s)")
        filters.append("(c.effective_start IS NULL OR c.effective_start < %s)")
        params.extend([start, end])
    if sales_type:
        filters.append("c.sales_type = %s")
        params.append(sales_type)
    if product_code:
        filters.append("p.product_code = %s")
        params.append(product_code)
    if customer:
        filters.append("(c.raw_customer_name LIKE %s OR cu.standard_name LIKE %s)")
        params.extend([f"%{customer}%", f"%{customer}%"])
    where = "WHERE " + " AND ".join(filters) if filters else ""
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            f"""
            SELECT
                c.id,
                c.raw_customer_name,
                COALESCE(cu.standard_name, c.raw_customer_name) AS customer_name,
                c.contract_no,
                c.appendix_no,
                c.signed_date,
                c.effective_start,
                c.effective_end,
                p.product_code,
                COALESCE(p.display_name, c.raw_product_name) AS product_name,
                c.raw_product_name,
                c.destination,
                c.sales_type,
                c.price_with_logistics,
                c.price_without_logistics,
                c.contract_quantity_tons,
                c.contract_amount_original,
                c.currency_code,
                c.remark,
                COUNT(d.id) AS matched_rows,
                COALESCE(SUM(d.weight_tons), 0) AS matched_tons,
                COALESCE(SUM(cr.amount_cny), 0) AS amount_cny
            FROM contracts c
            LEFT JOIN customers cu ON cu.id = c.customer_id
            LEFT JOIN products p ON p.id = c.product_id
            LEFT JOIN wholesale_detail_records d ON d.contract_id = c.id
            LEFT JOIN calculation_results cr ON cr.detail_id = d.id
            {where}
            GROUP BY
                c.id, c.raw_customer_name, cu.standard_name, c.contract_no, c.appendix_no,
                c.signed_date, c.effective_start, c.effective_end, p.product_code, p.display_name,
                c.raw_product_name, c.destination, c.sales_type, c.price_with_logistics,
                c.price_without_logistics, c.contract_quantity_tons, c.contract_amount_original,
                c.currency_code, c.remark
            ORDER BY COALESCE(c.effective_start, c.signed_date, c.created_at) DESC, c.id DESC
            LIMIT 500
            """,
            tuple(params),
        )
        items = [
            {
                "id": row[0],
                "rawCustomerName": row[1],
                "customerName": row[2],
                "contractNo": row[3],
                "appendixNo": row[4],
                "signedDate": row[5].isoformat() if row[5] else None,
                "effectiveStart": row[6].isoformat() if row[6] else None,
                "effectiveEnd": row[7].isoformat() if row[7] else None,
                "productCode": row[8],
                "productName": row[9],
                "rawProductName": row[10],
                "destination": row[11],
                "salesType": row[12],
                "priceWithLogistics": float(row[13]) if row[13] is not None else None,
                "priceWithoutLogistics": float(row[14]) if row[14] is not None else None,
                "quantityTons": float(row[15]) if row[15] is not None else None,
                "amountOriginal": float(row[16]) if row[16] is not None else None,
                "currencyCode": row[17],
                "remark": row[18] or "",
                "matchedRows": int(row[19] or 0),
                "matchedTons": round(float(row[20] or 0), 3),
                "amountCny": round(float(row[21] or 0), 2),
            }
            for row in cur.fetchall()
        ]
        return jsonify({"items": items, "total": len(items)})
    finally:
        conn.close()


@app.route("/api/admin/contracts/<int:contract_id>", methods=["PATCH"])
@admin_required
def api_admin_update_contract(contract_id):
    payload = request.get_json(force=True)
    customer = clean_str(payload.get("customerName"))
    contract_no = clean_str(payload.get("contractNo"))
    raw_product = clean_str(payload.get("productName"))
    product_code = clean_str(payload.get("productCode"))
    sales_type = payload.get("salesType") if payload.get("salesType") in ("local_wholesale", "export_wholesale", "internal", "unknown") else "unknown"
    if not customer or not contract_no or (not raw_product and not product_code):
        return jsonify({"error": "客户、合同号、油品必填"}), 400
    conn = get_connection()
    try:
        cur = conn.cursor()
        product = None
        if product_code:
            cur.execute("SELECT id, display_name FROM products WHERE product_code = %s", (product_code,))
            row = cur.fetchone()
            if row:
                product = {"id": row[0], "display_name": row[1]}
                raw_product = raw_product or row[1]
        if not product:
            product = lookup_product(load_product_lookup(cur), raw_product)
        customer_id = get_or_create_customer(cur, customer, sales_type if sales_type != "internal" else "unknown")
        cur.execute(
            """
            UPDATE contracts
            SET customer_id = %s,
                raw_customer_name = %s,
                contract_no = %s,
                appendix_no = %s,
                signed_date = %s,
                effective_start = %s,
                effective_end = %s,
                product_id = %s,
                raw_product_name = %s,
                destination = %s,
                sales_type = %s,
                price_with_logistics = %s,
                price_without_logistics = %s,
                contract_quantity_tons = %s,
                contract_amount_original = %s,
                currency_code = %s,
                remark = %s
            WHERE id = %s
            """,
            (
                customer_id,
                customer,
                contract_no,
                clean_str(payload.get("appendixNo")),
                parse_date(payload.get("signedDate")),
                parse_date(payload.get("effectiveStart")),
                parse_date(payload.get("effectiveEnd")),
                product["id"] if product else None,
                raw_product,
                clean_str(payload.get("destination")),
                sales_type,
                parse_decimal(payload.get("priceWithLogistics")),
                parse_decimal(payload.get("priceWithoutLogistics")),
                parse_decimal(payload.get("quantityTons")),
                parse_decimal(payload.get("amountOriginal")),
                (payload.get("currencyCode") or "KGS").upper(),
                clean_str(payload.get("remark")),
                contract_id,
            ),
        )
        if cur.rowcount == 0:
            conn.rollback()
            return jsonify({"error": "合同不存在"}), 404
        result = recalculate_all_wholesale(cur)
        log_operation(
            cur,
            "update_contract",
            "contracts",
            contract_id,
            {"customerName": customer, "contractNo": contract_no, "salesType": sales_type, "productCode": product_code},
        )
        conn.commit()
        return jsonify({"ok": True, "recalculation": result})
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@app.route("/api/admin/operation-logs")
@admin_required
def api_admin_operation_logs():
    conn = get_connection()
    try:
        cur = conn.cursor()
        ensure_operation_log_schema(cur)
        cur.execute(
            """
            SELECT id, username, action, target_type, target_id, detail, created_at
            FROM operation_logs
            ORDER BY created_at DESC, id DESC
            LIMIT 120
            """
        )
        items = [
            {
                "id": row[0],
                "username": row[1] or "system",
                "action": row[2],
                "targetType": row[3],
                "targetId": row[4],
                "detail": row[5],
                "createdAt": row[6].isoformat() if row[6] else None,
            }
            for row in cur.fetchall()
        ]
        return jsonify({"items": items})
    finally:
        conn.close()


@app.route("/api/periods")
def api_periods():
    conn = get_connection()
    try:
        cur = conn.cursor()
        return jsonify(available_periods(cur))
    finally:
        conn.close()


@app.route("/api/summary")
def api_summary():
    year = int(request.args.get("year", "2026"))
    month_raw = request.args.get("month")
    month = int(month_raw) if month_raw else None
    conn = get_connection()
    try:
        cur = conn.cursor()
        return jsonify(summarize_combined(cur, year, month))
    finally:
        conn.close()


@app.route("/api/home/exchange-current")
def api_home_exchange_current():
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT rate_date, usd_kgs, cny_kgs, source, is_manual, note
            FROM exchange_rates
            WHERE rate_date <= CURDATE()
              AND (usd_kgs IS NOT NULL OR cny_kgs IS NOT NULL)
            ORDER BY rate_date DESC
            LIMIT 1
            """
        )
        row = cur.fetchone()
        if not row:
            return jsonify({
                "available": False,
                "message": "暂无汇率快照，请在后台补录。",
            })
        rate_date, usd_kgs, cny_kgs, source, is_manual, note = row
        return jsonify({
            "available": True,
            "date": rate_date.isoformat() if rate_date else None,
            "usdKgs": float(usd_kgs) if usd_kgs is not None else None,
            "cnyKgs": float(cny_kgs) if cny_kgs is not None else None,
            "source": source,
            "isManual": bool(is_manual),
            "isToday": rate_date == date.today() if rate_date else False,
        })
    finally:
        conn.close()


@app.route("/api/home/market-watch")
def api_home_market_watch():
    conn = get_connection()
    try:
        cur = conn.cursor()
        ensure_market_watch_schema(cur)
        cur.execute(
            """
            SELECT id, market_region, market_type, title, summary,
                   price_text, source_name, source_url, observed_at
            FROM market_watch_items
            WHERE is_active = TRUE
              AND market_region IN ('Russia', 'Uzbekistan', 'Global')
            ORDER BY observed_at DESC, priority ASC, id DESC
            LIMIT 30
            """
        )
        rows = [
            {
                "id": row[0],
                "region": row[1],
                "type": row[2],
                "typeLabel": market_type_label(row[2]),
                "title": row[3],
                "summary": row[4],
                "priceText": row[5] or "",
                "sourceName": row[6] or "",
                "sourceUrl": row[7] or "",
                "observedAt": row[8].isoformat() if row[8] else None,
            }
            for row in cur.fetchall()
        ]
        rows = [chinese_market_item(item) for item in rows]
        by_region = {}
        grouped = {}
        for item in rows:
            grouped.setdefault(item["region"], []).append(item)
            by_region.setdefault(item["region"], item)
        items = []
        for key in ("Russia", "Uzbekistan", "Global"):
            if key not in by_region:
                continue
            region_rows = grouped.get(key, [])
            display_rows = [row for row in region_rows if row.get("displayReady")]
            item = dict(display_rows[0] if display_rows else by_region[key])
            detail_rows = [
                {
                    "title": detail.get("title") or "",
                    "summary": detail.get("summary") or "",
                    "priceText": detail.get("priceText") or "",
                    "typeLabel": detail.get("typeLabel") or market_type_label(detail.get("type")),
                    "sourceName": detail.get("sourceName") or "",
                    "sourceUrl": detail.get("sourceUrl") or "",
                    "observedAt": detail.get("observedAt"),
                }
                for detail in (display_rows or region_rows)[:5]
            ]
            seen_detail_labels = {detail.get("typeLabel") for detail in detail_rows if detail.get("typeLabel")}
            for fallback_detail in default_market_details(key):
                if fallback_detail["typeLabel"] not in seen_detail_labels:
                    detail_rows.append(fallback_detail)
                    seen_detail_labels.add(fallback_detail["typeLabel"])
            item["details"] = detail_rows[:5]
            items.append(item)
        latest_observed_at = None
        for item in rows:
            observed_at = item.get("observedAt")
            if observed_at and (latest_observed_at is None or observed_at > latest_observed_at):
                latest_observed_at = observed_at
        if not items:
            items = [
                {
                    "id": 0,
                    "region": "Russia",
                    "type": "refined",
                    "typeLabel": "成品油",
                    "title": "俄罗斯成品油市场",
                    "summary": "市场观察数据尚未维护，建议由后台录入俄罗斯炼厂开工、出口限制、铁路运输、批发价格和汽柴油供应变化。",
                    "sourceName": "",
                    "sourceUrl": "",
                    "observedAt": None,
                    "details": [
                        {"typeLabel": "市场行情", "title": "俄罗斯市场行情", "summary": "跟踪批发价格、区域供应和运输变化。"},
                        {"typeLabel": "原料油", "title": "俄罗斯原料油", "summary": "关注 Urals 原油、炼厂开工和出口流向。"},
                        {"typeLabel": "成品油", "title": "俄罗斯成品油", "summary": "关注汽油、柴油库存、出口政策和铁路发运。"},
                    ],
                },
                {
                    "id": 0,
                    "region": "Uzbekistan",
                    "type": "refined",
                    "typeLabel": "成品油",
                    "title": "乌兹别克成品油市场",
                    "summary": "市场观察数据尚未维护，建议跟踪当地汽柴油生产、进口、库存、批发价格和邻国供应变化。",
                    "sourceName": "",
                    "sourceUrl": "",
                    "observedAt": None,
                    "details": [
                        {"typeLabel": "市场行情", "title": "乌兹别克市场行情", "summary": "跟踪国内供应、进口节奏和批发成交。"},
                        {"typeLabel": "原料油", "title": "乌兹别克原料油", "summary": "关注凝析油、炼厂原料供应和生产负荷。"},
                        {"typeLabel": "成品油", "title": "乌兹别克成品油", "summary": "关注汽油、柴油供应、库存和周边价格变化。"},
                    ],
                },
                {
                    "id": 0,
                    "region": "Global",
                    "type": "crude",
                    "typeLabel": "原油",
                    "title": "国际原油参考",
                    "summary": "待录入 Brent、WTI、OPEC+、国际库存、裂解价差和炼厂利润等参考信息，用于辅助批发与零售经营判断。",
                    "sourceName": "",
                    "sourceUrl": "",
                    "observedAt": None,
                    "details": [
                        {"typeLabel": "市场行情", "title": "国际市场行情", "summary": "关注宏观需求、航运、库存和风险事件。"},
                        {"typeLabel": "原料油", "title": "国际原油", "summary": "关注 Brent、WTI、OPEC+ 和库存变化。"},
                        {"typeLabel": "成品油", "title": "国际成品油", "summary": "关注汽柴油裂解价差、炼厂利润和区域套利。"},
                    ],
                },
            ]
        return jsonify({"items": items, "updatedAt": latest_observed_at or date.today().isoformat()})
    finally:
        conn.close()


@app.route("/api/exchange-rates")
def api_exchange_rates():
    year = int(request.args.get("year", str(date.today().year)))
    month_raw = request.args.get("month")
    month = int(month_raw) if month_raw else None
    start, end = period_bounds(year, month)
    conn = get_connection()
    try:
        cur = conn.cursor()
        return jsonify(exchange_rates(cur, start, end))
    finally:
        conn.close()


@app.route("/api/wholesale")
def api_wholesale():
    year = int(request.args.get("year", "2026"))
    month_raw = request.args.get("month")
    month = int(month_raw) if month_raw else None
    conn = get_connection()
    try:
        cur = conn.cursor()
        return jsonify(summarize_wholesale_page(cur, year, month))
    finally:
        conn.close()


@app.route("/api/review/details")
def api_review_details():
    kind = request.args.get("kind", "pending")
    year = int(request.args.get("year", "2026"))
    month_raw = request.args.get("month")
    month = int(month_raw) if month_raw else None
    start, end = period_bounds(year, month)
    conn = get_connection()
    try:
        cur = conn.cursor()
        if kind == "issues":
            items = review_issue_rows(cur, start, end)
        else:
            items = review_detail_rows(cur, start, end, kind)
        return jsonify({"kind": kind, "items": items})
    finally:
        conn.close()


@app.route("/api/review/correct", methods=["POST"])
@admin_required
def api_review_correct():
    payload = request.get_json(force=True)
    detail_ids = [int(item) for item in payload.get("detailIds", []) if item]
    issue_ids = [int(item) for item in payload.get("issueIds", []) if item]
    classification = payload.get("classification") or None
    product_code = payload.get("productCode") or None
    entity_name = clean_str(payload.get("entityName"))
    issue_status = payload.get("issueStatus") if payload.get("issueStatus") in ("resolved", "ignored") else None
    conn = get_connection()
    try:
        cur = conn.cursor()
        batch_ids = set()
        product = None
        if product_code:
            cur.execute("SELECT id, display_name FROM products WHERE product_code = %s", (product_code,))
            row = cur.fetchone()
            if not row:
                return jsonify({"error": "未知油品"}), 400
            product = {"id": row[0], "display_name": row[1]}
        entity = lookup_entity(load_entity_lookup(cur), entity_name) if entity_name else None
        if detail_ids:
            cur.execute(
                "SELECT DISTINCT batch_id FROM wholesale_detail_records WHERE id IN %s",
                (tuple(detail_ids),),
            )
            batch_ids = {row[0] for row in cur.fetchall() if row[0]}
            updates = []
            params = []
            if classification:
                updates.append("classification = %s")
                params.append(classification)
            if product:
                updates.extend(["product_id = %s", "raw_product_name = %s"])
                params.extend([product["id"], product["display_name"]])
            if entity_name:
                updates.extend([
                    "raw_entity_name = %s",
                    "entity_type = %s",
                    "entity_code = %s",
                    "standard_entity_name = %s",
                ])
                params.extend([
                    entity_name,
                    entity["entity_type"] if entity else "unknown",
                    entity["entity_code"] if entity else None,
                    entity["standard_name"] if entity else None,
                ])
            if updates:
                params.append(tuple(detail_ids))
                cur.execute(
                    f"UPDATE wholesale_detail_records SET {', '.join(updates)} WHERE id IN %s",
                    tuple(params),
                )
                if entity_name and classification in ("local_wholesale", "export_wholesale", "unknown", None):
                    get_or_create_customer(cur, entity_name, classification or "unknown")
        if issue_ids and issue_status:
            cur.execute(
                """
                UPDATE data_quality_issues
                SET status = %s,
                    resolved_at = CURRENT_TIMESTAMP
                WHERE id IN %s
                """,
                (issue_status, tuple(issue_ids)),
            )
        recalc = {"matched": 0, "unmatched": 0, "missing_rates": 0}
        for batch_id in batch_ids:
            result = recalculate_wholesale_batch(cur, batch_id)
            for key in recalc:
                recalc[key] += result[key]
        log_operation(
            cur,
            "review_correct",
            "review",
            ",".join([str(item) for item in detail_ids + issue_ids]) or None,
            {
                "detailIds": detail_ids,
                "issueIds": issue_ids,
                "classification": classification,
                "productCode": product_code,
                "entityName": entity_name,
                "issueStatus": issue_status,
            },
        )
        conn.commit()
        return jsonify({
            "ok": True,
            "updatedDetails": len(detail_ids),
            "updatedIssues": len(issue_ids) if issue_status else 0,
            "recalculation": recalc,
        })
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@app.route("/api/admin/import-batches")
@admin_required
def api_admin_import_batches():
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT
                b.id,
                b.source_type,
                b.original_filename,
                b.period_year,
                b.period_month,
                b.status,
                b.created_at,
                b.imported_at,
                b.note,
                COALESCE(rr.raw_rows, 0),
                COALESCE(d.detail_rows, 0),
                COALESCE(d.detail_tons, 0),
                COALESCE(s.summary_rows, 0),
                COALESCE(c.contract_rows, 0)
            FROM import_batches b
            LEFT JOIN (
                SELECT batch_id, COUNT(*) raw_rows
                FROM import_raw_rows
                GROUP BY batch_id
            ) rr ON rr.batch_id = b.id
            LEFT JOIN (
                SELECT batch_id, COUNT(*) detail_rows, COALESCE(SUM(weight_tons), 0) detail_tons
                FROM wholesale_detail_records
                GROUP BY batch_id
            ) d ON d.batch_id = b.id
            LEFT JOIN (
                SELECT batch_id, COUNT(*) summary_rows
                FROM wholesale_summary_records
                GROUP BY batch_id
            ) s ON s.batch_id = b.id
            LEFT JOIN (
                SELECT batch_id, COUNT(*) contract_rows
                FROM contracts
                GROUP BY batch_id
            ) c ON c.batch_id = b.id
            ORDER BY COALESCE(b.imported_at, b.created_at) DESC, b.id DESC
            LIMIT 200
            """
        )
        items = []
        for row in cur.fetchall():
            period = ""
            if row[3]:
                period = f"{row[3]}-{int(row[4] or 0):02d}" if row[4] else str(row[3])
            items.append(
                {
                    "id": row[0],
                    "sourceType": row[1],
                    "sourceLabel": import_source_label(row[1]),
                    "filename": row[2],
                    "period": period,
                    "status": row[5],
                    "createdAt": row[6].isoformat() if row[6] else None,
                    "importedAt": row[7].isoformat() if row[7] else None,
                    "note": row[8],
                    "rawRows": int(row[9] or 0),
                    "detailRows": int(row[10] or 0),
                    "detailTons": round(float(row[11] or 0), 3),
                    "summaryRows": int(row[12] or 0),
                    "contractRows": int(row[13] or 0),
                }
            )
        return jsonify({"items": items})
    finally:
        conn.close()


@app.route("/api/admin/import-batches/<int:batch_id>/rollback", methods=["POST"])
@admin_required
def api_admin_rollback_import_batch(batch_id):
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            "SELECT source_type, original_filename, status FROM import_batches WHERE id = %s",
            (batch_id,),
        )
        batch = cur.fetchone()
        if not batch:
            return jsonify({"error": "导入批次不存在"}), 404

        source_type, filename, status = batch
        cur.execute("SELECT COUNT(*) FROM wholesale_detail_records WHERE batch_id = %s", (batch_id,))
        detail_rows = int(cur.fetchone()[0] or 0)
        cur.execute("SELECT COUNT(*) FROM wholesale_summary_records WHERE batch_id = %s", (batch_id,))
        summary_rows = int(cur.fetchone()[0] or 0)
        cur.execute("SELECT COUNT(*) FROM contracts WHERE batch_id = %s", (batch_id,))
        contract_rows = int(cur.fetchone()[0] or 0)
        cur.execute("SELECT COUNT(*) FROM import_raw_rows WHERE batch_id = %s", (batch_id,))
        raw_rows = int(cur.fetchone()[0] or 0)

        cur.execute(
            """
            DELETE cr
            FROM calculation_results cr
            JOIN wholesale_detail_records d ON d.id = cr.detail_id
            WHERE d.batch_id = %s
            """,
            (batch_id,),
        )
        cur.execute(
            """
            DELETE qi
            FROM data_quality_issues qi
            JOIN wholesale_detail_records d ON d.id = qi.related_id
            WHERE d.batch_id = %s
              AND qi.related_table = 'wholesale_detail_records'
            """,
            (batch_id,),
        )
        cur.execute("DELETE FROM wholesale_summary_records WHERE batch_id = %s", (batch_id,))
        cur.execute("DELETE FROM wholesale_detail_records WHERE batch_id = %s", (batch_id,))
        cur.execute("DELETE FROM contracts WHERE batch_id = %s", (batch_id,))
        cur.execute("DELETE FROM import_raw_rows WHERE batch_id = %s", (batch_id,))
        rollback_note = f"\nRolled back by {session_user().get('username', 'system')}."
        cur.execute(
            """
            UPDATE import_batches
            SET status = 'failed',
                note = CONCAT(COALESCE(note, ''), %s)
            WHERE id = %s
            """,
            (rollback_note, batch_id),
        )
        recalculation = recalculate_all_wholesale(cur)
        log_operation(
            cur,
            "rollback_import_batch",
            "import_batches",
            batch_id,
            {
                "sourceType": source_type,
                "filename": filename,
                "previousStatus": status,
                "rawRows": raw_rows,
                "detailRows": detail_rows,
                "summaryRows": summary_rows,
                "contractRows": contract_rows,
            },
        )
        conn.commit()
        return jsonify(
            {
                "ok": True,
                "batchId": batch_id,
                "removed": {
                    "rawRows": raw_rows,
                    "detailRows": detail_rows,
                    "summaryRows": summary_rows,
                    "contractRows": contract_rows,
                },
                "recalculation": recalculation,
            }
        )
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def rollback_import_batch_records(cur, batch_id, *, reason="Rolled back by organizer."):
    cur.execute(
        "SELECT source_type, original_filename, status, period_year, period_month FROM import_batches WHERE id = %s",
        (batch_id,),
    )
    batch = cur.fetchone()
    if not batch:
        return None

    source_type, filename, status, period_year, period_month = batch
    period = f"{period_year}-{int(period_month or 0):02d}" if period_year and period_month else (str(period_year) if period_year else "")
    cur.execute("SELECT COUNT(*) FROM wholesale_detail_records WHERE batch_id = %s", (batch_id,))
    detail_rows = int(cur.fetchone()[0] or 0)
    cur.execute("SELECT COUNT(*) FROM wholesale_summary_records WHERE batch_id = %s", (batch_id,))
    summary_rows = int(cur.fetchone()[0] or 0)
    cur.execute("SELECT COUNT(*) FROM contracts WHERE batch_id = %s", (batch_id,))
    contract_rows = int(cur.fetchone()[0] or 0)
    cur.execute("SELECT COUNT(*) FROM import_raw_rows WHERE batch_id = %s", (batch_id,))
    raw_rows = int(cur.fetchone()[0] or 0)
    cur.execute(
        """
        DELETE cr
        FROM calculation_results cr
        JOIN wholesale_detail_records d ON d.id = cr.detail_id
        WHERE d.batch_id = %s
        """,
        (batch_id,),
    )
    cur.execute(
        """
        DELETE qi
        FROM data_quality_issues qi
        JOIN wholesale_detail_records d ON d.id = qi.related_id
        WHERE d.batch_id = %s
          AND qi.related_table = 'wholesale_detail_records'
        """,
        (batch_id,),
    )
    cur.execute("DELETE FROM wholesale_summary_records WHERE batch_id = %s", (batch_id,))
    cur.execute("DELETE FROM wholesale_detail_records WHERE batch_id = %s", (batch_id,))
    cur.execute("DELETE FROM contracts WHERE batch_id = %s", (batch_id,))
    cur.execute("DELETE FROM import_raw_rows WHERE batch_id = %s", (batch_id,))
    cur.execute(
        """
        UPDATE import_batches
        SET status = 'failed',
            note = CONCAT(COALESCE(note, ''), %s)
        WHERE id = %s
        """,
        (f"\n{reason}", batch_id),
    )
    return {
        "batchId": batch_id,
        "sourceType": source_type,
        "filename": filename,
        "period": period,
        "previousStatus": status,
        "rawRows": raw_rows,
        "detailRows": detail_rows,
        "summaryRows": summary_rows,
        "contractRows": contract_rows,
    }


@app.route("/api/admin/wholesale/organize", methods=["POST"])
@admin_required
def api_admin_organize_wholesale():
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT id, period_year, period_month, original_filename, imported_at, created_at
            FROM import_batches
            WHERE source_type = 'export_sales_excel'
              AND status = 'imported'
              AND period_year IS NOT NULL
            ORDER BY period_year, period_month, imported_at DESC, created_at DESC, id DESC
            """
        )
        keep_by_period = {}
        rolled_back = []
        for batch_id, period_year, period_month, filename, imported_at, created_at in cur.fetchall():
            period = f"{period_year}-{int(period_month or 0):02d}" if period_month else str(period_year)
            if period not in keep_by_period:
                keep_by_period[period] = {
                    "batchId": batch_id,
                    "period": period,
                    "filename": filename,
                    "importedAt": imported_at.isoformat() if imported_at else None,
                    "createdAt": created_at.isoformat() if created_at else None,
                }
                continue
            rollback = rollback_import_batch_records(
                cur,
                batch_id,
                reason=f"Rolled back duplicate period {period} by wholesale organizer.",
            )
            if rollback:
                rolled_back.append(rollback)

        product_reconciliation = reconcile_wholesale_product_ids(cur)
        recalculation = recalculate_all_wholesale(cur)
        cur.execute(
            """
            SELECT match_status, classification, COUNT(*) AS rows_count, COALESCE(SUM(weight_tons), 0) AS tons
            FROM wholesale_detail_records
            GROUP BY match_status, classification
            ORDER BY match_status, classification
            """
        )
        status_rows = [
            {
                "matchStatus": row[0] or "unknown",
                "classification": row[1] or "unknown",
                "rows": int(row[2] or 0),
                "tons": round(float(row[3] or 0), 3),
            }
            for row in cur.fetchall()
        ]
        cur.execute(
            """
            SELECT issue_type, COUNT(*) AS rows_count
            FROM data_quality_issues
            WHERE related_table = 'wholesale_detail_records'
              AND status = 'open'
            GROUP BY issue_type
            ORDER BY rows_count DESC, issue_type
            """
        )
        issue_rows = [
            {"issueType": row[0] or "unknown", "rows": int(row[1] or 0)}
            for row in cur.fetchall()
        ]
        log_operation(
            cur,
            "organize_wholesale_data",
            "wholesale",
            None,
            {
                "keptBatches": list(keep_by_period.values()),
                "rolledBack": rolled_back,
                "productReconciliation": product_reconciliation,
                "recalculation": recalculation,
                "statusRows": status_rows,
                "issueRows": issue_rows,
            },
        )
        conn.commit()
        return jsonify({
            "ok": True,
            "keptBatches": list(keep_by_period.values()),
            "rolledBack": rolled_back,
            "productReconciliation": product_reconciliation,
            "recalculation": recalculation,
            "statusRows": status_rows,
            "issueRows": issue_rows,
        })
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def scan_wholesale_duplicate_product_conflicts(cur):
    issue_type = "possible_duplicate_conflicting_product"
    ensure_data_quality_issue_schema(cur)
    cur.execute(
        """
        UPDATE data_quality_issues
        SET status = 'resolved',
            resolved_at = CURRENT_TIMESTAMP
        WHERE issue_type = %s
          AND related_table = 'wholesale_detail_records'
          AND status = 'open'
        """,
        (issue_type,),
    )
    resolved_previous = cur.rowcount
    cur.execute(
        """
        SELECT
            business_date,
            entity_key,
            vehicle_key,
            tons_key,
            MIN(entity_name) AS entity_name,
            MIN(vehicle_no) AS vehicle_no,
            COUNT(*) AS row_count,
            COUNT(DISTINCT product_key) AS product_count,
            GROUP_CONCAT(DISTINCT raw_product_name ORDER BY raw_product_name SEPARATOR ' / ') AS products
        FROM (
            SELECT
                id,
                business_date,
                COALESCE(NULLIF(TRIM(standard_entity_name), ''), TRIM(raw_entity_name)) AS entity_name,
                LOWER(COALESCE(NULLIF(TRIM(standard_entity_name), ''), TRIM(raw_entity_name))) AS entity_key,
                LOWER(REPLACE(REPLACE(TRIM(COALESCE(vehicle_no, '')), ' ', ''), '-', '')) AS vehicle_key,
                vehicle_no,
                raw_product_name,
                LOWER(TRIM(raw_product_name)) AS product_key,
                ROUND(weight_tons, 3) AS tons_key
            FROM wholesale_detail_records
            WHERE business_date IS NOT NULL
              AND weight_tons IS NOT NULL
              AND vehicle_no IS NOT NULL
              AND TRIM(vehicle_no) <> ''
              AND raw_entity_name IS NOT NULL
              AND TRIM(raw_entity_name) <> ''
        ) grouped
        WHERE entity_key <> ''
          AND vehicle_key <> ''
        GROUP BY business_date, entity_key, vehicle_key, tons_key
        HAVING row_count > 1
           AND product_count > 1
        ORDER BY business_date DESC, row_count DESC
        LIMIT 500
        """
    )
    groups = []
    issue_rows = []
    for row in cur.fetchall():
        business_date, entity_key, vehicle_key, tons_key, entity_name, vehicle_no, row_count, product_count, products = row
        cur.execute(
            """
            SELECT id, raw_product_name, classification, match_status, batch_id
            FROM wholesale_detail_records
            WHERE business_date = %s
              AND LOWER(COALESCE(NULLIF(TRIM(standard_entity_name), ''), TRIM(raw_entity_name))) = %s
              AND LOWER(REPLACE(REPLACE(TRIM(COALESCE(vehicle_no, '')), ' ', ''), '-', '')) = %s
              AND ROUND(weight_tons, 3) = %s
            ORDER BY id
            """,
            (business_date, entity_key, vehicle_key, tons_key),
        )
        details = [
            {
                "id": detail[0],
                "product": detail[1],
                "classification": detail[2],
                "matchStatus": detail[3],
                "batchId": detail[4],
            }
            for detail in cur.fetchall()
        ]
        message = (
            f"同一天、同一家公司、同一个车号与同一重量存在重复记录，但油品名称不一致。"
            f"请核对是否重复录入或油品录入错误。客户/去向：{entity_name or '--'}；"
            f"车号：{vehicle_no or '--'}；重量：{float(tons_key or 0):.3f} 吨；油品：{products or '--'}。"
        )
        for detail in details:
            cur.execute(
                """
                INSERT INTO data_quality_issues
                    (issue_type, severity, related_table, related_id, issue_date, message)
                VALUES (%s, 'warning', 'wholesale_detail_records', %s, %s, %s)
                """,
                (issue_type, detail["id"], business_date, message),
            )
            issue_rows.append({"issueId": cur.lastrowid, "detailId": detail["id"]})
        groups.append(
            {
                "date": business_date.isoformat() if business_date else None,
                "entity": entity_name,
                "vehicleNo": vehicle_no,
                "tons": round(float(tons_key or 0), 3),
                "rows": int(row_count or 0),
                "productCount": int(product_count or 0),
                "products": products,
                "details": details,
            }
        )
    return {
        "groups": groups,
        "groupCount": len(groups),
        "issueRows": len(issue_rows),
        "resolvedPrevious": int(resolved_previous or 0),
    }


@app.route("/api/admin/wholesale/deduplicate", methods=["POST"])
@admin_required
def api_admin_deduplicate_wholesale():
    conn = get_connection()
    try:
        cur = conn.cursor()
        result = scan_wholesale_duplicate_product_conflicts(cur)
        log_operation(
            cur,
            "detect_wholesale_duplicates",
            "wholesale",
            None,
            {
                "groupCount": result["groupCount"],
                "issueRows": result["issueRows"],
                "resolvedPrevious": result["resolvedPrevious"],
                "sampleGroups": result["groups"][:10],
            },
        )
        conn.commit()
        return jsonify({"ok": True, **result})
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@app.route("/api/reference/options")
def api_reference_options():
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute("SELECT product_code, display_name FROM products WHERE is_active = TRUE ORDER BY display_name")
        products = [{"code": row[0], "name": row[1]} for row in cur.fetchall()]
        cur.execute("SELECT standard_name, customer_type FROM customers ORDER BY standard_name LIMIT 500")
        customers = [{"name": row[0], "type": row[1]} for row in cur.fetchall()]
        return jsonify({"products": products, "customers": customers})
    finally:
        conn.close()


@app.route("/api/import/upload", methods=["POST"])
@admin_required
def api_import_upload():
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify({"error": "请选择 Excel 文件"}), 400
    UPLOAD_ROOT.mkdir(exist_ok=True)
    filename = secure_filename(file.filename) or "upload.xlsx"
    suffix = Path(filename).suffix or ".xlsx"
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix, dir=UPLOAD_ROOT) as tmp:
        file.save(tmp.name)
        saved_path = tmp.name
    try:
        upload_type = detect_upload_type(saved_path)
        if upload_type == "wholesale":
            batch_id = import_wholesale(saved_path)
            recalc = None
            conn = get_connection()
            try:
                cur = conn.cursor()
                cur.execute("UPDATE import_batches SET original_filename = %s WHERE id = %s", (file.filename, batch_id))
                log_operation(cur, "upload_wholesale_excel", "import_batches", batch_id, {"filename": file.filename})
                conn.commit()
            except Exception:
                conn.rollback()
                raise
            finally:
                conn.close()
        else:
            batch_id = import_contracts(saved_path)
            conn = get_connection()
            try:
                cur = conn.cursor()
                cur.execute("UPDATE import_batches SET original_filename = %s WHERE id = %s", (file.filename, batch_id))
                recalc = recalculate_all_wholesale(cur)
                log_operation(cur, "upload_contract_excel", "import_batches", batch_id, {"filename": file.filename})
                conn.commit()
            except Exception:
                conn.rollback()
                raise
            finally:
                conn.close()
        return jsonify({
            "ok": True,
            "batchId": batch_id,
            "uploadType": upload_type,
            "uploadLabel": upload_type_label(upload_type),
            "recalculation": recalc,
        })
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500


@app.route("/api/manual/wholesale-detail", methods=["POST"])
@admin_required
def api_manual_wholesale_detail():
    payload = request.get_json(force=True)
    business_date = parse_date(payload.get("businessDate"))
    raw_product = clean_str(payload.get("productName"))
    product_code = clean_str(payload.get("productCode"))
    raw_entity = clean_str(payload.get("entityName"))
    classification = payload.get("classification") or "unknown"
    weight_kg = parse_decimal(payload.get("weightKg"))
    if classification not in ("local_wholesale", "export_wholesale", "transfer", "processing", "other", "unknown"):
        classification = "unknown"
    if not business_date or not raw_entity or weight_kg is None or (not raw_product and not product_code):
        return jsonify({"error": "日期、油品、客户/去向、重量必填"}), 400
    conn = get_connection()
    try:
        cur = conn.cursor()
        product = None
        if product_code:
            cur.execute("SELECT id, display_name FROM products WHERE product_code = %s", (product_code,))
            row = cur.fetchone()
            if row:
                product = {"id": row[0], "display_name": row[1]}
                raw_product = raw_product or row[1]
        if not product:
            product = lookup_product(load_product_lookup(cur), raw_product)
        entity = lookup_entity(load_entity_lookup(cur), raw_entity)
        if classification in ("local_wholesale", "export_wholesale", "unknown"):
            get_or_create_customer(cur, raw_entity, classification if classification != "unknown" else "unknown")
        batch_id = create_manual_batch(cur, "Manual wholesale detail entry.", business_date)
        match_status = "matched" if classification in ("transfer", "processing", "other") else "unmatched"
        cur.execute(
            """
            INSERT INTO wholesale_detail_records
                (batch_id, source_row_number, business_date, raw_product_name, product_id,
                 vehicle_no, raw_entity_name, entity_type, entity_code, standard_entity_name,
                 classification, weight_kg, remark, match_status)
            VALUES (%s, 1, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            """,
            (
                batch_id,
                business_date,
                raw_product,
                product["id"] if product else None,
                clean_str(payload.get("vehicleNo")),
                raw_entity,
                entity["entity_type"] if entity else "unknown",
                entity["entity_code"] if entity else None,
                entity["standard_name"] if entity else None,
                classification,
                weight_kg,
                clean_str(payload.get("remark")),
                match_status,
            ),
        )
        if not product:
            cur.execute(
                """
                INSERT INTO data_quality_issues
                    (issue_type, severity, related_table, issue_date, message)
                VALUES ('unknown_product', 'warning', 'wholesale_detail_records', %s, %s)
                """,
                (business_date, f"Manual detail product is not mapped: {raw_product}"),
            )
        result = recalculate_wholesale_batch(cur, batch_id)
        log_operation(
            cur,
            "manual_wholesale_detail",
            "import_batches",
            batch_id,
            {"businessDate": business_date.isoformat(), "classification": classification, "entityName": raw_entity},
        )
        conn.commit()
        return jsonify({"ok": True, "batchId": batch_id, "recalculation": result})
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@app.route("/api/manual/contract", methods=["POST"])
@admin_required
def api_manual_contract():
    payload = request.get_json(force=True)
    customer = clean_str(payload.get("customerName"))
    contract_no = clean_str(payload.get("contractNo"))
    raw_product = clean_str(payload.get("productName"))
    product_code = clean_str(payload.get("productCode"))
    if not customer or not contract_no or (not raw_product and not product_code):
        return jsonify({"error": "客户、合同号、油品必填"}), 400
    conn = get_connection()
    try:
        cur = conn.cursor()
        product = None
        if product_code:
            cur.execute("SELECT id, display_name FROM products WHERE product_code = %s", (product_code,))
            row = cur.fetchone()
            if row:
                product = {"id": row[0], "display_name": row[1]}
                raw_product = raw_product or row[1]
        if not product:
            product = lookup_product(load_product_lookup(cur), raw_product)
        sales_type = payload.get("salesType") if payload.get("salesType") in ("local_wholesale", "export_wholesale", "internal", "unknown") else "unknown"
        customer_id = get_or_create_customer(cur, customer, sales_type if sales_type != "internal" else "unknown")
        signed_date = parse_date(payload.get("signedDate"))
        batch_id = create_manual_batch(cur, "Manual contract entry.", signed_date or date.today())
        cur.execute(
            """
            INSERT INTO contracts
                (batch_id, customer_id, raw_customer_name, contract_no, appendix_no,
                 signed_date, effective_start, effective_end, product_id, raw_product_name,
                 destination, sales_type, price_with_logistics, price_without_logistics,
                 contract_quantity_tons, contract_amount_original, currency_code, remark)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            """,
            (
                batch_id,
                customer_id,
                customer,
                contract_no,
                clean_str(payload.get("appendixNo")),
                signed_date,
                parse_date(payload.get("effectiveStart")),
                parse_date(payload.get("effectiveEnd")),
                product["id"] if product else None,
                raw_product,
                clean_str(payload.get("destination")),
                sales_type,
                parse_decimal(payload.get("priceWithLogistics")),
                parse_decimal(payload.get("priceWithoutLogistics")),
                parse_decimal(payload.get("quantityTons")),
                parse_decimal(payload.get("amountOriginal")),
                (payload.get("currencyCode") or "KGS").upper(),
                clean_str(payload.get("remark")),
            ),
        )
        result = recalculate_all_wholesale(cur)
        log_operation(
            cur,
            "manual_contract",
            "import_batches",
            batch_id,
            {"customerName": customer, "contractNo": contract_no, "salesType": sales_type},
        )
        conn.commit()
        return jsonify({"ok": True, "batchId": batch_id, "recalculation": result})
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@app.route("/api/densities")
def api_densities():
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT p.product_code, p.display_name, d.density, d.effective_date, d.source, d.created_at
            FROM product_density_history d
            JOIN products p ON p.id = d.product_id
            ORDER BY p.product_code, d.effective_date DESC
            """
        )
        rows = [
            {
                "productCode": row[0],
                "productName": row[1],
                "density": float(row[2]),
                "effectiveDate": row[3].isoformat(),
                "source": row[4],
                "createdAt": row[5].isoformat() if row[5] else None,
            }
            for row in cur.fetchall()
        ]
        return jsonify({"items": rows})
    finally:
        conn.close()


@app.route("/api/densities", methods=["POST"])
@admin_required
def api_create_density():
    payload = request.get_json(force=True)
    product_code = payload["productCode"]
    density = payload["density"]
    effective_date = payload["effectiveDate"]
    note = payload.get("note")
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute("SELECT id FROM products WHERE product_code = %s", (product_code,))
        row = cur.fetchone()
        if not row:
            return jsonify({"error": "unknown_product"}), 400
        cur.execute(
            """
            INSERT INTO product_density_history
                (product_id, density, effective_date, source, created_by, note)
            VALUES (%s, %s, %s, 'manual', 'user', %s)
            ON DUPLICATE KEY UPDATE
                density = VALUES(density),
                source = VALUES(source),
                created_by = VALUES(created_by),
                note = VALUES(note)
            """,
            (row[0], density, effective_date, note),
        )
        log_operation(
            cur,
            "save_density",
            "product_density_history",
            f"{product_code}:{effective_date}",
            {"productCode": product_code, "density": str(density), "effectiveDate": effective_date},
        )
        conn.commit()
        return jsonify({"ok": True})
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


if __name__ == "__main__":
    logging.getLogger("werkzeug").disabled = True
    host = os.getenv("SALES_HOST", "0.0.0.0")
    port = int(os.getenv("SALES_PORT", "8900"))
    try:
        from waitress import serve

        serve(app, host=host, port=port)
    except ImportError:
        app.run(host=host, port=port, debug=False, use_reloader=False)
