from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import pymysql
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from db import get_connection
from text_utils import normalize_text


SOURCE_TAG = "customer_manual_review_20260814"
MANUAL_XLSX = Path("outputs/customer_standard_review/customer_standard_review_after_rules_20260814.xlsx")

TARGET_OVERRIDES = {
    normalize_text("CUSTOMER-T"): "CJSC CUSTOMER-T",
    normalize_text("CJSC Мега CUSTOMER-F"): "CJSC Мега CUSTOMER-F",
    normalize_text("CUSTOMER-N"): "CJSC CUSTOMER-N",
}
OVERRIDE_TARGET_NORMS = {normalize_text(value) for value in TARGET_OVERRIDES.values()}

INTERNAL_HISTORY_STATIONS = {
    normalize_text("АЗС CTY-33"): ("retail_station", "historical_CTY-33", "АЗС CTY-33"),
    normalize_text("АЗС CTY-34"): ("retail_station", "historical_CTY-33", "АЗС CTY-34"),
}

PUNCT_RE = re.compile(r"[\s\-_.·,，。/\\()（）№#\"'`«»“”‘’]+", re.UNICODE)


def json_default(value):
    if isinstance(value, Decimal):
        return float(value)
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def compact(value) -> str:
    return PUNCT_RE.sub("", normalize_text(value).replace("Ё", "Е").replace("C", "С"))


def is_sever_yug_oil(value) -> bool:
    text = normalize_text(value).replace("Ё", "Е").replace("C", "С")
    text_key = compact(text)
    return "СЕВЕР" in text_key and "ЮГ" in text_key and "ОЙЛ" in text_key


def is_baiel(value) -> bool:
    return "БАЙЭЛ" in normalize_text(value).replace("Ё", "Е")


def is_asman_oil(value) -> bool:
    text_key = compact(value)
    return "АСМАН" in text_key and "ОЙЛ" in text_key


RECURRING_RULES = [
    ("sever_yug_oil", 'CJSC "CUSTOMER-O"', is_sever_yug_oil),
    ("baiel", "CJSC НПЗ CUSTOMER-H-3", is_baiel),
    ("asman_oil", "CJSC CUSTOMER-P", is_asman_oil),
]


def placeholders(values) -> str:
    return ",".join(["%s"] * len(values))


def fetch_customer_maps(cur):
    cur.execute("SELECT id, standard_name, normalized_name, customer_type FROM customers")
    rows = [dict(row) for row in cur.fetchall()]
    return {int(row["id"]): row for row in rows}, {row["normalized_name"]: row for row in rows}


def choose_customer_type(rows, existing_type=None) -> str:
    if existing_type and existing_type != "unknown":
        return existing_type
    counts = Counter(row.get("customer_type") or "unknown" for row in rows)
    for customer_type, _count in counts.most_common():
        if customer_type in {"local_wholesale", "export_wholesale", "mixed"}:
            return customer_type
    return "local_wholesale"


def get_or_create_customer(cur, target_name: str, customer_type: str):
    normalized = normalize_text(target_name)
    cur.execute(
        "SELECT id, standard_name, normalized_name, customer_type FROM customers WHERE normalized_name = %s",
        (normalized,),
    )
    row = cur.fetchone()
    if row:
        customer_id = int(row["id"])
        cur.execute(
            """
            UPDATE customers
            SET standard_name = %s,
                customer_type = IF(customer_type = 'unknown', %s, customer_type)
            WHERE id = %s
            """,
            (target_name, customer_type, customer_id),
        )
        return customer_id, False, dict(row)

    cur.execute(
        "INSERT INTO customers (standard_name, normalized_name, customer_type) VALUES (%s, %s, %s)",
        (target_name, normalized, customer_type),
    )
    customer_id = int(cur.lastrowid)
    return customer_id, True, {"id": customer_id, "standard_name": target_name, "normalized_name": normalized, "customer_type": customer_type}


def upsert_entity_alias(cur, entity_type: str, entity_code: str, standard_name: str, alias_name: str):
    cur.execute(
        """
        INSERT INTO entity_aliases
            (entity_type, entity_code, standard_name, alias_name, normalized_alias, confidence, source)
        VALUES (%s, %s, %s, %s, %s, 100, %s)
        ON DUPLICATE KEY UPDATE
            entity_type = VALUES(entity_type),
            entity_code = VALUES(entity_code),
            standard_name = VALUES(standard_name),
            alias_name = VALUES(alias_name),
            confidence = VALUES(confidence),
            source = VALUES(source)
        """,
        (entity_type, entity_code, standard_name, alias_name, normalize_text(alias_name), SOURCE_TAG),
    )
    return cur.rowcount


def non_empty_aliases(*values):
    result = []
    seen = set()
    for value in values:
        if value is None:
            continue
        text = str(value).strip()
        if not text:
            continue
        key = normalize_text(text)
        if key in seen:
            continue
        seen.add(key)
        result.append(text)
    return result


def fetch_related_alias_names(cur, source_customer_ids, seed_names, predicate=None):
    alias_names = set(seed_names)
    source_codes = {str(int(item)) for item in source_customer_ids if item is not None}

    if source_codes:
        cur.execute(
            f"""
            SELECT alias_name, standard_name
            FROM entity_aliases
            WHERE entity_type = 'wholesale_customer'
              AND entity_code IN ({placeholders(source_codes)})
            """,
            list(source_codes),
        )
        for row in cur.fetchall():
            alias_names.add(row.get("alias_name"))
            alias_names.add(row.get("standard_name"))

    if predicate:
        cur.execute(
            """
            SELECT alias_name, standard_name
            FROM entity_aliases
            WHERE entity_type = 'wholesale_customer'
            """
        )
        for row in cur.fetchall():
            if predicate(row.get("alias_name")) or predicate(row.get("standard_name")):
                alias_names.add(row.get("alias_name"))
                alias_names.add(row.get("standard_name"))

    source_code_list = list(source_codes)
    conditions = []
    params = []
    if source_code_list:
        conditions.append(f"entity_code IN ({placeholders(source_code_list)})")
        params.extend(source_code_list)
    for name in list(alias_names):
        if name:
            conditions.append("raw_entity_name = %s")
            params.append(name)
            conditions.append("standard_entity_name = %s")
            params.append(name)
    if predicate:
        conditions.append("1=1")
    if conditions:
        cur.execute(
            f"""
            SELECT raw_entity_name, standard_entity_name
            FROM wholesale_detail_records
            WHERE {" OR ".join(conditions)}
            GROUP BY raw_entity_name, standard_entity_name
            """,
            params,
        )
        for row in cur.fetchall():
            for value in (row.get("raw_entity_name"), row.get("standard_entity_name")):
                if predicate is None or predicate(value):
                    alias_names.add(value)

    if source_code_list:
        cur.execute(
            f"""
            SELECT raw_customer_name
            FROM contracts
            WHERE customer_id IN ({placeholders(source_code_list)})
               OR raw_customer_name IN ({placeholders([name for name in alias_names if name])})
            GROUP BY raw_customer_name
            """,
            [*source_code_list, *[name for name in alias_names if name]],
        )
        for row in cur.fetchall():
            alias_names.add(row.get("raw_customer_name"))

    return non_empty_aliases(*alias_names)


def sync_contracts(cur, target_customer_id: int, source_customer_ids, alias_names):
    rowcount = 0
    source_ids = sorted({int(item) for item in source_customer_ids if int(item) != target_customer_id})
    if source_ids:
        cur.execute(
            f"UPDATE contracts SET customer_id = %s WHERE customer_id IN ({placeholders(source_ids)})",
            [target_customer_id, *source_ids],
        )
        rowcount += cur.rowcount
    if alias_names:
        cur.execute(
            f"UPDATE contracts SET customer_id = %s WHERE raw_customer_name IN ({placeholders(alias_names)})",
            [target_customer_id, *alias_names],
        )
        rowcount += cur.rowcount
    return rowcount


def sync_details(cur, entity_type: str, entity_code: str, standard_name: str, source_customer_ids, alias_names):
    conditions = []
    params = [entity_type, str(entity_code), standard_name]
    source_codes = sorted({str(int(item)) for item in source_customer_ids if str(item).isdigit() and str(int(item)) != str(entity_code)})
    if source_codes:
        conditions.append(f"entity_code IN ({placeholders(source_codes)})")
        params.extend(source_codes)
    if alias_names:
        conditions.append(f"raw_entity_name IN ({placeholders(alias_names)})")
        params.extend(alias_names)
        conditions.append(f"standard_entity_name IN ({placeholders(alias_names)})")
        params.extend(alias_names)
    if not conditions:
        return 0
    cur.execute(
        f"""
        UPDATE wholesale_detail_records
        SET entity_type = %s,
            entity_code = %s,
            standard_entity_name = %s
        WHERE {" OR ".join(conditions)}
        """,
        params,
    )
    return cur.rowcount


def source_ids_from_alias_hits(cur, target_name: str):
    normalized = normalize_text(target_name)
    cur.execute(
        """
        SELECT entity_code
        FROM entity_aliases
        WHERE entity_type = 'wholesale_customer'
          AND normalized_alias = %s
          AND entity_code REGEXP '^[0-9]+$'
        """,
        (normalized,),
    )
    return {int(row["entity_code"]) for row in cur.fetchall()}


def apply_wholesale_group(cur, target_name: str, row_items, extra_source_ids=None, predicate=None, slug="manual"):
    customers_by_id, customers_by_norm = fetch_customer_maps(cur)
    target_norm = normalize_text(target_name)
    existing = customers_by_norm.get(target_norm)
    source_ids = {int(item["customer_id"]) for item in row_items}
    source_ids.update(extra_source_ids or set())
    source_ids.update(source_ids_from_alias_hits(cur, target_name))
    source_rows = [customers_by_id[source_id] for source_id in source_ids if source_id in customers_by_id]
    chosen_type = choose_customer_type(source_rows, existing.get("customer_type") if existing else None)
    target_customer_id, created, existing_row = get_or_create_customer(cur, target_name, chosen_type)
    source_ids.add(target_customer_id)

    seed_names = [target_name]
    for item in row_items:
        seed_names.extend([item.get("db_standard_name"), item.get("manual_standard_name")])
    for source_id in source_ids:
        row = customers_by_id.get(source_id)
        if row:
            seed_names.extend([row.get("standard_name"), row.get("normalized_name")])
    alias_names = fetch_related_alias_names(cur, source_ids, seed_names, predicate=predicate)

    report = {
        "slug": slug,
        "targetName": target_name,
        "targetCustomerId": target_customer_id,
        "targetCreated": created,
        "sourceCustomerIds": sorted(source_ids),
        "aliasCount": len(alias_names),
        "entityAliasesUpserted": 0,
        "contractsUpdated": 0,
        "detailReferencesUpdated": 0,
    }
    for alias_name in alias_names:
        report["entityAliasesUpserted"] += upsert_entity_alias(
            cur,
            "wholesale_customer",
            str(target_customer_id),
            target_name,
            alias_name,
        )
    report["contractsUpdated"] = sync_contracts(cur, target_customer_id, source_ids, alias_names)
    report["detailReferencesUpdated"] = sync_details(
        cur,
        "wholesale_customer",
        str(target_customer_id),
        target_name,
        source_ids,
        alias_names,
    )
    return report


def apply_internal_group(cur, entity_type: str, entity_code: str, standard_name: str, row_items):
    source_ids = {int(item["customer_id"]) for item in row_items}
    seed_names = [standard_name]
    for item in row_items:
        seed_names.extend([item.get("db_standard_name"), item.get("manual_standard_name")])
    alias_names = fetch_related_alias_names(cur, source_ids, seed_names)
    report = {
        "slug": f"{entity_type}:{entity_code}",
        "targetName": standard_name,
        "targetCustomerId": entity_code,
        "targetCreated": False,
        "sourceCustomerIds": sorted(source_ids),
        "aliasCount": len(alias_names),
        "entityAliasesUpserted": 0,
        "contractsUpdated": 0,
        "detailReferencesUpdated": 0,
    }
    for alias_name in alias_names:
        report["entityAliasesUpserted"] += upsert_entity_alias(
            cur,
            entity_type,
            entity_code,
            standard_name,
            alias_name,
        )
    report["detailReferencesUpdated"] = sync_details(
        cur,
        entity_type,
        entity_code,
        standard_name,
        source_ids,
        alias_names,
    )
    return report


def load_manual_groups():
    workbook = load_workbook(MANUAL_XLSX, data_only=True)
    ws = workbook["标准客户名单"]
    groups = defaultdict(list)
    for row_idx in range(2, ws.max_row + 1):
        customer_id = ws.cell(row_idx, 1).value
        manual_name = ws.cell(row_idx, 2).value
        if not customer_id or not manual_name:
            continue
        manual_name = str(manual_name).strip()
        target_name = TARGET_OVERRIDES.get(normalize_text(manual_name), manual_name)
        db_standard_name = ws.cell(row_idx, 2).value
        groups[normalize_text(target_name)].append(
            {
                "customer_id": int(customer_id),
                "manual_standard_name": target_name,
                "db_standard_name": str(db_standard_name or ""),
            }
        )

    ws_dup = workbook["疑似重复"]
    for row_idx in range(2, ws_dup.max_row + 1):
        target_name = ws_dup.cell(row_idx, 3).value
        customer_id = ws_dup.cell(row_idx, 5).value
        current_name = ws_dup.cell(row_idx, 6).value
        if not target_name or not customer_id:
            continue
        target_name = TARGET_OVERRIDES.get(normalize_text(str(target_name).strip()), str(target_name).strip())
        groups[normalize_text(target_name)].append(
            {
                "customer_id": int(customer_id),
                "manual_standard_name": target_name,
                "db_standard_name": str(current_name or ""),
            }
        )
    return groups


def should_apply_group(cur, target_name: str, row_items):
    target_norm = normalize_text(target_name)
    if target_norm in INTERNAL_HISTORY_STATIONS:
        return True
    if len({int(item["customer_id"]) for item in row_items}) > 1:
        return True
    if target_norm in OVERRIDE_TARGET_NORMS:
        return True
    source_ids = [int(item["customer_id"]) for item in row_items]
    if source_ids:
        cur.execute(
            f"SELECT normalized_name FROM customers WHERE id IN ({placeholders(source_ids)})",
            source_ids,
        )
        if any((row.get("normalized_name") or "") != target_norm for row in cur.fetchall()):
            return True
    for item in row_items:
        if normalize_text(item.get("db_standard_name")) != target_norm:
            return True
    return bool(source_ids_from_alias_hits(cur, target_name))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--output-dir", default="D:/SalesVisualization/outputs/customer_standard_review")
    args = parser.parse_args()

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    mode = "apply" if args.apply else "dry_run"
    report_path = output_dir / f"customer_manual_review_{mode}_{stamp}.json"

    manual_groups = load_manual_groups()
    reports = []
    conn = get_connection()
    try:
        with conn.cursor(pymysql.cursors.DictCursor) as cur:
            for target_norm, row_items in sorted(manual_groups.items(), key=lambda item: item[0]):
                target_name = row_items[0]["manual_standard_name"]
                if not should_apply_group(cur, target_name, row_items):
                    continue
                if target_norm in INTERNAL_HISTORY_STATIONS:
                    entity_type, entity_code, standard_name = INTERNAL_HISTORY_STATIONS[target_norm]
                    reports.append(apply_internal_group(cur, entity_type, entity_code, standard_name, row_items))
                else:
                    reports.append(apply_wholesale_group(cur, target_name, row_items, slug="manual_b_c"))

            for slug, target_name, predicate in RECURRING_RULES:
                reports.append(apply_wholesale_group(cur, target_name, [], predicate=predicate, slug=slug))

        if args.apply:
            conn.commit()
        else:
            conn.rollback()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()

    payload = {
        "mode": mode,
        "source": SOURCE_TAG,
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "reportCount": len(reports),
        "reports": reports,
    }
    report_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=json_default), encoding="utf-8")
    print(json.dumps({
        "mode": mode,
        "reportPath": str(report_path),
        "reportCount": len(reports),
        "detailReferencesUpdated": sum(int(item.get("detailReferencesUpdated") or 0) for item in reports),
        "contractsUpdated": sum(int(item.get("contractsUpdated") or 0) for item in reports),
        "entityAliasesUpserted": sum(int(item.get("entityAliasesUpserted") or 0) for item in reports),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
