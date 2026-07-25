from __future__ import annotations

import argparse
import json
import re
import sys
from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import pymysql
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from db import get_connection
from scripts.export_standard_customers import (
    _attach_contract_and_detail_stats,
    _collect_alias_names,
    _fetch_aliases,
    _fetch_customers,
    _parse_dt,
)
from text_utils import normalize_text


STRICT_PUNCT_RE = re.compile(r"[\s\-_.·,，。/\\()（）№#]+", re.UNICODE)
LOOSE_KEEP_RE = re.compile(r"[^\w]+", re.UNICODE)
PREFIX_RE = re.compile(
    r"^(?:АЗС|AЗС|AZS|CJSC|ООО|ЗАО|ИП|ЧП|ТОО|АО|LLC|LLP|CJSC|ТД)\s+",
    re.IGNORECASE,
)


def _json_default(value):
    if isinstance(value, Decimal):
        return float(value)
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _strict_key(value: str) -> str:
    text = normalize_text(value).replace("Ё", "Е")
    return STRICT_PUNCT_RE.sub("", text)


def _loose_key(value: str) -> str:
    text = normalize_text(value).replace("Ё", "Е")
    text = PREFIX_RE.sub("", text).strip()
    return LOOSE_KEEP_RE.sub("", text)


def _row_for_customer(row, aliases):
    customer_id = int(row["id"])
    alias_names = _collect_alias_names(aliases.get(customer_id, []))
    return {
        "customer_id": customer_id,
        "standard_name": row["standard_name"],
        "normalized_name": row["normalized_name"],
        "customer_type": row["customer_type"],
        "contract_count": int(row["contract_count"] or 0),
        "detail_count": int(row["detail_count"] or 0),
        "detail_tons": round(float(row["detail_tons"] or 0), 3),
        "alias_count": len(alias_names),
        "aliases": "; ".join(alias_names),
        "created_at": _parse_dt(row.get("created_at")),
        "updated_at": _parse_dt(row.get("updated_at")),
    }


def _collect_groups(customers, aliases):
    groups = []
    seen_customer_sets = set()

    for group_type, key_func in (
        ("严格重复：大小写/空格/符号差异", _strict_key),
        ("可能重复：去除常见前缀后相同", _loose_key),
    ):
        buckets = defaultdict(list)
        for row in customers:
            key = key_func(row["standard_name"])
            if key:
                buckets[key].append(row)

        for key, rows in sorted(buckets.items(), key=lambda item: (-len(item[1]), item[0])):
            if len(rows) < 2:
                continue
            customer_set = tuple(sorted(int(row["id"]) for row in rows))
            if customer_set in seen_customer_sets:
                continue
            seen_customer_sets.add(customer_set)
            active_rows = [row for row in rows if int(row.get("contract_count") or 0) or int(row.get("detail_count") or 0)]
            inactive_rows = [row for row in rows if row not in active_rows]
            if active_rows and inactive_rows and group_type.startswith("严格"):
                suggestion = "建议保留有合同/装车数据的名称，无数据名称作为别名并入"
            elif len(active_rows) > 1:
                suggestion = "多条都有业务数据，需人工确认主标准名"
            else:
                suggestion = "疑似格式重复，需人工确认"
            groups.append(
                {
                    "group_type": group_type,
                    "match_key": key,
                    "count": len(rows),
                    "suggestion": suggestion,
                    "rows": [_row_for_customer(row, aliases) for row in rows],
                }
            )

    return groups


def _fetch_alias_conflicts(cur):
    cur.execute(
        """
        SELECT ea.id, ea.entity_code, ea.standard_name AS alias_standard_name,
               ea.alias_name, ea.normalized_alias, ea.source, c.standard_name AS customer_standard_name,
               c.customer_type
        FROM entity_aliases ea
        LEFT JOIN customers c
          ON ea.entity_code REGEXP '^[0-9]+$'
         AND CAST(ea.entity_code AS UNSIGNED) = c.id
        WHERE ea.entity_type = 'wholesale_customer'
        ORDER BY ea.normalized_alias, ea.id
        """
    )
    rows = [dict(row) for row in cur.fetchall()]

    conflicts = []
    alias_buckets = defaultdict(list)
    for row in rows:
        alias_buckets[row.get("normalized_alias") or normalize_text(row.get("alias_name"))].append(row)
        customer_name = row.get("customer_standard_name")
        alias_standard_name = row.get("alias_standard_name")
        if customer_name and alias_standard_name and normalize_text(customer_name) != normalize_text(alias_standard_name):
            conflicts.append({**row, "issue": "别名标准名与客户ID对应标准名不一致"})

    for normalized_alias, items in alias_buckets.items():
        customer_codes = {str(item.get("entity_code") or "") for item in items}
        standard_names = {normalize_text(item.get("alias_standard_name")) for item in items}
        if len(customer_codes) > 1 or len(standard_names) > 1:
            for item in items:
                conflicts.append({**item, "issue": f"同一别名存在多个指向：{normalized_alias}"})

    deduped = {}
    for item in conflicts:
        deduped[(item.get("id"), item.get("issue"))] = item
    return list(deduped.values())


def _fetch_merged_customer_alias_names(cur):
    cur.execute(
        """
        SELECT ea.normalized_alias, c.normalized_name AS target_normalized
        FROM entity_aliases ea
        LEFT JOIN customers c
          ON ea.entity_code REGEXP '^[0-9]+$'
         AND CAST(ea.entity_code AS UNSIGNED) = c.id
        WHERE ea.entity_type = 'wholesale_customer'
          AND ea.entity_code REGEXP '^[0-9]+$'
        """
    )
    merged = set()
    for row in cur.fetchall():
        normalized_alias = row.get("normalized_alias")
        target_normalized = row.get("target_normalized")
        if normalized_alias and target_normalized and normalized_alias != target_normalized:
            merged.add(normalized_alias)
    return merged


def _auto_size(ws):
    for column in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column)
        width = min(max(max_len + 2, 12), 70)
        ws.column_dimensions[get_column_letter(column[0].column)].width = width


def _write_sheet(ws, headers, rows):
    ws.append(headers)
    for item in rows:
        ws.append([item.get(header, "") for header in headers])
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    _auto_size(ws)


def _write_excel(path: Path, customers, aliases, groups, alias_conflicts):
    wb = Workbook()

    summary = wb.active
    summary.title = "标准客户名单"
    summary_rows = [_row_for_customer(row, aliases) for row in customers]
    _write_sheet(
        summary,
        [
            "customer_id",
            "standard_name",
            "normalized_name",
            "customer_type",
            "contract_count",
            "detail_count",
            "detail_tons",
            "alias_count",
            "aliases",
            "created_at",
            "updated_at",
        ],
        summary_rows,
    )

    duplicate_sheet = wb.create_sheet("疑似重复")
    duplicate_headers = [
        "group_no",
        "group_type",
        "match_key",
        "suggestion",
        "customer_id",
        "standard_name",
        "customer_type",
        "contract_count",
        "detail_count",
        "detail_tons",
        "alias_count",
        "aliases",
    ]
    duplicate_sheet.append(duplicate_headers)
    for idx, group in enumerate(groups, 1):
        for row in group["rows"]:
            duplicate_sheet.append(
                [
                    idx,
                    group["group_type"],
                    group["match_key"],
                    group["suggestion"],
                    row["customer_id"],
                    row["standard_name"],
                    row["customer_type"],
                    row["contract_count"],
                    row["detail_count"],
                    row["detail_tons"],
                    row["alias_count"],
                    row["aliases"],
                ]
            )
    duplicate_sheet.freeze_panes = "A2"
    for cell in duplicate_sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FCE4D6")
    _auto_size(duplicate_sheet)

    aliases_sheet = wb.create_sheet("别名明细")
    alias_rows = []
    for row in customers:
        customer_id = int(row["id"])
        for item in aliases.get(customer_id, []):
            alias_rows.append(
                {
                    "customer_id": customer_id,
                    "standard_name": row["standard_name"],
                    "alias_id": int(item["aliasId"]),
                    "alias_name": item["aliasName"],
                    "confidence": item["confidence"],
                    "source": item["source"],
                    "created_at": item.get("createdAt"),
                }
            )
    _write_sheet(
        aliases_sheet,
        ["customer_id", "standard_name", "alias_id", "alias_name", "confidence", "source", "created_at"],
        alias_rows,
    )

    conflict_sheet = wb.create_sheet("别名异常")
    _write_sheet(
        conflict_sheet,
        [
            "issue",
            "id",
            "entity_code",
            "customer_standard_name",
            "alias_standard_name",
            "alias_name",
            "normalized_alias",
            "customer_type",
            "source",
        ],
        alias_conflicts,
    )

    wb.save(path)


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--type", default="classified", choices=["all", "classified", "local_wholesale", "export_wholesale", "mixed", "unknown"])
    parser.add_argument("--keyword", default="")
    parser.add_argument("--output-dir", default="D:/SalesVisualization/outputs/customer_standard_review")
    parser.add_argument("--prefix", default=None)
    return parser.parse_args()


def main():
    args = parse_args()
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    prefix = args.prefix or f"customer_standard_review_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    xlsx_path = output_dir / f"{prefix}.xlsx"
    json_path = output_dir / f"{prefix}.json"

    conn = get_connection()
    try:
        with conn.cursor(pymysql.cursors.DictCursor) as cur:
            customers = _fetch_customers(cur, args.type, args.keyword.strip())
            merged_alias_names = _fetch_merged_customer_alias_names(cur)
            customers = [
                row for row in customers
                if normalize_text(row.get("standard_name")) not in merged_alias_names
                and row.get("normalized_name") not in merged_alias_names
            ]
            customers = _attach_contract_and_detail_stats(cur, customers)
            aliases = _fetch_aliases(cur, [int(row["id"]) for row in customers])
            groups = _collect_groups(customers, aliases)
            alias_conflicts = _fetch_alias_conflicts(cur)
    finally:
        conn.close()

    _write_excel(xlsx_path, customers, aliases, groups, alias_conflicts)
    payload = {
        "customer_count": len(customers),
        "duplicate_group_count": len(groups),
        "duplicate_row_count": sum(len(group["rows"]) for group in groups),
        "alias_conflict_count": len(alias_conflicts),
        "duplicate_groups": groups,
        "alias_conflicts": alias_conflicts,
        "xlsx_path": str(xlsx_path),
    }
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=_json_default), encoding="utf-8")
    print(json.dumps({k: payload[k] for k in ("customer_count", "duplicate_group_count", "duplicate_row_count", "alias_conflict_count", "xlsx_path")}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
