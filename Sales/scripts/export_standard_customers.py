from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from decimal import Decimal
import sys

import pymysql
from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from db import get_connection
from text_utils import normalize_text


REVIEW_TYPES = {
    "all",
    "classified",
    "local_wholesale",
    "export_wholesale",
    "mixed",
    "unknown",
}


def _normalize(value):
    if value is None:
        return ""
    if isinstance(value, (datetime,)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def _parse_dt(value):
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _build_filter(customer_type: str, keyword: str):
    clauses = []
    params = []
    if customer_type and customer_type in REVIEW_TYPES:
        if customer_type == "classified":
            clauses.append("c.customer_type IN ('local_wholesale', 'export_wholesale')")
        elif customer_type == "unknown":
            clauses.append("c.customer_type NOT IN ('local_wholesale', 'export_wholesale')")
        elif customer_type != "all":
            clauses.append("c.customer_type = %s")
            params.append(customer_type)
    if keyword:
        like = f"%{keyword}%"
        clauses.append("(c.standard_name LIKE %s OR c.normalized_name LIKE %s)")
        params.extend([like, like])
    where_clause = f"WHERE {' AND '.join(clauses)}" if clauses else ""
    return where_clause, params


def _fetch_customers(cur, customer_type: str, keyword: str):
    where, params = _build_filter(customer_type, keyword)
    cur.execute(
        f"""
        SELECT
            c.id,
            c.standard_name,
            c.normalized_name,
            c.customer_type,
            c.created_at,
            c.updated_at
        FROM customers c
        {where}
        ORDER BY
            CASE WHEN c.customer_type IN ('local_wholesale', 'export_wholesale') THEN 0 ELSE 1 END,
            c.updated_at DESC,
            c.created_at DESC,
            c.standard_name
        LIMIT 10000
        """,
        tuple(params),
    )
    rows = [dict(row) for row in cur.fetchall()]
    for row in rows:
        row["contract_count"] = 0
        row["detail_count"] = 0
        row["detail_tons"] = 0.0
    return rows


def _attach_contract_and_detail_stats(cur, customers):
    if not customers:
        return customers

    id_by_record = {int(row["id"]): row for row in customers}
    ids = list(id_by_record.keys())
    placeholders = ",".join(["%s"] * len(ids))

    cur.execute(
        f"""
        SELECT customer_id, COUNT(*) AS contract_count
        FROM contracts
        WHERE customer_id IN ({placeholders})
        GROUP BY customer_id
        """,
        ids,
    )
    for row in cur.fetchall():
        customer_id = int(row["customer_id"])
        if customer_id in id_by_record:
            id_by_record[customer_id]["contract_count"] = int(row["contract_count"] or 0)

    name_to_id = {}
    for row in customers:
        if row.get("standard_name"):
            name_to_id[normalize_text(row["standard_name"])] = int(row["id"])
        if row.get("normalized_name"):
            name_to_id[row["normalized_name"]] = int(row["id"])

    cur.execute(
        """
        SELECT entity_code, standard_entity_name, raw_entity_name,
               COUNT(*) AS detail_count,
               COALESCE(SUM(weight_tons), 0) AS detail_tons
        FROM wholesale_detail_records
        GROUP BY entity_code, standard_entity_name, raw_entity_name
        """
    )
    for row in cur.fetchall():
        customer_id = None
        entity_code = str(row.get("entity_code") or "").strip()
        if entity_code.isdigit():
            entity_id = int(entity_code)
            if entity_id in id_by_record:
                customer_id = entity_id
        if customer_id is None:
            for name in (row.get("standard_entity_name"), row.get("raw_entity_name")):
                mapped = name_to_id.get(normalize_text(name))
                if mapped is not None:
                    customer_id = mapped
                    break
        if customer_id is None:
            continue
        target = id_by_record[customer_id]
        target["detail_count"] += int(row["detail_count"] or 0)
        target["detail_tons"] += float(row["detail_tons"] or 0)

    return customers


def _fetch_aliases(cur, customer_ids):
    if not customer_ids:
        return {}
    customer_id_set = set(customer_ids)
    placeholders = ",".join(["%s"] * len(customer_ids))
    cur.execute(
        f"""
        SELECT id, entity_type, entity_code, standard_name, alias_name,
               confidence, source, created_at
        FROM entity_aliases
        WHERE entity_type = 'wholesale_customer'
          AND entity_code REGEXP '^[0-9]+$'
          AND CAST(entity_code AS UNSIGNED) IN ({placeholders})
        ORDER BY entity_code, id
        """,
        tuple(str(i) for i in customer_ids),
    )
    aliases = defaultdict(list)
    for row in cur.fetchall():
        entity_code = row.get("entity_code") or ""
        if not str(entity_code).isdigit():
            continue
        alias_code = int(entity_code)
        if alias_code in customer_id_set:
            aliases[alias_code].append(
                {
                    "aliasId": row["id"],
                    "standardName": row["standard_name"],
                    "aliasName": row["alias_name"],
                    "confidence": row.get("confidence"),
                    "source": row.get("source"),
                    "createdAt": _parse_dt(row.get("created_at")),
                }
            )
    return aliases


def _collect_alias_names(alias_rows):
    alias_names = []
    seen = set()
    for item in alias_rows:
        alias_name = item.get("aliasName")
        if not alias_name:
            continue
        normalized_alias = normalize_text(alias_name)
        if normalized_alias in seen:
            continue
        seen.add(normalized_alias)
        alias_names.append(alias_name)
    return alias_names


def _auto_size_columns(ws):
    for column in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column)
        ws.column_dimensions[column[0].column_letter].width = min(max(max_len + 2, 12), 58)


def _write_excel(path: Path, customers, aliases):
    wb = Workbook()

    summary_sheet = wb.active
    summary_sheet.title = "标准客户名单"
    summary_headers = [
        "客户ID",
        "标准名称",
        "规范化名称",
        "客户类型",
        "合同数",
        "装车明细数",
        "装车吨数",
        "创建时间",
        "更新时间",
        "别名数",
        "别名列表",
    ]
    summary_sheet.append(summary_headers)
    for row in customers:
        alias_names = _collect_alias_names(aliases.get(int(row["id"]), []))
        summary_sheet.append([
            int(row["id"]),
            row["standard_name"],
            row["normalized_name"],
            row["customer_type"],
            int(row["contract_count"]),
            int(row["detail_count"]),
            round(float(row["detail_tons"] or 0), 3),
            _parse_dt(row.get("created_at")),
            _parse_dt(row.get("updated_at")),
            len(alias_names),
            "; ".join(alias_names),
        ])

    alias_sheet = wb.create_sheet("别名明细")
    alias_headers = [
        "客户ID",
        "标准名称",
        "别名ID",
        "别名",
        "置信度",
        "来源",
        "创建时间",
    ]
    alias_sheet.append(alias_headers)
    for row in customers:
        customer_id = int(row["id"])
        for item in aliases.get(customer_id, []):
            alias_sheet.append([
                customer_id,
                row["standard_name"],
                int(item["aliasId"]),
                item["aliasName"],
                item["confidence"],
                item["source"],
                item.get("createdAt"),
            ])
    for sheet in (summary_sheet, alias_sheet):
        sheet.freeze_panes = "A2"
        _auto_size_columns(sheet)
        for cell in sheet[1]:
            cell.font = Font(bold=True)

    wb.save(path)


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--type", default="all", choices=sorted(REVIEW_TYPES), help="客户类型筛选")
    parser.add_argument("--keyword", default="", help="按名称关键字模糊筛选")
    parser.add_argument(
        "--output-dir",
        default="D:/SalesVisualization/outputs/standard_customer_export",
        help="输出目录，默认 outputs/standard_customer_export",
    )
    parser.add_argument(
        "--prefix",
        default=None,
        help="输出文件前缀；默认按当前日期",
    )
    return parser.parse_args()


def main():
    args = parse_args()
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    now = datetime.now().strftime("%Y%m%d")
    prefix = args.prefix or f"standard_customers_{now}"
    xlsx_path = output_dir / f"{prefix}.xlsx"

    conn = get_connection()
    try:
        with conn.cursor(pymysql.cursors.DictCursor) as cur:
            customers = _fetch_customers(cur, args.type, args.keyword.strip())
            customers = _attach_contract_and_detail_stats(cur, customers)
            aliases = _fetch_aliases(cur, [int(row["id"]) for row in customers])
        conn.close()
    except Exception:
        conn.close()
        raise

    _write_excel(xlsx_path, customers, aliases)
    print(f"exported customers={len(customers)} to:")
    print(f"  Excel: {xlsx_path}")


if __name__ == "__main__":
    main()
