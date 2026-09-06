from datetime import date, datetime, timedelta
from io import BytesIO
import re
import unicodedata

import pymysql
from flask import current_app, jsonify, request, send_file
from flask_jwt_extended import get_jwt_identity, jwt_required
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from api_v2 import bp_v2
from calculations import recalculate_wholesale_batch
from db import get_connection
from import_helpers import (
    classify_wholesale_movement,
    detail_transport_mode,
    format_lpg_remark,
    get_or_create_customer,
    load_entity_lookup,
    load_product_lookup,
    load_sender_standard_lookup,
    lookup_entity,
    lookup_product,
    resolve_sender_standard_name,
)
from text_utils import clean_str, normalize_text, parse_date, parse_decimal
from wholesale_source_policy import (
    WHOLESALE_SOURCE_CUTOFF,
    detail_source_policy_clause,
    detail_source_policy_params,
)

from .authz import permission_required
from .data_scope import custom_resource_filter, custom_resource_values


DETAIL_DATA_SOURCES = {"business", "weighing", "uploaded"}
VALID_TRANSPORT_MODES = {"truck", "train"}
VALID_CLASSIFICATIONS = {"local_wholesale", "export_wholesale", "transfer", "processing", "procurement", "other", "unknown"}

CLASSIFICATION_LABELS = {
    "local_wholesale": "本地批发",
    "export_wholesale": "出口批发",
    "transfer": "外运",
    "processing": "代加工",
    "procurement": "采购",
    "other": "其他",
    "unknown": "待识别批发",
}
MATCH_STATUS_LABELS = {
    "matched": "已匹配",
    "unmatched": "未匹配",
    "needs_review": "待审核",
}
TRANSPORT_LABELS = {"truck": "汽车", "train": "火车"}
SOURCE_LABELS = {
    "manual": "手工录入",
    "wholesale_detail_single": "批发明细导入",
    "contract_excel": "批发明细导入",
    "weighing": "称重生成",
    "storagedistrict": "储运区",
    "kj_weighing": "储运区",
    "storage_district": "储运区",
    "lpg": "液化气罐区",
    "lpg_storage": "液化气罐区",
    "lpg_weighing": "液化气罐区",
    "储运区": "储运区",
    "液化气罐区": "液化气罐区",
}
example_RUSSIAN_NAME = "Example Petroleum CJSC"
example_CHINESE_NAME = "Example石油公司"


def _parse_date_arg(value, default_value):
    if not value:
        return default_value
    return datetime.strptime(value, "%Y-%m-%d").date()


def _normalize_classification(value):
    if value == "external_transport":
        return "transfer"
    if value == "toll_processing":
        return "processing"
    return value or None


def _normalize_data_source(value):
    value = (value or "business").strip()
    return value if value in DETAIL_DATA_SOURCES else "business"


def _source_clause(data_source):
    if data_source == "weighing":
        return "d.batch_id IS NULL", []
    if data_source == "uploaded":
        return "d.batch_id IS NOT NULL", []
    return detail_source_policy_clause("d"), list(detail_source_policy_params())


def _option_key(value):
    normalized = normalize_text(value)
    return re.sub(r"[\s\-_.]+", "", normalized)


ENTRY_CUSTOMER_TYPES = ("local_wholesale", "export_wholesale", "unknown")


def _normalize_entry_customer_type(value):
    if not value:
        return ""
    return value if value in ENTRY_CUSTOMER_TYPES else "unknown"


def _append_option(options, seen, value, customer_type=""):
    label = str(value or "").strip()
    if not label:
        return
    key = _option_key(label)
    if not key:
        return
    if key in seen:
        for item in options:
            if _option_key(item["value"]) == key:
                if not item.get("customerType") and customer_type:
                    item["customerType"] = customer_type
                break
        return
    seen.add(key)
    options.append({"value": label, "label": label, "customerType": customer_type or ""})


def _payload_text(payload, key, current):
    if key not in payload:
        return current
    return clean_str(payload.get(key))


def _payload_weight_kg(payload, current):
    if payload.get("weightTons") not in (None, ""):
        value = parse_decimal(payload.get("weightTons"))
        return value * 1000 if value is not None else None
    if "weightKg" in payload:
        return parse_decimal(payload.get("weightKg"))
    return current


def _empty_details(page, page_size, data_source):
    return {
        "items": [],
        "total": 0,
        "totalWeight": 0,
        "page": page,
        "page_size": page_size,
        "dataSource": data_source,
        "sourcePolicy": {
            "cutoffDate": WHOLESALE_SOURCE_CUTOFF.isoformat(),
            "beforeCutoff": "weighing",
            "fromCutoff": "uploaded",
        },
    }


def _detail_request_filters():
    date_from = request.args.get("date_from")
    date_to = request.args.get("date_to")
    classification = _normalize_classification(request.args.get("classification"))
    product_name = request.args.get("product_name")
    entity_name = request.args.get("entity_name")
    vehicle_no = request.args.get("vehicle_no")
    source = request.args.get("source")
    transport_mode = request.args.get("transport_mode")
    data_source = _normalize_data_source(request.args.get("data_source") or request.args.get("dataSource"))

    today = date.today()
    start = _parse_date_arg(date_from, today - timedelta(days=7))
    end = _parse_date_arg(date_to, today)
    if transport_mode not in VALID_TRANSPORT_MODES:
        transport_mode = None
    return {
        "start": start,
        "end": end,
        "classification": classification,
        "product_name": product_name,
        "entity_name": entity_name,
        "vehicle_no": vehicle_no,
        "source": source,
        "transport_mode": transport_mode,
        "data_source": data_source,
    }


def _detail_from_clause():
    return """
        FROM wholesale_detail_records d
        LEFT JOIN products p ON p.id = d.product_id
        LEFT JOIN import_batches b ON b.id = d.batch_id
    """


def _detail_select_sql(from_clause, where_sql, *, order_by=True, limit=False):
    suffix = "ORDER BY d.business_date DESC, d.id DESC" if order_by else ""
    if limit:
        suffix += "\nLIMIT %s OFFSET %s"
    return f"""
        SELECT d.id, d.batch_id, d.source_row_number, d.business_date,
               d.raw_product_name, p.display_name AS standard_product_name,
               COALESCE(NULLIF(p.display_name, ''), d.raw_product_name) AS display_product_name,
               d.vehicle_no, d.weigh_serial_no, d.raw_entity_name, d.sender_name,
               d.entity_type, d.entity_code, d.standard_entity_name,
               COALESCE(NULLIF(d.standard_entity_name, ''), d.raw_entity_name) AS display_entity_name,
               d.classification, d.weight_kg, d.weight_tons, d.match_status,
               d.weigh_source, d.transport_mode, d.remark, d.created_at,
               b.source_type AS batch_source_type,
               CASE
                 WHEN d.batch_id IS NULL THEN COALESCE(NULLIF(d.weigh_source, ''), 'weighing')
                 WHEN b.note LIKE 'Imported %% wholesale detail rows from single sheet.%%'
                   THEN 'wholesale_detail_single'
                 WHEN b.source_type = 'contract_excel'
                   THEN 'wholesale_detail_single'
                 ELSE COALESCE(b.source_type, 'manual')
               END AS display_source_type
        {from_clause}
        WHERE {where_sql}
        {suffix}
    """


def _build_detail_query(cur):
    filters = _detail_request_filters()
    where = ["d.business_date >= %s", "d.business_date <= %s"]
    params = [filters["start"], filters["end"]]
    source_filter, source_params = _source_clause(filters["data_source"])
    where.append(source_filter)
    params.extend(source_params)
    if filters["classification"]:
        where.append("d.classification = %s")
        params.append(filters["classification"])
    if filters["product_name"]:
        where.append("(d.raw_product_name LIKE %s OR p.display_name LIKE %s)")
        params.extend([f"%{filters['product_name']}%", f"%{filters['product_name']}%"])
    if filters["entity_name"]:
        like = f"%{filters['entity_name']}%"
        where.append(
            """
            (
                d.raw_entity_name LIKE %s
                OR d.standard_entity_name LIKE %s
                OR d.sender_name LIKE %s
                OR EXISTS (
                    SELECT 1
                    FROM entity_aliases ea
                    WHERE ea.entity_type IN ('wholesale_customer', 'retail_station', 'own_depot')
                      AND ea.standard_name LIKE %s
                      AND (
                          ea.alias_name IN (d.raw_entity_name, d.sender_name)
                          OR ea.standard_name IN (d.raw_entity_name, d.standard_entity_name, d.sender_name)
                      )
                )
            )
            """
        )
        params.extend([like, like, like, like])
    if filters["vehicle_no"]:
        where.append("d.vehicle_no LIKE %s")
        params.append(f"%{filters['vehicle_no']}%")
    if filters["source"]:
        where.append("(d.weigh_source = %s OR b.source_type = %s)")
        params.extend([filters["source"], filters["source"]])
    if filters["transport_mode"]:
        where.append("d.transport_mode = %s")
        params.append(filters["transport_mode"])

    scope_values = custom_resource_values(cur, int(get_jwt_identity()), "wholesale", "customer")
    allowed_entity_names = scope_values["labels"] if scope_values is not None else None
    if allowed_entity_names is not None:
        if not allowed_entity_names:
            return {**filters, "empty_scope": True}
        clauses = []
        for label in allowed_entity_names:
            clauses.append("(d.raw_entity_name LIKE %s OR d.standard_entity_name LIKE %s OR d.sender_name LIKE %s)")
            params.extend([f"%{label}%", f"%{label}%", f"%{label}%"])
        where.append(f"({' OR '.join(clauses)})")

    from_clause = _detail_from_clause()
    return {
        **filters,
        "empty_scope": False,
        "where_sql": " AND ".join(where),
        "params": params,
        "from_clause": from_clause,
    }


def load_detail_entity_display_lookup(cur):
    cur.execute(
        """
        SELECT entity_type, entity_code, standard_name, alias_name, normalized_alias
        FROM entity_aliases
        WHERE entity_type IN ('wholesale_customer', 'retail_station', 'own_depot')
          AND standard_name IS NOT NULL
          AND TRIM(standard_name) <> ''
        ORDER BY confidence DESC, id DESC
        """
    )
    lookup = {}
    for row in cur.fetchall():
        entry = {
            "entity_type": row.get("entity_type"),
            "entity_code": row.get("entity_code"),
            "standard_name": row.get("standard_name"),
        }
        for value in (row.get("normalized_alias"), normalize_text(row.get("alias_name")), normalize_text(row.get("standard_name"))):
            if value and value not in lookup:
                lookup[value] = entry
    return lookup


def resolve_detail_entity_display(lookup, *names):
    if not lookup:
        return None
    for name in names:
        key = normalize_text(name)
        if key and key in lookup:
            return lookup[key]
    return None


def _normalize_customer_display_name(name):
    return normalize_text(name or "").strip()


def _processing_customer_reference_maps(cur):
    def _get_row_value(row, index, *names):
        if row is None:
            return None
        if isinstance(row, dict):
            for name in names:
                if name in row:
                    return row.get(name)
            return None
        return row[index] if len(row) > index else None

    cur.execute(
        """
        SELECT id, standard_name, normalized_name
        FROM customers
        WHERE standard_name IS NOT NULL
          AND TRIM(standard_name) <> ''
        """
    )
    name_by_customer_id = {
        str(_get_row_value(row, 0, "id")): _get_row_value(row, 1, "standard_name")
        for row in cur.fetchall()
        if _get_row_value(row, 0, "id") is not None and _get_row_value(row, 1, "standard_name")
    }

    cur.execute(
        """
        SELECT normalized_name, standard_name
        FROM customers
        WHERE standard_name IS NOT NULL
          AND TRIM(standard_name) <> ''
          AND normalized_name IS NOT NULL
          AND TRIM(normalized_name) <> ''
        """
    )
    normalized_name_by_alias = {
        _normalize_customer_display_name(_get_row_value(row, 0, "normalized_name")): _get_row_value(row, 1, "standard_name")
        for row in cur.fetchall()
        if _get_row_value(row, 0, "normalized_name") and _get_row_value(row, 1, "standard_name")
    }

    cur.execute(
        """
        SELECT normalized_alias, alias_name, standard_name, entity_type, entity_code
        FROM entity_aliases
        WHERE entity_type = 'wholesale_customer'
          AND standard_name IS NOT NULL
          AND TRIM(standard_name) <> ''
        ORDER BY confidence DESC, id DESC
        """
    )
    name_by_alias = {}
    for row in cur.fetchall():
        for value in (
            _get_row_value(row, 0, "normalized_alias"),
            _get_row_value(row, 1, "alias_name"),
            _get_row_value(row, 2, "standard_name"),
        ):
            normalized = _normalize_customer_display_name(value)
            if not normalized:
                continue
            name_by_alias.setdefault(normalized, _get_row_value(row, 2, "standard_name"))
    name_by_alias.update({
        normalized: standard_name
        for normalized, standard_name in normalized_name_by_alias.items()
        if normalized and normalized not in name_by_alias
    })

    return name_by_customer_id, name_by_alias

def _resolve_export_customer_name(name_by_customer_id, name_by_alias, row):
    entity_type, entity_code, raw_entity_name, source_entity_name = (
        row.get("entity_type"),
        row.get("entity_code"),
        row.get("raw_entity_name"),
        row.get("source_entity_name"),
    )
    if entity_type == "wholesale_customer" and entity_code is not None:
        standard_name = name_by_customer_id.get(str(entity_code))
        if standard_name:
            return standard_name

    for candidate in (source_entity_name, raw_entity_name):
        normalized_candidate = _normalize_customer_display_name(candidate)
        if not normalized_candidate:
            continue
        mapped = name_by_alias.get(normalized_candidate)
        if mapped:
            return mapped

    return source_entity_name


def _row_to_detail_item(row, sender_lookup=None, entity_display_lookup=None):
    data_source = "uploaded" if row.get("batch_id") is not None else "weighing"
    weight_kg = float(row.get("weight_kg") or 0)
    display_source_type = row.get("display_source_type") or row.get("batch_source_type")
    source_value = row.get("weigh_source") or display_source_type or ("uploaded" if data_source == "uploaded" else "weighing")
    raw_weigh_source = str(row.get("weigh_source") or source_value or "").strip()
    source_key = raw_weigh_source.lower()
    remark = row.get("remark") or ""
    if raw_weigh_source == "液化气罐区" or source_key in {"lpg", "lpg_weighing", "lpg_storage"}:
        remark = format_lpg_remark(remark)
    standard_sender_name = row.get("standard_sender_name")
    if standard_sender_name is None and sender_lookup is not None:
        standard_sender_name = resolve_sender_standard_name(sender_lookup, row.get("sender_name"))
    sender_resolved = resolve_detail_entity_display(entity_display_lookup, standard_sender_name, row.get("sender_name"))
    if sender_resolved:
        standard_sender_name = sender_resolved.get("standard_name")

    entity_resolved = resolve_detail_entity_display(
        entity_display_lookup,
        row.get("standard_entity_name"),
        row.get("display_entity_name"),
        row.get("raw_entity_name"),
    )
    display_entity_name = row.get("display_entity_name")
    standard_entity_name = row.get("standard_entity_name")
    entity_type = row.get("entity_type")
    entity_code = row.get("entity_code")
    if entity_resolved:
        display_entity_name = entity_resolved.get("standard_name")
        standard_entity_name = entity_resolved.get("standard_name")
        entity_type = entity_resolved.get("entity_type") or entity_type
        entity_code = entity_resolved.get("entity_code") or entity_code

    return {
        "id": row.get("id"),
        "sourceId": row.get("id"),
        "dataSource": data_source,
        "importBatchId": row.get("batch_id"),
        "sourceRowNumber": row.get("source_row_number"),
        "businessDate": row["business_date"].isoformat() if row.get("business_date") else None,
        "productName": row.get("display_product_name"),
        "rawProductName": row.get("raw_product_name"),
        "standardProductName": row.get("standard_product_name"),
        "vehicleNo": row.get("vehicle_no") or "",
        "weighSerialNo": row.get("weigh_serial_no") or "",
        "senderName": row.get("sender_name") or "",
        "standardSenderName": standard_sender_name,
        "entityName": display_entity_name,
        "rawEntityName": row.get("raw_entity_name"),
        "standardEntityName": standard_entity_name,
        "entityCode": entity_code,
        "entityType": entity_type,
        "classification": row.get("classification") or "unknown",
        "weightKg": weight_kg,
        "weightTons": float(row.get("weight_tons") or (weight_kg / 1000)),
        "matchStatus": row.get("match_status") or "unmatched",
        "weighSource": source_value,
        "sourceType": display_source_type or source_value,
        "transportMode": row.get("transport_mode") or "truck",
        "remark": remark,
        "createdAt": row["created_at"].isoformat() if row.get("created_at") else None,
    }


def _label(mapping, value):
    key = str(value or "").strip()
    return mapping.get(key, key or "")


def _contains_chinese(text):
    return any("\u4e00" <= ch <= "\u9fff" for ch in text)


def _contains_cyrillic(text):
    return any("\u0400" <= ch <= "\u04ff" for ch in text)


def _cell_width(text):
    return sum(2 if unicodedata.east_asian_width(ch) in {"W", "F"} else 1 for ch in text)


def _is_example_party(raw_name, standard_name):
    text = f"{raw_name or ''} {standard_name or ''}".strip().lower()
    return any(keyword in text for keyword in ("Example", "ExampleCo", "ExampleCo公司", "ExampleLatin", "china petrol", "ExampleCo", "中油"))


def _contains_station_keyword(text):
    value = str(text or "").lower()
    return bool(re.search(r"азс|azs|加油站|油库", value))


def _extract_russian_name(text):
    candidates = [part.strip() for part in re.split(r"[|;/]", text)]
    for candidate in candidates:
        if _contains_cyrillic(candidate):
            russian = _extract_language_line(candidate, "cyrillic")
            if russian:
                return russian
    for candidate in candidates:
        if _contains_cyrillic(candidate):
            russian = _extract_language_line(candidate, "cyrillic")
            if russian:
                return russian
    return ""


def _normalize_whitespace(text):
    return " ".join(str(text).strip().split())


def _extract_language_line(text, script):
    value = _normalize_whitespace(text)
    if script == "chinese":
        if not _contains_chinese(value):
            return ""
        value = re.sub(r"[\u0400-\u04ff]+", " ", value)
    elif script == "cyrillic":
        if not _contains_cyrillic(value):
            return ""
        value = re.sub(r"[\u4e00-\u9fff]+", " ", value)
    else:
        return ""
    value = re.sub(r"[\|;/]+", " ", value)
    return _normalize_whitespace(value).strip(" -_/:：,，;|")


def _extract_bilingual_display(raw_name, standard_name):
    candidate_text = " | ".join(part for part in (str(raw_name or "").strip(), str(standard_name or "").strip()) if part)
    parts = [part.strip() for part in re.split(r"[|;/\n]", candidate_text) if part.strip()]
    chinese_parts = []
    russian_parts = []
    for part in parts:
        chinese = _extract_language_line(part, "chinese")
        russian = _extract_language_line(part, "cyrillic")
        if chinese:
            if chinese not in chinese_parts:
                chinese_parts.append(chinese)
        if russian:
            if russian not in russian_parts:
                russian_parts.append(russian)

    chinese = _normalize_whitespace(" ".join(chinese_parts))
    russian = _normalize_whitespace(" ".join(russian_parts))
    if chinese and russian and chinese != russian:
        return f"{chinese}\n{russian}"
    if chinese:
        return chinese
    return russian


def _normalize_multiline_text(value):
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    lines = [_normalize_whitespace(line) for line in text.split("\n") if line.strip()]
    if not lines:
        return ""

    expanded_lines = []
    for line in lines:
        if _contains_chinese(line) and _contains_cyrillic(line):
            bilingual = _extract_bilingual_display(line, "")
            if bilingual:
                for chunk in bilingual.split("\n"):
                    normalized_chunk = _normalize_whitespace(chunk)
                    if normalized_chunk:
                        expanded_lines.append(normalized_chunk)
                continue
        expanded_lines.append(line)

    deduped = []
    for line in expanded_lines:
        if line not in deduped:
            deduped.append(line)

    chinese_lines = [line for line in deduped if _contains_chinese(line)]
    russian_lines = [line for line in deduped if _contains_cyrillic(line)]
    if chinese_lines and russian_lines:
        return f"{chinese_lines[0]}\n{russian_lines[0]}"
    if len(deduped) > 2:
        return deduped[0] + "\n" + deduped[1]
    return "\n".join(deduped)


def _count_display_lines(value):
    text = _normalize_multiline_text(value)
    if not text:
        return 1
    return max(1, len(text.splitlines()))


def _preferred_party_name(raw_name, standard_name, *, prefer_internal_full=False, prefer_russian_station=False):
    raw = str(raw_name or "").strip()
    standard = str(standard_name or "").strip()
    names = [raw, standard]
    if prefer_internal_full:
        zh = example_CHINESE_NAME if _is_example_party(raw, standard) else next((name for name in names if _contains_chinese(name)), "")
        ru = example_RUSSIAN_NAME if _is_example_party(raw, standard) else _extract_russian_name(" | ".join(names))
        return f"{zh}\n{ru}" if zh and ru else (zh or ru)
    if prefer_russian_station:
        bilingual = _extract_bilingual_display(raw, standard)
        if bilingual:
            standard_russian = _extract_language_line(standard, "cyrillic")
            standard_chinese = _extract_language_line(standard, "chinese")
            if "\n" not in bilingual:
                if standard_russian:
                    return standard_russian
                if standard_chinese:
                    return standard_chinese
            return _normalize_multiline_text(bilingual)
        return _extract_russian_name(" | ".join(names)) or (standard or raw)
    bilingual = _extract_bilingual_display(raw, standard)
    if bilingual:
        # Keep bilingual split only when scripts differ; when both values are same script,
        # prefer the standard name to avoid duplicate alias+standard rendering.
        if "\n" not in bilingual and standard:
            standard_cyr = _extract_language_line(standard, "cyrillic")
            standard_zh = _extract_language_line(standard, "chinese")
            if standard_cyr:
                return standard_cyr
            if standard_zh:
                return standard_zh
        return _normalize_multiline_text(bilingual)
    return _extract_russian_name(" | ".join(names)) or (standard or raw)


def _station_line_with_prefix(value):
    text = _normalize_whitespace(str(value or ""))
    if not text:
        return ""
    if re.search(r"^\s*(АЗС|AЗС|AZS)\s", text, re.IGNORECASE):
        return text
    return f"АЗС {text}"


_DEPOT_CATALOG_CACHE = None


def _depot_catalog_cached():
    """自有油库（own_depot）slug -> 零售板块标准油库名称。"""
    global _DEPOT_CATALOG_CACHE
    if _DEPOT_CATALOG_CACHE is None:
        catalog = {}
        try:
            from reference_data import ENTITY_ALIASES

            for entry in ENTITY_ALIASES:
                if len(entry) < 4 or entry[0] != "own_depot":
                    continue
                code = entry[1]
                name = _normalize_whitespace(entry[2])
                if code and name and code not in catalog:
                    catalog[code] = name
        except Exception:
            catalog = {}
        _DEPOT_CATALOG_CACHE = catalog
    return _DEPOT_CATALOG_CACHE


def _match_transfer_entity_to_catalog(source):
    """按名称回退匹配零售板块加油站/油库，返回带 Азс 前缀的站名或油库名。

    覆盖实体未被解析（entityType/entityCode 缺失）但名称已能对应到零售板块
    主数据的情况，避免外运汇总表把这批发自有的站点/油库当成未知客户丢掉。
    """
    text = _normalize_whitespace(str(source or "")).lower()
    if not text:
        return ""
    try:
        from shared_utils import retail_station_catalog_cached

        station_catalog = retail_station_catalog_cached()
    except Exception:
        station_catalog = {}
    for info in (station_catalog or {}).values():
        for cand in (info.get("short"), info.get("name")):
            cand = _normalize_whitespace(str(cand or "")).lower()
            if not cand:
                continue
            if text == cand or text.endswith(cand):
                return _station_line_with_prefix(info.get("short") or info.get("name"))
    for name in _depot_catalog_cached().values():
        cand = _normalize_whitespace(str(name or "")).lower()
        if text == cand or text.endswith(cand):
            return _normalize_whitespace(name)
    return ""


def _build_transfer_summary_name(raw_name, standard_name, *, entity_type=None, entity_code=None):
    raw = _normalize_whitespace(str(raw_name or ""))
    standard = _normalize_whitespace(str(standard_name or ""))
    source = standard or raw
    if not source:
        return ""

    # 1) 已解析到零售板块实体：站点按 shop_key，油库按 slug。
    if entity_type == "retail_station" and entity_code not in (None, ""):
        station_key = int(entity_code) if str(entity_code).isdigit() else None
        if station_key is not None:
            try:
                from shared_utils import retail_station_catalog_cached

                station = retail_station_catalog_cached().get(station_key)
            except Exception:
                station = None
            if station:
                return _station_line_with_prefix(station.get("short") or station.get("name"))
    elif entity_type == "own_depot" and entity_code not in (None, ""):
        depot_name = _depot_catalog_cached().get(entity_code)
        if depot_name:
            return _normalize_whitespace(depot_name)

    # 2) 按名称回退匹配零售板块加油站/油库。
    matched = _match_transfer_entity_to_catalog(source)
    if matched:
        return matched

    # 3) 含站点关键词则补 Азс 前缀。
    if _contains_station_keyword(source):
        return _station_line_with_prefix(source)
    return ""


def _is_transfer_party(item):
    """收/发货方是否为零售板块自有站点/油库（外运目的地）。

    实体已解析为零售站点即算；未解析时按名称精确匹配站点目录兜底；
    自有油库仅站/库性质（目录名含 DemoDepot/нефтебаза/油库 等）算，
    内部用油/Example等非站库实体不算，避免误归外运。
    """
    entity_type = item.get("entityType")
    if entity_type == "retail_station":
        return True
    if entity_type == "own_depot":
        depot_name = _depot_catalog_cached().get(item.get("entityCode"))
        if depot_name and re.search(
            r"DemoDepot|нефтебаза|油库",
            _normalize_whitespace(depot_name).lower(),
        ):
            return True
        return False
    source = (
        item.get("standardEntityName")
        or item.get("entityName")
        or item.get("rawEntityName")
        or ""
    )
    text = _normalize_whitespace(str(source)).lower()
    if not text:
        return False
    try:
        from shared_utils import retail_station_catalog_cached

        station_catalog = retail_station_catalog_cached()
    except Exception:
        station_catalog = {}
    for info in (station_catalog or {}).values():
        for cand in (info.get("short"), info.get("name")):
            cand = _normalize_whitespace(str(cand or "")).lower()
            if cand and text == cand:
                return True
    return False


def _format_export_product(raw_product, standard_product):
    raw = str(raw_product or "").strip()
    standard = str(standard_product or "").strip()
    bilingual = _extract_bilingual_display(standard, raw)
    if bilingual and raw != standard:
        return _normalize_multiline_text(bilingual)
    if raw and standard and raw != standard:
        return _normalize_multiline_text(f"{raw}\n{standard}")
    return _normalize_multiline_text(standard or raw)


def _is_external_station_or_depot(item):
    if item.get("classification") == "transfer":
        return True
    if item.get("entityType") == "retail_station":
        return True
    return (
        _contains_station_keyword((item.get("rawEntityName") or item.get("entityName") or item.get("standardEntityName") or "").lower())
        or _contains_station_keyword((item.get("senderName") or item.get("standardSenderName") or "").lower())
        or _contains_station_keyword((item.get("rawEntityName") or item.get("entityName") or "").lower())
    )


def _compute_display_cell_width(value):
    text = "" if value is None else str(value)
    if not text:
        return 0
    lines = text.splitlines()
    return max(_cell_width(line) for line in lines) if lines else _cell_width(text)


def _export_column_width(column_index, content_width):
    if column_index in (4, 6, 7):
        return min(max(content_width * 1.2 + 8, 18), 100)
    return min(max(content_width + 2, 12), 80)


def _detail_export_filename(start, end, data_source):
    source_label = {
        "business": "业务口径",
        "uploaded": "上传手工",
        "weighing": "称重生成",
    }.get(data_source, data_source or "明细")
    return f"明细表_{source_label}_{start.isoformat()}_{end.isoformat()}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

def _resolve_summary_classification(classification):
    if classification in {"local_wholesale", "export_wholesale"}:
        return "wholesale"
    if classification in {"processing"}:
        return "processing"
    if classification in {"transfer"}:
        return "transfer"
    return None


def _summary_row_sort_key(group):
    """外运汇总表按加油站编号（shop_key）数值排序，其余按名称排序。"""
    if group == "transfer":
        def key(row):
            sn = row.get("stationNumber")
            sn_key = sn if isinstance(sn, int) else 10 ** 9
            return (sn_key, row["companyName"], row["productName"])
        return key
    return lambda row: (row["companyName"], row["productName"])


def _build_summary_rows(items):
    buckets = {
        "wholesale": {},
        "processing": {},
        "transfer": {},
    }
    for item in items:
        group = _resolve_summary_classification(item.get("classification"))
        # 收/发货方是零售板块自有站点/油库 → 一律按外运归口（分类可能未标 transfer）
        if _is_transfer_party(item):
            group = "transfer"
        if group is None:
            continue

        company_name = (
            item.get("standardEntityName")
            or item.get("entityName")
            or item.get("rawEntityName")
            or "未命名客户"
        ).strip()
        if group == "transfer":
            company_name = _build_transfer_summary_name(
                item.get("rawEntityName"),
                item.get("standardEntityName"),
                entity_type=item.get("entityType"),
                entity_code=item.get("entityCode"),
            )
            if not company_name:
                continue

        product_group_name = (
            item.get("standardProductName")
            or item.get("productName")
            or item.get("rawProductName")
            or "未命名油品"
        ).strip()
        product_name = _format_export_product(
            item.get("rawProductName"),
            item.get("standardProductName") or item.get("productName"),
        ).strip()
        if not company_name:
            company_name = "未命名客户"
        if not product_name:
            product_name = "未命名油品"
        if not product_group_name:
            product_group_name = product_name

        key = (company_name, product_group_name)
        bucket = buckets[group].setdefault(
            key,
            {
                "companyName": company_name,
                "productName": product_name,
                "totalTons": 0.0,
                "totalWeightKg": 0.0,
                "recordCount": 0,
                "stationNumber": None,
            },
        )
        bucket["totalTons"] += float(item.get("weightTons") or 0)
        bucket["totalWeightKg"] += float(item.get("weightKg") or 0)
        bucket["recordCount"] += 1
        # 外运汇总表按加油站编号排序/显示：仅零售站点记录 shop_key（兼容 int 与数值字符串）
        if group == "transfer" and bucket["stationNumber"] is None:
            ec = item.get("entityCode")
            if item.get("entityType") == "retail_station" and ec is not None and str(ec).isdigit():
                bucket["stationNumber"] = int(ec)

    return {
        key: [
            {
                "companyName": row["companyName"],
                "productName": row["productName"],
                "totalTons": float(row["totalTons"]),
                "totalWeightKg": float(row["totalWeightKg"]),
                "recordCount": int(row["recordCount"]),
                "stationNumber": row["stationNumber"],
            }
            for row in sorted(rows.values(), key=_summary_row_sort_key(key))
        ]
        for key, rows in buckets.items()
    }


def _merge_company_cells(ws, start_row, end_row, merge_column=1):
    if end_row < start_row + 1:
        return

    start = start_row
    prev_value = ws.cell(row=start_row, column=merge_column).value
    for row in range(start_row + 1, end_row + 1):
        current = ws.cell(row=row, column=merge_column).value
        if current != prev_value:
            if row - start > 1:
                ws.merge_cells(start_row=start, start_column=merge_column, end_row=row - 1, end_column=merge_column)
                ws.cell(row=start, column=merge_column).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            start = row
            prev_value = current

    if end_row - start + 1 > 1:
        ws.merge_cells(start_row=start, start_column=merge_column, end_row=end_row, end_column=merge_column)
        ws.cell(row=start, column=merge_column).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _format_station_number_label(station_number):
    """外运汇总表「编号」列显示：Азс - 01（两位数，数值不足前补零）。"""
    if station_number is None:
        return ""
    try:
        return "Азс - %02d" % int(station_number)
    except (TypeError, ValueError):
        return ""


def _build_summary_sheet(wb, title, rows, *, include_grand_total=True, first_column_name="公司", include_station_number=False):
    ws = wb.create_sheet(title)
    if include_station_number:
        headers = ["编号", first_column_name, "油品", "合计量(吨)", "合计量(kg)", "记录数"]
    else:
        headers = [first_column_name, "油品", "合计量(吨)", "合计量(kg)", "记录数"]
    ws.append(headers)
    name_col = 2 if include_station_number else 1

    total_tons = 0.0
    total_weight_kg = 0.0
    total_records = 0
    for row in rows:
        total_tons += float(row["totalTons"])
        total_weight_kg += float(row["totalWeightKg"])
        total_records += int(row["recordCount"])
        if include_station_number:
            ws.append([
                _format_station_number_label(row.get("stationNumber")),
                row["companyName"],
                row["productName"],
                row["totalTons"],
                row["totalWeightKg"],
                row["recordCount"],
            ])
        else:
            ws.append([
                row["companyName"],
                row["productName"],
                row["totalTons"],
                row["totalWeightKg"],
                row["recordCount"],
            ])
        data_row_index = ws.max_row
        multiline_lines = max(
            _count_display_lines(row["companyName"]),
            _count_display_lines(row["productName"]),
        )
        if multiline_lines > 1:
            ws.row_dimensions[data_row_index].height = 16.0 * multiline_lines

    if include_grand_total:
        if include_station_number:
            ws.append(["", "合计", "", total_tons, total_weight_kg, total_records])
        else:
            ws.append(["合计", "", total_tons, total_weight_kg, total_records])

    _merge_company_cells(ws, 2, ws.max_row - (1 if include_grand_total else 0), merge_column=name_col)

    fill = PatternFill("solid", fgColor="D9EAF7")
    bold_font = Font(name='Times New Roman', bold=True)
    center = Alignment(horizontal="center", vertical="center")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = bold_font
        cell.alignment = center

    for row in ws.iter_rows(min_row=2):
        for idx, cell in enumerate(row, start=1):
            if idx == name_col:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif idx == name_col + 1:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            else:
                cell.alignment = center

    num_cols = ("D", "E") if include_station_number else ("C", "D")
    for column in num_cols:
        for cell in ws[column][1:]:
            cell.number_format = '#,##0.000'

    for column_cells in ws.iter_cols(min_col=1, max_col=ws.max_column):
        max_len = 0
        for cell in column_cells:
            value_width = _compute_display_cell_width(cell.value)
            if value_width > max_len:
                max_len = value_width
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = _export_column_width(
            column_cells[0].column,
            max_len,
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def _build_detail_export_workbook(items, summary_rows_by_category):
    wb = Workbook()
    wb._fonts[0] = Font(name='Times New Roman', size=11)
    detail_ws = wb.active
    detail_ws.title = "明细表"

    summary_specs = [
        ("批发汇总表", "wholesale", "公司", False),
        ("代加工汇总表", "processing", "公司", False),
        ("外运汇总表", "transfer", "站点", True),
    ]
    for title, group_key, first_column_name, include_station_number in summary_specs:
        _build_summary_sheet(
            wb,
            title,
            summary_rows_by_category.get(group_key, []),
            first_column_name=first_column_name,
            include_station_number=include_station_number,
        )

    ws = detail_ws
    headers = [
        "业务日期", "运输方式", "磅单号", "油品",
        "车号", "发货单位", "收货单位",
        "重量kg", "重量吨", "备注",
    ]
    ws.append(headers)
    for item in items:
        raw_sender = item.get("senderName") or ""
        standard_sender = item.get("standardSenderName") or ""
        sender_is_example = _is_example_party(raw_sender, standard_sender)
        if sender_is_example:
            sender_display = _preferred_party_name(raw_sender, standard_sender, prefer_internal_full=True)
        elif _is_external_station_or_depot(item):
            sender_display = _preferred_party_name(
                raw_sender,
                standard_sender,
                prefer_russian_station=True,
            )
        else:
            sender_display = _preferred_party_name(raw_sender, standard_sender)

        raw_entity = item.get("rawEntityName") or ""
        standard_entity = item.get("standardEntityName") or item.get("entityName") or ""
        entity_is_example = _is_example_party(raw_entity, standard_entity)
        if entity_is_example:
            entity_display = _preferred_party_name(raw_entity, standard_entity, prefer_internal_full=True)
        elif _is_transfer_party(item):
            entity_display = _build_transfer_summary_name(
                raw_entity,
                standard_entity,
                entity_type=item.get("entityType"),
                entity_code=item.get("entityCode"),
            )
            if not entity_display:
                entity_display = _preferred_party_name(
                    raw_entity,
                    standard_entity,
                    prefer_russian_station=True,
                )
        elif _is_external_station_or_depot(item):
            entity_display = _preferred_party_name(
                raw_entity,
                standard_entity,
                prefer_russian_station=True,
            )
        else:
            entity_display = _preferred_party_name(raw_entity, standard_entity)

        product_display = _format_export_product(
            item.get("rawProductName"),
            item.get("standardProductName") or item.get("productName"),
        )
        ws.append([
            item.get("businessDate") or "",
            _label(TRANSPORT_LABELS, item.get("transportMode")),
            item.get("weighSerialNo") or "",
            product_display,
            item.get("vehicleNo") or "",
            sender_display,
            entity_display,
            item.get("weightKg") or 0,
            item.get("weightTons") or 0,
            item.get("remark") or "",
        ])
        data_row_index = ws.max_row
        multiline_lines = max(
            _count_display_lines(product_display),
            _count_display_lines(sender_display),
            _count_display_lines(entity_display),
        )
        if multiline_lines > 1:
            ws.row_dimensions[data_row_index].height = 16.0 * multiline_lines

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(name='Times New Roman', bold=True)
    center = Alignment(horizontal="center", vertical="center")
    wrap_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
    for row in ws.iter_rows(min_row=2):
        for index, cell in enumerate(row, start=1):
            if index in (4, 6, 7):
                cell.alignment = wrap_center
            else:
                cell.alignment = center
    for column in ("H", "I"):
        for cell in ws[column][1:]:
            cell.number_format = '#,##0.000'

    max_widths = []
    for column_cells in ws.iter_cols(min_col=1, max_col=ws.max_column):
        max_len = 0
        for cell in column_cells:
            value_width = _compute_display_cell_width(cell.value)
            if value_width > max_len:
                max_len = value_width
        max_widths.append(max_len)
    for index, width in enumerate(max_widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = _export_column_width(index, width)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return wb



@bp_v2.route("/wholesale/details/filter-options", methods=["GET"])
@jwt_required()
def wholesale_detail_filter_options():
    conn = get_connection()
    try:
        with conn.cursor(pymysql.cursors.DictCursor) as cur:
            product_options = []
            product_seen = set()
            cur.execute(
                """
                SELECT display_name
                FROM products
                WHERE is_active = 1
                ORDER BY display_name
                """
            )
            for row in cur.fetchall():
                _append_option(product_options, product_seen, row.get("display_name"))

            company_options = []
            company_seen = set()
            for value in (example_CHINESE_NAME, example_RUSSIAN_NAME):
                _append_option(company_options, company_seen, value)
            cur.execute(
                """
                SELECT c.standard_name, c.customer_type
                FROM customers c
                WHERE c.standard_name IS NOT NULL
                  AND TRIM(c.standard_name) <> ''
                  AND NOT EXISTS (
                      SELECT 1
                      FROM entity_aliases merged
                      LEFT JOIN customers target
                        ON merged.entity_code REGEXP '^[0-9]+$'
                       AND CAST(merged.entity_code AS UNSIGNED) = target.id
                      WHERE merged.entity_type = 'wholesale_customer'
                        AND merged.normalized_alias = c.normalized_name
                        AND target.id IS NOT NULL
                        AND target.normalized_name <> c.normalized_name
                  )
                ORDER BY
                    CASE WHEN c.customer_type IN ('local_wholesale', 'export_wholesale') THEN 0 ELSE 1 END,
                    c.standard_name
                """
            )
            for row in cur.fetchall():
                _append_option(
                    company_options,
                    company_seen,
                    row.get("standard_name"),
                    _normalize_entry_customer_type(row.get("customer_type")),
                )

            cur.execute(
                """
                SELECT DISTINCT
                    CASE
                        WHEN ea.entity_type = 'wholesale_customer' AND c.standard_name IS NOT NULL
                            THEN c.standard_name
                        ELSE ea.standard_name
                    END AS standard_name,
                    c.customer_type
                FROM entity_aliases ea
                LEFT JOIN customers c
                  ON ea.entity_type = 'wholesale_customer'
                 AND ea.entity_code REGEXP '^[0-9]+$'
                 AND CAST(ea.entity_code AS UNSIGNED) = c.id
                WHERE ea.entity_type IN ('wholesale_customer', 'retail_station', 'own_depot')
                  AND ea.standard_name IS NOT NULL
                  AND TRIM(ea.standard_name) <> ''
                  AND (
                      ea.entity_type IN ('retail_station', 'own_depot')
                      OR c.id IS NOT NULL
                  )
                ORDER BY standard_name
                """
            )
            for row in cur.fetchall():
                _append_option(
                    company_options,
                    company_seen,
                    row.get("standard_name"),
                    _normalize_entry_customer_type(row.get("customer_type")),
                )

            return jsonify({
                "code": 0,
                "message": "success",
                "data": {
                    "companies": company_options,
                    "products": product_options,
                },
            })
    finally:
        conn.close()


@bp_v2.route("/wholesale/details", methods=["GET"])
@jwt_required()
@permission_required("wholesale:view")
def wholesale_details():
    page = max(int(request.args.get("page", 1)), 1)
    page_size = min(max(int(request.args.get("page_size", 50)), 1), 2000)

    conn = get_connection()
    try:
        with conn.cursor(pymysql.cursors.DictCursor) as cur:
            query = _build_detail_query(cur)
            if query.get("empty_scope"):
                return jsonify({"code": 0, "message": "success", "data": _empty_details(page, page_size, query["data_source"])})
            cur.execute(
                f"""
                SELECT COUNT(*) AS total,
                       COALESCE(SUM(d.weight_kg), 0) AS total_weight
                {query["from_clause"]}
                WHERE {query["where_sql"]}
                """,
                query["params"],
            )
            total_row = cur.fetchone() or {}
            offset = (page - 1) * page_size
            cur.execute(
                _detail_select_sql(query["from_clause"], query["where_sql"], limit=True),
                [*query["params"], page_size, offset],
            )
            detail_rows = cur.fetchall()
            sender_lookup = load_sender_standard_lookup(cur)
            entity_display_lookup = load_detail_entity_display_lookup(cur)
            items = [_row_to_detail_item(row, sender_lookup, entity_display_lookup) for row in detail_rows]
            return jsonify({
                "code": 0,
                "message": "success",
                "data": {
                    "items": items,
                    "total": int(total_row.get("total") or 0),
                    "totalWeight": float(total_row.get("total_weight") or 0),
                    "page": page,
                    "page_size": page_size,
                    "dataSource": query["data_source"],
                    "sourcePolicy": {
                        "cutoffDate": WHOLESALE_SOURCE_CUTOFF.isoformat(),
                        "beforeCutoff": "weighing",
                        "fromCutoff": "uploaded",
                    },
                },
            })
    finally:
        conn.close()


@bp_v2.route("/wholesale/details/export", methods=["GET"])
@jwt_required()
@permission_required("wholesale:view")
def export_wholesale_details():
    query = None
    conn = get_connection()
    try:
        with conn.cursor(pymysql.cursors.DictCursor) as cur:
            query = _build_detail_query(cur)
            if query.get("empty_scope"):
                items = []
                summary_rows_by_category = {"wholesale": [], "processing": [], "transfer": []}
            else:
                cur.execute(
                    _detail_select_sql(query["from_clause"], query["where_sql"]),
                    query["params"],
                )
                rows = cur.fetchall()
                sender_lookup = load_sender_standard_lookup(cur)
                entity_display_lookup = load_detail_entity_display_lookup(cur)
                items = [_row_to_detail_item(row, sender_lookup, entity_display_lookup) for row in rows]
                summary_rows_by_category = _build_summary_rows(items)

        wb = _build_detail_export_workbook(items, summary_rows_by_category)
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        filename = _detail_export_filename(query["start"], query["end"], query["data_source"])
        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=filename,
        )
    except Exception as exc:
        current_app.logger.exception("Failed to generate wholesale detail export")
        return jsonify({"code": 5001, "message": "导出失败，请稍后重试", "data": None}), 500
    finally:
        conn.close()


@bp_v2.route("/wholesale/details/<int:detail_id>", methods=["PUT"])
@jwt_required()
@permission_required("wholesale_entry:update")
def update_wholesale_detail(detail_id):
    payload = request.get_json(silent=True) or {}
    conn = get_connection()
    try:
        with conn.cursor(pymysql.cursors.DictCursor) as cur:
            cur.execute(
                """
                SELECT id, batch_id, business_date, raw_product_name, vehicle_no,
                       weigh_serial_no, raw_entity_name, sender_name, classification,
                       weight_kg, transport_mode, remark
                FROM wholesale_detail_records
                WHERE id = %s
                """,
                (detail_id,),
            )
            current = cur.fetchone()
            if not current:
                return jsonify({"code": 4041, "message": "detail_not_found", "data": None}), 404
            if current.get("batch_id") is None:
                return jsonify({"code": 4221, "message": "weighing_detail_is_read_only", "data": None}), 422

            business_date = parse_date(payload.get("businessDate")) if "businessDate" in payload else current.get("business_date")
            raw_product = _payload_text(payload, "productName", current.get("raw_product_name"))
            raw_entity = _payload_text(payload, "entityName", current.get("raw_entity_name"))
            sender_name = _payload_text(payload, "senderName", current.get("sender_name"))
            vehicle_no = _payload_text(payload, "vehicleNo", current.get("vehicle_no"))
            weigh_serial_no = _payload_text(payload, "weighSerialNo", current.get("weigh_serial_no"))
            remark = _payload_text(payload, "remark", current.get("remark"))
            weight_kg = _payload_weight_kg(payload, current.get("weight_kg"))
            requested_transport_mode = payload.get("transportMode") if "transportMode" in payload else current.get("transport_mode")

            classification = payload.get("classification") if "classification" in payload else current.get("classification")
            if classification not in VALID_CLASSIFICATIONS:
                classification = "unknown"
            should_reclassify = (
                classification == "unknown"
                or "senderName" in payload
                or "entityName" in payload
            )

            if not business_date or not raw_product or not raw_entity or weight_kg is None or not (float(weight_kg) > 0):
                return jsonify({"code": 4222, "message": "日期、油品、客户/去向、重量必填", "data": None}), 422

            with conn.cursor() as tcur:
                product = lookup_product(load_product_lookup(tcur), raw_product)
                entity_lookup = load_entity_lookup(tcur)
                entity = lookup_entity(entity_lookup, raw_entity)
                if should_reclassify:
                    classification = classify_wholesale_movement(
                        tcur,
                        entity_lookup,
                        raw_entity,
                        sender_name=sender_name,
                        current="unknown",
                    )
                if classification == "processing" or requested_transport_mode not in VALID_TRANSPORT_MODES:
                    transport_mode = detail_transport_mode(classification, vehicle_no, weigh_serial_no, remark)
                else:
                    transport_mode = requested_transport_mode
                if classification in ("local_wholesale", "export_wholesale", "unknown"):
                    get_or_create_customer(tcur, raw_entity, classification if classification != "unknown" else "unknown")

                match_status = "matched" if classification in ("transfer", "processing", "procurement", "other") else "unmatched"
                tcur.execute(
                    """
                    UPDATE wholesale_detail_records
                    SET business_date = %s,
                        raw_product_name = %s,
                        product_id = %s,
                        vehicle_no = %s,
                        weigh_serial_no = %s,
                        raw_entity_name = %s,
                        sender_name = %s,
                        entity_type = %s,
                        entity_code = %s,
                        standard_entity_name = %s,
                        classification = %s,
                        weight_kg = %s,
                        remark = %s,
                        contract_id = NULL,
                        match_status = %s,
                        transport_mode = %s
                    WHERE id = %s
                      AND batch_id IS NOT NULL
                    """,
                    (
                        business_date,
                        raw_product,
                        product["id"] if product else None,
                        vehicle_no,
                        weigh_serial_no,
                        raw_entity,
                        sender_name,
                        entity["entity_type"] if entity else "unknown",
                        entity["entity_code"] if entity else None,
                        entity["standard_name"] if entity else None,
                        classification,
                        weight_kg,
                        remark,
                        match_status,
                        transport_mode,
                        detail_id,
                    ),
                )
                recalculation = recalculate_wholesale_batch(tcur, current["batch_id"])
        conn.commit()
        return jsonify({"code": 0, "message": "success", "data": {"id": detail_id, "recalculation": recalculation}})
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@bp_v2.route("/wholesale/export-detail", methods=["GET"])
@jwt_required()
@permission_required("wholesale:export")
def wholesale_export_detail():
    classification = _normalize_classification(request.args.get("classification", "export_wholesale"))
    date_from = request.args.get("date_from")
    date_to = request.args.get("date_to")

    conn = get_connection()
    try:
        with conn.cursor(pymysql.cursors.DictCursor) as cur:
            where = ["d.classification = %s", detail_source_policy_clause("d")]
            params = [classification, *detail_source_policy_params()]
            if date_from:
                where.append("d.business_date >= %s")
                params.append(date_from)
            if date_to:
                where.append("d.business_date <= %s")
                params.append(date_to)
            scope_clause, scope_params = custom_resource_filter(
                cur,
                int(get_jwt_identity()),
                "wholesale",
                {"customer": {"labels": ["d.sender_name", "d.raw_entity_name", "d.standard_entity_name"]}},
            )
            where.append(f"1=1{scope_clause}")
            params.extend(scope_params)

            cur.execute(
                f"""
                SELECT d.raw_product_name, d.standard_entity_name,
                       SUM(d.weight_tons) AS total_tons,
                       COUNT(*) AS record_count
                FROM wholesale_detail_records d
                WHERE {" AND ".join(where)}
                GROUP BY d.raw_product_name, d.standard_entity_name
                ORDER BY total_tons DESC
                """,
                params,
            )

            items = []
            for row in cur.fetchall():
                items.append({
                    "productName": row["raw_product_name"],
                    "entityName": row["standard_entity_name"],
                    "totalTons": float(row["total_tons"]),
                    "recordCount": row["record_count"],
                })
            return jsonify({"code": 0, "message": "success", "data": items})
    finally:
        conn.close()
