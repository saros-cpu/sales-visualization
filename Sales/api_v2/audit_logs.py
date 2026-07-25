from __future__ import annotations

import json
import time
from datetime import datetime
from io import BytesIO
from typing import Any

from flask import g, jsonify, request, send_file
from flask_jwt_extended import get_jwt_identity, jwt_required
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import pymysql

from api_v2 import bp_v2
from db import get_connection
from .authz import permission_required
from .permission_service import ensure_permission_schema


SENSITIVE_KEYWORDS = ("password", "token", "secret", "authorization", "credential")
PERMISSION_AUDIT_MODULES = ("users", "departments", "roles", "organization", "audit_logs")
PERMISSION_AUDIT_ACTIONS = ("permissions", "data_scopes", "reset_password", "manage")
AUDITED_GET_PREFIXES = (
    "/api/v2/dashboard",
    "/api/v2/retail",
    "/api/v2/wholesale",
    "/api/v2/procurement",
    "/api/v2/contracts",
    "/api/v2/materials",
    "/api/v2/transfer",
    "/api/v2/train-entry",
    "/api/v2/weighing",
    "/api/v2/aliases",
    "/api/v2/organization",
    "/api/v2/users",
    "/api/v2/map",
    "/api/v2/exchange",
    "/api/v2/assistant",
    "/api/v2/system",
    "/api/v2/retail/alert-insights",
)


def _ensure_column(cur, table: str, column: str, ddl: str) -> None:
    cur.execute(f"SHOW COLUMNS FROM {table} LIKE %s", (column,))
    if not cur.fetchone():
        cur.execute(f"ALTER TABLE {table} ADD COLUMN {ddl}")


def ensure_audit_schema(cur) -> None:
    cur.execute("""
        CREATE TABLE IF NOT EXISTS system_operation_logs (
            id BIGINT PRIMARY KEY AUTO_INCREMENT,
            user_id BIGINT NULL,
            username VARCHAR(128) NULL,
            display_name VARCHAR(128) NULL,
            module_key VARCHAR(64) NOT NULL,
            action_key VARCHAR(64) NOT NULL,
            method VARCHAR(16) NOT NULL,
            path VARCHAR(255) NOT NULL,
            endpoint VARCHAR(128) NULL,
            resource_type VARCHAR(64) NULL,
            resource_id VARCHAR(128) NULL,
            role_code VARCHAR(64) NULL,
            role_name VARCHAR(128) NULL,
            department_id BIGINT NULL,
            department_name VARCHAR(128) NULL,
            status_code INT NOT NULL,
            success BOOLEAN NOT NULL DEFAULT TRUE,
            request_query MEDIUMTEXT NULL,
            request_body MEDIUMTEXT NULL,
            duration_ms INT NULL,
            ip_address VARCHAR(64) NULL,
            user_agent VARCHAR(255) NULL,
            created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
            KEY idx_created_at (created_at),
            KEY idx_user_created (user_id, created_at),
            KEY idx_module_action (module_key, action_key, created_at),
            KEY idx_status_created (success, created_at)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
    """)
    _ensure_column(cur, "system_operation_logs", "role_code", "role_code VARCHAR(64) NULL")
    _ensure_column(cur, "system_operation_logs", "role_name", "role_name VARCHAR(128) NULL")
    _ensure_column(cur, "system_operation_logs", "department_id", "department_id BIGINT NULL")
    _ensure_column(cur, "system_operation_logs", "department_name", "department_name VARCHAR(128) NULL")
    _ensure_column(cur, "system_operation_logs", "request_query", "request_query MEDIUMTEXT NULL")
    _ensure_column(cur, "system_operation_logs", "duration_ms", "duration_ms INT NULL")


def _redact(value: Any) -> Any:
    if isinstance(value, dict):
        result = {}
        for key, item in value.items():
            lower_key = str(key).lower()
            if any(word in lower_key for word in SENSITIVE_KEYWORDS):
                result[key] = "***"
            else:
                result[key] = _redact(item)
        return result
    if isinstance(value, list):
        return [_redact(item) for item in value[:100]]
    return value


def _request_body() -> str | None:
    payload = request.get_json(silent=True)
    if payload is None:
        return None
    text = json.dumps(_redact(payload), ensure_ascii=False, default=str)
    return text[:12000]


def _request_query() -> str | None:
    if not request.args:
        return None
    payload = {key: request.args.getlist(key) for key in request.args.keys()}
    text = json.dumps(_redact(payload), ensure_ascii=False, default=str)
    return text[:12000]


def _path_segments() -> list[str]:
    prefix = "/api/v2/"
    path = request.path[len(prefix):] if request.path.startswith(prefix) else request.path.strip("/")
    return [segment for segment in path.split("/") if segment]


def _derive_module_action() -> tuple[str, str, str | None, str | None]:
    segments = _path_segments()
    first = segments[0] if segments else "unknown"
    module_key = first.replace("-", "_")
    resource_type = module_key
    resource_id = None

    if first == "organization" and len(segments) > 1:
        if segments[1] == "departments":
            module_key = "departments"
            resource_type = "department"
        elif segments[1] == "roles":
            module_key = "roles"
            resource_type = "role"
        if len(segments) > 2:
            resource_id = segments[2]
    elif first == "train-entry":
        module_key = "train_entry"
        resource_type = "train_entry"
        if len(segments) > 1 and segments[1] != "batch":
            resource_id = segments[1]
    elif first == "transfer":
        module_key = "transfer_loss"
        resource_type = "transfer_loss"
    elif first == "audit-logs":
        module_key = "audit_logs"
        resource_type = "audit_log"
    elif first == "assistant":
        module_key = "assistant"
        resource_type = "ai_agent"
    elif first == "system":
        module_key = "system_health"
        resource_type = "system_health"
    elif first == "retail" and len(segments) > 1 and segments[1] == "alert-insights":
        module_key = "retail"
        resource_type = "retail_alert"
    else:
        if len(segments) > 1:
            resource_id = segments[1]

    method_action = {
        "GET": "view",
        "POST": "create",
        "PUT": "update",
        "PATCH": "update",
        "DELETE": "delete",
    }.get(request.method, request.method.lower())

    if first == "auth":
        module_key = "auth"
        resource_type = "auth"
        method_action = segments[1].replace("-", "_") if len(segments) > 1 else "auth"
    elif first == "audit-logs" and len(segments) > 1:
        method_action = segments[1].replace("-", "_")
    elif first == "assistant":
        method_action = segments[1].replace("-", "_") if len(segments) > 1 else "query"
    elif first == "system":
        method_action = segments[1].replace("-", "_") if len(segments) > 1 else "view"
    elif first == "weighing":
        method_action = segments[1].replace("-", "_") if len(segments) > 1 else method_action
    elif len(segments) > 1 and segments[-1] in {"export"}:
        method_action = segments[-1].replace("-", "_")
    elif len(segments) > 2 and segments[-1] in {"approve", "status", "permissions", "data-scopes", "reset-password"}:
        method_action = segments[-1].replace("-", "_")

    return module_key, method_action, resource_type, resource_id


def _identity() -> int | None:
    try:
        identity = get_jwt_identity()
        return int(identity) if identity is not None else None
    except Exception:
        return None


def _user_snapshot(cur, user_id: int | None, body: str | None) -> dict[str, Any]:
    username = None
    display_name = None
    role_code = None
    role_name = None
    department_id = None
    department_name = None
    if user_id:
        cur.execute("SELECT username, display_name FROM system_users WHERE id = %s", (user_id,))
        row = cur.fetchone()
        if row:
            username = row["username"] if isinstance(row, dict) else row[0]
            display_name = row["display_name"] if isinstance(row, dict) else row[1]
        try:
            cur.execute("""
                SELECT r.role_code, r.role_name
                FROM system_user_roles ur
                JOIN system_roles r ON r.id = ur.role_id
                WHERE ur.user_id = %s AND r.is_active = TRUE
                ORDER BY ur.is_primary DESC, r.sort_order, r.id
                LIMIT 1
            """, (user_id,))
            role = cur.fetchone()
            if role:
                role_code = role["role_code"] if isinstance(role, dict) else role[0]
                role_name = role["role_name"] if isinstance(role, dict) else role[1]
            cur.execute("""
                SELECT d.id, d.dept_name
                FROM system_user_departments ud
                JOIN system_departments d ON d.id = ud.dept_id
                WHERE ud.user_id = %s AND d.is_active = TRUE
                ORDER BY ud.is_primary DESC, d.sort_order, d.id
                LIMIT 1
            """, (user_id,))
            department = cur.fetchone()
            if department:
                department_id = department["id"] if isinstance(department, dict) else department[0]
                department_name = department["dept_name"] if isinstance(department, dict) else department[1]
        except Exception:
            pass
    if not username and body:
        try:
            payload = json.loads(body)
            username = payload.get("username") if isinstance(payload, dict) else None
        except (TypeError, ValueError):
            username = None
    return {
        "username": username,
        "displayName": display_name,
        "roleCode": role_code,
        "roleName": role_name,
        "departmentId": department_id,
        "departmentName": department_name,
    }


def _should_record() -> bool:
    if not request.path.startswith("/api/v2/"):
        return False
    if request.path == "/api/v2/audit-logs/export":
        return True
    if request.path.startswith("/api/v2/audit-logs"):
        return False
    if request.method in {"OPTIONS", "HEAD"}:
        return False
    if request.method == "GET":
        return any(request.path.startswith(prefix) for prefix in AUDITED_GET_PREFIXES)
    return True


def _int_arg(name: str, default: int) -> int:
    try:
        return int(request.args.get(name) or default)
    except (TypeError, ValueError):
        return default


def _audit_filters() -> tuple[str, list[Any]]:
    where = []
    params: list[Any] = []

    preset = request.args.get("preset")
    module_key = request.args.get("moduleKey")
    action_key = request.args.get("actionKey")
    role_code = request.args.get("roleCode")
    department_name = request.args.get("departmentName")
    username = request.args.get("username")
    date_from = request.args.get("dateFrom")
    date_to = request.args.get("dateTo")
    success = request.args.get("success")

    if preset == "today":
        where.append("created_at >= CURDATE()")
    elif preset == "failed":
        where.append("success = 0")
    elif preset == "data_view":
        where.append("action_key = 'view'")
    elif preset == "permission":
        where.append(
            f"(module_key IN ({','.join(['%s'] * len(PERMISSION_AUDIT_MODULES))}) "
            f"OR action_key IN ({','.join(['%s'] * len(PERMISSION_AUDIT_ACTIONS))}))"
        )
        params.extend(PERMISSION_AUDIT_MODULES)
        params.extend(PERMISSION_AUDIT_ACTIONS)
    elif preset == "export":
        where.append("action_key = 'export'")

    if module_key:
        where.append("module_key = %s")
        params.append(module_key)
    if action_key:
        where.append("action_key = %s")
        params.append(action_key)
    if role_code:
        where.append("role_code = %s")
        params.append(role_code)
    if department_name:
        where.append("department_name LIKE %s")
        params.append(f"%{department_name}%")
    if username:
        where.append("(username LIKE %s OR display_name LIKE %s)")
        keyword = f"%{username}%"
        params.extend([keyword, keyword])
    if date_from:
        where.append("created_at >= %s")
        params.append(date_from)
    if date_to:
        where.append("created_at < DATE_ADD(%s, INTERVAL 1 DAY)")
        params.append(date_to)
    if success in {"0", "1"}:
        where.append("success = %s")
        params.append(int(success))

    where_sql = f"WHERE {' AND '.join(where)}" if where else ""
    return where_sql, params


def _select_audit_rows(cur, where_sql: str, params: list[Any], limit: int, offset: int = 0):
    cur.execute(
        f"""
        SELECT id, user_id, username, display_name, module_key, action_key, method, path,
               endpoint, resource_type, resource_id, role_code, role_name, department_id,
               department_name, status_code, success, request_query, request_body,
               duration_ms, ip_address, user_agent, created_at
        FROM system_operation_logs
        {where_sql}
        ORDER BY created_at DESC, id DESC
        LIMIT %s OFFSET %s
        """,
        tuple(params + [limit, offset]),
    )
    return cur.fetchall()


def _row_to_item(row) -> dict[str, Any]:
    return {
        "id": row["id"],
        "userId": row["user_id"],
        "username": row["username"],
        "displayName": row["display_name"],
        "moduleKey": row["module_key"],
        "actionKey": row["action_key"],
        "method": row["method"],
        "path": row["path"],
        "endpoint": row["endpoint"],
        "resourceType": row["resource_type"],
        "resourceId": row["resource_id"],
        "roleCode": row["role_code"],
        "roleName": row["role_name"],
        "departmentId": row["department_id"],
        "departmentName": row["department_name"],
        "statusCode": row["status_code"],
        "success": bool(row["success"]),
        "requestQuery": row["request_query"],
        "requestBody": row["request_body"],
        "durationMs": row["duration_ms"],
        "ipAddress": row["ip_address"],
        "userAgent": row["user_agent"],
        "createdAt": row["created_at"].isoformat(sep=" ") if isinstance(row["created_at"], datetime) else str(row["created_at"]),
    }


@bp_v2.before_request
def start_operation_timer():
    g.audit_started_at = time.perf_counter()


@bp_v2.after_request
def record_operation_log(response):
    if not _should_record():
        return response
    body = _request_body()
    query = _request_query()
    module_key, action_key, resource_type, resource_id = _derive_module_action()
    user_id = _identity()
    duration_ms = None
    started_at = getattr(g, "audit_started_at", None)
    if started_at:
        duration_ms = int((time.perf_counter() - started_at) * 1000)
    try:
        conn = get_connection(autocommit=True)
        try:
            with conn.cursor(pymysql.cursors.DictCursor) as cur:
                ensure_audit_schema(cur)
                user_snapshot = _user_snapshot(cur, user_id, body)
                cur.execute("""
                    INSERT INTO system_operation_logs
                        (user_id, username, display_name, module_key, action_key, method, path,
                         endpoint, resource_type, resource_id, role_code, role_name, department_id,
                         department_name, status_code, success, request_query, request_body,
                         duration_ms, ip_address, user_agent)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """, (
                    user_id,
                    user_snapshot["username"],
                    user_snapshot["displayName"],
                    module_key,
                    action_key,
                    request.method,
                    request.path,
                    request.endpoint,
                    resource_type,
                    resource_id,
                    user_snapshot["roleCode"],
                    user_snapshot["roleName"],
                    user_snapshot["departmentId"],
                    user_snapshot["departmentName"],
                    response.status_code,
                    1 if response.status_code < 400 else 0,
                    query,
                    body,
                    duration_ms,
                    request.headers.get("X-Forwarded-For", request.remote_addr),
                    (request.headers.get("User-Agent") or "")[:255],
                ))
        finally:
            conn.close()
    except Exception:
        pass
    return response


@bp_v2.route("/audit-logs", methods=["GET"])
@jwt_required()
@permission_required("audit_logs:view")
def list_audit_logs():
    page = max(_int_arg("page", 1), 1)
    page_size = min(max(_int_arg("pageSize", 50), 10), 200)
    offset = (page - 1) * page_size
    where_sql, params = _audit_filters()
    conn = get_connection()
    try:
        with conn.cursor(pymysql.cursors.DictCursor) as cur:
            ensure_permission_schema(cur)
            ensure_audit_schema(cur)
            conn.commit()
            cur.execute(f"SELECT COUNT(*) AS total FROM system_operation_logs {where_sql}", tuple(params))
            total = int(cur.fetchone()["total"])
            items = [_row_to_item(row) for row in _select_audit_rows(cur, where_sql, params, page_size, offset)]
            return jsonify({
                "code": 0,
                "message": "success",
                "data": {"items": items, "total": total, "page": page, "pageSize": page_size},
            })
    finally:
        conn.close()


@bp_v2.route("/audit-logs/export", methods=["GET"])
@jwt_required()
@permission_required("audit_logs:export")
def export_audit_logs():
    where_sql, params = _audit_filters()
    conn = get_connection()
    try:
        with conn.cursor(pymysql.cursors.DictCursor) as cur:
            ensure_permission_schema(cur)
            ensure_audit_schema(cur)
            conn.commit()
            rows = _select_audit_rows(cur, where_sql, params, 10000, 0)
    finally:
        conn.close()

    wb = Workbook()
    wb._fonts[0] = Font(name='Times New Roman', size=11)
    ws = wb.active
    ws.title = "审计日志"
    headers = [
        "时间", "操作人", "用户名", "部门", "角色", "模块", "动作", "方法", "路径",
        "资源", "状态", "状态码", "耗时(ms)", "IP", "查询条件", "请求内容",
    ]
    ws.append(headers)
    for row in rows:
        item = _row_to_item(row)
        ws.append([
            item["createdAt"],
            item.get("displayName") or item.get("username") or "",
            item.get("username") or "",
            item.get("departmentName") or "",
            item.get("roleName") or item.get("roleCode") or "",
            item["moduleKey"],
            item["actionKey"],
            item["method"],
            item["path"],
            item.get("resourceId") or "",
            "成功" if item["success"] else "失败",
            item["statusCode"],
            item.get("durationMs") or "",
            item.get("ipAddress") or "",
            (item.get("requestQuery") or "")[:30000],
            (item.get("requestBody") or "")[:30000],
        ])

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(name='Times New Roman', bold=True)
    wrap = Alignment(vertical="top", wrap_text=True)
    widths = [21, 18, 18, 18, 18, 16, 16, 10, 44, 18, 10, 10, 12, 18, 48, 70]
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"审计日志_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )
