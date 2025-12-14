from __future__ import annotations

from contextlib import contextmanager
import msvcrt
import os
import sys
import time
from datetime import date, datetime, timedelta
import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from api import app, calc_yoy_mom, get_report_month_inventory, query_month_summary  # noqa: E402
from business_time import business_period_condition_with_pruning, business_period_params, shift_join  # noqa: E402
from cache_store import set_cached  # noqa: E402
from constants import CLOSED_STATIONS, DEPOT_COORDS, SHOP_COORDS, SHOP_KEY_MAP  # noqa: E402
from db import get_connection  # noqa: E402
from endpoints.v1.retail.analytics.controller import (  # noqa: E402
    _apply_self_use_generator_rules,
    _map_resource_key_to_name as _self_use_resource_name,
    _self_use_report_resource_key,
    _self_use_number_for_record,
    _normalize_self_use_plate,
)
from export_excel import (  # noqa: E402
    EXPORT_DIR,
    _allow_fixed_density_fallback,
    _missing_density_fallback_density,
    _selling_mass_expr,
)
from routes.analysis_report import _inventory_contract_totals  # noqa: E402
from routes.analysis_report import _find_report_workbook_for_ym, _station_key_from_cell, _to_float  # noqa: E402
from routes.analytics import (  # noqa: E402
    _dashboard_yearly_from_monthly_snapshots,
    _resources_yearly_from_monthly_snapshots,
    add_resource_bucket,
    _attach_monthly_finance_metrics,
    format_resource_rows,
    report_amount_expr,
    report_bonus_amount_expr,
    report_card_amount_expr,
    report_card_volume_expr,
    report_cash_amount_expr,
    report_cash_volume_expr,
    report_resource_key_expr,
    report_volume_expr,
    station_status_for,
)
from routes.kpi import _yearly_kpi_from_monthly_snapshots  # noqa: E402

DAY_SNAPSHOT_DIR = ROOT / "cache_snapshots" / "daily"
REFRESH_LOCK_PATH = ROOT / "cache_snapshots" / ".refresh_current_snapshots.lock"


def _map_resource_key_to_name(rk, shop_key, base_price=None):
    """将ResourceKey映射为油品名称，使用与export_excel.py相同的映射逻辑"""
    # STN-55(SK=28)特殊油罐映射
    if shop_key == 28:
        from export_excel import SK28_RSK_MAP
        return SK28_RSK_MAP.get(rk, f'Other-{rk}')
    
    # STN-11(SK=6)映射
    if shop_key == 6:
        from export_excel import SK6_RSK_MAP
        return SK6_RSK_MAP.get(rk, f'Other-{rk}')
    
    # STN-07(SK=4)映射
    if shop_key == 4:
        from export_excel import SK4_RSK_MAP
        return SK4_RSK_MAP.get(rk, f'Other-{rk}')
    
    # 特殊站点映射
    if shop_key in (7, 15) and rk == 104:
        return 'АИ 80'
    if shop_key == 29:
        from export_excel import SK29_RSK_MAP
        return SK29_RSK_MAP.get(rk, f'Other-{rk}')
    
    # ResourceKey >= 100 的直接返回
    if rk >= 100:
        return f'Other-{rk}'
    
    # 基于价格的映射
    if base_price is not None:
        if base_price <= 45:
            return 'Д/Т'
        if base_price <= 65:
            return 'АИ 92'
        if base_price <= 78:
            return 'АИ 92'
        if base_price <= 83:
            return 'АИ 95'
    
    # 默认映射
    from export_excel import RESOURCE_NAMES
    return RESOURCE_NAMES.get(rk, f'Other-{rk}')


@contextmanager
def _single_refresh_lock():
    REFRESH_LOCK_PATH.parent.mkdir(parents=True, exist_ok=True)
    handle = REFRESH_LOCK_PATH.open("a+", encoding="utf-8")
    handle.seek(0, os.SEEK_END)
    if handle.tell() == 0:
        handle.write(" ")
        handle.flush()
    handle.seek(0)
    try:
        msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
    except OSError:
        handle.close()
        yield False
        return
    try:
        handle.seek(0)
        handle.write(f"{os.getpid()} {datetime.now().isoformat(timespec='seconds')}".ljust(64))
        handle.flush()
        yield True
    finally:
        handle.seek(0)
        msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
        handle.close()


def _month_end_exclusive(year: int, month: int) -> str:
    if month == 12:
        return f"{year + 1}-01-01"
    return f"{year}-{month + 1:02d}-01"


def _request(client, path: str, query_string: dict | None = None) -> bool:
    t0 = time.time()
    response = client.get(path, query_string=query_string or {})
    elapsed = time.time() - t0
    if response.status_code >= 400:
        print(f"failed {path}: {response.status_code} {response.get_data(as_text=True)[:300]}")
        return False
    print(f"ok {path}: {elapsed:.2f}s")
    return True


def _write_json_snapshot(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = path.with_suffix(path.suffix + ".tmp")
    content = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    for attempt in range(5):
        try:
            temp_path.write_text(content, encoding="utf-8")
            temp_path.replace(path)
            return
        except PermissionError:
            if attempt == 4:
                raise
            time.sleep(0.1 * (attempt + 1))


def _rebuild_yearly_snapshots(year: int, month: int) -> bool:
    written = []
    dashboard = _dashboard_yearly_from_monthly_snapshots(year)
    if dashboard:
        set_cached(f"dashboard_yearly:{year}", dashboard)
        written.append(f"dashboard_yearly:{year}")

    yearly_resources = _resources_yearly_from_monthly_snapshots(year)
    if yearly_resources:
        set_cached(f"resources_yearly:{year}:year", yearly_resources)
        written.append(f"resources_yearly:{year}:year")

    monthly_resources = _resources_yearly_from_monthly_snapshots(year, month)
    if monthly_resources:
        set_cached(f"resources_yearly:{year}:{month}", monthly_resources)
        written.append(f"resources_yearly:{year}:{month}")

    for include_inventory, label in ((False, "summary"), (True, "with_inventory")):
        kpi = _yearly_kpi_from_monthly_snapshots(year, include_inventory, get_report_month_inventory)
        if kpi:
            set_cached(f"kpi_yearly:{year}:{label}", kpi)
            written.append(f"kpi_yearly:{year}:{label}")

    print(f"ok yearly snapshot rebuild {year}: {', '.join(written) or 'no monthly data'}")
    return bool(written)


def _resource_key_from_name(value) -> int:
    name = str(value or "").strip().lower()
    if not name:
        return 0
    if "-35" in name:
        return 110
    if "спбт" in name:
        return 108
    if "95" in name:
        return 104
    if "92" in name:
        return 103
    if "80" in name:
        return 102
    if "д/т" in name or "диз" in name:
        return 106
    return 0


def _station_template(shop_key: int) -> dict:
    info = SHOP_KEY_MAP[shop_key]
    coords = SHOP_COORDS.get(shop_key, {"lng": 0, "lat": 0})
    return {
        "shopKey": shop_key,
        "name": info["name"],
        "short": info["short"],
        "region": info["region"],
        "lng": coords.get("lng", 0),
        "lat": coords.get("lat", 0),
        "txnCount": 0,
        "volume": 0.0,
        "tons": 0.0,
        "amount": 0.0,
        "cash": 0.0,
        "card": 0.0,
        "cashLiters": 0.0,
        "cardLiters": 0.0,
        "bonus": 0.0,
    }


def _round_station(item: dict) -> dict:
    amount = round(float(item.get("amount") or 0), 2)
    item["volume"] = round(float(item.get("volume") or 0), 2)
    item["tons"] = round(float(item.get("tons") or 0), 3)
    item["amount"] = amount
    item["cash"] = round(float(item.get("cash") or 0), 2)
    item["card"] = round(float(item.get("card") or 0), 2)
    item["cashLiters"] = round(float(item.get("cashLiters") or 0), 2)
    item["cardLiters"] = round(float(item.get("cardLiters") or 0), 2)
    item["bonus"] = round(float(item.get("bonus") or 0), 2)
    item["status"] = station_status_for(item["shopKey"], amount)
    item["closedReason"] = CLOSED_STATIONS.get(item["shopKey"], "")
    return item


def _add_aggregate(
    *,
    stations: dict,
    region_totals: dict,
    resource_totals: dict,
    resource_by_region: dict,
    daily_totals: dict,
    station_daily: dict,
    day_totals: dict | None = None,
    shop_key: int,
    current_date: str,
    resource_key: int,
    cash_volume: float,
    cash_amount: float,
    card_volume: float,
    card_amount: float,
    bonus_amount: float,
    total_volume: float,
    total_amount: float,
    total_mass: float | None = None,
    txn_count: int = 0,
) -> None:
    if shop_key not in stations:
        return
    station = stations[shop_key]
    mass = float(total_mass or 0)
    if not mass and total_volume:
        density = _missing_density_fallback_density(resource_key, current_date)
        if density:
            mass = total_volume * density
    tons = mass / 1000
    station["txnCount"] += int(txn_count or 0)
    station["volume"] += total_volume
    station["tons"] += tons
    station["amount"] += total_amount
    station["cash"] += cash_amount
    station["card"] += card_amount
    station["cashLiters"] += cash_volume
    station["cardLiters"] += card_volume
    station["bonus"] += bonus_amount

    daily_item = daily_totals.setdefault(current_date, {"volume": 0.0, "amount": 0.0})
    daily_item["volume"] += total_volume
    daily_item["amount"] += total_amount

    per_station = station_daily.setdefault(shop_key, {})
    per_station_day = per_station.setdefault(current_date, {"volume": 0.0, "amount": 0.0})
    per_station_day["volume"] += total_volume
    per_station_day["amount"] += total_amount

    region = station["region"]
    region_item = region_totals.setdefault(region, {
        "amount": 0.0,
        "volume": 0.0,
        "tons": 0.0,
        "cash": 0.0,
        "card": 0.0,
        "bonus": 0.0,
        "stationCount": 0,
        "activeCount": 0,
    })
    region_item["amount"] += total_amount
    region_item["volume"] += total_volume
    region_item["tons"] += tons
    region_item["cash"] += cash_amount
    region_item["card"] += card_amount
    region_item["bonus"] += bonus_amount

    add_resource_bucket(resource_totals, resource_key, total_volume, total_amount)
    region_resource = resource_by_region.setdefault(region, {})
    add_resource_bucket(region_resource, resource_key, total_volume, total_amount)

    if day_totals is not None:
        day_item = day_totals.setdefault(current_date, {
            "stations": {},
            "resources": {},
            "daily": {"date": current_date, "volume": 0.0, "amount": 0.0},
            "totalTxn": 0,
        })
        day_station = day_item["stations"].setdefault(shop_key, _station_template(shop_key))
        day_station["txnCount"] += int(txn_count or 0)
        day_station["volume"] += total_volume
        day_station["tons"] += tons
        day_station["amount"] += total_amount
        day_station["cash"] += cash_amount
        day_station["card"] += card_amount
        day_station["cashLiters"] += cash_volume
        day_station["cardLiters"] += card_volume
        day_station["bonus"] += bonus_amount
        day_item["daily"]["volume"] += total_volume
        day_item["daily"]["amount"] += total_amount
        day_item["totalTxn"] += int(txn_count or 0)
        add_resource_bucket(day_item["resources"], resource_key, total_volume, total_amount)


def _merge_db_missing_days(
    ym: str,
    stations: dict,
    region_totals: dict,
    resource_totals: dict,
    resource_by_region: dict,
    daily_totals: dict,
    station_daily: dict,
    day_totals: dict | None = None,
) -> None:
    year, month = [int(part) for part in ym.split("-")]
    loaded_days = [
        int(current_date[-2:])
        for current_date, values in daily_totals.items()
        if current_date.startswith(ym) and (values.get("volume") or values.get("amount"))
    ]
    start_day = max(loaded_days, default=0) + 1
    month_end = date.fromisoformat(_month_end_exclusive(year, month))
    if start_day > (month_end - timedelta(days=1)).day:
        return
    end_day = min(date.today(), month_end)
    start_date = date(year, month, start_day)
    if start_date >= end_day:
        return

    start = start_date.isoformat()
    end = end_day.isoformat()
    condition = business_period_condition_with_pruning("s", "sh")
    business_day = (
        "CASE "
        "WHEN sh.DatetimeShiftBegin IS NOT NULL "
        "AND sh.DatetimeShiftEnd IS NOT NULL "
        "AND DATE(sh.DatetimeShiftBegin) <> DATE(sh.DatetimeShiftEnd) "
        "AND TIME(sh.DatetimeShiftBegin) >= '18:00:00' "
        "THEN DATE(sh.DatetimeShiftEnd) "
        "WHEN sh.DatetimeShiftBegin IS NOT NULL "
        "THEN DATE(sh.DatetimeShiftBegin) "
        "ELSE DATE(s.TransactionDatetime) "
        "END"
    )
    resource_key = report_resource_key_expr("s")
    selling_mass = _selling_mass_expr(
        "s",
        use_ai80_missing_density_fallback=not _allow_fixed_density_fallback(start),
    )
    total_mass_expr = (
        "SUM(CASE "
        f"WHEN (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 3, 4)) "
        f" AND s.BasePaymentTypeKey IN (1, 2) THEN {selling_mass} ELSE 0 END) "
        f"+ SUM(CASE WHEN s.DerivedPaymentTypeKey = 3 THEN {selling_mass} ELSE 0 END)"
    )
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute(f"""
            SELECT
                s.ShopKey,
                {business_day} AS business_day,
                {resource_key} AS resource_key,
                COUNT(*) AS txn_count,
                {report_volume_expr("s")} AS total_volume,
                {report_amount_expr("s")} AS total_amount,
                {report_cash_volume_expr("s")} AS cash_volume,
                {report_cash_amount_expr("s")} AS cash_amount,
                {report_card_volume_expr("s")} AS card_volume,
                {report_card_amount_expr("s")} AS card_amount,
                {report_bonus_amount_expr("s")} AS bonus_amount,
                {total_mass_expr} AS total_mass
            FROM selling s
            {shift_join("s", "sh")}
            WHERE {condition}
              AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
            GROUP BY s.ShopKey, business_day, resource_key
        """, business_period_params(start, end))
        merged_rows = 0
        for row in cur.fetchall():
            (
                shop_key,
                current_date,
                rk,
                txn_count,
                total_volume,
                total_amount,
                cash_volume,
                cash_amount,
                card_volume,
                card_amount,
                bonus_amount,
                total_mass,
            ) = row
            _add_aggregate(
                stations=stations,
                region_totals=region_totals,
                resource_totals=resource_totals,
                resource_by_region=resource_by_region,
                daily_totals=daily_totals,
                station_daily=station_daily,
                day_totals=day_totals,
                shop_key=int(shop_key),
                current_date=str(current_date),
                resource_key=int(rk or 0),
                cash_volume=float(cash_volume or 0),
                cash_amount=float(cash_amount or 0),
                card_volume=float(card_volume or 0),
                card_amount=float(card_amount or 0),
                bonus_amount=float(bonus_amount or 0),
                total_volume=float(total_volume or 0),
                total_amount=float(total_amount or 0),
                total_mass=float(total_mass or 0),
                txn_count=int(txn_count or 0),
            )
            merged_rows += 1
        if merged_rows:
            print(f"ok db supplement {start}..{end}: {merged_rows} rows")
    finally:
        conn.close()


def _regions_from_station_rows(station_rows: list[dict]) -> dict:
    regions = {}
    for station in station_rows:
        region = station["region"]
        item = regions.setdefault(region, {
            "amount": 0.0,
            "volume": 0.0,
            "tons": 0.0,
            "cashSales": 0.0,
            "cardSales": 0.0,
            "cash": 0.0,
            "card": 0.0,
            "bonus": 0.0,
            "stationCount": 0,
            "activeCount": 0,
        })
        item["amount"] += station["amount"]
        item["volume"] += station["volume"]
        item["tons"] += station.get("tons", 0)
        item["cashSales"] += station["cash"]
        item["cardSales"] += station["card"]
        item["cash"] += station["cash"]
        item["card"] += station["card"]
        item["bonus"] += station["bonus"]
        item["stationCount"] += 1
        if station["amount"] > 0:
            item["activeCount"] += 1
    for item in regions.values():
        for key in ("amount", "volume", "cashSales", "cardSales", "cash", "card", "bonus"):
            item[key] = round(item[key], 2)
        item["tons"] = round(item["tons"], 3)
    return regions


def _write_day_snapshots(day_totals: dict) -> int:
    written = 0
    for current_date, values in sorted(day_totals.items()):
        station_map = {
            shop_key: dict(values["stations"].get(shop_key) or _station_template(shop_key))
            for shop_key in SHOP_KEY_MAP
        }
        station_rows = [_round_station(station_map[shop_key]) for shop_key in sorted(station_map)]
        by_resource = format_resource_rows([
            (key, item["volume"], item["amount"])
            for key, item in values["resources"].items()
        ])
        daily_item = {
            "date": current_date,
            "volume": round(values["daily"]["volume"], 2),
            "amount": round(values["daily"]["amount"], 2),
        }
        next_day = (date.fromisoformat(current_date) + timedelta(days=1)).isoformat()
        payload = {
            "date": current_date,
            "summary": {
                "totalTxn": int(values.get("totalTxn") or 0),
                "totalShops": len(SHOP_KEY_MAP),
                "totalLiters": daily_item["volume"],
                "totalSales": daily_item["amount"],
                "cashSales": round(sum(item["cash"] for item in station_rows), 2),
                "cardSales": round(sum(item["card"] for item in station_rows), 2),
                "cashLiters": round(sum(item["cashLiters"] for item in station_rows), 2),
                "cardLiters": round(sum(item["cardLiters"] for item in station_rows), 2),
                "activeShops": sum(1 for item in station_rows if item["amount"] > 0),
                "queryStart": current_date,
                "queryEnd": next_day,
            },
            "stations": station_rows,
            "regions": _regions_from_station_rows(station_rows),
            "byResource": by_resource,
            "daily": [daily_item],
        }
        _attach_monthly_finance_metrics(payload)
        year, month, day = current_date.split("-")
        day_dir = DAY_SNAPSHOT_DIR / year / month
        day_dir.mkdir(parents=True, exist_ok=True)
        _write_json_snapshot(day_dir / f"day_{int(day):02d}.json", payload)
        written += 1
    return written


def _report_snapshots_from_workbook(ym: str) -> bool:
    workbook_path = _find_report_workbook_for_ym(ym)
    if not workbook_path:
        print(f"skip workbook snapshots {ym}: no daily report workbook in {EXPORT_DIR}")
        return False

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        stations = {shop_key: _station_template(shop_key) for shop_key in SHOP_KEY_MAP}
        region_totals = {}
        resource_totals = {}
        resource_by_region = {}
        daily_totals = {}
        station_daily = {}
        day_totals = {}

        for sheet_name in workbook.sheetnames:
            if not sheet_name.endswith("日") or not sheet_name[:2].isdigit():
                continue
            day = int(sheet_name[:2])
            current_date = f"{ym}-{day:02d}"
            if date.fromisoformat(current_date) >= date.today():
                continue
            current_shop_key = None
            for row in workbook[sheet_name].iter_rows(min_row=5, values_only=True):
                station_key = _station_key_from_cell(row[0] if len(row) > 0 else None)
                if station_key is not None:
                    current_shop_key = station_key
                if current_shop_key is None or current_shop_key not in stations:
                    continue
                product_name = str(row[3] if len(row) > 3 and row[3] is not None else "").strip()
                if not product_name or "合计" in product_name or "итого" in product_name.lower():
                    continue

                resource_key = _resource_key_from_name(product_name)
                cash_volume = _to_float(row[10] if len(row) > 10 else None)
                bonus_volume = _to_float(row[12] if len(row) > 12 else None)
                bonus_amount = _to_float(row[13] if len(row) > 13 else None)
                cash_amount = _to_float(row[14] if len(row) > 14 else None)
                card_volume = _to_float(row[15] if len(row) > 15 else None)
                card_amount = _to_float(row[16] if len(row) > 16 else None)
                total_volume = _to_float(row[18] if len(row) > 18 else None)
                total_mass = _to_float(row[19] if len(row) > 19 else None) * 1000
                total_amount = _to_float(row[20] if len(row) > 20 else None)
                if not total_volume and (cash_volume or card_volume or bonus_volume):
                    total_volume = cash_volume + card_volume + bonus_volume
                if not total_amount and (cash_amount or card_amount or bonus_amount):
                    total_amount = cash_amount + card_amount + bonus_amount
                if not any((total_volume, total_amount, cash_amount, card_amount, bonus_amount)):
                    continue

                _add_aggregate(
                    stations=stations,
                    region_totals=region_totals,
                    resource_totals=resource_totals,
                    resource_by_region=resource_by_region,
                    daily_totals=daily_totals,
                    station_daily=station_daily,
                    day_totals=day_totals,
                    shop_key=current_shop_key,
                    current_date=current_date,
                    resource_key=resource_key,
                    cash_volume=cash_volume,
                    cash_amount=cash_amount,
                    card_volume=card_volume,
                    card_amount=card_amount,
                    bonus_amount=bonus_amount,
                    total_volume=total_volume,
                    total_amount=total_amount,
                    total_mass=total_mass,
                )

        _merge_db_missing_days(
            ym,
            stations,
            region_totals,
            resource_totals,
            resource_by_region,
            daily_totals,
            station_daily,
            day_totals,
        )

        for shop_key, station in stations.items():
            region = station["region"]
            region_item = region_totals.setdefault(region, {
                "amount": 0.0,
                "volume": 0.0,
                "tons": 0.0,
                "cash": 0.0,
                "card": 0.0,
                "bonus": 0.0,
                "stationCount": 0,
                "activeCount": 0,
            })
            region_item["stationCount"] += 1
            if station["amount"] > 0:
                region_item["activeCount"] += 1

        station_rows = [_round_station(stations[shop_key]) for shop_key in sorted(stations)]
        regions = {
            region: {
                **values,
                "amount": round(values["amount"], 2),
                "volume": round(values["volume"], 2),
                "tons": round(values.get("tons", 0), 3),
                "cashSales": round(values["cash"], 2),
                "cardSales": round(values["card"], 2),
                "cash": round(values["cash"], 2),
                "card": round(values["card"], 2),
                "bonus": round(values["bonus"], 2),
            }
            for region, values in region_totals.items()
        }
        daily_rows = [
            {"date": current_date, "volume": round(values["volume"], 2), "amount": round(values["amount"], 2)}
            for current_date, values in sorted(daily_totals.items())
        ]
        station_daily_rows = {
            shop_key: [
                {"date": current_date, "volume": round(values["volume"], 2), "amount": round(values["amount"], 2)}
                for current_date, values in sorted(days.items())
            ]
            for shop_key, days in station_daily.items()
        }
        existing_monthly = {}
        y, m = ym.split("-")
        monthly_path = ROOT / "cache_snapshots" / "monthly" / y / f"monthly_{int(m):02d}.json"
        if monthly_path.exists():
            try:
                existing_monthly = json.loads(monthly_path.read_text(encoding="utf-8-sig"))
            except (OSError, json.JSONDecodeError):
                existing_monthly = {}

        query_start = f"{ym}-01"
        if daily_rows:
            last_day = max(int(item["date"][-2:]) for item in daily_rows)
            year, month = [int(part) for part in ym.split("-")]
            if month == 12 and last_day >= 31:
                query_end = f"{year + 1}-01-01"
            elif last_day >= 31:
                query_end = f"{year}-{month + 1:02d}-01"
            else:
                query_end = f"{ym}-{last_day + 1:02d}"
        else:
            query_end = query_start
        monthly_summary = {
            "totalTxn": int((existing_monthly.get("summary") or {}).get("totalTxn") or 0),
            "totalShops": len(SHOP_KEY_MAP),
            "totalLiters": round(sum(item["volume"] for item in station_rows), 2),
            "totalSales": round(sum(item["amount"] for item in station_rows), 2),
            "cashSales": round(sum(item["cash"] for item in station_rows), 2),
            "cardSales": round(sum(item["card"] for item in station_rows), 2),
            "cashLiters": round(sum(item["cashLiters"] for item in station_rows), 2),
            "cardLiters": round(sum(item["cardLiters"] for item in station_rows), 2),
            "activeShops": sum(1 for item in station_rows if item["amount"] > 0),
            "queryStart": query_start,
            "queryEnd": query_end,
        }
        # 重新计算同比环比，避免沿用旧快照的错误数据
        y, m = int(ym.split('-')[0]), int(ym.split('-')[1])
        dates = calc_yoy_mom(y, m)
        conn_cmp = get_connection()
        try:
            cur_cmp = conn_cmp.cursor()
            prev_summary = query_month_summary(cur_cmp, dates['prev'])
            yoy_summary = query_month_summary(cur_cmp, dates['yoy'])
        finally:
            conn_cmp.close()

        def calc_change(current, previous):
            if not previous or previous <= 0:
                return None
            return round((current - previous) / previous * 100, 1)

        comparison = {
            'mom': {
                'sales': calc_change(monthly_summary['totalSales'], prev_summary['totalSales']),
                'volume': calc_change(monthly_summary['totalLiters'], prev_summary['totalLiters']),
                'cashSales': calc_change(monthly_summary['cashSales'], prev_summary['cashSales']),
                'cardSales': calc_change(monthly_summary['cardSales'], prev_summary['cardSales']),
                'activeShops': (monthly_summary['activeShops'] - prev_summary.get('activeShops', monthly_summary['activeShops'])) if prev_summary.get('activeShops') else None,
                'prevMonth': dates['prev'],
            },
            'yoy': {
                'sales': calc_change(monthly_summary['totalSales'], yoy_summary['totalSales']),
                'volume': calc_change(monthly_summary['totalLiters'], yoy_summary['totalLiters']),
                'cashSales': calc_change(monthly_summary['cashSales'], yoy_summary['cashSales']),
                'cardSales': calc_change(monthly_summary['cardSales'], yoy_summary['cardSales']),
                'activeShops': (monthly_summary['activeShops'] - yoy_summary.get('activeShops', monthly_summary['activeShops'])) if yoy_summary.get('activeShops') else None,
                'yoyMonth': dates['yoy'],
            },
            'prev': {'summary': prev_summary},
            'yoyRef': {'summary': yoy_summary},
        }

        monthly_payload = {
            "month": ym,
            "summary": monthly_summary,
            "comparison": comparison,
            "stations": station_rows,
            "regions": {
                region: {
                    "amount": values["amount"],
                    "volume": values["volume"],
                    "cashSales": values["cash"],
                    "cardSales": values["card"],
                    "stationCount": values["stationCount"],
                    "activeCount": values["activeCount"],
                    "momVolume": (existing_monthly.get("regions") or {}).get(region, {}).get("momVolume"),
                    "momSales": (existing_monthly.get("regions") or {}).get(region, {}).get("momSales"),
                    "yoyVolume": (existing_monthly.get("regions") or {}).get(region, {}).get("yoyVolume"),
                    "yoySales": (existing_monthly.get("regions") or {}).get(region, {}).get("yoySales"),
                }
                for region, values in regions.items()
            },
            "byResource": format_resource_rows([
                (key, values["volume"], values["amount"]) for key, values in resource_totals.items()
            ]),
            "daily": daily_rows,
        }
        _attach_monthly_finance_metrics(monthly_payload)

        station_payload = {
            "month": ym,
            "stations": station_rows,
            "depots": [
                {
                    "name": name,
                    "short": item["short"],
                    "region": item["region"],
                    "lng": item["lng"],
                    "lat": item["lat"],
                    "type": "depot",
                    "status": item.get("status", "normal"),
                    "closedReason": item.get("closedReason", ""),
                }
                for name, item in DEPOT_COORDS.items()
            ],
        }
        region_payload = {
            "month": ym,
            "regions": regions,
            "byResource": format_resource_rows([
                (key, values["volume"], values["amount"]) for key, values in resource_totals.items()
            ]),
            "byResourceByRegion": {
                region: format_resource_rows([
                    (key, values["volume"], values["amount"]) for key, values in resource_map.items()
                ])
                for region, resource_map in resource_by_region.items()
            },
        }
        daily_payload = {"month": ym, "daily": daily_rows, "stations": station_daily_rows}

        set_cached(f"monthly:{ym}", monthly_payload)
        set_cached(f"stations:{ym}", station_payload)
        set_cached(f"regions:{ym}", region_payload)
        set_cached(f"daily:{ym}", daily_payload)
        day_count = _write_day_snapshots(day_totals)
        print(f"ok workbook snapshots {ym}: {workbook_path.name}; day snapshots={day_count}")
        return True
    finally:
        workbook.close()


def _refresh_snapshots() -> int:
    today = date.today()
    completed_day = today - timedelta(days=1)
    ym = f"{completed_day.year}-{completed_day.month:02d}"
    client = app.test_client()
    ok = True

    ok = _request(client, f"/api/export/monthly/{ym}", {"refresh": "1"}) and ok
    ok = _report_snapshots_from_workbook(ym) and ok
    self_use_snapshot = _build_self_use_monthly_snapshot(ym)
    if self_use_snapshot:
        set_cached(f"self_use:{ym}", self_use_snapshot)
        print(f"ok self_use snapshot {ym}: {len(self_use_snapshot['records'])} records")
    ok = _rebuild_yearly_snapshots(completed_day.year, completed_day.month) and ok
    requests = [
        (f"/api/regions/{ym}", {}),
        ("/api/dashboard/yearly", {"year": str(completed_day.year)}),
        ("/api/kpi/yearly", {"year": str(completed_day.year), "inventory": "0"}),
        ("/api/resources/yearly", {"year": str(completed_day.year)}),
        (
            "/api/resources/yearly",
            {"year": str(completed_day.year), "month": str(completed_day.month)},
        ),
        (f"/api/retail/self-use/{ym}", {}),
    ]
    for path, query in requests:
        ok = _request(client, path, query) and ok

    start = f"{ym}-01"
    end = _month_end_exclusive(completed_day.year, completed_day.month)
    t0 = time.time()
    extras = _inventory_contract_totals(start, end)
    elapsed = time.time() - t0
    print(
        f"ok report_extras {ym}: inventory={extras.get('totalEndInventory', 0):.2f}, "
        f"contract={extras.get('totalContractVolume', 0):.2f}, elapsed={elapsed:.2f}s"
    )
    return 0 if ok else 1


def _build_self_use_monthly_snapshot(ym: str) -> dict | None:
    """构建自用油月度快照"""
    conn = get_connection()
    try:
        cur = conn.cursor()
        y, m = int(ym.split('-')[0]), int(ym.split('-')[1])
        ym_start = f'{ym}-01'
        ym_end = f'{y + 1}-01-01' if m == 12 else f'{y}-{m + 1:02d}-01'

        cur.execute("""
            SELECT s.TransactionKey, s.ShopKey, s.ResourceKey, s.Volume, s.ShopCost,
                   DATE(s.TransactionDatetime) as trx_date,
                   TIME(s.TransactionDatetime) as trx_time,
                   s.PersonKey, CAST(p.Name AS BINARY) as plate_number,
                   s.ResourceShopKey, s.BasePrice, p.COD_A, p.COD_OWN,
                   s.Mass, s.Density
            FROM selling s
            LEFT JOIN person p ON s.PersonKey = p.PersonKey
            WHERE s.DerivedPaymentTypeKey = 2
              AND s.TransactionDatetime >= %s AND s.TransactionDatetime < %s
            ORDER BY s.TransactionDatetime DESC
        """, (ym_start, ym_end))

        records = []
        total_volume = 0.0
        total_amount = 0.0
        total_mass = 0.0
        by_station = {}
        by_resource = {}
        by_region = {}
        daily = {}

        for r in cur.fetchall():
            sk = int(r[1]) if r[1] else 0
            rk = int(r[2]) if r[2] else 0
            vol = float(r[3] or 0)
            amt = float(r[4] or 0)
            date_str = str(r[5]) if r[5] else ''
            time_str = str(r[6]) if r[6] else ''
            pk = int(r[7]) if r[7] else 0
            rsk = int(r[9]) if r[9] else 0
            base_price = float(r[10]) if len(r) > 10 and r[10] else None
            group_key = int(r[11]) if len(r) > 11 and r[11] else 0
            self_use_number = _self_use_number_for_record(pk, group_key, int(r[12]) if len(r) > 12 and r[12] else 0)
            mass = float(r[13]) if len(r) > 13 and r[13] else 0.0
            density = float(r[14]) if len(r) > 14 and r[14] else 0.0
            price = amt / vol if vol else base_price
            report_rk = _self_use_report_resource_key(rk, sk, price, rsk)
            if not mass and not density and vol:
                density = _missing_density_fallback_density(report_rk, date_str)
            if not mass and vol and density:
                mass = vol * density

            total_volume += vol
            total_amount += amt
            total_mass += mass

            info = SHOP_KEY_MAP.get(sk, {'name': f'Station-{sk}', 'short': '', 'region': ''})
            region = info.get('region', '')
            product_name = _self_use_resource_name(report_rk, sk, price, rsk)
            raw_name = r[8]
            if raw_name:
                if isinstance(raw_name, bytes):
                    plate_number = raw_name.decode('cp1251', errors='replace').strip()
                else:
                    plate_number = str(raw_name).strip()
            else:
                plate_number = ''
            tk = int(r[0])
            plate_number = _normalize_self_use_plate(plate_number, pk, tk, group_key, self_use_number)
            plate_number, group_key, self_use_number = _apply_self_use_generator_rules(
                plate_number,
                info,
                transaction_key=tk,
                group_key=group_key,
                self_use_number=self_use_number,
            )

            records.append({
                'transactionKey': int(r[0]),
                'date': date_str,
                'time': time_str,
                'shopKey': sk,
                'stationName': info['name'],
                'stationShort': info.get('short', ''),
                'resourceKey': report_rk if report_rk >= 100 else rk,
                'resourceShopKey': rsk,
                'productName': product_name,
                'volume': round(vol, 3),
                'mass': round(mass, 3),
                'density': round(density or (mass / vol if vol else 0), 3),
                'amount': round(amt, 2),
                'plateNumber': plate_number,
                'personKey': pk,
                'personGroup': group_key,
                'selfUseNumber': self_use_number,
            })

            by_station.setdefault(sk, {'shopKey': sk, 'name': info['name'], 'volume': 0.0, 'amount': 0.0, 'txnCount': 0})
            by_station[sk]['volume'] += vol
            by_station[sk]['amount'] += amt
            by_station[sk]['txnCount'] += 1

            resource_bucket = report_rk if report_rk >= 100 else rk
            by_resource.setdefault(resource_bucket, {'resourceKey': resource_bucket, 'name': product_name, 'volume': 0.0, 'amount': 0.0})
            by_resource[resource_bucket]['volume'] += vol
            by_resource[resource_bucket]['amount'] += amt

            by_region.setdefault(region, {'volume': 0.0, 'amount': 0.0, 'activeCount': 0})
            by_region[region]['volume'] += vol
            by_region[region]['amount'] += amt

            daily.setdefault(date_str, {'volume': 0.0, 'amount': 0.0})
            daily[date_str]['volume'] += vol
            daily[date_str]['amount'] += amt

        station_rows = [
            {'shopKey': sk, 'name': item['name'], 'volume': round(item['volume'], 2), 'amount': round(item['amount'], 2), 'txnCount': item['txnCount']}
            for sk, item in sorted(by_station.items())
        ]

        resource_rows = [
            {'resourceKey': rk, 'name': item['name'], 'volume': round(item['volume'], 2), 'amount': round(item['amount'], 2)}
            for rk, item in sorted(by_resource.items())
        ]

        region_rows = {
            region: {'volume': round(item['volume'], 2), 'amount': round(item['amount'], 2), 'activeCount': item['activeCount']}
            for region, item in sorted(by_region.items())
        }

        daily_rows = [
            {'date': date_str, 'volume': round(item['volume'], 2), 'amount': round(item['amount'], 2)}
            for date_str, item in sorted(daily.items())
        ]

        dates = calc_yoy_mom(y, m)
        prev_summary = query_month_summary(cur, dates['prev'])
        yoy_summary = query_month_summary(cur, dates['yoy'])

        def calc_change(current, previous):
            if not previous or previous <= 0:
                return None
            return round((current - previous) / previous * 100, 1)

        comparison = {
            'mom': {
                'volume': calc_change(total_volume, prev_summary.get('totalVolume', 0)),
                'amount': calc_change(total_amount, prev_summary.get('totalAmount', 0)),
                'prevMonth': dates['prev']
            },
            'yoy': {
                'volume': calc_change(total_volume, yoy_summary.get('totalVolume', 0)),
                'amount': calc_change(total_amount, yoy_summary.get('totalAmount', 0)),
                'yoyMonth': dates['yoy']
            }
        }

        return {
            'month': ym,
            'records': records,
            'summary': {
                'totalVolume': round(total_volume, 2),
                'totalWeightKg': round(total_mass, 2),
                'totalAmount': round(total_amount, 2),
                'totalRecords': len(records),
                'activeStations': len(station_rows)
            },
            'comparison': comparison,
            'byStation': station_rows,
            'byResource': resource_rows,
            'byRegion': region_rows,
            'daily': daily_rows
        }
    except Exception as e:
        print(f"error building self_use snapshot for {ym}: {e}")
        return None
    finally:
        conn.close()


def main() -> int:
    with _single_refresh_lock() as acquired:
        if not acquired:
            print("skip retail snapshot refresh: another refresh is already running")
            return 0
        return _refresh_snapshots()


if __name__ == "__main__":
    raise SystemExit(main())
