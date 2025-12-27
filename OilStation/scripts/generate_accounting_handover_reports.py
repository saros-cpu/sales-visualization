# -*- coding: utf-8 -*-
from __future__ import annotations

import argparse
from collections import defaultdict
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
import json
from pathlib import Path
import re
import shutil
import subprocess
import sys
import zipfile
from xml.etree import ElementTree as ET

import openpyxl
import pymysql
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from business_time import shift_business_date_expr  # noqa: E402
from config import DB_CONFIG  # noqa: E402
from export_excel import _infer_resource_code_daily  # noqa: E402


REPORT_ROOT = Path(r"D:\每日销量日报")
DEFAULT_TEMPLATE_GLOB = r"20\d{6}"
COUNTER_TOLERANCE = 0.05
SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
XML_NS = "http://www.w3.org/XML/1998/namespace"
RESOURCE_ORDER = {
    102: 10,
    103: 20,
    104: 30,
    108: 40,
    106: 50,
    109: 51,
    110: 52,
}
SALES_FIELDS = ("cash_volume", "cash_amount", "talon_sber", "talon_zenit", "card_cnk", "vedomosti", "card_ts")
STATION_RESOURCE_SHOP_MAP = {
    21: {
        1: 103,
        2: 102,
        3: 102,
        5: 106,
    },
}

ET.register_namespace("", SHEET_NS)
ET.register_namespace("r", OFFICE_REL_NS)


@dataclass(frozen=True)
class ShiftRange:
    shop_key: int
    shift_keys: tuple[int, ...]
    begin: datetime
    end: datetime


@dataclass(frozen=True)
class ReportTemplate:
    path: Path
    shop_key: int
    sequence: int
    shift_begin: datetime | None = None
    shift_end: datetime | None = None


def parse_day(value: str | None) -> date:
    if not value:
        return date.today() - timedelta(days=1)
    return datetime.strptime(value, "%Y-%m-%d").date()


def day_folder_name(day: date) -> str:
    return day.strftime("%Y%m%d")


def report_file_name(shop_key: int, sequence: int = 0) -> str:
    suffix = "" if sequence == 0 else f"-{sequence}"
    return f"{shop_key:03d}{suffix}.xlsx"


def parse_report_file_name(name: str) -> tuple[int, int] | None:
    match = re.fullmatch(r"(\d{3})(?:-(\d+))?\.xlsx", name, re.IGNORECASE)
    if match:
        return int(match.group(1)), int(match.group(2) or 0)
    dated_match = re.fullmatch(r"20\d{6}(\d{3})(?:-(\d+))?\.xlsx", name, re.IGNORECASE)
    if dated_match:
        return int(dated_match.group(1)), int(dated_match.group(2) or 0)
    return None


def parse_shift_header(value) -> tuple[datetime, datetime] | None:
    matches = re.findall(r"(\d{2}\.\d{2}\.\d{4})\s+(\d{1,2}:\d{2}:\d{2})", str(value or ""))
    if len(matches) < 2:
        return None
    try:
        begin = datetime.strptime(" ".join(matches[0]), "%d.%m.%Y %H:%M:%S")
        end = datetime.strptime(" ".join(matches[1]), "%d.%m.%Y %H:%M:%S")
    except ValueError:
        return None
    return begin, end


def connect():
    return pymysql.connect(**DB_CONFIG, cursorclass=pymysql.cursors.DictCursor)


def find_template_dir(report_root: Path, target_day: date, explicit: str | None) -> Path:
    if explicit:
        template_dir = Path(explicit)
        if not template_dir.exists():
            raise FileNotFoundError(f"模板目录不存在: {template_dir}")
        return template_dir

    target_name = day_folder_name(target_day)
    candidates: list[Path] = []
    for item in report_root.iterdir():
        if not item.is_dir() or not re.fullmatch(DEFAULT_TEMPLATE_GLOB, item.name):
            continue
        if item.name >= target_name:
            continue
        if any(item.glob("*.xlsx")):
            candidates.append(item)
    if not candidates:
        raise FileNotFoundError(f"未找到 {target_name} 之前的日报模板目录")
    return sorted(candidates, key=lambda p: p.name)[-1]


def infer_resource(shop_key: int, resource_key, resource_shop_key=None, price=None) -> int:
    try:
        rk = int(resource_key or 0)
    except (TypeError, ValueError):
        rk = 0
    try:
        rsk = int(resource_shop_key) if resource_shop_key not in (None, "") else None
    except (TypeError, ValueError):
        rsk = None
    mapped = STATION_RESOURCE_SHOP_MAP.get(int(shop_key), {}).get(rsk)
    if mapped:
        return mapped
    try:
        base_price = float(price or 0)
    except (TypeError, ValueError):
        base_price = 0.0
    return int(_infer_resource_code_daily(rk, base_price, rsk=rsk, sk=int(shop_key)))


def resource_from_name(value) -> int | None:
    text = str(value or "").upper().replace(" ", "").replace("_", "-")
    if not text:
        return None
    if "ГАЗ" in text or "ГAЗ" in text:
        return 108
    if "95" in text:
        return 104
    if "92" in text:
        return 103
    if "80" in text:
        return 102
    if "ДТ-35" in text:
        return 110
    if "ДТ-15" in text:
        return 109
    if "ДТ" in text:
        return 106
    return None


def resource_sort_key(resource_key: int) -> int:
    return RESOURCE_ORDER.get(int(resource_key or 0), 999)


def as_float(value, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "").replace(",", ".")
    if not text:
        return default
    try:
        return float(text)
    except ValueError:
        return default


def money_decimal(value) -> Decimal:
    return Decimal(str(value or 0)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def round_float(value, digits: int = 2, adjust_float_tail: bool = False) -> float:
    decimal_value = Decimal(str(value or 0))
    if adjust_float_tail:
        if decimal_value > 0:
            decimal_value += Decimal("0.000000001")
        elif decimal_value < 0:
            decimal_value -= Decimal("0.000000001")
    return float(decimal_value.quantize(Decimal("1." + "0" * digits), rounding=ROUND_HALF_UP))


def value_or_blank(existing, value: float, digits: int = 2):
    if abs(float(value or 0)) < 0.0000001 and existing in (None, ""):
        return existing
    return round_float(value, digits)


def comma_number(value, digits: int = 2, trim: bool = False) -> str:
    quant = Decimal("1." + "0" * digits)
    text = f"{Decimal(str(value or 0)).quantize(quant, rounding=ROUND_HALF_UP):f}"
    if "." not in text and digits:
        text += "." + ("0" * digits)
    if trim and "." in text:
        text = text.rstrip("0").rstrip(".")
    return text.replace(".", ",")


UNITS_MALE = {
    1: "один",
    2: "два",
    3: "три",
    4: "четыре",
    5: "пять",
    6: "шесть",
    7: "семь",
    8: "восемь",
    9: "девять",
}
UNITS_FEMALE = UNITS_MALE.copy()
UNITS_FEMALE.update({1: "одна", 2: "две"})
TEENS = {
    10: "десять",
    11: "одиннадцать",
    12: "двенадцать",
    13: "тринадцать",
    14: "четырнадцать",
    15: "пятнадцать",
    16: "шестнадцать",
    17: "семнадцать",
    18: "восемнадцать",
    19: "девятнадцать",
}
TENS = {
    2: "двадцать",
    3: "тридцать",
    4: "сорок",
    5: "пятьдесят",
    6: "шестьдесят",
    7: "семьдесят",
    8: "восемьдесят",
    9: "девяносто",
}
HUNDREDS = {
    1: "сто",
    2: "двести",
    3: "триста",
    4: "четыреста",
    5: "пятьсот",
    6: "шестьсот",
    7: "семьсот",
    8: "восемьсот",
    9: "девятьсот",
}


def plural_ru(number: int, one: str, few: str, many: str) -> str:
    n = abs(number) % 100
    if 11 <= n <= 14:
        return many
    last = n % 10
    if last == 1:
        return one
    if 2 <= last <= 4:
        return few
    return many


def triad_words(number: int, female: bool = False) -> list[str]:
    words: list[str] = []
    hundreds = number // 100
    tens_units = number % 100
    tens = tens_units // 10
    units = tens_units % 10
    if hundreds:
        words.append(HUNDREDS[hundreds])
    if 10 <= tens_units <= 19:
        words.append(TEENS[tens_units])
        return words
    if tens:
        words.append(TENS[tens])
    if units:
        words.append((UNITS_FEMALE if female else UNITS_MALE)[units])
    return words


def integer_words(number: int) -> str:
    if number == 0:
        return "ноль"
    groups: list[tuple[int, bool, tuple[str, str, str] | None]] = [
        (number % 1000, False, None),
        ((number // 1000) % 1000, True, ("тысяча", "тысячи", "тысяч")),
        ((number // 1_000_000) % 1000, False, ("миллион", "миллиона", "миллионов")),
        ((number // 1_000_000_000) % 1000, False, ("миллиард", "миллиарда", "миллиардов")),
    ]
    words: list[str] = []
    for amount, female, names in reversed(groups):
        if not amount:
            continue
        words.extend(triad_words(amount, female=female))
        if names:
            words.append(plural_ru(amount, *names))
    return " ".join(words)


def amount_words_ru(value) -> str:
    amount = money_decimal(value)
    negative = amount < 0
    amount = abs(amount)
    rubles = int(amount)
    kopecks = int((amount - Decimal(rubles)) * 100)
    parts = [
        integer_words(rubles),
        plural_ru(rubles, "рубль", "рубля", "рублей"),
        f"{kopecks:02d}",
        plural_ru(kopecks, "копейка", "копейки", "копеек"),
    ]
    text = " ".join(parts)
    if negative:
        text = "минус " + text
    return text[:1].upper() + text[1:]


def dt_text(value: datetime) -> str:
    return f"{value.day:02d}.{value.month:02d}.{value.year} {value.hour}:{value.minute:02d}:{value.second:02d}"


def find_label_row(ws, label: str, column: int = 1) -> int:
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, column).value == label:
            return row
    raise ValueError(f"未找到表格标记: {label}")


def copy_cell_style(src, dst) -> None:
    if src.has_style:
        dst._style = copy(src._style)
    if src.number_format:
        dst.number_format = src.number_format
    if src.alignment:
        dst.alignment = copy(src.alignment)


def sheet_qn(name: str) -> str:
    return f"{{{SHEET_NS}}}{name}"


def rel_qn(name: str) -> str:
    return f"{{{PACKAGE_REL_NS}}}{name}"


def active_sheet_xml_name(xlsx_path: Path) -> str:
    with zipfile.ZipFile(xlsx_path) as archive:
        workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
        sheets = workbook_root.find(sheet_qn("sheets"))
        if sheets is None or not list(sheets):
            raise ValueError(f"Workbook has no sheets: {xlsx_path}")
        first_sheet = list(sheets)[0]
        rel_id = first_sheet.get(f"{{{OFFICE_REL_NS}}}id")
        if not rel_id:
            raise ValueError(f"Workbook sheet has no relationship id: {xlsx_path}")

        rels_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        target = None
        for rel in rels_root.findall(rel_qn("Relationship")):
            if rel.get("Id") == rel_id:
                target = rel.get("Target")
                break
        if not target:
            raise ValueError(f"Workbook sheet relationship not found: {xlsx_path}")
        if target.startswith("/"):
            return target.lstrip("/")
        return "xl/" + target.lstrip("/")


def worksheet_value_snapshot(ws) -> dict[tuple[int, int], object]:
    return {
        (row, col): ws.cell(row, col).value
        for row in range(1, ws.max_row + 1)
        for col in range(1, ws.max_column + 1)
    }


def cell_col_index(cell_ref: str) -> int:
    match = re.match(r"([A-Z]+)", cell_ref)
    if not match:
        return 0
    result = 0
    for char in match.group(1):
        result = result * 26 + ord(char) - ord("A") + 1
    return result


def ensure_sheet_row(sheet_data, row_number: int):
    rows = list(sheet_data.findall(sheet_qn("row")))
    for row in rows:
        if int(row.get("r", "0")) == row_number:
            return row
    row = ET.Element(sheet_qn("row"), {"r": str(row_number)})
    insert_at = len(rows)
    for index, existing in enumerate(rows):
        if int(existing.get("r", "0")) > row_number:
            insert_at = index
            break
    sheet_data.insert(insert_at, row)
    return row


def ensure_sheet_cell(row_element, row_number: int, column_number: int):
    coordinate = f"{get_column_letter(column_number)}{row_number}"
    cells = list(row_element.findall(sheet_qn("c")))
    for cell in cells:
        if cell.get("r") == coordinate:
            return cell
    cell = ET.Element(sheet_qn("c"), {"r": coordinate})
    insert_at = len(cells)
    for index, existing in enumerate(cells):
        if cell_col_index(existing.get("r", "")) > column_number:
            insert_at = index
            break
    row_element.insert(insert_at, cell)
    return cell


def remove_cell_value_nodes(cell) -> None:
    for child in list(cell):
        if child.tag in {sheet_qn("v"), sheet_qn("is"), sheet_qn("f")}:
            cell.remove(child)


def number_xml_text(value) -> str:
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, float):
        return f"{value:.15g}"
    return str(value)


def comparable_decimal(value) -> Decimal | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    if isinstance(value, str):
        text = value.strip().replace("\xa0", "").replace(" ", "").replace(",", ".")
        if re.fullmatch(r"[-+]?\d+(?:\.\d+)?(?:[Ee][-+]?\d+)?", text):
            return Decimal(text)
    return None


def worksheet_values_equivalent(old, new) -> bool:
    if old == new:
        return True
    old_decimal = comparable_decimal(old)
    new_decimal = comparable_decimal(new)
    if old_decimal is None or new_decimal is None:
        return False
    return abs(old_decimal - new_decimal) < Decimal("0.0000001")


def set_cell_xml_value(cell, value) -> None:
    remove_cell_value_nodes(cell)
    if value is None:
        cell.attrib.pop("t", None)
        return

    if isinstance(value, bool):
        cell.set("t", "b")
        value_element = ET.SubElement(cell, sheet_qn("v"))
        value_element.text = "1" if value else "0"
        return

    if isinstance(value, (int, float, Decimal)):
        cell.attrib.pop("t", None)
        value_element = ET.SubElement(cell, sheet_qn("v"))
        value_element.text = number_xml_text(value)
        return

    text = str(value)
    cell.set("t", "inlineStr")
    inline_string = ET.SubElement(cell, sheet_qn("is"))
    text_element = ET.SubElement(inline_string, sheet_qn("t"))
    if text != text.strip() or "\n" in text:
        text_element.set(f"{{{XML_NS}}}space", "preserve")
    text_element.text = text


def write_values_preserving_template_package(
    template_path: Path,
    output_path: Path,
    ws,
    original_values: dict[tuple[int, int], object],
) -> None:
    changed_values: dict[tuple[int, int], object] = {}
    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            value = ws.cell(row, col).value
            if not worksheet_values_equivalent(original_values.get((row, col)), value):
                changed_values[(row, col)] = value

    output_path.parent.mkdir(parents=True, exist_ok=True)
    sheet_xml_name = active_sheet_xml_name(template_path)
    with zipfile.ZipFile(template_path, "r") as source:
        sheet_root = ET.fromstring(source.read(sheet_xml_name))
        sheet_data = sheet_root.find(sheet_qn("sheetData"))
        if sheet_data is None:
            raise ValueError(f"Worksheet has no sheetData: {template_path}")

        for (row_number, column_number), value in changed_values.items():
            row_element = ensure_sheet_row(sheet_data, row_number)
            cell = ensure_sheet_cell(row_element, row_number, column_number)
            set_cell_xml_value(cell, value)

        sheet_xml = ET.tostring(sheet_root, encoding="utf-8", xml_declaration=True)
        with zipfile.ZipFile(output_path, "w") as target:
            for item in source.infolist():
                data = sheet_xml if item.filename == sheet_xml_name else source.read(item.filename)
                target.writestr(item, data)


def template_files(template_dir: Path) -> list[Path]:
    return sorted(template_dir.glob("*.xlsx"), key=lambda p: p.name)


def read_template_shift(path: Path) -> tuple[datetime, datetime] | None:
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            return parse_shift_header(workbook.active["K3"].value)
        finally:
            workbook.close()
    except Exception:
        return None


def report_templates(template_dir: Path) -> dict[int, list[ReportTemplate]]:
    result: dict[int, list[ReportTemplate]] = defaultdict(list)
    for path in template_files(template_dir):
        parsed = parse_report_file_name(path.name)
        if not parsed:
            continue
        shop_key, sequence = parsed
        shift = read_template_shift(path)
        shift_begin = shift[0] if shift else None
        shift_end = shift[1] if shift else None
        result[shop_key].append(ReportTemplate(path, shop_key, sequence, shift_begin, shift_end))
    return {shop_key: sorted(items, key=lambda item: (item.sequence, item.path.name)) for shop_key, items in result.items()}


def template_shift_match_score(template: ReportTemplate, shift_range: ShiftRange) -> float | None:
    if template.shift_begin is None or template.shift_end is None:
        return None
    return (
        abs((template.shift_begin - shift_range.begin).total_seconds())
        + abs((template.shift_end - shift_range.end).total_seconds())
    )


def match_templates_to_shifts(
    templates: list[ReportTemplate],
    shift_ranges: list[ShiftRange],
    tolerance_seconds: int = 300,
) -> tuple[list[tuple[ReportTemplate, ShiftRange, int]], list[ReportTemplate], set[int]]:
    matched: list[tuple[ReportTemplate, ShiftRange, int]] = []
    unmatched_templates: list[ReportTemplate] = []
    used_shift_indexes: set[int] = set()
    for template in templates:
        candidates: list[tuple[float, int, ShiftRange]] = []
        for index, shift_range in enumerate(shift_ranges):
            if index in used_shift_indexes:
                continue
            score = template_shift_match_score(template, shift_range)
            if score is not None:
                candidates.append((score, index, shift_range))
        if not candidates:
            unmatched_templates.append(template)
            continue
        score, index, shift_range = min(candidates, key=lambda item: item[0])
        if score <= tolerance_seconds:
            matched.append((template, shift_range, index))
            used_shift_indexes.add(index)
        else:
            unmatched_templates.append(template)
    return matched, unmatched_templates, used_shift_indexes


def load_shift_range_for_template(
    cur,
    shop_key: int,
    shift_begin: datetime | None,
    shift_end: datetime | None,
    tolerance_seconds: int = 300,
) -> ShiftRange | None:
    if shift_begin is None or shift_end is None:
        return None
    window_begin = shift_begin - timedelta(days=3)
    window_end = shift_end + timedelta(days=3)
    cur.execute(
        """
        SELECT ShopKey, ShiftKey, DatetimeShiftBegin, DatetimeShiftEnd,
               ABS(TIMESTAMPDIFF(SECOND, DatetimeShiftBegin, %s))
             + ABS(TIMESTAMPDIFF(SECOND, DatetimeShiftEnd, %s)) AS MatchScore
        FROM shift
        WHERE ShopKey = %s
          AND DatetimeShiftBegin >= %s
          AND DatetimeShiftEnd <= %s
        ORDER BY MatchScore, DatetimeShiftBegin, ShiftKey
        LIMIT 1
        """,
        (shift_begin, shift_end, shop_key, window_begin, window_end),
    )
    row = cur.fetchone()
    if not row or int(row.get("MatchScore") or 0) > tolerance_seconds:
        return None
    return ShiftRange(
        shop_key=int(row["ShopKey"]),
        shift_keys=(int(row["ShiftKey"]),),
        begin=row["DatetimeShiftBegin"],
        end=row["DatetimeShiftEnd"],
    )


def load_shift_ranges(cur, target_day: date) -> dict[int, list[ShiftRange]]:
    business_day = shift_business_date_expr("s")
    start = target_day.isoformat()
    end = (target_day + timedelta(days=1)).isoformat()
    cur.execute(
        f"""
        SELECT ShopKey, ShiftKey, DatetimeShiftBegin, DatetimeShiftEnd,
               TIMESTAMPDIFF(MINUTE, DatetimeShiftBegin, DatetimeShiftEnd) AS ShiftMinutes,
               (
                   SELECT COUNT(*)
                   FROM selling ss
                   WHERE ss.ShopKey = s.ShopKey
                     AND ss.ShiftKey = s.ShiftKey
               ) AS SellingRows
        FROM shift s
        WHERE {business_day} >= %s
          AND {business_day} < %s
          AND DatetimeShiftBegin IS NOT NULL
          AND DatetimeShiftEnd IS NOT NULL
          AND DatetimeShiftEnd > DatetimeShiftBegin
        ORDER BY ShopKey, DatetimeShiftBegin, ShiftKey
        """,
        (start, end),
    )
    grouped: dict[int, list[dict]] = defaultdict(list)
    for row in cur.fetchall():
        grouped[int(row["ShopKey"])].append(row)
    result: dict[int, list[ShiftRange]] = {}
    for shop_key, rows in grouped.items():
        if len(rows) > 1:
            filtered = [
                row
                for row in rows
                if not (int(row.get("ShiftMinutes") or 0) < 30 and int(row.get("SellingRows") or 0) == 0)
            ]
            if filtered:
                rows = filtered
        result[shop_key] = [
            ShiftRange(
                shop_key=shop_key,
                shift_keys=(int(row["ShiftKey"]),),
                begin=row["DatetimeShiftBegin"],
                end=row["DatetimeShiftEnd"],
            )
            for row in rows
        ]
    return result


def sql_in(values: tuple[int, ...]) -> str:
    return ",".join(["%s"] * len(values))


def load_balances(cur, shift_range: ShiftRange) -> tuple[list[dict], list[dict]]:
    placeholders = sql_in(shift_range.shift_keys)
    cur.execute(
        f"""
        SELECT BalanceShopShiftKey, ShopKey, ShiftKey, ResourceKey, ResourceShopKey,
               BeginBalance, EndBalance, EndBalanceCalculation,
               BeginMass, EndMass, BeginDensity, EndDensity,
               BeginPrice, EndPrice, BeginLevel, EndLevel,
               BeginTemperature, EndTemperature, WarehouseCode, VolumeType
        FROM balance
        WHERE ShopKey = %s
          AND ShiftKey IN ({placeholders})
          AND ResourceKey IS NOT NULL
        ORDER BY ShiftKey, BalanceShopShiftKey
        """,
        (shift_range.shop_key, *shift_range.shift_keys),
    )
    counter_rows: list[dict] = []
    tank_rows: list[dict] = []
    for row in cur.fetchall():
        row["resource"] = infer_resource(
            shift_range.shop_key,
            row.get("ResourceKey"),
            row.get("ResourceShopKey"),
            row.get("EndPrice") or row.get("BeginPrice"),
        )
        if row.get("BeginMass") is None and row.get("EndMass") is None:
            counter_rows.append(row)
        else:
            tank_rows.append(row)
    return counter_rows, tank_rows


def new_supply_bucket() -> dict:
    return {"volume": 0.0, "mass": 0.0, "density_volume": 0.0, "temperature_volume": 0.0, "docs": [], "entries": []}


def supply_doc_number(row: dict) -> str:
    doc_number = str(row.get("TTNNumberExt") or row.get("InvoiceNumberExt") or "").strip()
    if not doc_number:
        ttn = row.get("TTNNumber")
        if ttn not in (None, 0, 2147483647):
            doc_number = str(ttn)
    if not doc_number:
        supply_number = row.get("SupplyNumber")
        if supply_number not in (None, 0):
            doc_number = str(supply_number)
    return doc_number


def add_supply_entry(result: dict[int, dict], resource: int, entry: dict) -> None:
    item = result[resource]
    volume = float(entry.get("volume") or 0)
    mass = float(entry.get("mass") or 0)
    density = float(entry.get("density") or 0)
    temperature = float(entry.get("temperature") or 0)
    item["volume"] += volume
    item["mass"] += mass
    if density and volume:
        item["density_volume"] += density * volume
    if temperature and volume:
        item["temperature_volume"] += temperature * volume
    doc_number = str(entry.get("doc") or "").strip()
    if doc_number and doc_number not in item["docs"]:
        item["docs"].append(doc_number)
    item["entries"].append(entry)


def finalize_supply(result: dict[int, dict]) -> dict[int, dict]:
    for item in result.values():
        volume = item["volume"]
        item["density"] = item["mass"] / volume if volume else 0.0
        if item["density_volume"]:
            item["density"] = item["density_volume"] / volume
        item["temperature"] = item["temperature_volume"] / volume if item["temperature_volume"] and volume else 0.0
    return dict(result)


def load_supply(cur, shift_range: ShiftRange) -> dict[int, dict]:
    placeholders = sql_in(shift_range.shift_keys)
    result: dict[int, dict] = defaultdict(new_supply_bucket)
    cur.execute(
        f"""
        SELECT fc.ConsignmentPetrolKey, fc.ResourceKey, fc.ResourceShopKey,
               fc.VolumePN, fc.MassPN, fc.DensityPN, fc.TemperaturePN,
               fc.DatetimeBegin, fc.WarehouseCode,
               s.SupplyNumber, s.TTNNumber, s.TTNNumberExt, s.InvoiceNumberExt
        FROM fuel_consignment fc
        LEFT JOIN supply s
          ON s.ShopKey = fc.ShopKey
         AND s.ShiftKey = fc.ShiftKey
         AND s.SupplyCode = fc.ConsignmentPetrolKey
         AND COALESCE(s.Deleted, 0) = 0
        WHERE fc.ShopKey = %s
          AND fc.ShiftKey IN ({placeholders})
          AND COALESCE(fc.Deleted, 0) = 0
          AND COALESCE(fc.VolumePN, fc.Volume, 0) <> 0
        ORDER BY fc.DatetimeBegin, fc.ConsignmentPetrolKey
        """,
        (shift_range.shop_key, *shift_range.shift_keys),
    )
    consignment_rows = cur.fetchall()
    for row in consignment_rows:
        resource = infer_resource(shift_range.shop_key, row.get("ResourceKey"), row.get("ResourceShopKey"))
        entry = {
            "resource": resource,
            "resource_key": int(row.get("ResourceKey") or 0),
            "resource_shop_key": int(row.get("ResourceShopKey") or 0),
            "doc": supply_doc_number(row),
            "volume": float(row.get("VolumePN") or 0),
            "mass": float(row.get("MassPN") or 0),
            "density": float(row.get("DensityPN") or 0),
            "temperature": float(row.get("TemperaturePN") or 0),
            "warehouse_code": int(row.get("WarehouseCode") or 0),
            "sort_key": (
                row.get("DatetimeBegin") or datetime.min,
                int(row.get("ConsignmentPetrolKey") or 0),
                int(row.get("WarehouseCode") or 0),
            ),
        }
        add_supply_entry(result, resource, entry)
    if consignment_rows:
        return finalize_supply(result)

    cur.execute(
        f"""
        SELECT ResourceKey, ResourceShopKey, SupplyNumber, TTNNumber, TTNNumberExt, InvoiceNumberExt,
               Volume, Mass, MassTTN, Density, DensityTTN, DocDensity,
               Temperature, TemperatureTTN, SupplyCode, AcceptDateTime
        FROM supply
        WHERE ShopKey = %s
          AND ShiftKey IN ({placeholders})
          AND COALESCE(Deleted, 0) = 0
        ORDER BY AcceptDateTime, SupplyCode
        """,
        (shift_range.shop_key, *shift_range.shift_keys),
    )
    for row in cur.fetchall():
        resource = infer_resource(shift_range.shop_key, row.get("ResourceKey"), row.get("ResourceShopKey"))
        volume = float(row.get("Volume") or 0)
        mass = float(row.get("Mass") or row.get("MassTTN") or 0)
        density = float(row.get("Density") or row.get("DensityTTN") or row.get("DocDensity") or 0)
        temperature = float(row.get("Temperature") or row.get("TemperatureTTN") or 0)
        entry = {
            "resource": resource,
            "resource_key": int(row.get("ResourceKey") or 0),
            "resource_shop_key": int(row.get("ResourceShopKey") or 0),
            "doc": supply_doc_number(row),
            "volume": volume,
            "mass": mass,
            "density": density,
            "temperature": temperature,
            "warehouse_code": 0,
            "sort_key": (row.get("AcceptDateTime") or datetime.min, int(row.get("SupplyCode") or 0), 0),
        }
        add_supply_entry(result, resource, entry)
    return finalize_supply(result)


def new_sales_bucket() -> dict:
    return {field: 0.0 for field in SALES_FIELDS} | {"entries": [], "_entry_map": {}}


def add_sales_amounts(target: dict, base: int, derived: int, volume: float, amount: float) -> None:
    if base == 1:
        target["cash_volume"] += volume
        target["cash_amount"] += amount
    elif derived == 2:
        target["vedomosti"] += volume


def finalize_sales(result: dict[int, dict]) -> dict[int, dict]:
    for item in result.values():
        entry_map = item.pop("_entry_map", {})
        item["entries"] = sorted(
            entry_map.values(),
            key=lambda entry: (
                int(entry.get("resource_shop_key") or 0),
                int(entry.get("resource_key") or 0),
                float(entry.get("price") or 0),
            ),
        )
    return dict(result)


def load_sales(cur, shift_range: ShiftRange) -> dict[int, dict]:
    placeholders = sql_in(shift_range.shift_keys)
    cur.execute(
        f"""
        SELECT ResourceKey, ResourceShopKey, BasePaymentTypeKey, DerivedPaymentTypeKey,
               MAX(BasePrice) AS Price, SUM(Volume) AS Volume, SUM(ShopCost) AS Amount
        FROM selling
        WHERE ShopKey = %s
          AND ShiftKey IN ({placeholders})
        GROUP BY ResourceKey, ResourceShopKey, BasePaymentTypeKey, DerivedPaymentTypeKey
        """,
        (shift_range.shop_key, *shift_range.shift_keys),
    )
    result: dict[int, dict] = defaultdict(new_sales_bucket)
    for row in cur.fetchall():
        resource = infer_resource(shift_range.shop_key, row.get("ResourceKey"), row.get("ResourceShopKey"), row.get("Price"))
        base = int(row.get("BasePaymentTypeKey") or 0)
        derived = int(row.get("DerivedPaymentTypeKey") or 0)
        volume = float(row.get("Volume") or 0)
        amount = float(row.get("Amount") or 0)
        item = result[resource]
        add_sales_amounts(item, base, derived, volume, amount)
        identity = (
            int(row.get("ResourceKey") or 0),
            int(row.get("ResourceShopKey") or 0),
            round_float(row.get("Price") or 0, 2),
        )
        entry = item["_entry_map"].setdefault(
            identity,
            {
                **{field: 0.0 for field in SALES_FIELDS},
                "resource": resource,
                "resource_key": identity[0],
                "resource_shop_key": identity[1],
                "price": identity[2],
            },
        )
        add_sales_amounts(entry, base, derived, volume, amount)
    return finalize_sales(result)


def sales_cash_amount(sales: dict[int, dict]) -> float:
    return round_float(sum(float(item.get("cash_amount") or 0) for item in sales.values()), 2, adjust_float_tail=True)


def load_cash_carryover(cur, shift_range: ShiftRange) -> float:
    cur.execute(
        """
        SELECT COALESCE(SUM(s.ShopCost), 0) AS CashSales
        FROM selling s
        JOIN shift sh
          ON sh.ShopKey = s.ShopKey
         AND sh.ShiftKey = s.ShiftKey
        WHERE s.ShopKey = %s
          AND sh.DatetimeShiftEnd <= %s
          AND s.BasePaymentTypeKey = 1
        """,
        (shift_range.shop_key, shift_range.begin),
    )
    sales_row = cur.fetchone() or {}
    cur.execute(
        """
        SELECT COALESCE(SUM(pc.Collect), 0) AS CashCollect
        FROM payment_collection pc
        JOIN shift sh
          ON sh.ShopKey = pc.ShopKey
         AND sh.ShiftKey = pc.ShiftKey
        WHERE pc.ShopKey = %s
          AND sh.DatetimeShiftEnd <= %s
        """,
        (shift_range.shop_key, shift_range.begin),
    )
    collect_row = cur.fetchone() or {}
    cash_sales = float(sales_row.get("CashSales") or 0)
    cash_collect = float(collect_row.get("CashCollect") or 0)
    return round_float(cash_sales - cash_collect, 2, adjust_float_tail=True)


def shift_range_key(shift_range: ShiftRange) -> tuple[int, tuple[int, ...]]:
    return shift_range.shop_key, shift_range.shift_keys


def load_cash_carryovers(cur, shift_ranges: list[ShiftRange]) -> dict[tuple[int, tuple[int, ...]], float]:
    if not shift_ranges:
        return {}
    shop_keys = tuple(sorted({shift_range.shop_key for shift_range in shift_ranges}))
    max_begin = max(shift_range.begin for shift_range in shift_ranges)
    placeholders = sql_in(shop_keys)

    cur.execute(
        f"""
        SELECT sh.ShopKey, sh.ShiftKey, sh.DatetimeShiftEnd
        FROM shift sh
        WHERE sh.ShopKey IN ({placeholders})
          AND sh.DatetimeShiftEnd <= %s
        ORDER BY sh.ShopKey, sh.DatetimeShiftEnd, sh.ShiftKey
        """,
        (*shop_keys, max_begin),
    )
    prior_shifts = cur.fetchall()
    max_shift_by_shop: dict[int, int] = {}
    for row in prior_shifts:
        shop_key = int(row["ShopKey"])
        shift_key = int(row["ShiftKey"])
        max_shift_by_shop[shop_key] = max(max_shift_by_shop.get(shop_key, 0), shift_key)

    sales_by_shift: dict[tuple[int, int], float] = {}
    if max_shift_by_shop:
        range_clauses = " OR ".join(["(s.ShopKey = %s AND s.ShiftKey <= %s)" for _ in max_shift_by_shop])
        range_params: list[int] = []
        for shop_key, shift_key in sorted(max_shift_by_shop.items()):
            range_params.extend([shop_key, shift_key])
        cur.execute(
            f"""
            SELECT s.ShopKey, s.ShiftKey, COALESCE(SUM(s.ShopCost), 0) AS CashSales
            FROM selling s
            WHERE ({range_clauses})
              AND s.BasePaymentTypeKey = 1
            GROUP BY s.ShopKey, s.ShiftKey
            """,
            range_params,
        )
        sales_by_shift = {
            (int(row["ShopKey"]), int(row["ShiftKey"])): float(row.get("CashSales") or 0)
            for row in cur.fetchall()
        }

    cur.execute(
        f"""
        SELECT pc.ShopKey, pc.ShiftKey, COALESCE(SUM(pc.Collect), 0) AS CashCollect
        FROM payment_collection pc
        JOIN shift sh
          ON sh.ShopKey = pc.ShopKey
         AND sh.ShiftKey = pc.ShiftKey
        WHERE pc.ShopKey IN ({placeholders})
          AND sh.DatetimeShiftEnd <= %s
        GROUP BY pc.ShopKey, pc.ShiftKey
        """,
        (*shop_keys, max_begin),
    )
    collect_by_shift = {
        (int(row["ShopKey"]), int(row["ShiftKey"])): float(row.get("CashCollect") or 0)
        for row in cur.fetchall()
    }

    targets_by_shop: dict[int, list[ShiftRange]] = defaultdict(list)
    for shift_range in shift_ranges:
        targets_by_shop[shift_range.shop_key].append(shift_range)
    for items in targets_by_shop.values():
        items.sort(key=lambda item: (item.begin, item.shift_keys))

    result: dict[tuple[int, tuple[int, ...]], float] = {}
    cumulative_by_shop: dict[int, float] = defaultdict(float)
    index_by_shop: dict[int, int] = defaultdict(int)
    shifts_by_shop: dict[int, list[dict]] = defaultdict(list)
    for row in prior_shifts:
        shifts_by_shop[int(row["ShopKey"])].append(row)

    for shop_key, targets in targets_by_shop.items():
        shop_shifts = shifts_by_shop.get(shop_key, [])
        for target in targets:
            index = index_by_shop[shop_key]
            while index < len(shop_shifts) and shop_shifts[index]["DatetimeShiftEnd"] <= target.begin:
                shift_key = int(shop_shifts[index]["ShiftKey"])
                movement_key = (shop_key, shift_key)
                cumulative_by_shop[shop_key] += sales_by_shift.get(movement_key, 0.0) - collect_by_shift.get(movement_key, 0.0)
                index += 1
            index_by_shop[shop_key] = index
            result[shift_range_key(target)] = round_float(cumulative_by_shop[shop_key], 2, adjust_float_tail=True)
    return result


def load_collection(cur, shift_range: ShiftRange) -> dict:
    placeholders = sql_in(shift_range.shift_keys)
    cur.execute(
        f"""
        SELECT SUM(Collect) AS Collect, MIN(CashBoxNumber) AS CashBoxNumber
        FROM payment_collection
        WHERE ShopKey = %s
          AND ShiftKey IN ({placeholders})
        """,
        (shift_range.shop_key, *shift_range.shift_keys),
    )
    row = cur.fetchone() or {}
    return {
        "collect": float(row.get("Collect") or 0),
        "cash_box": row.get("CashBoxNumber") or 1,
    }


def grouped_by_resource(rows: list[dict]) -> dict[int, list[dict]]:
    grouped: dict[int, list[dict]] = defaultdict(list)
    for row in rows:
        grouped[int(row["resource"])].append(row)
    return grouped


def best_match(candidates: list[dict], used: set[int], checks: list[tuple[str, float]]) -> dict | None:
    remaining = [row for row in candidates if id(row) not in used]
    for field, expected in checks:
        if expected is None:
            continue
        matches = sorted(
            (
                (abs(float(row.get(field) or 0) - expected), row)
                for row in remaining
            ),
            key=lambda item: item[0],
        )
        if matches and matches[0][0] <= COUNTER_TOLERANCE:
            used.add(id(matches[0][1]))
            return matches[0][1]
    return None


def iter_main_specs(ws) -> list[dict]:
    receipt_row = find_label_row(ws, "Расшифровка поступления")
    specs: list[dict] = []
    current_resource = None
    for row in range(7, receipt_row - 1):
        name = ws.cell(row, 1).value
        resource = resource_from_name(name)
        if resource:
            current_resource = resource
        if current_resource is None:
            continue
        if ws.cell(row, 3).value not in ("л", "л."):
            continue
        specs.append(
            {
                "row": row,
                "resource": current_resource,
                "is_product_row": resource is not None,
                "template_begin": as_float(ws.cell(row, 13).value, None),
                "template_end": as_float(ws.cell(row, 10).value, None),
                "template_tank_begin": as_float(ws.cell(row, 5).value, None),
                "template_tank_end": as_float(ws.cell(row, 26).value, None),
            }
        )
    return specs


def iter_detail_specs(ws) -> list[dict]:
    receipt_row = find_label_row(ws, "Расшифровка поступления")
    cash_row = find_label_row(ws, "Движение наличных денег")
    specs: list[dict] = []
    for row in range(receipt_row + 4, cash_row - 1):
        resource = resource_from_name(ws.cell(row, 1).value)
        if resource:
            specs.append({"row": row, "resource": resource})
    return specs


def is_zero_like(value) -> bool:
    return value is not None and abs(float(value)) < 0.000001


def counter_match_checks(spec: dict, template_matches_shift: bool) -> list[tuple[str, float]]:
    template_begin = spec["template_begin"]
    template_end = spec["template_end"]
    if template_matches_shift:
        if is_zero_like(template_begin) and is_zero_like(template_end):
            return [("EndBalance", template_end)]
        return [
            ("EndBalance", template_end),
            ("BeginBalance", template_begin),
            ("BeginBalance", template_end),
            ("EndBalance", template_begin),
        ]
    return [
        ("BeginBalance", template_end),
        ("BeginBalance", template_begin),
        ("EndBalance", template_end),
        ("EndBalance", template_begin),
    ]


def assign_counter_rows(ws, counter_rows: list[dict], template_matches_shift: bool = False) -> dict[int, dict | None]:
    grouped = grouped_by_resource(counter_rows)
    for rows in grouped.values():
        rows.sort(key=lambda row: int(row.get("BalanceShopShiftKey") or 0))
    all_rows = sorted(counter_rows, key=lambda row: int(row.get("BalanceShopShiftKey") or 0))
    used: set[int] = set()
    assigned: dict[int, dict | None] = {}
    for spec in iter_main_specs(ws):
        row = spec["row"]
        checks = counter_match_checks(spec, template_matches_shift)
        candidates = grouped.get(spec["resource"], [])
        match = best_match(
            candidates,
            used,
            checks,
        )
        if match is None:
            match = best_match(all_rows, used, checks)
        assigned[row] = match
    return assigned


def assign_tank_rows(ws, tank_rows: list[dict]) -> dict[int, dict | None]:
    grouped = grouped_by_resource(tank_rows)
    for rows in grouped.values():
        rows.sort(key=lambda row: int(row.get("BalanceShopShiftKey") or 0))
    used: set[int] = set()
    assigned: dict[int, dict | None] = {}
    for spec in iter_main_specs(ws):
        if not spec["is_product_row"]:
            continue
        row = spec["row"]
        candidates = grouped.get(spec["resource"], [])
        assigned[row] = best_match(
            candidates,
            used,
            [
                ("BeginBalance", spec["template_tank_end"]),
                ("BeginBalance", spec["template_tank_begin"]),
                ("EndBalance", spec["template_tank_end"]),
                ("EndBalance", spec["template_tank_begin"]),
            ],
        )
    return assigned


def tank_density_by_resource(ws, tank_rows: list[dict]) -> dict[int, float]:
    result: dict[int, float] = {}
    assignments = assign_tank_rows(ws, tank_rows)
    for spec in iter_main_specs(ws):
        if not spec["is_product_row"]:
            continue
        tank = assignments.get(spec["row"])
        if not tank:
            continue
        density = tank.get("EndDensity") or tank.get("EndDensityCalculation")
        if density:
            result.setdefault(spec["resource"], float(density))
    return result


def supply_entry_score(ws, row_number: int, entry: dict) -> float:
    score = 0.0
    template_doc = str(ws.cell(row_number, 7).value or "").strip()
    entry_doc = str(entry.get("doc") or "").strip()
    if template_doc and entry_doc:
        score += 0 if template_doc == entry_doc else 1000000
    template_volume = as_float(ws.cell(row_number, 10).value, None)
    template_density = as_float(ws.cell(row_number, 13).value, None)
    template_mass = as_float(ws.cell(row_number, 14).value, None)
    template_temperature = as_float(ws.cell(row_number, 16).value, None)
    if template_volume is not None:
        score += abs(template_volume - float(entry.get("volume") or 0)) * 10
    if template_mass is not None:
        score += abs(template_mass - float(entry.get("mass") or 0))
    if template_density is not None:
        score += abs(template_density - float(entry.get("density") or 0)) * 10000
    if template_temperature is not None:
        score += abs(template_temperature - float(entry.get("temperature") or 0)) * 100
    sort_key = entry.get("sort_key") or (datetime.min, 0, 0)
    return score + (sort_key[1] if len(sort_key) > 1 else 0) / 100000000


def assign_supply_entries(ws, supply: dict[int, dict], template_matches_shift: bool = False) -> dict[int, dict | None]:
    specs = iter_detail_specs(ws)
    if template_matches_shift:
        specs_by_resource: dict[int, list[dict]] = defaultdict(list)
        for spec in specs:
            specs_by_resource[spec["resource"]].append(spec)
        assigned: dict[int, dict | None] = {spec["row"]: None for spec in specs}
        for resource, item in supply.items():
            rows = specs_by_resource.get(resource) or []
            if not rows:
                continue
            used_rows: set[int] = set()
            entries = sorted(item.get("entries") or [], key=lambda entry: entry.get("sort_key") or (datetime.min, 0, 0))
            for entry in entries:
                candidates = [spec for spec in rows if spec["row"] not in used_rows]
                if not candidates:
                    break
                spec = min(candidates, key=lambda candidate: supply_entry_score(ws, candidate["row"], entry))
                used_rows.add(spec["row"])
                assigned[spec["row"]] = entry
        return assigned

    grouped: dict[int, list[dict]] = {}
    for resource, item in supply.items():
        grouped[resource] = sorted(item.get("entries") or [], key=lambda entry: entry.get("sort_key") or (datetime.min, 0, 0))
    used: dict[int, set[int]] = defaultdict(set)
    assigned: dict[int, dict | None] = {}
    for spec in specs:
        resource = spec["resource"]
        entries = grouped.get(resource) or []
        candidates = [(index, entry) for index, entry in enumerate(entries) if index not in used[resource]]
        if not candidates:
            assigned[spec["row"]] = None
            continue
        index, entry = min(candidates, key=lambda item: supply_entry_score(ws, spec["row"], item[1]))
        used[resource].add(index)
        assigned[spec["row"]] = entry
    return assigned


def clear_generated_cells(ws) -> None:
    receipt_row = find_label_row(ws, "Расшифровка поступления")
    cash_row = find_label_row(ws, "Движение наличных денег")
    for row in range(7, receipt_row - 1):
        for col in (5, 6, 10, 13, 15, 19, 20, 23, 26, 30, 32, 36, 38):
            ws.cell(row, col).value = None
    for row in range(receipt_row + 4, cash_row - 1):
        for col in (7, 10, 13, 14, 16, 19, 21, 24, 27, 30, 31, 35, 37, 39):
            ws.cell(row, col).value = None


def fill_main_table(
    ws,
    shift_range: ShiftRange,
    counter_rows: list[dict],
    tank_rows: list[dict],
    supply: dict[int, dict],
    preserve_supply_values: dict[int, object] | None = None,
    preserve_template_main_display: bool = False,
) -> list[str]:
    warnings: list[str] = []
    counter_assignments = assign_counter_rows(ws, counter_rows, template_matches_shift=preserve_template_main_display)
    tank_assignments = assign_tank_rows(ws, tank_rows)
    supplied_once_without_tank: set[int] = set()
    all_supply_volume = sum(float(item.get("volume") or 0) for item in supply.values())
    for spec in iter_main_specs(ws):
        row_number = spec["row"]
        counter = counter_assignments.get(row_number)
        if counter:
            begin = float(counter.get("BeginBalance") or 0)
            end = float(counter.get("EndBalance") or 0)
            ws.cell(row_number, 13).value = round_float(begin, 2)
            ws.cell(row_number, 10).value = round_float(end, 2)
            ws.cell(row_number, 15).value = round_float(end - begin, 2)
        elif abs(float(spec.get("template_begin") or 0)) < 0.000001 and abs(float(spec.get("template_end") or 0)) < 0.000001:
            ws.cell(row_number, 13).value = 0.0
            ws.cell(row_number, 10).value = 0.0
            ws.cell(row_number, 15).value = 0.0
        else:
            warnings.append(f"{shift_range.shop_key:03d}: 未匹配机表行 {row_number}")

        if not spec["is_product_row"]:
            continue

        resource = spec["resource"]
        tank = tank_assignments.get(row_number)
        if tank:
            begin = float(tank.get("BeginBalance") or 0)
            end = float(tank.get("EndBalance") or 0)
            calc = float(tank.get("EndBalanceCalculation") or end)
            diff = 0.0 if abs(end) < 0.000001 and abs(calc) > 0.000001 else round_float(end - calc, 2)
            display_end = calc if (abs(end) < 0.000001 and abs(calc) > 0.000001) or abs(end - calc) < 0.001 else end
            begin_value = round_float(begin, 2)
            if preserve_template_main_display and abs(begin_value) < 0.000001 and ws.cell(row_number, 5).value in (None, ""):
                pass
            else:
                ws.cell(row_number, 5).value = begin_value

            existing_display_end = parse_cash_cell(ws.cell(row_number, 26).value)
            if (
                preserve_template_main_display
                and existing_display_end is not None
                and abs(existing_display_end - float(display_end)) < 0.001
            ):
                pass
            else:
                ws.cell(row_number, 26).value = comma_number(display_end, 6, trim=True)
            ws.cell(row_number, 30).value = round_float(calc, 2)
            ws.cell(row_number, 32).value = value_or_blank(ws.cell(row_number, 32).value, max(diff, 0), 2)
            ws.cell(row_number, 36).value = value_or_blank(ws.cell(row_number, 36).value, min(diff, 0), 2)
            ws.cell(row_number, 38).value = 0.0
        else:
            warnings.append(f"{shift_range.shop_key:03d}: 未匹配油罐行 {row_number}")

        item = supply.get(resource, {})
        entries = item.get("entries") or []
        entries_have_warehouse = any(int(entry.get("warehouse_code") or 0) for entry in entries)
        volume = 0.0
        matched_by_tank = False
        if tank and entries_have_warehouse:
            warehouse_code = int(tank.get("WarehouseCode") or 0)
            tank_entries = [
                entry
                for entry in entries
                if int(entry.get("warehouse_code") or 0) == warehouse_code
            ]
            matched_by_tank = True
            volume = sum(float(entry.get("volume") or 0) for entry in tank_entries)
        if not matched_by_tank and resource not in supplied_once_without_tank:
            supplied_once_without_tank.add(resource)
            volume = float(item.get("volume") or 0)
        if preserve_supply_values is not None and row_number in preserve_supply_values:
            ws.cell(row_number, 6).value = preserve_supply_values[row_number]
        elif preserve_template_main_display:
            existing_supply = as_float(ws.cell(row_number, 6).value, None)
            if (
                existing_supply is not None
                and abs(volume) < 0.000001
                and abs(existing_supply - all_supply_volume) < 0.01
            ):
                pass
            else:
                ws.cell(row_number, 6).value = value_or_blank(ws.cell(row_number, 6).value, volume, 2)
        else:
            ws.cell(row_number, 6).value = value_or_blank(ws.cell(row_number, 6).value, volume, 2)
    return warnings


def empty_sales_item() -> dict:
    return {field: 0.0 for field in SALES_FIELDS}


def sales_entry_score(ws, row_number: int, entry: dict) -> float:
    score = 0.0
    checks = [
        (19, "talon_sber", 10),
        (21, "talon_zenit", 10),
        (24, "card_cnk", 10),
        (27, "vedomosti", 10),
        (30, "cash_volume", 10),
        (31, "cash_amount", 1),
        (39, "card_ts", 10),
    ]
    for column, field, weight in checks:
        template_value = as_float(ws.cell(row_number, column).value, None)
        if template_value is not None:
            score += abs(template_value - float(entry.get(field) or 0)) * weight
    return score


def aggregate_sales_items(items: list[dict]) -> dict:
    result = empty_sales_item()
    for item in items:
        for field in SALES_FIELDS:
            result[field] += float(item.get(field) or 0)
    return result


def entry_identity(entry: dict) -> tuple[int, int]:
    return int(entry.get("resource_key") or 0), int(entry.get("resource_shop_key") or 0)


def assign_sales_entries(
    ws,
    sales: dict[int, dict],
    supply_assignments: dict[int, dict | None] | None = None,
    template_matches_shift: bool = False,
) -> dict[int, dict]:
    specs_by_resource: dict[int, list[dict]] = defaultdict(list)
    for spec in iter_detail_specs(ws):
        specs_by_resource[spec["resource"]].append(spec)

    assigned: dict[int, dict] = {}
    for resource, specs in specs_by_resource.items():
        original_spec_count = len(specs)
        item = sales.get(resource) or empty_sales_item()
        entries = list(item.get("entries") or [])
        if supply_assignments and entries:
            entries_by_identity = {entry_identity(entry): entry for entry in entries}
            for spec in specs:
                supply_entry = supply_assignments.get(spec["row"])
                if not supply_entry:
                    continue
                matching_entry = entries_by_identity.get(entry_identity(supply_entry))
                if matching_entry:
                    assigned[spec["row"]] = matching_entry

        specs = [spec for spec in specs if spec["row"] not in assigned]
        if not specs:
            continue
        if len(specs) == 1:
            if template_matches_shift and original_spec_count > 1 and entries:
                assigned[specs[0]["row"]] = min(entries, key=lambda entry: sales_entry_score(ws, specs[0]["row"], entry))
            else:
                assigned[specs[0]["row"]] = item
            continue
        if not entries:
            for spec in specs:
                assigned[spec["row"]] = empty_sales_item()
            continue
        if len(entries) == 1:
            target_spec = min(specs, key=lambda spec: sales_entry_score(ws, spec["row"], entries[0])) if template_matches_shift else specs[0]
            for spec in specs:
                assigned[spec["row"]] = entries[0] if spec["row"] == target_spec["row"] else empty_sales_item()
            continue
        if template_matches_shift:
            used_entries: set[int] = set()
            for spec in specs:
                candidates = [(index, entry) for index, entry in enumerate(entries) if index not in used_entries]
                if not candidates:
                    assigned[spec["row"]] = empty_sales_item()
                    continue
                index, entry = min(candidates, key=lambda candidate: sales_entry_score(ws, spec["row"], candidate[1]))
                used_entries.add(index)
                assigned[spec["row"]] = entry
            continue

        for index, spec in enumerate(specs):
            if index < len(entries) - 1:
                assigned[spec["row"]] = entries[index]
            elif index == len(specs) - 1:
                assigned[spec["row"]] = aggregate_sales_items(entries[index:])
            elif index < len(entries):
                assigned[spec["row"]] = entries[index]
            else:
                assigned[spec["row"]] = empty_sales_item()
    return assigned


def detail_doc_value(existing, docs: list[str], template_matches_shift: bool):
    value = "\n".join([doc for doc in docs if doc])
    if template_matches_shift and isinstance(existing, str) and existing.strip() == value:
        return existing
    return value


def fill_detail_table(ws, supply: dict[int, dict], sales: dict[int, dict], template_matches_shift: bool = False) -> None:
    supply_assignments = assign_supply_entries(ws, supply, template_matches_shift=template_matches_shift)
    sales_assignments = assign_sales_entries(
        ws,
        sales,
        supply_assignments=supply_assignments,
        template_matches_shift=template_matches_shift,
    )
    for spec in iter_detail_specs(ws):
        row_number = spec["row"]
        sales_item = sales_assignments.get(row_number) or empty_sales_item()
        supply_entry = supply_assignments.get(row_number)
        if supply_entry:
            supply_volume = float(supply_entry.get("volume") or 0)
            supply_mass = float(supply_entry.get("mass") or 0)
            supply_density = float(supply_entry.get("density") or 0)
            supply_temperature = float(supply_entry.get("temperature") or 0)
            supply_docs = [str(supply_entry.get("doc") or "").strip()]
        else:
            supply_volume = 0.0
            supply_mass = 0.0
            supply_density = 0.0
            supply_temperature = 0.0
            supply_docs = []

        ws.cell(row_number, 7).value = detail_doc_value(ws.cell(row_number, 7).value, supply_docs, template_matches_shift)
        ws.cell(row_number, 10).value = round_float(supply_volume, 2)
        ws.cell(row_number, 13).value = comma_number(supply_density, 15, trim=True) if supply_density else 0
        ws.cell(row_number, 14).value = round_float(supply_mass, 2)
        ws.cell(row_number, 16).value = comma_number(supply_temperature, 13, trim=True) if supply_temperature else 0
        ws.cell(row_number, 19).value = round_float(sales_item.get("talon_sber", 0), 2, adjust_float_tail=True)
        ws.cell(row_number, 21).value = round_float(sales_item.get("talon_zenit", 0), 2, adjust_float_tail=True)
        ws.cell(row_number, 24).value = round_float(sales_item.get("card_cnk", 0), 2, adjust_float_tail=True)
        ws.cell(row_number, 27).value = round_float(sales_item.get("vedomosti", 0), 2, adjust_float_tail=True)
        ws.cell(row_number, 30).value = round_float(sales_item.get("cash_volume", 0), 2, adjust_float_tail=True)
        ws.cell(row_number, 31).value = round_float(sales_item.get("cash_amount", 0), 2)
        ws.cell(row_number, 39).value = round_float(sales_item.get("card_ts", 0), 2, adjust_float_tail=True)


def zero_amount_words_like_template(existing, fallback: str) -> str:
    if not isinstance(existing, str):
        return fallback
    first_word = fallback.split(" ", 1)[0]
    if existing.strip().casefold().startswith(first_word.casefold()):
        return existing
    return fallback


def parse_cash_cell(value) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float, Decimal)):
        return float(value)
    text = str(value).strip().replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def template_cash_incoming(ws, template_matches_shift: bool) -> float | None:
    try:
        cash_row = find_label_row(ws, "Движение наличных денег")
    except ValueError:
        return None
    source_row = cash_row + 2 if template_matches_shift else cash_row + 12
    return parse_cash_cell(ws.cell(source_row, 18).value)


def cash_amount_cell_value(amount: float):
    if abs(amount) < 0.000001:
        return 0
    return comma_number(amount, 2, trim=True)


def fill_cash_table(ws, collection: dict, preserve_template_zero_cash: bool = False) -> None:
    cash_row = find_label_row(ws, "Движение наличных денег")
    incoming = float(collection.get("incoming") or 0)
    cash_sales = float(collection.get("cash_sales") or 0)
    collect = float(collection.get("collect") or 0)
    depot_cash = float(collection.get("depot_cash") or 0)
    total_cash = round_float(incoming + cash_sales, 2, adjust_float_tail=True)
    handover = round_float(total_cash - collect - depot_cash, 2, adjust_float_tail=True)
    zero_words = amount_words_ru(0)

    amounts = {
        cash_row + 2: incoming,
        cash_row + 4: cash_sales,
        cash_row + 6: total_cash,
        cash_row + 8: collect,
        cash_row + 10: depot_cash,
        cash_row + 12: handover,
    }
    for row_number, amount in amounts.items():
        existing_words = ws.cell(row_number, 4).value
        existing_cash_value = ws.cell(row_number, 18).value
        existing_amount = parse_cash_cell(existing_cash_value)
        preserve_existing_cash = (
            preserve_template_zero_cash
            and existing_amount is not None
            and abs(round_float(existing_amount, 2, adjust_float_tail=True) - round_float(amount, 2, adjust_float_tail=True)) < 0.005
        )
        if abs(amount) < 0.000001:
            if preserve_existing_cash and isinstance(existing_words, str):
                words = existing_words
            else:
                words = zero_amount_words_like_template(existing_words, zero_words)
        else:
            words = existing_words if preserve_existing_cash and isinstance(existing_words, str) else amount_words_ru(amount)
        ws.cell(row_number, 4).value = words
        if preserve_existing_cash:
            continue
        ws.cell(row_number, 18).value = cash_amount_cell_value(amount)


def update_shift_header(ws, shift_range: ShiftRange) -> None:
    ws["K3"].value = f"Смена с {dt_text(shift_range.begin)} до {dt_text(shift_range.end)}"


def generate_station_report(
    cur,
    template_path: Path,
    output_path: Path,
    shift_range: ShiftRange,
    preserve_template_supply_blanks: bool = False,
    template_matches_shift: bool = False,
    cash_carryover: float | None = None,
) -> dict:
    wb = openpyxl.load_workbook(template_path)
    ws = wb.active
    original_values = worksheet_value_snapshot(ws)
    preserve_supply_values = None
    if preserve_template_supply_blanks:
        preserve_supply_values = {
            spec["row"]: ws.cell(spec["row"], 6).value
            for spec in iter_main_specs(ws)
            if ws.cell(spec["row"], 6).value in (None, "")
        }
    counter_rows, tank_rows = load_balances(cur, shift_range)
    supply = load_supply(cur, shift_range)
    sales = load_sales(cur, shift_range)
    collection = load_collection(cur, shift_range)
    collection["cash_sales"] = sales_cash_amount(sales)
    template_incoming = template_cash_incoming(ws, template_matches_shift)
    if template_incoming is not None:
        collection["incoming"] = template_incoming
    else:
        collection["incoming"] = load_cash_carryover(cur, shift_range) if cash_carryover is None else cash_carryover

    if not template_matches_shift:
        update_shift_header(ws, shift_range)
    warnings = fill_main_table(
        ws,
        shift_range,
        counter_rows,
        tank_rows,
        supply,
        preserve_supply_values,
        preserve_template_main_display=template_matches_shift,
    )
    fill_detail_table(ws, supply, sales, template_matches_shift=template_matches_shift)
    fill_cash_table(ws, collection, preserve_template_zero_cash=template_matches_shift)

    write_values_preserving_template_package(template_path, output_path, ws, original_values)
    return {
        "shopKey": shift_range.shop_key,
        "file": str(output_path),
        "shiftKeys": list(shift_range.shift_keys),
        "shiftBegin": shift_range.begin.isoformat(sep=" "),
        "shiftEnd": shift_range.end.isoformat(sep=" "),
        "cashCollect": round_float(collection.get("collect", 0), 2),
        "warnings": warnings,
    }


def find_rar_exe() -> Path | None:
    candidates = [
        Path(r"C:\Program Files\WinRAR\Rar.exe"),
        Path(r"C:\Program Files (x86)\WinRAR\Rar.exe"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    found = shutil.which("rar.exe")
    return Path(found) if found else None


def archive_folder(output_dir: Path, archive_format: str = "rar") -> Path:
    output_root = output_dir.parent
    if archive_format == "rar":
        rar_exe = find_rar_exe()
        if rar_exe:
            archive_path = output_root / f"{output_dir.name}.rar"
            if archive_path.exists():
                archive_path.unlink()
            completed = subprocess.run(
                [str(rar_exe), "a", "-r", "-idq", "-y", str(archive_path), output_dir.name],
                cwd=str(output_root),
                capture_output=True,
                text=True,
                timeout=300,
            )
            if completed.returncode != 0:
                raise RuntimeError((completed.stderr or completed.stdout or "").strip())
            return archive_path

    archive_path = output_root / f"{output_dir.name}.zip"
    if archive_path.exists():
        archive_path.unlink()
    with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for file_path in sorted(output_dir.rglob("*")):
            if file_path.is_file():
                archive.write(file_path, file_path.relative_to(output_root))
    return archive_path


def generate_reports(
    target_day: date,
    report_root: Path = REPORT_ROOT,
    template_dir: Path | None = None,
    output_dir: Path | None = None,
    force: bool = False,
    archive: bool = False,
    archive_format: str = "rar",
) -> dict:
    template_path = template_dir or find_template_dir(report_root, target_day, None)
    output_path = output_dir or report_root / day_folder_name(target_day)
    if output_path.exists() and any(output_path.glob("*.xlsx")) and not force:
        raise FileExistsError(f"输出目录已存在日报文件，如需重生成请使用 --force: {output_path}")
    if output_path.exists() and force:
        shutil.rmtree(output_path)
    output_path.mkdir(parents=True, exist_ok=True)

    templates_by_shop = report_templates(template_path)
    if not templates_by_shop:
        raise FileNotFoundError(f"模板目录没有 xlsx 文件: {template_path}")

    same_day_template_dir = template_path.name == day_folder_name(target_day)
    generated: list[dict] = []
    skipped: list[dict] = []
    with connect() as conn:
        cur = conn.cursor()
        shift_ranges = load_shift_ranges(cur, target_day)
        cash_carryovers: dict[tuple[int, tuple[int, ...]], float] = {}
        for shop_key in sorted(templates_by_shop):
            shop_templates = templates_by_shop[shop_key]
            shop_shifts = shift_ranges.get(shop_key) or []
            if not shop_shifts:
                if same_day_template_dir:
                    used_shift_keys: set[tuple[int, ...]] = set()
                    remaining_templates: list[ReportTemplate] = []
                    for template in shop_templates:
                        template_shift = load_shift_range_for_template(cur, shop_key, template.shift_begin, template.shift_end)
                        if template_shift and template_shift.shift_keys not in used_shift_keys:
                            used_shift_keys.add(template_shift.shift_keys)
                            generated.append(generate_station_report(
                                cur,
                                template.path,
                                output_path / template.path.name,
                                template_shift,
                                preserve_template_supply_blanks=True,
                                template_matches_shift=True,
                                cash_carryover=cash_carryovers.get(shift_range_key(template_shift)),
                            ))
                            continue
                        remaining_templates.append(template)
                    if used_shift_keys:
                        for template in remaining_templates:
                            skipped.append({
                                "shopKey": shop_key,
                                "file": template.path.name,
                                "reason": "template has no matching completed shift",
                            })
                        continue
                skipped.append({"shopKey": shop_key, "reason": "target day has no completed shift"})
                continue

            matched_templates, unmatched_templates, used_shift_indexes = match_templates_to_shifts(shop_templates, shop_shifts)
            if matched_templates:
                used_shift_keys = {shift_range.shift_keys for _, shift_range, _ in matched_templates}
                for template, shift_range, _ in matched_templates:
                    generated.append(generate_station_report(
                        cur,
                        template.path,
                        output_path / template.path.name,
                        shift_range,
                        preserve_template_supply_blanks=True,
                        template_matches_shift=True,
                        cash_carryover=cash_carryovers.get(shift_range_key(shift_range)),
                    ))
                remaining_unmatched_templates: list[ReportTemplate] = []
                for template in unmatched_templates:
                    template_shift = load_shift_range_for_template(cur, shop_key, template.shift_begin, template.shift_end) if same_day_template_dir else None
                    if template_shift and template_shift.shift_keys not in used_shift_keys:
                        used_shift_keys.add(template_shift.shift_keys)
                        generated.append(generate_station_report(
                            cur,
                            template.path,
                            output_path / template.path.name,
                            template_shift,
                            preserve_template_supply_blanks=True,
                            template_matches_shift=True,
                            cash_carryover=cash_carryovers.get(shift_range_key(template_shift)),
                        ))
                        continue
                    remaining_unmatched_templates.append(template)
                for template in remaining_unmatched_templates:
                    skipped.append({
                        "shopKey": shop_key,
                        "file": template.path.name,
                        "reason": "template shift time does not match target day",
                    })
                for sequence, shift_range in enumerate(shop_shifts):
                    if sequence not in used_shift_indexes:
                        skipped.append({
                            "shopKey": shop_key,
                            "shiftKeys": list(shift_range.shift_keys),
                            "shiftBegin": shift_range.begin.isoformat(sep=" "),
                            "shiftEnd": shift_range.end.isoformat(sep=" "),
                            "reason": "target shift has no matching template file",
                        })
                continue

            if same_day_template_dir:
                used_shift_keys: set[tuple[int, ...]] = set()
                remaining_templates: list[ReportTemplate] = []
                for template in shop_templates:
                    template_shift = load_shift_range_for_template(cur, shop_key, template.shift_begin, template.shift_end)
                    if template_shift and template_shift.shift_keys not in used_shift_keys:
                        used_shift_keys.add(template_shift.shift_keys)
                        generated.append(generate_station_report(
                            cur,
                            template.path,
                            output_path / template.path.name,
                            template_shift,
                            preserve_template_supply_blanks=True,
                            template_matches_shift=True,
                            cash_carryover=cash_carryovers.get(shift_range_key(template_shift)),
                        ))
                        continue
                    remaining_templates.append(template)
                if used_shift_keys:
                    for template in remaining_templates:
                        skipped.append({
                            "shopKey": shop_key,
                            "file": template.path.name,
                            "reason": "template has no matching completed shift",
                        })
                    for sequence, shift_range in enumerate(shop_shifts):
                        if shift_range.shift_keys not in used_shift_keys:
                            skipped.append({
                                "shopKey": shop_key,
                                "shiftKeys": list(shift_range.shift_keys),
                                "shiftBegin": shift_range.begin.isoformat(sep=" "),
                                "shiftEnd": shift_range.end.isoformat(sep=" "),
                                "reason": "target shift has no matching template file",
                            })
                    continue

            template_by_sequence = {template.sequence: template for template in shop_templates}
            fallback_template = template_by_sequence.get(0) or shop_templates[0]
            for sequence, shift_range in enumerate(shop_shifts):
                template = template_by_sequence.get(sequence) or fallback_template
                output_name = report_file_name(shop_key, sequence)
                generated.append(generate_station_report(
                    cur,
                    template.path,
                    output_path / output_name,
                    shift_range,
                    cash_carryover=cash_carryovers.get(shift_range_key(shift_range)),
                ))
            for template in shop_templates:
                if template.sequence >= len(shop_shifts):
                    skipped.append({
                        "shopKey": shop_key,
                        "file": template.path.name,
                        "reason": "template has no matching completed shift",
                    })

    archive_path = None
    if archive and generated:
        archive_path = archive_folder(output_path, archive_format=archive_format)

    warnings = []
    for item in generated:
        for warning in item.get("warnings") or []:
            warnings.append(warning)

    return {
        "targetDate": target_day.isoformat(),
        "templateDir": str(template_path),
        "outputDir": str(output_path),
        "generatedCount": len(generated),
        "skippedCount": len(skipped),
        "generated": generated,
        "skipped": skipped,
        "warnings": warnings,
        "archive": str(archive_path) if archive_path else None,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate daily station accounting handover reports.")
    parser.add_argument("--date", help="业务日期，格式 YYYY-MM-DD；默认生成昨天")
    parser.add_argument("--report-root", default=str(REPORT_ROOT), help="日报根目录")
    parser.add_argument("--template-dir", help="模板目录，默认取目标日期之前最近的日报文件夹")
    parser.add_argument("--output-dir", help="输出目录，默认 D:\\每日销量日报\\YYYYMMDD")
    parser.add_argument("--force", action="store_true", help="先清空输出目录再生成")
    parser.add_argument("--archive", action="store_true", help="生成后打包")
    parser.add_argument("--archive-format", choices=["rar", "zip"], default="rar", help="优先打包格式")
    args = parser.parse_args()

    target_day = parse_day(args.date)
    result = generate_reports(
        target_day=target_day,
        report_root=Path(args.report_root),
        template_dir=Path(args.template_dir) if args.template_dir else None,
        output_dir=Path(args.output_dir) if args.output_dir else None,
        force=args.force,
        archive=args.archive,
        archive_format=args.archive_format,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
