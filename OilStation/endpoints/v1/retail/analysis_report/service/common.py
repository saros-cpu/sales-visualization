import html
import json
import os
from pathlib import Path
import re
import zipfile
from datetime import date, datetime, timedelta
from urllib.parse import quote
from xml.sax.saxutils import escape

from flask import Blueprint, jsonify, send_file
from openpyxl import load_workbook
import pymysql

from business_time import (
    business_date_expr,
    business_period_condition_with_pruning,
    business_period_params,
    shift_join,
)
from cache_store import get_cached, set_cached
from constants import SHOP_KEY_MAP
from db import get_connection
from export_excel import EXPORT_DIR, RESOURCE_DENSITY, _build_sheet_data, _round_half_up
from endpoints.v1.retail.analytics.service.shared import (
    format_resource_rows,
    report_amount_expr,
    report_bonus_amount_expr,
    report_card_amount_expr,
    report_cash_amount_expr,
    report_resource_key_expr,
    report_volume_expr,
)


from endpoints.v1.retail.common.query import get_query_date_range, query_month_summary

analysis_report_api = Blueprint("analysis_report_api", __name__)
PROJECT_ROOT = Path(__file__).resolve().parents[6]
OILSTATION_ROOT = Path(__file__).resolve().parents[5]
REPORT_EXTRAS_DIR = OILSTATION_ROOT / "cache_snapshots" / "report_extras"
REPORT_MONTH_NAMES_RU = {
    1: "Январь",
    2: "Февраль",
    3: "Март",
    4: "Апрель",
    5: "Май",
    6: "Июнь",
    7: "Июль",
    8: "Август",
    9: "Сентябрь",
    10: "Октябрь",
    11: "Ноябрь",
    12: "Декабрь",
}
REPORT_MONTH_NAME_TO_NUMBER = {
    name.lower(): month for month, name in REPORT_MONTH_NAMES_RU.items()
}


def _empty_report_extras():
    return {
        "totalEndInventory": 0.0,
        "totalEndInventoryTons": 0.0,
        "totalContractVolume": 0.0,
        "totalContractAmount": 0.0,
        "stations": {},
    }


def register_analysis_report_routes(app):
    pass  # registered via new controller


def _calc_prev_month(year, month):
    if month == 1:
        return year - 1, 12
    return year, month - 1


def _pct(current, previous):
    if not previous:
        return None
    return _round_half_up((float(current or 0) - float(previous or 0)) / float(previous) * 100, 1)


def _fmt(value, digits=2):
    digits = int(digits or 0)
    return f"{_round_half_up(value, digits):,.{digits}f}"


def _fmt_pct(value):
    if value is None:
        return "--"
    sign = "+" if value > 0 else ""
    return f"{sign}{_round_half_up(value, 1):.1f}%"


def _date_label(value):
    try:
        day = datetime.strptime(str(value), "%Y-%m-%d").date()
        return f"{day.year}年{day.month}月{day.day}日"
    except (TypeError, ValueError):
        return str(value or "")


def _period_label(data):
    start = data.get("queryStart")
    end_exclusive = data.get("queryEndExclusive")
    try:
        end_day = datetime.strptime(str(end_exclusive), "%Y-%m-%d").date() - timedelta(days=1)
        end = end_day.isoformat()
    except (TypeError, ValueError):
        end = end_exclusive
    return f"{_date_label(start)} — {_date_label(end)}"


def _docx_report_title(data):
    if data.get("kind") == "monthly":
        return f"{data.get('year')}年{data.get('month')}月份零售数据分析报告"
    return f"{data.get('year')}年零售数据分析报告"


def _load_env_file(path):
    if not path.exists():
        return
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def _sales_db_config():
    _load_env_file(PROJECT_ROOT / "Sales" / ".env")
    _load_env_file(PROJECT_ROOT / "OilStation" / ".env")
    database = os.getenv("SALES_DB_NAME", os.getenv("DB_NAME", "sales_visualization"))
    if database == "ocw":
        database = "sales_visualization"
    return {
        "host": os.getenv("DB_HOST", "127.0.0.1"),
        "port": int(os.getenv("DB_PORT", "3306")),
        "user": os.getenv("DB_USER", "reader"),
        "password": os.getenv("DB_PASSWORD", ""),
        "database": database,
        "charset": os.getenv("DB_CHARSET", "utf8mb4"),
        "use_unicode": True,
        "autocommit": True,
    }


def _exchange_rates_for_period(start, end):
    empty = {
        "rows": [],
        "count": 0,
        "totalDays": 0,
        "missingCount": 0,
        "avgUsd": None,
        "avgCny": None,
        "first": None,
        "last": None,
        "error": "",
    }
    try:
        start_day = datetime.strptime(start, "%Y-%m-%d").date()
        end_day = datetime.strptime(end, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return empty
    if end_day <= start_day:
        return empty

    try:
        conn = pymysql.connect(**_sales_db_config())
    except Exception as exc:
        result = dict(empty)
        result["error"] = f"汇率数据库连接失败：{exc}"
        return result

    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT rate_date, usd_kgs, cny_kgs, source, is_manual
            FROM exchange_rates
            WHERE rate_date >= %s AND rate_date < %s
            ORDER BY rate_date
            """,
            (start_day, end_day),
        )
        rows = []
        for rate_date, usd_kgs, cny_kgs, source, is_manual in cur.fetchall():
            rows.append({
                "date": rate_date.isoformat(),
                "usdKgs": float(usd_kgs) if usd_kgs is not None else None,
                "cnyKgs": float(cny_kgs) if cny_kgs is not None else None,
                "source": "手工录入" if is_manual else (source or "NBKR"),
            })
    finally:
        conn.close()

    present = {item["date"] for item in rows}
    total_days = (end_day - start_day).days
    missing_count = sum(
        1 for offset in range(total_days)
        if (start_day + timedelta(days=offset)).isoformat() not in present
    )
    usd_values = [item["usdKgs"] for item in rows if item["usdKgs"] is not None]
    cny_values = [item["cnyKgs"] for item in rows if item["cnyKgs"] is not None]
    return {
        "rows": rows,
        "count": len(rows),
        "totalDays": total_days,
        "missingCount": missing_count,
        "avgUsd": _round_half_up(sum(usd_values) / len(usd_values), 4) if usd_values else None,
        "avgCny": _round_half_up(sum(cny_values) / len(cny_values), 4) if cny_values else None,
        "first": rows[0] if rows else None,
        "last": rows[-1] if rows else None,
        "error": "",
    }


def _resource_tons(resources):
    total = 0.0
    for item in resources:
        if item.get("tons") is not None:
            total += float(item.get("tons") or 0)
            continue
        if item.get("mass") is not None:
            total += float(item.get("mass") or 0) / 1000
            continue
        density = RESOURCE_DENSITY.get(int(item.get("resourceKey") or item.get("key") or 0), 0)
        total += float(item.get("volume") or 0) * density / 1000
    return _round_half_up(total, 3)


def _month_after(year, month):
    return (year + 1, 1) if month == 12 else (year, month + 1)


def _is_completed_ym(year, month):
    today = datetime.now().date()
    return (year, month) < (today.year, today.month)


def _is_snapshot_writable_ym(year, month):
    today = datetime.now().date()
    return (year, month) <= (today.year, today.month)


def _ym_for_full_month(start, end):
    try:
        start_day = datetime.strptime(start, "%Y-%m-%d").date()
        end_day = datetime.strptime(end, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None
    if start_day.day != 1:
        return None
    next_year, next_month = _month_after(start_day.year, start_day.month)
    if end_day.year == next_year and end_day.month == next_month and end_day.day == 1:
        return f"{start_day.year}-{start_day.month:02d}"
    return None


def _report_extras_snapshot_path(ym):
    year, month = ym.split("-")
    return REPORT_EXTRAS_DIR / year / f"report_extras_{int(month):02d}.json"


def _read_report_extras_snapshot(ym):
    path = _report_extras_snapshot_path(ym)
    if not path.exists():
        return None
    try:
        data = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError):
        return None
    workbook_path = _find_report_workbook_for_ym(ym)
    if workbook_path and data.get("source") == "daily_report_workbook":
        if data.get("sourceFile") != workbook_path.name:
            return None
        source_mtime = data.get("sourceMtime")
        try:
            if source_mtime is not None and workbook_path.stat().st_mtime > float(source_mtime) + 0.5:
                return None
        except (OSError, TypeError, ValueError):
            return None
    return data


def _write_report_extras_snapshot(ym, extras):
    year, month = [int(part) for part in ym.split("-")]
    if not _is_snapshot_writable_ym(year, month):
        return
    try:
        path = _report_extras_snapshot_path(ym)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(
            json.dumps(extras, ensure_ascii=False, separators=(",", ":")),
            encoding="utf-8",
        )
    except (OSError, TypeError):
        pass


def _to_float(value):
    if value is None or value == "":
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _station_key_from_cell(value):
    try:
        station_key = int(float(value))
    except (TypeError, ValueError):
        return None
    return station_key if station_key in SHOP_KEY_MAP else None


def _workbook_report_ym(path):
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        if "总表" not in workbook.sheetnames:
            return None
        title = str(workbook["总表"].cell(1, 1).value or "").lower()
        year_match = re.search(r"(20\d{2})", title)
        month = next(
            (value for name, value in REPORT_MONTH_NAME_TO_NUMBER.items() if name in title),
            None,
        )
        if year_match and month:
            return f"{int(year_match.group(1))}-{month:02d}"
    except Exception:
        return None
    finally:
        if workbook is not None:
            workbook.close()
    return None


def _find_report_workbook_for_ym(ym):
    export_dir = Path(EXPORT_DIR)
    if not export_dir.exists():
        return None
    for path in sorted(export_dir.glob("*.xlsx"), key=lambda item: item.stat().st_mtime, reverse=True):
        if _workbook_report_ym(path) == ym:
            return path
    return None


def _report_extras_from_workbook(ym):
    workbook_path = _find_report_workbook_for_ym(ym)
    if not workbook_path:
        return None
    workbook = None
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        if "总表" not in workbook.sheetnames:
            return None
        sheet = workbook["总表"]
        result = _empty_report_extras()
        current_shop_key = None
        for row in sheet.iter_rows(min_row=5, values_only=True):
            station_key = _station_key_from_cell(row[0] if len(row) > 0 else None)
            if station_key is not None:
                current_shop_key = station_key
            if current_shop_key is None:
                continue

            product_name = str(row[3] if len(row) > 3 and row[3] is not None else "").strip()
            if not product_name or "合计" in product_name or "итого" in product_name.lower():
                continue

            density = _to_float(row[4] if len(row) > 4 else None)
            contract_volume = _to_float(row[16] if len(row) > 16 else None)
            contract_amount = _to_float(row[18] if len(row) > 18 else None)
            end_inventory = _to_float(row[23] if len(row) > 23 else None)
            end_inventory_tons = _to_float(row[24] if len(row) > 24 else None)
            if not end_inventory_tons and end_inventory and density:
                end_inventory_tons = end_inventory * density / 1000

            if not any((contract_volume, contract_amount, end_inventory, end_inventory_tons)):
                continue

            item = result["stations"].setdefault(current_shop_key, {
                "endInventory": 0.0,
                "endInventoryTons": 0.0,
                "contractVolume": 0.0,
                "contractAmount": 0.0,
            })
            item["endInventory"] += end_inventory
            item["endInventoryTons"] += end_inventory_tons
            item["contractVolume"] += contract_volume
            item["contractAmount"] += contract_amount
            result["totalEndInventory"] += end_inventory
            result["totalEndInventoryTons"] += end_inventory_tons
            result["totalContractVolume"] += contract_volume
            result["totalContractAmount"] += contract_amount

        result = _finalize_report_extras(result)
        result["source"] = "daily_report_workbook"
        result["sourceFile"] = workbook_path.name
        result["sourceMtime"] = workbook_path.stat().st_mtime
        return result
    except Exception:
        return None
    finally:
        if workbook is not None:
            workbook.close()


def _merge_report_extras(target, source):
    target["totalEndInventory"] = float(source.get("totalEndInventory") or 0)
    target["totalEndInventoryTons"] = float(source.get("totalEndInventoryTons") or 0)
    target["totalContractVolume"] += float(source.get("totalContractVolume") or 0)
    target["totalContractAmount"] += float(source.get("totalContractAmount") or 0)
    for raw_shop_key, source_item in (source.get("stations") or {}).items():
        shop_key = int(raw_shop_key)
        item = target["stations"].setdefault(shop_key, {
            "endInventory": 0.0,
            "endInventoryTons": 0.0,
            "contractVolume": 0.0,
            "contractAmount": 0.0,
        })
        item["endInventory"] = float(source_item.get("endInventory") or 0)
        item["endInventoryTons"] = float(source_item.get("endInventoryTons") or 0)
        item["contractVolume"] += float(source_item.get("contractVolume") or 0)
        item["contractAmount"] += float(source_item.get("contractAmount") or 0)


def _finalize_report_extras(result):
    for item in result["stations"].values():
        item["endInventory"] = _round_half_up(item["endInventory"], 2)
        item["endInventoryTons"] = _round_half_up(item["endInventoryTons"], 3)
        item["contractVolume"] = _round_half_up(item["contractVolume"], 2)
        item["contractAmount"] = _round_half_up(item["contractAmount"], 2)
    result["totalEndInventory"] = _round_half_up(result["totalEndInventory"], 2)
    result["totalEndInventoryTons"] = _round_half_up(result["totalEndInventoryTons"], 3)
    result["totalContractVolume"] = _round_half_up(result["totalContractVolume"], 2)
    result["totalContractAmount"] = _round_half_up(result["totalContractAmount"], 2)
    return result


def _month_slices(start, end):
    start_day = datetime.strptime(start, "%Y-%m-%d").date()
    end_day = datetime.strptime(end, "%Y-%m-%d").date()
    cursor = start_day
    while cursor < end_day:
        next_year, next_month = _month_after(cursor.year, cursor.month)
        next_month_start = datetime(next_year, next_month, 1).date()
        slice_end = min(next_month_start, end_day)
        yield cursor.isoformat(), slice_end.isoformat()
        cursor = slice_end


def _inventory_contract_totals(start, end):
    """Use the Excel report data source so analysis reports match generated reports."""
    if not start or not end:
        return _empty_report_extras()
    ym = _ym_for_full_month(start, end)
    if ym:
        cached = _read_report_extras_snapshot(ym)
        if cached is not None:
            return cached
        workbook_result = _report_extras_from_workbook(ym)
        if workbook_result is not None:
            _write_report_extras_snapshot(ym, workbook_result)
            return workbook_result
    else:
        cache_key = f"report_extras:{start}:{end}"
        cached = get_cached(cache_key)
        if cached is not None:
            return cached
    station_keys = sorted(SHOP_KEY_MAP.keys())
    result = _empty_report_extras()
    conn = get_connection()
    try:
        cur = conn.cursor()
        _, _, _, contract_data, _, end_balance, _ = _build_sheet_data(cur, start, end, station_keys)
    finally:
        conn.close()

    for (shop_key, resource_key), volume in (end_balance or {}).items():
        item = result["stations"].setdefault(int(shop_key), {
            "endInventory": 0.0,
            "endInventoryTons": 0.0,
            "contractVolume": 0.0,
            "contractAmount": 0.0,
        })
        liters = float(volume or 0)
        mass = float(getattr(end_balance, "mass", {}).get((shop_key, resource_key), 0) or 0)
        if mass:
            tons = mass / 1000
        else:
            density = RESOURCE_DENSITY.get(int(resource_key or 0), 0)
            tons = liters * density / 1000
        item["endInventory"] += liters
        item["endInventoryTons"] += tons
        result["totalEndInventory"] += liters
        result["totalEndInventoryTons"] += tons

    for shop_key, by_resource in (contract_data or {}).items():
        item = result["stations"].setdefault(int(shop_key), {
            "endInventory": 0.0,
            "endInventoryTons": 0.0,
            "contractVolume": 0.0,
            "contractAmount": 0.0,
        })
        for contract_item in by_resource.values():
            volume = float(contract_item.get("vol") or 0)
            amount = float(contract_item.get("amt") or 0)
            item["contractVolume"] += volume
            item["contractAmount"] += amount
            result["totalContractVolume"] += volume
            result["totalContractAmount"] += amount

    result = _finalize_report_extras(result)
    if ym:
        _write_report_extras_snapshot(ym, result)
    else:
        set_cached(f"report_extras:{start}:{end}", result, ttl=300)
    return result


def _inventory_contract_totals_for_period(start, end):
    slices = list(_month_slices(start, end))
    if len(slices) <= 1:
        return _inventory_contract_totals(start, end)
    result = _empty_report_extras()
    for slice_start, slice_end in slices:
        _merge_report_extras(result, _inventory_contract_totals(slice_start, slice_end))
    return _finalize_report_extras(result)


def _merge_period_dataset(target, part):
    summary = part.get("summary") or {}
    target_summary = target["summary"]
    target_summary["txnCount"] += int(summary.get("txnCount") or summary.get("totalTxn") or 0)
    target_summary["totalLiters"] += float(summary.get("totalLiters") or 0)
    target_summary["totalSales"] += float(summary.get("totalSales") or 0)
    target_summary["cashSales"] += float(summary.get("cashSales") or 0)
    target_summary["cardSales"] += float(summary.get("cardSales") or 0)
    target_summary["totalTons"] += float(summary.get("totalTons") or 0)

    for station in part.get("stations") or []:
        shop_key = station.get("shopKey")
        if shop_key is None:
            continue
        shop_key = int(shop_key)
        info = SHOP_KEY_MAP.get(shop_key, {})
        bucket = target["stations"].setdefault(shop_key, {
            "shopKey": shop_key,
            "name": station.get("name") or info.get("name") or f"Station-{shop_key}",
            "short": station.get("short") or info.get("short") or "",
            "region": station.get("region") or info.get("region") or "",
            "txnCount": 0,
            "volume": 0.0,
            "amount": 0.0,
            "cashSales": 0.0,
            "cardSales": 0.0,
            "bonus": 0.0,
        })
        bucket["txnCount"] += int(station.get("txnCount") or 0)
        bucket["volume"] += float(station.get("volume") or 0)
        bucket["amount"] += float(station.get("amount") or 0)
        bucket["cashSales"] += float(station.get("cashSales") or station.get("cash") or 0)
        bucket["cardSales"] += float(station.get("cardSales") or station.get("card") or 0)
        bucket["bonus"] += float(station.get("bonus") or station.get("bonusCost") or 0)

    for resource in part.get("resources") or part.get("byResource") or []:
        key = int(resource.get("resourceKey") or resource.get("key") or 0)
        if not key:
            continue
        bucket = target["resources"].setdefault(key, {
            "resourceKey": key,
            "key": key,
            "name": resource.get("name") or f"Other-{key}",
            "volume": 0.0,
            "amount": 0.0,
            "color": resource.get("color") or "#999999",
        })
        bucket["volume"] += float(resource.get("volume") or 0)
        bucket["amount"] += float(resource.get("amount") or 0)


def _finalize_period_dataset(aggregate, start, end):
    stations = []
    for station in aggregate["stations"].values():
        for key in ("volume", "amount", "cashSales", "cardSales", "bonus"):
            station[key] = _round_half_up(station[key], 2)
        stations.append(station)
    stations.sort(key=lambda item: item.get("amount") or 0, reverse=True)

    regions = {}
    for station in stations:
        region = station.get("region") or "--"
        bucket = regions.setdefault(region, {
            "region": region,
            "volume": 0.0,
            "amount": 0.0,
            "cash": 0.0,
            "card": 0.0,
            "bonus": 0.0,
            "stationCount": 0,
        })
        bucket["volume"] += station.get("volume") or 0
        bucket["amount"] += station.get("amount") or 0
        bucket["cash"] += station.get("cashSales") or 0
        bucket["card"] += station.get("cardSales") or 0
        bucket["bonus"] += station.get("bonus") or 0
        bucket["stationCount"] += 1

    resources = []
    for resource in aggregate["resources"].values():
        resource["volume"] = _round_half_up(resource["volume"], 2)
        resource["amount"] = _round_half_up(resource["amount"], 2)
        resources.append(resource)
    resources.sort(key=lambda item: item.get("volume") or 0, reverse=True)

    summary = aggregate["summary"]
    summary["activeShops"] = len([item for item in stations if (item.get("amount") or 0) > 0])
    summary["totalLiters"] = _round_half_up(summary["totalLiters"], 2)
    summary["totalSales"] = _round_half_up(summary["totalSales"], 2)
    summary["cashSales"] = _round_half_up(summary["cashSales"], 2)
    summary["cardSales"] = _round_half_up(summary["cardSales"], 2)
    summary["totalTons"] = _round_half_up(summary["totalTons"], 3)

    return {
        "queryStart": start,
        "queryEndExclusive": end,
        "summary": summary,
        "resources": resources,
        "stations": stations,
        "regions": sorted(regions.values(), key=lambda item: item.get("amount") or 0, reverse=True),
        "daily": [],
    }


def _period_dataset_from_month_slices(start, end, include_monthly=False):
    aggregate = {
        "summary": {
            "txnCount": 0,
            "activeShops": 0,
            "totalLiters": 0.0,
            "totalSales": 0.0,
            "cashSales": 0.0,
            "cardSales": 0.0,
            "totalTons": 0.0,
        },
        "stations": {},
        "resources": {},
    }
    monthly = []
    for slice_start, slice_end in _month_slices(start, end):
        ym = _ym_for_full_month(slice_start, slice_end)
        if ym:
            part = _month_dataset(ym)
        else:
            cache_key = f"period:{slice_start}:{slice_end}"
            part = get_cached(cache_key)
            if part is None:
                conn = get_connection()
                try:
                    cur = conn.cursor()
                    part = _query_period_dataset(cur, slice_start, slice_end)
                finally:
                    conn.close()
                set_cached(cache_key, part, ttl=300)
        _merge_period_dataset(aggregate, part)
        if include_monthly:
            summary = part.get("summary") or {}
            monthly.append({
                "month": int(slice_start[5:7]),
                "volume": _round_half_up(float(summary.get("totalLiters") or 0), 2),
                "amount": _round_half_up(float(summary.get("totalSales") or 0), 2),
                "tons": _round_half_up(float(summary.get("totalTons") or 0), 3),
            })
    result = _finalize_period_dataset(aggregate, start, end)
    if include_monthly:
        result["monthly"] = monthly
    return result


def _attach_report_extras(data):
    extras = _inventory_contract_totals_for_period(data.get("queryStart"), data.get("queryEndExclusive"))
    data["reportExtras"] = extras
    summary = data.setdefault("summary", {})
    summary["endInventory"] = extras["totalEndInventory"]
    summary["endInventoryTons"] = extras["totalEndInventoryTons"]
    summary["contractVolume"] = extras["totalContractVolume"]
    summary["contractAmount"] = extras["totalContractAmount"]
    by_shop = {int(shop_key): value for shop_key, value in (extras.get("stations") or {}).items()}
    rows = list(data.get("stations") or [])
    seen = set()
    for row in rows:
        shop_key = row.get("shopKey")
        if shop_key is None:
            continue
        shop_key = int(shop_key)
        seen.add(shop_key)
        row.update(by_shop.get(shop_key, {}))
    for shop_key, extra in by_shop.items():
        if shop_key in seen:
            continue
        info = SHOP_KEY_MAP.get(shop_key, {})
        rows.append({
            "shopKey": shop_key,
            "name": info.get("name") or f"Station-{shop_key}",
            "short": info.get("short") or "",
            "region": info.get("region") or "",
            "volume": 0,
            "amount": 0,
            **extra,
        })
    data["stations"] = rows
    return data


def _query_period_dataset(cur, start, end):
    period_condition = business_period_condition_with_pruning("s", "sh")
    total_volume = report_volume_expr("s")
    total_amount = report_amount_expr("s")
    cash_amount = report_cash_amount_expr("s")
    card_amount = report_card_amount_expr("s")
    bonus_amount = report_bonus_amount_expr("s")
    resource_key = report_resource_key_expr("s")
    business_day = business_date_expr("s", "sh")

    cur.execute(f"""
        SELECT
            COUNT(*) AS txn_count,
            COUNT(DISTINCT s.ShopKey) AS active_shops,
            {total_volume} AS total_liters,
            {total_amount} AS total_sales,
            SUM(CASE WHEN (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 3, 4)) AND s.BasePaymentTypeKey = 1 AND COALESCE(s.DerivedPaymentTypeKey, 0) <> 5 THEN s.ShopCost ELSE 0 END) AS cash_sales,
            SUM(CASE WHEN (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 3, 4)) AND (s.BasePaymentTypeKey = 2 OR s.DerivedPaymentTypeKey = 5) THEN s.ShopCost ELSE 0 END) AS card_sales
        FROM selling s
        {shift_join("s", "sh")}
        WHERE {period_condition}
          AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
    """, business_period_params(start, end))
    row = cur.fetchone()
    summary = {
        "txnCount": int(row[0] or 0),
        "activeShops": int(row[1] or 0),
        "totalLiters": _round_half_up(float(row[2] or 0), 2),
        "totalSales": _round_half_up(float(row[3] or 0), 2),
        "cashSales": _round_half_up(float(row[4] or 0), 2),
        "cardSales": _round_half_up(float(row[5] or 0), 2),
    }

    cur.execute(f"""
        SELECT {resource_key} AS resource_key, {total_volume} AS total_liters, {total_amount} AS total_sales
        FROM selling s
        {shift_join("s", "sh")}
        WHERE {period_condition}
          AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
        GROUP BY resource_key
    """, business_period_params(start, end))
    resources = format_resource_rows(cur.fetchall())

    cur.execute(f"""
        SELECT
            s.ShopKey,
            {total_volume} AS total_liters,
            {total_amount} AS total_sales,
            {cash_amount} AS cash_sales,
            {card_amount} AS card_sales,
            SUM(CASE WHEN s.DerivedPaymentTypeKey = 3 THEN s.Volume ELSE 0 END) AS contract_volume,
            SUM(CASE WHEN s.DerivedPaymentTypeKey = 3 THEN s.ShopCost ELSE 0 END) AS contract_amount,
            {bonus_amount} AS bonus_cost
        FROM selling s
        {shift_join("s", "sh")}
        WHERE {period_condition}
          AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
        GROUP BY s.ShopKey
        ORDER BY total_sales DESC
    """, business_period_params(start, end))
    stations = []
    for shop_key, volume, amount, cash, card, contract_volume, contract_amount, bonus in cur.fetchall():
        info = SHOP_KEY_MAP.get(shop_key, {"name": f"Station-{shop_key}", "short": "", "region": "unknown"})
        stations.append({
            "shopKey": int(shop_key),
            "name": info.get("name") or info.get("short") or f"Station-{shop_key}",
            "short": info.get("short") or "",
            "region": info.get("region") or "unknown",
            "volume": _round_half_up(float(volume or 0), 2),
            "amount": _round_half_up(float(amount or 0), 2),
            "cashSales": _round_half_up(float(cash or 0), 2),
            "cardSales": _round_half_up(float(card or 0), 2),
            "bonus": _round_half_up(float(bonus or 0), 2),
            "contractVolume": _round_half_up(float(contract_volume or 0), 2),
            "contractAmount": _round_half_up(float(contract_amount or 0), 2),
        })

    regions = {}
    for station in stations:
        region = station["region"]
        bucket = regions.setdefault(region, {
            "region": region,
            "volume": 0.0,
            "amount": 0.0,
            "cash": 0.0,
            "card": 0.0,
            "bonus": 0.0,
            "stationCount": 0,
        })
        bucket["volume"] += station["volume"]
        bucket["amount"] += station["amount"]
        bucket["cash"] += station.get("cashSales", 0)
        bucket["card"] += station.get("cardSales", 0)
        bucket["bonus"] += station.get("bonus", 0)
        bucket["stationCount"] += 1
    region_rows = sorted(regions.values(), key=lambda item: item["amount"], reverse=True)

    cur.execute(f"""
        SELECT {business_day} AS day, {total_volume} AS total_liters, {total_amount} AS total_sales
        FROM selling s
        {shift_join("s", "sh")}
        WHERE {period_condition}
          AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
        GROUP BY day
        ORDER BY day
    """, business_period_params(start, end))
    daily = [
        {"date": str(day), "volume": _round_half_up(float(volume or 0), 2), "amount": _round_half_up(float(amount or 0), 2)}
        for day, volume, amount in cur.fetchall()
    ]

    summary["totalTons"] = _resource_tons(resources)
    return {
        "summary": summary,
        "resources": resources,
        "stations": stations,
        "regions": region_rows,
        "daily": daily,
    }


def _same_period_previous_year(start, end):
    start_day = datetime.strptime(start, "%Y-%m-%d").date()
    end_day = datetime.strptime(end, "%Y-%m-%d").date()
    try:
        prev_start = start_day.replace(year=start_day.year - 1)
    except ValueError:
        prev_start = start_day.replace(year=start_day.year - 1, day=28)
    try:
        prev_end = end_day.replace(year=end_day.year - 1)
    except ValueError:
        prev_end = end_day.replace(year=end_day.year - 1, day=28)
    return prev_start.isoformat(), prev_end.isoformat()


def _comparison_payload(current, previous, label):
    return {
        "yoyLabel": label,
        "yoySales": _pct(current.get("totalSales"), previous.get("totalSales")),
        "yoyVolume": _pct(current.get("totalLiters"), previous.get("totalLiters")),
        "yoyCash": _pct(current.get("cashSales"), previous.get("cashSales")),
        "yoyCard": _pct(current.get("cardSales"), previous.get("cardSales")),
    }


def _ensure_report_comparison(data):
    comparison = dict(data.get("comparison") or {})
    if data.get("kind") == "monthly":
        required = {"momSales", "momVolume", "momCash", "momCard", "yoySales", "yoyVolume", "yoyCash", "yoyCard"}
        if required.issubset(comparison):
            return data
        year = int(data.get("year"))
        month = int(data.get("month"))
        prev_year, prev_month = _calc_prev_month(year, month)
        prev_label = f"{prev_year}-{prev_month:02d}"
        yoy_label = f"{year - 1}-{month:02d}"
        prev_start, prev_end = get_query_date_range(prev_label, exclude_today=False)
        yoy_start, yoy_end = get_query_date_range(yoy_label, exclude_today=False)
        conn = get_connection()
        try:
            cur = conn.cursor()
            prev = _query_period_dataset(cur, prev_start, prev_end)["summary"]
            yoy = _query_period_dataset(cur, yoy_start, yoy_end)["summary"]
        finally:
            conn.close()
        summary = data.get("summary") or {}
        comparison.update({
            "momSales": _pct(summary.get("totalSales"), prev.get("totalSales")),
            "momVolume": _pct(summary.get("totalLiters"), prev.get("totalLiters")),
            "momCash": _pct(summary.get("cashSales"), prev.get("cashSales")),
            "momCard": _pct(summary.get("cardSales"), prev.get("cardSales")),
            "yoySales": _pct(summary.get("totalSales"), yoy.get("totalSales")),
            "yoyVolume": _pct(summary.get("totalLiters"), yoy.get("totalLiters")),
            "yoyCash": _pct(summary.get("cashSales"), yoy.get("cashSales")),
            "yoyCard": _pct(summary.get("cardSales"), yoy.get("cardSales")),
            "prevLabel": comparison.get("prevLabel") or prev_label,
            "yoyLabel": comparison.get("yoyLabel") or yoy_label,
        })
    elif data.get("kind") == "yearly":
        if {"yoySales", "yoyVolume", "yoyCash", "yoyCard"}.issubset(comparison):
            return data
        prev_start, prev_end = _same_period_previous_year(data.get("queryStart"), data.get("queryEndExclusive"))
        previous = _period_dataset_from_month_slices(prev_start, prev_end).get("summary") or {}
        end_label = _period_label({"queryStart": prev_start, "queryEndExclusive": prev_end})
        comparison.update(_comparison_payload(data.get("summary") or {}, previous, end_label))
    data["comparison"] = comparison
    return data


def _month_exclusive_end(year, month):
    today = date.today()
    if year == today.year and month == today.month:
        return today.isoformat()
    if month == 12:
        return f"{year + 1}-01-01"
    return f"{year}-{month + 1:02d}-01"


def _month_dataset(year_month):
    year, month = [int(part) for part in year_month.split("-")]
    cached = get_cached(f"monthly:{year_month}")
    if cached:
        summary = dict(cached.get("summary") or {})
        resources = cached.get("byResource") or []
        summary["txnCount"] = summary.get("totalTxn", summary.get("txnCount", 0))
        summary["activeShops"] = summary.get("activeShops", 0)
        summary["totalLiters"] = summary.get("totalLiters", 0)
        summary["totalSales"] = summary.get("totalSales", 0)
        summary["totalTons"] = _resource_tons(resources)
        regions_raw = cached.get("regions") or {}
        regions = []
        if isinstance(regions_raw, dict):
            for name, item in regions_raw.items():
                region_item = dict(item)
                region_item.setdefault("region", name)
                regions.append(region_item)
        else:
            regions = list(regions_raw)
        regions.sort(key=lambda item: item.get("amount", 0), reverse=True)
        comparison_raw = cached.get("comparison") or {}
        mom = comparison_raw.get("mom") or {}
        yoy = comparison_raw.get("yoy") or {}
        return {
            "kind": "monthly",
            "year": year,
            "month": month,
            "title": f"加油站{year}年{month}月份分析报告",
            "periodLabel": f"{year}年{month}月",
            "queryStart": summary.get("queryStart", f"{year_month}-01"),
            "queryEndExclusive": _month_exclusive_end(year, month),
            "summary": summary,
            "resources": resources,
            "stations": cached.get("stations") or [],
            "regions": regions,
            "daily": cached.get("daily") or [],
            "comparison": {
                "momSales": mom.get("sales"),
                "momVolume": mom.get("volume"),
                "momCash": mom.get("cashSales"),
                "momCard": mom.get("cardSales"),
                "yoySales": yoy.get("sales"),
                "yoyVolume": yoy.get("volume"),
                "yoyCash": yoy.get("cashSales"),
                "yoyCard": yoy.get("cardSales"),
                "prevLabel": mom.get("prevMonth", ""),
                "yoyLabel": yoy.get("yoyMonth", ""),
            },
        }

    start, end = get_query_date_range(year_month)
    prev_year, prev_month = _calc_prev_month(year, month)
    prev_label = f"{prev_year}-{prev_month:02d}"
    yoy_label = f"{year - 1}-{month:02d}"
    prev_start, prev_end = get_query_date_range(prev_label, exclude_today=False)
    yoy_start, yoy_end = get_query_date_range(yoy_label, exclude_today=False)

    conn = get_connection()
    try:
        cur = conn.cursor()
        data = _query_period_dataset(cur, start, end)
        prev = _query_period_dataset(cur, prev_start, prev_end)["summary"]
        yoy = _query_period_dataset(cur, yoy_start, yoy_end)["summary"]
        data.update({
            "kind": "monthly",
            "year": year,
            "month": month,
            "title": f"加油站{year}年{month}月份分析报告",
            "periodLabel": f"{year}年{month}月",
            "queryStart": start,
            "queryEndExclusive": end,
            "comparison": {
                "momSales": _pct(data["summary"]["totalSales"], prev["totalSales"]),
                "momVolume": _pct(data["summary"]["totalLiters"], prev["totalLiters"]),
                "momCash": _pct(data["summary"]["cashSales"], prev["cashSales"]),
                "momCard": _pct(data["summary"]["cardSales"], prev["cardSales"]),
                "yoySales": _pct(data["summary"]["totalSales"], yoy["totalSales"]),
                "yoyVolume": _pct(data["summary"]["totalLiters"], yoy["totalLiters"]),
                "yoyCash": _pct(data["summary"]["cashSales"], yoy["cashSales"]),
                "yoyCard": _pct(data["summary"]["cardSales"], yoy["cardSales"]),
                "prevLabel": prev_label,
                "yoyLabel": yoy_label,
            },
        })
        return data
    finally:
        conn.close()


def _year_dataset(year):
    start = f"{year}-01-01"
    current_year = datetime.now().year
    end = datetime.now().strftime("%Y-%m-%d") if year == current_year else f"{year + 1}-01-01"
    data = _period_dataset_from_month_slices(start, end, include_monthly=True)
    data.update({
        "kind": "yearly",
        "year": year,
        "month": None,
        "title": f"加油站{year}年年度分析报告",
        "periodLabel": f"{year}年",
        "queryStart": start,
        "queryEndExclusive": end,
        "comparison": {},
    })
    return data


def _table_rows(items, columns, class_name="data-table", total_row=None):
    rows = []
    for item in items:
        rows.append("<tr>" + "".join(f"<td>{html.escape(str(getter(item)))}</td>" for _, getter in columns) + "</tr>")
    if total_row is not None:
        rows.append("<tr class=\"total-row\">" + "".join(f"<td>{html.escape(str(getter(total_row)))}</td>" for _, getter in columns) + "</tr>")
    head = "".join(f"<th>{html.escape(title)}</th>" for title, _ in columns)
    return f"<table class=\"{html.escape(class_name)}\"><thead><tr>{head}</tr></thead><tbody>{''.join(rows)}</tbody></table>"


def _sum_key(items, key):
    total = 0.0
    for item in items or []:
        try:
            total += float(item.get(key) or 0)
        except (TypeError, ValueError, AttributeError):
            continue
    return total


def _complete_station_rows(stations):
    by_key = {int(item.get("shopKey")): dict(item) for item in stations if item.get("shopKey") is not None}
    rows = []
    for shop_key, info in SHOP_KEY_MAP.items():
        item = by_key.get(shop_key, {})
        rows.append({
            "shopKey": shop_key,
            "name": item.get("name") or info.get("name") or f"Station-{shop_key}",
            "short": item.get("short") or info.get("short") or "",
            "region": item.get("region") or info.get("region") or "",
            "txnCount": int(item.get("txnCount") or 0),
            "volume": float(item.get("volume") or 0),
            "amount": float(item.get("amount") or 0),
            "cashSales": float(item.get("cashSales") or item.get("cash") or 0),
            "cardSales": float(item.get("cardSales") or item.get("card") or 0),
            "bonus": float(item.get("bonus") or item.get("bonusCost") or 0),
            "contractVolume": float(item.get("contractVolume") or item.get("contract") or 0),
            "contractAmount": float(item.get("contractAmount") or 0),
            "endInventory": float(item.get("endInventory") or item.get("end") or 0),
            "endInventoryTons": float(item.get("endInventoryTons") or 0),
            "momVolume": item.get("momVolume"),
            "momSales": item.get("momSales"),
            "yoyVolume": item.get("yoyVolume"),
            "yoySales": item.get("yoySales"),
        })
    rows.sort(key=lambda item: item["amount"], reverse=True)
    return rows


def _top_name(items, value_key="amount"):
    if not items:
        return "--"
    item = max(items, key=lambda x: float(x.get(value_key) or 0))
    return item.get("name") or item.get("region") or "--"


def _money_short(value):
    value = float(value or 0)
    if value >= 100000000:
        return f"{_round_half_up(value / 100000000, 2):.2f}亿"
    if value >= 10000:
        return f"{_round_half_up(value / 10000, 2):.2f}万"
    return _fmt(value, 2)


def _volume_short(value):
    value = float(value or 0)
    if value >= 10000:
        return f"{_round_half_up(value / 10000, 2):.2f}万"
    return _fmt(value, 2)


def _percent(value, total):
    total = float(total or 0)
    if total <= 0:
        return "0.00%"
    return f"{_round_half_up(float(value or 0) / total * 100, 2):.2f}%"


def _rank_class(index):
    return "rank-1" if index == 0 else ("rank-2" if index == 1 else ("rank-3" if index == 2 else ""))


def _resource_js_rows(resources):
    rows = []
    for item in resources:
        rows.append({
            "name": item.get("name") or "--",
            "volume": _round_half_up(float(item.get("volume") or 0), 2),
            "amount": _round_half_up(float(item.get("amount") or 0), 2),
            "amountCnyAvg": _round_half_up(float(item.get("amountCnyAvg") or 0), 2),
            "color": item.get("color") or "#2e6da4",
        })
    return rows


def _region_js_rows(regions):
    rows = []
    for item in regions:
        rows.append({
            "region": item.get("region") or "--",
            "volume": _round_half_up(float(item.get("volume") or 0), 2),
            "amount": _round_half_up(float(item.get("amount") or 0), 2),
            "cash": _round_half_up(float(item.get("cashSales") or item.get("cash") or 0), 2),
            "card": _round_half_up(float(item.get("cardSales") or item.get("card") or 0), 2),
            "bonus": _round_half_up(float(item.get("bonus") or item.get("bonusCost") or 0), 2),
            "stationCount": int(item.get("stationCount") or 0),
        })
    return rows


def _station_js_rows(stations):
    rows = []
    for item in stations:
        rows.append({
            "name": item.get("name") or "--",
            "short": item.get("short") or "",
            "region": item.get("region") or "--",
            "volume": _round_half_up(float(item.get("volume") or 0), 2),
            "amount": _round_half_up(float(item.get("amount") or 0), 2),
            "cash": _round_half_up(float(item.get("cashSales") or item.get("cash") or 0), 2),
            "card": _round_half_up(float(item.get("cardSales") or item.get("card") or 0), 2),
            "bonus": _round_half_up(float(item.get("bonus") or item.get("bonusCost") or 0), 2),
            "contractVolume": _round_half_up(float(item.get("contractVolume") or item.get("contract") or 0), 2),
            "contractAmount": _round_half_up(float(item.get("contractAmount") or 0), 2),
            "endInventory": _round_half_up(float(item.get("endInventory") or item.get("end") or 0), 2),
            "endInventoryTons": _round_half_up(float(item.get("endInventoryTons") or 0), 3),
            "momVolume": item.get("momVolume"),
            "momSales": item.get("momSales"),
            "yoyVolume": item.get("yoyVolume"),
            "yoySales": item.get("yoySales"),
        })
    return rows


def _render_html_report(data):
    summary = data["summary"]
    comparison = data.get("comparison", {})
    stations = _complete_station_rows(data.get("stations") or [])
    active_stations = [item for item in stations if item["amount"] > 0 or item["volume"] > 0]
    resources = sorted(data.get("resources") or [], key=lambda item: item.get("volume", 0), reverse=True)
    regions = sorted(data.get("regions") or [], key=lambda item: item.get("amount", 0), reverse=True)
    monthly_table = ""
    period_label = _period_label(data)
    if data["kind"] == "yearly":
        monthly_items = data.get("monthly", [])
        monthly_table = _table_rows(monthly_items, [
            ("月份", lambda x: "合计" if x.get("month") == "合计" else f"{x.get('month')}月"),
            ("销量(L)", lambda x: _fmt(x.get("volume"), 2)),
            ("销量(吨)", lambda x: _fmt(x.get("tons"), 3)),
            ("销售金额(索姆)", lambda x: _fmt(x.get("amount"), 2)),
        ], total_row={
            "month": "合计",
            "volume": _sum_key(monthly_items, "volume"),
            "tons": _sum_key(monthly_items, "tons"),
            "amount": _sum_key(monthly_items, "amount"),
        })

    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M")
    cash_sales = float(summary.get("cashSales") or 0)
    card_sales = float(summary.get("cardSales") or 0)
    total_sales = float(summary.get("totalSales") or 0)
    total_liters = float(summary.get("totalLiters") or 0)
    end_inventory = float(summary.get("endInventory") or 0)
    end_inventory_tons = float(summary.get("endInventoryTons") or 0)
    contract_volume = float(summary.get("contractVolume") or 0)
    cny_conversion = data.get("cnyConversion") or {}
    has_cny_conversion = bool(cny_conversion.get("available"))
    avg_cny_kgs = cny_conversion.get("avgCnyKgs")
    total_sales_cny = summary.get("totalSalesCnyAvg")
    cny_amount_phrase = (
        f"，按本月平均 CNY/KGS {_fmt(avg_cny_kgs, 4)} 折算人民币 {_fmt(total_sales_cny, 2)} 元"
        if has_cny_conversion else ""
    )
    sales_kpi_sub = f"精确：{_fmt(summary['totalSales'], 2)} 索姆"
    if has_cny_conversion:
        sales_kpi_sub += f"<br>折人民币：{_fmt(total_sales_cny, 2)} 元"
    cny_conclusion = (
        f"人民币金额按本月平均 CNY/KGS {_fmt(avg_cny_kgs, 4)} 统一折算，便于中国管理口径复核。"
        if has_cny_conversion else "后续可继续补充汇率折人民币口径。"
    )
    top_station = _top_name(active_stations)
    top_region = _top_name(regions)
    top_resource = _top_name(resources, "volume")
    analysis = [
        f"统计期内销售总量为 {_fmt(summary['totalLiters'], 2)} 升，折算 {_fmt(summary['totalTons'], 3)} 吨；销售金额为 {_fmt(summary['totalSales'], 2)} 索姆{cny_amount_phrase}。",
        f"当前有销售记录站点 {len(active_stations)} 个，系统登记站点 {len(stations)} 个；销售金额最高站点为 {top_station}，销售贡献最高区域为 {top_region}。",
        f"油品结构中销量最高为 {top_resource}；现金支付占 {_percent(cash_sales, total_sales)}，刷卡支付占 {_percent(card_sales, total_sales)}。",
    ]
    if data["kind"] == "monthly":
        analysis.append(
            f"环比销售金额 {_fmt_pct(comparison.get('momSales'))}，环比销量 {_fmt_pct(comparison.get('momVolume'))}；"
            f"同比销售金额 {_fmt_pct(comparison.get('yoySales'))}，同比销量 {_fmt_pct(comparison.get('yoyVolume'))}。"
        )
    else:
        best_month = max(data.get("monthly") or [], key=lambda item: item.get("amount") or 0, default=None)
        if best_month:
            analysis.append(f"年度月度趋势中，{best_month.get('month')}月销售金额最高，为 {_fmt(best_month.get('amount'), 2)} 索姆。")

    resource_rows = _resource_js_rows(resources)
    region_rows = _region_js_rows(regions)
    station_rows = _station_js_rows(stations)
    station_chart_height = max(460, min(1100, len(station_rows) * 28 + 80))
    exchange_info = data.get("exchangeRates") or {}
    exchange_rows = exchange_info.get("rows") or []
    rate_line = (
        f" | 月均汇率：1 人民币 = {_fmt(exchange_info.get('avgCny'), 4)} 索姆"
        if exchange_info.get("avgCny") is not None else ""
    )
    exchange_chart = (
        '<div id="chartRate" class="chart"></div>'
        if exchange_rows else '<p class="empty-note">本统计周期暂未查询到汇率快照。</p>'
    )
    exchange_summary = (
        f"已获取 {exchange_info.get('count', 0)} / {exchange_info.get('totalDays', 0)} 天汇率；"
        f"平均 USD/KGS：{_fmt(exchange_info.get('avgUsd'), 4) if exchange_info.get('avgUsd') is not None else '--'}，"
        f"平均 CNY/KGS：{_fmt(exchange_info.get('avgCny'), 4) if exchange_info.get('avgCny') is not None else '--'}。"
    )
    if has_cny_conversion:
        exchange_summary += (
            f"本报告人民币金额按本月平均 CNY/KGS {_fmt(avg_cny_kgs, 4)} 统一折算，"
            f"公式为索姆金额 ÷ 月平均汇率；本期销售额折算人民币 {_fmt(total_sales_cny, 2)} 元。"
        )
    if exchange_info.get("missingCount"):
        exchange_summary += f"缺少 {exchange_info.get('missingCount')} 天快照，请在后台财务维护补齐。"
    if exchange_info.get("error"):
        exchange_summary = exchange_info.get("error")
    month_compare_block = ""
    year_compare_block = ""
    def compare_tile(label, current_text, mom_value=None, yoy_value=None, mom_label="环比", yoy_label="同比"):
        mom_color = "#27ae60" if (mom_value or 0) >= 0 else "#c0392b"
        yoy_color = "#27ae60" if (yoy_value or 0) >= 0 else "#c0392b"
        mom_line = f"<div class=\"compare-line\"><span>{html.escape(mom_label)}</span><strong style=\"color:{mom_color}\">{html.escape(_fmt_pct(mom_value))}</strong></div>" if mom_value is not None else ""
        yoy_line = f"<div class=\"compare-line\"><span>{html.escape(yoy_label)}</span><strong style=\"color:{yoy_color}\">{html.escape(_fmt_pct(yoy_value))}</strong></div>" if yoy_value is not None else ""
        return (
            "<div class=\"compare-tile\">"
            f"<div class=\"compare-label\">{html.escape(label)}</div>"
            f"<div class=\"compare-current\">{html.escape(current_text)}</div>"
            f"{mom_line}{yoy_line}"
            "</div>"
        )
    if data["kind"] == "monthly":
        prev_label = comparison.get("prevLabel") or "上月"
        yoy_label = comparison.get("yoyLabel") or "上年同期"
        month_compare_block = f"""
    <div class="card">
      <div class="card-title">📈 同比环比分析（{data.get('month')}月 vs {html.escape(prev_label)} / {html.escape(yoy_label)}）</div>
      <div class="compare-grid">
        {compare_tile("销售总量", _volume_short(total_liters) + "升", comparison.get("momVolume"), comparison.get("yoyVolume"))}
        {compare_tile("销售金额", _money_short(total_sales) + "索姆", comparison.get("momSales"), comparison.get("yoySales"))}
        {compare_tile("现金支付", _money_short(cash_sales) + "索姆", comparison.get("momCash"), comparison.get("yoyCash"))}
        {compare_tile("刷卡支付", _money_short(card_sales) + "索姆", comparison.get("momCard"), comparison.get("yoyCard"))}
      </div>
      <div class="analysis-note">
        <strong>📊 同比环比分析要点：</strong><br>
        本月销售总量环比 {html.escape(_fmt_pct(comparison.get('momVolume')))}、同比 {html.escape(_fmt_pct(comparison.get('yoyVolume')))}；
        销售金额环比 {html.escape(_fmt_pct(comparison.get('momSales')))}、同比 {html.escape(_fmt_pct(comparison.get('yoySales')))}。
        刷卡支付占比 {html.escape(_percent(card_sales, total_sales))}，现金支付占比 {html.escape(_percent(cash_sales, total_sales))}。
        合同销售量 {_fmt(contract_volume, 2)} 升，已计入总销售量。
      </div>
    </div>
"""
    if data["kind"] == "yearly":
        yoy_label = comparison.get("yoyLabel") or "上年同期"
        year_compare_block = f"""
    <div class="card">
      <div class="card-title">📈 年度同比分析（本年累计 vs {html.escape(yoy_label)}）</div>
      <div class="compare-grid">
        {compare_tile("销售总量", _volume_short(total_liters) + "升", None, comparison.get("yoyVolume"), yoy_label="同比")}
        {compare_tile("销售金额", _money_short(total_sales) + "索姆", None, comparison.get("yoySales"), yoy_label="同比")}
        {compare_tile("现金支付", _money_short(cash_sales) + "索姆", None, comparison.get("yoyCash"), yoy_label="同比")}
        {compare_tile("刷卡支付", _money_short(card_sales) + "索姆", None, comparison.get("yoyCard"), yoy_label="同比")}
      </div>
      <div class="analysis-note">
        <strong>📊 年度同比分析要点：</strong><br>
        本年累计销售总量同比 {html.escape(_fmt_pct(comparison.get('yoyVolume')))}，销售金额同比 {html.escape(_fmt_pct(comparison.get('yoySales')))}。
      </div>
    </div>
"""
    station_table_rows = []
    station_rows_by_volume = sorted(station_rows, key=lambda row: row.get("volume") or 0, reverse=True)
    station_total = {
        "volume": _sum_key(station_rows_by_volume, "volume"),
        "contractVolume": _sum_key(station_rows_by_volume, "contractVolume"),
        "amount": _sum_key(station_rows_by_volume, "amount"),
        "cash": _sum_key(station_rows_by_volume, "cash"),
        "card": _sum_key(station_rows_by_volume, "card"),
        "endInventory": _sum_key(station_rows_by_volume, "endInventory"),
    }
    for index, item in enumerate(station_rows_by_volume):
        station_table_rows.append(
            f"<tr><td class=\"{_rank_class(index)}\">{index + 1}</td>"
            f"<td>{html.escape(item['name'])}<br><small>{html.escape(item['short'])}</small></td>"
            f"<td>{html.escape(item['region'])}</td>"
            f"<td>{_fmt(item['volume'], 2)}</td>"
            f"<td>{_fmt(item['contractVolume'], 2) if item['contractVolume'] else '--'}</td>"
            f"<td>{_fmt(item['amount'], 2)}</td>"
            f"<td>{_fmt(item['cash'], 2)}</td>"
            f"<td>{_fmt(item['card'], 2)}</td>"
            f"<td>{_fmt(item['endInventory'], 2) if item['endInventory'] else '--'}</td></tr>"
        )
    station_table_rows.append(
        "<tr class=\"total-row\"><td>合计</td>"
        "<td>全部加油站</td><td>全部区域</td>"
        f"<td>{_fmt(station_total['volume'], 2)}</td>"
        f"<td>{_fmt(station_total['contractVolume'], 2)}</td>"
        f"<td>{_fmt(station_total['amount'], 2)}</td>"
        f"<td>{_fmt(station_total['cash'], 2)}</td>"
        f"<td>{_fmt(station_total['card'], 2)}</td>"
        f"<td>{_fmt(station_total['endInventory'], 2)}</td></tr>"
    )
    region_total = {
        "region": "合计",
        "volume": _sum_key(region_rows, "volume"),
        "amount": _sum_key(region_rows, "amount"),
        "stationCount": sum(int(item.get("stationCount") or 0) for item in region_rows),
    }
    resource_total = {
        "name": "合计",
        "volume": _sum_key(resource_rows, "volume"),
        "amount": _sum_key(resource_rows, "amount"),
        "amountCnyAvg": _sum_key(resource_rows, "amountCnyAvg"),
    }
    resource_columns = [
        ("油品", lambda x: x.get("name", "")),
        ("销量(L)", lambda x: _fmt(x.get("volume"), 2)),
        ("销售金额(索姆)", lambda x: _fmt(x.get("amount"), 2)),
    ]
    if has_cny_conversion:
        resource_columns.append(("折算人民币(元)", lambda x: _fmt(x.get("amountCnyAvg"), 2)))
    resource_columns.append(("销量占比", lambda x: _percent(x.get("volume"), summary.get("totalLiters"))))
    region_table = _table_rows(region_rows, [
        ("区域", lambda x: x.get("region", "")),
        ("销量(L)", lambda x: _fmt(x.get("volume"), 2)),
        ("销售金额(索姆)", lambda x: _fmt(x.get("amount"), 2)),
        ("金额占比", lambda x: _percent(x.get("amount"), total_sales)),
        ("站点数", lambda x: x.get("stationCount", 0)),
    ], "data-table region-table", total_row=region_total)
    resource_table = _table_rows(resource_rows, resource_columns, "data-table resource-table", total_row=resource_total)

    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>{html.escape(data['title'])}</title>
  <script src="https://cdn.jsdelivr.net/npm/echarts@5.4.3/dist/echarts.min.js"></script>
  <style>
    * {{ margin:0; padding:0; box-sizing:border-box; }}
    body {{ font-family:'Microsoft YaHei', Arial, sans-serif; background:#f0f2f5; color:#333; }}
    .header {{ background:linear-gradient(135deg,#1a3a5c 0%,#2e6da4 60%,#1e8bc3 100%); color:#fff; padding:32px 0 24px; }}
    .header-inner {{ max-width:1400px; margin:0 auto; padding:0 20px; text-align:center; }}
    .header h1 {{ font-size:26px; font-weight:700; letter-spacing:2px; }}
    .header p {{ margin-top:8px; font-size:14px; opacity:.86; }}
    .container {{ max-width:1400px; margin:24px auto; padding:0 20px; }}
    .kpi-grid {{ display:grid; grid-template-columns:repeat(4,1fr); gap:16px; margin-bottom:24px; }}
    .kpi-card {{ background:#fff; border-radius:12px; padding:20px 22px; box-shadow:0 2px 8px rgba(0,0,0,.07); border-left:4px solid #2e6da4; }}
    .kpi-label {{ font-size:13px; color:#888; margin-bottom:8px; }}
    .kpi-value {{ font-size:28px; font-weight:700; color:#1a3a5c; }}
    .kpi-sub {{ font-size:12px; color:#999; margin-top:4px; }}
    .compare-grid {{ display:grid; grid-template-columns:repeat(4,1fr); gap:16px; margin-bottom:20px; }}
    .compare-tile {{ background:#f8faff; border-radius:8px; padding:16px; text-align:center; }}
    .compare-label {{ font-size:13px; color:#888; margin-bottom:6px; }}
    .compare-value {{ font-size:22px; font-weight:700; }}
    .compare-current {{ font-size:17px; font-weight:700; color:#1a3a5c; margin-bottom:8px; }}
    .compare-line {{ display:flex; justify-content:center; gap:8px; font-size:13px; line-height:1.7; }}
    .compare-sub {{ font-size:12px; color:#aaa; margin-top:4px; }}
    .analysis-note {{ background:#f0f6ff; border-left:3px solid #2e6da4; padding:14px 18px; border-radius:0 8px 8px 0; font-size:13px; line-height:1.9; }}
    .card {{ background:#fff; border-radius:12px; padding:24px; box-shadow:0 2px 8px rgba(0,0,0,.07); margin-bottom:24px; }}
    .card-title {{ font-size:16px; font-weight:700; color:#1a3a5c; margin-bottom:18px; padding-bottom:10px; border-bottom:2px solid #e8edf3; display:flex; align-items:center; gap:8px; }}
    .card-title::before {{ content:''; display:inline-block; width:4px; height:18px; background:#2e6da4; border-radius:2px; }}
    .two-col {{ display:grid; grid-template-columns:1fr 1fr; gap:20px; margin-bottom:24px; }}
    .three-col {{ display:grid; grid-template-columns:1fr 1fr 1fr; gap:20px; margin-bottom:24px; }}
    .chart {{ width:100%; height:360px; }}
    table {{ width:100%; border-collapse:collapse; font-size:13px; table-layout:fixed; }}
    th {{ background:#1a3a5c; color:#fff; padding:10px 12px; text-align:center; font-weight:600; }}
    td {{ padding:9px 12px; border-bottom:1px solid #f0f2f5; text-align:right; vertical-align:middle; word-break:break-word; }}
    .data-table th, .data-table td {{ text-align:right; }}
    .data-table th:first-child, .data-table td:first-child {{ text-align:left; }}
    .station-table th:nth-child(-n+3), .station-table td:nth-child(-n+3) {{ text-align:left; }}
    .station-table th:nth-child(n+4), .station-table td:nth-child(n+4) {{ text-align:right; }}
    .note {{ font-size:12px; color:#8a96a8; margin-top:10px; font-style:italic; }}
    .empty-note {{ color:#7b8798; font-size:13px; line-height:1.8; }}
    tr:hover td {{ background:#f5f8ff; }}
    tr.total-row td {{ background:#edf4fb; color:#1a3a5c; font-weight:700; }}
    tr.total-row:hover td {{ background:#e7f0fa; }}
    small {{ color:#8a96a8; }}
    .rank-1 {{ color:#c0392b; font-weight:700; }}
    .rank-2 {{ color:#e67e22; font-weight:700; }}
    .rank-3 {{ color:#27ae60; font-weight:700; }}
    .insight-card {{ background:linear-gradient(135deg,#f0f6ff 0%,#f8faff 100%); border-left:3px solid #2e6da4; padding:16px 20px; border-radius:0 8px 8px 0; font-size:13px; line-height:1.8; }}
    .insight-card h4 {{ font-size:14px; color:#1a3a5c; margin-bottom:8px; }}
    @media(max-width:900px) {{ .kpi-grid,.compare-grid {{ grid-template-columns:repeat(2,1fr); }} .two-col,.three-col {{ grid-template-columns:1fr; }} }}
    @media(max-width:600px) {{ .kpi-grid,.compare-grid {{ grid-template-columns:1fr; }} }}
    @media print {{ body {{ background:#fff; }} .card,.kpi-card {{ box-shadow:none; }} }}
  </style>
</head>
<body>
  <div class="header"><div class="header-inner">
    <h1>⛽ {html.escape(data['title'])}</h1>
    <p>吉尔吉斯斯坦 | 统计周期：{html.escape(period_label)} | 共{len(stations)}个加油站 | 货币单位：索姆（KGS）{html.escape(rate_line)}</p>
  </div></div>
  <div class="container">
    <div class="kpi-grid">
      <div class="kpi-card"><div class="kpi-label">销售总量</div><div class="kpi-value">{_volume_short(summary['totalLiters'])}升</div><div class="kpi-sub">折算：{_fmt(summary['totalTons'], 3)} 吨</div></div>
      <div class="kpi-card"><div class="kpi-label">销售金额</div><div class="kpi-value">{_money_short(summary['totalSales'])}</div><div class="kpi-sub">{sales_kpi_sub}</div></div>
      <div class="kpi-card"><div class="kpi-label">期末库存</div><div class="kpi-value">{_volume_short(end_inventory)}升</div><div class="kpi-sub">精确：{_fmt(end_inventory, 2)} 升 ≈ {_fmt(end_inventory_tons, 3)} 吨</div></div>
      <div class="kpi-card"><div class="kpi-label">活跃站点</div><div class="kpi-value">{len(active_stations)} / {len(stations)}</div><div class="kpi-sub">以本期实际销售记录统计</div></div>
    </div>

    {month_compare_block}
    {year_compare_block}

    <div class="card"><div class="card-title">📈 经营概览与对比分析</div>
      <div class="three-col">
        <div class="insight-card"><h4>销售表现</h4>{html.escape(analysis[0])}</div>
        <div class="insight-card" style="border-left-color:#27ae60"><h4>区域与站点</h4>{html.escape(analysis[1])}</div>
        <div class="insight-card" style="border-left-color:#e67e22"><h4>结构变化</h4>{html.escape(' '.join(analysis[2:]))}</div>
      </div>
    </div>

    <div class="two-col">
      <div class="card"><div class="card-title">💳 支付方式分布</div><div id="chartPay" class="chart"></div></div>
      <div class="card"><div class="card-title">⛽ 油品销量占比</div><div id="chartOil" class="chart"></div></div>
    </div>
    <div class="card"><div class="card-title">🗺️ 各区域销售量对比</div><div id="chartRegion" class="chart"></div></div>
    <div class="card"><div class="card-title">🏪 加油站销售排行（完整站点）</div><div id="chartStation" class="chart" style="height:{station_chart_height}px"></div></div>
    <div class="two-col">
      <div class="card"><div class="card-title">🗺️ 各区域销售金额占比</div><div id="chartRegionAmt" class="chart"></div></div>
      <div class="card"><div class="card-title">💳 各区域支付方式分布</div><div id="chartRegionPay" class="chart"></div></div>
    </div>
    {f'<div class="card"><div class="card-title">📅 月度趋势</div>{monthly_table}</div>' if monthly_table else ''}
    <div class="card"><div class="card-title">⛽ 油品结构明细</div>{resource_table}</div>
    <div class="card"><div class="card-title">🗺️ 区域表现明细</div>{region_table}</div>
    <div class="card"><div class="card-title">📋 加油站销售数据明细（按销售量排序）</div><div style="overflow-x:auto"><table class="data-table station-table">
      <thead><tr><th>排名</th><th>加油站</th><th>区域</th><th>销售总量(升)</th><th>合同量(升)</th><th>销售金额(索姆)</th><th>现金(索姆)</th><th>刷卡(索姆)</th><th>期末库存(升)</th></tr></thead>
      <tbody>{''.join(station_table_rows)}</tbody>
    </table></div><div class="note">注：数据来源于零售报表同一口径，销售总量含合同销售；只统计已结束班组数据。</div></div>

    <div class="card"><div class="card-title">💱 汇率信息</div>
      <div class="insight-card" style="margin-bottom:16px">{html.escape(exchange_summary)}</div>
      {exchange_chart}
    </div>

    <div class="three-col">
      <div class="insight-card"><h4>✅ 经营亮点</h4>本期销售最高站点为 {html.escape(top_station)}；销售最高区域为 {html.escape(top_region)}；主力油品为 {html.escape(top_resource)}。</div>
      <div class="insight-card" style="border-left-color:#e67e22"><h4>⚠️ 关注事项</h4>低销量或无销量站点已保留在完整明细中，建议结合后台站点状态、库存和班组完成情况复核。</div>
      <div class="insight-card" style="border-left-color:#27ae60"><h4>📊 综合评价</h4>报告仅统计已结束班组数据，油品折吨按系统密度快照计算；{html.escape(cny_conclusion)}</div>
    </div>
  </div>
  <script>
    const resources = {json.dumps(resource_rows, ensure_ascii=False)};
    const regions = {json.dumps(region_rows, ensure_ascii=False)};
    const stations = {json.dumps(station_rows, ensure_ascii=False)};
    const exchangeRows = {json.dumps(exchange_rows, ensure_ascii=False)};
    const totalCash = {cash_sales};
    const totalCard = {card_sales};
    echarts.init(document.getElementById('chartPay')).setOption({{
      tooltip: {{ trigger:'item', formatter:'{{b}}: {{c}} ({{d}}%)' }},
      legend: {{ bottom:0 }},
      series:[{{ type:'pie', radius:['42%','70%'], data:[{{value:totalCash,name:'现金支付',itemStyle:{{color:'#3498db'}}}},{{value:totalCard,name:'刷卡支付',itemStyle:{{color:'#2ecc71'}}}}], label:{{formatter:'{{b}}\\n{{d}}%'}} }}]
    }});
    echarts.init(document.getElementById('chartOil')).setOption({{
      tooltip: {{ trigger:'item', formatter:'{{b}}: {{c}} L ({{d}}%)' }},
      legend: {{ bottom:0 }},
      series:[{{ type:'pie', radius:'68%', data:resources.map(x=>({{value:x.volume,name:x.name,itemStyle:{{color:x.color}}}})), label:{{formatter:'{{b}}\\n{{d}}%'}} }}]
    }});
    echarts.init(document.getElementById('chartRegion')).setOption({{
      tooltip: {{ trigger:'axis' }},
      grid: {{ left:80, right:30, top:20, bottom:40 }},
      xAxis: {{ type:'value' }},
      yAxis: {{ type:'category', data:regions.map(x=>x.region).reverse() }},
      series:[{{ type:'bar', data:regions.map(x=>x.volume).reverse(), itemStyle:{{color:'#2e6da4'}}, label:{{show:true, position:'right', formatter:p=>(p.value/10000).toFixed(1)+'万L'}} }}]
    }});
    echarts.init(document.getElementById('chartRegionAmt')).setOption({{
      tooltip: {{ trigger:'item', formatter:'{{b}}: {{c}} ({{d}}%)' }},
      legend: {{ bottom:0 }},
      series:[{{ type:'pie', radius:['42%','70%'], data:regions.map(x=>({{value:x.amount,name:x.region}})), label:{{formatter:'{{b}}\\n{{d}}%'}} }}]
    }});
    echarts.init(document.getElementById('chartRegionPay')).setOption({{
      tooltip: {{ trigger:'axis', axisPointer:{{type:'shadow'}} }},
      legend: {{ bottom:0 }},
      grid: {{ left:70, right:30, top:30, bottom:60 }},
      xAxis: {{ type:'value' }},
      yAxis: {{ type:'category', data:regions.map(x=>x.region).reverse() }},
      series:[
        {{ name:'现金', type:'bar', stack:'pay', data:regions.map(x=>x.cash||0).reverse(), itemStyle:{{color:'#3498db'}} }},
        {{ name:'刷卡', type:'bar', stack:'pay', data:regions.map(x=>x.card||0).reverse(), itemStyle:{{color:'#2ecc71'}} }}
      ]
    }});
    echarts.init(document.getElementById('chartStation')).setOption({{
      tooltip: {{ trigger:'axis' }},
      grid: {{ left:120, right:30, top:20, bottom:40 }},
      xAxis: {{ type:'value' }},
      yAxis: {{ type:'category', data:stations.map(x=>x.name).reverse() }},
      series:[{{ type:'bar', data:stations.map(x=>x.amount).reverse(), itemStyle:{{color:'#1e8bc3'}}, label:{{show:true, position:'right', formatter:p=>(p.value/10000).toFixed(1)+'万'}} }}]
    }});
    if (exchangeRows.length && document.getElementById('chartRate')) {{
      echarts.init(document.getElementById('chartRate')).setOption({{
        tooltip: {{ trigger:'axis' }},
        legend: {{ top:0 }},
        grid: {{ left:60, right:24, top:42, bottom:42 }},
        xAxis: {{ type:'category', data:exchangeRows.map(x=>x.date) }},
        yAxis: [
          {{ type:'value', name:'USD/KGS', scale:true }},
          {{ type:'value', name:'CNY/KGS', scale:true }}
        ],
        series: [
          {{ name:'USD/KGS', type:'line', smooth:true, symbolSize:6, data:exchangeRows.map(x=>x.usdKgs), itemStyle:{{color:'#1e8bc3'}}, areaStyle:{{color:'rgba(30,139,195,.10)'}}, markLine:{{silent:true, data:[{{type:'average', name:'月均'}}], lineStyle:{{color:'#1e8bc3', type:'dashed'}}, label:{{formatter:'月均 {{c}}'}}}} }},
          {{ name:'CNY/KGS', type:'line', smooth:true, symbolSize:6, yAxisIndex:1, data:exchangeRows.map(x=>x.cnyKgs), itemStyle:{{color:'#27ae60'}}, areaStyle:{{color:'rgba(39,174,96,.10)'}}, markLine:{{silent:true, data:[{{type:'average', name:'月均'}}], lineStyle:{{color:'#27ae60', type:'dashed'}}, label:{{formatter:'月均 {{c}}'}}}} }}
        ]
      }});
    }}
  </script>
</body>
</html>"""
