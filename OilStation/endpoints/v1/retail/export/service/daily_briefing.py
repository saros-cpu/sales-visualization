import json
from datetime import date, datetime, timedelta
from pathlib import Path

import pymysql
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side

from business_time import business_period_condition_with_pruning, business_period_params, shift_join
from cache_store import get_cached
from endpoints.v1.retail.analytics.service.shared import (
    report_amount_expr,
    report_mass_expr,
    report_resource_key_expr,
    report_volume_expr,
    retail_mass_from_item,
)
from endpoints.v1.retail.common.currency import sales_db_config
from export_excel import EXPORT_DIR, _round_half_up

DAILY_BRIEFING_DIR = Path(EXPORT_DIR) / "daily_briefing"
SNAPSHOT_ROOT = Path(EXPORT_DIR).resolve().parent / "cache_snapshots"
DAILY_SNAPSHOT_ROOT = SNAPSHOT_ROOT / "daily"
DAILY_BRIEFING_TITLE = "ProductTeam销售日报表"
DAILY_BRIEFING_EXCHANGE_RATE = 87.50
DAILY_BRIEFING_REMARK = (
    "备注：\n"
    "1.销量单位：吨；价格单位：美元/吨。当日统计数据为前一天实际销售数据。\n"
    "2.柴油无库存，批发未销售；加油站限量销售。待外购柴油到厂后，恢复批发销售，解禁零售限量。\n"
    "3.加油站价格：日销售价格按当日实际价格显示；同油品存在多区域、多价格，按加权均价显示。"
)
TONS_DIGITS = 3
TONS_NUMBER_FORMAT = "0.000"
PRICE_DIGITS = 2
PRICE_NUMBER_FORMAT = "0.00_ "
# OpenXML width that Excel displays as ColumnWidth 16.0.
DAILY_BRIEFING_DATA_COLUMN_WIDTH = 16.7109375
DAILY_BRIEFING_TEMPLATE_VERSION = "daily_briefing_20260815_wholesale_db_v2"

PRODUCT_ROWS = [
    (5, "95#汽油", {104}),
    (6, "92#汽油", {103}),
    (7, "80#汽油", {102}),
    (8, "柴油", {106, 109, 110}),
    (9, "催化柴油", set()),
    (10, "油浆", set()),
    (11, "液化气", {108}),
]

# 批发明细 product_id → 模板行号；柴油涵盖 K5/常柴等全部柴油产品。
WHOLESALE_PRODUCT_TO_ROW = {
    3: 5,   # 95#汽油
    2: 6,   # 92#汽油
    1: 7,   # 80#汽油
    5: 8,   # 柴油
    6: 8,
    7: 8,
    8: 9,   # 催化柴油
    9: 10,  # 油浆
    4: 11,  # 液化气
}

# 批发来源边界，与 Sales/wholesale_source_policy.py 的 detail_source_policy_clause 保持同步。
WHOLESALE_SOURCE_CUTOFF = date(2026, 6, 1)
# 批发量口径：本地/出口/待识别批发；外运、代加工、采购、内部用油不统计。
WHOLESALE_CLASSIFICATIONS = ("local_wholesale", "export_wholesale", "unknown")


def normalize_daily_briefing_date(value=None):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        return datetime.now().date()
    for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def daily_briefing_workbook_path(report_date):
    report_date = normalize_daily_briefing_date(report_date)
    if not report_date:
        return None, None
    filename = f"{DAILY_BRIEFING_TITLE}{report_date.strftime('%Y.%m.%d')}.xlsx"
    return filename, DAILY_BRIEFING_DIR / filename


def _daily_briefing_file_matches_current_template(filepath):
    try:
        wb = load_workbook(filepath, read_only=False, data_only=False)
        try:
            ws = wb["报表"] if "报表" in wb.sheetnames else wb.active
            merges = {str(item) for item in ws.merged_cells.ranges}
            return (
                wb.properties.keywords == DAILY_BRIEFING_TEMPLATE_VERSION
                and ws.max_column >= 11
                and ws["A1"].value == DAILY_BRIEFING_TITLE
                and ws["G4"].value == "零售价"
                and ws["H2"].value == f"汇率：{DAILY_BRIEFING_EXCHANGE_RATE:.2f}"
                and "A1:K1" in merges
                and "C3:G3" in merges
                and "A14:K14" in merges
                and all(
                    abs((ws.column_dimensions[column].width or 0) - DAILY_BRIEFING_DATA_COLUMN_WIDTH) < 0.01
                    for column in "CDEFGHIJK"
                )
                and all(ws[cell].font.name == "黑体" and float(ws[cell].font.sz or 0) == 12 for cell in ("D9", "E9"))
                and _top_rows_have_outline_borders(ws)
            )
        finally:
            wb.close()
    except Exception:
        return False


def _row_value(row, key, default=None):
    if isinstance(row, dict):
        return row.get(key, default)
    fields = {
        "resource_key": 0,
        "total_volume": 1,
        "total_amount": 2,
        "total_mass": 3,
    }
    index = fields.get(key)
    if index is None or index >= len(row):
        return default
    return row[index]


def _first_row_value(row, *keys, default=None):
    for key in keys:
        value = _row_value(row, key, None)
        if value is not None:
            return value
    return default


def _safe_float(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _query_retail_resource_totals(conn, start, end):
    condition = business_period_condition_with_pruning("s", "sh")
    resource_key = report_resource_key_expr("s")
    volume_expr = report_volume_expr("s")
    amount_expr = report_amount_expr("s")
    mass_expr = report_mass_expr("s", start)
    query = f"""
        SELECT
            {resource_key} AS resource_key,
            {volume_expr} AS total_volume,
            {amount_expr} AS total_amount,
            {mass_expr} AS total_mass
        FROM selling s
        {shift_join("s", "sh")}
        WHERE {condition}
          AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
        GROUP BY resource_key
    """
    with conn.cursor() as cur:
        cur.execute(query, business_period_params(start, end))
        return cur.fetchall()


def _template_totals_from_rows(rows):
    totals = _empty_template_totals()
    key_to_row = {}
    for row, _, resource_keys in PRODUCT_ROWS:
        for resource_key in resource_keys:
            key_to_row[int(resource_key)] = row

    for item in rows or []:
        try:
            resource_key = int(_first_row_value(item, "resource_key", "resourceKey", "key", default=0) or 0)
        except (TypeError, ValueError):
            continue
        target_row = key_to_row.get(resource_key)
        if not target_row:
            continue
        volume = _first_row_value(item, "total_volume", "volume", default=0)
        amount = _first_row_value(item, "total_amount", "amount", default=0)
        mass = _first_row_value(item, "total_mass", "mass", default=None)
        if mass is None:
            mass = retail_mass_from_item(item) if isinstance(item, dict) else None
        if mass is None:
            mass = _safe_float(_first_row_value(item, "tons", default=0)) * 1000
        totals[target_row]["volume"] += _safe_float(volume)
        totals[target_row]["amount"] += _safe_float(amount)
        totals[target_row]["mass"] += _safe_float(mass)

    return totals


def _empty_template_totals():
    return {row: {"volume": 0.0, "amount": 0.0, "mass": 0.0} for row, _, _ in PRODUCT_ROWS}


def _template_tons_from_totals(totals):
    return {
        row: _round_half_up(_safe_float((values or {}).get("mass")) / 1000, TONS_DIGITS)
        for row, values in (totals or {}).items()
    }


def _template_tons_from_resource_rows(rows):
    totals = _template_totals_from_rows(rows)
    return _template_tons_from_totals(totals)


def _merge_template_totals(target, source):
    for row, values in (source or {}).items():
        if row not in target:
            continue
        target[row]["volume"] += _safe_float((values or {}).get("volume"))
        target[row]["amount"] += _safe_float((values or {}).get("amount"))
        target[row]["mass"] += _safe_float((values or {}).get("mass"))
    return target


def _empty_tons():
    return {row: 0.0 for row, _, _ in PRODUCT_ROWS}


def _merge_tons(target, source):
    for row, value in (source or {}).items():
        target[row] = _round_half_up(_safe_float(target.get(row)) + _safe_float(value), TONS_DIGITS)
    return target


def _month_label(day):
    return f"{day.year}-{day.month:02d}"


def _next_month(day):
    if day.month == 12:
        return date(day.year + 1, 1, 1)
    return date(day.year, day.month + 1, 1)


def _iter_days(start, end):
    cursor = start
    while cursor < end:
        yield cursor
        cursor += timedelta(days=1)


def _load_daily_snapshot(day):
    path = DAILY_SNAPSHOT_ROOT / f"{day.year}" / f"{day.month:02d}" / f"day_{day.day:02d}.json"
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError):
        return None


def _daily_snapshot_totals(day):
    payload = _load_daily_snapshot(day)
    if not isinstance(payload, dict):
        return None
    rows = payload.get("byResource")
    if not isinstance(rows, list):
        return None
    return _template_totals_from_rows(rows)


def _daily_snapshot_tons(day):
    totals = _daily_snapshot_totals(day)
    if totals is None:
        return None
    return _template_tons_from_totals(totals)


def _day_totals_by_cache_or_query(conn, sales_day, report_date):
    totals = _daily_snapshot_totals(sales_day)
    if totals is not None:
        return totals
    return _template_totals_from_rows(
        _query_retail_resource_totals(conn, sales_day.isoformat(), report_date.isoformat())
    )


def _period_tons_from_daily_snapshots(start, end):
    totals = _empty_tons()
    missing = []
    for day in _iter_days(start, end):
        day_tons = _daily_snapshot_tons(day)
        if day_tons is None:
            missing.append(day)
            continue
        _merge_tons(totals, day_tons)
    return totals, missing


def _monthly_snapshot_tons(year_month, expected_end=None):
    payload = get_cached(f"monthly:{year_month}")
    if not isinstance(payload, dict):
        return None
    summary = payload.get("summary") or {}
    query_end = summary.get("queryEnd") or summary.get("queryEndExclusive")
    if expected_end and query_end and str(query_end)[:10] != expected_end.isoformat():
        return None
    rows = payload.get("byResource")
    if not isinstance(rows, list):
        return None
    return _template_tons_from_resource_rows(rows)


def _yearly_snapshot_tons(year, expected_end=None):
    payload = get_cached(f"resources_yearly:{year}:year")
    if not isinstance(payload, dict):
        return None

    if expected_end:
        kpi_payload = get_cached(f"kpi_yearly:{year}:summary")
        kpi_summary = kpi_payload.get("summary") if isinstance(kpi_payload, dict) else {}
        query_end = (kpi_summary or {}).get("queryEnd") or (kpi_summary or {}).get("queryEndExclusive")
        if query_end and str(query_end)[:10] != expected_end.isoformat():
            return None

    rows = payload.get("resources") or payload.get("byResource")
    if not isinstance(rows, list):
        return None
    return _template_tons_from_resource_rows(rows)


def _retail_tons_by_template_row(conn, start, end):
    rows = _query_retail_resource_totals(conn, start, end)
    return _template_tons_from_resource_rows(rows)


def _period_tons_by_cache_or_query(conn, start, end):
    if start >= end:
        return _empty_tons()

    if start.day == 1 and end <= _next_month(start):
        month_tons = _monthly_snapshot_tons(_month_label(start), expected_end=end)
        if month_tons is not None:
            return month_tons

    day_tons, missing_days = _period_tons_from_daily_snapshots(start, end)
    if not missing_days:
        return day_tons

    if start.day == 1 and end <= _next_month(start):
        month_tons = _monthly_snapshot_tons(_month_label(start), expected_end=end)
        if month_tons is not None:
            return month_tons

    return _retail_tons_by_template_row(conn, start.isoformat(), end.isoformat())


def _year_tons_by_cache_or_query(conn, report_date):
    year_tons = _yearly_snapshot_tons(report_date.year, expected_end=report_date)
    if year_tons is not None:
        return year_tons

    totals = _empty_tons()
    cursor = date(report_date.year, 1, 1)
    while cursor < report_date:
        period_end = min(_next_month(cursor), report_date)
        _merge_tons(totals, _period_tons_by_cache_or_query(conn, cursor, period_end))
        cursor = _next_month(cursor)
    return totals


def _apply_template_layout(ws):
    ws.title = "报表"
    ws.sheet_view.showGridLines = False

    widths = {
        "A": 10.375,
        "B": 12.5,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    for column in "CDEFGHIJK":
        ws.column_dimensions[column].width = DAILY_BRIEFING_DATA_COLUMN_WIDTH

    for row in range(1, 15):
        ws.row_dimensions[row].height = 32
    ws.row_dimensions[14].height = 87

    for cells in (
        "A1:K1",
        "A2:G2",
        "H2:I2",
        "J2:K2",
        "A3:A4",
        "B3:B4",
        "C3:G3",
        "H3:I3",
        "J3:K3",
        "A13:B13",
        "C13:G13",
        "H13:I13",
        "J13:K13",
        "A14:K14",
    ):
        ws.merge_cells(cells)


def _apply_template_styles(ws):
    thin_side = Side(style="thin", color="000000")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    center = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=1, max_row=14, min_col=1, max_col=11):
        for cell in row:
            cell.font = Font(name="黑体", size=12)
            cell.alignment = center
            cell.border = thin_border

    ws["A1"].font = Font(name="黑体", size=20, bold=True)
    ws["H2"].font = Font(name="黑体", size=16, bold=True)
    ws["J2"].font = Font(name="黑体", size=16, bold=True)
    ws["H2"].alignment = Alignment(horizontal="right", vertical="center")
    ws["J2"].alignment = Alignment(horizontal="right", vertical="center")
    ws["J2"].number_format = "mm-dd-yy"

    ws["A14"].font = Font(name="方正仿宋_GB2312", size=12)
    ws["A14"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for cell in ("C3", "H3", "J3"):
        ws[cell].font = Font(name="黑体", size=12, bold=True)
        ws[cell].number_format = "mm-dd-yy"

    for row in range(5, 13):
        for column in (3, 6, 8, 9, 10, 11):
            ws.cell(row=row, column=column).number_format = TONS_NUMBER_FORMAT
        for column in (4, 5, 7):
            ws.cell(row=row, column=column).number_format = PRICE_NUMBER_FORMAT
    for cell in ("C13", "H13", "J13"):
        ws[cell].number_format = TONS_NUMBER_FORMAT
    _apply_top_rows_outline_borders(ws, thin_side)


def _outline_border(left=False, right=False, top=False, bottom=False, side=None):
    side = side or Side(style="thin", color="000000")
    return Border(
        left=side if left else Side(),
        right=side if right else Side(),
        top=side if top else Side(),
        bottom=side if bottom else Side(),
    )


def _apply_outline_border(ws, cell_range, side):
    rows = list(ws[cell_range])
    if not rows:
        return
    min_row = rows[0][0].row
    max_row = rows[-1][0].row
    min_col = rows[0][0].column
    max_col = rows[0][-1].column
    for row in rows:
        for cell in row:
            cell.border = _outline_border(
                left=cell.column == min_col,
                right=cell.column == max_col,
                top=cell.row == min_row,
                bottom=cell.row == max_row,
                side=side,
            )


def _apply_top_rows_outline_borders(ws, side):
    _apply_outline_border(ws, "A1:K2", side)


def _has_side(cell, side_name):
    return getattr(cell.border, side_name).style == "thin"


def _has_no_side(cell, side_name):
    return getattr(cell.border, side_name).style is None


def _top_rows_have_outline_borders(ws):
    return (
        _has_side(ws["A1"], "left")
        and _has_side(ws["A1"], "top")
        and _has_side(ws["K1"], "right")
        and _has_side(ws["K1"], "top")
        and _has_side(ws["A2"], "left")
        and _has_side(ws["A2"], "bottom")
        and _has_side(ws["K2"], "right")
        and _has_side(ws["K2"], "bottom")
        and _has_no_side(ws["A1"], "bottom")
        and _has_no_side(ws["K1"], "bottom")
        and _has_no_side(ws["A2"], "top")
        and _has_no_side(ws["K2"], "top")
        and _has_no_side(ws["B1"], "left")
        and _has_no_side(ws["J1"], "right")
        and _has_no_side(ws["B2"], "left")
        and _has_no_side(ws["G2"], "right")
        and _has_no_side(ws["H2"], "left")
        and _has_no_side(ws["H2"], "right")
        and _has_no_side(ws["I2"], "left")
        and _has_no_side(ws["I2"], "right")
        and _has_no_side(ws["J2"], "left")
        and _has_no_side(ws["K2"], "left")
    )


def _retail_price_usd_per_ton(values, exchange_rate=DAILY_BRIEFING_EXCHANGE_RATE):
    tons = _safe_float((values or {}).get("mass")) / 1000
    amount = _safe_float((values or {}).get("amount"))
    rate = _safe_float(exchange_rate)
    if tons <= 0 or amount <= 0 or rate <= 0:
        return None
    return _round_half_up(amount / tons / rate, PRICE_DIGITS)


def _wholesale_briefing_tons(start, end):
    """批发量按模板行汇总；来源边界与 Sales/wholesale_source_policy.py 一致。"""
    totals = _empty_tons()
    if not start or not end or start >= end:
        return totals
    placeholders = ", ".join(["%s"] * len(WHOLESALE_CLASSIFICATIONS))
    query = f"""
        SELECT d.product_id, SUM(d.weight_tons) AS total_tons
        FROM wholesale_detail_records d
        WHERE d.business_date >= %s AND d.business_date < %s
          AND d.classification IN ({placeholders})
          AND (
              (d.business_date < %s AND d.batch_id IS NULL)
              OR (d.business_date >= %s AND d.batch_id IS NOT NULL)
          )
        GROUP BY d.product_id
    """
    params = [
        start.isoformat(), end.isoformat(),
        *WHOLESALE_CLASSIFICATIONS,
        WHOLESALE_SOURCE_CUTOFF.isoformat(), WHOLESALE_SOURCE_CUTOFF.isoformat(),
    ]
    try:
        conn = pymysql.connect(**sales_db_config())
        try:
            with conn.cursor() as cur:
                cur.execute(query, params)
                rows = cur.fetchall()
        finally:
            conn.close()
    except Exception:
        return totals
    raw = {}
    for product_id, total_tons in rows or []:
        try:
            product_id = int(product_id)
        except (TypeError, ValueError):
            continue
        row_no = WHOLESALE_PRODUCT_TO_ROW.get(product_id)
        if not row_no:
            continue
        raw[row_no] = _safe_float(raw.get(row_no)) + _safe_float(total_tons)
    for row_no, value in raw.items():
        totals[row_no] = _round_half_up(value, TONS_DIGITS)
    return totals


def _load_wholesale_briefing_prices():
    """批发价维护值 {模板行: (大客户价, 散户价)}，未维护或读取失败时返回空。"""
    prices = {}
    try:
        conn = pymysql.connect(**sales_db_config())
        try:
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT product_key, large_price, retail_price FROM wholesale_briefing_prices"
                )
                rows = cur.fetchall()
        finally:
            conn.close()
    except Exception:
        return prices
    for product_key, large_price, retail_price in rows or []:
        try:
            product_key = int(product_key)
        except (TypeError, ValueError):
            continue
        prices[product_key] = (
            _safe_float(large_price) if large_price is not None else None,
            _safe_float(retail_price) if retail_price is not None else None,
        )
    return prices


def _write_daily_briefing_sheet(ws, report_date, day_totals, month_tons, year_tons,
                                wholesale_day=None, wholesale_month=None, wholesale_year=None,
                                wholesale_prices=None):
    _apply_template_layout(ws)
    day_tons = _template_tons_from_totals(day_totals)

    ws["A1"] = DAILY_BRIEFING_TITLE
    ws["H2"] = f"汇率：{DAILY_BRIEFING_EXCHANGE_RATE:.2f}"
    ws["J2"] = report_date
    ws["J2"].number_format = "mm-dd-yy"

    ws["A3"] = "序号"
    ws["B3"] = "油品"
    ws["C3"] = "日销售数据"
    ws["H3"] = "月销售数据"
    ws["J3"] = "年销售数据"
    ws["C4"] = "批发量"
    ws["D4"] = "批发价(大客户)"
    ws["E4"] = "批发价(散户)"
    ws["F4"] = "零售量"
    ws["G4"] = "零售价"
    ws["H4"] = "批发量"
    ws["I4"] = "零售量"
    ws["J4"] = "批发量"
    ws["K4"] = "零售量"

    for row, product_name, resource_keys in PRODUCT_ROWS:
        ws.cell(row=row, column=1, value=f"=ROW()-2")
        ws.cell(row=row, column=2, value=product_name)
        ws.cell(row=row, column=3, value=wholesale_day.get(row, 0))
        large_price, retail_price = wholesale_prices.get(row, (None, None))
        if large_price is not None:
            ws.cell(row=row, column=4, value=large_price)
        if retail_price is not None:
            ws.cell(row=row, column=5, value=retail_price)
        if resource_keys:
            ws.cell(row=row, column=6, value=day_tons.get(row, 0))
            daily_price = _retail_price_usd_per_ton((day_totals or {}).get(row))
            if daily_price is not None:
                ws.cell(row=row, column=7, value=daily_price)
            ws.cell(row=row, column=9, value=month_tons.get(row, 0))
            ws.cell(row=row, column=11, value=year_tons.get(row, 0))
        ws.cell(row=row, column=8, value=wholesale_month.get(row, 0))
        ws.cell(row=row, column=10, value=wholesale_year.get(row, 0))

    ws["B12"] = "小计"
    for column in ("C", "F", "H", "I", "J", "K"):
        ws[f"{column}12"] = f"=SUM({column}5:{column}11)"

    ws["A13"] = "合计"
    ws["C13"] = "=C12+F12"
    ws["H13"] = "=H12+I12"
    ws["J13"] = "=J12+K12"
    ws["A14"] = DAILY_BRIEFING_REMARK

    _apply_template_styles(ws)


def generate_daily_briefing_excel(conn, report_date_value=None, force_refresh=False):
    report_date = normalize_daily_briefing_date(report_date_value)
    if not report_date:
        raise ValueError("invalid report_date, expected YYYY-MM-DD")

    filename, filepath = daily_briefing_workbook_path(report_date)
    if filepath.exists() and not force_refresh and _daily_briefing_file_matches_current_template(filepath):
        return {
            "status": "ok",
            "file": filename,
            "path": str(filepath),
            "cached": True,
            "reportDate": report_date.isoformat(),
            "salesDate": (report_date - timedelta(days=1)).isoformat(),
        }

    DAILY_BRIEFING_DIR.mkdir(parents=True, exist_ok=True)

    sales_day = report_date - timedelta(days=1)
    month_start = report_date.replace(day=1).isoformat()
    year_start = date(report_date.year, 1, 1).isoformat()

    day_totals = _day_totals_by_cache_or_query(conn, sales_day, report_date)
    month_tons = _period_tons_by_cache_or_query(conn, report_date.replace(day=1), report_date)
    year_tons = _year_tons_by_cache_or_query(conn, report_date)

    wholesale_day = _wholesale_briefing_tons(sales_day, report_date)
    wholesale_month = _wholesale_briefing_tons(report_date.replace(day=1), report_date)
    wholesale_year = _wholesale_briefing_tons(date(report_date.year, 1, 1), report_date)
    wholesale_prices = _load_wholesale_briefing_prices()

    wb = Workbook()
    wb._fonts[0] = Font(name='Times New Roman', size=11)
    wb.properties.keywords = DAILY_BRIEFING_TEMPLATE_VERSION
    ws = wb.active
    _write_daily_briefing_sheet(ws, report_date, day_totals, month_tons, year_tons,
                                wholesale_day, wholesale_month, wholesale_year, wholesale_prices)
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except AttributeError:
        pass
    wb.save(filepath)

    return {
        "status": "ok",
        "file": filename,
        "path": str(filepath),
        "cached": False,
        "reportDate": report_date.isoformat(),
        "salesDate": sales_day.isoformat(),
        "monthStart": month_start,
        "yearStart": year_start,
    }
