# -*- coding: utf-8 -*-
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import sys
from collections import defaultdict
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Side
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[5]
_root_str = str(ROOT)
_original_sys_path = list(sys.path)
if _root_str in sys.path:
    sys.path.remove(_root_str)
sys.path.insert(0, _root_str)

try:
    from export_excel import (  # noqa: E402
        EXPORT_DIR,
        _allow_fixed_density_fallback,
        _missing_density_fallback_density,
        _selling_mass_expr,
    )
    from scripts.generate_accounting_handover_reports import (  # noqa: E402
        ShiftRange,
        as_float,
        connect,
        infer_resource as _fallback_infer_resource,
        load_shift_operator,
        round_float,
        sql_in,
    )
finally:
    sys.path[:] = _original_sys_path

ARCHIVE_ROOT = Path(os.getenv("SALES_RETAIL_254_ARCHIVE_ROOT", r"D:\DevTeam文件夹\АЗС Общее Данных Документ"))
STAGING_ROOT = Path(os.getenv("SALES_RETAIL_254_STAGING_ROOT", str(Path(EXPORT_DIR) / "retail_254_monthly")))

RU_MONTH_DIR = {
    1: "Январь",
    2: "Февраль",
    3: "Март",
    4: "Апрель",
    5: "Май",
    6: "Июнь",
    7: "Июль",
    8: "Август",
    9: "Сентябрь",
    10: "Октябрь",
    11: "Ноябрь",
    12: "Декабрь",
}

RESOURCE_LABELS = {
    102: "АИ-80",
    103: "АИ-92",
    104: "АИ-95",
    106: "ДТ",
    108: "ГАЗ",
    109: "ДТ-15",
    110: "ДТ-35",
}
RESOURCE_ORDER = {104: 10, 103: 20, 102: 30, 106: 40, 108: 50, 109: 60, 110: 70}
STATION_RSK_MAP = {
    6: {1: 103, 2: 104, 3: 102, 5: 106},
    28: {7: 102, 8: 103, 9: 104, 10: 109, 11: 110, 12: 108},
    29: {7: 103, 8: 102, 9: 106, 10: 108},
    30: {7: 103, 8: 102, 9: 106},
}
STATION_RESOURCE_CORRECTIONS = {
    (7, 104): 102,
    (15, 104): 102,
}
CASH_TALL_ROW_SHOPS = {1, 5, 6, 7, 8, 9, 10, 13, 14, 15, 16, 17, 18, 19, 20, 22, 23, 24, 25, 27, 29, 30}
BLACK_THIN = Side(style="thin", color="FF000000")
NO_SIDE = Side(style=None)


def infer_report_resource(shop_key: int, resource_key, resource_shop_key=None, price=None) -> int:
    try:
        rk = int(resource_key or 0)
    except (TypeError, ValueError):
        rk = 0
    try:
        rsk = int(resource_shop_key) if resource_shop_key not in (None, "") else None
    except (TypeError, ValueError):
        rsk = None
    mapped = STATION_RSK_MAP.get(int(shop_key), {}).get(rsk)
    if mapped:
        return mapped
    corrected = STATION_RESOURCE_CORRECTIONS.get((int(shop_key), rk))
    if corrected:
        return corrected
    return int(_fallback_infer_resource(shop_key, rk, rsk, price))


def infer_supply_resource(shop_key: int, resource_key, resource_shop_key=None, price=None) -> int:
    try:
        rk = int(resource_key or 0)
    except (TypeError, ValueError):
        rk = 0
    if int(shop_key) == 6 and rk == 4:
        return 106
    return infer_report_resource(shop_key, rk, resource_shop_key, price)

MAIN_COLS = {
    "begin": 4,
    "supply": 10,
    "realization": 15,
    "realization_net": 17,
    "cash": 21,
    "fuel_card": 25,
    "bank_card": 29,
    "vedomosti": 31,
    "talon": 36,
    "bonus": 39,
    "return_supplier": 43,
    "end_fact": 47,
    "end_calc": 52,
    "balance": 57,
}


@dataclass
class PeriodData:
    shift_range: ShiftRange
    first_shift: dict
    last_shift: dict
    main_rows: list[dict]
    sales: dict[int, dict]
    checks: dict[int, float]
    supplies: dict[int, dict]
    supply_entries: list[dict]
    tanks: list[dict]
    tank_totals: dict[int, dict]
    counters: list[dict]
    cash_rows: list[dict]
    cash_amount: float
    collection: dict
    operator: str


def q(value: float | int | Decimal | None, digits: int) -> Decimal:
    decimal_value = Decimal(str(value or 0))
    if isinstance(value, float) and decimal_value:
        epsilon = Decimal(1).scaleb(-(digits + 6))
        decimal_value += epsilon if decimal_value > 0 else -epsilon
    return decimal_value.quantize(Decimal("1." + ("0" * digits)), rounding=ROUND_HALF_UP)


def fmt(value, digits: int = 3, dash: bool = False) -> str:
    if dash:
        return "-"
    rounded = q(value, digits)
    if rounded == 0:
        rounded = abs(rounded)
    text = f"{rounded:f}"
    if "." not in text and digits:
        text += "." + ("0" * digits)
    if digits:
        whole, frac = text.split(".", 1)
        frac = (frac + ("0" * digits))[:digits]
        text = whole + "." + frac
    return text.replace(".", ",")


def density(mass: float, volume: float) -> float:
    return float(mass or 0) / float(volume or 0) if abs(float(volume or 0)) > 1e-9 else 0.0


def money(volume: float, price: float) -> float:
    return round_float(float(volume or 0) * float(price or 0), 2)


def label_row(ws, label: str, start: int = 1) -> int:
    for row in range(start, ws.max_row + 1):
        if str(ws.cell(row, 1).value or "").strip() == label:
            return row
    raise ValueError(f"label not found: {label}")


def row_contains(ws, text: str, start: int = 1) -> int:
    for row in range(start, ws.max_row + 1):
        value = str(ws.cell(row, 1).value or "")
        if text in value:
            return row
    raise ValueError(f"text not found: {text}")


def decode_db_text(value) -> str:
    if value in (None, ""):
        return ""
    text = str(value)
    try:
        fixed = text.encode("latin1").decode("cp1251")
    except UnicodeError:
        return text
    cyr = re.compile(r"[\u0400-\u04FF]")
    return fixed if len(cyr.findall(fixed)) > len(cyr.findall(text)) else text


def resource_from_text(value) -> int | None:
    text = str(value or "").upper().replace(" ", "").replace("_", "-")
    if not text:
        return None
    if "ГАЗ" in text or "ГAЗ" in text:
        return 108
    if "95" in text:
        return 104
    if "92" in text:
        return 103
    if "80" in text:
        return 102
    if "ДТ-35" in text:
        return 110
    if "ДТ-15" in text:
        return 109
    if "ДТ" in text or "Д/Т" in text:
        return 106
    return None


def parse_shop_key(path: Path) -> int | None:
    if re.fullmatch(r"\d{3}", path.stem):
        return int(path.stem)
    match = re.search(r"АЗС\s*-\s*(\d+)", path.name)
    return int(match.group(1)) if match else None


def clear_row(ws, row: int, max_col: int = 58) -> None:
    for col in range(1, max_col + 1):
        cell = ws.cell(row, col)
        if not isinstance(cell, MergedCell):
            cell.value = None


def row_dimension_state(ws, row: int) -> dict:
    dim = ws.row_dimensions[row]
    return {"height": dim.height, "hidden": dim.hidden, "outlineLevel": dim.outlineLevel}


def apply_row_dimension_state(ws, row: int, state: dict) -> None:
    dim = ws.row_dimensions[row]
    dim.height = state.get("height")
    dim.hidden = bool(state.get("hidden", False))
    dim.outlineLevel = int(state.get("outlineLevel") or 0)


def unmerge_all(ws) -> list[CellRange]:
    ranges = [CellRange(str(rng)) for rng in ws.merged_cells.ranges]
    for rng in list(ws.merged_cells.ranges):
        try:
            ws.unmerge_cells(str(rng))
        except KeyError:
            ws.merged_cells.ranges = {
                existing for existing in ws.merged_cells.ranges if str(existing) != str(rng)
            }
            for row in range(rng.min_row, rng.max_row + 1):
                for col in range(rng.min_col, rng.max_col + 1):
                    key = (row, col)
                    if isinstance(ws._cells.get(key), MergedCell):
                        del ws._cells[key]
    return ranges


def remerge_ranges(ws, ranges: list[CellRange]) -> None:
    for rng in ranges:
        ws.merge_cells(str(rng))


def insert_rows_preserving_layout(ws, row: int, amount: int) -> None:
    if amount <= 0:
        return
    max_row = ws.max_row
    merges = unmerge_all(ws)
    row_states = {idx: row_dimension_state(ws, idx) for idx in range(row, max_row + 1)}
    ws.insert_rows(row, amount)
    shifted_merges: list[CellRange] = []
    for rng in merges:
        if rng.min_row >= row:
            rng.shift(row_shift=amount)
        elif rng.max_row >= row:
            rng.max_row += amount
        shifted_merges.append(rng)
    for idx, state in sorted(row_states.items(), reverse=True):
        apply_row_dimension_state(ws, idx + amount, state)
    for idx in range(row, row + amount):
        apply_row_dimension_state(ws, idx, {"height": None, "hidden": False, "outlineLevel": 0})
    remerge_ranges(ws, shifted_merges)


def delete_rows_preserving_layout(ws, row: int, amount: int) -> None:
    if amount <= 0:
        return
    max_row = ws.max_row
    delete_end = row + amount - 1
    merges = unmerge_all(ws)
    row_states = {idx: row_dimension_state(ws, idx) for idx in range(delete_end + 1, max_row + 1)}
    ws.delete_rows(row, amount)
    shifted_merges: list[CellRange] = []
    for rng in merges:
        if rng.max_row < row:
            shifted_merges.append(rng)
        elif rng.min_row > delete_end:
            rng.shift(row_shift=-amount)
            shifted_merges.append(rng)
        elif rng.min_row < row and rng.max_row > delete_end:
            rng.max_row -= amount
            shifted_merges.append(rng)
    for idx, state in sorted(row_states.items()):
        apply_row_dimension_state(ws, idx - amount, state)
    remerge_ranges(ws, shifted_merges)


def trim_rows_after(ws, final_row: int) -> None:
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row > final_row:
            try:
                ws.unmerge_cells(str(rng))
            except KeyError:
                ws.merged_cells.ranges = {
                    existing for existing in ws.merged_cells.ranges if str(existing) != str(rng)
                }
    for key in list(ws._cells):
        if key[0] > final_row:
            del ws._cells[key]
    for idx in list(ws.row_dimensions):
        if idx > final_row:
            del ws.row_dimensions[idx]
    ws._current_row = min(ws.max_row, final_row)


def copy_cell_style(src, dst) -> None:
    if src.has_style:
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)
        dst.number_format = src.number_format


def copy_row_format(ws, src_row: int, dst_row: int, max_col: int = 58) -> None:
    src_merges = [
        (rng.min_col, rng.max_col)
        for rng in ws.merged_cells.ranges
        if rng.min_row == src_row and rng.max_row == src_row
    ]
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row <= dst_row <= rng.max_row:
            try:
                ws.unmerge_cells(str(rng))
            except KeyError:
                ws.merged_cells.ranges = {
                    existing for existing in ws.merged_cells.ranges if str(existing) != str(rng)
                }
                for row in range(rng.min_row, rng.max_row + 1):
                    for col in range(rng.min_col, rng.max_col + 1):
                        key = (row, col)
                        if isinstance(ws._cells.get(key), MergedCell):
                            del ws._cells[key]
    ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height
    for col in range(1, max_col + 1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if isinstance(dst, MergedCell):
            continue
        copy_cell_style(src, dst)
    for min_col, max_col_ in src_merges:
        ws.merge_cells(
            start_row=dst_row,
            start_column=min_col,
            end_row=dst_row,
            end_column=max_col_,
        )
        copy_cell_style(ws.cell(src_row, min_col), ws.cell(dst_row, min_col))


def ensure_capacity(ws, start_row: int, current_rows: int, needed_rows: int, template_row: int) -> int:
    if needed_rows < current_rows:
        delete_rows_preserving_layout(ws, start_row + needed_rows, current_rows - needed_rows)
        return needed_rows
    if needed_rows == current_rows:
        return current_rows
    add_count = needed_rows - current_rows
    insert_at = start_row + current_rows
    insert_rows_preserving_layout(ws, insert_at, add_count)
    for offset in range(add_count):
        copy_row_format(ws, template_row, insert_at + offset)
    return needed_rows


def copy_block_merges(ws, src_start: int, dst_start: int, row_count: int) -> None:
    src_end = src_start + row_count - 1
    row_shift = dst_start - src_start
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row < src_start or rng.max_row > src_end:
            continue
        if rng.min_row == rng.max_row:
            continue
        ws.merge_cells(
            start_row=rng.min_row + row_shift,
            start_column=rng.min_col,
            end_row=rng.max_row + row_shift,
            end_column=rng.max_col,
        )


def ensure_main_capacity(ws, start_row: int, current_rows: int, needed_rows: int, template_row: int) -> int:
    group_size = 4
    if needed_rows < current_rows:
        delete_rows_preserving_layout(ws, start_row + needed_rows, current_rows - needed_rows)
        return needed_rows
    if needed_rows == current_rows:
        return current_rows
    add_count = needed_rows - current_rows
    insert_at = start_row + current_rows
    insert_rows_preserving_layout(ws, insert_at, add_count)
    for offset in range(add_count):
        copy_row_format(ws, template_row + (offset % group_size), insert_at + offset)
    for offset in range(0, add_count, group_size):
        if offset + group_size <= add_count:
            copy_block_merges(ws, template_row, insert_at + offset, group_size)
    return needed_rows


def set_cell(ws, row: int, col: int, value) -> None:
    cell = ws.cell(row, col)
    if not isinstance(cell, MergedCell):
        cell.value = value


def set_horizontal_alignment(ws, row: int, cols: tuple[int, ...], horizontal: str) -> None:
    for col in cols:
        cell = ws.cell(row, col)
        if isinstance(cell, MergedCell):
            continue
        alignment = cell.alignment
        cell.alignment = Alignment(
            horizontal=horizontal,
            vertical=alignment.vertical,
            text_rotation=alignment.textRotation,
            wrap_text=alignment.wrap_text,
            shrink_to_fit=alignment.shrink_to_fit,
        )


def apply_supply_row_alignment(ws, row: int) -> None:
    set_horizontal_alignment(ws, row, (1, 8, 26), "left")
    set_horizontal_alignment(ws, row, (33, 40, 48, 54), "right")


def apply_supply_row_borders(ws, row: int, resource_start: bool) -> None:
    ws.cell(row, 1).border = Border(
        left=BLACK_THIN,
        right=BLACK_THIN,
        top=BLACK_THIN if resource_start else NO_SIDE,
        bottom=NO_SIDE,
    )
    ws.cell(row, 8).border = Border(left=BLACK_THIN, right=BLACK_THIN, top=NO_SIDE, bottom=NO_SIDE)


def normalize_cash_signature_borders(ws, cash_label: int, chairman_label: int) -> None:
    ws.cell(cash_label, 1).border = Border(left=NO_SIDE, right=NO_SIDE, top=NO_SIDE, bottom=NO_SIDE)
    for row in range(cash_label + 2, max(cash_label + 2, chairman_label - 1)):
        ws.cell(row, 2).border = Border(left=BLACK_THIN, right=NO_SIDE, top=BLACK_THIN, bottom=NO_SIDE)
    for row in range(chairman_label, chairman_label + 4):
        ws.cell(row, 1).border = Border(left=NO_SIDE, right=NO_SIDE, top=NO_SIDE, bottom=NO_SIDE)
        ws.cell(row, 11).border = Border(
            left=NO_SIDE,
            right=NO_SIDE,
            top=NO_SIDE,
            bottom=BLACK_THIN if row == chairman_label else NO_SIDE,
        )
        ws.cell(row, 23).border = Border(left=NO_SIDE, right=NO_SIDE, top=NO_SIDE, bottom=NO_SIDE)
        ws.cell(row, 32).border = Border(left=NO_SIDE, right=NO_SIDE, top=NO_SIDE, bottom=NO_SIDE)
        ws.cell(row, 45).border = Border(left=NO_SIDE, right=NO_SIDE, top=NO_SIDE, bottom=BLACK_THIN)


def normalize_section_label_borders(ws) -> None:
    for text in (
        "Отчет по счетчикам ТРК",
        "*Переход счетчика через максимальное значение",
        "Отчет по резервуарам",
        "Поставки топлива/товара",
    ):
        try:
            row = row_contains(ws, text)
        except ValueError:
            continue
        ws.cell(row, 1).border = Border(left=NO_SIDE, right=NO_SIDE, top=NO_SIDE, bottom=NO_SIDE)


def set_metric_group(ws, row_l: int, col: int, volume: float, mass: float, amount: float, amount_digits: int = 2) -> None:
    set_cell(ws, row_l, col, fmt(volume, 3))
    set_cell(ws, row_l + 1, col, fmt(mass, 3))
    set_cell(ws, row_l + 2, col, fmt(density(mass, volume), 4) if abs(volume or 0) > 1e-9 else "0,0000")
    set_cell(ws, row_l + 3, col, fmt(amount, amount_digits))


def set_metric_group_exact(
    ws,
    row_l: int,
    col: int,
    volume: float,
    mass: float,
    density_value: float,
    amount: float,
    amount_digits: int = 2,
) -> None:
    set_cell(ws, row_l, col, fmt(volume, 3))
    set_cell(ws, row_l + 1, col, fmt(mass, 3))
    set_cell(ws, row_l + 2, col, fmt(density_value, 4))
    set_cell(ws, row_l + 3, col, fmt(amount, amount_digits))


def set_dash_group(ws, row_l: int, col: int, volume=0.0, mass=0.0, amount=0.0) -> None:
    set_cell(ws, row_l, col, fmt(volume, 3))
    set_cell(ws, row_l + 1, col, fmt(mass, 3))
    set_cell(ws, row_l + 2, col, "-")
    set_cell(ws, row_l + 3, col, fmt(amount, 2))


def archive_year_dir(year: int, archive_root: Path | None = None) -> Path:
    return (archive_root or ARCHIVE_ROOT) / f"{year}г"


def archive_monthly_root(year: int, archive_root: Path | None = None) -> Path:
    return archive_year_dir(year, archive_root=archive_root) / "По Месяц"


def archive_yearly_root(year: int, archive_root: Path | None = None) -> Path:
    return archive_year_dir(year, archive_root=archive_root) / "По Год"


def month_dir(year_month: str, archive_root: Path | None = None) -> Path:
    year, month = [int(part) for part in year_month.split("-")]
    return archive_monthly_root(year, archive_root=archive_root) / RU_MONTH_DIR[month]


def is_254_template(path: Path) -> bool:
    wb = None
    try:
        wb = load_workbook(path, read_only=False, data_only=False)
        ws = wb.active
        label_row(ws, "Итого по всем типам топлива, сом")
        wb.close()
        return True
    except Exception:
        if wb is not None:
            wb.close()
        return False


def collect_templates(
    year_month: str,
    template_dirs: list[Path] | None = None,
    archive_root: Path | None = None,
) -> list[Path]:
    year, month = [int(part) for part in year_month.split("-")]
    root = archive_root or ARCHIVE_ROOT
    templates_by_shop: dict[int, Path] = {}
    for candidate_dir in template_dirs or []:
        if not candidate_dir.exists():
            continue
        for path in candidate_dir.glob("*.xlsx"):
            shop_key = parse_shop_key(path)
            if shop_key is not None and shop_key not in templates_by_shop and is_254_template(path):
                templates_by_shop[shop_key] = path
    for candidate_month in range(month, 0, -1):
        candidate_dir = archive_monthly_root(year, archive_root=root) / RU_MONTH_DIR[candidate_month]
        if not candidate_dir.exists():
            continue
        for path in candidate_dir.glob("*.xlsx"):
            shop_key = parse_shop_key(path)
            if shop_key is not None and shop_key not in templates_by_shop and is_254_template(path):
                templates_by_shop[shop_key] = path
    return [templates_by_shop[key] for key in sorted(templates_by_shop)]


def load_period_shifts(cur, shop_key: int, start: date, end_exclusive: date) -> list[dict]:
    completed_by_archive_day = end_exclusive + timedelta(days=1)
    cur.execute(
        """
        SELECT ShiftKey, DatetimeShiftBegin, DatetimeShiftEnd
        FROM shift
        WHERE ShopKey = %s
          AND DatetimeShiftBegin >= %s
          AND DatetimeShiftBegin < %s
          AND DatetimeShiftBegin IS NOT NULL
          AND DatetimeShiftEnd IS NOT NULL
          AND DatetimeShiftEnd > DatetimeShiftBegin
          AND DatetimeShiftEnd < %s
        ORDER BY DatetimeShiftBegin, ShiftKey
        """,
        (shop_key, start.isoformat(), end_exclusive.isoformat(), completed_by_archive_day.isoformat()),
    )
    return cur.fetchall()


def load_collection(cur, shift_range: ShiftRange) -> dict:
    placeholders = sql_in(shift_range.shift_keys)
    cur.execute(
        f"""
        SELECT SUM(Collect) AS Collect, MIN(CashBoxNumber) AS CashBoxNumber
        FROM payment_collection
        WHERE ShopKey = %s
          AND ShiftKey IN ({placeholders})
        """,
        (shift_range.shop_key, *shift_range.shift_keys),
    )
    row = cur.fetchone() or {}
    return {"collect": float(row.get("Collect") or 0), "cash_box": row.get("CashBoxNumber") or 1}


def load_report_operator(cur, shift_range: ShiftRange) -> str:
    placeholders = sql_in(shift_range.shift_keys)
    cur.execute(
        f"""
        SELECT GROUP_CONCAT(DISTINCT cl1.Parameters SEPARATOR ', ') AS operators
        FROM balance bl
        LEFT JOIN (
            SELECT ShiftKey AS ShKey, ShopWorkerKey
            FROM Shift_Participant
            WHERE ShopKey = %s
              AND ShiftKey IN ({placeholders})
            GROUP BY ShiftKey
        ) shu ON shu.ShKey = bl.ShiftKey
        LEFT JOIN (
            SELECT c.SeekKey,
                   SUBSTR(MAX(CONCAT(CAST(c.DateTimeEnd AS CHAR), c.Parameters)), 20) AS Parameters
            FROM change_log c
            WHERE c.ShopKey = %s
              AND c.GuideKey = 598272
            GROUP BY c.SeekKey
        ) cl1 ON cl1.SeekKey = shu.ShopWorkerKey
        WHERE bl.ShopKey = %s
          AND bl.ShiftKey IN ({placeholders})
        """,
        (shift_range.shop_key, *shift_range.shift_keys, shift_range.shop_key, shift_range.shop_key, *shift_range.shift_keys),
    )
    row = cur.fetchone() or {}
    return decode_db_text(row.get("operators"))


def load_main_rows(cur, shift_range: ShiftRange) -> list[dict]:
    placeholders = sql_in(shift_range.shift_keys)
    first_shift = shift_range.shift_keys[0]
    last_shift = shift_range.shift_keys[-1]
    cur.execute(
        f"""
        SELECT
            IF(bl.CollectionKey = 1, -1, bl.ResourceShopKey) AS ResourceKeyGroup,
            IFNULL(sr.Name, '') AS Resource,
            bl.ResourceKey,
            bl.ResourceShopKey,

            IFNULL(bl2.BeginBalance, 0) AS BeginVolume,
            IFNULL(bl2.BeginMass, 0) AS BeginMass,
            IFNULL(bl2.BeginDensity, 0) AS BeginDensity,
            IFNULL(bl2.BeginSum, 0) AS BeginSum,

            IFNULL(sup.SumSupplier, 0) AS SumSupplier,
            IFNULL(sup.VolumeSupplier, 0) AS VolumeSupplier,
            IFNULL(sup.MassSupplier, 0) AS MassSupplier,
            IF(IFNULL(sup.VolumeSupplier, 0) = 0, 0, IFNULL(sup.MassSupplier, 0) / IFNULL(sup.VolumeSupplier, 0)) AS DensitySupplier,

            IFNULL(sl.TotalVolume, 0) AS TotalVolume,
            IFNULL(sl.TotalSum, 0) AS TotalSum,
            IFNULL(sl.TotalMass, 0) AS TotalMass,
            IF(IFNULL(sl.TotalVolume, 0) = 0, 0, IFNULL(sl.TotalMass, 0) / IFNULL(sl.TotalVolume, 0)) AS TotalDensity,

            IFNULL(sl.BankBonusVolume, 0) AS BankBonusVolume,
            IFNULL(sl.BankBonusSum, 0) AS BankBonusSum,
            IFNULL(sl.BankBonusMass, 0) AS BankBonusMass,
            IF(IFNULL(sl.BankBonusVolume, 0) = 0, 0, IFNULL(sl.BankBonusMass, 0) / IFNULL(sl.BankBonusVolume, 0)) AS BankBonusDensity,

            IFNULL(sl.CashBonusVolume, 0) AS CashBonusVolume,
            IFNULL(sl.CashBonusSum, 0) AS CashBonusSum,
            IFNULL(sl.CashBonusMass, 0) AS CashBonusMass,
            IF(IFNULL(sl.CashBonusVolume, 0) = 0, 0, IFNULL(sl.CashBonusMass, 0) / IFNULL(sl.CashBonusVolume, 0)) AS CashBonusDensity,

            IFNULL(sl.CashVolume, 0) AS CashVolume,
            IFNULL(sl.CashSum, 0) AS CashSum,
            IFNULL(sl.CashMass, 0) AS CashMass,
            IF(IFNULL(sl.CashVolume, 0) = 0, 0, IFNULL(sl.CashMass, 0) / IFNULL(sl.CashVolume, 0)) AS CashDensity,

            IFNULL(sl.BonusVolume, 0) AS BonusVolume,
            IFNULL(sl.BonusSum, 0) AS BonusSum,
            IFNULL(sl.BonusMass, 0) AS BonusMass,
            IF(IFNULL(sl.BonusVolume, 0) = 0, 0, IFNULL(sl.BonusMass, 0) / IFNULL(sl.BonusVolume, 0)) AS BonusDensity,

            IFNULL(sl.BankCardVolume, 0) AS BankCardVolume,
            IFNULL(sl.BankCardSum, 0) AS BankCardSum,
            IFNULL(sl.BankCardMass, 0) AS BankCardMass,
            IF(IFNULL(sl.BankCardVolume, 0) = 0, 0, IFNULL(IFNULL(sl.BankCardMass, 0) / sl.BankCardVolume, 0)) AS BankCardDensity,

            IFNULL(sl.StatementVolume, 0) AS StatementVolume,
            IFNULL(sl.StatementSum, 0) AS StatementSum,
            IFNULL(sl.StatementMass, 0) AS StatementMass,
            IF(IFNULL(sl.StatementVolume, 0) = 0, 0, IFNULL(sl.StatementMass, 0) / IFNULL(sl.StatementVolume, 0)) AS StatementDensity,

            IFNULL(sl.CardVolume, 0) AS CardVolume,
            IFNULL(sl.CardSum, 0) AS CardSum,
            IFNULL(sl.CardMass, 0) AS CardMass,
            IF(IFNULL(sl.CardVolume, 0) = 0, 0, IFNULL(sl.CardMass, 0) / IFNULL(sl.CardVolume, 0)) AS CardDensity,

            IFNULL(sl.CuponVolume, 0) AS CuponVolume,
            IFNULL(sl.CuponSum, 0) AS CuponSum,
            IFNULL(sl.CuponMass, 0) AS CuponMass,
            IF(IFNULL(sl.CuponVolume, 0) = 0, 0, IFNULL(sl.CuponMass, 0) / IFNULL(sl.CuponVolume, 0)) AS CuponDensity,

            IFNULL(sup.ReturnSumSupplier, 0) AS ReturnSumSupplier,
            IFNULL(sup.ReturnVolumeSupplier, 0) AS ReturnVolumeSupplier,
            IFNULL(sup.ReturnMassSupplier, 0) AS ReturnMassSupplier,
            IF(IFNULL(sup.ReturnVolumeSupplier, 0) = 0, 0, IFNULL(sup.ReturnMassSupplier, 0) / IFNULL(sup.ReturnVolumeSupplier, 0)) AS ReturnDensitySupplier,

            SUM(bl.EndBalance) AS EndVolume,
            SUM(IFNULL(bl.EndMass, 0)) AS EndMass,
            AVG(IFNULL(bl.EndDensity, 0)) AS EndDensity,
            SUM(bl.EndBalance * IFNULL(bl.EndPrice, 0)) AS EndSum,

            IFNULL(bl2.BeginBalance, 0) + IFNULL(sup.VolumeSupplier, 0) - IFNULL(sup.ReturnVolumeSupplier, 0) - IFNULL(sl.TotalVolume, 0) AS CalcVolume,
            IFNULL(bl2.BeginMass, 0) + IFNULL(sup.MassSupplier, 0) - IFNULL(sup.ReturnMassSupplier, 0) - IFNULL(sl.TotalMass, 0) AS CalcMass,
            IFNULL(bl2.BeginSum, 0) + IFNULL(sup.SumSupplier, 0) - IFNULL(sup.ReturnSumSupplier, 0) - IFNULL(sl.TotalSum, 0) AS CalcSum,

            SUM(bl.EndBalance) - (IFNULL(bl2.BeginBalance, 0) + IFNULL(sup.VolumeSupplier, 0) - IFNULL(sup.ReturnVolumeSupplier, 0) - IFNULL(sl.TotalVolume, 0)) AS DeltaVolume,
            SUM(IFNULL(bl.EndMass, 0)) - (IFNULL(bl2.BeginMass, 0) + IFNULL(sup.MassSupplier, 0) - IFNULL(sup.ReturnMassSupplier, 0) - IFNULL(sl.TotalMass, 0)) AS DeltaMass,
            SUM(bl.EndBalance * IFNULL(bl.EndPrice, 0)) - (IFNULL(bl2.BeginSum, 0) + IFNULL(sup.SumSupplier, 0) - IFNULL(sup.ReturnSumSupplier, 0) - IFNULL(sl.TotalSum, 0)) AS DeltaSum
        FROM balance bl
        LEFT JOIN (
            SELECT
                IF(CollectionKey = 1, -1, ResourceShopKey) AS ResourceKeyGroup3,
                SUM(BeginBalance) AS BeginBalance,
                SUM(IFNULL(BeginMass, 0)) AS BeginMass,
                AVG(IFNULL(BeginDensity, 0)) AS BeginDensity,
                SUM(BeginBalance * IFNULL(BeginPrice, 0)) AS BeginSum
            FROM balance
            WHERE ShopKey = %s
              AND ShiftKey = %s
              AND (WarehouseKey > 0 OR CollectionKey = 1)
            GROUP BY ResourceKeyGroup3
        ) bl2 ON bl2.ResourceKeyGroup3 = IF(bl.CollectionKey = 1, -1, bl.ResourceShopKey)
        LEFT JOIN ShopResource sr
          ON sr.ShopKey = bl.ShopKey
         AND sr.ResourceShopCode = bl.ResourceShopKey
        LEFT JOIN (
            SELECT
                IF(sp.CollectionKey = 1, -1, sp.ResourceShopKey) AS ResourceKeyGroup1,
                SUM(CASE WHEN (IFNULL(DocType, IF(VolumeTTN < 0, 1, 0)) = 0) THEN (sp.SellingPrice * sp.VolumeTTN) ELSE 0 END) AS SumSupplier,
                SUM(CASE WHEN (IFNULL(DocType, IF(VolumeTTN < 0, 1, 0)) = 0) THEN sp.VolumeTTN ELSE 0 END) AS VolumeSupplier,
                SUM(CASE WHEN (IFNULL(DocType, IF(VolumeTTN < 0, 1, 0)) = 0) THEN sp.MassTTN ELSE 0 END) AS MassSupplier,
                AVG(CASE WHEN (IFNULL(DocType, IF(VolumeTTN < 0, 1, 0)) = 0) THEN sp.DensityTTN ELSE 0 END) AS DensitySupplier,
                SUM(CASE WHEN (IFNULL(DocType, IF(VolumeTTN < 0, 1, 0)) = 1) THEN ABS(sp.SellingPrice * sp.VolumeTTN) ELSE 0 END) AS ReturnSumSupplier,
                SUM(CASE WHEN (IFNULL(DocType, IF(VolumeTTN < 0, 1, 0)) = 1) THEN ABS(sp.VolumeTTN) ELSE 0 END) AS ReturnVolumeSupplier,
                SUM(CASE WHEN (IFNULL(DocType, IF(VolumeTTN < 0, 1, 0)) = 1) THEN ABS(sp.MassTTN) ELSE 0 END) AS ReturnMassSupplier,
                AVG(CASE WHEN (IFNULL(DocType, IF(VolumeTTN < 0, 1, 0)) = 1) THEN ABS(sp.DensityTTN) ELSE 0 END) AS ReturnDensitySupplier
            FROM supply sp
            WHERE sp.ShopKey = %s
              AND sp.ShiftKey IN ({placeholders})
              AND IFNULL(sp.Deleted, 0) = 0
            GROUP BY ResourceKeyGroup1
        ) sup ON sup.ResourceKeyGroup1 = IF(bl.CollectionKey = 1, -1, bl.ResourceShopKey)
        LEFT JOIN (
            SELECT
                ResourceShopKey,
                IF(CollectionKey = 1, -1, ResourceShopKey) AS ResourceKeyGroup2,
                SUM(IF(ShopBaseCost != 0, BonusOutCost / ShopBaseCost * Volume, 0)) AS BonusVolume,
                SUM(BonusOutCost) AS BonusSum,
                SUM(IF(ShopBaseCost != 0, BonusOutCost / ShopBaseCost * Mass, 0)) AS BonusMass,
                AVG(Density) AS BonusDensity,
                IF(DerivedPaymentTypeKey = 5, SUM(IF(ShopBaseCost != 0, BonusOutCost / ShopBaseCost * Volume, 0)), 0) AS BankBonusVolume,
                IF(DerivedPaymentTypeKey = 5, SUM(BonusOutCost), 0) AS BankBonusSum,
                IF(DerivedPaymentTypeKey = 5, SUM(IF(ShopBaseCost != 0, BonusOutCost / ShopBaseCost * Mass, 0)), 0) AS BankBonusMass,
                IF(DerivedPaymentTypeKey = 5, AVG(Density), 0) AS BankBonusDensity,
                IF(DerivedPaymentTypeKey != 5, SUM(IF(ShopBaseCost != 0, BonusOutCost / ShopBaseCost * Volume, 0)), 0) AS CashBonusVolume,
                IF(DerivedPaymentTypeKey != 5, SUM(BonusOutCost), 0) AS CashBonusSum,
                IF(DerivedPaymentTypeKey != 5, SUM(IF(ShopBaseCost != 0, BonusOutCost / ShopBaseCost * Mass, 0)), 0) AS CashBonusMass,
                IF(DerivedPaymentTypeKey != 5, AVG(Density), 0) AS CashBonusDensity,
                SUM(CASE WHEN (BasePaymentTypeKey = 2 AND DerivedPaymentTypeKey = 2) THEN (IFNULL(Piece, 0) + IFNULL(Volume, 0)) ELSE 0 END) AS StatementVolume,
                SUM(CASE WHEN (BasePaymentTypeKey = 2 AND DerivedPaymentTypeKey = 2) THEN ShopCost ELSE 0 END) AS StatementSum,
                SUM(CASE WHEN (BasePaymentTypeKey = 2 AND DerivedPaymentTypeKey = 2) THEN Mass ELSE 0 END) AS StatementMass,
                AVG(CASE WHEN (BasePaymentTypeKey = 2 AND DerivedPaymentTypeKey = 2) THEN Density ELSE 0 END) AS StatementDensity,
                SUM(CASE WHEN (BasePaymentTypeKey = 1) THEN (IFNULL(Piece, 0) + IFNULL(Volume, 0)) ELSE 0 END) AS CashVolume,
                SUM(CASE WHEN (BasePaymentTypeKey = 1) THEN ShopCost ELSE 0 END) AS CashSum,
                SUM(CASE WHEN (BasePaymentTypeKey = 1) THEN Mass ELSE 0 END) AS CashMass,
                AVG(CASE WHEN (BasePaymentTypeKey = 1) THEN Density ELSE 0 END) AS CashDensity,
                SUM(CASE WHEN (BasePaymentTypeKey = 2 AND DerivedPaymentTypeKey = 1) THEN (IFNULL(Piece, 0) + IFNULL(Volume, 0)) ELSE 0 END) AS CardVolume,
                SUM(CASE WHEN (BasePaymentTypeKey = 2 AND DerivedPaymentTypeKey = 1) THEN ShopCost ELSE 0 END) AS CardSum,
                SUM(CASE WHEN (BasePaymentTypeKey = 2 AND DerivedPaymentTypeKey = 1) THEN Mass ELSE 0 END) AS CardMass,
                AVG(CASE WHEN (BasePaymentTypeKey = 2 AND DerivedPaymentTypeKey = 1) THEN Density ELSE 0 END) AS CardDensity,
                SUM(CASE WHEN (BasePaymentTypeKey = 2 AND DerivedPaymentTypeKey = 5) THEN (IFNULL(Piece, 0) + IFNULL(Volume, 0)) ELSE 0 END) AS BankCardVolume,
                SUM(CASE WHEN (BasePaymentTypeKey = 2 AND DerivedPaymentTypeKey = 5) THEN ShopCost ELSE 0 END) AS BankCardSum,
                SUM(CASE WHEN (BasePaymentTypeKey = 2 AND DerivedPaymentTypeKey = 5) THEN Mass ELSE 0 END) AS BankCardMass,
                AVG(CASE WHEN (BasePaymentTypeKey = 2 AND DerivedPaymentTypeKey = 5) THEN Density ELSE 0 END) AS BankCardDensity,
                SUM(CASE WHEN (BasePaymentTypeKey = 2 AND DerivedPaymentTypeKey = 3) THEN (IFNULL(Piece, 0) + IFNULL(Volume, 0)) ELSE 0 END) AS CuponVolume,
                SUM(CASE WHEN (BasePaymentTypeKey = 2 AND DerivedPaymentTypeKey = 3) THEN ShopCost ELSE 0 END) AS CuponSum,
                SUM(CASE WHEN (BasePaymentTypeKey = 2 AND DerivedPaymentTypeKey = 3) THEN Mass ELSE 0 END) AS CuponMass,
                AVG(CASE WHEN (BasePaymentTypeKey = 2 AND DerivedPaymentTypeKey = 3) THEN Density ELSE 0 END) AS CuponDensity,
                SUM(CASE WHEN (BasePaymentTypeKey != 2 OR DerivedPaymentTypeKey != 4 OR PourOutWarehouse = 0) THEN (IFNULL(Piece, 0) + IFNULL(Volume, 0)) ELSE 0 END) AS TotalVolume,
                SUM(CASE WHEN (BasePaymentTypeKey != 2 OR DerivedPaymentTypeKey != 4 OR PourOutWarehouse = 0) THEN ShopCost ELSE 0 END) AS TotalSum,
                SUM(CASE WHEN (BasePaymentTypeKey != 2 OR DerivedPaymentTypeKey != 4 OR PourOutWarehouse = 0) THEN Mass ELSE 0 END) AS TotalMass,
                AVG(CASE WHEN (BasePaymentTypeKey != 2 OR DerivedPaymentTypeKey != 4 OR PourOutWarehouse = 0) THEN Density ELSE 0 END) AS TotalDensity
            FROM Selling
            WHERE ShopKey = %s
              AND ShiftKey IN ({placeholders})
            GROUP BY ResourceKeyGroup2
        ) sl ON sl.ResourceKeyGroup2 = IF(bl.CollectionKey = 1, -1, bl.ResourceShopKey)
        WHERE bl.ShopKey = %s
          AND bl.ShiftKey = %s
          AND (bl.WarehouseKey > 0 OR bl.CollectionKey = 1)
          AND (IFNULL(bl.Flags, 0) & 0x00000040) = 0
        GROUP BY ResourceKeyGroup
        """,
        (
            shift_range.shop_key,
            first_shift,
            shift_range.shop_key,
            *shift_range.shift_keys,
            shift_range.shop_key,
            *shift_range.shift_keys,
            shift_range.shop_key,
            last_shift,
        ),
    )
    rows = []
    for row in cur.fetchall():
        if int(row.get("ResourceKeyGroup") or 0) == -1:
            resource = -1
            label = decode_db_text(row.get("Resource")) or "Соп. тов."
        else:
            label = decode_db_text(row.get("Resource"))
            resource = resource_from_text(label) or infer_report_resource(
                shift_range.shop_key,
                row.get("ResourceKey"),
                row.get("ResourceShopKey"),
            )
            label = label or RESOURCE_LABELS.get(resource, str(resource))
        row["resource"] = resource
        row["resource_label"] = label
        rows.append(row)
    return rows


def load_sales(cur, shift_range: ShiftRange) -> tuple[dict[int, dict], dict[int, float], float]:
    placeholders = sql_in(shift_range.shift_keys)
    mass_expr = _selling_mass_expr(
        "s",
        use_ai80_missing_density_fallback=not _allow_fixed_density_fallback(shift_range.begin),
    )
    cur.execute(
        f"""
        SELECT ResourceKey, ResourceShopKey, BasePaymentTypeKey, DerivedPaymentTypeKey,
               COALESCE(BasePrice, 0) AS Price,
               SUM(Volume) AS Volume,
               SUM(ShopCost) AS Amount,
               SUM(BonusOutCost) AS BonusCost,
               SUM({mass_expr}) AS Mass
        FROM selling s
        WHERE ShopKey = %s
          AND ShiftKey IN ({placeholders})
        GROUP BY ResourceKey, ResourceShopKey, BasePaymentTypeKey, DerivedPaymentTypeKey,
                 COALESCE(BasePrice, 0)
        """,
        (shift_range.shop_key, *shift_range.shift_keys),
    )
    result: dict[int, dict] = defaultdict(lambda: defaultdict(float))
    checks: dict[int, float] = defaultdict(float)
    cash_amount = 0.0
    for row in cur.fetchall():
        resource = infer_report_resource(
            shift_range.shop_key,
            row.get("ResourceKey"),
            row.get("ResourceShopKey"),
            row.get("Price"),
        )
        base = int(row.get("BasePaymentTypeKey") or 0)
        derived = int(row.get("DerivedPaymentTypeKey") or 0)
        price = float(row.get("Price") or 0)
        volume = float(row.get("Volume") or 0)
        amount = float(row.get("Amount") or 0)
        bonus_amount = float(row.get("BonusCost") or 0)
        mass = float(row.get("Mass") or 0)
        row_density = density(mass, volume)
        bonus_volume = bonus_amount / price if price else 0.0
        bonus_mass = bonus_volume * row_density
        item = result[resource]
        if base == 1:
            cash_amount += amount
            item["cash_l"] += volume - bonus_volume
            item["cash_amt"] += amount - bonus_amount
            item["cash_mass"] += mass - bonus_mass
            item["bonus_l"] += bonus_volume
            item["bonus_amt"] += bonus_amount
            item["bonus_mass"] += bonus_mass
        elif derived == 2:
            item["ved_l"] += volume
            item["ved_amt"] += amount
            item["ved_mass"] += mass
        elif derived == 3:
            item["talon_l"] += volume
            item["talon_amt"] += amount
            item["talon_mass"] += mass
        elif derived == 4:
            checks[resource] += volume
        elif base == 2 or derived == 5:
            item["card_l"] += volume
            item["card_amt"] += amount
            item["card_mass"] += mass
    for item in result.values():
        for prefix in ("cash", "card", "ved", "talon", "bonus"):
            item["total_l"] += item.get(f"{prefix}_l", 0.0)
            item["total_amt"] += item.get(f"{prefix}_amt", 0.0)
            item["total_mass"] += item.get(f"{prefix}_mass", 0.0)
    return {key: dict(value) for key, value in result.items()}, dict(checks), cash_amount


def supply_doc(row: dict) -> str:
    for key in ("TTNNumberExt", "InvoiceNumberExt", "TTNNumber", "SupplyNumber"):
        value = row.get(key)
        if value not in (None, "", 0, 2147483647):
            return str(value).strip()
    return ""


def load_supply(cur, shift_range: ShiftRange) -> tuple[dict[int, dict], list[dict]]:
    placeholders = sql_in(shift_range.shift_keys)
    cur.execute(
        f"""
        SELECT
            IF(sp.CollectionKey = 1, -1, sp.ResourceShopKey) AS ResourceKeyGroup,
            sp.ResourceKey,
            sp.ResourceShopKey,
            IFNULL(sr.Name, '') AS Resource,
            c3.Parameters AS Supplier,
            sp.SupplyNumber,
            SUM(sp.VolumeTTN) AS VolumeTTN,
            SUM(sp.MassTTN) AS MassTTN,
            AVG(sp.DensityTTN) AS DensityTTN,
            SUM(sp.SellingPrice * sp.VolumeTTN) AS SumTTN
        FROM supply sp
        LEFT JOIN ShopResource sr
          ON sr.ShopKey = sp.ShopKey
         AND sr.ResourceShopCode = sp.ResourceShopKey
        LEFT JOIN (
            SELECT c.SeekKey,
                   SUBSTR(MAX(CONCAT(CAST(c.DateTimeEnd AS CHAR), c.Parameters)), 20) AS Parameters
            FROM change_log c
            WHERE c.ShopKey = %s
              AND c.GuideKey = 600320
            GROUP BY c.SeekKey
        ) c3 ON c3.SeekKey = sp.SupplierShopKey
        WHERE sp.ShopKey = %s
          AND sp.ShiftKey IN ({placeholders})
          AND IFNULL(DocType, IF(VolumeTTN < 0, 1, 0)) = 0
          AND IFNULL(sp.Deleted, 0) = 0
          AND sp.CollectionKey > 0
        GROUP BY ResourceKeyGroup, Supplier, SupplyNumber
        """,
        (shift_range.shop_key, shift_range.shop_key, *shift_range.shift_keys),
    )
    rows: list[dict] = []
    for row in cur.fetchall():
        if int(row.get("ResourceKeyGroup") or 0) == -1:
            resource = -1
            label = decode_db_text(row.get("Resource")) or "Соп. тов."
        else:
            label = decode_db_text(row.get("Resource"))
            resource = resource_from_text(label) or infer_supply_resource(
                shift_range.shop_key,
                row.get("ResourceKey"),
                row.get("ResourceShopKey"),
            )
            label = label or RESOURCE_LABELS.get(resource, str(resource))
        rows.append(
            {
                "resource": resource,
                "resource_label": label,
                "supplier": decode_db_text(row.get("Supplier")),
                "doc": "" if row.get("SupplyNumber") is None else str(row.get("SupplyNumber")),
                "volume": float(row.get("VolumeTTN") or 0),
                "mass": float(row.get("MassTTN") or 0),
                "density": float(row.get("DensityTTN") or 0),
                "amount": float(row.get("SumTTN") or 0),
            }
        )
    totals: dict[int, dict] = defaultdict(lambda: defaultdict(float))
    for row in rows:
        item = totals[row["resource"]]
        item["volume"] += row["volume"]
        item["mass"] += row["mass"]
        item["amount"] += row["amount"]
    for item in totals.values():
        item["density"] = density(item["mass"], item["volume"])
    return {key: dict(value) for key, value in totals.items()}, rows


def load_balances(cur, shift_range: ShiftRange) -> tuple[list[dict], dict[int, dict], list[dict]]:
    first_shift = shift_range.shift_keys[0]
    last_shift = shift_range.shift_keys[-1]
    placeholders = sql_in(shift_range.shift_keys)
    cur.execute(
        f"""
        SELECT
            IFNULL(bl.Flags, 0) & 0x00000010 AS isMassCounter,
            bl.VendorKey,
            bl.ResourceKey,
            bl.ResourceShopKey,
            IFNULL(sr.Name, '') AS Resource,
            IFNULL(bl2.BeginBalance, 0) AS BeginCounter,
            bl.EndBalance AS EndCounter,
            IF(bl.EndBalance < IFNULL(bl2.BeginBalance, 0), '*', '') AS CounterChange,
            IFNULL(ss.SaleCheckingVolume, 0) AS SaleCheckingVolume,
            IFNULL(ss.TotalVolume, 0) AS TotalVolume,
            CASE
                WHEN bl.EndBalance < IFNULL(bl2.BeginBalance, 0)
                THEN bl.EndBalance - IFNULL(bl2.BeginBalance, 0) - IFNULL(ss.SaleCheckingVolume, 0) + IFNULL(chv.MaxCounter, 0)
                ELSE bl.EndBalance - IFNULL(bl2.BeginBalance, 0) - IFNULL(ss.SaleCheckingVolume, 0)
            END AS DeltaCounter,
            CASE
                WHEN bl.EndBalance < IFNULL(bl2.BeginBalance, 0)
                THEN ROUND(bl.EndBalance - IFNULL(bl2.BeginBalance, 0) - IFNULL(ss.SaleCheckingVolume, 0) + IFNULL(chv.MaxCounter, 0) - IFNULL(ss.TotalVolume, 0), 3)
                ELSE ROUND(bl.EndBalance - IFNULL(bl2.BeginBalance, 0) - IFNULL(ss.SaleCheckingVolume, 0) - IFNULL(ss.TotalVolume, 0), 3)
            END AS BalanceCounter
        FROM balance bl
        LEFT JOIN ShopResource sr
          ON sr.ShopKey = bl.ShopKey
         AND sr.ResourceShopCode = bl.ResourceShopKey
        LEFT JOIN (
            SELECT VendorKey, BeginBalance
            FROM balance
            WHERE ShopKey = %s
              AND ShiftKey = %s
              AND CollectionKey > 1
              AND VendorKey != 0
        ) bl2 ON bl2.VendorKey = bl.VendorKey
        LEFT JOIN (
            SELECT
                c.SeekKey,
                CAST(SUBSTR(MAX(IF(GuideKey = 3684355, CONCAT(CAST(c.DateTimeEnd AS CHAR), c.Parameters), '')), 20) AS SIGNED) AS VendorKey,
                CAST(SUBSTR(MAX(IF(GuideKey = 3684357, CONCAT(CAST(c.DateTimeEnd AS CHAR), c.Parameters), '0')), 20) AS SIGNED) / 1000 AS MaxCounter
            FROM change_log c
            WHERE c.ShopKey = %s
              AND c.GuideKey IN (3684355, 3684357)
              AND IFNULL(flags, 0) = 0
            GROUP BY c.SeekKey
        ) chv ON chv.VendorKey = bl.VendorKey
        LEFT JOIN (
            SELECT
                s.ShopKey,
                s.VendorKey,
                SUM(CASE
                    WHEN (s.BasePaymentTypeKey = 2 AND s.DerivedPaymentTypeKey = 4 AND s.PourOutWarehouse != 0)
                    THEN IF((s.Flags & 0x00400000) > 0, Mass, Volume)
                    ELSE 0
                END) AS SaleCheckingVolume,
                SUM(CASE
                    WHEN (s.BasePaymentTypeKey != 2 OR s.DerivedPaymentTypeKey != 4 OR s.PourOutWarehouse = 0)
                    THEN IF((s.Flags & 0x00400000) > 0, Mass, Volume)
                    ELSE 0
                END) AS TotalVolume
            FROM Selling s
            WHERE s.ShopKey = %s
              AND s.ShiftKey IN ({placeholders})
              AND s.VendorKey > 0
            GROUP BY s.VendorKey
        ) ss ON ss.ShopKey = bl.ShopKey AND ss.VendorKey = bl.VendorKey
        WHERE bl.ShopKey = %s
          AND bl.ShiftKey = %s
          AND bl.CollectionKey > 1
          AND bl.VendorKey != 0
        GROUP BY bl.VendorKey
        """,
        (
            shift_range.shop_key,
            first_shift,
            shift_range.shop_key,
            shift_range.shop_key,
            *shift_range.shift_keys,
            shift_range.shop_key,
            last_shift,
        ),
    )
    counters = []
    for row in cur.fetchall():
        label = decode_db_text(row.get("Resource"))
        resource = resource_from_text(label) or infer_report_resource(
            shift_range.shop_key,
            row.get("ResourceKey"),
            row.get("ResourceShopKey"),
        )
        counters.append(
            {
                "number": int(row.get("VendorKey") or 0),
                "unit": "кг" if int(row.get("isMassCounter") or 0) else "л",
                "resource": resource,
                "resource_label": label or RESOURCE_LABELS.get(resource, str(resource)),
                "begin": float(row.get("BeginCounter") or 0),
                "delta_counter": float(row.get("DeltaCounter") or 0),
                "total_volume": float(row.get("TotalVolume") or 0),
                "sale_check": float(row.get("SaleCheckingVolume") or 0),
                "end": float(row.get("EndCounter") or 0),
                "counter_change": decode_db_text(row.get("CounterChange")),
                "balance_counter": float(row.get("BalanceCounter") or 0),
            }
        )

    cur.execute(
        """
        SELECT
            bl.WarehouseKey,
            bl.ResourceKey,
            bl.ResourceShopKey,
            IFNULL(sr.Name, '') AS Resource,
            SUM(bl.EndBalance) AS Volume,
            SUM(IFNULL(bl.EndMass, 0)) AS Mass,
            AVG(bl.EndDensity) AS Density,
            bl.EndTemperature AS Temperature,
            bl.EndLevel AS Level
        FROM balance bl
        LEFT JOIN ShopResource sr
          ON sr.ShopKey = bl.ShopKey
         AND sr.ResourceShopCode = bl.ResourceShopKey
        WHERE bl.ShopKey = %s
          AND bl.ShiftKey = %s
          AND bl.CollectionKey > 1
          AND bl.WarehouseKey != 0
        GROUP BY bl.WarehouseKey
        """,
        (shift_range.shop_key, last_shift),
    )
    tanks: list[dict] = []
    for row in cur.fetchall():
        label = decode_db_text(row.get("Resource"))
        resource = resource_from_text(label) or infer_report_resource(
            shift_range.shop_key,
            row.get("ResourceKey"),
            row.get("ResourceShopKey"),
        )
        tanks.append(
            {
                "warehouse": int(row.get("WarehouseKey") or 0),
                "resource": resource,
                "resource_label": label or RESOURCE_LABELS.get(resource, str(resource)),
                "end_l": float(row.get("Volume") or 0),
                "end_mass": float(row.get("Mass") or 0),
                "end_density": float(row.get("Density") or 0),
                "end_temperature": float(row.get("Temperature") or 0),
                "end_level": float(row.get("Level") or 0),
            }
        )
    return counters, {}, tanks


def load_cash_rows(cur, shift_range: ShiftRange) -> list[dict]:
    placeholders = sql_in(shift_range.shift_keys)
    query = f"""
        SELECT
            GROUP_CONCAT(DISTINCT oper.UserName SEPARATOR ', ') AS UserName,
            CAST(MIN(oper.DatetimeShiftBegin) AS CHAR) AS DatetimeShiftBegin,
            CAST(MAX(oper.DatetimeShiftEnd) AS CHAR) AS DatetimeShiftEnd,
            CAST(CONCAT(MIN(bl.ShiftKey), ' - ', MAX(bl.ShiftKey)) AS CHAR) AS ShiftKey,
            bl.CashBoxNumber,
            SUM(bl.BeginBalanceCalculation) AS StartVal,
            SUM(bl.EndBalance) AS EndVal,
            SUM(bl.EndBalanceCalculation) AS BookEndVal,
            SUM(bl.EndBalance) - SUM(bl.EndBalanceCalculation) AS EndDelta,
            SUM(sel.TotalValue) AS TotalValue,
            SUM(sel.FuelValue) AS FuelValue,
            SUM(sel.GoodsValue) AS GoodsValue,
            SUM(sel.PaymentValue) AS PaymentValue,
            SUM(sel.ServiceValue) AS ServiceValue,
            SUM(sel.AcceptRemainsValue) AS AcceptRemainsValue,
            SUM(sel.CardPayment) AS CardPayment,
            SUM(sel.RepaymentValue) AS RepaymentValue,
            SUM(IFNULL(pc.Amount, 0)) AS Amount,
            SUM(IFNULL(pc.BankAmount, 0)) AS BankAmount,
            SUM(IFNULL(pc.CashBoxAmount, 0)) AS CashBoxAmount,
            SUM(ss.Delta) AS Delta
        FROM balance bl
        LEFT JOIN (
            SELECT
                s.ShopKey,
                s.ShiftKey,
                s.CashBoxNumber,
                s.TotalValue,
                s.FuelValue,
                s.GoodsValue,
                s.PaymentValue,
                s.ServiceValue,
                s.AcceptRemainsValue,
                r.RepaymentValue,
                pm.CardPayment
            FROM (
                SELECT
                    sl.CashBoxNumber,
                    sl.ShopKey,
                    sl.ShiftKey,
                    SUM(CASE WHEN (sl.CollectionKey != 0) THEN sl.ShopCost ELSE 0 END) AS TotalValue,
                    SUM(CASE WHEN (sl.CollectionKey != 9 AND sl.CollectionKey != 1 AND sl.CollectionKey != 0) THEN sl.ShopCost ELSE 0 END) AS FuelValue,
                    SUM(CASE WHEN (sl.CollectionKey = 1) THEN sl.ShopCost ELSE 0 END) AS GoodsValue,
                    SUM(sl.PayingSum) AS PaymentValue,
                    SUM(CASE WHEN (sl.CollectionKey = 9) THEN sl.ShopCost ELSE 0 END) AS ServiceValue,
                    SUM(IF((sl.BasePaymentTypeKey = 1 AND sl.TransactionSign = 3), sl.ShopRepaymentCost, 0)) AS AcceptRemainsValue,
                    SUM(IF((sl.BasePaymentTypeKey = 1 AND sl.RepaymentCheckNumber != 0), sl.ShopCost + sl.ShopRepaymentCost, 0)) AS _RepaymentValue
                FROM selling sl
                WHERE sl.ShopKey = %s
                  AND sl.ShiftKey IN ({placeholders})
                  AND (sl.BasePaymentTypeKey = 1 OR sl.TransactionType = 10)
                GROUP BY sl.ShopKey, sl.ShiftKey, sl.CashBoxNumber
            ) s
            LEFT JOIN (
                SELECT
                    el.CashBoxNumber,
                    el.ShopKey,
                    el.ShiftKey,
                    SUM(ABS(el.ShopCost)) AS RepaymentValue
                FROM EventLog el
                WHERE el.ShopKey = %s
                  AND el.ShiftKey IN ({placeholders})
                  AND el.EventType = 196610
                  AND el.BasePaymentTypeKey = 1
                GROUP BY el.ShopKey, el.ShiftKey, el.CashBoxNumber
            ) r ON r.ShopKey = s.ShopKey AND r.ShiftKey = s.ShiftKey AND r.CashBoxNumber = s.CashBoxNumber
            LEFT JOIN (
                SELECT
                    p.CashBoxNumber,
                    p.ShopKey,
                    p.ShiftKey,
                    SUM(p.AmountMoney) AS CardPayment
                FROM payment p
                WHERE p.ShopKey = %s
                  AND p.ShiftKey IN ({placeholders})
                  AND p.CashBoxNumber > 0
                GROUP BY p.ShopKey, p.ShiftKey, p.CashBoxNumber
            ) pm ON pm.ShopKey = s.ShopKey AND pm.ShiftKey = s.ShiftKey AND pm.CashBoxNumber = s.CashBoxNumber
        ) sel ON sel.CashBoxNumber = bl.CashBoxNumber AND sel.ShiftKey = bl.ShiftKey AND sel.ShopKey = bl.ShopKey
        LEFT JOIN (
            SELECT
                ShopKey,
                ShiftKey,
                CashBoxNumber,
                SUM(IF(s.CostDelta < 0 OR IFNULL(s.ShopCost, 0) < 0, 0, s.CostDelta)) AS Delta
            FROM selling s
            WHERE ShopKey = %s
              AND ShiftKey IN ({placeholders})
              AND s.CollectionKey > 1
              AND s.BasePaymentTypeKey = 1
              AND s.ShopRepaymentCost = 0
            GROUP BY ShopKey, ShiftKey, CashBoxNumber
        ) ss ON ss.CashBoxNumber = bl.CashBoxNumber AND ss.ShiftKey = bl.ShiftKey AND ss.ShopKey = bl.ShopKey
        LEFT JOIN (
            SELECT
                cl.Parameters AS UserName,
                sh.DatetimeShiftBegin,
                sh.DatetimeShiftEnd,
                shp.ShopKey,
                shp.ShiftKey
            FROM Shift_Participant shp
            LEFT JOIN (
                SELECT
                    c.SeekKey,
                    SUBSTR(MAX(CONCAT(CAST(c.DateTimeEnd AS CHAR), c.Parameters)), 20) AS Parameters
                FROM change_log c
                WHERE c.ShopKey = %s
                  AND c.GuideKey = 598272
                GROUP BY c.SeekKey
            ) cl ON cl.SeekKey = shp.ShopWorkerKey
            LEFT JOIN shift sh ON sh.ShopKey = shp.ShopKey AND sh.ShiftKey = shp.ShiftKey
            WHERE shp.ShopKey = %s
              AND shp.ShiftKey IN ({placeholders})
        ) oper ON oper.ShiftKey = bl.ShiftKey AND oper.ShopKey = bl.ShopKey
        LEFT JOIN (
            SELECT
                ShopKey,
                ShiftKey,
                CashBoxNumber,
                SUM(Collect) AS Amount,
                SUM(CASE WHEN (CollectionType != 2) THEN Collect ELSE 0 END) AS BankAmount,
                SUM(CASE WHEN (CollectionType = 2) THEN Collect ELSE 0 END) AS CashBoxAmount
            FROM payment_collection
            WHERE ShopKey = %s
              AND ShiftKey IN ({placeholders})
            GROUP BY ShopKey, ShiftKey, CashBoxNumber
        ) pc ON pc.CashBoxNumber = bl.CashBoxNumber AND pc.ShiftKey = bl.ShiftKey AND pc.ShopKey = bl.ShopKey
        WHERE bl.ShopKey = %s
          AND bl.ShiftKey IN ({placeholders})
          AND bl.TurnoverType = 1
        GROUP BY bl.CashBoxNumber
        ORDER BY bl.CashBoxNumber
    """
    params = []
    for _ in range(4):
        params.extend([shift_range.shop_key, *shift_range.shift_keys])
    params.append(shift_range.shop_key)
    params.extend([shift_range.shop_key, *shift_range.shift_keys])
    params.extend([shift_range.shop_key, *shift_range.shift_keys])
    params.extend([shift_range.shop_key, *shift_range.shift_keys])
    cur.execute(query, tuple(params))
    rows = []
    for row in cur.fetchall():
        row["UserName"] = decode_db_text(row.get("UserName"))
        rows.append(row)
    return rows


def load_period_data(cur, shop_key: int, start: date, end_exclusive: date) -> PeriodData | None:
    shifts = load_period_shifts(cur, shop_key, start, end_exclusive)
    if not shifts:
        return None
    shift_range = ShiftRange(
        shop_key=shop_key,
        shift_keys=tuple(int(row["ShiftKey"]) for row in shifts),
        begin=shifts[0]["DatetimeShiftBegin"],
        end=shifts[-1]["DatetimeShiftEnd"],
    )
    main_rows = load_main_rows(cur, shift_range)
    sales, checks, cash_amount = load_sales(cur, shift_range)
    supplies, supply_entries = load_supply(cur, shift_range)
    counters, tank_totals, tanks = load_balances(cur, shift_range)
    cash_rows = load_cash_rows(cur, shift_range)
    collection = load_collection(cur, shift_range)
    operator = load_report_operator(cur, shift_range)
    return PeriodData(
        shift_range=shift_range,
        first_shift=shifts[0],
        last_shift=shifts[-1],
        main_rows=main_rows,
        sales=sales,
        checks=checks,
        supplies=supplies,
        supply_entries=supply_entries,
        tanks=tanks,
        tank_totals=tank_totals,
        counters=counters,
        cash_rows=cash_rows,
        cash_amount=cash_amount,
        collection=collection,
        operator=operator,
    )


def template_resources(ws) -> list[int]:
    total_row = label_row(ws, "Итого по всем типам топлива, сом")
    result = []
    row = 10
    while row < total_row:
        resource = resource_from_text(ws.cell(row, 1).value)
        if resource:
            result.append(resource)
            row += 4
        else:
            row += 1
    return result


def section_range(ws, header: str, data_offset: int, next_header: str) -> tuple[int, int]:
    start_label = label_row(ws, header)
    next_label = label_row(ws, next_header, start_label + 1)
    start = start_label + data_offset
    end = next_label - 2
    return start, max(start - 1, end)


def write_main_table(ws, period: PeriodData) -> list[int]:
    total_row = label_row(ws, "Итого по всем типам топлива, сом")
    existing_count = max(0, total_row - 10)
    rows = period.main_rows
    resources = [int(row["resource"]) for row in rows]
    needed_count = len(rows) * 4
    ensure_main_capacity(ws, 10, existing_count, needed_count, 10)
    total_row = label_row(ws, "Итого по всем типам топлива, сом")
    for row in range(10, total_row):
        clear_row(ws, row)
    total_amounts = defaultdict(float)
    num = lambda row, key: float(row.get(key) or 0)
    for idx, data in enumerate(rows):
        resource = int(data["resource"])
        row = 10 + idx * 4
        set_cell(ws, row, 1, data.get("resource_label") or RESOURCE_LABELS.get(resource, str(resource)))
        set_cell(ws, row, 3, "шт." if resource == -1 else "л.")
        set_cell(ws, row + 1, 3, "кг.")
        set_cell(ws, row + 2, 3, "кг/л")
        set_cell(ws, row + 3, 3, "сом")
        if period.shift_range.shop_key == 15 and resource == 102:
            ws.row_dimensions[row + 3].height = 19.4
        elif period.shift_range.shop_key == 5 and resource == 102:
            ws.row_dimensions[row + 3].height = 19.4 if period.shift_range.begin.month == 5 else 11.21

        set_metric_group_exact(
            ws,
            row,
            MAIN_COLS["begin"],
            num(data, "BeginVolume"),
            num(data, "BeginMass"),
            num(data, "BeginDensity"),
            num(data, "BeginSum"),
        )
        set_metric_group_exact(
            ws,
            row,
            MAIN_COLS["supply"],
            num(data, "VolumeSupplier"),
            num(data, "MassSupplier"),
            num(data, "DensitySupplier"),
            num(data, "SumSupplier"),
        )
        set_metric_group_exact(
            ws,
            row,
            MAIN_COLS["realization"],
            num(data, "TotalVolume"),
            num(data, "TotalMass"),
            num(data, "TotalDensity"),
            num(data, "TotalSum"),
        )
        set_metric_group_exact(
            ws,
            row,
            MAIN_COLS["realization_net"],
            num(data, "TotalVolume") + num(data, "ReturnVolumeSupplier"),
            num(data, "TotalMass") + num(data, "ReturnMassSupplier"),
            num(data, "TotalDensity"),
            num(data, "TotalSum") + num(data, "ReturnSumSupplier"),
        )
        set_metric_group_exact(
            ws,
            row,
            MAIN_COLS["cash"],
            num(data, "CashVolume") - num(data, "CashBonusVolume"),
            num(data, "CashMass") - num(data, "CashBonusMass"),
            num(data, "CashDensity"),
            num(data, "CashSum") - num(data, "CashBonusSum"),
        )
        set_metric_group_exact(
            ws,
            row,
            MAIN_COLS["fuel_card"],
            num(data, "CardVolume"),
            num(data, "CardMass"),
            num(data, "CardDensity"),
            num(data, "CardSum"),
        )
        set_metric_group_exact(
            ws,
            row,
            MAIN_COLS["bank_card"],
            num(data, "BankCardVolume") - num(data, "BankBonusVolume"),
            num(data, "BankCardMass") - num(data, "BankBonusMass"),
            num(data, "BankCardDensity"),
            num(data, "BankCardSum") - num(data, "BankBonusSum"),
        )
        set_metric_group_exact(
            ws,
            row,
            MAIN_COLS["vedomosti"],
            num(data, "StatementVolume"),
            num(data, "StatementMass"),
            num(data, "StatementDensity"),
            num(data, "StatementSum"),
        )
        set_metric_group_exact(
            ws,
            row,
            MAIN_COLS["talon"],
            num(data, "CuponVolume"),
            num(data, "CuponMass"),
            num(data, "CuponDensity"),
            num(data, "CuponSum"),
        )
        set_metric_group_exact(
            ws,
            row,
            MAIN_COLS["bonus"],
            num(data, "BonusVolume"),
            num(data, "BonusMass"),
            num(data, "BonusDensity"),
            num(data, "BonusSum"),
        )
        set_metric_group_exact(
            ws,
            row,
            MAIN_COLS["return_supplier"],
            num(data, "ReturnVolumeSupplier"),
            num(data, "ReturnMassSupplier"),
            num(data, "ReturnDensitySupplier"),
            num(data, "ReturnSumSupplier"),
        )
        set_metric_group_exact(
            ws,
            row,
            MAIN_COLS["end_fact"],
            num(data, "EndVolume"),
            num(data, "EndMass"),
            num(data, "EndDensity"),
            num(data, "EndSum"),
        )
        set_dash_group(ws, row, MAIN_COLS["end_calc"], num(data, "CalcVolume"), num(data, "CalcMass"), num(data, "CalcSum"))
        set_dash_group(ws, row, MAIN_COLS["balance"], num(data, "DeltaVolume"), num(data, "DeltaMass"), num(data, "DeltaSum"))

        for name, col in MAIN_COLS.items():
            if name in {"begin", "end_fact", "end_calc", "balance"}:
                continue
            total_amounts[col] += as_float(ws.cell(row + 3, col).value, 0)

    clear_row(ws, total_row)
    set_cell(ws, total_row, 1, "Итого по всем типам топлива, сом")
    if period.shift_range.shop_key in {1, 2, 6}:
        ws.row_dimensions[total_row].height = 19.4
    if period.shift_range.begin.month == 6 and period.shift_range.shop_key in {14, 21}:
        ws.row_dimensions[total_row].height = 11.21
    for col in (10, 15, 17, 21, 25, 29, 31, 36, 39, 43):
        set_cell(ws, total_row, col, fmt(total_amounts[col], 2))
    for col in (47, 52, 57):
        set_cell(ws, total_row, col, "-")
    return resources


def write_counter_table(ws, period: PeriodData) -> None:
    counter_label = label_row(ws, "Отчет по счетчикам ТРК")
    star_label = row_contains(ws, "*Переход счетчика через максимальное значение", counter_label + 1)
    start, end = counter_label + 3, star_label - 1
    current = max(0, end - start + 1)
    needed = len(period.counters)
    ensure_capacity(ws, start, current, needed, start if current else start - 1)
    for row in range(start, start + needed):
        clear_row(ws, row)
    for idx, counter in enumerate(period.counters):
        row = start + idx
        set_cell(ws, row, 1, str(counter["number"]))
        set_cell(ws, row, 6, counter.get("unit") or "л")
        set_cell(ws, row, 13, counter.get("resource_label") or RESOURCE_LABELS.get(counter["resource"], ""))
        set_cell(ws, row, 20, fmt(counter["begin"], 3))
        set_cell(ws, row, 27, fmt(counter["delta_counter"], 3))
        set_cell(ws, row, 35, fmt(counter["total_volume"], 3))
        set_cell(ws, row, 42, fmt(counter["sale_check"], 3))
        set_cell(ws, row, 50, f"{fmt(counter['end'], 3)}{counter.get('counter_change') or ''}")
        set_cell(ws, row, 56, fmt(counter["balance_counter"], 3))


def write_tank_table(ws, period: PeriodData) -> None:
    start, end = section_range(ws, "Отчет по резервуарам", 3, "Поставки топлива/товара")
    current = max(0, end - start + 1)
    needed = len(period.tanks)
    ensure_capacity(ws, start, current, needed, start if current else start - 1)
    for row in range(start, start + needed):
        clear_row(ws, row)
    for idx, tank in enumerate(period.tanks):
        row = start + idx
        set_cell(ws, row, 1, str(tank["warehouse"] or idx + 1))
        set_cell(ws, row, 9, tank.get("resource_label") or RESOURCE_LABELS.get(tank["resource"], ""))
        set_cell(ws, row, 19, fmt(tank["end_l"], 3))
        set_cell(ws, row, 29, fmt(tank["end_mass"], 3))
        set_cell(ws, row, 38, fmt(tank["end_density"], 4))
        set_cell(ws, row, 46, round_float(tank["end_temperature"], 2))
        set_cell(ws, row, 53, int(round(tank["end_level"] or 0)))


def write_supply_table(ws, period: PeriodData) -> None:
    start, end = section_range(ws, "Поставки топлива/товара", 2, "Движение денежных средств")
    current = max(0, end - start + 1)
    entries = period.supply_entries
    ensure_capacity(ws, start, current, len(entries), start if current else start - 1)
    for row in range(start, start + len(entries)):
        clear_row(ws, row)
        apply_supply_row_alignment(ws, row)
    previous = None
    for idx, entry in enumerate(entries):
        row = start + idx
        resource = entry["resource"]
        resource_start = resource != previous
        apply_supply_row_borders(ws, row, resource_start)
        if resource_start:
            set_cell(ws, row, 1, entry.get("resource_label") or RESOURCE_LABELS.get(resource, ""))
            previous = resource
        set_cell(ws, row, 8, entry.get("supplier") or "")
        set_cell(ws, row, 26, entry.get("doc") or "")
        set_cell(ws, row, 33, fmt(entry.get("volume", 0), 3))
        set_cell(ws, row, 40, fmt(entry.get("mass", 0), 3))
        set_cell(ws, row, 48, fmt(entry.get("density", 0), 4))
        set_cell(ws, row, 54, fmt(entry.get("amount", 0), 2))


def write_cash_table(ws, period: PeriodData) -> None:
    cash_label = label_row(ws, "Движение денежных средств")
    chairman_label = row_contains(ws, "Председатель комиссии:", cash_label + 1)
    start = cash_label + 2
    current = max(0, chairman_label - start - 1)
    rows = period.cash_rows
    ensure_capacity(ws, start, current, len(rows), start if current else start - 1)
    chairman_label = row_contains(ws, "Председатель комиссии:", cash_label + 1)
    if rows:
        for row in range(start + 1, start + len(rows)):
            copy_row_format(ws, start, row)
    for row in range(start, start + len(rows)):
        clear_row(ws, row)

    def dt_text(value) -> str:
        if isinstance(value, datetime):
            return f"{value:%Y-%m-%d %H:%M:%S}"
        return "" if value is None else str(value)

    for idx, item in enumerate(rows):
        row = start + idx
        shift_key = item.get("ShiftKey") or f"{period.first_shift['ShiftKey']} - {period.last_shift['ShiftKey']}"
        shift_parts = [part.strip() for part in str(shift_key).split("-", 1)]
        single_shift = len(shift_parts) == 2 and shift_parts[0] == shift_parts[1]
        tall_cash_row = (not single_shift) and period.shift_range.shop_key in CASH_TALL_ROW_SHOPS
        ws.row_dimensions[row].height = 43.96 if tall_cash_row else 35.78
        set_cell(ws, row, 1, shift_key)
        set_cell(ws, row, 2, f"{dt_text(item.get('DatetimeShiftBegin'))}\n{dt_text(item.get('DatetimeShiftEnd'))}")
        set_cell(ws, row, 5, "" if item.get("CashBoxNumber") in (None, "") else str(item.get("CashBoxNumber")))
        set_cell(ws, row, 7, item.get("UserName") or "")
        set_cell(ws, row, 12, fmt(item.get("StartVal"), 2))
        set_cell(ws, row, 14, fmt(item.get("TotalValue"), 2))
        set_cell(ws, row, 16, fmt(item.get("FuelValue"), 2))
        set_cell(ws, row, 18, fmt(item.get("GoodsValue"), 2))
        set_cell(ws, row, 22, fmt(item.get("ServiceValue"), 2))
        set_cell(ws, row, 24, fmt(item.get("PaymentValue"), 2))
        set_cell(ws, row, 28, fmt(item.get("CardPayment"), 2))
        set_cell(ws, row, 30, fmt(item.get("RepaymentValue"), 2))
        set_cell(ws, row, 34, fmt(item.get("AcceptRemainsValue"), 2))
        set_cell(ws, row, 37, fmt(item.get("Amount"), 2))
        set_cell(ws, row, 41, fmt(item.get("BankAmount"), 2))
        set_cell(ws, row, 44, fmt(item.get("CashBoxAmount"), 2))
        set_cell(ws, row, 49, fmt(item.get("Delta"), 2))
        set_cell(ws, row, 51, fmt(item.get("BookEndVal"), 2))
        set_cell(ws, row, 55, fmt(item.get("EndVal"), 2))
        set_cell(ws, row, 58, fmt(item.get("EndDelta"), 2))
    normalize_cash_signature_borders(ws, cash_label, chairman_label)
    trim_rows_after(ws, chairman_label + 3)


def update_header(ws, period: PeriodData, save_date: date) -> None:
    set_cell(ws, 2, 1, f"Дата составления: {save_date:%d.%m.%Y}")
    set_cell(ws, 3, 1, f"Материально ответственные лица: {period.operator or ''}")
    set_cell(
        ws,
        6,
        1,
        (
            f"за период смен с {period.first_shift['ShiftKey']} "
            f"[{period.shift_range.begin:%d.%m.%Y %H:%M:%S} - {period.first_shift['DatetimeShiftEnd']:%d.%m.%Y %H:%M:%S}] "
            f"по {period.last_shift['ShiftKey']} "
            f"[{period.last_shift['DatetimeShiftBegin']:%d.%m.%Y %H:%M:%S} - {period.shift_range.end:%d.%m.%Y %H:%M:%S}]"
        ),
    )


def generate_workbook(cur, template_path: Path, output_path: Path, year_month: str, save_date: date) -> dict:
    shop_key = parse_shop_key(template_path)
    if shop_key is None:
        return {"file": str(template_path), "status": "skipped", "reason": "shop key not parsed"}
    year, month = [int(part) for part in year_month.split("-")]
    start = date(year, month, 1)
    end = date(year + (1 if month == 12 else 0), 1 if month == 12 else month + 1, 1)
    period = load_period_data(cur, shop_key, start, end)
    if period is None:
        return {"file": str(template_path), "shopKey": shop_key, "status": "skipped", "reason": "no shifts"}
    wb = load_workbook(template_path)
    ws = wb.active
    update_header(ws, period, save_date)
    resources = write_main_table(ws, period)
    write_counter_table(ws, period)
    write_tank_table(ws, period)
    write_supply_table(ws, period)
    write_cash_table(ws, period)
    normalize_section_label_borders(ws)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()
    return {
        "file": str(output_path),
        "shopKey": shop_key,
        "status": "generated",
        "resources": resources,
        "shiftFrom": period.first_shift["ShiftKey"],
        "shiftTo": period.last_shift["ShiftKey"],
    }


def extract_main(path: Path) -> dict:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    result = {}
    total_row = label_row(ws, "Итого по всем типам топлива, сом")
    row = 10
    while row < total_row:
        resource = resource_from_text(ws.cell(row, 1).value)
        if not resource:
            row += 1
            continue
        item = {}
        for name, col in MAIN_COLS.items():
            item[f"{name}_l"] = round_float(as_float(ws.cell(row, col).value, 0), 3)
            item[f"{name}_kg"] = round_float(as_float(ws.cell(row + 1, col).value, 0), 3)
            item[f"{name}_amt"] = round_float(as_float(ws.cell(row + 3, col).value, 0), 2)
        result[resource] = item
        row += 4
    wb.close()
    return result


def compare_main(reference: Path, generated: Path, tolerance_l: float = 0.005, tolerance_amt: float = 0.01) -> dict:
    ref = extract_main(reference)
    gen = extract_main(generated)
    diffs = []
    for resource in sorted(set(ref) | set(gen), key=lambda r: RESOURCE_ORDER.get(r, 999)):
        ref_item = ref.get(resource, {})
        gen_item = gen.get(resource, {})
        for key in sorted(set(ref_item) | set(gen_item)):
            a = float(ref_item.get(key) or 0)
            b = float(gen_item.get(key) or 0)
            tolerance = tolerance_amt if key.endswith("_amt") else tolerance_l
            if abs(a - b) > tolerance:
                diffs.append({"resource": resource, "field": key, "reference": a, "generated": b, "diff": round_float(b - a, 3)})
    return {"reference": str(reference), "generated": str(generated), "diffCount": len(diffs), "diffs": diffs[:200]}


def _parse_year_month(year_month: str) -> tuple[int, int]:
    if not re.fullmatch(r"\d{4}-\d{2}", str(year_month or "")):
        raise ValueError("invalid month format, expected YYYY-MM")
    year, month = [int(part) for part in year_month.split("-")]
    if month not in RU_MONTH_DIR:
        raise ValueError("invalid month")
    return year, month


def _parse_only(only: str | list[int] | tuple[int, ...] | set[int] | None) -> set[int] | None:
    if only is None or only == "":
        return None
    if isinstance(only, str):
        return {int(part) for part in re.split(r"[, ]+", only.strip()) if part}
    return {int(part) for part in only}


def generate_monthly_254_archive(
    year_month: str,
    *,
    apply: bool = True,
    only: str | list[int] | tuple[int, ...] | set[int] | None = None,
    template_dirs: list[Path] | None = None,
    archive_root: Path | None = None,
    staging_root: Path | None = None,
) -> dict:
    year, month = _parse_year_month(year_month)
    save_date = date(year + (1 if month == 12 else 0), 1 if month == 12 else month + 1, 1)
    root = archive_root or ARCHIVE_ROOT
    source_dir = month_dir(year_month, archive_root=root)
    work_root = staging_root or STAGING_ROOT
    generated_dir = work_root / "generated" / year_month
    backup_dir = work_root / "backup" / f"{year_month}_{datetime.now():%Y%m%d_%H%M%S}"
    only_set = _parse_only(only)
    templates = collect_templates(year_month, template_dirs or [], archive_root=root)
    summary = {
        "month": year_month,
        "sourceDir": str(source_dir),
        "generatedDir": str(generated_dir),
        "apply": apply,
        "generated": [],
        "compares": [],
    }

    conn = connect()
    try:
        cur = conn.cursor()
        for template in templates:
            shop_key = parse_shop_key(template)
            if only_set and shop_key not in only_set:
                continue
            output = generated_dir / template.name
            result = generate_workbook(cur, template, output, year_month, save_date)
            if template.parent != source_dir:
                result["templateSource"] = str(template)
            summary["generated"].append(result)
            if shop_key in (1, 2) and output.exists():
                summary["compares"].append(compare_main(template, output))
            if apply and output.exists():
                target = source_dir / output.name
                backup_dir.mkdir(parents=True, exist_ok=True)
                if target.exists():
                    shutil.copy2(target, backup_dir / target.name)
                    result["backup"] = str(backup_dir / target.name)
                else:
                    result["backup"] = None
                    result["targetWasMissing"] = True
                source_dir.mkdir(parents=True, exist_ok=True)
                shutil.copy2(output, target)
                result["appliedTo"] = str(target)
    finally:
        conn.close()

    summary_path = work_root / f"summary_{year_month}_{'apply' if apply else 'stage'}.json"
    summary_path.parent.mkdir(parents=True, exist_ok=True)
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    generated_count = sum(1 for result in summary["generated"] if result.get("status") == "generated")
    skipped_count = sum(1 for result in summary["generated"] if result.get("status") == "skipped")
    applied_count = sum(1 for result in summary["generated"] if result.get("appliedTo"))
    return {
        "status": "ok",
        "month": year_month,
        "sourceDir": str(source_dir),
        "generatedDir": str(generated_dir),
        "backupDir": str(backup_dir) if applied_count else None,
        "summaryPath": str(summary_path),
        "generatedCount": generated_count,
        "skippedCount": skipped_count,
        "applied": apply,
        "appliedCount": applied_count,
        "items": summary["generated"],
        "compares": summary["compares"],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--month", default="2026-07")
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--only", default="")
    parser.add_argument("--template-dir", action="append", default=[])
    args = parser.parse_args()
    result = generate_monthly_254_archive(
        args.month,
        apply=args.apply,
        only=args.only,
        template_dirs=[Path(value) for value in args.template_dir],
    )
    print(
        json.dumps(
            {
                "summary": result["summaryPath"],
                "generated": result["generatedCount"],
                "skipped": result["skippedCount"],
                "applied": result["applied"],
                "appliedCount": result["appliedCount"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
