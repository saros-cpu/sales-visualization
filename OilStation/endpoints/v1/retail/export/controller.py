from copy import copy
from pathlib import Path

from flask import Blueprint, jsonify, request, send_file
from loguru import logger

from db import get_connection
from export_excel import (
    RESOURCE_NAMES,
    _round_half_up,
)
from endpoints.v1.retail.export.service.workbook import (
    _can_reuse_report_workbook,
    _latest_report_workbook,
    _weekly_report_workbook,
)
from endpoints.v1.retail.export.service.generator import generate_monthly_excel, generate_range_summary_excel
from datetime import datetime, timedelta

from endpoints.v1.retail.analytics.controller import _build_self_use_month_response

export_api = Blueprint("export_api_v2", __name__)

SELF_USE_COMPANY_IDS = {
    1, 2, 3, 4, 5, 6, 7, 8, 9, 10,
    11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21,
    35, 40, 43, 44, 45, 46, 55, 56,
}
def _self_use_float(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _self_use_weight_kg(record):
    return _self_use_float(record.get('mass'))


def _group_self_use_totals_by_product(totals):
    groups = {}
    for (product, plate), values in totals.items():
        product_name = product or '-'
        groups.setdefault(product_name, []).append({
            'plate': plate or '-',
            'volume': values['volume'],
            'weightKg': values['weightKg'],
        })

    for items in groups.values():
        items.sort(key=lambda item: item['plate'])
    return groups


def _compact_self_use_text(value):
    return ''.join(str(value or '').upper().split())


def _self_use_is_generator(record):
    plate = _compact_self_use_text(record.get('plateNumber'))
    if plate in {'ГЕНЕРАТОР', '发电机'}:
        return True
    try:
        group_key = int(record.get('personGroup') or 0)
        self_use_number = int(record.get('selfUseNumber') or 0)
    except (TypeError, ValueError):
        return False
    return group_key == 9002 and self_use_number == 1


def _self_use_vehicle_remark(record):
    if _self_use_is_generator(record):
        return '发电机'
    try:
        person_key = int(record.get('personKey') or 0)
    except (TypeError, ValueError):
        person_key = 0
    return '公司车辆' if person_key in SELF_USE_COMPANY_IDS else '员工车辆'


def _group_self_use_totals_by_station(records):
    totals = {}
    for rec in records:
        try:
            station_order = int(rec.get('shopKey') or 999999)
        except (TypeError, ValueError):
            station_order = 999999
        key = (
            station_order,
            str(rec.get('stationName') or '-'),
            str(rec.get('stationShort') or ''),
            str(rec.get('plateNumber') or '-'),
            str(rec.get('productName') or '-'),
            _self_use_vehicle_remark(rec),
        )
        item = totals.setdefault(key, {'volume': 0.0, 'weightKg': 0.0, 'amount': 0.0})
        item['volume'] += _self_use_float(rec.get('volume'))
        item['weightKg'] += _self_use_weight_kg(rec)
        item['amount'] += _self_use_float(rec.get('amount'))
    return sorted(totals.items(), key=lambda item: item[0])


def _group_self_use_station_blocks(records):
    blocks = {}
    for (station_order, station_name, station_short, plate, product, remark), values in _group_self_use_totals_by_station(records):
        station_key = (station_order, station_name, station_short)
        blocks.setdefault(station_key, []).append({
            'plate': plate,
            'product': product,
            'remark': remark,
            'volume': values['volume'],
            'weightKg': values['weightKg'],
            'amount': values['amount'],
        })
    return [
        ((station_name, station_short), items)
        for (station_order, station_name, station_short), items in sorted(blocks.items(), key=lambda item: item[0])
    ]


def _has_cyrillic(value):
    return any('\u0400' <= char <= '\u052f' for char in str(value or ''))


def _apply_self_use_russian_font(cell):
    if not _has_cyrillic(cell.value):
        return
    font = copy(cell.font)
    font.name = 'Times New Roman'
    cell.font = font


def _apply_self_use_russian_fonts(ws):
    for row in ws.iter_rows():
        for cell in row:
            _apply_self_use_russian_font(cell)


def _apply_self_use_filter(ws, ref):
    ws.auto_filter.ref = ref
    ws.freeze_panes = 'A2'


def _text_width(value):
    text = str(value or '')
    if not text:
        return 0
    return max(sum(2 if ord(char) > 127 else 1 for char in part) for part in text.splitlines())


def _auto_fit_self_use_columns(ws, min_width=8, max_width=60):
    for col_cells in ws.columns:
        col_letter = col_cells[0].column_letter
        max_len = 0
        for cell in col_cells:
            if cell.value is None:
                continue
            max_len = max(max_len, _text_width(cell.value))
        if max_len:
            ws.column_dimensions[col_letter].width = min(max(max_len + 5, min_width), max_width)


def _write_self_use_station_sheet(wb, records, styles):
    ws = wb.create_sheet('按加油站统计')
    headers = ['加油站', '加油站俄语名称', '车牌号', '油品名称', '加油量公斤(kg)', '加油量升(L)', '金额(сом)', '备注']
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = styles['header_font']
        cell.fill = styles['header_fill']
        cell.alignment = styles['center_align']
        cell.border = styles['thin_border']

    row_idx = 2
    grand_weight = 0.0
    grand_volume = 0.0
    grand_amount = 0.0
    for (station_name, station_short), items in _group_self_use_station_blocks(records):
        block_start = row_idx
        station_weight = sum(item['weightKg'] for item in items)
        station_volume = sum(item['volume'] for item in items)
        station_amount = sum(item['amount'] for item in items if item['remark'] == '员工车辆')
        grand_weight += station_weight
        grand_volume += station_volume
        grand_amount += station_amount

        for item in items:
            employee_amount = _round_half_up(item['amount'], 2) if item['remark'] == '员工车辆' else ''
            row_values = [
                station_name,
                station_short,
                item['plate'],
                item['product'],
                _round_half_up(item['weightKg'], 2),
                _round_half_up(item['volume'], 2),
                employee_amount,
                item['remark'],
            ]
            for col_idx, value in enumerate(row_values, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = styles['center_align']
                cell.border = styles['thin_border']
                if col_idx in (5, 6, 7) and value != '':
                    cell.number_format = '#,##0.00'
            row_idx += 1

        subtotal_row = row_idx
        subtotal_values = [
            station_name,
            station_short,
            '小计',
            '',
            _round_half_up(station_weight, 2),
            _round_half_up(station_volume, 2),
            _round_half_up(station_amount, 2) if station_amount else '',
            '',
        ]
        for col_idx, value in enumerate(subtotal_values, 1):
            cell = ws.cell(row=subtotal_row, column=col_idx, value=value)
            cell.font = styles['bold_font']
            cell.fill = styles['subtotal_fill']
            cell.alignment = styles['center_align']
            cell.border = styles['thin_border']
            if col_idx in (5, 6, 7) and value != '':
                cell.number_format = '#,##0.00'

        if subtotal_row > block_start:
            ws.merge_cells(start_row=block_start, start_column=1, end_row=subtotal_row, end_column=1)
            ws.merge_cells(start_row=block_start, start_column=2, end_row=subtotal_row, end_column=2)
            ws.cell(row=block_start, column=1).alignment = styles['center_align']
            ws.cell(row=block_start, column=2).alignment = styles['center_align']
        row_idx += 1

    total_row = row_idx
    total_values = {
        1: '总合计',
        5: _round_half_up(grand_weight, 2),
        6: _round_half_up(grand_volume, 2),
        7: _round_half_up(grand_amount, 2) if grand_amount else '',
    }
    for col_idx in range(1, 9):
        cell = ws.cell(row=total_row, column=col_idx, value=total_values.get(col_idx, ''))
        cell.font = styles['bold_font']
        cell.fill = styles['total_fill']
        cell.alignment = styles['center_align']
        cell.border = styles['thin_border']
        if col_idx in (5, 6, 7) and cell.value != '':
            cell.number_format = '#,##0.00'
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=4)

    _apply_self_use_russian_fonts(ws)
    _auto_fit_self_use_columns(ws)
    _apply_self_use_filter(ws, ws.dimensions)
    return ws


def register_export_routes(app):
    app.register_blueprint(export_api)


@export_api.route('/api/export/monthly/<year_month>')
def api_export_monthly_excel(year_month):
    force_refresh = str(request.args.get("refresh") or "").lower() in {"1", "true", "yes"}
    conn = get_connection()
    try:
        result = generate_monthly_excel(conn, year_month, force_refresh=force_refresh)
        return jsonify(result)
    finally:
        conn.close()


@export_api.route('/api/export/monthly/<year_month>/download')
def api_export_download(year_month):
    cached_workbook = _latest_report_workbook(year_month)
    if cached_workbook and cached_workbook.exists() and _can_reuse_report_workbook(year_month, cached_workbook):
        return send_file(cached_workbook, as_attachment=True, download_name=cached_workbook.name)

    conn = get_connection()
    try:
        result = generate_monthly_excel(conn, year_month, force_refresh=False)
        workbook_path = Path(result.get('path') or '')
    finally:
        conn.close()

    if not workbook_path.exists():
        return jsonify({'error': 'File not found, please generate first'}), 404
    return send_file(workbook_path, as_attachment=True, download_name=workbook_path.name)


def _parse_report_range():
    date_from = request.args.get('date_from')
    date_to = request.args.get('date_to')
    if not date_from or not date_to:
        return None, None, jsonify({'error': 'date_from and date_to are required'}), 400
    try:
        start_dt = datetime.strptime(date_from, '%Y-%m-%d')
        end_dt = datetime.strptime(date_to, '%Y-%m-%d')
    except ValueError:
        return None, None, jsonify({'error': 'invalid date format, expected YYYY-MM-DD'}), 400
    if end_dt < start_dt:
        return None, None, jsonify({'error': 'date_to must be greater than or equal to date_from'}), 400
    return date_from, date_to, None, None


@export_api.route('/api/export/retail/range')
def api_export_retail_range():
    date_from, date_to, error_response, status_code = _parse_report_range()
    if error_response is not None:
        return error_response, status_code
    force_refresh = str(request.args.get('refresh') or '').lower() in {'1', 'true', 'yes'}
    conn = get_connection()
    try:
        return jsonify(generate_range_summary_excel(conn, date_from, date_to, force_refresh=force_refresh))
    finally:
        conn.close()


@export_api.route('/api/export/retail/range/download')
def api_export_retail_range_download():
    date_from, date_to, error_response, status_code = _parse_report_range()
    if error_response is not None:
        return error_response, status_code

    workbook_path = _weekly_report_workbook(date_from, date_to)
    if not workbook_path or not workbook_path.exists():
        conn = get_connection()
        try:
            result = generate_range_summary_excel(conn, date_from, date_to, force_refresh=False)
            workbook_path = Path(result['path'])
        finally:
            conn.close()

    if not workbook_path.exists():
        return jsonify({'error': 'File not found, please generate first'}), 404
    return send_file(workbook_path, as_attachment=True, download_name=workbook_path.name)


@export_api.route('/api/export/self-use/<year_month>')
def api_export_self_use(year_month):
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    parts = year_month.split('-')
    if len(parts) != 2 or len(parts[0]) != 4 or len(parts[1]) != 2:
        return jsonify({'error': 'invalid year_month'}), 400

    snapshot = _build_self_use_month_response(year_month)
    if not snapshot:
        return jsonify({'error': 'self-use snapshot not found'}), 404

    wb = Workbook()
    ws = wb.active
    ws.title = '自用油'

    headers = ['日期', '站点', '加油站俄语名称', '油品', '车牌号', '数量(L)', '公斤(kg)', '金额(сом)']
    header_fill = PatternFill(start_color='CC0000', end_color='CC0000', fill_type='solid')
    header_font_white = Font(name='Times New Roman', bold=True, size=11, color='FFFFFF')
    subtotal_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
    total_fill = PatternFill(start_color='D9EAF7', end_color='D9EAF7', fill_type='solid')
    center_align = Alignment(horizontal='center', vertical='center')
    bold_font = Font(name='Times New Roman', bold=True, size=11)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    for i, rec in enumerate(snapshot.get('records', []), 2):
        values = [
            rec.get('date', ''),
            rec.get('stationName', ''),
            rec.get('stationShort', ''),
            rec.get('productName', ''),
            rec.get('plateNumber', ''),
            _round_half_up(rec.get('volume'), 2),
            _round_half_up(_self_use_weight_kg(rec), 2),
            _round_half_up(rec.get('amount'), 2),
        ]
        for col_idx, v in enumerate(values, 1):
            cell = ws.cell(row=i, column=col_idx, value=v)
            cell.border = thin_border
            if col_idx in (6, 7, 8):
                cell.number_format = '#,##0.00'

    _apply_self_use_russian_fonts(ws)
    _auto_fit_self_use_columns(ws)
    _apply_self_use_filter(ws, ws.dimensions)
    _write_self_use_station_sheet(
        wb,
        snapshot.get('records', []),
        {
            'header_font': header_font_white,
            'header_fill': header_fill,
            'center_align': center_align,
            'thin_border': thin_border,
            'bold_font': bold_font,
            'subtotal_fill': subtotal_fill,
            'total_fill': total_fill,
        },
    )

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)

    return send_file(bio, as_attachment=True, download_name=f'自用油_{year_month}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@export_api.route('/api/export/self-use')
def api_export_self_use_range():
    date_from = request.args.get('date_from')
    date_to = request.args.get('date_to')
    if not date_from or not date_to:
        return jsonify({'error': 'date_from and date_to are required'}), 400

    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from collections import defaultdict

    try:
        d = datetime.strptime(date_from, '%Y-%m-%d')
        end_dt = datetime.strptime(date_to, '%Y-%m-%d') + timedelta(days=1)
    except ValueError:
        return jsonify({'error': 'invalid date format, expected YYYY-MM-DD'}), 400

    months = []
    while d < end_dt:
        months.append(d.strftime('%Y-%m'))
        if d.month == 12:
            d = d.replace(year=d.year + 1, month=1, day=1)
        else:
            d = d.replace(month=d.month + 1, day=1)

    records = []
    for ym in months:
        snap = _build_self_use_month_response(ym)
        if not snap:
            continue
        for rec in snap.get('records', []):
            if date_from <= rec['date'] <= date_to:
                records.append(rec)

    header_fill = PatternFill(start_color='CC0000', end_color='CC0000', fill_type='solid')
    header_font = Font(name='Times New Roman', bold=True, size=11, color='FFFFFF')
    subtotal_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
    total_fill = PatternFill(start_color='D9EAF7', end_color='D9EAF7', fill_type='solid')
    center_align = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    wb = Workbook()

    # Sheet 1: 自用油 (detail)
    ws = wb.active
    ws.title = '自用油'

    headers = ['日期', '时间', '站点', '加油站俄语名称', '油品', '车牌号', '数量(L)', '公斤(kg)']
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    for i, rec in enumerate(records, 2):
        values = [
            rec['date'],
            rec['time'],
            rec['stationName'],
            rec.get('stationShort', ''),
            rec['productName'],
            rec['plateNumber'],
            _round_half_up(rec['volume'], 2),
            _round_half_up(_self_use_weight_kg(rec), 2),
        ]
        for col_idx, v in enumerate(values, 1):
            cell = ws.cell(row=i, column=col_idx, value=v)
            cell.alignment = center_align
            cell.border = thin_border
            if col_idx in (7, 8):
                cell.number_format = '#,##0.00'

    _apply_self_use_russian_fonts(ws)
    _auto_fit_self_use_columns(ws)
    _apply_self_use_filter(ws, ws.dimensions)

    # Sheet 2: 总表
    company = defaultdict(lambda: {'volume': 0.0, 'weightKg': 0.0})
    employee = defaultdict(lambda: {'volume': 0.0, 'weightKg': 0.0})

    for rec in records:
        pk = int(rec.get('personKey') or 0)
        key = (str(rec.get('productName') or '-'), str(rec.get('plateNumber') or '-'))
        target = company if pk in SELF_USE_COMPANY_IDS else employee
        target[key]['volume'] += _self_use_float(rec.get('volume'))
        target[key]['weightKg'] += _self_use_weight_kg(rec)

    company_groups = _group_self_use_totals_by_product(company)
    employee_groups = _group_self_use_totals_by_product(employee)
    products = sorted(set(company_groups) | set(employee_groups))

    def write_total_cell(row, col, value, *, font=None, fill=None, num_fmt=False):
        cell = ws2.cell(row=row, column=col, value=value)
        cell.alignment = center_align
        cell.border = thin_border
        if font:
            cell.font = font
        if fill:
            cell.fill = fill
        if num_fmt:
            cell.number_format = '#,##0.00'
        return cell

    def write_side_row(row, start_col, product, plate, volume, weight_kg, *, font=None, fill=None):
        values = [
            (start_col, product),
            (start_col + 1, plate),
            (start_col + 2, _round_half_up(volume, 2)),
            (start_col + 3, _round_half_up(weight_kg, 2)),
        ]
        for col, value in values:
            write_total_cell(row, col, value, font=font, fill=fill, num_fmt=col in (start_col + 2, start_col + 3))

    def merge_product_rows(start_row, end_row, product_col):
        if start_row and end_row > start_row:
            ws2.merge_cells(start_row=start_row, start_column=product_col, end_row=end_row, end_column=product_col)
            ws2.cell(row=start_row, column=product_col).alignment = center_align

    def side_sums(items):
        return (
            sum(item['volume'] for item in items),
            sum(item['weightKg'] for item in items),
        )

    company_total_volume = sum(values['volume'] for values in company.values())
    company_total_weight = sum(values['weightKg'] for values in company.values())
    employee_total_volume = sum(values['volume'] for values in employee.values())
    employee_total_weight = sum(values['weightKg'] for values in employee.values())

    ws2 = wb.create_sheet('总表')

    col_headers = {
        1: '公司车辆油品',
        2: '公司车辆车牌号',
        3: '公司车辆加油量(L)',
        4: '公司车辆公斤(kg)',
        5: '员工车辆油品',
        6: '员工车辆车牌号',
        7: '员工车辆加油量(L)',
        8: '员工车辆公斤(kg)',
    }
    for col, h in col_headers.items():
        cell = ws2.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    bold_font = Font(name='Times New Roman', bold=True, size=11)
    row = 2
    for product in products:
        company_items = company_groups.get(product, [])
        employee_items = employee_groups.get(product, [])
        company_start_row = row if company_items else None
        employee_start_row = row if employee_items else None
        max_rows = max(len(company_items), len(employee_items))
        for idx in range(max_rows):
            if idx < len(company_items):
                item = company_items[idx]
                write_side_row(row, 1, product, item['plate'], item['volume'], item['weightKg'])
            if idx < len(employee_items):
                item = employee_items[idx]
                write_side_row(row, 5, product, item['plate'], item['volume'], item['weightKg'])
            row += 1

        if company_items:
            volume, weight_kg = side_sums(company_items)
            write_side_row(row, 1, product, '小计', volume, weight_kg, font=bold_font, fill=subtotal_fill)
            merge_product_rows(company_start_row, row, 1)
        if employee_items:
            volume, weight_kg = side_sums(employee_items)
            write_side_row(row, 5, product, '小计', volume, weight_kg, font=bold_font, fill=subtotal_fill)
            merge_product_rows(employee_start_row, row, 5)
        row += 1

    write_side_row(row, 1, '合计', '', company_total_volume, company_total_weight, font=bold_font, fill=total_fill)
    write_side_row(row, 5, '合计', '', employee_total_volume, employee_total_weight, font=bold_font, fill=total_fill)

    _apply_self_use_russian_fonts(ws2)
    _auto_fit_self_use_columns(ws2)
    _apply_self_use_filter(ws2, f"A1:H{row}")
    _write_self_use_station_sheet(
        wb,
        records,
        {
            'header_font': header_font,
            'header_fill': header_fill,
            'center_align': center_align,
            'thin_border': thin_border,
            'bold_font': bold_font,
            'subtotal_fill': subtotal_fill,
            'total_fill': total_fill,
        },
    )

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)

    return send_file(bio, as_attachment=True,
                     download_name=f'自用油_{date_from}_至_{date_to}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
