import os
import json
from datetime import date, datetime, timedelta
from pathlib import Path

from flask import Blueprint, jsonify, request
from openpyxl import load_workbook
import pymysql

from business_time import (
    business_date_expr,
    business_period_condition_with_pruning,
    business_period_params,
    shift_business_date_expr,
    shift_join,
)
from cache_store import get_cached, set_cached
from constants import (
    CLOSED_STATIONS,
    DEPOT_COORDS,
    FORCE_NORMAL_STATIONS,
    MAP_ZERO_INVENTORY_STATIONS,
    SHOP_COORDS,
    SHOP_KEY_MAP,
)
from db import get_connection
from export_excel import RESOURCE_DENSITY, SK4_RSK_MAP, SK6_RSK_MAP, SK28_RSK_MAP, SK29_RSK_MAP, STATION_RK_CORRECTIONS

from endpoints.v1.retail.common.query import calc_yoy_mom, get_query_date_range, query_month_summary
from endpoints.v1.retail.common.currency import (
    attach_amount_cny_average,
    average_cny_kgs_for_period,
    cny_from_kgs_average,
    load_cny_rates as _common_load_cny_rates,
)
from endpoints.v1.retail.analytics.service.shared import (
    _resource_key_from_report_name,
    _station_key_from_report_cell,
    _to_float,
)
from endpoints.v1.retail.analytics.service.shared import *

PROJECT_ROOT = Path(__file__).resolve().parents[6]
OILSTATION_ROOT = Path(__file__).resolve().parents[5]

def cache_allowed():
    return request.args.get("refresh") not in {"1", "true", "yes"}


def _comparison_month(year: int, month: int, offset: int) -> str:
    value = year * 12 + (month - 1) + offset
    return f"{value // 12}-{value % 12 + 1:02d}"


def _snapshot_region_values(year_month: str) -> dict:
    payload = get_cached(f"regions:{year_month}") or get_cached(f"monthly:{year_month}") or {}
    return payload.get("regions") or {}


def _attach_region_snapshot_comparisons(year_month: str, payload: dict) -> bool:
    regions = payload.get("regions") or {}
    if not regions:
        return False
    try:
        year, month = [int(part) for part in year_month.split("-")]
    except (TypeError, ValueError):
        return False

    previous = _snapshot_region_values(_comparison_month(year, month, -1))
    last_year = _snapshot_region_values(_comparison_month(year, month, -12))
    changed = False
    for name, current in regions.items():
        for label, reference in (("mom", previous), ("yoy", last_year)):
            previous_item = reference.get(name) or {}
            previous_amount = float(previous_item.get("amount") or 0)
            previous_volume = float(previous_item.get("volume") or 0)
            current_amount = float(current.get("amount") or 0)
            current_volume = float(current.get("volume") or 0)
            sales = round((current_amount - previous_amount) / previous_amount * 100, 1) if previous_amount else 0
            volume = round((current_volume - previous_volume) / previous_volume * 100, 1) if previous_volume else 0
            if current.get(f"{label}Sales") != sales or current.get(f"{label}Volume") != volume:
                current[f"{label}Sales"] = sales
                current[f"{label}Volume"] = volume
                changed = True
    return changed


STATION_MAP_DETAILS_TTL = 600
STATION_MAP_DETAILS_CACHE_KEY = "station_map_details:latest"
FUELBALANCE_WAREHOUSE_RESOURCE_MAP = {
    6: {1: 106, 2: 102, 3: 103, 4: 104},
    19: {4: 108},
    22: {4: 108},
    24: {4: 108},
    28: {1: 102, 2: 103, 3: 104, 4: 109, 5: 110, 6: 108},
    29: {1: 102, 2: 103, 3: 106, 4: 108},
    30: {1: 102, 2: 103},
}
STATION_RESOURCE_ALLOWLIST = {
    30: {102, 103},
}


def _load_env_file(path):
    try:
        lines = path.read_text(encoding="utf-8-sig").splitlines()
    except OSError:
        return
    for line in lines:
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def _sales_db_config():
    _load_env_file(PROJECT_ROOT / "Sales" / ".env")
    _load_env_file(PROJECT_ROOT / "OilStation" / ".env")
    database = os.getenv("SALES_DB_NAME", os.getenv("DB_NAME", "sales_visualization"))
    if database == "ocw":
        database = "sales_visualization"
    return {
        "host": os.getenv("DB_HOST", "127.0.0.1"),
        "port": int(os.getenv("DB_PORT", "3306")),
        "user": os.getenv("DB_USER", "reader"),
        "password": os.getenv("DB_PASSWORD", ""),
        "database": database,
        "charset": os.getenv("DB_CHARSET", "utf8mb4"),
        "use_unicode": True,
        "autocommit": True,
    }


def _parse_day(value):
    if isinstance(value, date):
        return value
    return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()


def _load_cny_rates(start, end):
    return _common_load_cny_rates(start, end)


def _load_density_history():
    densities = {}
    try:
        conn = pymysql.connect(**_sales_db_config())
    except Exception:
        return densities
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT p.retail_resource_key, d.effective_date, d.density
            FROM product_density_history d
            JOIN products p ON p.id = d.product_id
            WHERE p.retail_resource_key IS NOT NULL
            ORDER BY p.retail_resource_key, d.effective_date
            """
        )
        for resource_key, effective_date, density in cur.fetchall():
            densities.setdefault(canonical_resource_key(resource_key), []).append(
                (_parse_day(effective_date), float(density))
            )
    except Exception:
        return {}
    finally:
        conn.close()
    return densities


def _density_for(density_history, resource_key, day):
    rk = canonical_resource_key(resource_key)
    selected = None
    for effective_date, density in density_history.get(rk, []):
        if effective_date <= day:
            selected = density
        else:
            break
    if selected is not None:
        return selected
    return RESOURCE_DENSITY.get(rk, 0)


def _monthly_cny_from_daily(daily_rows, start, end):
    rates = _load_cny_rates(start, end)
    total_cny = 0.0
    total_amount = 0.0
    missing = []
    for item in daily_rows or []:
        amount = float(item.get("amount") or 0)
        if not amount:
            continue
        total_amount += amount
        day = str(item.get("date") or "")[:10]
        rate = rates.get(day)
        if rate:
            total_cny += amount / rate
        else:
            missing.append(day)
    avg_cny = average_cny_kgs_for_period(start, end)
    return {
        "totalSalesCny": round(total_cny, 2),
        "totalSalesCnyAvg": cny_from_kgs_average(total_amount, avg_cny),
        "exchangeMissingDates": sorted(set(missing)),
        "exchangeComplete": not missing,
        "avgCnyKgs": avg_cny,
    }


def _total_tons_from_daily_resource(daily_resource_rows):
    density_history = _load_density_history()
    total = 0.0
    missing = []
    for item in daily_resource_rows or []:
        mass = retail_mass_from_item(item)
        if mass is not None:
            total += mass / 1000
            continue
        volume = float(item.get("volume") or 0)
        if not volume:
            continue
        try:
            business_day = _parse_day(item.get("date"))
        except (TypeError, ValueError):
            continue
        density = _density_for(density_history, item.get("resourceKey") or item.get("key"), business_day)
        if not density:
            missing.append(f"{item.get('resourceKey')}:{business_day.isoformat()}")
            continue
        total += volume * density / 1000
    return {
        "totalTons": round(total, 3),
        "densityMissingKeys": sorted(set(missing)),
        "densityComplete": not missing,
    }


def _total_tons_from_resource_summary(resources, business_day=None):
    density_history = _load_density_history()
    if business_day is None:
        business_day = date.today()
    total = 0.0
    missing = []
    for item in resources or []:
        mass = retail_mass_from_item(item)
        if mass is not None:
            total += mass / 1000
            continue
        volume = float(item.get("volume") or 0)
        if not volume:
            continue
        density = _density_for(density_history, item.get("resourceKey") or item.get("key"), business_day)
        if not density:
            missing.append(str(item.get("resourceKey") or item.get("key") or "unknown"))
            continue
        total += volume * density / 1000
    return {
        "totalTons": round(total, 3),
        "densityMissingKeys": sorted(set(missing)),
        "densityComplete": not missing,
    }


def _attach_monthly_finance_metrics(payload, daily_resource_rows=None):
    summary = payload.get("summary") or {}
    daily = payload.get("daily") or []
    start = summary.get("queryStart") or payload.get("queryStart") or f"{payload.get('month')}-01"
    end = summary.get("queryEnd") or payload.get("queryEndExclusive")
    if not end and daily:
        last_day = max(str(item.get("date"))[:10] for item in daily)
        end = (datetime.strptime(last_day, "%Y-%m-%d").date() + timedelta(days=1)).isoformat()
    if start and end:
        summary.update(_monthly_cny_from_daily(daily, start, end))
        avg_cny = summary.get("avgCnyKgs")
        for group_name in ("stations", "regions", "byResource"):
            rows = payload.get(group_name)
            if isinstance(rows, dict):
                rows = rows.values()
            if not isinstance(rows, (list, tuple)):
                continue
            for item in rows:
                attach_amount_cny_average(item, "amount", "amountCnyAvg", avg_cny)
    if daily_resource_rows:
        summary.update(_total_tons_from_daily_resource(daily_resource_rows))
    elif "totalTons" not in summary:
        try:
            business_day = _parse_day(start)
        except (TypeError, ValueError):
            business_day = None
        summary.update(_total_tons_from_resource_summary(payload.get("byResource") or [], business_day))
    payload["summary"] = summary
    return payload


def _attach_summary_cny_average(summary, start=None, end=None, amount_key="totalSales"):
    if not isinstance(summary, dict):
        return False
    start = start or summary.get("queryStart")
    end = end or summary.get("queryEndExclusive") or summary.get("queryEnd")
    avg_cny = average_cny_kgs_for_period(start, end)
    amount = summary.get(amount_key)
    if amount is None and amount_key != "totalAmount":
        amount = summary.get("totalAmount")
    changed = False
    amount_cny = cny_from_kgs_average(amount, avg_cny)
    for key in ("totalSalesCnyAvg", "totalAmountCnyAvg"):
        if summary.get(key) != amount_cny:
            summary[key] = amount_cny
            changed = True
    if avg_cny is not None and summary.get("avgCnyKgs") != avg_cny:
        summary["avgCnyKgs"] = avg_cny
        changed = True
    return changed


def _attach_period_cny_average(payload, row_groups=()):
    if not isinstance(payload, dict):
        return False
    start = payload.get("queryStart") or (payload.get("summary") or {}).get("queryStart")
    end = payload.get("queryEndExclusive") or payload.get("queryEnd") or (payload.get("summary") or {}).get("queryEnd")
    avg_cny = average_cny_kgs_for_period(start, end)
    changed = False
    if avg_cny is not None and payload.get("avgCnyKgs") != avg_cny:
        payload["avgCnyKgs"] = avg_cny
        changed = True
    summary = payload.get("summary")
    if isinstance(summary, dict):
        changed = _attach_summary_cny_average(summary, start, end, "totalAmount") or changed
    for group_name in row_groups:
        rows = payload.get(group_name)
        if isinstance(rows, dict):
            rows = rows.values()
        if not isinstance(rows, (list, tuple)):
            continue
        for item in rows:
            changed = attach_amount_cny_average(item, "amount", "amountCnyAvg", avg_cny) or changed
    return changed
def _latest_report_workbook_for_map():
    index_path = OILSTATION_ROOT / "exports" / ".monthly_report_index.json"
    if not index_path.exists():
        return None, None
    try:
        report_index = json.loads(index_path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError):
        return None, None
    if not isinstance(report_index, dict):
        return None, None
    for ym in sorted(report_index.keys(), reverse=True):
        item = report_index.get(ym) or {}
        stored_path = item.get("path") or item.get("file")
        if not stored_path:
            continue
        path = OILSTATION_ROOT / "exports" / stored_path
        if path.exists():
            return ym, path
    return None, None


def _report_workbook_snapshot_state(workbook_path):
    stat = workbook_path.stat()
    return {
        "source": "retail_report_snapshot",
        "sourceFile": workbook_path.name,
        "sourceMtime": stat.st_mtime,
        "snapshotGeneratedAt": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="minutes"),
    }


def _latest_completed_shift_marker_for_map():
    station_keys = sorted(SHOP_KEY_MAP.keys())
    if not station_keys:
        return None
    station_key_sql = ",".join(str(key) for key in station_keys)
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            f"""
            SELECT MAX(DatetimeShiftEnd)
            FROM shift
            WHERE ShopKey IN ({station_key_sql})
              AND DatetimeShiftBegin IS NOT NULL
              AND DatetimeShiftEnd IS NOT NULL
              AND DatetimeShiftEnd > DatetimeShiftBegin
              AND DatetimeShiftEnd >= DATE_SUB(NOW(), INTERVAL 180 DAY)
            """
        )
        row = cur.fetchone()
        return row[0] if row else None
    finally:
        conn.close()


def _station_map_details_cache_current(cached):
    if not isinstance(cached, dict):
        return True
    if not _station_map_details_has_depots(cached):
        return False
    if not _station_map_details_has_capacity(cached):
        return False
    if cached.get("source") == "live_database":
        # The in-memory cache TTL already bounds live DB staleness. Re-querying
        # the latest shift marker here made every warm map-details request as
        # slow as a cold request.
        return bool(cached.get("snapshotGeneratedAt") or cached.get("updatedAt"))
    _, workbook_path = _latest_report_workbook_for_map()
    if not workbook_path:
        return False
    try:
        snapshot_state = _report_workbook_snapshot_state(workbook_path)
    except OSError:
        return False
    try:
        cached_mtime = float(cached.get("sourceMtime"))
    except (TypeError, ValueError):
        cached_mtime = None
    return (
        cached.get("source") == snapshot_state["source"]
        and cached.get("sourceFile") == snapshot_state["sourceFile"]
        and cached_mtime is not None
        and abs(cached_mtime - snapshot_state["sourceMtime"]) < 0.001
    )


def _station_map_depots():
    return [
        {
            "name": name,
            "short": data.get("short", ""),
            "region": data.get("region", ""),
            "lng": data.get("lng"),
            "lat": data.get("lat"),
            "type": data.get("type", "depot"),
            "status": data.get("status", "normal"),
            "closedReason": data.get("closedReason", ""),
        }
        for name, data in DEPOT_COORDS.items()
    ]


def _station_map_details_has_depots(payload):
    expected = len(DEPOT_COORDS or {})
    if expected <= 0:
        return True
    depots = payload.get("depots") if isinstance(payload, dict) else []
    if not isinstance(depots, list) or len(depots) < expected:
        return False
    return all(
        isinstance(item, dict)
        and item.get("name")
        and item.get("lng") is not None
        and item.get("lat") is not None
        for item in depots
    )


def _station_map_details_has_capacity(payload):
    items = payload.get("items") if isinstance(payload, dict) else []
    if not isinstance(items, list):
        return False
    has_inventory_row = False
    has_known_capacity = False
    for item in items:
        rows = item.get("inventoryByResource") if isinstance(item, dict) else []
        for row in rows or []:
            if not isinstance(row, dict):
                continue
            has_inventory_row = True
            if "fillStatus" not in row:
                return False
            if row.get("fillStatus") in {"ok", "warning", "critical"}:
                has_known_capacity = True
    return not has_inventory_row or has_known_capacity


def _inventory_fill_status(percent):
    try:
        value = float(percent)
    except (TypeError, ValueError):
        return "unknown"
    if value < 30:
        return "critical"
    if value < 50:
        return "warning"
    return "ok"


def _station_resource_allowed(shop_key, resource_key):
    allowlist = STATION_RESOURCE_ALLOWLIST.get(int(shop_key or 0))
    return not allowlist or int(resource_key or 0) in allowlist


def _infer_current_fuelbalance_price_resource_key(price):
    try:
        p = float(price or 0)
    except (TypeError, ValueError):
        return 0
    if p <= 0:
        return 0
    if p <= 60:
        return 108
    if p <= 80:
        return 102
    if p <= 95:
        return 103
    if p <= 105:
        return 106
    if p <= 115:
        return 104
    return 106


def _infer_fuelbalance_resource_key(shop_key, resource_code, resource_name, price, warehouse_key=None):
    # Resource names avoid price-threshold drift after fuel price changes.
    sk = int(shop_key or 0)
    name_key = _resource_key_from_report_name(resource_name)
    if name_key in KNOWN_RESOURCE_KEYS:
        return canonical_resource_key(name_key)
    try:
        wk = int(warehouse_key or 0)
    except (TypeError, ValueError):
        wk = 0
    mapped = FUELBALANCE_WAREHOUSE_RESOURCE_MAP.get(sk, {}).get(wk)
    if mapped:
        return canonical_resource_key(mapped)
    try:
        rc = int(resource_code or 0)
    except (TypeError, ValueError):
        rc = 0
    if rc < 100:
        return _infer_current_fuelbalance_price_resource_key(price)
    return infer_balance_resource_key(shop_key, resource_code, price, None)


def _fuelbalance_observed_capacity_by_tank(rows):
    values_by_tank = {}
    for shop_key, warehouse_key, resource_key, resource_name, price, volume in rows:
        try:
            sk = int(shop_key)
            wk = str(warehouse_key)
            amount = float(volume or 0)
        except (TypeError, ValueError):
            continue
        if amount <= 0:
            continue
        rk = _infer_fuelbalance_resource_key(sk, resource_key, resource_name, price, warehouse_key)
        if rk < 100 or rk not in KNOWN_RESOURCE_KEYS or not _station_resource_allowed(sk, rk):
            continue
        values_by_tank.setdefault((sk, wk, rk), []).append(amount)

    observed = {}
    for key, values in values_by_tank.items():
        values.sort()
        percentile_index = int((len(values) - 1) * 0.95)
        observed[key] = round(values[percentile_index], 2)
    return observed


def _latest_fuelbalance_tank_rows(rows):
    latest = {}
    anonymous_index = 0
    for row in rows:
        shop_key, warehouse_key, resource_key, resource_name, price, volume, capacity, created_at = row
        try:
            sk = int(shop_key)
        except (TypeError, ValueError):
            continue
        if warehouse_key is None:
            anonymous_index += 1
            tank_key = (sk, f"row:{anonymous_index}")
        else:
            tank_key = (sk, str(warehouse_key))
        score = (
            1 if float(capacity or 0) > 0 else 0,
            float(volume or 0),
            float(price or 0),
        )
        existing = latest.get(tank_key)
        if existing is None or score > existing[0]:
            latest[tank_key] = (score, row)
    return [item[1] for item in latest.values()]


def _fuelbalance_capacity_by_resource_from_rows(rows, warehouseinfo_keys=None, observed_by_tank=None):
    warehouseinfo_keys = warehouseinfo_keys or set()
    observed_by_tank = observed_by_tank or {}
    latest_tanks = {}
    anonymous_index = 0
    for row in rows:
        shop_key, warehouse_key, resource_key, resource_name, price, capacity = row[:6]
        volume = row[6] if len(row) > 6 else None
        try:
            sk = int(shop_key)
            cap = float(capacity or 0)
        except (TypeError, ValueError):
            continue
        if warehouse_key is None:
            anonymous_index += 1
            tank_key = (sk, f"row:{anonymous_index}")
        else:
            tank_key = (sk, str(warehouse_key))
        rk = _infer_fuelbalance_resource_key(sk, resource_key, resource_name, price, warehouse_key)
        if rk < 100 or rk not in KNOWN_RESOURCE_KEYS or not _station_resource_allowed(sk, rk):
            continue
        if cap <= 0:
            cap = max(float(volume or 0), float(observed_by_tank.get((sk, tank_key[1], rk)) or 0))
        if cap <= 0:
            continue
        existing = latest_tanks.get(tank_key)
        if existing is None or cap > existing[-1]:
            latest_tanks[tank_key] = (sk, rk, cap)

    result = {}
    for sk, rk, capacity in latest_tanks.values():
        key = (sk, rk)
        if key in warehouseinfo_keys:
            continue
        result[key] = round(result.get(key, 0.0) + capacity, 2)
    return result


def _station_inventory_capacity_by_resource():
    station_keys = sorted(SHOP_KEY_MAP.keys())
    if not station_keys:
        return {}
    station_key_sql = ",".join(str(key) for key in station_keys)
    result = {}
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            f"""
            SELECT ShopKey, ResourceKey, ResourceShopKey,
                   SUM(COALESCE(NULLIF(Capacity, 0), NULLIF(RefCapacity, 0), 0)) AS capacity
            FROM warehouseinfo
            WHERE ShopKey IN ({station_key_sql})
              AND COALESCE(NULLIF(Capacity, 0), NULLIF(RefCapacity, 0), 0) > 0
            GROUP BY ShopKey, ResourceKey, ResourceShopKey
            """
        )
        for shop_key, resource_key, resource_shop_key, capacity in cur.fetchall():
            sk = int(shop_key)
            rk = infer_balance_resource_key(sk, resource_key, None, resource_shop_key)
            if rk < 100 or rk not in KNOWN_RESOURCE_KEYS:
                continue
            key = (sk, rk)
            result[key] = round(result.get(key, 0.0) + float(capacity or 0), 2)
        warehouseinfo_keys = set(result)
        cur.execute(
            f"""
            SELECT fb.ShopKey, fb.WarehouseKey, fb.ResourceCode, fb.ResourceName,
                   fb.CurrentPrice, fb.VolumeMaximum, fb.VolumeBalance
            FROM fuelbalance fb
            JOIN (
                SELECT ShopKey, WarehouseKey, MAX(CreateDatetime) AS max_dt
                FROM fuelbalance
                WHERE ShopKey IN ({station_key_sql})
                  AND CreateDatetime >= DATE_SUB(NOW(), INTERVAL 365 DAY)
                  AND (VolumeMaximum > 0 OR VolumeBalance > 0 OR CurrentPrice > 0)
                GROUP BY ShopKey, WarehouseKey
            ) latest
              ON latest.ShopKey = fb.ShopKey
             AND latest.WarehouseKey = fb.WarehouseKey
             AND latest.max_dt = fb.CreateDatetime
            WHERE fb.ShopKey IN ({station_key_sql})
              AND (fb.VolumeMaximum > 0 OR fb.VolumeBalance > 0 OR fb.CurrentPrice > 0)
            """
        )
        latest_rows = cur.fetchall()
        cur.execute(
            f"""
            SELECT ShopKey, WarehouseKey, ResourceCode, ResourceName, CurrentPrice, VolumeBalance
            FROM fuelbalance
            WHERE ShopKey IN ({station_key_sql})
              AND VolumeBalance > 0
              AND CreateDatetime >= DATE_SUB(NOW(), INTERVAL 365 DAY)
            """
        )
        observed = _fuelbalance_observed_capacity_by_tank(cur.fetchall())
        result.update(_fuelbalance_capacity_by_resource_from_rows(latest_rows, warehouseinfo_keys, observed))
    finally:
        conn.close()
    return result


def _attach_station_inventory_levels(details):
    capacity_by_resource = _station_inventory_capacity_by_resource()
    iterable = details.values() if isinstance(details, dict) else details
    for detail in iterable:
        if not isinstance(detail, dict):
            continue
        shop_key = int(detail.get("shopKey") or 0)
        total_capacity = 0.0
        total_inventory = 0.0
        for row in detail.get("inventoryByResource") or []:
            if not isinstance(row, dict):
                continue
            resource_key = int(row.get("resourceKey") or 0)
            inventory = round(float(row.get("inventory") or 0), 2)
            capacity = round(float(capacity_by_resource.get((shop_key, resource_key)) or 0), 2)
            total_inventory += inventory
            if capacity <= 0:
                row["fillStatus"] = "unknown"
                continue
            fill_percent = round(inventory / capacity * 100, 1)
            row["capacity"] = capacity
            row["fillPercent"] = fill_percent
            row["fillStatus"] = _inventory_fill_status(fill_percent)
            total_capacity += capacity
        if total_capacity > 0:
            station_percent = round(total_inventory / total_capacity * 100, 1)
            detail["inventoryCapacity"] = round(total_capacity, 2)
            detail["inventoryFillPercent"] = station_percent
            detail["inventoryFillStatus"] = _inventory_fill_status(station_percent)
        else:
            detail["inventoryFillStatus"] = "unknown"
    return details


def _apply_station_map_inventory_overrides(details):
    zero_keys = set(MAP_ZERO_INVENTORY_STATIONS or set())
    if not zero_keys:
        return False
    changed = False
    iterable = details.values() if isinstance(details, dict) else details
    for detail in iterable or []:
        if not isinstance(detail, dict):
            continue
        try:
            shop_key = int(detail.get("shopKey") or 0)
        except (TypeError, ValueError):
            continue
        if shop_key not in zero_keys:
            continue
        reason = CLOSED_STATIONS.get(shop_key) or "停业，库存清零"
        if float(detail.get("inventory") or 0) != 0:
            changed = True
        detail["inventory"] = 0.0
        detail["inventoryZeroedByMapRule"] = True
        detail["inventoryZeroReason"] = reason
        total_capacity = 0.0
        for row in detail.get("inventoryByResource") or []:
            if not isinstance(row, dict):
                continue
            if float(row.get("inventory") or 0) != 0:
                changed = True
            row["inventory"] = 0.0
            capacity = float(row.get("capacity") or 0)
            if capacity > 0:
                row["fillPercent"] = 0.0
                row["fillStatus"] = "critical"
                total_capacity += capacity
            elif "fillPercent" in row:
                row["fillPercent"] = 0.0
                row["fillStatus"] = "unknown"
        if total_capacity <= 0:
            total_capacity = float(detail.get("inventoryCapacity") or 0)
        if total_capacity > 0:
            detail["inventoryCapacity"] = round(total_capacity, 2)
            detail["inventoryFillPercent"] = 0.0
            detail["inventoryFillStatus"] = "critical"
        else:
            detail["inventoryFillStatus"] = "unknown"
    return changed


def _report_sheet_last_nonempty_day(workbook):
    last_day = 0
    for sheet_name in workbook.sheetnames:
        if not sheet_name.endswith("日") or not sheet_name[:2].isdigit():
            continue
        sheet = workbook[sheet_name]
        for row in sheet.iter_rows(min_row=5, values_only=True):
            if _to_float(row[18] if len(row) > 18 else None) or _to_float(row[20] if len(row) > 20 else None):
                last_day = max(last_day, int(sheet_name[:2]))
                break
    return last_day


def _station_map_details_from_report_snapshot():
    ym, workbook_path = _latest_report_workbook_for_map()
    if not workbook_path:
        return None
    workbook = None
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        last_day = _report_sheet_last_nonempty_day(workbook)
        sheet_name = f"{last_day:02d}日" if last_day else None
        if not sheet_name or sheet_name not in workbook.sheetnames:
            sheet_name = "总表" if "总表" in workbook.sheetnames else workbook.sheetnames[0]
        sheet = workbook[sheet_name]
        snapshot_date = f"{ym}-{last_day:02d}" if ym and last_day else ym
        snapshot_state = _report_workbook_snapshot_state(workbook_path)
        snapshot_generated_at = snapshot_state["snapshotGeneratedAt"]
        details = {
            shop_key: {
                "shopKey": shop_key,
                "inventory": 0.0,
                "inventoryByResource": [],
                "prices": [],
                "inventoryAsOf": snapshot_date,
                "priceAsOf": snapshot_date,
                "snapshotGeneratedAt": snapshot_generated_at,
            }
            for shop_key in sorted(SHOP_KEY_MAP.keys())
        }
        rows_by_station = {shop_key: {} for shop_key in details}
        current_shop_key = None
        for row in sheet.iter_rows(min_row=5, values_only=True):
            station_key = _station_key_from_report_cell(row[0] if len(row) > 0 else None)
            if station_key is not None:
                current_shop_key = station_key
            if current_shop_key is None or current_shop_key not in rows_by_station:
                continue
            product_name = str(row[3] if len(row) > 3 and row[3] is not None else "").strip()
            if not product_name or "合计" in product_name or "итого" in product_name.lower():
                continue
            rk = _resource_key_from_report_name(product_name)
            if rk not in KNOWN_RESOURCE_KEYS:
                continue
            if sheet_name == "总表":
                inventory = _to_float(row[23] if len(row) > 23 else None)
                price = _to_float(row[22] if len(row) > 22 else None)
            else:
                inventory = _to_float(row[21] if len(row) > 21 else None)
                direct_price = _to_float(row[11] if len(row) > 11 else None)
                sale_volume = _to_float(row[18] if len(row) > 18 else None)
                sale_amount = _to_float(row[20] if len(row) > 20 else None)
                price = direct_price if direct_price else (round(sale_amount / sale_volume, 2) if sale_volume else None)
            if not inventory and not price:
                continue
            bucket = rows_by_station[current_shop_key].setdefault(rk, {
                "resourceKey": rk,
                "name": RESOURCE_NAMES.get(rk, f"Other-{rk}"),
                "inventory": 0.0,
                "price": None,
            })
            bucket["inventory"] += inventory
            if price is not None:
                bucket["price"] = round(price, 2)

        for shop_key, by_resource in rows_by_station.items():
            rows = sorted(
                (
                    {**item, "inventory": round(float(item.get("inventory") or 0), 2)}
                    for item in by_resource.values()
                    if float(item.get("inventory") or 0) or item.get("price") is not None
                ),
                key=lambda item: RESOURCE_ORDER_INDEX.get(item["resourceKey"], 999),
            )
            detail = details[shop_key]
            detail["inventoryByResource"] = rows
            detail["inventory"] = round(sum(item["inventory"] for item in rows), 2)
            detail["prices"] = [
                {"resourceKey": item["resourceKey"], "name": item["name"], "price": item["price"]}
                for item in rows
                if item.get("price") is not None
            ]
        _attach_station_inventory_levels(details)
        _apply_station_map_inventory_overrides(details)
        return {
            "items": list(details.values()),
            "depots": _station_map_depots(),
            "updatedAt": datetime.now().isoformat(timespec="seconds"),
            "sourceSheet": sheet_name,
            **snapshot_state,
        }
    except Exception:
        return None
    finally:
        if workbook is not None:
            workbook.close()


def _station_map_details_payload():
    cached = get_cached(STATION_MAP_DETAILS_CACHE_KEY)
    if cached and cache_allowed() and _station_map_details_cache_current(cached):
        if _apply_station_map_inventory_overrides(cached.get("items") or []):
            set_cached(STATION_MAP_DETAILS_CACHE_KEY, cached, ttl=STATION_MAP_DETAILS_TTL)
        return cached

    snapshot_payload = _station_map_details_from_report_snapshot()
    if snapshot_payload:
        set_cached(STATION_MAP_DETAILS_CACHE_KEY, snapshot_payload, ttl=STATION_MAP_DETAILS_TTL)
        return snapshot_payload

    station_keys = sorted(SHOP_KEY_MAP.keys())
    station_key_sql = ",".join(str(key) for key in station_keys)
    details = {
        shop_key: {
            "shopKey": shop_key,
            "inventory": 0.0,
            "inventoryByResource": [],
            "prices": [],
            "inventoryAsOf": None,
            "priceAsOf": None,
        }
        for shop_key in station_keys
    }
    inventory = {}
    prices = {}
    price_dates = {}
    inventory_dates = {}
    latest_shifts = {}
    latest_inventory_marker = None

    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            f"""
            SELECT ShopKey, ShiftKey, DatetimeShiftEnd
            FROM shift
            WHERE ShopKey IN ({station_key_sql})
              AND DatetimeShiftBegin IS NOT NULL
              AND DatetimeShiftEnd IS NOT NULL
              AND DatetimeShiftEnd > DatetimeShiftBegin
              AND DatetimeShiftEnd >= DATE_SUB(NOW(), INTERVAL 180 DAY)
            ORDER BY ShopKey, DatetimeShiftEnd DESC
            """
        )
        for shop_key, shift_key, shift_end in cur.fetchall():
            latest_shifts.setdefault(int(shop_key), (int(shift_key), shift_end))

        cur.execute(
            f"""
            SELECT fb.ShopKey, fb.WarehouseKey, fb.ResourceCode, fb.ResourceName,
                   fb.CurrentPrice, fb.VolumeBalance, fb.VolumeMaximum, fb.CreateDatetime
            FROM fuelbalance fb
            JOIN (
                SELECT ShopKey, WarehouseKey, MAX(CreateDatetime) AS max_dt
                FROM fuelbalance
                WHERE ShopKey IN ({station_key_sql})
                  AND CreateDatetime >= DATE_SUB(NOW(), INTERVAL 365 DAY)
                  AND (VolumeBalance > 0 OR CurrentPrice > 0 OR VolumeMaximum > 0)
                GROUP BY ShopKey, WarehouseKey
            ) latest
              ON latest.ShopKey = fb.ShopKey
             AND latest.WarehouseKey = fb.WarehouseKey
             AND latest.max_dt = fb.CreateDatetime
            WHERE fb.ShopKey IN ({station_key_sql})
              AND (fb.VolumeBalance > 0 OR fb.CurrentPrice > 0 OR fb.VolumeMaximum > 0)
            ORDER BY fb.ShopKey, fb.WarehouseKey
            """
        )
        for shop_key, warehouse_key, resource_key, resource_name, price, volume_balance, volume_maximum, created_at in _latest_fuelbalance_tank_rows(cur.fetchall()):
            sk = int(shop_key)
            rk = _infer_fuelbalance_resource_key(sk, resource_key, resource_name, price, warehouse_key)
            if rk < 100 or rk not in KNOWN_RESOURCE_KEYS or not _station_resource_allowed(sk, rk):
                continue
            volume = round(float(volume_balance or 0), 2)
            key = (sk, rk)
            inventory[key] = round(inventory.get(key, 0.0) + volume, 2)
            if price:
                prices[key] = round(float(price), 2)
                price_dates[key] = created_at
            previous_dt = inventory_dates.get(sk)
            if previous_dt is None or (created_at and created_at > previous_dt):
                inventory_dates[sk] = created_at
            if latest_inventory_marker is None or (created_at and created_at > latest_inventory_marker):
                latest_inventory_marker = created_at

        if not inventory and latest_shifts:
            conditions = " OR ".join(["(ShopKey = %s AND ShiftKey = %s)" for _ in latest_shifts])
            params = []
            for shop_key, (shift_key, _) in latest_shifts.items():
                params.extend([shop_key, shift_key])
            cur.execute(
                f"""
                SELECT ShopKey, ResourceKey, EndBalance, EndPrice, ResourceShopKey, ShiftKey
                FROM balance
                WHERE BeginMass IS NOT NULL
                  AND ({conditions})
                ORDER BY ShopKey, ResourceKey, ResourceShopKey
                """,
                params,
            )
            for shop_key, resource_key, end_balance, end_price, resource_shop_key, shift_key in cur.fetchall():
                sk = int(shop_key)
                rk = infer_balance_resource_key(sk, resource_key, end_price, resource_shop_key)
                if rk < 100 or rk not in KNOWN_RESOURCE_KEYS or not _station_resource_allowed(sk, rk):
                    continue
                key = (sk, rk)
                inventory[key] = round(inventory.get(key, 0.0) + float(end_balance or 0), 2)
                as_of = latest_shifts.get(sk, (None, None))[1]
                if end_price:
                    prices[key] = round(float(end_price), 2)
                    price_dates[key] = as_of
                inventory_dates[sk] = as_of
    finally:
        conn.close()

    for shop_key in station_keys:
        resource_keys = sorted(
            {rk for sk, rk in inventory if sk == shop_key} | {rk for sk, rk in prices if sk == shop_key},
            key=lambda rk: RESOURCE_ORDER_INDEX.get(rk, 999),
        )
        rows = []
        for rk in resource_keys:
            volume = round(float(inventory.get((shop_key, rk), 0) or 0), 2)
            price = prices.get((shop_key, rk))
            if not volume and price is None:
                continue
            rows.append({
                "resourceKey": rk,
                "name": RESOURCE_NAMES.get(rk, f"Other-{rk}"),
                "inventory": volume,
                "price": price,
            })
        total_inventory = round(sum(item["inventory"] for item in rows), 2)
        detail = details[shop_key]
        detail["inventory"] = total_inventory
        detail["inventoryByResource"] = rows
        detail["prices"] = [
            {"resourceKey": item["resourceKey"], "name": item["name"], "price": item["price"]}
            for item in rows
            if item.get("price") is not None
        ]
        inv_dt = inventory_dates.get(shop_key)
        price_dt = max((price_dates.get((shop_key, rk)) for rk in resource_keys if price_dates.get((shop_key, rk))), default=None)
        detail["inventoryAsOf"] = inv_dt.isoformat(sep=" ", timespec="minutes") if hasattr(inv_dt, "isoformat") else None
        detail["priceAsOf"] = price_dt.isoformat(sep=" ", timespec="minutes") if hasattr(price_dt, "isoformat") else None
        detail["snapshotGeneratedAt"] = detail["inventoryAsOf"] or detail["priceAsOf"]

    _attach_station_inventory_levels(details)
    _apply_station_map_inventory_overrides(details)
    latest_shift_end = max((shift_end for _, shift_end in latest_shifts.values() if shift_end), default=None)
    latest_marker = latest_inventory_marker or latest_shift_end
    generated_at = (
        latest_marker.isoformat(sep=" ", timespec="minutes")
        if hasattr(latest_marker, "isoformat")
        else datetime.now().isoformat(timespec="minutes")
    )
    source_mtime = latest_marker.timestamp() if hasattr(latest_marker, "timestamp") else None

    payload = {
        "items": list(details.values()),
        "depots": _station_map_depots(),
        "updatedAt": datetime.now().isoformat(timespec="seconds"),
        "source": "live_database",
        "snapshotGeneratedAt": generated_at,
        "sourceMtime": source_mtime,
    }
    set_cached(STATION_MAP_DETAILS_CACHE_KEY, payload, ttl=STATION_MAP_DETAILS_TTL)
    return payload
def _snapshot_month_labels(year):
    now = datetime.now()
    max_month = 12
    if year == now.year:
        max_month = now.month
    elif year > now.year:
        max_month = 0
    return [f"{year}-{month:02d}" for month in range(1, max_month + 1)]


def _monthly_snapshots_for_year(year):
    snapshots = []
    for label in _snapshot_month_labels(year):
        payload = get_cached(f"monthly:{label}")
        if payload:
            if retail_repair_payload_mass_from_report(label, payload):
                set_cached(f"monthly:{label}", payload)
            snapshots.append(payload)
    return snapshots


def _resource_map_from_rows(rows):
    bucket = {}
    for item in rows or []:
        add_resource_bucket(
            bucket,
            item.get("resourceKey") or item.get("key"),
            item.get("volume"),
            item.get("amount"),
            retail_mass_from_item(item),
        )
    return bucket


def _sum_resource_maps(*maps):
    merged = {}
    for resource_map in maps:
        for rk, values in (resource_map or {}).items():
            add_resource_bucket(merged, rk, values.get("volume"), values.get("amount"), values.get("mass"))
    return merged


def _resource_items_from_map(resource_map, previous=None, month_data=None, prev_month_data=None):
    previous = previous or {}
    month_data = month_data or {}
    prev_month_data = prev_month_data or {}
    resource_keys = sorted(set(resource_map) | set(previous) | set(month_data) | set(prev_month_data))
    resources = []
    for rk in resource_keys:
        cur_d = resource_map.get(rk, {"volume": 0, "amount": 0, "mass": 0})
        prev_d = previous.get(rk, {"volume": 0, "amount": 0, "mass": 0})
        mon_d = month_data.get(rk, {"volume": 0, "amount": 0, "mass": 0})
        prev_mon_d = prev_month_data.get(rk, {"volume": 0, "amount": 0, "mass": 0})
        mass = round(float(cur_d.get("mass") or 0), 3)
        latest_month_mass = round(float(mon_d.get("mass") or 0), 3)
        resources.append({
            "resourceKey": rk,
            "key": rk,
            "name": RESOURCE_NAMES.get(rk, f"Other-{rk}"),
            "volume": round(cur_d["volume"], 2),
            "mass": mass,
            "tons": round(mass / 1000, 3),
            "amount": round(cur_d["amount"], 2),
            "color": RESOURCE_COLORS.get(rk, "#999999"),
            "yoyVolume": calc_change(cur_d["volume"], prev_d["volume"]),
            "yoySales": calc_change(cur_d["amount"], prev_d["amount"]),
            "momVolume": calc_change(mon_d["volume"], prev_mon_d["volume"]),
            "momSales": calc_change(mon_d["amount"], prev_mon_d["amount"]),
            "latestMonthVolume": round(mon_d["volume"], 2),
            "latestMonthMass": latest_month_mass,
            "latestMonthTons": round(latest_month_mass / 1000, 3),
            "latestMonthAmount": round(mon_d["amount"], 2),
        })
    resources.sort(key=lambda item: RESOURCE_ORDER_INDEX.get(item["key"], 999))
    return resources


def _dashboard_yearly_from_monthly_snapshots(year):
    months = _monthly_snapshots_for_year(year)
    if not months:
        return None

    stations = {}
    resource_totals = {}
    query_start = f"{year}-01-01"
    query_end = query_start

    for payload in months:
        summary = payload.get("summary") or {}
        query_start = min(query_start, summary.get("queryStart") or query_start)
        query_end = max(query_end, summary.get("queryEnd") or query_end)
        for station in payload.get("stations") or []:
            shop_key = station.get("shopKey")
            if shop_key is None:
                continue
            bucket = stations.setdefault(int(shop_key), {
                "txnCount": 0,
                "volume": 0.0,
                "mass": 0.0,
                "amount": 0.0,
                "cash": 0.0,
                "card": 0.0,
                "cashLiters": 0.0,
                "cardLiters": 0.0,
                "bonus": 0.0,
            })
            bucket["txnCount"] += int(station.get("txnCount") or 0)
            bucket["volume"] += float(station.get("volume") or 0)
            bucket["mass"] += float(station.get("mass") or 0) or float(station.get("tons") or 0) * 1000
            bucket["amount"] += float(station.get("amount") or station.get("totalSales") or 0)
            bucket["cash"] += float(station.get("cash") or station.get("cashSales") or 0)
            bucket["card"] += float(station.get("card") or station.get("cardSales") or 0)
            bucket["cashLiters"] += float(station.get("cashLiters") or 0)
            bucket["cardLiters"] += float(station.get("cardLiters") or 0)
            bucket["bonus"] += float(station.get("bonus") or station.get("bonusCost") or 0)
        resource_totals = _sum_resource_maps(resource_totals, _resource_map_from_rows(payload.get("byResource")))

    stations_list = []
    for shop_key, info in SHOP_KEY_MAP.items():
        coords = SHOP_COORDS.get(shop_key, {"lng": 0, "lat": 0})
        data = stations.get(shop_key, {})
        amount = round(float(data.get("amount") or 0), 2)
        stations_list.append({
            "shopKey": shop_key,
            "name": info["name"],
            "short": info["short"],
            "region": info["region"],
            "lng": coords.get("lng", 0),
            "lat": coords.get("lat", 0),
            "txnCount": int(data.get("txnCount") or 0),
            "volume": round(float(data.get("volume") or 0), 2),
            "mass": round(float(data.get("mass") or 0), 3),
            "tons": round(float(data.get("mass") or 0) / 1000, 3),
            "amount": amount,
            "cash": round(float(data.get("cash") or 0), 2),
            "card": round(float(data.get("card") or 0), 2),
            "cashLiters": round(float(data.get("cashLiters") or 0), 2),
            "cardLiters": round(float(data.get("cardLiters") or 0), 2),
            "bonus": round(float(data.get("bonus") or 0), 2),
            "status": station_status_for(shop_key, amount),
            "closedReason": CLOSED_STATIONS.get(shop_key, ""),
        })

    regions = {}
    for station in stations_list:
        region = station["region"]
        bucket = regions.setdefault(region, {
            "amount": 0.0,
            "volume": 0.0,
            "tons": 0.0,
            "cash": 0.0,
            "card": 0.0,
            "bonus": 0.0,
            "stationCount": 0,
            "activeCount": 0,
        })
        bucket["stationCount"] += 1
        bucket["amount"] += station["amount"]
        bucket["volume"] += station["volume"]
        bucket["tons"] += station.get("tons") or 0
        bucket["cash"] += station["cash"]
        bucket["card"] += station["card"]
        bucket["bonus"] += station["bonus"]
        if station["amount"] > 0:
            bucket["activeCount"] += 1
    for data in regions.values():
        for key in ("amount", "volume", "cash", "card", "bonus"):
            data[key] = round(data[key], 2)
        data["tons"] = round(data.get("tons") or 0, 3)

    result = {
        "year": year,
        "queryStart": query_start,
        "queryEndExclusive": query_end,
        "stations": stations_list,
        "depots": [
            {
                "name": name,
                "short": data["short"],
                "region": data["region"],
                "lng": data["lng"],
                "lat": data["lat"],
                "type": data.get("type", "depot"),
                "status": data.get("status", "normal"),
                "closedReason": data.get("closedReason", ""),
            }
            for name, data in DEPOT_COORDS.items()
        ],
        "regions": regions,
        "byResource": format_resource_rows([
            (rk, values["volume"], values["amount"], values.get("mass", 0))
            for rk, values in resource_totals.items()
        ]),
        "snapshot": True,
    }
    _attach_period_cny_average(result, ("stations", "regions", "byResource"))
    return result


def _resources_yearly_from_monthly_snapshots(year, month=None):
    if month:
        current_label = f"{year}-{month:02d}"
        current_payload = get_cached(f"monthly:{current_label}")
        if not current_payload:
            return None
        if retail_repair_payload_mass_from_report(current_label, current_payload):
            set_cached(f"monthly:{current_label}", current_payload)
        prev_label = f"{year - 1}-12" if month == 1 else f"{year}-{month - 1:02d}"
        yoy_label = f"{year - 1}-{month:02d}"
        current = _resource_map_from_rows(current_payload.get("byResource"))
        previous = _resource_map_from_rows((get_cached(f"monthly:{yoy_label}") or {}).get("byResource"))
        prev_month_data = _resource_map_from_rows((get_cached(f"monthly:{prev_label}") or {}).get("byResource"))
        latest_month = current_label
        prev_month = prev_label
        period_start = current_payload.get("summary", {}).get("queryStart", f"{year}-{month:02d}-01")
        period_end = current_payload.get("summary", {}).get("queryEnd", period_start)
        resources = _resource_items_from_map(current, previous, current, prev_month_data)
    else:
        months = _monthly_snapshots_for_year(year)
        if not months:
            return None
        current = {}
        for payload in months:
            current = _sum_resource_maps(current, _resource_map_from_rows(payload.get("byResource")))
        previous = {}
        for payload in _monthly_snapshots_for_year(year - 1):
            previous = _sum_resource_maps(previous, _resource_map_from_rows(payload.get("byResource")))
        latest_payload = months[-1]
        latest_month = latest_payload.get("month")
        prev_month = None
        if latest_month:
            y, m = int(latest_month[:4]), int(latest_month[5:7])
            prev_month = f"{y - 1}-12" if m == 1 else f"{y}-{m - 1:02d}"
        month_data = _resource_map_from_rows(latest_payload.get("byResource"))
        prev_month_data = _resource_map_from_rows((get_cached(f"monthly:{prev_month}") or {}).get("byResource")) if prev_month else {}
        period_start = f"{year}-01-01"
        period_end = max((payload.get("summary", {}).get("queryEnd") or period_start for payload in months), default=period_start)
        resources = _resource_items_from_map(current, previous, month_data, prev_month_data)

    result = {
        "year": year,
        "month": month,
        "periodLabel": f"{year}-{month:02d}" if month else f"{year} 全年",
        "queryStart": period_start,
        "queryEndExclusive": period_end,
        "latestMonth": latest_month,
        "prevMonth": prev_month,
        "resources": resources,
        "summary": {
            "totalVolume": round(sum(item["volume"] for item in resources), 2),
            "totalTons": round(sum(item["tons"] for item in resources), 3),
            "totalAmount": round(sum(item["amount"] for item in resources), 2),
            "resourceCount": len([item for item in resources if item["amount"] > 0 or item["volume"] > 0]),
        },
        "snapshot": True,
    }
    _attach_period_cny_average(result, ("resources",))
    return result
