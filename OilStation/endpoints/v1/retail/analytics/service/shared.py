import os
import json
from datetime import date, datetime, timedelta
from pathlib import Path

from flask import Blueprint, jsonify, request
from openpyxl import load_workbook
import pymysql

from business_time import (
    business_date_expr,
    business_period_condition_with_pruning,
    business_period_params,
    shift_business_date_expr,
    shift_join,
)
from cache_store import get_cached, set_cached
from constants import (
    CLOSED_STATIONS,
    DEPOT_COORDS,
    FORCE_NORMAL_STATIONS,
    SHOP_COORDS,
    SHOP_KEY_MAP,
)
from db import get_connection
from export_excel import (
    EXPORT_DIR,
    RESOURCE_DENSITY,
    SK4_RSK_MAP,
    SK6_RSK_MAP,
    SK19_22_24_RSK_MAP,
    SK28_RSK_MAP,
    SK29_RSK_MAP,
    SK30_RSK_MAP,
    STATION_RK_CORRECTIONS,
    STATION_TANK_RESOURCE_MAP,
    STATION_TANK_RESOURCE_MAP_PRIMARY,
    _allow_fixed_density_fallback,
    _selling_mass_expr,
)

from endpoints.v1.retail.common.query import calc_yoy_mom, get_query_date_range, query_month_summary

RESOURCE_NAMES = {
    102: 'АИ 80',
    103: 'АИ 92',
    104: 'АИ 95',
    106: 'Д/Т',
    108: 'СПБТ',
    109: 'Д/Т',
    110: 'ДТ -35',
}
RESOURCE_COLORS = {
    102: '#00e5a0',
    103: '#00d4ff',
    104: '#a78bfa',
    106: '#ffa726',
    108: '#ff6b6b',
    109: '#ffb74d',
    110: '#ff8a65',
}
RESOURCE_ORDER = [102, 103, 104, 106, 110, 108]
RESOURCE_ORDER_INDEX = {rk: index for index, rk in enumerate(RESOURCE_ORDER)}
RESOURCE_ALIASES = {
    109: 106,
}
KNOWN_RESOURCE_KEYS = set(RESOURCE_NAMES)

def station_status_for(shop_key, amount):
    if shop_key in FORCE_NORMAL_STATIONS:
        return 'normal'
    if shop_key in CLOSED_STATIONS:
        return 'closed'
    return 'normal' if amount > 0 else 'closed'


def resource_key_expr(alias="s"):
    return report_resource_key_expr(alias)


def canonical_resource_key(resource_key):
    try:
        rk = int(resource_key or 0)
    except (TypeError, ValueError):
        return 0
    return RESOURCE_ALIASES.get(rk, rk)


def add_resource_bucket(target, resource_key, volume, amount, mass=None):
    rk = canonical_resource_key(resource_key)
    if rk < 100 or rk not in KNOWN_RESOURCE_KEYS:
        return
    bucket = target.setdefault(rk, {'volume': 0.0, 'amount': 0.0, 'mass': 0.0})
    bucket['volume'] += float(volume or 0)
    bucket['amount'] += float(amount or 0)
    if mass is not None:
        bucket['mass'] += float(mass or 0)


def report_volume_expr(alias="s"):
    return (
        "SUM(CASE "
        f"WHEN ({alias}.DerivedPaymentTypeKey IS NULL OR {alias}.DerivedPaymentTypeKey NOT IN (2, 3, 4)) "
        f" AND {alias}.BasePaymentTypeKey IN (1, 2) THEN {alias}.Volume ELSE 0 END) "
        f"+ SUM(CASE WHEN {alias}.DerivedPaymentTypeKey = 3 THEN {alias}.Volume ELSE 0 END)"
    )


def report_amount_expr(alias="s"):
    return (
        "SUM(CASE "
        f"WHEN ({alias}.DerivedPaymentTypeKey IS NULL OR {alias}.DerivedPaymentTypeKey NOT IN (2, 3, 4)) "
        f" AND {alias}.BasePaymentTypeKey IN (1, 2) THEN {alias}.ShopCost ELSE 0 END) "
        f"+ SUM(CASE WHEN {alias}.DerivedPaymentTypeKey = 3 THEN {alias}.ShopCost ELSE 0 END)"
    )


def report_mass_expr(alias="s", period_start=None):
    use_ai80_fallback = not _allow_fixed_density_fallback(period_start)
    selling_mass = _selling_mass_expr(alias, use_ai80_missing_density_fallback=use_ai80_fallback)
    return (
        "SUM(CASE "
        f"WHEN ({alias}.DerivedPaymentTypeKey IS NULL OR {alias}.DerivedPaymentTypeKey NOT IN (2, 3, 4)) "
        f" AND {alias}.BasePaymentTypeKey IN (1, 2) THEN {selling_mass} ELSE 0 END) "
        f"+ SUM(CASE WHEN {alias}.DerivedPaymentTypeKey = 3 THEN {selling_mass} ELSE 0 END)"
    )


def report_cash_amount_expr(alias="s"):
    return (
        "SUM(CASE "
        f"WHEN ({alias}.DerivedPaymentTypeKey IS NULL OR {alias}.DerivedPaymentTypeKey NOT IN (2, 3, 4)) "
        f" AND {alias}.BasePaymentTypeKey = 1 "
        f" AND COALESCE({alias}.DerivedPaymentTypeKey, 0) <> 5 THEN {alias}.ShopCost ELSE 0 END)"
    )


def report_card_amount_expr(alias="s"):
    return (
        "SUM(CASE "
        f"WHEN ({alias}.DerivedPaymentTypeKey IS NULL OR {alias}.DerivedPaymentTypeKey NOT IN (2, 3, 4)) "
        f" AND ({alias}.BasePaymentTypeKey = 2 OR {alias}.DerivedPaymentTypeKey = 5) "
        f"THEN {alias}.ShopCost ELSE 0 END)"
    )


def report_cash_volume_expr(alias="s"):
    return (
        "SUM(CASE "
        f"WHEN ({alias}.DerivedPaymentTypeKey IS NULL OR {alias}.DerivedPaymentTypeKey NOT IN (2, 3, 4)) "
        f" AND {alias}.BasePaymentTypeKey = 1 "
        f" AND COALESCE({alias}.DerivedPaymentTypeKey, 0) <> 5 THEN {alias}.Volume ELSE 0 END)"
    )


def report_card_volume_expr(alias="s"):
    return (
        "SUM(CASE "
        f"WHEN ({alias}.DerivedPaymentTypeKey IS NULL OR {alias}.DerivedPaymentTypeKey NOT IN (2, 3, 4)) "
        f" AND ({alias}.BasePaymentTypeKey = 2 OR {alias}.DerivedPaymentTypeKey = 5) "
        f"THEN {alias}.Volume ELSE 0 END)"
    )


def report_bonus_amount_expr(alias="s"):
    return (
        "SUM(CASE "
        f"WHEN ({alias}.DerivedPaymentTypeKey IS NULL OR {alias}.DerivedPaymentTypeKey NOT IN (2, 3, 4)) "
        f"THEN COALESCE({alias}.BonusOutCost, 0) ELSE 0 END)"
    )


def report_resource_key_expr(alias="s"):
    """快照/KPI 侧油品推断表达式（与 export_excel 固化映射字典同源）。

    前置权威层由 STATION_TANK_RESOURCE_MAP_PRIMARY 生成（含全部液化气罐）；
    SK4 汽柴油罐（t1/t2/t5）沿用快照侧既有无条件口径；尾部注入固化兜底层
    （带生效日），与月报 selling_inferred_resource_key_sql 同一套罐位事实。
    """
    a = f"{alias}."
    front = []
    for (sk, rsk), code in sorted(STATION_TANK_RESOURCE_MAP_PRIMARY.items()):
        front.append(f"WHEN {a}ShopKey = {sk} AND {a}ResourceShopKey = {rsk} THEN {code} ")
    # SK4 汽柴油罐：快照侧既有无条件口径（与日销表 _infer_resource_code_daily 一致）
    front.append(f"WHEN {a}ShopKey = 4 AND {a}ResourceShopKey = 1 THEN 103 ")
    front.append(f"WHEN {a}ShopKey = 4 AND {a}ResourceShopKey = 2 THEN 104 ")
    front.append(f"WHEN {a}ShopKey = 4 AND {a}ResourceShopKey = 5 THEN 106 ")
    pins = []
    for (sk, rsk), (code, dfrom) in sorted(STATION_TANK_RESOURCE_MAP.items()):
        cond = f"WHEN {a}ShopKey = {sk} AND {a}ResourceShopKey = {rsk}"
        if dfrom:
            cond += f" AND {a}TransactionDatetime >= '{dfrom}'"
        pins.append(f"{cond} THEN {code} ")
    return (
        "CASE "
        + "".join(front)
        + f"WHEN {alias}.ShopKey IN (7, 15) AND {alias}.ResourceKey = 104 THEN 102 "
        f"WHEN {alias}.ResourceKey IN (102, 103, 104, 106, 108, 109, 110) THEN {alias}.ResourceKey "
        f"WHEN {alias}.BasePrice <= 0 THEN {alias}.ResourceKey "
        + "".join(pins)
        + f"WHEN {alias}.BasePrice <= 45 THEN 108 "
        f"WHEN {alias}.BasePrice <= 65 THEN 102 "
        f"WHEN {alias}.BasePrice <= 78 THEN 103 "
        f"WHEN {alias}.BasePrice <= 83 THEN 104 "
        "ELSE 106 "
        "END"
    )


def calc_change(current, previous):
    if not previous or previous <= 0:
        return None
    return round((float(current or 0) - float(previous or 0)) / float(previous) * 100, 1)


def format_resource_rows(rows):
    merged = {}
    for row in rows:
        mass = row[3] if len(row) > 3 else None
        add_resource_bucket(merged, row[0], row[1], row[2], mass)
    result = []
    for rk, values in merged.items():
        volume = round(values['volume'], 2)
        amount = round(values['amount'], 2)
        mass = round(float(values.get('mass') or 0), 3)
        result.append({
            'resourceKey': rk,
            'key': rk,
            'name': RESOURCE_NAMES.get(rk, f'Other-{rk}'),
            'volume': volume,
            'mass': mass,
            'tons': round(mass / 1000, 3),
            'amount': amount,
            'color': RESOURCE_COLORS.get(rk, '#999999'),
        })
    result.sort(key=lambda item: RESOURCE_ORDER_INDEX.get(item['key'], 999))
    return result


def infer_balance_resource_key(shop_key, resource_key, price=None, resource_shop_key=None):
    try:
        sk = int(shop_key or 0)
        rk = int(resource_key or 0)
    except (TypeError, ValueError):
        return 0
    try:
        rsk = int(resource_shop_key or 0)
    except (TypeError, ValueError):
        rsk = 0

    if sk == 28 and rsk in SK28_RSK_MAP:
        return canonical_resource_key(SK28_RSK_MAP[rsk])
    if sk == 6 and rsk in SK6_RSK_MAP:
        return canonical_resource_key(SK6_RSK_MAP[rsk])
    if sk == 4 and rsk in SK4_RSK_MAP:
        return canonical_resource_key(SK4_RSK_MAP[rsk])
    if sk == 30 and rsk:
        return canonical_resource_key(SK30_RSK_MAP.get(rsk, 0))
    if sk == 29 and rsk in SK29_RSK_MAP:
        return canonical_resource_key(SK29_RSK_MAP[rsk])
    if sk in (19, 22, 24) and rsk in SK19_22_24_RSK_MAP:
        return canonical_resource_key(SK19_22_24_RSK_MAP[rsk])
    # 液化气罐前置权威（2026-09-04 固化映射）：覆盖既有链未含的罐（SK2/3/5/8/13/14/26）
    if rsk and (sk, rsk) in STATION_TANK_RESOURCE_MAP_PRIMARY:
        return canonical_resource_key(STATION_TANK_RESOURCE_MAP_PRIMARY[(sk, rsk)])
    if sk == 30 and rk not in {102, 103}:
        return 0
    corrected = STATION_RK_CORRECTIONS.get((sk, rk))
    if corrected:
        return canonical_resource_key(corrected)
    if rk in KNOWN_RESOURCE_KEYS:
        return canonical_resource_key(rk)

    try:
        p = float(price or 0)
    except (TypeError, ValueError):
        p = 0
    if p <= 0:
        return 0
    if p <= 45:
        return 108
    if p <= 65:
        return 102
    if p <= 78:
        return 103
    if p <= 83:
        return 104
    return 106


def _to_float(value):
    if value is None or value == "":
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _station_key_from_report_cell(value):
    try:
        station_key = int(float(value))
    except (TypeError, ValueError):
        return None
    return station_key if station_key in SHOP_KEY_MAP else None


def _resource_key_from_report_name(value):
    name = str(value or "").strip().lower()
    if not name:
        return 0
    if "-35" in name:
        return 110
    if "спбт" in name or "сжиж" in name or "газ" in name or "ãàç" in name or "gaz" in name:
        return 108
    if "95" in name:
        return 104
    if "92" in name:
        return 103
    if "80" in name:
        return 102
    if "д/т" in name or "диз" in name:
        return 106
    return 0
RETAIL_DENSITY_MIN = 0.4
RETAIL_DENSITY_MAX = 1.1


def retail_mass_is_plausible(volume, mass):
    try:
        volume_value = float(volume or 0)
        mass_value = float(mass or 0)
    except (TypeError, ValueError):
        return False
    if volume_value <= 0:
        return mass_value >= 0
    if mass_value <= 0:
        return False
    density = mass_value / volume_value
    return RETAIL_DENSITY_MIN <= density <= RETAIL_DENSITY_MAX


def retail_mass_from_item(item):
    volume = item.get("volume") if isinstance(item, dict) else 0
    if isinstance(item, dict) and item.get("tons") is not None:
        try:
            mass = float(item.get("tons") or 0) * 1000
        except (TypeError, ValueError):
            mass = 0
        if retail_mass_is_plausible(volume, mass):
            return mass
    if isinstance(item, dict) and item.get("mass") is not None:
        try:
            mass = float(item.get("mass") or 0)
        except (TypeError, ValueError):
            mass = 0
        if retail_mass_is_plausible(volume, mass):
            return mass
    return None


def retail_payload_mass_needs_refresh(payload):
    if not isinstance(payload, dict):
        return False
    resource_items = []
    station_items = []
    for key in ("byResource", "resources"):
        values = payload.get(key)
        if isinstance(values, list):
            resource_items.extend(values)
    values = payload.get("stations")
    if isinstance(values, list):
        station_items.extend(values)
    items = [*resource_items, *station_items]
    if not items:
        return False
    resource_tons = 0.0
    has_volume_without_mass = False
    for item in items:
        if not isinstance(item, dict):
            continue
        try:
            volume = float(item.get("volume") or item.get("latestMonthVolume") or 0)
        except (TypeError, ValueError):
            volume = 0
        if volume <= 0:
            continue
        if "mass" not in item:
            has_volume_without_mass = True
            continue
        mass = retail_mass_from_item(item)
        if mass is None:
            has_volume_without_mass = True
            continue
        if item in resource_items:
            resource_tons += mass / 1000
    if has_volume_without_mass:
        return True
    summary = payload.get("summary") or {}
    try:
        summary_tons = float(summary.get("totalTons") or 0)
    except (TypeError, ValueError):
        summary_tons = 0
    if summary_tons > 0 and resource_tons > 0:
        return abs(summary_tons - resource_tons) > max(1.0, summary_tons * 0.05)
    return False


def retail_daily_payload_mass_needs_refresh(payload):
    if not isinstance(payload, dict):
        return False
    items = []
    daily = payload.get("daily")
    if isinstance(daily, list):
        items.extend(daily)
    stations = payload.get("stations")
    if isinstance(stations, dict):
        for values in stations.values():
            if isinstance(values, list):
                items.extend(values)
    if not items:
        return False
    for item in items:
        if not isinstance(item, dict):
            continue
        try:
            volume = float(item.get("volume") or 0)
        except (TypeError, ValueError):
            volume = 0
        if volume <= 0:
            continue
        has_mass_field = "mass" in item or "tons" in item
        try:
            mass = float(item.get("mass") or 0)
        except (TypeError, ValueError):
            mass = 0
        try:
            tons = float(item.get("tons") or 0)
        except (TypeError, ValueError):
            tons = 0
        if not has_mass_field or (mass <= 0 and tons <= 0):
            return True
    return False


def _retail_report_workbook_for_month(year_month):
    try:
        export_dir = Path(EXPORT_DIR)
        index_path = export_dir / ".monthly_report_index.json"
        if index_path.exists():
            index = json.loads(index_path.read_text(encoding="utf-8-sig"))
            item = index.get(year_month) if isinstance(index, dict) else None
            stored_path = (item.get("path") or item.get("file")) if isinstance(item, dict) else None
            if stored_path:
                path = export_dir / stored_path
                if path.exists():
                    return path
    except (OSError, json.JSONDecodeError, TypeError, ValueError):
        return None
    return None


def _retail_report_column_map(sheet):
    header = next(sheet.iter_rows(min_row=4, max_row=4, values_only=True), ())
    labels = [str(value or "").strip() for value in header]

    def find(*candidates, default=0):
        for candidate in candidates:
            for index, label in enumerate(labels):
                if candidate in label:
                    return index
        return default

    return {
        "product": find("油品名", default=3),
        "sales_tons": find("销售总吨", "销售量吨", default=20),
    }


def _retail_report_sales_tons(year_month):
    workbook_path = _retail_report_workbook_for_month(year_month)
    if not workbook_path:
        return {}, {}
    workbook = None
    resource_tons = {}
    station_tons = {}
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        sheet_name = "总表" if "总表" in workbook.sheetnames else workbook.sheetnames[0]
        sheet = workbook[sheet_name]
        columns = _retail_report_column_map(sheet)
        product_idx = columns["product"]
        tons_idx = columns["sales_tons"]
        current_shop_key = None
        for row in sheet.iter_rows(min_row=5, values_only=True):
            station_key = _station_key_from_report_cell(row[0] if len(row) > 0 else None)
            if station_key is not None:
                current_shop_key = station_key
            product_name = str(row[product_idx] if len(row) > product_idx and row[product_idx] is not None else "").strip()
            if not product_name or "合计" in product_name or "итого" in product_name.lower():
                continue
            resource_key = _resource_key_from_report_name(product_name)
            if not resource_key:
                continue
            tons = _to_float(row[tons_idx] if len(row) > tons_idx else None)
            resource_tons[resource_key] = resource_tons.get(resource_key, 0.0) + tons
            if current_shop_key is not None:
                station_tons[current_shop_key] = station_tons.get(current_shop_key, 0.0) + tons
    except Exception:
        return {}, {}
    finally:
        if workbook is not None:
            workbook.close()
    return (
        {key: round(value, 3) for key, value in resource_tons.items()},
        {key: round(value, 3) for key, value in station_tons.items()},
    )


def retail_repair_payload_mass_from_report(year_month, payload):
    if not year_month or not isinstance(payload, dict):
        return False
    if payload.get("stationResources"):
        return False
    resource_tons, station_tons = _retail_report_sales_tons(year_month)
    if not resource_tons and not station_tons:
        return False
    changed = False
    for item in payload.get("byResource") or payload.get("resources") or []:
        if not isinstance(item, dict):
            continue
        try:
            resource_key = int(item.get("resourceKey") or item.get("key") or 0)
        except (TypeError, ValueError):
            resource_key = 0
        if resource_key in resource_tons:
            tons = resource_tons[resource_key]
            item["mass"] = round(tons * 1000, 3)
            item["tons"] = tons
            changed = True
    for item in payload.get("stations") or []:
        if not isinstance(item, dict):
            continue
        try:
            shop_key = int(item.get("shopKey") or 0)
        except (TypeError, ValueError):
            shop_key = 0
        if shop_key in station_tons:
            tons = station_tons[shop_key]
            item["mass"] = round(tons * 1000, 3)
            item["tons"] = tons
            changed = True
    if station_tons:
        region_tons = {}
        for shop_key, tons in station_tons.items():
            region = (SHOP_KEY_MAP.get(shop_key) or {}).get("region") or "--"
            region_tons[region] = region_tons.get(region, 0.0) + float(tons or 0)
        regions = payload.get("regions")
        if isinstance(regions, dict):
            for region, item in regions.items():
                if isinstance(item, dict) and region in region_tons:
                    item["mass"] = round(region_tons[region] * 1000, 3)
                    item["tons"] = round(region_tons[region], 3)
                    changed = True
        elif isinstance(regions, list):
            for item in regions:
                if not isinstance(item, dict):
                    continue
                region = item.get("region") or item.get("name") or "--"
                if region in region_tons:
                    item["mass"] = round(region_tons[region] * 1000, 3)
                    item["tons"] = round(region_tons[region], 3)
                    changed = True
        region_items = regions.values() if isinstance(regions, dict) else regions if isinstance(regions, list) else []
        for item in region_items:
            if not isinstance(item, dict):
                continue
            for key in ("amount", "volume", "cash", "card", "bonus"):
                if key in item:
                    item[key] = round(_to_float(item.get(key)), 2)
            for key in ("mass", "tons"):
                if key in item:
                    item[key] = round(_to_float(item.get(key)), 3)
    if changed:
        summary = payload.setdefault("summary", {})
        total_tons = round(sum(resource_tons.values()) if resource_tons else sum(station_tons.values()), 3)
        if total_tons:
            summary["totalTons"] = total_tons
        payload["massSource"] = "retail_report_workbook"
    return changed


def add_resource_tons(resources):
    total = 0.0
    for item in resources:
        mass = retail_mass_from_item(item)
        if mass is not None:
            tons = round(mass / 1000, 3)
        else:
            density = RESOURCE_DENSITY.get(int(item.get("resourceKey") or item.get("key") or 0), 0)
            mass = float(item.get("volume") or 0) * density
            tons = round(mass / 1000, 3)
        item["mass"] = round(mass, 3)
        item["tons"] = tons
        total += tons
    return round(total, 3)
