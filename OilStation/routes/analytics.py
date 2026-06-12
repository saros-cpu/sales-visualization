import os
import json
from datetime import date, datetime, timedelta
from pathlib import Path

from flask import Blueprint, jsonify, request
from openpyxl import load_workbook
import pymysql

from business_time import (
    business_date_expr,
    business_period_condition_with_pruning,
    business_period_params,
    shift_business_date_expr,
    shift_join,
)
from cache_store import get_cached, set_cached
from constants import (
    CLOSED_STATIONS,
    DEPOT_COORDS,
    FORCE_NORMAL_STATIONS,
    SHOP_COORDS,
    SHOP_KEY_MAP,
)
from db import get_connection
from export_excel import (
    RESOURCE_DENSITY,
    SK4_RSK_MAP,
    SK6_RSK_MAP,
    SK28_RSK_MAP,
    SK29_RSK_MAP,
    STATION_RK_CORRECTIONS,
    _allow_fixed_density_fallback,
    _selling_mass_expr,
)

from endpoints.v1.retail.common.query import calc_yoy_mom, get_query_date_range, get_report_month_inventory, query_month_summary

analytics_api = Blueprint("analytics_api", __name__)
PROJECT_ROOT = Path(__file__).resolve().parents[2]
OILSTATION_ROOT = Path(__file__).resolve().parents[1]

RESOURCE_NAMES = {
    102: 'АИ 80',
    103: 'АИ 92',
    104: 'АИ 95',
    106: 'Д/Т',
    108: 'СПБТ',
    109: 'Д/Т',
    110: 'ДТ -35',
}


def cache_allowed():
    return request.args.get("refresh") not in {"1", "true", "yes"}


def _comparison_month(year: int, month: int, offset: int) -> str:
    value = year * 12 + (month - 1) + offset
    return f"{value // 12}-{value % 12 + 1:02d}"


def _snapshot_region_values(year_month: str) -> dict:
    payload = get_cached(f"regions:{year_month}") or get_cached(f"monthly:{year_month}") or {}
    return payload.get("regions") or {}


def _attach_region_snapshot_comparisons(year_month: str, payload: dict) -> bool:
    regions = payload.get("regions") or {}
    if not regions:
        return False
    try:
        year, month = [int(part) for part in year_month.split("-")]
    except (TypeError, ValueError):
        return False

    previous = _snapshot_region_values(_comparison_month(year, month, -1))
    last_year = _snapshot_region_values(_comparison_month(year, month, -12))
    changed = False
    for name, current in regions.items():
        for label, reference in (("mom", previous), ("yoy", last_year)):
            previous_item = reference.get(name) or {}
            previous_amount = float(previous_item.get("amount") or 0)
            previous_volume = float(previous_item.get("volume") or 0)
            current_amount = float(current.get("amount") or 0)
            current_volume = float(current.get("volume") or 0)
            sales = round((current_amount - previous_amount) / previous_amount * 100, 1) if previous_amount else 0
            volume = round((current_volume - previous_volume) / previous_volume * 100, 1) if previous_volume else 0
            if current.get(f"{label}Sales") != sales or current.get(f"{label}Volume") != volume:
                current[f"{label}Sales"] = sales
                current[f"{label}Volume"] = volume
                changed = True
    return changed


RESOURCE_COLORS = {
    102: '#00e5a0',
    103: '#00d4ff',
    104: '#a78bfa',
    106: '#ffa726',
    108: '#ff6b6b',
    109: '#ffb74d',
    110: '#ff8a65',
}
RESOURCE_ORDER = [102, 103, 104, 106, 110, 108]
RESOURCE_ORDER_INDEX = {rk: index for index, rk in enumerate(RESOURCE_ORDER)}
RESOURCE_ALIASES = {
    109: 106,
}
KNOWN_RESOURCE_KEYS = set(RESOURCE_NAMES)
STATION_MAP_DETAILS_TTL = 600
STATION_MAP_DETAILS_CACHE_KEY = "station_map_details:latest"


def _load_env_file(path):
    try:
        lines = path.read_text(encoding="utf-8-sig").splitlines()
    except OSError:
        return
    for line in lines:
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def _sales_db_config():
    _load_env_file(PROJECT_ROOT / "Sales" / ".env")
    _load_env_file(PROJECT_ROOT / "OilStation" / ".env")
    database = os.getenv("SALES_DB_NAME", os.getenv("DB_NAME", "sales_visualization"))
    if database == "ocw":
        database = "sales_visualization"
    return {
        "host": os.getenv("DB_HOST", "127.0.0.1"),
        "port": int(os.getenv("DB_PORT", "3306")),
        "user": os.getenv("DB_USER", "reader"),
        "password": os.getenv("DB_PASSWORD", ""),
        "database": database,
        "charset": os.getenv("DB_CHARSET", "utf8mb4"),
        "use_unicode": True,
        "autocommit": True,
    }


def _parse_day(value):
    if isinstance(value, date):
        return value
    return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()


def _load_cny_rates(start, end):
    rates = {}
    try:
        start_day = _parse_day(start)
        end_day = _parse_day(end)
    except (TypeError, ValueError):
        return rates
    if end_day <= start_day:
        return rates
    try:
        conn = pymysql.connect(**_sales_db_config())
    except Exception:
        return rates
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT rate_date, cny_kgs
            FROM exchange_rates
            WHERE rate_date >= %s AND rate_date < %s
              AND cny_kgs IS NOT NULL
            """,
            (start_day, end_day),
        )
        for rate_date, cny_kgs in cur.fetchall():
            rates[_parse_day(rate_date).isoformat()] = float(cny_kgs)
    finally:
        conn.close()
    return rates


def _load_density_history():
    densities = {}
    try:
        conn = pymysql.connect(**_sales_db_config())
    except Exception:
        return densities
    try:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT p.retail_resource_key, d.effective_date, d.density
            FROM product_density_history d
            JOIN products p ON p.id = d.product_id
            WHERE p.retail_resource_key IS NOT NULL
            ORDER BY p.retail_resource_key, d.effective_date
            """
        )
        for resource_key, effective_date, density in cur.fetchall():
            densities.setdefault(canonical_resource_key(resource_key), []).append(
                (_parse_day(effective_date), float(density))
            )
    except Exception:
        return {}
    finally:
        conn.close()
    return densities


def _density_for(density_history, resource_key, day):
    rk = canonical_resource_key(resource_key)
    selected = None
    for effective_date, density in density_history.get(rk, []):
        if effective_date <= day:
            selected = density
        else:
            break
    if selected is not None:
        return selected
    return RESOURCE_DENSITY.get(rk, 0)


def _monthly_cny_from_daily(daily_rows, start, end):
    rates = _load_cny_rates(start, end)
    total_cny = 0.0
    missing = []
    rate_values = []
    for item in daily_rows or []:
        amount = float(item.get("amount") or 0)
        if not amount:
            continue
        day = str(item.get("date") or "")[:10]
        rate = rates.get(day)
        if rate:
            total_cny += amount / rate
            rate_values.append(rate)
        else:
            missing.append(day)
    return {
        "totalSalesCny": round(total_cny, 2),
        "exchangeMissingDates": sorted(set(missing)),
        "exchangeComplete": not missing,
        "avgCnyKgs": round(sum(rate_values) / len(rate_values), 4) if rate_values else None,
    }


def _total_tons_from_daily_resource(daily_resource_rows):
    density_history = _load_density_history()
    total = 0.0
    missing = []
    for item in daily_resource_rows or []:
        mass = retail_mass_from_item(item)
        if mass is not None:
            total += mass / 1000
            continue
        volume = float(item.get("volume") or 0)
        if not volume:
            continue
        try:
            business_day = _parse_day(item.get("date"))
        except (TypeError, ValueError):
            continue
        density = _density_for(density_history, item.get("resourceKey") or item.get("key"), business_day)
        if not density:
            missing.append(f"{item.get('resourceKey')}:{business_day.isoformat()}")
            continue
        total += volume * density / 1000
    return {
        "totalTons": round(total, 3),
        "densityMissingKeys": sorted(set(missing)),
        "densityComplete": not missing,
    }


def _total_tons_from_resource_summary(resources, business_day=None):
    density_history = _load_density_history()
    if business_day is None:
        business_day = date.today()
    total = 0.0
    missing = []
    for item in resources or []:
        mass = retail_mass_from_item(item)
        if mass is not None:
            total += mass / 1000
            continue
        volume = float(item.get("volume") or 0)
        if not volume:
            continue
        density = _density_for(density_history, item.get("resourceKey") or item.get("key"), business_day)
        if not density:
            missing.append(str(item.get("resourceKey") or item.get("key") or "unknown"))
            continue
        total += volume * density / 1000
    return {
        "totalTons": round(total, 3),
        "densityMissingKeys": sorted(set(missing)),
        "densityComplete": not missing,
    }


def _attach_monthly_finance_metrics(payload, daily_resource_rows=None):
    summary = payload.get("summary") or {}
    daily = payload.get("daily") or []
    start = summary.get("queryStart") or payload.get("queryStart") or f"{payload.get('month')}-01"
    end = summary.get("queryEnd") or payload.get("queryEndExclusive")
    if not end and daily:
        last_day = max(str(item.get("date"))[:10] for item in daily)
        end = (datetime.strptime(last_day, "%Y-%m-%d").date() + timedelta(days=1)).isoformat()
    if start and end:
        summary.update(_monthly_cny_from_daily(daily, start, end))
    if daily_resource_rows:
        summary.update(_total_tons_from_daily_resource(daily_resource_rows))
    elif "totalTons" not in summary:
        try:
            business_day = _parse_day(start)
        except (TypeError, ValueError):
            business_day = None
        summary.update(_total_tons_from_resource_summary(payload.get("byResource") or [], business_day))
    payload["summary"] = summary
    return payload


def station_status_for(shop_key, amount):
    if shop_key in FORCE_NORMAL_STATIONS:
        return 'normal'
    if shop_key in CLOSED_STATIONS:
        return 'closed'
    return 'normal' if amount > 0 else 'closed'


def resource_key_expr(alias="s"):
    return report_resource_key_expr(alias)


def canonical_resource_key(resource_key):
    try:
        rk = int(resource_key or 0)
    except (TypeError, ValueError):
        return 0
    return RESOURCE_ALIASES.get(rk, rk)


def add_resource_bucket(target, resource_key, volume, amount, mass=None):
    rk = canonical_resource_key(resource_key)
    if rk < 100 or rk not in KNOWN_RESOURCE_KEYS:
        return
    bucket = target.setdefault(rk, {'volume': 0.0, 'amount': 0.0, 'mass': 0.0})
    bucket['volume'] += float(volume or 0)
    bucket['amount'] += float(amount or 0)
    if mass is not None:
        bucket['mass'] += float(mass or 0)


def report_volume_expr(alias="s"):
    return (
        "SUM(CASE "
        f"WHEN ({alias}.DerivedPaymentTypeKey IS NULL OR {alias}.DerivedPaymentTypeKey NOT IN (2, 3, 4)) "
        f" AND {alias}.BasePaymentTypeKey IN (1, 2) THEN {alias}.Volume ELSE 0 END) "
        f"+ SUM(CASE WHEN {alias}.DerivedPaymentTypeKey = 3 THEN {alias}.Volume ELSE 0 END)"
    )


def report_amount_expr(alias="s"):
    return (
        "SUM(CASE "
        f"WHEN ({alias}.DerivedPaymentTypeKey IS NULL OR {alias}.DerivedPaymentTypeKey NOT IN (2, 3, 4)) "
        f" AND {alias}.BasePaymentTypeKey IN (1, 2) THEN {alias}.ShopCost ELSE 0 END) "
        f"+ SUM(CASE WHEN {alias}.DerivedPaymentTypeKey = 3 THEN {alias}.ShopCost ELSE 0 END)"
    )


def report_mass_expr(alias="s", period_start=None):
    use_ai80_fallback = not _allow_fixed_density_fallback(period_start)
    selling_mass = _selling_mass_expr(alias, use_ai80_missing_density_fallback=use_ai80_fallback)
    return (
        "SUM(CASE "
        f"WHEN ({alias}.DerivedPaymentTypeKey IS NULL OR {alias}.DerivedPaymentTypeKey NOT IN (2, 3, 4)) "
        f" AND {alias}.BasePaymentTypeKey IN (1, 2) THEN {selling_mass} ELSE 0 END) "
        f"+ SUM(CASE WHEN {alias}.DerivedPaymentTypeKey = 3 THEN {selling_mass} ELSE 0 END)"
    )


def report_cash_amount_expr(alias="s"):
    return (
        "SUM(CASE "
        f"WHEN ({alias}.DerivedPaymentTypeKey IS NULL OR {alias}.DerivedPaymentTypeKey NOT IN (2, 3, 4)) "
        f" AND {alias}.BasePaymentTypeKey = 1 "
        f" AND COALESCE({alias}.DerivedPaymentTypeKey, 0) <> 5 THEN {alias}.ShopCost ELSE 0 END)"
    )


def report_card_amount_expr(alias="s"):
    return (
        "SUM(CASE "
        f"WHEN ({alias}.DerivedPaymentTypeKey IS NULL OR {alias}.DerivedPaymentTypeKey NOT IN (2, 3, 4)) "
        f" AND ({alias}.BasePaymentTypeKey = 2 OR {alias}.DerivedPaymentTypeKey = 5) "
        f"THEN {alias}.ShopCost ELSE 0 END)"
    )


def report_cash_volume_expr(alias="s"):
    return (
        "SUM(CASE "
        f"WHEN ({alias}.DerivedPaymentTypeKey IS NULL OR {alias}.DerivedPaymentTypeKey NOT IN (2, 3, 4)) "
        f" AND {alias}.BasePaymentTypeKey = 1 "
        f" AND COALESCE({alias}.DerivedPaymentTypeKey, 0) <> 5 THEN {alias}.Volume ELSE 0 END)"
    )


def report_card_volume_expr(alias="s"):
    return (
        "SUM(CASE "
        f"WHEN ({alias}.DerivedPaymentTypeKey IS NULL OR {alias}.DerivedPaymentTypeKey NOT IN (2, 3, 4)) "
        f" AND ({alias}.BasePaymentTypeKey = 2 OR {alias}.DerivedPaymentTypeKey = 5) "
        f"THEN {alias}.Volume ELSE 0 END)"
    )


def report_bonus_amount_expr(alias="s"):
    return (
        "SUM(CASE "
        f"WHEN ({alias}.DerivedPaymentTypeKey IS NULL OR {alias}.DerivedPaymentTypeKey NOT IN (2, 3, 4)) "
        f"THEN COALESCE({alias}.BonusOutCost, 0) ELSE 0 END)"
    )


def report_resource_key_expr(alias="s"):
    return (
        "CASE "
        f"WHEN {alias}.ShopKey = 28 AND {alias}.ResourceShopKey = 7 THEN 102 "
        f"WHEN {alias}.ShopKey = 28 AND {alias}.ResourceShopKey = 8 THEN 103 "
        f"WHEN {alias}.ShopKey = 28 AND {alias}.ResourceShopKey = 9 THEN 104 "
        f"WHEN {alias}.ShopKey = 28 AND {alias}.ResourceShopKey = 10 THEN 109 "
        f"WHEN {alias}.ShopKey = 28 AND {alias}.ResourceShopKey = 11 THEN 110 "
        f"WHEN {alias}.ShopKey = 28 AND {alias}.ResourceShopKey = 12 THEN 108 "
        f"WHEN {alias}.ShopKey = 6 AND {alias}.ResourceShopKey = 1 THEN 103 "
        f"WHEN {alias}.ShopKey = 6 AND {alias}.ResourceShopKey = 2 THEN 104 "
        f"WHEN {alias}.ShopKey = 6 AND {alias}.ResourceShopKey = 3 THEN 102 "
        f"WHEN {alias}.ShopKey = 6 AND {alias}.ResourceShopKey = 5 THEN 106 "
        f"WHEN {alias}.ShopKey = 4 AND {alias}.ResourceShopKey = 1 THEN 103 "
        f"WHEN {alias}.ShopKey = 4 AND {alias}.ResourceShopKey = 2 THEN 104 "
        f"WHEN {alias}.ShopKey = 4 AND {alias}.ResourceShopKey = 5 THEN 106 "
        f"WHEN {alias}.ShopKey = 4 AND {alias}.ResourceShopKey = 7 THEN 108 "
        f"WHEN {alias}.ShopKey = 29 AND {alias}.ResourceShopKey = 7 THEN 103 "
        f"WHEN {alias}.ShopKey = 29 AND {alias}.ResourceShopKey = 8 THEN 102 "
        f"WHEN {alias}.ShopKey = 29 AND {alias}.ResourceShopKey = 9 THEN 106 "
        f"WHEN {alias}.ShopKey = 29 AND {alias}.ResourceShopKey = 10 THEN 108 "
        f"WHEN {alias}.ShopKey = 30 AND {alias}.ResourceShopKey = 7 THEN 103 "
        f"WHEN {alias}.ShopKey = 30 AND {alias}.ResourceShopKey = 8 THEN 102 "
        f"WHEN {alias}.ShopKey = 30 AND {alias}.ResourceShopKey = 9 THEN 106 "
        f"WHEN {alias}.ShopKey IN (7, 15) AND {alias}.ResourceKey = 104 THEN 102 "
        f"WHEN {alias}.ResourceKey IN (102, 103, 104, 106, 108, 109, 110) THEN {alias}.ResourceKey "
        f"WHEN {alias}.BasePrice <= 0 THEN {alias}.ResourceKey "
        f"WHEN {alias}.BasePrice <= 45 THEN 108 "
        f"WHEN {alias}.BasePrice <= 65 THEN 102 "
        f"WHEN {alias}.BasePrice <= 78 THEN 103 "
        f"WHEN {alias}.BasePrice <= 83 THEN 104 "
        "ELSE 106 "
        "END"
    )


def calc_change(current, previous):
    if not previous or previous <= 0:
        return None
    return round((float(current or 0) - float(previous or 0)) / float(previous) * 100, 1)


def format_resource_rows(rows):
    merged = {}
    for row in rows:
        mass = row[3] if len(row) > 3 else None
        add_resource_bucket(merged, row[0], row[1], row[2], mass)
    result = []
    for rk, values in merged.items():
        volume = round(values['volume'], 2)
        amount = round(values['amount'], 2)
        mass = round(float(values.get('mass') or 0), 3)
        result.append({
            'resourceKey': rk,
            'key': rk,
            'name': RESOURCE_NAMES.get(rk, f'Other-{rk}'),
            'volume': volume,
            'mass': mass,
            'tons': round(mass / 1000, 3),
            'amount': amount,
            'color': RESOURCE_COLORS.get(rk, '#999999'),
        })
    result.sort(key=lambda item: RESOURCE_ORDER_INDEX.get(item['key'], 999))
    return result


def infer_balance_resource_key(shop_key, resource_key, price=None, resource_shop_key=None):
    try:
        sk = int(shop_key or 0)
        rk = int(resource_key or 0)
    except (TypeError, ValueError):
        return 0
    try:
        rsk = int(resource_shop_key or 0)
    except (TypeError, ValueError):
        rsk = 0

    if sk == 28 and rsk in SK28_RSK_MAP:
        return canonical_resource_key(SK28_RSK_MAP[rsk])
    if sk == 6 and rsk in SK6_RSK_MAP:
        return canonical_resource_key(SK6_RSK_MAP[rsk])
    if sk == 4 and rsk in SK4_RSK_MAP:
        return canonical_resource_key(SK4_RSK_MAP[rsk])
    if sk == 29 and rsk in SK29_RSK_MAP:
        return canonical_resource_key(SK29_RSK_MAP[rsk])
    corrected = STATION_RK_CORRECTIONS.get((sk, rk))
    if corrected:
        return canonical_resource_key(corrected)
    if rk in KNOWN_RESOURCE_KEYS:
        return canonical_resource_key(rk)

    try:
        p = float(price or 0)
    except (TypeError, ValueError):
        p = 0
    if p <= 0:
        return 0
    if p <= 45:
        return 108
    if p <= 65:
        return 102
    if p <= 78:
        return 103
    if p <= 83:
        return 104
    return 106


def _to_float(value):
    if value is None or value == "":
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _station_key_from_report_cell(value):
    try:
        station_key = int(float(value))
    except (TypeError, ValueError):
        return None
    return station_key if station_key in SHOP_KEY_MAP else None


def _resource_key_from_report_name(value):
    name = str(value or "").strip().lower()
    if not name:
        return 0
    if "-35" in name:
        return 110
    if "спбт" in name or "сжиж" in name:
        return 108
    if "95" in name:
        return 104
    if "92" in name:
        return 103
    if "80" in name:
        return 102
    if "д/т" in name or "диз" in name:
        return 106
    return 0


def _latest_report_workbook_for_map():
    index_path = OILSTATION_ROOT / "exports" / ".monthly_report_index.json"
    if not index_path.exists():
        return None, None
    try:
        report_index = json.loads(index_path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError):
        return None, None
    if not isinstance(report_index, dict):
        return None, None
    for ym in sorted(report_index.keys(), reverse=True):
        filename = (report_index.get(ym) or {}).get("file")
        if not filename:
            continue
        path = OILSTATION_ROOT / "exports" / filename
        if path.exists():
            return ym, path
    return None, None


def _report_workbook_snapshot_state(workbook_path):
    stat = workbook_path.stat()
    return {
        "source": "retail_report_snapshot",
        "sourceFile": workbook_path.name,
        "sourceMtime": stat.st_mtime,
        "snapshotGeneratedAt": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="minutes"),
    }


def _latest_completed_shift_marker_for_map():
    station_keys = sorted(SHOP_KEY_MAP.keys())
    if not station_keys:
        return None
    station_key_sql = ",".join(str(key) for key in station_keys)
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            f"""
            SELECT MAX(DatetimeShiftEnd)
            FROM shift
            WHERE ShopKey IN ({station_key_sql})
              AND DatetimeShiftBegin IS NOT NULL
              AND DatetimeShiftEnd IS NOT NULL
              AND DatetimeShiftEnd > DatetimeShiftBegin
              AND DatetimeShiftEnd >= DATE_SUB(NOW(), INTERVAL 180 DAY)
            """
        )
        row = cur.fetchone()
        return row[0] if row else None
    finally:
        conn.close()


def _station_map_details_cache_current(cached):
    if not isinstance(cached, dict):
        return True
    _, workbook_path = _latest_report_workbook_for_map()
    if not workbook_path:
        latest_shift_end = _latest_completed_shift_marker_for_map()
        if not latest_shift_end:
            return cached.get("source") == "live_database" and bool(cached.get("snapshotGeneratedAt"))
        try:
            cached_mtime = float(cached.get("sourceMtime"))
        except (TypeError, ValueError):
            cached_mtime = None
        return (
            cached.get("source") == "live_database"
            and cached_mtime is not None
            and abs(cached_mtime - latest_shift_end.timestamp()) < 0.001
        )
    try:
        snapshot_state = _report_workbook_snapshot_state(workbook_path)
    except OSError:
        return False
    try:
        cached_mtime = float(cached.get("sourceMtime"))
    except (TypeError, ValueError):
        cached_mtime = None
    return (
        cached.get("source") == snapshot_state["source"]
        and cached.get("sourceFile") == snapshot_state["sourceFile"]
        and cached_mtime is not None
        and abs(cached_mtime - snapshot_state["sourceMtime"]) < 0.001
    )


def _report_sheet_last_nonempty_day(workbook):
    last_day = 0
    for sheet_name in workbook.sheetnames:
        if not sheet_name.endswith("日") or not sheet_name[:2].isdigit():
            continue
        sheet = workbook[sheet_name]
        for row in sheet.iter_rows(min_row=5, values_only=True):
            if _to_float(row[18] if len(row) > 18 else None) or _to_float(row[20] if len(row) > 20 else None):
                last_day = max(last_day, int(sheet_name[:2]))
                break
    return last_day


def _station_map_details_from_report_snapshot():
    ym, workbook_path = _latest_report_workbook_for_map()
    if not workbook_path:
        return None
    workbook = None
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        last_day = _report_sheet_last_nonempty_day(workbook)
        sheet_name = f"{last_day:02d}日" if last_day else None
        if not sheet_name or sheet_name not in workbook.sheetnames:
            sheet_name = "总表" if "总表" in workbook.sheetnames else workbook.sheetnames[0]
        sheet = workbook[sheet_name]
        snapshot_date = f"{ym}-{last_day:02d}" if ym and last_day else ym
        snapshot_state = _report_workbook_snapshot_state(workbook_path)
        snapshot_generated_at = snapshot_state["snapshotGeneratedAt"]
        details = {
            shop_key: {
                "shopKey": shop_key,
                "inventory": 0.0,
                "inventoryByResource": [],
                "prices": [],
                "inventoryAsOf": snapshot_date,
                "priceAsOf": snapshot_date,
                "snapshotGeneratedAt": snapshot_generated_at,
            }
            for shop_key in sorted(SHOP_KEY_MAP.keys())
        }
        rows_by_station = {shop_key: {} for shop_key in details}
        current_shop_key = None
        for row in sheet.iter_rows(min_row=5, values_only=True):
            station_key = _station_key_from_report_cell(row[0] if len(row) > 0 else None)
            if station_key is not None:
                current_shop_key = station_key
            if current_shop_key is None or current_shop_key not in rows_by_station:
                continue
            product_name = str(row[3] if len(row) > 3 and row[3] is not None else "").strip()
            if not product_name or "合计" in product_name or "итого" in product_name.lower():
                continue
            rk = _resource_key_from_report_name(product_name)
            if rk not in KNOWN_RESOURCE_KEYS:
                continue
            if sheet_name == "总表":
                inventory = _to_float(row[23] if len(row) > 23 else None)
                price = _to_float(row[22] if len(row) > 22 else None)
            else:
                inventory = _to_float(row[21] if len(row) > 21 else None)
                sale_volume = _to_float(row[18] if len(row) > 18 else None)
                sale_amount = _to_float(row[20] if len(row) > 20 else None)
                price = round(sale_amount / sale_volume, 2) if sale_volume else 0.0
            if not inventory and not price:
                continue
            bucket = rows_by_station[current_shop_key].setdefault(rk, {
                "resourceKey": rk,
                "name": RESOURCE_NAMES.get(rk, f"Other-{rk}"),
                "inventory": 0.0,
                "price": None,
            })
            bucket["inventory"] += inventory
            if price:
                bucket["price"] = round(price, 2)

        for shop_key, by_resource in rows_by_station.items():
            rows = sorted(
                (
                    {**item, "inventory": round(float(item.get("inventory") or 0), 2)}
                    for item in by_resource.values()
                    if float(item.get("inventory") or 0) or item.get("price") is not None
                ),
                key=lambda item: RESOURCE_ORDER_INDEX.get(item["resourceKey"], 999),
            )
            detail = details[shop_key]
            detail["inventoryByResource"] = rows
            detail["inventory"] = round(sum(item["inventory"] for item in rows), 2)
            detail["prices"] = [
                {"resourceKey": item["resourceKey"], "name": item["name"], "price": item["price"]}
                for item in rows
                if item.get("price") is not None
            ]
        return {
            "items": list(details.values()),
            "updatedAt": datetime.now().isoformat(timespec="seconds"),
            "sourceSheet": sheet_name,
            **snapshot_state,
        }
    except Exception:
        return None
    finally:
        if workbook is not None:
            workbook.close()


def _station_map_details_payload():
    cached = get_cached(STATION_MAP_DETAILS_CACHE_KEY)
    if cached and cache_allowed() and _station_map_details_cache_current(cached):
        return cached

    snapshot_payload = _station_map_details_from_report_snapshot()
    if snapshot_payload:
        set_cached(STATION_MAP_DETAILS_CACHE_KEY, snapshot_payload, ttl=STATION_MAP_DETAILS_TTL)
        return snapshot_payload

    station_keys = sorted(SHOP_KEY_MAP.keys())
    station_key_sql = ",".join(str(key) for key in station_keys)
    details = {
        shop_key: {
            "shopKey": shop_key,
            "inventory": 0.0,
            "inventoryByResource": [],
            "prices": [],
            "inventoryAsOf": None,
            "priceAsOf": None,
        }
        for shop_key in station_keys
    }
    inventory = {}
    prices = {}
    price_dates = {}
    inventory_dates = {}
    latest_shifts = {}

    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute(
            f"""
            SELECT ShopKey, ShiftKey, DatetimeShiftEnd
            FROM shift
            WHERE ShopKey IN ({station_key_sql})
              AND DatetimeShiftBegin IS NOT NULL
              AND DatetimeShiftEnd IS NOT NULL
              AND DatetimeShiftEnd > DatetimeShiftBegin
              AND DatetimeShiftEnd >= DATE_SUB(NOW(), INTERVAL 180 DAY)
            ORDER BY ShopKey, DatetimeShiftEnd DESC
            """
        )
        for shop_key, shift_key, shift_end in cur.fetchall():
            latest_shifts.setdefault(int(shop_key), (int(shift_key), shift_end))

        if latest_shifts:
            conditions = " OR ".join(["(ShopKey = %s AND ShiftKey = %s)" for _ in latest_shifts])
            params = []
            for shop_key, (shift_key, _) in latest_shifts.items():
                params.extend([shop_key, shift_key])
            cur.execute(
                f"""
                SELECT ShopKey, ResourceKey, EndBalance, EndPrice, ResourceShopKey, ShiftKey
                FROM balance
                WHERE BeginMass IS NOT NULL
                  AND ({conditions})
                ORDER BY ShopKey, ResourceKey, ResourceShopKey
                """,
                params,
            )
            for shop_key, resource_key, end_balance, end_price, resource_shop_key, shift_key in cur.fetchall():
                sk = int(shop_key)
                rk = infer_balance_resource_key(sk, resource_key, end_price, resource_shop_key)
                if rk < 100 or rk not in KNOWN_RESOURCE_KEYS:
                    continue
                key = (sk, rk)
                inventory[key] = round(inventory.get(key, 0.0) + float(end_balance or 0), 2)
                as_of = latest_shifts.get(sk, (None, None))[1]
                if end_price:
                    prices[key] = round(float(end_price), 2)
                    price_dates[key] = as_of
                inventory_dates[sk] = as_of
    finally:
        conn.close()

    for shop_key in station_keys:
        resource_keys = sorted(
            {rk for sk, rk in inventory if sk == shop_key} | {rk for sk, rk in prices if sk == shop_key},
            key=lambda rk: RESOURCE_ORDER_INDEX.get(rk, 999),
        )
        rows = []
        for rk in resource_keys:
            volume = round(float(inventory.get((shop_key, rk), 0) or 0), 2)
            price = prices.get((shop_key, rk))
            if not volume and price is None:
                continue
            rows.append({
                "resourceKey": rk,
                "name": RESOURCE_NAMES.get(rk, f"Other-{rk}"),
                "inventory": volume,
                "price": price,
            })
        total_inventory = round(sum(item["inventory"] for item in rows), 2)
        detail = details[shop_key]
        detail["inventory"] = total_inventory
        detail["inventoryByResource"] = rows
        detail["prices"] = [
            {"resourceKey": item["resourceKey"], "name": item["name"], "price": item["price"]}
            for item in rows
            if item.get("price") is not None
        ]
        inv_dt = inventory_dates.get(shop_key)
        price_dt = max((price_dates.get((shop_key, rk)) for rk in resource_keys if price_dates.get((shop_key, rk))), default=None)
        detail["inventoryAsOf"] = inv_dt.isoformat(sep=" ", timespec="minutes") if hasattr(inv_dt, "isoformat") else None
        detail["priceAsOf"] = price_dt.isoformat(sep=" ", timespec="minutes") if hasattr(price_dt, "isoformat") else None
        detail["snapshotGeneratedAt"] = detail["inventoryAsOf"] or detail["priceAsOf"]

    latest_shift_end = max((shift_end for _, shift_end in latest_shifts.values() if shift_end), default=None)
    generated_at = (
        latest_shift_end.isoformat(sep=" ", timespec="minutes")
        if hasattr(latest_shift_end, "isoformat")
        else datetime.now().isoformat(timespec="minutes")
    )
    source_mtime = latest_shift_end.timestamp() if hasattr(latest_shift_end, "timestamp") else None

    payload = {
        "items": list(details.values()),
        "updatedAt": datetime.now().isoformat(timespec="seconds"),
        "source": "live_database",
        "snapshotGeneratedAt": generated_at,
        "sourceMtime": source_mtime,
    }
    set_cached(STATION_MAP_DETAILS_CACHE_KEY, payload, ttl=STATION_MAP_DETAILS_TTL)
    return payload


RETAIL_DENSITY_MIN = 0.4
RETAIL_DENSITY_MAX = 1.1


def retail_mass_is_plausible(volume, mass):
    try:
        volume_value = float(volume or 0)
        mass_value = float(mass or 0)
    except (TypeError, ValueError):
        return False
    if volume_value <= 0:
        return mass_value >= 0
    if mass_value <= 0:
        return False
    density = mass_value / volume_value
    return RETAIL_DENSITY_MIN <= density <= RETAIL_DENSITY_MAX


def retail_mass_from_item(item):
    volume = item.get("volume") if isinstance(item, dict) else 0
    if isinstance(item, dict) and item.get("tons") is not None:
        try:
            mass = float(item.get("tons") or 0) * 1000
        except (TypeError, ValueError):
            mass = 0
        if retail_mass_is_plausible(volume, mass):
            return mass
    if isinstance(item, dict) and item.get("mass") is not None:
        try:
            mass = float(item.get("mass") or 0)
        except (TypeError, ValueError):
            mass = 0
        if retail_mass_is_plausible(volume, mass):
            return mass
    return None


def retail_payload_mass_needs_refresh(payload):
    if not isinstance(payload, dict):
        return False
    items = []
    for key in ("byResource", "resources", "stations"):
        values = payload.get(key)
        if isinstance(values, list):
            items.extend(values)
    if not items:
        return False
    resource_tons = 0.0
    has_volume_without_mass = False
    for item in items:
        if not isinstance(item, dict):
            continue
        try:
            volume = float(item.get("volume") or item.get("latestMonthVolume") or 0)
        except (TypeError, ValueError):
            volume = 0
        if volume <= 0:
            continue
        if "mass" not in item:
            has_volume_without_mass = True
            continue
        mass = retail_mass_from_item(item)
        if mass is None:
            has_volume_without_mass = True
            continue
        resource_tons += mass / 1000
    if has_volume_without_mass:
        return True
    summary = payload.get("summary") or {}
    try:
        summary_tons = float(summary.get("totalTons") or 0)
    except (TypeError, ValueError):
        summary_tons = 0
    if summary_tons > 0 and resource_tons > 0:
        return abs(summary_tons - resource_tons) > max(1.0, summary_tons * 0.05)
    return False


def add_resource_tons(resources):
    total = 0.0
    for item in resources:
        mass = retail_mass_from_item(item)
        if mass is not None:
            tons = round(mass / 1000, 3)
        else:
            density = RESOURCE_DENSITY.get(int(item.get("resourceKey") or item.get("key") or 0), 0)
            mass = float(item.get("volume") or 0) * density
            tons = round(mass / 1000, 3)
        item["mass"] = round(mass, 3)
        item["tons"] = tons
        total += tons
    return round(total, 3)


def _snapshot_month_labels(year):
    now = datetime.now()
    max_month = 12
    if year == now.year:
        max_month = now.month
    elif year > now.year:
        max_month = 0
    return [f"{year}-{month:02d}" for month in range(1, max_month + 1)]


def _monthly_snapshots_for_year(year):
    snapshots = []
    for label in _snapshot_month_labels(year):
        payload = get_cached(f"monthly:{label}")
        if payload:
            snapshots.append(payload)
    return snapshots


def _resource_map_from_rows(rows):
    bucket = {}
    for item in rows or []:
        add_resource_bucket(
            bucket,
            item.get("resourceKey") or item.get("key"),
            item.get("volume"),
            item.get("amount"),
            retail_mass_from_item(item),
        )
    return bucket


def _sum_resource_maps(*maps):
    merged = {}
    for resource_map in maps:
        for rk, values in (resource_map or {}).items():
            add_resource_bucket(merged, rk, values.get("volume"), values.get("amount"), values.get("mass"))
    return merged


def _resource_items_from_map(resource_map, previous=None, month_data=None, prev_month_data=None):
    previous = previous or {}
    month_data = month_data or {}
    prev_month_data = prev_month_data or {}
    resource_keys = sorted(set(resource_map) | set(previous) | set(month_data) | set(prev_month_data))
    resources = []
    for rk in resource_keys:
        cur_d = resource_map.get(rk, {"volume": 0, "amount": 0, "mass": 0})
        prev_d = previous.get(rk, {"volume": 0, "amount": 0, "mass": 0})
        mon_d = month_data.get(rk, {"volume": 0, "amount": 0, "mass": 0})
        prev_mon_d = prev_month_data.get(rk, {"volume": 0, "amount": 0, "mass": 0})
        mass = round(float(cur_d.get("mass") or 0), 3)
        latest_month_mass = round(float(mon_d.get("mass") or 0), 3)
        resources.append({
            "resourceKey": rk,
            "key": rk,
            "name": RESOURCE_NAMES.get(rk, f"Other-{rk}"),
            "volume": round(cur_d["volume"], 2),
            "mass": mass,
            "tons": round(mass / 1000, 3),
            "amount": round(cur_d["amount"], 2),
            "color": RESOURCE_COLORS.get(rk, "#999999"),
            "yoyVolume": calc_change(cur_d["volume"], prev_d["volume"]),
            "yoySales": calc_change(cur_d["amount"], prev_d["amount"]),
            "momVolume": calc_change(mon_d["volume"], prev_mon_d["volume"]),
            "momSales": calc_change(mon_d["amount"], prev_mon_d["amount"]),
            "latestMonthVolume": round(mon_d["volume"], 2),
            "latestMonthMass": latest_month_mass,
            "latestMonthTons": round(latest_month_mass / 1000, 3),
            "latestMonthAmount": round(mon_d["amount"], 2),
        })
    resources.sort(key=lambda item: RESOURCE_ORDER_INDEX.get(item["key"], 999))
    return resources


def _dashboard_yearly_from_monthly_snapshots(year):
    months = _monthly_snapshots_for_year(year)
    if not months:
        return None

    stations = {}
    resource_totals = {}
    query_start = f"{year}-01-01"
    query_end = query_start

    for payload in months:
        summary = payload.get("summary") or {}
        query_start = min(query_start, summary.get("queryStart") or query_start)
        query_end = max(query_end, summary.get("queryEnd") or query_end)
        for station in payload.get("stations") or []:
            shop_key = station.get("shopKey")
            if shop_key is None:
                continue
            bucket = stations.setdefault(int(shop_key), {
                "txnCount": 0,
                "volume": 0.0,
                "mass": 0.0,
                "amount": 0.0,
                "cash": 0.0,
                "card": 0.0,
                "cashLiters": 0.0,
                "cardLiters": 0.0,
                "bonus": 0.0,
            })
            bucket["txnCount"] += int(station.get("txnCount") or 0)
            bucket["volume"] += float(station.get("volume") or 0)
            bucket["mass"] += float(station.get("mass") or 0) or float(station.get("tons") or 0) * 1000
            bucket["amount"] += float(station.get("amount") or station.get("totalSales") or 0)
            bucket["cash"] += float(station.get("cash") or station.get("cashSales") or 0)
            bucket["card"] += float(station.get("card") or station.get("cardSales") or 0)
            bucket["cashLiters"] += float(station.get("cashLiters") or 0)
            bucket["cardLiters"] += float(station.get("cardLiters") or 0)
            bucket["bonus"] += float(station.get("bonus") or station.get("bonusCost") or 0)
        resource_totals = _sum_resource_maps(resource_totals, _resource_map_from_rows(payload.get("byResource")))

    stations_list = []
    for shop_key, info in SHOP_KEY_MAP.items():
        coords = SHOP_COORDS.get(shop_key, {"lng": 0, "lat": 0})
        data = stations.get(shop_key, {})
        amount = round(float(data.get("amount") or 0), 2)
        stations_list.append({
            "shopKey": shop_key,
            "name": info["name"],
            "short": info["short"],
            "region": info["region"],
            "lng": coords.get("lng", 0),
            "lat": coords.get("lat", 0),
            "txnCount": int(data.get("txnCount") or 0),
            "volume": round(float(data.get("volume") or 0), 2),
            "mass": round(float(data.get("mass") or 0), 3),
            "tons": round(float(data.get("mass") or 0) / 1000, 3),
            "amount": amount,
            "cash": round(float(data.get("cash") or 0), 2),
            "card": round(float(data.get("card") or 0), 2),
            "cashLiters": round(float(data.get("cashLiters") or 0), 2),
            "cardLiters": round(float(data.get("cardLiters") or 0), 2),
            "bonus": round(float(data.get("bonus") or 0), 2),
            "status": station_status_for(shop_key, amount),
            "closedReason": CLOSED_STATIONS.get(shop_key, ""),
        })

    regions = {}
    for station in stations_list:
        region = station["region"]
        bucket = regions.setdefault(region, {
            "amount": 0.0,
            "volume": 0.0,
            "tons": 0.0,
            "cash": 0.0,
            "card": 0.0,
            "bonus": 0.0,
            "stationCount": 0,
            "activeCount": 0,
        })
        bucket["stationCount"] += 1
        bucket["amount"] += station["amount"]
        bucket["volume"] += station["volume"]
        bucket["tons"] += station.get("tons") or 0
        bucket["cash"] += station["cash"]
        bucket["card"] += station["card"]
        bucket["bonus"] += station["bonus"]
        if station["amount"] > 0:
            bucket["activeCount"] += 1
    for data in regions.values():
        for key in ("amount", "volume", "cash", "card", "bonus"):
            data[key] = round(data[key], 2)
        data["tons"] = round(data.get("tons") or 0, 3)

    return {
        "year": year,
        "queryStart": query_start,
        "queryEndExclusive": query_end,
        "stations": stations_list,
        "depots": [
            {
                "name": name,
                "short": data["short"],
                "region": data["region"],
                "lng": data["lng"],
                "lat": data["lat"],
                "type": "depot",
                "status": data.get("status", "normal"),
                "closedReason": data.get("closedReason", ""),
            }
            for name, data in DEPOT_COORDS.items()
        ],
        "regions": regions,
        "byResource": format_resource_rows([
            (rk, values["volume"], values["amount"], values.get("mass", 0))
            for rk, values in resource_totals.items()
        ]),
        "snapshot": True,
    }


def _resources_yearly_from_monthly_snapshots(year, month=None):
    if month:
        current_label = f"{year}-{month:02d}"
        current_payload = get_cached(f"monthly:{current_label}")
        if not current_payload:
            return None
        prev_label = f"{year - 1}-12" if month == 1 else f"{year}-{month - 1:02d}"
        yoy_label = f"{year - 1}-{month:02d}"
        current = _resource_map_from_rows(current_payload.get("byResource"))
        previous = _resource_map_from_rows((get_cached(f"monthly:{yoy_label}") or {}).get("byResource"))
        prev_month_data = _resource_map_from_rows((get_cached(f"monthly:{prev_label}") or {}).get("byResource"))
        latest_month = current_label
        prev_month = prev_label
        period_start = current_payload.get("summary", {}).get("queryStart", f"{year}-{month:02d}-01")
        period_end = current_payload.get("summary", {}).get("queryEnd", period_start)
        resources = _resource_items_from_map(current, previous, current, prev_month_data)
    else:
        months = _monthly_snapshots_for_year(year)
        if not months:
            return None
        current = {}
        for payload in months:
            current = _sum_resource_maps(current, _resource_map_from_rows(payload.get("byResource")))
        previous = {}
        for payload in _monthly_snapshots_for_year(year - 1):
            previous = _sum_resource_maps(previous, _resource_map_from_rows(payload.get("byResource")))
        latest_payload = months[-1]
        latest_month = latest_payload.get("month")
        prev_month = None
        if latest_month:
            y, m = int(latest_month[:4]), int(latest_month[5:7])
            prev_month = f"{y - 1}-12" if m == 1 else f"{y}-{m - 1:02d}"
        month_data = _resource_map_from_rows(latest_payload.get("byResource"))
        prev_month_data = _resource_map_from_rows((get_cached(f"monthly:{prev_month}") or {}).get("byResource")) if prev_month else {}
        period_start = f"{year}-01-01"
        period_end = max((payload.get("summary", {}).get("queryEnd") or period_start for payload in months), default=period_start)
        resources = _resource_items_from_map(current, previous, month_data, prev_month_data)

    return {
        "year": year,
        "month": month,
        "periodLabel": f"{year}-{month:02d}" if month else f"{year} 全年",
        "queryStart": period_start,
        "queryEndExclusive": period_end,
        "latestMonth": latest_month,
        "prevMonth": prev_month,
        "resources": resources,
        "summary": {
            "totalVolume": round(sum(item["volume"] for item in resources), 2),
            "totalTons": round(sum(item["tons"] for item in resources), 3),
            "totalAmount": round(sum(item["amount"] for item in resources), 2),
            "resourceCount": len([item for item in resources if item["amount"] > 0 or item["volume"] > 0]),
        },
        "snapshot": True,
    }


def register_analytics_routes(app):
    app.register_blueprint(analytics_api)


@analytics_api.route('/api/dashboard/yearly')
def api_dashboard_yearly():
    """Dashboard yearly cumulative data for map, rankings, regions, oil mix, and payment charts."""
    year = request.args.get("year", type=int) or datetime.now().year
    cache_key = f'dashboard_yearly:{year}'
    cached = get_cached(cache_key)
    if cached and cache_allowed():
        return jsonify(cached)
    if cache_allowed():
        snapshot_result = _dashboard_yearly_from_monthly_snapshots(year)
        if snapshot_result:
            set_cached(cache_key, snapshot_result)
            return jsonify(snapshot_result)

    year_start = f'{year}-01-01'
    current_year = datetime.now().year
    if year == current_year:
        year_end = datetime.now().strftime('%Y-%m-%d')
    elif year < current_year:
        year_end = f'{year + 1}-01-01'
    else:
        year_end = year_start

    conn = get_connection()
    try:
        cur = conn.cursor()

        year_condition = business_period_condition_with_pruning("s", "sh")
        total_volume = report_volume_expr("s")
        total_amount = report_amount_expr("s")
        resource_totals = {}
        resource_key = resource_key_expr("s")
        cur.execute(f"""
            SELECT
                s.ShopKey,
                {resource_key} as resource_key,
                COUNT(*) as txn_count,
                {total_volume} as total_liters,
                {total_amount} as total_sales,
                {report_cash_amount_expr("s")} as cash_sales,
                {report_card_amount_expr("s")} as card_sales,
                {report_cash_volume_expr("s")} as cash_liters,
                {report_card_volume_expr("s")} as card_liters,
                SUM(s.BonusOutCost) as bonus_cost
            FROM selling s
            {shift_join("s", "sh")}
            WHERE {year_condition}
              AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
            GROUP BY s.ShopKey, resource_key
        """, business_period_params(year_start, year_end))

        sales_data = {}
        for row in cur.fetchall():
            shop_key = row[0]
            resource_key_value = row[1]
            shop_bucket = sales_data.setdefault(shop_key, {
                'txnCount': 0,
                'volume': 0.0,
                'amount': 0.0,
                'cash': 0.0,
                'card': 0.0,
                'cashLiters': 0.0,
                'cardLiters': 0.0,
                'bonus': 0.0,
            })
            shop_bucket['txnCount'] += int(row[2] or 0)
            shop_bucket['volume'] += float(row[3] or 0)
            shop_bucket['amount'] += float(row[4] or 0)
            shop_bucket['cash'] += float(row[5] or 0)
            shop_bucket['card'] += float(row[6] or 0)
            shop_bucket['cashLiters'] += float(row[7] or 0)
            shop_bucket['cardLiters'] += float(row[8] or 0)
            shop_bucket['bonus'] += float(row[9] or 0)
            add_resource_bucket(resource_totals, resource_key_value, row[3], row[4])

        for data in sales_data.values():
            for key in ('volume', 'amount', 'cash', 'card', 'cashLiters', 'cardLiters', 'bonus'):
                data[key] = round(data[key], 2)

        by_resource = []
        for rk, values in resource_totals.items():
            by_resource.append((
                rk,
                round(values['volume'], 2),
                round(values['amount'], 2),
            ))
        by_resource = format_resource_rows(by_resource)

        stations_list = []
        for shop_key, info in SHOP_KEY_MAP.items():
            coords = SHOP_COORDS.get(shop_key, {'lng': 0, 'lat': 0})
            data = sales_data.get(shop_key, {})
            amount = data.get('amount', 0)
            stations_list.append({
                'shopKey': shop_key,
                'name': info['name'],
                'short': info['short'],
                'region': info['region'],
                'lng': coords.get('lng', 0),
                'lat': coords.get('lat', 0),
                'txnCount': data.get('txnCount', 0),
                'volume': data.get('volume', 0),
                'amount': amount,
                'cash': data.get('cash', 0),
                'card': data.get('card', 0),
                'cashLiters': data.get('cashLiters', 0),
                'cardLiters': data.get('cardLiters', 0),
                'bonus': data.get('bonus', 0),
                'status': station_status_for(shop_key, amount),
                'closedReason': CLOSED_STATIONS.get(shop_key, ''),
            })

        regions = {}
        for shop_key, info in SHOP_KEY_MAP.items():
            region_name = info['region']
            regions.setdefault(region_name, {
                'amount': 0,
                'volume': 0,
                'cash': 0,
                'card': 0,
                'bonus': 0,
                'stationCount': 0,
                'activeCount': 0,
            })
            regions[region_name]['stationCount'] += 1

            data = sales_data.get(shop_key, {})
            amount = float(data.get('amount', 0) or 0)
            regions[region_name]['amount'] += amount
            regions[region_name]['volume'] += float(data.get('volume', 0) or 0)
            regions[region_name]['cash'] += float(data.get('cash', 0) or 0)
            regions[region_name]['card'] += float(data.get('card', 0) or 0)
            regions[region_name]['bonus'] += float(data.get('bonus', 0) or 0)
            if amount > 0:
                regions[region_name]['activeCount'] += 1

        for data in regions.values():
            for key in ('amount', 'volume', 'cash', 'card', 'bonus'):
                data[key] = round(data[key], 2)

        result = {
            'year': year,
            'queryStart': year_start,
            'queryEndExclusive': year_end,
            'stations': stations_list,
            'depots': [
                {
                    'name': name,
                    'short': data['short'],
                    'region': data['region'],
                    'lng': data['lng'],
                    'lat': data['lat'],
                    'type': 'depot',
                    'status': data.get('status', 'normal'),
                    'closedReason': data.get('closedReason', ''),
                }
                for name, data in DEPOT_COORDS.items()
            ],
            'regions': regions,
            'byResource': by_resource,
        }
        set_cached(cache_key, result)
        return jsonify(result)
    finally:
        conn.close()


@analytics_api.route('/api/resources/yearly')
def api_resources_yearly():
    year = request.args.get("year", type=int) or datetime.now().year
    month = request.args.get("month", type=int)
    if month is not None and (month < 1 or month > 12):
        return jsonify({"error": "month must be between 1 and 12"}), 400
    cache_key = f'resources_yearly:{year}:{month or "year"}'
    cached = get_cached(cache_key)
    if cached and cache_allowed() and not retail_payload_mass_needs_refresh(cached):
        return jsonify(cached)
    if cache_allowed():
        snapshot_result = _resources_yearly_from_monthly_snapshots(year, month)
        if snapshot_result and not retail_payload_mass_needs_refresh(snapshot_result):
            set_cached(cache_key, snapshot_result)
            return jsonify(snapshot_result)

    now = datetime.now()
    current_year = now.year
    if month:
        period_label = f"{year}-{month:02d}"
        period_start, period_end = get_query_date_range(period_label)
        prev_label = f"{year - 1}-12" if month == 1 else f"{year}-{month - 1:02d}"
        prev_start, prev_end = get_query_date_range(prev_label, exclude_today=False)
        yoy_label = f"{year - 1}-{month:02d}"
        yoy_start, yoy_end = get_query_date_range(yoy_label, exclude_today=False)
        latest_month = period_label
        prev_month = prev_label
    else:
        period_start = f'{year}-01-01'
        if year == current_year:
            period_end = now.strftime('%Y-%m-%d')
        elif year < current_year:
            period_end = f'{year + 1}-01-01'
        else:
            period_end = period_start

        yoy_start = f'{year - 1}-01-01'
        if year == current_year:
            elapsed_days = (now.date() - datetime(year, 1, 1).date()).days
            yoy_end = (datetime(year - 1, 1, 1).date() + timedelta(days=elapsed_days)).strftime('%Y-%m-%d')
        else:
            yoy_end = f'{year}-01-01'
        latest_month = None
        prev_month = None
        prev_start = prev_end = None

    conn = get_connection()
    try:
        cur = conn.cursor()
        resource_key = resource_key_expr("s")
        business_day = shift_business_date_expr("sh")
        volume_expr = report_volume_expr("s")
        amount_expr = report_amount_expr("s")

        def query_period(start, end):
            cur.execute(f"""
                SELECT {resource_key} as resource_key,
                       {volume_expr} as vol,
                       {amount_expr} as amt
                FROM shift sh
                JOIN selling s FORCE INDEX (Index_SS)
                  ON s.ShopKey = sh.ShopKey
                 AND s.ShiftKey = sh.ShiftKey
                WHERE sh.ShopKey IN ({",".join(str(k) for k in sorted(SHOP_KEY_MAP.keys()))})
                  AND sh.DatetimeShiftBegin IS NOT NULL
                  AND sh.DatetimeShiftEnd IS NOT NULL
                  AND sh.DatetimeShiftEnd > sh.DatetimeShiftBegin
                  AND sh.DatetimeShiftEnd < %s
                  AND {business_day} >= %s
                  AND {business_day} < %s
                  AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
                GROUP BY resource_key
            """, (end, start, end))
            data = {}
            for r in cur.fetchall():
                add_resource_bucket(data, r[0], r[1], r[2])
            return data

        current = query_period(period_start, period_end)
        previous = query_period(yoy_start, yoy_end)

        month_data = current if month else {}
        prev_month_data = query_period(prev_start, prev_end) if month and prev_start and prev_end else {}
        if not month:
            cur.execute(f"""
                SELECT DATE_FORMAT(MAX({business_day}), '%%Y-%%m')
                FROM shift sh
                JOIN selling s FORCE INDEX (Index_SS)
                  ON s.ShopKey = sh.ShopKey
                 AND s.ShiftKey = sh.ShiftKey
                WHERE sh.ShopKey IN ({",".join(str(k) for k in sorted(SHOP_KEY_MAP.keys()))})
                  AND sh.DatetimeShiftBegin IS NOT NULL
                  AND sh.DatetimeShiftEnd IS NOT NULL
                  AND sh.DatetimeShiftEnd > sh.DatetimeShiftBegin
                  AND sh.DatetimeShiftEnd < %s
                  AND {business_day} >= %s
                  AND {business_day} < %s
                  AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
            """, (period_end, period_start, period_end))
            latest_month_row = cur.fetchone()
            latest_month = latest_month_row[0] if latest_month_row and latest_month_row[0] else None
            if latest_month:
                y, m = int(latest_month[:4]), int(latest_month[5:7])
                prev_month = f"{y - 1}-12" if m == 1 else f"{y}-{m - 1:02d}"
                latest_start, latest_end = get_query_date_range(latest_month, exclude_today=False)
                prev_latest_start, prev_latest_end = get_query_date_range(prev_month, exclude_today=False)
                month_data = query_period(latest_start, latest_end)
                prev_month_data = query_period(prev_latest_start, prev_latest_end)

        resource_keys = sorted(set(current) | set(previous) | set(month_data) | set(prev_month_data))
        resources = []
        for rk in resource_keys:
            cur_d = current.get(rk, {'volume': 0, 'amount': 0})
            prev_d = previous.get(rk, {'volume': 0, 'amount': 0})
            mon_d = month_data.get(rk, {'volume': 0, 'amount': 0})
            prev_mon_d = prev_month_data.get(rk, {'volume': 0, 'amount': 0})
            resources.append({
                'resourceKey': rk,
                'key': rk,
                'name': RESOURCE_NAMES.get(rk, f'Other-{rk}'),
                'volume': round(cur_d['volume'], 2),
                'amount': round(cur_d['amount'], 2),
                'color': RESOURCE_COLORS.get(rk, '#999999'),
                'yoyVolume': calc_change(cur_d['volume'], prev_d['volume']),
                'yoySales': calc_change(cur_d['amount'], prev_d['amount']),
                'momVolume': calc_change(mon_d['volume'], prev_mon_d['volume']),
                'momSales': calc_change(mon_d['amount'], prev_mon_d['amount']),
                'latestMonthVolume': round(mon_d['volume'], 2),
                'latestMonthAmount': round(mon_d['amount'], 2),
            })
        resources.sort(key=lambda item: RESOURCE_ORDER_INDEX.get(item['key'], 999))

        result = {
            'year': year,
            'month': month,
            'periodLabel': f"{year}-{month:02d}" if month else f"{year} 全年",
            'queryStart': period_start,
            'queryEndExclusive': period_end,
            'latestMonth': latest_month,
            'prevMonth': prev_month,
            'resources': resources,
            'summary': {
                'totalVolume': round(sum(item['volume'] for item in resources), 2),
                'totalAmount': round(sum(item['amount'] for item in resources), 2),
                'resourceCount': len([item for item in resources if item['amount'] > 0 or item['volume'] > 0]),
            },
        }
        set_cached(cache_key, result)
        return jsonify(result)
    finally:
        conn.close()


@analytics_api.route('/api/station-map-details')
def api_station_map_details():
    """Latest per-station inventory and retail prices for map hover cards."""
    return jsonify(_station_map_details_payload())


@analytics_api.route('/api/integration/retail-summary')
def api_integration_retail_summary():
    year = request.args.get("year", type=int) or datetime.now().year
    month = request.args.get("month", type=int)
    if month is not None and (month < 1 or month > 12):
        return jsonify({"error": "month must be between 1 and 12"}), 400

    cache_key = f"integration:retail-summary:{year}:{month or 'year'}"
    cached = get_cached(cache_key)
    if cached and cache_allowed():
        return jsonify(cached)

    if month:
        period_label = f"{year}-{month:02d}"
        start, end = get_query_date_range(period_label)
    else:
        start = f"{year}-01-01"
        current_year = datetime.now().year
        if year == current_year:
            end = datetime.now().strftime("%Y-%m-%d")
        elif year < current_year:
            end = f"{year + 1}-01-01"
        else:
            end = start

    conn = get_connection()
    try:
        cur = conn.cursor()
        business_day = shift_business_date_expr("sh")
        resource_key = report_resource_key_expr("s")
        daily_rows = {}
        by_resource_rows = {}
        daily_resource = []
        active_shops = set()
        total_txn = 0

        for shop_key in sorted(SHOP_KEY_MAP.keys()):
            cur.execute(f"""
                SELECT {business_day} as business_day,
                       {resource_key} as resource_key,
                       COUNT(*) as txn_count,
                       SUM(CASE
                           WHEN (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 3, 4))
                            AND s.BasePaymentTypeKey IN (1, 2)
                           THEN s.Volume ELSE 0 END)
                       + SUM(CASE WHEN s.DerivedPaymentTypeKey = 3 THEN s.Volume ELSE 0 END) as total_liters,
                       SUM(CASE
                           WHEN (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 3, 4))
                            AND s.BasePaymentTypeKey IN (1, 2)
                           THEN s.ShopCost ELSE 0 END)
                       + SUM(CASE WHEN s.DerivedPaymentTypeKey = 3 THEN s.ShopCost ELSE 0 END) as total_sales
                FROM shift sh
                JOIN selling s FORCE INDEX (Index_SS)
                  ON s.ShopKey = sh.ShopKey
                 AND s.ShiftKey = sh.ShiftKey
                WHERE sh.ShopKey = %s
                  AND sh.DatetimeShiftBegin IS NOT NULL
                  AND sh.DatetimeShiftEnd IS NOT NULL
                  AND sh.DatetimeShiftEnd > sh.DatetimeShiftBegin
                  AND sh.DatetimeShiftEnd < %s
                  AND {business_day} >= %s
                  AND {business_day} < %s
                  AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
                GROUP BY business_day, resource_key
                ORDER BY business_day, resource_key
            """, (shop_key, end, start, end))

            shop_has_sales = False
            for row in cur.fetchall():
                business_day_value = str(row[0])
                rk = int(row[1] or 0)
                if rk < 100:
                    continue
                txn_count = int(row[2] or 0)
                volume = float(row[3] or 0)
                amount = float(row[4] or 0)
                if not volume and not amount:
                    continue
                total_txn += txn_count
                if amount > 0 or volume > 0:
                    shop_has_sales = True

                daily_resource.append({
                    "date": business_day_value,
                    "resourceKey": rk,
                    "volume": round(volume, 3),
                    "amount": round(amount, 2),
                })
                daily_bucket = daily_rows.setdefault(business_day_value, [0.0, 0.0])
                daily_bucket[0] += volume
                daily_bucket[1] += amount
                resource_bucket = by_resource_rows.setdefault(rk, [0.0, 0.0])
                resource_bucket[0] += volume
                resource_bucket[1] += amount
            if shop_has_sales:
                active_shops.add(shop_key)

        daily = [
            {
                "date": day,
                "volume": round(values[0], 3),
                "amount": round(values[1], 2),
            }
            for day, values in sorted(daily_rows.items())
        ]
        total_liters = sum(values[0] for values in daily_rows.values())
        total_sales = sum(values[1] for values in daily_rows.values())

        result = {
            "source": "OilStationReport",
            "year": year,
            "month": month,
            "queryStart": start,
            "queryEndExclusive": end,
            "summary": {
                "totalTxn": total_txn,
                "activeShops": len(active_shops),
                "totalLiters": round(total_liters, 3),
                "totalSales": round(total_sales, 2),
            },
            "daily": daily,
            "dailyResource": daily_resource,
            "byResource": format_resource_rows([
                (rk, values[0], values[1])
                for rk, values in by_resource_rows.items()
            ]),
        }
        set_cached(cache_key, result)
        return jsonify(result)
    finally:
        conn.close()


@analytics_api.route('/api/assistant/retail-summary')
def api_assistant_retail_summary():
    start = request.args.get("start")
    end = request.args.get("end")
    try:
        datetime.strptime(start or "", "%Y-%m-%d")
        datetime.strptime(end or "", "%Y-%m-%d")
    except ValueError:
        return jsonify({"error": "start and end must use YYYY-MM-DD"}), 400
    if end <= start:
        return jsonify({"error": "end must be greater than start"}), 400

    cache_key = f"assistant:retail-summary:{start}:{end}"
    cached = get_cached(cache_key)
    if cached and cache_allowed():
        return jsonify(cached)

    conn = get_connection()
    try:
        cur = conn.cursor()
        period_condition = business_period_condition_with_pruning("s", "sh")
        total_volume = report_volume_expr("s")
        total_amount = report_amount_expr("s")
        resource_key = report_resource_key_expr("s")
        cur.execute(f"""
            SELECT {resource_key} AS resource_key,
                   COUNT(*) AS txn_count,
                   {total_volume} AS total_liters,
                   {total_amount} AS total_sales
            FROM selling s
            {shift_join("s", "sh")}
            WHERE {period_condition}
              AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
            GROUP BY resource_key
            ORDER BY total_liters DESC
        """, business_period_params(start, end))
        rows = cur.fetchall()
        resources = format_resource_rows([(row[0], row[2], row[3]) for row in rows])
        total_tons = add_resource_tons(resources)
        cur.execute(f"""
            SELECT s.ShopKey,
                   {resource_key} AS resource_key,
                   COUNT(*) AS txn_count,
                   {total_volume} AS total_liters,
                   {total_amount} AS total_sales
            FROM selling s
            {shift_join("s", "sh")}
            WHERE {period_condition}
              AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
              AND s.ShopKey IN ({",".join(str(k) for k in sorted(SHOP_KEY_MAP.keys()))})
            GROUP BY s.ShopKey, resource_key
            ORDER BY total_liters DESC
        """, business_period_params(start, end))
        station_resource_rows = cur.fetchall()
        station_buckets = {}
        for shop_key, rk, txn_count, liters, amount in station_resource_rows:
            shop_key = int(shop_key)
            bucket = station_buckets.setdefault(shop_key, {"txn": 0, "resources": []})
            bucket["txn"] += int(txn_count or 0)
            bucket["resources"].append((rk, liters, amount))
        station_rows = []
        region_buckets = {}
        for shop_key, bucket in station_buckets.items():
            info = SHOP_KEY_MAP.get(shop_key, {"name": f"Station-{shop_key}", "short": "", "region": "未知区域"})
            station_resources = format_resource_rows(bucket["resources"])
            station_tons = add_resource_tons(station_resources)
            station_liters = round(sum(float(item.get("volume") or 0) for item in station_resources), 2)
            station_amount = round(sum(float(item.get("amount") or 0) for item in station_resources), 2)
            station_rows.append({
                "shopKey": shop_key,
                "name": info.get("name") or f"Station-{shop_key}",
                "short": info.get("short") or "",
                "region": info.get("region") or "未知区域",
                "txn": bucket["txn"],
                "volume": station_liters,
                "tons": station_tons,
                "amount": station_amount,
            })
            region = info.get("region") or "未知区域"
            region_bucket = region_buckets.setdefault(region, {"region": region, "txn": 0, "volume": 0, "tons": 0, "amount": 0, "stationCount": 0})
            region_bucket["txn"] += bucket["txn"]
            region_bucket["volume"] += station_liters
            region_bucket["tons"] += station_tons
            region_bucket["amount"] += station_amount
            region_bucket["stationCount"] += 1
        station_rows.sort(key=lambda item: item["amount"], reverse=True)
        region_rows = sorted(
            (
                {
                    **item,
                    "volume": round(item["volume"], 2),
                    "tons": round(item["tons"], 3),
                    "amount": round(item["amount"], 2),
                }
                for item in region_buckets.values()
            ),
            key=lambda item: item["amount"],
            reverse=True,
        )
        result = {
            "source": "OilStationReport",
            "queryStart": start,
            "queryEndExclusive": end,
            "summary": {
                "totalTxn": int(sum(row[1] or 0 for row in rows)),
                "totalLiters": round(sum(float(row[2] or 0) for row in rows), 2),
                "totalTons": total_tons,
                "totalSales": round(sum(float(row[3] or 0) for row in rows), 2),
            },
            "byResource": resources,
            "stations": station_rows,
            "regions": region_rows,
        }
        set_cached(cache_key, result)
        return jsonify(result)
    finally:
        conn.close()


def _monthly_inventory_needs_refresh(payload):
    summary = (payload or {}).get("summary") or {}
    return "inventory" not in summary or "inventoryMonth" not in summary


def _attach_monthly_inventory(year_month, payload, cur=None):
    summary = payload.setdefault("summary", {})
    if not _monthly_inventory_needs_refresh(payload):
        return False

    if cur is not None:
        summary["inventory"] = get_report_month_inventory(cur, year_month)
        summary["inventoryMonth"] = year_month
        return True

    conn = get_connection()
    try:
        local_cur = conn.cursor()
        summary["inventory"] = get_report_month_inventory(local_cur, year_month)
        summary["inventoryMonth"] = year_month
        return True
    finally:
        conn.close()


@analytics_api.route('/api/monthly/<year_month>')
def api_monthly(year_month):
    cache_key = f'monthly:{year_month}'
    cached = get_cached(cache_key)
    if cached and cache_allowed() and not retail_payload_mass_needs_refresh(cached):
        cache_changed = False
        if "totalTons" not in (cached.get("summary") or {}) or "totalSalesCny" not in (cached.get("summary") or {}):
            cached = _attach_monthly_finance_metrics(cached)
            cache_changed = True
        if _attach_monthly_inventory(year_month, cached):
            cache_changed = True
        if cache_changed:
            set_cached(cache_key, cached)
        return jsonify(cached)

    conn = get_connection()
    try:
        cur = conn.cursor()
        y, m = int(year_month.split('-')[0]), int(year_month.split('-')[1])
        ym_start, ym_end = get_query_date_range(year_month)

        # ---- 当月全网汇总 ----
        summary = query_month_summary(cur, year_month)
        summary['inventory'] = get_report_month_inventory(cur, year_month)
        summary['inventoryMonth'] = year_month
        # activeShops 已在 query_month_summary 中合并计算，无需额外查询

        # ---- 同比环比数据 ----
        dates = calc_yoy_mom(y, m)
        prev_summary = query_month_summary(cur, dates['prev'])
        yoy_summary = query_month_summary(cur, dates['yoy'])

        def calc_change(current, previous):
            if not previous or previous <= 0:
                return None
            return round((current - previous) / previous * 100, 1)

        comparison = {
            'mom': {  # 环比
                'sales': calc_change(summary['totalSales'], prev_summary['totalSales']),
                'volume': calc_change(summary['totalLiters'], prev_summary['totalLiters']),
                'cashSales': calc_change(summary['cashSales'], prev_summary['cashSales']),
                'cardSales': calc_change(summary['cardSales'], prev_summary['cardSales']),
                'activeShops': (summary['activeShops'] - prev_summary.get('activeShops', summary['activeShops'])) if prev_summary.get('activeShops') else None,
                'prevMonth': dates['prev'],
            },
            'yoy': {  # 同比
                'sales': calc_change(summary['totalSales'], yoy_summary['totalSales']),
                'volume': calc_change(summary['totalLiters'], yoy_summary['totalLiters']),
                'cashSales': calc_change(summary['cashSales'], yoy_summary['cashSales']),
                'cardSales': calc_change(summary['cardSales'], yoy_summary['cardSales']),
                'activeShops': (summary['activeShops'] - yoy_summary.get('activeShops', summary['activeShops'])) if yoy_summary.get('activeShops') else None,
                'yoyMonth': dates['yoy'],
            },
            'prev': {'summary': prev_summary},
            'yoyRef': {'summary': yoy_summary},
        }

        # ---- 按站点汇总 ----
        period_condition = business_period_condition_with_pruning("s", "sh")
        cur.execute(f"""
            SELECT
                s.ShopKey,
                COUNT(*) as txn_count,
                {report_volume_expr("s")} as total_liters,
                {report_amount_expr("s")} as total_sales,
                SUM(s.ShopBaseCost) as total_base,
                {report_cash_amount_expr("s")} as cash_sales,
                {report_card_amount_expr("s")} as card_sales,
                {report_cash_volume_expr("s")} as cash_liters,
                {report_card_volume_expr("s")} as card_liters,
                {report_bonus_amount_expr("s")} as bonus_cost
            FROM selling s
            {shift_join("s", "sh")}
            WHERE {period_condition}
              AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
            GROUP BY s.ShopKey
            ORDER BY total_sales DESC
        """, business_period_params(ym_start, ym_end))
        station_rows = cur.fetchall()
        stations_data = []
        for r in station_rows:
            shop_key = r[0]
            info = SHOP_KEY_MAP.get(shop_key, {'name': f'Station-{shop_key}', 'short': '', 'region': 'unknown'})
            stations_data.append({
                'shopKey': shop_key,
                'name': info['name'],
                'short': info['short'],
                'region': info['region'],
                'txnCount': int(r[1] or 0),
                'volume': round(float(r[2] or 0), 2),
                'amount': round(float(r[3] or 0), 2),
                'baseCost': round(float(r[4] or 0), 2),
                'cashSales': round(float(r[5] or 0), 2),
                'cardSales': round(float(r[6] or 0), 2),
                'cashLiters': round(float(r[7] or 0), 2),
                'cardLiters': round(float(r[8] or 0), 2),
                'bonus': round(float(r[9] or 0), 2),
                'bonusCost': round(float(r[9] or 0), 2),
            })

        # 站点同比环比 - 批量查询（2次代替 N*2 次）
        # 上月所有站点
        prev_start, prev_end = get_query_date_range(dates['prev'], exclude_today=False)
        yoy_start, yoy_end = get_query_date_range(dates['yoy'], exclude_today=False)
        cur.execute(f"""
            SELECT s.ShopKey, {report_volume_expr("s")} as vol, {report_amount_expr("s")} as amt
            FROM selling s
            {shift_join("s", "sh")}
            WHERE {business_period_condition_with_pruning("s", "sh")}
              AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
            GROUP BY s.ShopKey
        """, business_period_params(prev_start, prev_end))
        prev_shop_data = {r[0]: {'vol': float(r[1] or 0), 'amt': float(r[2] or 0)} for r in cur.fetchall()}
        # 去年同月所有站点
        cur.execute(f"""
            SELECT s.ShopKey, {report_volume_expr("s")} as vol, {report_amount_expr("s")} as amt
            FROM selling s
            {shift_join("s", "sh")}
            WHERE {business_period_condition_with_pruning("s", "sh")}
              AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
            GROUP BY s.ShopKey
        """, business_period_params(yoy_start, yoy_end))
        yoy_shop_data = {r[0]: {'vol': float(r[1] or 0), 'amt': float(r[2] or 0)} for r in cur.fetchall()}

        for st in stations_data:
            sk = st['shopKey']
            prev_d = prev_shop_data.get(sk, {'vol': 0, 'amt': 0})
            yoy_d = yoy_shop_data.get(sk, {'vol': 0, 'amt': 0})
            st['momVolume'] = calc_change(st['volume'], prev_d['vol'])
            st['momSales'] = calc_change(st['amount'], prev_d['amt'])
            st['yoyVolume'] = calc_change(st['volume'], yoy_d['vol'])
            st['yoySales'] = calc_change(st['amount'], yoy_d['amt'])

        # ---- 按区域汇总 ----
        regions = {}
        for s in stations_data:
            rgn = s['region']
            if rgn not in regions:
                regions[rgn] = {'amount': 0, 'volume': 0, 'cashSales': 0, 'cardSales': 0, 'stationCount': 0, 'activeCount': 0}
            regions[rgn]['amount'] += s['amount']
            regions[rgn]['volume'] += s['volume']
            regions[rgn]['cashSales'] += s['cashSales']
            regions[rgn]['cardSales'] += s['cardSales']
            regions[rgn]['stationCount'] += 1
            if s['amount'] > 0:
                regions[rgn]['activeCount'] += 1

        # 区域同比环比 - 直接复用已查询的站点同比环比数据，无需重复查询
        for rgn_name, rgn_data in regions.items():
            rgn_shopkeys = [sk for sk, info in SHOP_KEY_MAP.items() if info['region'] == rgn_name]
            prev_vol = sum(prev_shop_data.get(sk, {}).get('vol', 0) for sk in rgn_shopkeys)
            prev_amt = sum(prev_shop_data.get(sk, {}).get('amt', 0) for sk in rgn_shopkeys)
            yoy_vol = sum(yoy_shop_data.get(sk, {}).get('vol', 0) for sk in rgn_shopkeys)
            yoy_amt = sum(yoy_shop_data.get(sk, {}).get('amt', 0) for sk in rgn_shopkeys)
            regions[rgn_name]['momVolume'] = calc_change(rgn_data['volume'], prev_vol)
            regions[rgn_name]['momSales'] = calc_change(rgn_data['amount'], prev_amt)
            regions[rgn_name]['yoyVolume'] = calc_change(rgn_data['volume'], yoy_vol)
            regions[rgn_name]['yoySales'] = calc_change(rgn_data['amount'], yoy_amt)

        # ---- 油品维度汇总 ----
        resource_key = resource_key_expr("s")
        total_volume = report_volume_expr("s")
        total_amount = report_amount_expr("s")
        cur.execute(f"""
            SELECT {resource_key} as resource_key, {total_volume} as total_vol, {total_amount} as total_amt
            FROM selling s
            {shift_join("s", "sh")}
            WHERE {period_condition}
              AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
            GROUP BY resource_key
        """, business_period_params(ym_start, ym_end))
        by_resource = format_resource_rows(cur.fetchall())
        business_day = business_date_expr("s", "sh")
        cur.execute(f"""
            SELECT
                {business_day} as day,
                {resource_key} as resource_key,
                {report_volume_expr("s")} as total_liters,
                {report_amount_expr("s")} as total_sales
            FROM selling s
            {shift_join("s", "sh")}
            WHERE {period_condition}
              AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
            GROUP BY day, resource_key
            ORDER BY day, resource_key
        """, business_period_params(ym_start, ym_end))
        daily_resource = []
        for r in cur.fetchall():
            rk = int(r[1] or 0)
            if rk < 100:
                continue
            daily_resource.append({
                'date': str(r[0]),
                'resourceKey': canonical_resource_key(rk),
                'volume': round(float(r[2] or 0), 3),
                'amount': round(float(r[3] or 0), 2),
            })

        # ---- 每日趋势 ----
        business_day = business_date_expr("s", "sh")
        cur.execute(f"""
            SELECT
                {business_day} as day,
                {report_volume_expr("s")} as total_liters,
                {report_amount_expr("s")} as total_sales
            FROM selling s
            {shift_join("s", "sh")}
            WHERE {period_condition}
              AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
            GROUP BY day
            ORDER BY day
        """, business_period_params(ym_start, ym_end))
        daily = []
        for r in cur.fetchall():
            daily.append({
                'date': str(r[0]),
                'volume': round(float(r[1] or 0), 2),
                'amount': round(float(r[2] or 0), 2),
            })

        result = {
            'month': year_month,
            'summary': summary,
            'comparison': comparison,
            'stations': stations_data,
            'regions': regions,
            'byResource': by_resource,
            'daily': daily,
        }
        _attach_monthly_finance_metrics(result, daily_resource)
        set_cached(cache_key, result)
        return jsonify(result)
    finally:
        conn.close()


# ===== API: 获取单个站点月度详细（含同比环比） =====
@analytics_api.route('/api/station/<year_month>/<shop_key>')
def api_station_monthly(year_month, shop_key):
    cache_key = f'station:{year_month}:{shop_key}'
    cached = get_cached(cache_key)
    if cached and cache_allowed():
        return jsonify(cached)

    conn = get_connection()
    try:
        cur = conn.cursor()
        shop_key = int(shop_key)
        y, m = int(year_month.split('-')[0]), int(year_month.split('-')[1])
        ym_start, ym_end = get_query_date_range(year_month)

        # 站点汇总
        period_condition = business_period_condition_with_pruning("s", "sh")
        cur.execute(f"""
            SELECT
                COUNT(*) as txn_count,
                {report_volume_expr("s")} as total_liters,
                {report_amount_expr("s")} as total_sales,
                SUM(s.ShopBaseCost) as total_base,
                {report_cash_amount_expr("s")} as cash_sales,
                {report_card_amount_expr("s")} as card_sales,
                {report_cash_volume_expr("s")} as cash_liters,
                {report_card_volume_expr("s")} as card_liters,
                {report_bonus_amount_expr("s")} as bonus_cost
            FROM selling s
            {shift_join("s", "sh")}
            WHERE s.ShopKey = %s
              AND {period_condition}
              AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
        """, (shop_key, *business_period_params(ym_start, ym_end)))
        r = cur.fetchone()
        info = SHOP_KEY_MAP.get(shop_key, {'name': f'Station-{shop_key}', 'short': '', 'region': 'unknown'})

        summary = {
            'txnCount': int(r[0] or 0),
            'volume': round(float(r[1] or 0), 2),
            'amount': round(float(r[2] or 0), 2),
            'baseCost': round(float(r[3] or 0), 2),
            'cashSales': round(float(r[4] or 0), 2),
            'cardSales': round(float(r[5] or 0), 2),
            'cashLiters': round(float(r[6] or 0), 2),
            'cardLiters': round(float(r[7] or 0), 2),
            'bonus': round(float(r[8] or 0), 2),
            'bonusCost': round(float(r[8] or 0), 2),
        }

        # 同比环比
        dates = calc_yoy_mom(y, m)
        prev_data = {}
        yoy_data = {}
        for label, dt in [('prev', dates['prev']), ('yoy', dates['yoy'])]:
            dt_start, dt_end = get_query_date_range(dt, exclude_today=False)
            cur.execute(f"""
                SELECT {report_volume_expr("s")}, {report_amount_expr("s")},
                       {report_cash_amount_expr("s")},
                       {report_card_amount_expr("s")}
                FROM selling s
                {shift_join("s", "sh")}
                WHERE s.ShopKey = %s
                  AND {business_period_condition_with_pruning("s", "sh")}
                  AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
            """, (shop_key, *business_period_params(dt_start, dt_end)))
            row = cur.fetchone()
            d = {
                'volume': round(float(row[0] or 0), 2),
                'amount': round(float(row[1] or 0), 2),
                'cashSales': round(float(row[2] or 0), 2),
                'cardSales': round(float(row[3] or 0), 2),
            }
            if label == 'prev':
                prev_data = d
            else:
                yoy_data = d

        def pct(cur, prev):
            if not prev or prev <= 0:
                return None
            return round((cur - prev) / prev * 100, 1)

        comparison = {
            'mom': {
                'volume': pct(summary['volume'], prev_data.get('volume')),
                'amount': pct(summary['amount'], prev_data.get('amount')),
                'cashSales': pct(summary['cashSales'], prev_data.get('cashSales')),
                'cardSales': pct(summary['cardSales'], prev_data.get('cardSales')),
                'prev': prev_data,
                'prevMonth': dates['prev'],
            },
            'yoy': {
                'volume': pct(summary['volume'], yoy_data.get('volume')),
                'amount': pct(summary['amount'], yoy_data.get('amount')),
                'cashSales': pct(summary['cashSales'], yoy_data.get('cashSales')),
                'cardSales': pct(summary['cardSales'], yoy_data.get('cardSales')),
                'yoy': yoy_data,
                'yoyMonth': dates['yoy'],
            },
        }

        # 每日趋势
        business_day = business_date_expr("s", "sh")
        cur.execute(f"""
            SELECT
                {business_day} as day,
                {report_volume_expr("s")} as total_liters,
                {report_amount_expr("s")} as total_sales
            FROM selling s
            {shift_join("s", "sh")}
            WHERE s.ShopKey = %s
              AND {period_condition}
              AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
            GROUP BY day
            ORDER BY day
        """, (shop_key, *business_period_params(ym_start, ym_end)))
        daily = [{'date': str(r2[0]), 'volume': round(float(r2[1] or 0), 2), 'amount': round(float(r2[2] or 0), 2)} for r2 in cur.fetchall()]

        resource_key = resource_key_expr("s")
        total_volume = report_volume_expr("s")
        total_amount = report_amount_expr("s")
        cur.execute(f"""
            SELECT {resource_key} as resource_key, {total_volume} as total_vol, {total_amount} as total_amt
            FROM selling s
            {shift_join("s", "sh")}
            WHERE s.ShopKey = %s
              AND {period_condition}
              AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
            GROUP BY resource_key
        """, (shop_key, *business_period_params(ym_start, ym_end)))
        by_resource = format_resource_rows(cur.fetchall())

        result = {
            'month': year_month,
            'shopKey': shop_key,
            'name': info['name'],
            'short': info['short'],
            'region': info['region'],
            'status': station_status_for(shop_key, summary['amount']),
            'closedReason': CLOSED_STATIONS.get(shop_key, ''),
            'summary': summary,
            'comparison': comparison,
            'byResource': by_resource,
            'daily': daily,
        }
        set_cached(cache_key, result)
        return jsonify(result)
    finally:
        conn.close()



@analytics_api.route('/api/stations/<year_month>')
def api_stations(year_month):
    """返回当月所有站点数据，用于首页地图和排行榜"""
    cache_key = f'stations:{year_month}'
    cached = get_cached(cache_key)
    if cached and cache_allowed():
        return jsonify(cached)

    conn = get_connection()
    try:
        cur = conn.cursor()
        ym_start, ym_end = get_query_date_range(year_month)

        period_condition = business_period_condition_with_pruning("s", "sh")
        cur.execute(f"""
            SELECT
                s.ShopKey,
                COUNT(*) as txn_count,
                {report_volume_expr("s")} as total_liters,
                {report_amount_expr("s")} as total_sales,
                {report_cash_amount_expr("s")} as cash_sales,
                {report_card_amount_expr("s")} as card_sales,
                {report_cash_volume_expr("s")} as cash_liters,
                {report_card_volume_expr("s")} as card_liters,
                {report_bonus_amount_expr("s")} as bonus_cost
            FROM selling s
            {shift_join("s", "sh")}
            WHERE {period_condition}
              AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
            GROUP BY s.ShopKey
        """, business_period_params(ym_start, ym_end))

        # 先收集有销售记录的站点
        sales_data = {}
        for r in cur.fetchall():
            sales_data[r[0]] = {
                'txnCount': int(r[1] or 0),
                'volume': round(float(r[2] or 0), 2),
                'amount': round(float(r[3] or 0), 2),
                'cash': round(float(r[4] or 0), 2),
                'card': round(float(r[5] or 0), 2),
                'cashLiters': round(float(r[6] or 0), 2),
                'cardLiters': round(float(r[7] or 0), 2),
                'bonus': round(float(r[8] or 0), 2),
            }

        # 遍历所有站点（包括停业的）
        stations_list = []
        for sk, info in SHOP_KEY_MAP.items():
            coords = SHOP_COORDS.get(sk, {'lng': 0, 'lat': 0})
            sd = sales_data.get(sk, {})
            stations_list.append({
                'shopKey': sk,
                'name': info['name'],
                'short': info['short'],
                'region': info['region'],
                'lng': coords.get('lng', 0),
                'lat': coords.get('lat', 0),
                'txnCount': sd.get('txnCount', 0),
                'volume': sd.get('volume', 0),
                'amount': sd.get('amount', 0),
                'cash': sd.get('cash', 0),
                'card': sd.get('card', 0),
                'cashLiters': sd.get('cashLiters', 0),
                'cardLiters': sd.get('cardLiters', 0),
                'bonus': sd.get('bonus', 0),
                'status': station_status_for(sk, sd.get('amount', 0)),
                'closedReason': CLOSED_STATIONS.get(sk, ''),
            })

        result = {
            'month': year_month,
            'stations': stations_list,
            'depots': [
                {'name': k, 'short': v['short'], 'region': v['region'], 'lng': v['lng'], 'lat': v['lat'],
                 'type': 'depot', 'status': v.get('status', 'normal'),
                 'closedReason': v.get('closedReason', '')}
                for k, v in DEPOT_COORDS.items()
            ],
        }
        set_cached(cache_key, result)
        return jsonify(result)
    finally:
        conn.close()


# ===== API: 每日汇总（按站点+按日期） =====
@analytics_api.route('/api/daily/<year_month>')
def api_daily(year_month):
    """返回当月所有站点的每日销售汇总，用于趋势图"""
    cache_key = f'daily:{year_month}'
    cached = get_cached(cache_key)
    if cached and cache_allowed():
        return jsonify(cached)

    conn = get_connection()
    try:
        cur = conn.cursor()
        ym_start, ym_end = get_query_date_range(year_month)

        # 全网每日汇总
        period_condition = business_period_condition_with_pruning("s", "sh")
        business_day = business_date_expr("s", "sh")
        cur.execute(f"""
            SELECT {business_day} as dt,
                   {report_volume_expr("s")} as vol, {report_amount_expr("s")} as amt
            FROM selling s
            {shift_join("s", "sh")}
            WHERE {period_condition}
              AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
            GROUP BY dt ORDER BY dt
        """, business_period_params(ym_start, ym_end))
        daily = [{'date': str(r[0]), 'volume': round(float(r[1] or 0), 2), 'amount': round(float(r[2] or 0), 2)} for r in cur.fetchall()]

        # 每个站点的每日汇总
        cur.execute(f"""
            SELECT s.ShopKey, {business_day} as dt,
                   {report_volume_expr("s")} as vol, {report_amount_expr("s")} as amt
            FROM selling s
            {shift_join("s", "sh")}
            WHERE {period_condition}
              AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
            GROUP BY s.ShopKey, dt ORDER BY s.ShopKey, dt
        """, business_period_params(ym_start, ym_end))

        stations_daily = {}
        for r in cur.fetchall():
            sk = r[0]
            if sk not in stations_daily:
                stations_daily[sk] = []
            stations_daily[sk].append({'date': str(r[1]), 'volume': round(float(r[2] or 0), 2), 'amount': round(float(r[3] or 0), 2)})

        result = {'month': year_month, 'daily': daily, 'stations': stations_daily}
        set_cached(cache_key, result)
        return jsonify(result)
    finally:
        conn.close()


# ===== API: 区域汇总 =====
@analytics_api.route('/api/regions/<year_month>')
def api_regions(year_month):
    """返回当月各区域汇总数据（含同比环比）"""
    cache_key = f'regions:{year_month}'
    cached = get_cached(cache_key)
    if cached and cache_allowed():
        if _attach_region_snapshot_comparisons(year_month, cached):
            set_cached(cache_key, cached)
        return jsonify(cached)

    conn = get_connection()
    try:
        cur = conn.cursor()
        y, m = int(year_month.split('-')[0]), int(year_month.split('-')[1])
        ym_start, ym_end = get_query_date_range(year_month)
        dates = calc_yoy_mom(y, m)

        # 当月区域汇总（含现金/刷卡分拆）
        period_condition = business_period_condition_with_pruning("s", "sh")
        cur.execute(f"""
            SELECT s.ShopKey,
                   {report_volume_expr("s")} as vol,
                   {report_amount_expr("s")} as amt,
                   {report_cash_amount_expr("s")} as cash_sales,
                   {report_card_amount_expr("s")} as card_sales,
                   {report_bonus_amount_expr("s")} as bonus_cost
            FROM selling s
            {shift_join("s", "sh")}
            WHERE {period_condition}
              AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
            GROUP BY s.ShopKey
        """, business_period_params(ym_start, ym_end))
        shop_data = cur.fetchall()

        regions = {}
        for r in shop_data:
            sk = r[0]
            info = SHOP_KEY_MAP.get(sk, {'name': f'Station-{sk}', 'region': 'unknown'})
            rgn = info['region']
            if rgn not in regions:
                regions[rgn] = {'amount': 0, 'volume': 0, 'cash': 0, 'card': 0, 'bonus': 0, 'stationCount': 0, 'activeCount': 0}
            regions[rgn]['amount'] += float(r[2] or 0)
            regions[rgn]['volume'] += float(r[1] or 0)
            regions[rgn]['cash'] += float(r[3] or 0)
            regions[rgn]['card'] += float(r[4] or 0)
            regions[rgn]['bonus'] += float(r[5] or 0)
            regions[rgn]['activeCount'] += 1

        # 计算每个区域的总站数（所有站点，不管有没有销售）
        for sk, info in SHOP_KEY_MAP.items():
            rgn = info['region']
            if rgn in regions:
                regions[rgn]['stationCount'] += 1
            else:
                regions[rgn] = {'amount': 0, 'volume': 0, 'cash': 0, 'card': 0, 'bonus': 0, 'stationCount': 1, 'activeCount': 0}

        # 环比同比
        def calc_change(cur_val, prev_val):
            if not prev_val or prev_val == 0:
                return 0
            return round((cur_val - prev_val) / prev_val * 100, 1)

        for label, dt in [('mom', dates['prev']), ('yoy', dates['yoy'])]:
            dt_start, dt_end = get_query_date_range(dt, exclude_today=False)
            cur.execute(f"""
                SELECT s.ShopKey, {report_volume_expr("s")} as vol, {report_amount_expr("s")} as amt
                FROM selling s
                {shift_join("s", "sh")}
                WHERE {business_period_condition_with_pruning("s", "sh")}
                  AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
                GROUP BY s.ShopKey
            """, business_period_params(dt_start, dt_end))
            prev_shop = cur.fetchall()

            prev_regions = {}
            for r in prev_shop:
                sk = r[0]
                info = SHOP_KEY_MAP.get(sk, {'region': 'unknown'})
                rgn = info['region']
                if rgn not in prev_regions:
                    prev_regions[rgn] = {'amount': 0, 'volume': 0}
                prev_regions[rgn]['amount'] += float(r[2] or 0)
                prev_regions[rgn]['volume'] += float(r[1] or 0)

            for rgn_name, rgn_data in regions.items():
                prev_rgn = prev_regions.get(rgn_name, {'amount': 0, 'volume': 0})
                amt_chg = calc_change(rgn_data['amount'], prev_rgn['amount'])
                vol_chg = calc_change(rgn_data['volume'], prev_rgn['volume'])
                if label == 'mom':
                    regions[rgn_name]['momSales'] = amt_chg
                    regions[rgn_name]['momVolume'] = vol_chg
                else:
                    regions[rgn_name]['yoySales'] = amt_chg
                    regions[rgn_name]['yoyVolume'] = vol_chg

        resource_key = resource_key_expr("s")
        total_volume = report_volume_expr("s")
        total_amount = report_amount_expr("s")
        cur.execute(f"""
            SELECT s.ShopKey, {resource_key} as resource_key, {total_volume} as vol, {total_amount} as amt
            FROM selling s
            {shift_join("s", "sh")}
            WHERE {period_condition}
              AND (s.DerivedPaymentTypeKey IS NULL OR s.DerivedPaymentTypeKey NOT IN (2, 4))
            GROUP BY s.ShopKey, resource_key
        """, business_period_params(ym_start, ym_end))
        region_resource_rows = {}
        for shop_key, rk, vol, amt in cur.fetchall():
            rk = canonical_resource_key(rk)
            if rk < 100 or rk not in KNOWN_RESOURCE_KEYS:
                continue
            info = SHOP_KEY_MAP.get(shop_key, {'region': 'unknown'})
            bucket = region_resource_rows.setdefault(info['region'], {})
            item = bucket.setdefault(rk, [0.0, 0.0])
            item[0] += float(vol or 0)
            item[1] += float(amt or 0)
        by_resource_by_region = {
            region: format_resource_rows([(rk, vals[0], vals[1]) for rk, vals in rows.items()])
            for region, rows in region_resource_rows.items()
        }
        by_resource = format_resource_rows([
            (rk, sum(rows.get(rk, [0, 0])[0] for rows in region_resource_rows.values()),
             sum(rows.get(rk, [0, 0])[1] for rows in region_resource_rows.values()))
            for rk in {rk for rows in region_resource_rows.values() for rk in rows}
        ])

        result = {
            'month': year_month,
            'regions': regions,
            'byResource': by_resource,
            'byResourceByRegion': by_resource_by_region,
        }
        set_cached(cache_key, result)
        return jsonify(result)
    finally:
        conn.close()


